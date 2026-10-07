"""Fórmulas VISIBLES en los libros de Excel, con el número del API en caché.

Pedido del cliente (5-oct-2026, sobre la hoja «reporte horas»): «¿me puedes
apoyar poniendo con fórmulas las celdas que sean por suma? Por ejemplo los
permisos de AFAC, queremos que tenga las fórmulas visibles, no que sea solo un
texto con la cantidad».

Tres piezas, puras (sin red ni disco):

1. REGISTRO (`escribir`): la celda recibe la fórmula y se anota el número que
   manda el API para ella (el que la celda mostraba antes de este cambio).
2. VERIFICACIÓN (`finalizar`): cada fórmula se evalúa como la evalúa Excel
   (mismas celdas del libro, `ROUND` medio-hacia-afuera, doble precisión) y
   debe dar EXACTAMENTE ese número a la vista. Si no (un redondeo intermedio
   del API que no está en el libro, un insumo vacío, una división entre
   cero), la celda vuelve a ser VALOR. Regla dura: los números del libro no
   cambian ni un centavo — al abrir, Excel recalcula (`fullCalcOnLoad`) y
   debe ver lo mismo que el panel.
3. CACHÉ (`inyectar_valores_cache`): openpyxl guarda `<f>…</f><v></v>` sin
   valor, y la vista previa del teléfono (Quick Look/iOS, WhatsApp, Gmail) NO
   calcula: mostraría la celda VACÍA. Tras guardar, se escribe el número del
   API en el `<v>` de cada fórmula de cada hoja (la hoja se resuelve por
   nombre vía `xl/workbook.xml` + `xl/_rels/workbook.xml.rels`, nunca por
   orden) y el resto del zip queda byte a byte igual.

`guardar(wb)` hace 2 + guardar + 3 y devuelve los bytes del .xlsx.
"""

from __future__ import annotations

import logging
import math
import posixpath
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# Atributo con el que el registro viaja PEGADO al libro (no hay estado de
# módulo: dos libros generados a la vez en hilos distintos no se mezclan).
_ATRIBUTO_REGISTRO = "_vt_formulas"


# ---------------------------------------------------------------------------
# Registro
# ---------------------------------------------------------------------------


@dataclass
class FormulaRegistrada:
    """Una fórmula escrita en una celda y el número del API que debe dar."""

    texto: str  # sin el «=» inicial
    valor: float | int
    decimales: int  # decimales que se VEN con el formato de la celda
    degradada: bool = False


@dataclass
class RegistroFormulas:
    """Fórmulas por hoja (objeto Worksheet → coordenada → fórmula)."""

    por_hoja: dict[Worksheet, dict[str, FormulaRegistrada]] = field(default_factory=dict)

    def de(self, ws: Worksheet, coord: str) -> FormulaRegistrada | None:
        return self.por_hoja.get(ws, {}).get(coord)

    def degradadas(self) -> list[tuple[str, str]]:
        """(hoja, celda) de las fórmulas que volvieron a ser valor."""
        return [
            (ws.title, coord)
            for ws, celdas in self.por_hoja.items()
            for coord, f in celdas.items()
            if f.degradada
        ]

    def cache_por_hoja(self) -> dict[str, dict[str, float | int]]:
        """Título de hoja → {celda: número del API} de las fórmulas vivas."""
        cache: dict[str, dict[str, float | int]] = {}
        for ws, celdas in self.por_hoja.items():
            vivas = {c: f.valor for c, f in celdas.items() if not f.degradada}
            if vivas:
                cache[ws.title] = vivas
        return cache


def registro(wb: Workbook) -> RegistroFormulas:
    """Registro de fórmulas del libro (se crea con la primera)."""
    reg = getattr(wb, _ATRIBUTO_REGISTRO, None)
    if reg is None:
        reg = RegistroFormulas()
        setattr(wb, _ATRIBUTO_REGISTRO, reg)
    return reg


def escribir(cell: Cell, formula: str, valor: float | int, decimales: int = 2) -> Cell:
    """Escribe `=formula` en la celda y anota el número del API que le toca.
    `decimales` = los que muestra el formato de la celda (2 en dinero/horas,
    6 en el T.C.): la verificación exige que la fórmula se VEA igual."""
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        raise TypeError(f"valor del API no numérico para {cell.coordinate}: {valor!r}")
    texto = formula[1:] if formula.startswith("=") else formula
    cell.value = f"={texto}"
    ws = cell.parent
    registro(ws.parent).por_hoja.setdefault(ws, {})[cell.coordinate] = FormulaRegistrada(
        texto=texto, valor=valor, decimales=decimales
    )
    return cell


# ---------------------------------------------------------------------------
# Redondeo y comparación «como Excel»
# ---------------------------------------------------------------------------


def redondear_excel(x: float, decimales: int) -> float:
    """ROUND de Excel: medio hacia AFUERA de cero sobre el número a 15 cifras
    significativas (Excel da ROUND(1.005, 2) = 1.01 aunque el binario de
    1.005 sea 1.00499…; el `Math.round` del API daría 1.00 — por eso existe
    la verificación)."""
    d = Decimal(format(x, ".15g"))
    return float(d.quantize(Decimal(1).scaleb(-decimales), rounding=ROUND_HALF_UP))


def coincide(calculado: float, api: float, decimales: int) -> bool:
    """¿La fórmula se VE igual que el número del API? Misma cifra con los
    decimales del formato y a menos de medio último dígito."""
    if not math.isfinite(calculado):
        return False
    if abs(calculado - api) >= 0.5 * 10.0**-decimales:
        return False
    return redondear_excel(calculado, decimales) == redondear_excel(float(api), decimales)


# ---------------------------------------------------------------------------
# Evaluador mínimo (solo la gramática que generan los libros)
# ---------------------------------------------------------------------------


class FormulaInvalidaError(Exception):
    """#VALOR!, #DIV/0!, referencia inválida… (la celda se degrada a valor)."""


class _Vacio:
    """Celda vacía: 0 en aritmética, ignorada en SUM/AVERAGE."""

    def __repr__(self) -> str:  # pragma: no cover - ayuda de depuración
        return "VACIO"


VACIO = _Vacio()


class _Arreglo(list):
    """ROUND sobre un RANGO (6-oct-2026): un número por celda, ya redondeado.
    Solo vale como argumento de SUMPRODUCT, que en Excel evalúa sus
    argumentos como arreglo (sin Ctrl+Mayús+Intro, desde siempre). En
    cualquier otro lugar Excel haría intersección implícita: aquí es
    #¡VALOR! y la celda se queda como valor."""


_TOKEN = re.compile(
    r"""\s*(?:
        (?P<hoja>'(?:[^']|'')+'!)
       |(?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)
       |(?P<txt>"(?:[^"]|"")*")
       |(?P<ref>\$?[A-Z]{1,3}\$?\d+)(?![A-Za-z0-9_(])
       |(?P<fn>[A-Z][A-Z0-9.]*)(?=\()
       |(?P<op>[-+*/(),:])
    )""",
    re.VERBOSE,
)


def _tokens(texto: str) -> list[tuple[str, str]]:
    pos, salida = 0, []
    texto = texto.rstrip()
    while pos < len(texto):
        m = _TOKEN.match(texto, pos)
        if m is None or m.end() == pos:
            raise FormulaInvalidaError(f"no se entiende la fórmula en {pos}: {texto!r}")
        tipo = m.lastgroup or ""
        salida.append((tipo, m.group(tipo)))
        pos = m.end()
    return salida


@dataclass(frozen=True)
class _Rango:
    ws: Worksheet
    fila_ini: int
    col_ini: int
    fila_fin: int
    col_fin: int


def _fila_col(ref: str) -> tuple[int, int]:
    letras, fila = coordinate_from_string(ref.replace("$", ""))
    return fila, column_index_from_string(letras)


_CRITERIO = re.compile(r"^(<=|>=|<>|<|>|=)?(.*)$", re.DOTALL)


def cumple_criterio(valor: object, criterio: str) -> bool:
    """Criterio de SUMIF/AVERAGEIF: «>0», «<>CANCELADO», «=X» o «X»."""
    m = _CRITERIO.match(criterio)
    op = (m.group(1) if m else None) or "="
    objetivo = m.group(2) if m else criterio
    try:
        numero = float(objetivo)
    except ValueError:
        numero = None
    if numero is not None:
        if not isinstance(valor, float):
            return op == "<>"
        return {
            "=": valor == numero,
            "<>": valor != numero,
            "<": valor < numero,
            ">": valor > numero,
            "<=": valor <= numero,
            ">=": valor >= numero,
        }[op]
    texto = "" if valor is VACIO else str(valor)
    if op == "=":
        return texto.casefold() == objetivo.casefold()
    if op == "<>":
        return texto.casefold() != objetivo.casefold()
    return False


class Evaluador:
    """Evalúa las fórmulas registradas del libro con la aritmética de Excel.
    Una referencia a otra celda con fórmula se evalúa recursivamente (Excel
    usa su valor COMPLETO, no el redondeado a la vista)."""

    def __init__(self, wb: Workbook, reg: RegistroFormulas) -> None:
        self.wb = wb
        self.reg = reg
        self._memo: dict[tuple[int, str], object] = {}
        self._en_curso: set[tuple[int, str]] = set()
        # Errores que NO son de la fórmula (otra versión de openpyxl sin
        # `_cells`, recursión…): la celda se degrada igual y `finalizar`
        # avisa UNA vez con el primero.
        self.inesperados = 0
        self.primer_inesperado: BaseException | None = None

    # -- celdas --------------------------------------------------------------

    def valor_celda(self, ws: Worksheet, fila: int, col: int) -> object:
        """Valor que Excel ve en la celda: número, texto o VACIO. Lee el
        diccionario interno para NO crear celdas vacías en la hoja (acceder
        con ws.cell() las crearía y movería la dimensión del libro)."""
        cell = ws._cells.get((fila, col))
        coord = cell.coordinate if cell is not None else None
        if coord is not None and self.reg.de(ws, coord) is not None:
            return self.verificar(ws, coord)
        valor = cell.value if cell is not None else None
        if valor is None:
            return VACIO
        if isinstance(valor, bool):
            return 1.0 if valor else 0.0
        if isinstance(valor, int):
            return float(valor)
        if isinstance(valor, float):
            # Excel lee el número que openpyxl ESCRIBE («%.16g»), no el
            # float de Python: se evalúa con ese mismo número.
            return float(format(valor, ".16g"))
        return str(valor)

    def verificar(self, ws: Worksheet, coord: str) -> object:
        """Valor de una celda con fórmula registrada; si no reproduce el número
        del API, la celda se DEGRADA a ese número (y ese es su valor)."""
        clave = (id(ws), coord)
        if clave in self._memo:
            return self._memo[clave]
        f = self.reg.de(ws, coord)
        if f is None:
            fila, col = _fila_col(coord)
            return self.valor_celda(ws, fila, col)
        if clave in self._en_curso:
            raise FormulaInvalidaError(f"referencia circular en {ws.title}!{coord}")
        self._en_curso.add(clave)
        try:
            calculado = self.evaluar(f.texto, ws)
            ok = isinstance(calculado, float) and coincide(calculado, f.valor, f.decimales)
        except (FormulaInvalidaError, ArithmeticError, ValueError, TypeError):
            # #DIV/0!, #¡VALOR!, desbordes del ROUND (Decimal)… ⇒ valor.
            calculado, ok = None, False
        except Exception as e:  # jamás un 500 por una fórmula
            # Fallo del evaluador, no de la fórmula (AttributeError de otra
            # versión de openpyxl, RecursionError…): solo ESTA celda vuelve a
            # valor y el balance sale.
            self.inesperados += 1
            if self.primer_inesperado is None:
                self.primer_inesperado = e
            calculado, ok = None, False
        finally:
            self._en_curso.discard(clave)
        if ok:
            resultado: object = calculado
        else:
            f.degradada = True
            ws[coord].value = f.valor
            # Quien la cite verá el número tal como openpyxl lo escribe.
            resultado = float(format(float(f.valor), ".16g"))
            logger.debug(
                "fórmula degradada a valor %s!%s: =%s → %r (API %r)",
                ws.title,
                coord,
                f.texto,
                calculado,
                f.valor,
            )
        self._memo[clave] = resultado
        return resultado

    # -- expresiones ---------------------------------------------------------

    def evaluar(self, texto: str, ws: Worksheet) -> object:
        p = _Parser(_tokens(texto), self, ws)
        valor = p.expr()
        if p.i != len(p.toks):
            raise FormulaInvalidaError(f"sobra texto en la fórmula: {texto!r}")
        if isinstance(valor, (_Rango, _Arreglo)):
            raise FormulaInvalidaError("un rango suelto no es un número")
        return valor

    def valores_rango(self, r: _Rango) -> list[object]:
        return [
            self.valor_celda(r.ws, fila, col)
            for fila in range(r.fila_ini, r.fila_fin + 1)
            for col in range(r.col_ini, r.col_fin + 1)
        ]


def _suma(nums: list[float]) -> float:
    """Suma SECUENCIAL en doble precisión, como el `reduce` del API y Excel
    (el `sum()` de Python 3.12 compensa errores y podría diferir en el último
    bit justo en un empate de redondeo)."""
    total = 0.0
    for x in nums:
        total += x
    return total


def _numero(x: object) -> float:
    """Operando aritmético: vacío = 0; texto o rango = #¡VALOR!."""
    if x is VACIO:
        return 0.0
    if isinstance(x, float):
        return x
    raise FormulaInvalidaError(f"operando no numérico: {x!r}")


class _Parser:
    """Descenso recursivo que evalúa mientras lee (+ − × ÷, paréntesis,
    menos unario, referencias A1/$A$1, 'hoja'!A1, rangos A1:B9 y las
    funciones SUM, AVERAGE, AVERAGEIF, SUMIF, ROUND y SUMPRODUCT — ROUND
    sobre un rango solo dentro de SUMPRODUCT, `_Arreglo`)."""

    def __init__(self, toks: list[tuple[str, str]], ev: Evaluador, ws: Worksheet) -> None:
        self.toks = toks
        self.i = 0
        self.ev = ev
        self.ws = ws

    def _ver(self) -> tuple[str, str] | None:
        return self.toks[self.i] if self.i < len(self.toks) else None

    def _tomar(self, op: str | None = None) -> tuple[str, str]:
        tok = self._ver()
        if tok is None or (op is not None and tok != ("op", op)):
            raise FormulaInvalidaError(f"se esperaba {op!r} y llegó {tok!r}")
        self.i += 1
        return tok

    def expr(self) -> object:
        valor = self.termino()
        while (tok := self._ver()) in (("op", "+"), ("op", "-")):
            self.i += 1
            derecho = _numero(self.termino())
            izq = _numero(valor)
            valor = izq + derecho if tok == ("op", "+") else izq - derecho
        return valor

    def termino(self) -> object:
        valor = self.unario()
        while (tok := self._ver()) in (("op", "*"), ("op", "/")):
            self.i += 1
            derecho = _numero(self.unario())
            izq = _numero(valor)
            if tok == ("op", "*"):
                valor = izq * derecho
            else:
                if derecho == 0:
                    raise FormulaInvalidaError("#DIV/0!")
                valor = izq / derecho
        return valor

    def unario(self) -> object:
        tok = self._ver()
        if tok == ("op", "-"):
            self.i += 1
            return -_numero(self.unario())
        if tok == ("op", "+"):
            self.i += 1
            return _numero(self.unario())
        return self.primario()

    def primario(self) -> object:
        tok = self._ver()
        if tok is None:
            raise FormulaInvalidaError("fórmula incompleta")
        tipo, txt = tok
        if tipo == "num":
            self.i += 1
            return float(txt)
        if tipo == "txt":
            self.i += 1
            return txt[1:-1].replace('""', '"')
        if tok == ("op", "("):
            self.i += 1
            valor = self.expr()
            self._tomar(")")
            return valor
        if tipo == "fn":
            self.i += 1
            return self._funcion(txt)
        if tipo in ("hoja", "ref"):
            return self._referencia()
        raise FormulaInvalidaError(f"token inesperado {tok!r}")

    def _referencia(self) -> object:
        ws = self.ws
        tipo, txt = self._tomar()
        if tipo == "hoja":
            nombre = txt[1:-2].replace("''", "'")
            if nombre not in self.ev.wb.sheetnames:
                raise FormulaInvalidaError(f"#¡REF! hoja inexistente {nombre!r}")
            ws = self.ev.wb[nombre]
            tipo, txt = self._tomar()
        if tipo != "ref":
            raise FormulaInvalidaError(f"se esperaba una celda y llegó {txt!r}")
        fila, col = _fila_col(txt)
        if self._ver() == ("op", ":"):
            self.i += 1
            tipo2, txt2 = self._tomar()
            if tipo2 != "ref":
                raise FormulaInvalidaError("rango incompleto")
            fila2, col2 = _fila_col(txt2)
            return _Rango(ws, min(fila, fila2), min(col, col2), max(fila, fila2), max(col, col2))
        return self.ev.valor_celda(ws, fila, col)

    def _argumentos(self) -> list[object]:
        self._tomar("(")
        args: list[object] = []
        if self._ver() == ("op", ")"):
            self.i += 1
            return args
        while True:
            args.append(self.expr())
            tok = self._tomar()
            if tok == ("op", ")"):
                return args
            if tok != ("op", ","):
                raise FormulaInvalidaError(f"se esperaba «,» o «)» y llegó {tok!r}")

    def _aplanar(self, args: list[object]) -> list[object]:
        valores: list[object] = []
        for a in args:
            if isinstance(a, _Arreglo):
                raise FormulaInvalidaError("ROUND de un rango solo vale dentro de SUMPRODUCT")
            if isinstance(a, _Rango):
                valores.extend(self.ev.valores_rango(a))
            else:
                valores.append(a)
        return valores

    def _sumproduct(self, args: list[object]) -> float:
        """SUMPRODUCT de Excel: arreglos (ROUND de un rango) o rangos del
        mismo tamaño; una entrada no numérica cuenta como 0. Con uno solo, la
        suma de sus números (`ROUND(SUMPRODUCT(ROUND(X3:X9,2)),2)` = la Σ de
        las filas tal como se ven)."""
        if not args:
            raise FormulaInvalidaError("SUMPRODUCT sin argumentos")
        columnas: list[list[object]] = []
        for a in args:
            if isinstance(a, _Arreglo):
                columnas.append(list(a))
            elif isinstance(a, _Rango):
                columnas.append(self.ev.valores_rango(a))
            else:
                raise FormulaInvalidaError("SUMPRODUCT: argumento que no es rango ni arreglo")
        if any(len(c) != len(columnas[0]) for c in columnas):
            raise FormulaInvalidaError("SUMPRODUCT: arreglos de distinto tamaño")
        productos: list[float] = []
        for fila in zip(*columnas, strict=True):
            numeros = [x for x in fila if isinstance(x, float)]
            if len(numeros) == len(fila):
                producto = 1.0
                for x in numeros:
                    producto *= x
                productos.append(producto)
        return _suma(productos)

    def _funcion(self, nombre: str) -> object:
        args = self._argumentos()
        if nombre == "SUM":
            return _suma([v for v in self._aplanar(args) if isinstance(v, float)])
        if nombre == "AVERAGE":
            nums = [v for v in self._aplanar(args) if isinstance(v, float)]
            if not nums:
                raise FormulaInvalidaError("#DIV/0! (AVERAGE sin números)")
            return _suma(nums) / len(nums)
        if nombre in ("AVERAGEIF", "SUMIF"):
            if len(args) not in (2, 3) or not isinstance(args[0], _Rango):
                raise FormulaInvalidaError(f"{nombre} mal formada")
            criterio = args[1]
            if isinstance(criterio, float):
                criterio = repr(criterio)
            if not isinstance(criterio, str):
                raise FormulaInvalidaError(f"criterio de {nombre} inválido")
            rango = args[0]
            objetivo = args[2] if len(args) == 3 else rango
            if not isinstance(objetivo, _Rango):
                raise FormulaInvalidaError(f"{nombre}: el rango a promediar/sumar no es rango")
            prueba = self.ev.valores_rango(rango)
            datos = self.ev.valores_rango(objetivo)
            if len(prueba) != len(datos):
                raise FormulaInvalidaError(f"{nombre}: rangos de distinto tamaño")
            nums = [
                d
                for p, d in zip(prueba, datos, strict=True)
                if cumple_criterio(p, criterio) and isinstance(d, float)
            ]
            if nombre == "SUMIF":
                return _suma(nums)
            if not nums:
                raise FormulaInvalidaError("#DIV/0! (AVERAGEIF sin números)")
            return _suma(nums) / len(nums)
        if nombre == "ROUND":
            if len(args) != 2:
                raise FormulaInvalidaError("ROUND lleva dos argumentos")
            decimales = int(_numero(args[1]))
            if isinstance(args[0], _Rango):
                # Celda vacía = 0; texto = #¡VALOR! (como Excel).
                return _Arreglo(
                    redondear_excel(_numero(v), decimales) for v in self.ev.valores_rango(args[0])
                )
            return redondear_excel(_numero(args[0]), decimales)
        if nombre == "SUMPRODUCT":
            return self._sumproduct(args)
        raise FormulaInvalidaError(f"función no soportada: {nombre}")


# ---------------------------------------------------------------------------
# Cierre: verificar, guardar e inyectar la caché
# ---------------------------------------------------------------------------


def finalizar(wb: Workbook) -> dict[str, dict[str, float | int]]:
    """Verifica TODAS las fórmulas del libro (degrada a valor las que no
    reproducen el número del API), pide a Excel recalcular al abrir y
    devuelve la caché por hoja para `inyectar_valores_cache`."""
    reg = getattr(wb, _ATRIBUTO_REGISTRO, None)
    if reg is None:
        return {}
    ev = Evaluador(wb, reg)
    for ws, celdas in list(reg.por_hoja.items()):
        for coord in list(celdas):
            ev.verificar(ws, coord)
    if ev.inesperados:
        logger.warning(
            "%d fórmula(s) quedaron como valor por un error del evaluador (primero: %r)",
            ev.inesperados,
            ev.primer_inesperado,
        )
    degradadas = reg.degradadas()
    if degradadas:
        logger.info(
            "%d fórmula(s) quedaron como valor (no reproducen al API al centavo)",
            len(degradadas),
        )
    wb.calculation.fullCalcOnLoad = True
    return reg.cache_por_hoja()


def _a_valores(wb: Workbook) -> None:
    """Respaldo: TODAS las fórmulas vivas vuelven a ser su número del API
    (el libro queda exactamente como antes de las fórmulas)."""
    reg = registro(wb)
    for ws, celdas in reg.por_hoja.items():
        for coord, f in celdas.items():
            if not f.degradada:
                f.degradada = True
                ws[coord].value = f.valor


def guardar(wb: Workbook) -> bytes:
    """Verifica las fórmulas, guarda el libro e inyecta su caché. Si la
    verificación falla (un error que no es de una fórmula: revisión
    5-oct-2026) o la caché no se puede escribir (p. ej. otra versión de
    openpyxl serializa distinto las celdas), el libro sale con VALORES —el de
    siempre— en vez de fórmulas sin verificar o que la vista previa del
    teléfono mostraría vacías; jamás un error 500."""
    try:
        cache = finalizar(wb)
    except Exception:
        logger.exception("no se pudieron verificar las fórmulas: el libro sale con valores")
        _a_valores(wb)
        cache = {}
    buf = BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if not cache:
        return data
    try:
        return inyectar_valores_cache(data, cache)
    except Exception:
        logger.exception("no se pudo escribir la caché de las fórmulas: el libro sale con valores")
        _a_valores(wb)
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()


_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL_DOC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_REL_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def rutas_de_hojas(z: zipfile.ZipFile) -> dict[str, str]:
    """Nombre de hoja → miembro del zip (`xl/worksheets/sheetN.xml`), leído de
    `xl/workbook.xml` + `xl/_rels/workbook.xml.rels` (jamás por orden)."""
    libro = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    destinos = {
        r.get("Id"): r.get("Target") or "" for r in rels.iter(f"{{{_NS_REL_PKG}}}Relationship")
    }
    rutas: dict[str, str] = {}
    for hoja in libro.iter(f"{{{_NS_MAIN}}}sheet"):
        nombre = hoja.get("name")
        destino = destinos.get(hoja.get(f"{{{_NS_REL_DOC}}}id"))
        if nombre is None or not destino:
            continue
        if destino.startswith("/"):
            ruta = destino.lstrip("/")
        else:
            ruta = posixpath.normpath(posixpath.join("xl", destino))
        rutas[nombre] = ruta
    return rutas


def numero_xml(x: float | int) -> str:
    """Número para `<v>`, escrito IGUAL que openpyxl escribía el valor antes
    de las fórmulas (`safe_string`: «%.16g», 16 cifras significativas) pero
    SIN notación científica. Un entero va sin «.0» (7000.0 → «7000»). Así,
    leído con `data_only=True`, el libro da los MISMOS números —bit a bit— y
    tipos que el generador anterior: con `repr` (17 cifras) un % COBRADO
    como 18500.50 / 52158.40 salía distinto en el último bit (revisión
    5-oct-2026)."""
    if isinstance(x, bool):
        raise TypeError("booleano como caché de fórmula")
    if isinstance(x, int):
        return str(x)
    f = float(x)
    if not math.isfinite(f):
        raise ValueError(f"caché no finita: {x!r}")
    if f.is_integer() and abs(f) < 1e16:
        return str(int(f))
    texto = format(f, ".16g")  # == "%.16g" % f de openpyxl
    if "e" in texto or "E" in texto:
        texto = format(Decimal(texto), "f")
    return texto


# Celda con fórmula tal como la escribe openpyxl: <c r="K3" s="5"><f>…</f><v></v></c>
# (con lxml puede salir <v/>; sin <v> también se acepta; el atributo r puede
# no ir primero).
_CELDA_FORMULA = re.compile(
    r"<c(?P<attrs>\s[^>]*)>"
    r"(?P<f><f(?:\s[^>]*)?>[^<]*</f>)"
    r"(?:<v\s*/>|<v>[^<]*</v>)?"
    r"</c>"
)
_ATRIBUTO_R = re.compile(r'\sr="(?P<ref>[A-Z]+[0-9]+)"')
_ATRIBUTO_T = re.compile(r'\s+t="[^"]*"')


def _inyectar_en_hoja(xml: str, celdas: dict[str, float | int], hoja: str) -> str:
    pendientes = set(celdas)

    def reemplazo(m: re.Match[str]) -> str:
        r = _ATRIBUTO_R.search(m.group("attrs"))
        ref = r.group("ref") if r else None
        if ref is None or ref not in celdas:
            return m.group(0)
        pendientes.discard(ref)
        # Resultado NUMÉRICO: sin t="str"/"e" (el tipo por omisión es n).
        attrs = _ATRIBUTO_T.sub("", m.group("attrs"))
        return f"<c{attrs}>{m.group('f')}<v>{numero_xml(celdas[ref])}</v></c>"

    nuevo = _CELDA_FORMULA.sub(reemplazo, xml)
    if pendientes:
        raise ValueError(
            f"hoja {hoja!r}: no se encontró la fórmula de {sorted(pendientes)[:5]} "
            "para escribir su valor en caché"
        )
    return nuevo


def inyectar_valores_cache(xlsx: bytes, cache_por_hoja: dict[str, dict[str, float | int]]) -> bytes:
    """Escribe en `<v>` el número del API de cada fórmula (por hoja y celda) y
    devuelve el .xlsx nuevo. Los demás miembros del zip se copian byte a
    byte, con su mismo orden, fecha y compresión."""
    entrada = zipfile.ZipFile(BytesIO(xlsx))
    rutas = rutas_de_hojas(entrada)
    nuevos: dict[str, bytes] = {}
    for hoja, celdas in cache_por_hoja.items():
        if not celdas:
            continue
        ruta = rutas.get(hoja)
        if ruta is None:
            raise ValueError(f"el libro no tiene la hoja {hoja!r}")
        xml = entrada.read(ruta).decode("utf-8")
        nuevos[ruta] = _inyectar_en_hoja(xml, celdas, hoja).encode("utf-8")
    salida = BytesIO()
    with zipfile.ZipFile(salida, "w") as z:
        for info in entrada.infolist():
            z.writestr(info, nuevos.get(info.filename, entrada.read(info.filename)))
    return salida.getvalue()
