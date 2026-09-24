"""Excel de una REPOSICIÓN de caja chica (24-sep-2026).

Pedido del cliente: «en el apartado de caja chica, quiero ver si se puede al
momento de reembolsar la caja de cada uno, me puede arrojar un Excel
descargable con la información de lo que estoy reembolsando».

Una hoja con: título y subtítulo; ENCABEZADO (responsable, caja, moneda,
fondo, fecha, monto repuesto, autorizó, registró, notas, periodo…); la TABLA
de lo que la reposición repone, en el orden del libro, con el saldo y lo por
reponer por fila; la fila TOTAL; el bloque de TOTALES (Σ gastos, reintegros
/ ajustes, repuesto, diferencia, saldo antes y después) y los AVISOS.

Aquí NO se calcula dinero: qué filas entran, sus saldos y los totales llegan
del API (fuente única `caja-chica-saldo.util.ts`). Solo se pinta.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.reportes import CajaChicaDato, CajaChicaFila, CajaChicaReposicionRequest

BRAND = "0F4C81"
LIGHT = "EEF2F7"
WHITE = "FFFFFF"
GRIS = "5B6470"
NARANJA = "ED7D31"
NARANJA_FONDO = "FFF3E6"
MONEY = '"$"#,##0.00'
FECHA = "dd/mm/yyyy"

_thin = Side(style="thin", color="D5DBE3")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

# (encabezado, ancho, es_monto) — el orden es el de `CajaChicaFila`.
COLUMNAS: list[tuple[str, float, bool]] = [
    ("Fecha", 12, False),
    ("Tipo", 20, False),
    ("Categoría", 22, False),
    ("Descripción / lugar", 44, False),
    ("Vuelo", 9, False),
    ("Matrícula", 11, False),
    ("Gasto", 14, True),
    ("Reintegro / ajuste (±)", 16, True),
    ("Comprobante", 17, False),
    ("Facturación", 19, False),
    ("Capturó", 20, False),
    ("Capturado (hora Cancún)", 19, False),
    ("Saldo del libro", 16, True),
    ("Por reponer", 14, True),
]
N_COLS = len(COLUMNAS)

# Excel prohíbe estos caracteres en el NOMBRE de la hoja.
_SHEET_ILEGALES = str.maketrans({c: "-" for c in "\\/*?:[]"})


def _texto_literal(celda) -> None:
    """Texto que EMPIEZA con «=» se queda como TEXTO (revisión 24-sep-2026).

    openpyxl convierte en FÓRMULA cualquier cadena que empiece con «=»: una
    nota de gasto como «= 3 taxis» daba un libro que Excel abre con «encontramos
    un problema con el contenido» (o, peor, una fórmula que se evalúa). Las
    notas y lugares los teclean pilotos y oficina: aquí son siempre texto.
    """
    if isinstance(celda.value, str) and celda.value.startswith("="):
        celda.data_type = "s"


def _nombre_hoja(nombre: str | None) -> str:
    limpio = (nombre or "Reposición").translate(_SHEET_ILEGALES).strip()
    return (limpio or "Reposición")[:31]


def _fecha(s: str | None) -> date | str:
    """'YYYY-MM-DD' → date (celda de FECHA real); otra cosa → texto tal cual."""
    if not s:
        return ""
    try:
        return date.fromisoformat(s[:10])
    except ValueError:
        return s


def _valores_fila(f: CajaChicaFila) -> list:
    return [
        _fecha(f.fecha),
        f.tipo,
        f.categoria,
        f.descripcion,
        f.vuelo,
        f.matricula,
        f.gasto,
        f.otro,
        f.comprobante,
        f.facturacion,
        f.capturo,
        f.capturado,
        f.saldo,
        f.por_reponer,
    ]


def _bloque_datos(
    ws: Worksheet, fila: int, titulo: str | None, datos: list[CajaChicaDato]
) -> int:
    """Pares etiqueta (A:C) → valor (D) → nota (E:H). Devuelve la fila libre."""
    if not datos:
        return fila
    if titulo:
        t = ws.cell(row=fila, column=1, value=titulo)
        t.font = Font(bold=True, size=11, color=BRAND)
        fila += 1
    for d in datos:
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=3)
        e = ws.cell(row=fila, column=1, value=d.etiqueta)
        _texto_literal(e)
        e.font = Font(bold=True)
        e.fill = PatternFill("solid", fgColor=LIGHT)
        e.alignment = Alignment(vertical="top", wrap_text=True)
        for col in (1, 2, 3, 4):
            ws.cell(row=fila, column=col).border = _border
        v = ws.cell(row=fila, column=4, value=d.valor)
        _texto_literal(v)
        v.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if isinstance(d.valor, (int, float)) and not isinstance(d.valor, bool):
            v.number_format = MONEY
        if d.destacado:
            v.font = Font(bold=True, size=12, color=BRAND)
            e.font = Font(bold=True, size=12, color=BRAND)
        if d.nota:
            ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=8)
            n = ws.cell(row=fila, column=5, value=d.nota)
            _texto_literal(n)
            n.font = Font(italic=True, color=GRIS)
            n.alignment = Alignment(vertical="top", wrap_text=True)
        fila += 1
    return fila


def render_caja_chica_reposicion_xlsx(req: CajaChicaReposicionRequest) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _nombre_hoja(req.hoja)

    for i, (_, ancho, _) in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    # Título y subtítulo.
    t = ws.cell(row=1, column=1, value=req.titulo)
    t.font = Font(bold=True, size=14, color=BRAND)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=N_COLS)
    fila = 2
    if req.subtitulo:
        s = ws.cell(row=2, column=1, value=req.subtitulo)
        s.font = Font(italic=True, size=10, color=GRIS)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=N_COLS)
        fila = 3

    # Encabezado (datos de la reposición).
    fila = _bloque_datos(ws, fila + 1, req.encabezado_titulo, req.encabezado)

    # Tabla.
    fila += 1
    for col, (label, _, _) in enumerate(COLUMNAS, start=1):
        c = ws.cell(row=fila, column=col, value=label)
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=BRAND)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border
    ws.row_dimensions[fila].height = 30
    fila += 1

    if not req.filas:
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=N_COLS)
        c = ws.cell(row=fila, column=1, value=req.sin_filas or "Sin movimientos en el periodo.")
        c.font = Font(italic=True, color=GRIS)
        fila += 1

    for f in req.filas:
        for col, (valor, (_, _, es_monto)) in enumerate(
            zip(_valores_fila(f), COLUMNAS, strict=True), start=1
        ):
            c = ws.cell(row=fila, column=col, value=valor)
            c.border = _border
            c.alignment = Alignment(vertical="top", wrap_text=col == 4)
            _texto_literal(c)
            if col == 1 and isinstance(valor, date):
                c.number_format = FECHA
            if es_monto and valor is not None:
                c.number_format = MONEY
        if f.resaltar:
            for col in (1, 12):
                c = ws.cell(row=fila, column=col)
                c.font = Font(bold=True, color=NARANJA)
                c.fill = PatternFill("solid", fgColor=NARANJA_FONDO)
        fila += 1

    # Fila TOTAL.
    totales = ["TOTAL", f"{req.n_gastos} {'gasto' if req.n_gastos == 1 else 'gastos'}"]
    totales += [None] * 4 + [req.total_gastos, req.total_otros] + [None] * 6
    for col, (valor, (_, _, es_monto)) in enumerate(zip(totales, COLUMNAS, strict=True), start=1):
        c = ws.cell(row=fila, column=col, value=valor)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.border = _border
        if es_monto and valor is not None:
            c.number_format = MONEY
    fila += 2

    # Totales de la reposición y avisos.
    fila = _bloque_datos(ws, fila, req.totales_titulo, req.totales)
    if req.avisos:
        fila += 1
        for aviso in req.avisos:
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=N_COLS)
            c = ws.cell(row=fila, column=1, value=aviso)
            _texto_literal(c)
            c.font = Font(bold=True, color=NARANJA)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[fila].height = 30
            fila += 1

    # Impresión: horizontal y a lo ancho de UNA hoja (la oficina lo imprime
    # para firmar la reposición).
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
