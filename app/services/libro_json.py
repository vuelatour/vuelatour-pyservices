"""Libro de Excel (openpyxl) → JSON para verlo en pantalla TAL CUAL (12-oct-2026).

Pedido del cliente: «ver el balance como el Excel sin necesidad de descargarlo
[…] con el mismo formato, los mismos colores y estructura, para que dé la
impresión de que está trabajando sobre el Excel» y, después, «editar la celda
y que afecte a lo relacionado». La vista del panel se arma del MISMO libro que
se descarga: aquí se recorre ese libro —después de verificar sus fórmulas igual
que `xlsx_formulas.guardar`— y se vuelca a un JSON con lo que el navegador
necesita para pintarlo como hoja de cálculo: hojas y color de pestaña, celdas
(valor + TEXTO ya formateado con el formato de número de Excel + la fórmula
visible), estilos (relleno, letra, alineación, bordes), celdas combinadas,
anchos de columna y altos de fila, paneles congelados, notas y el formato
condicional ya resuelto (el rojo de los negativos). No se calcula dinero: los
números son los que el libro ya trae del API.

Es GENÉRICO: sirve para cualquier libro openpyxl de este servicio (balances,
conciliación, caja chica…). Fase 2 (editar): cada celda admite un `vinculo`
(registro + campo del API) que quien construye el libro anota con
`vincular(cell, {...})`; este módulo solo lo transporta.

Contrato del JSON (versión 1; el panel lo tipa 1:1 en `types/libro-excel.ts`):
  {version, titulo, archivo, generado, ancho_columna_default,
   alto_fila_default, estilos: [Estilo], hojas: [Hoja], vinculos_gasto?}
  vinculos_gasto = true SOLO si quien pidió el libro lo pidió (el API, en el
           libro de la vista en pantalla): este libro marca las celdas de los
           gastos. Sin la bandera la raíz lleva EXACTAMENTE las 8 llaves de
           siempre.
  Hoja  = {nombre, pestana, cuadricula, inmovilizar: {filas, columnas} | null,
           max_fila, max_col, columnas: [{c, ancho?, oculta?}],
           filas: [{r, alto?, oculta?}], combinadas: [{r1, c1, r2, c2}],
           celdas: [Celda]}
  Celda = {r, c, v, t, f?, s?, n?, vinculo?, vinculos?}   v = valor crudo
           (número, texto, booleano o fecha ISO), t = texto como lo mostraría
           Excel, f = fórmula visible («=…»), s = índice en `estilos`, n =
           nota de la celda.
  Estilo = {bg?, co?, b?, i?, u?, sz?, ha?, va?, wr?, ind?, bd?, nf?}
           (relleno, color de letra, negrita, cursiva, subrayado, tamaño,
           alineación horizontal/vertical, ajuste de texto, sangría, bordes
           «lrtb», formato de número).
  Vinculo (en `vinculos`, 13-oct-2026, API 0.0.75: «abrir el balance en la
  celda del gasto» desde Conciliación) =
           {tipo: 'gasto', id, rol: 'fila'|'suma'|'nota'|'sin_tc',
            monto_mxn?, gasto_mxn?, parte?: 'reparto'|'remanente', fila?,
            columna?}
           rol 'fila' = la celda es el MONTO de la fila de ese gasto (hojas de
           gastos y combustible, VENTA AL AVIÓN de inventario); 'suma' = la
           celda lo suma junto con otros (OPERACIONES/PILOTO/OTROS de la hoja
           de vuelos); 'nota' = el gasto solo va en la nota de la celda (TUA o
           extensión de horario); 'sin_tc' = gasto en USD sin T.C. que todavía
           no suma. `monto_mxn` = lo que ese gasto aporta a la celda (pesos;
           None si no suma), `gasto_mxn` = el total del gasto cuando la celda
           lleva solo una parte, `parte` = parcial de un reparto manual,
           `fila` = clave del vuelo, `columna` = encabezado sin saltos.
           REGLA: aquí no se decide a qué libro ni a qué mes va un gasto; solo
           se transporta lo que el API ya acreditó (lo anota quien construye
           el libro con `vincular_lista(cell, [...])`).
"""

from __future__ import annotations

import math
import re
from datetime import date, datetime, time
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet

from app.services import xlsx_formulas

VERSION_LIBRO = 1
# Defaults de Excel (Calibri 11): los pinta el panel cuando la hoja no fija nada.
ANCHO_COLUMNA_DEFAULT = 8.43
ALTO_FILA_DEFAULT = 15.0
# Vínculos celda → registro del API (fase 2, edición). Las celdas de openpyxl
# no admiten atributos sueltos (`__slots__`), así que viajan en un registro
# PEGADO al libro, como las fórmulas de `xlsx_formulas` (dos libros a la vez
# en hilos distintos no se mezclan).
_ATRIBUTO_VINCULOS = "_vt_vinculos"
# Vínculos en LISTA (13-oct-2026, punto 5): los gastos que componen una celda.
# Mismo patrón: registro pegado al libro, por hoja y coordenada.
_ATRIBUTO_VINCULOS_LISTA = "_vt_vinculos_lista"


def vincular(cell: Cell, vinculo: dict[str, Any]) -> Cell:
    """Anota en la celda su vínculo al registro del API (p. ej.
    `{"tipo": "gasto", "id": "...", "campo": "monto"}`): la vista lo usa para
    saber qué editar. Solo se transporta; nada se interpreta aquí."""
    ws = cell.parent
    wb = ws.parent
    reg = getattr(wb, _ATRIBUTO_VINCULOS, None)
    if reg is None:
        reg = {}
        setattr(wb, _ATRIBUTO_VINCULOS, reg)
    reg.setdefault(id(ws), {})[cell.coordinate] = dict(vinculo)
    return cell


def _vinculo_de(cell: Cell) -> dict[str, Any] | None:
    reg = getattr(cell.parent.parent, _ATRIBUTO_VINCULOS, None)
    if not reg:
        return None
    return reg.get(id(cell.parent), {}).get(cell.coordinate)


def vincular_lista(cell: Cell, vinculos: list[dict[str, Any]]) -> Cell:
    """Anota en la celda los gastos que la componen (contrato `Vinculo` de la
    cabecera): la vista del panel ubica con ellos la celda de un gasto. Guarda
    COPIAS (quien llama puede reutilizar sus dicts); una segunda llamada sobre
    la misma celda REEMPLAZA la lista; una lista vacía no anota nada. Solo se
    transporta: aquí no se interpreta ni se decide nada del dinero."""
    if not vinculos:
        return cell
    ws = cell.parent
    wb = ws.parent
    reg = getattr(wb, _ATRIBUTO_VINCULOS_LISTA, None)
    if reg is None:
        reg = {}
        setattr(wb, _ATRIBUTO_VINCULOS_LISTA, reg)
    reg.setdefault(id(ws), {})[cell.coordinate] = [dict(v) for v in vinculos]
    return cell


def _vinculos_de(cell: Cell) -> list[dict[str, Any]] | None:
    reg = getattr(cell.parent.parent, _ATRIBUTO_VINCULOS_LISTA, None)
    if not reg:
        return None
    return reg.get(id(cell.parent), {}).get(cell.coordinate)

# ---------------------------------------------------------------------------
# Formato de número «como Excel»
# ---------------------------------------------------------------------------

_RE_COLOR_SECCION = re.compile(r"\[[^\]]*\]")
_CARACTERES_NUCLEO = set("0#?.,%")


def _partir_seccion(seccion: str) -> tuple[str, str, str]:
    """Una sección de formato → (prefijo literal, núcleo numérico, sufijo
    literal). Los literales van entre comillas o escapados con «\\»; «$»,
    «-» y espacios sueltos también son literales para Excel."""
    prefijo, nucleo, sufijo = [], [], []
    i, n = 0, len(seccion)
    while i < n:
        ch = seccion[i]
        if ch == '"':
            fin = seccion.find('"', i + 1)
            if fin < 0:
                fin = n
            literal = seccion[i + 1 : fin]
            (sufijo if nucleo else prefijo).append(literal)
            i = fin + 1
            continue
        if ch == "\\" and i + 1 < n:
            (sufijo if nucleo else prefijo).append(seccion[i + 1])
            i += 2
            continue
        if ch in _CARACTERES_NUCLEO:
            if sufijo:
                # Un segundo bloque numérico (p. ej. «0.00 "h" 0») no se
                # modela: se trata como literal para no inventar un número.
                sufijo.append(ch)
            else:
                nucleo.append(ch)
            i += 1
            continue
        if ch in "_*":
            # «_x» reserva el ancho de x y «*x» repite x: solo afectan el
            # espaciado en Excel; en pantalla no se pintan.
            i += 2
            continue
        (sufijo if nucleo else prefijo).append(ch)
        i += 1
    return "".join(prefijo), "".join(nucleo), "".join(sufijo)


def _general(valor: float | int) -> str:
    """Formato «General» de Excel: entero sin decimales, hasta 10 cifras
    significativas sin ceros de cola."""
    if isinstance(valor, int):
        return str(valor)
    if valor.is_integer() and abs(valor) < 1e15:
        return str(int(valor))
    texto = f"{valor:.10g}"
    if "e" in texto or "E" in texto:
        texto = f"{valor:.10f}".rstrip("0").rstrip(".")
    return texto


def _numero_con_seccion(valor: float, seccion: str) -> str:
    seccion = _RE_COLOR_SECCION.sub("", seccion)
    prefijo, nucleo, sufijo = _partir_seccion(seccion)
    if not nucleo:
        return f"{prefijo}{sufijo}"
    if "%" in nucleo:
        valor *= 100
        nucleo = nucleo.replace("%", "")
        sufijo = "%" + sufijo
    entero, _, decimales = nucleo.partition(".")
    min_dec = decimales.count("0")
    max_dec = min_dec + decimales.count("#") + decimales.count("?")
    miles = "," in entero
    redondeado = xlsx_formulas.redondear_excel(valor, max_dec)
    texto = f"{redondeado:,.{max_dec}f}" if miles else f"{redondeado:.{max_dec}f}"
    if max_dec > min_dec and "." in texto:
        entero_txt, _, dec_txt = texto.partition(".")
        dec_txt = dec_txt.rstrip("0")
        if len(dec_txt) < min_dec:
            dec_txt = dec_txt.ljust(min_dec, "0")
        texto = f"{entero_txt}.{dec_txt}" if dec_txt else entero_txt
    return f"{prefijo}{texto}{sufijo}"


def _fecha_texto(valor: datetime | date | time, fmt: str) -> str:
    if isinstance(valor, time):
        return valor.strftime("%H:%M")
    f = (fmt or "").lower()
    if "yyyy" in f and f.index("yyyy") < f.find("d"):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, datetime) and ("h" in f and ":" in f):
        return valor.strftime("%d/%m/%Y %H:%M")
    return valor.strftime("%d/%m/%Y")


def formatear(valor: Any, fmt: str | None) -> str:
    """Texto de una celda tal como lo MOSTRARÍA Excel con su formato de
    número (`#,##0.00`, `"$"#,##0.00`, `0.00####`, `0.0%`, `0.00"%"`,
    secciones «positivo;negativo;cero», `General`). Vacío ⇒ ''."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "VERDADERO" if valor else "FALSO"
    if isinstance(valor, datetime | date | time):
        return _fecha_texto(valor, fmt or "")
    if isinstance(valor, str):
        return valor
    if not isinstance(valor, int | float):
        return str(valor)
    if isinstance(valor, float) and (math.isnan(valor) or math.isinf(valor)):
        return ""
    fmt = fmt or "General"
    if fmt == "General":
        return _general(valor)
    secciones = fmt.split(";")
    negativo = valor < 0
    if negativo and len(secciones) >= 2 and secciones[1]:
        # La sección negativa trae su propio signo o paréntesis.
        return _numero_con_seccion(abs(float(valor)), secciones[1])
    if valor == 0 and len(secciones) >= 3 and secciones[2]:
        return _numero_con_seccion(0.0, secciones[2])
    texto = _numero_con_seccion(abs(float(valor)), secciones[0])
    return f"-{texto}" if negativo else texto


# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------


def _rgb(color: Any) -> str | None:
    """Color openpyxl → «RRGGBB» (solo colores RGB explícitos; un color de
    tema o indexado es el default y no se manda)."""
    if color is None:
        return None
    if getattr(color, "type", None) != "rgb":
        return None
    rgb = getattr(color, "rgb", None)
    if not isinstance(rgb, str) or len(rgb) < 6:
        return None
    return rgb[-6:].upper()


def _color_relleno(fill: Any) -> str | None:
    if fill is None or getattr(fill, "fill_type", None) != "solid":
        return None
    return _rgb(getattr(fill, "fgColor", None)) or _rgb(getattr(fill, "bgColor", None))


def _estilo(cell: Cell) -> dict[str, Any] | None:
    e: dict[str, Any] = {}
    bg = _color_relleno(cell.fill)
    if bg:
        e["bg"] = bg
    font = cell.font
    if font is not None:
        co = _rgb(font.color)
        if co and co != "000000":
            e["co"] = co
        if font.b:
            e["b"] = True
        if font.i:
            e["i"] = True
        if font.u:
            e["u"] = True
        if font.sz and float(font.sz) != 11.0:
            e["sz"] = float(font.sz)
    al = cell.alignment
    if al is not None:
        if al.horizontal and al.horizontal != "general":
            e["ha"] = al.horizontal
        if al.vertical and al.vertical != "bottom":
            e["va"] = al.vertical
        if al.wrap_text:
            e["wr"] = True
        if al.indent:
            e["ind"] = int(al.indent)
    bd = cell.border
    if bd is not None:
        lados = "".join(
            letra
            for letra, lado in (("l", bd.left), ("r", bd.right), ("t", bd.top), ("b", bd.bottom))
            if lado is not None and lado.style
        )
        if lados:
            e["bd"] = lados
    nf = cell.number_format
    if nf and nf != "General":
        e["nf"] = nf
    return e or None


class _Estilos:
    """Tabla de estilos sin repetir: el JSON lleva cada combinación UNA vez y
    las celdas la citan por índice (un libro grande repite el mismo estilo
    miles de veces)."""

    def __init__(self) -> None:
        self.lista: list[dict[str, Any]] = []
        self._indice: dict[tuple[tuple[str, Any], ...], int] = {}

    def indice(self, estilo: dict[str, Any]) -> int:
        clave = tuple(sorted(estilo.items()))
        idx = self._indice.get(clave)
        if idx is None:
            idx = len(self.lista)
            self.lista.append(dict(estilo))
            self._indice[clave] = idx
        return idx


# ---------------------------------------------------------------------------
# Celdas
# ---------------------------------------------------------------------------


def _json_valor(valor: Any) -> Any:
    if valor is None or isinstance(valor, bool | str):
        return valor
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return None if math.isnan(valor) or math.isinf(valor) else valor
    if isinstance(valor, datetime | date | time):
        return valor.isoformat()
    return str(valor)


def _celda(cell: Cell, reg: xlsx_formulas.RegistroFormulas) -> dict[str, Any] | None:
    valor = cell.value
    formula: str | None = None
    if isinstance(valor, str) and valor.startswith("="):
        registrada = reg.de(cell.parent, cell.coordinate)
        if registrada is not None and not registrada.degradada:
            formula = valor
            valor = registrada.valor
        else:
            # Fórmula sin número verificado: se muestra la fórmula, nunca un
            # número inventado (no debería ocurrir: `_formula` registra todas).
            formula = valor
            valor = None
    estilo = _estilo(cell)
    nota = cell.comment.text if cell.comment is not None else None
    vinculo = _vinculo_de(cell)
    vinculos = _vinculos_de(cell)
    # Una celda sin valor, estilo ni nota pero CON vínculos sí viaja: la vista
    # la necesita para ubicar el gasto.
    if (
        valor is None
        and formula is None
        and estilo is None
        and nota is None
        and vinculo is None
        and not vinculos
    ):
        return None
    d: dict[str, Any] = {
        "r": cell.row,
        "c": cell.column,
        "v": _json_valor(valor),
        "t": formatear(valor, cell.number_format),
    }
    if formula is not None:
        d["f"] = formula
    if nota:
        d["n"] = nota
    if vinculo:
        d["vinculo"] = vinculo
    if vinculos:
        d["vinculos"] = [dict(x) for x in vinculos]
    if estilo:
        d["_estilo"] = estilo
    return d


def _cumple(operador: str, v: float, umbrales: list[float]) -> bool:
    try:
        if operador == "lessThan":
            return v < umbrales[0]
        if operador == "lessThanOrEqual":
            return v <= umbrales[0]
        if operador == "greaterThan":
            return v > umbrales[0]
        if operador == "greaterThanOrEqual":
            return v >= umbrales[0]
        if operador == "equal":
            return v == umbrales[0]
        if operador == "notEqual":
            return v != umbrales[0]
        if operador == "between":
            return min(umbrales[:2]) <= v <= max(umbrales[:2])
        if operador == "notBetween":
            return not (min(umbrales[:2]) <= v <= max(umbrales[:2]))
    except IndexError:
        return False
    return False


def _aplicar_condicional(ws: Worksheet, por_coord: dict[str, dict[str, Any]]) -> None:
    """Resuelve el formato condicional «el valor es …» (cellIs) con el número
    que la celda muestra —como lo haría Excel al abrir— y lo deja escrito en
    el estilo de la celda. La primera regla que aplica gana."""
    for cf in ws.conditional_formatting:
        for regla in cf.rules:
            if getattr(regla, "type", None) != "cellIs" or not regla.operator:
                continue
            try:
                umbrales = [float(x) for x in (regla.formula or [])]
            except (TypeError, ValueError):
                continue
            dxf = regla.dxf
            bg = _color_relleno(dxf.fill) if dxf is not None and dxf.fill is not None else None
            co = _rgb(dxf.font.color) if dxf is not None and dxf.font is not None else None
            if not bg and not co:
                continue
            for rango in cf.sqref.ranges:
                for fila, col in rango.cells:
                    d = por_coord.get(f"{get_column_letter(col)}{fila}")
                    if d is None or d.get("_cf"):
                        continue
                    v = d.get("v")
                    if isinstance(v, bool) or not isinstance(v, int | float):
                        continue
                    if not _cumple(regla.operator, float(v), umbrales):
                        continue
                    d["_cf"] = True
                    estilo = dict(d.get("_estilo") or {})
                    if bg:
                        estilo["bg"] = bg
                    if co:
                        estilo["co"] = co
                    d["_estilo"] = estilo


# ---------------------------------------------------------------------------
# Hojas y libro
# ---------------------------------------------------------------------------


def _columnas(ws: Worksheet) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for letra, cd in ws.column_dimensions.items():
        if cd.width is None and not cd.hidden:
            continue
        ini = cd.min or column_index_from_string(letra)
        fin = cd.max or ini
        for c in range(ini, fin + 1):
            d: dict[str, Any] = {"c": c}
            if cd.width is not None:
                d["ancho"] = float(cd.width)
            if cd.hidden:
                d["oculta"] = True
            out.append(d)
    out.sort(key=lambda d: d["c"])
    return out


def _filas(ws: Worksheet) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for r, rd in ws.row_dimensions.items():
        if rd.height is None and not rd.hidden:
            continue
        d: dict[str, Any] = {"r": int(r)}
        if rd.height is not None:
            d["alto"] = float(rd.height)
        if rd.hidden:
            d["oculta"] = True
        out.append(d)
    out.sort(key=lambda d: d["r"])
    return out


def _inmovilizar(ws: Worksheet) -> dict[str, int] | None:
    ref = ws.freeze_panes
    if not ref:
        return None
    col_letra, fila = coordinate_from_string(str(ref))
    return {"filas": fila - 1, "columnas": column_index_from_string(col_letra) - 1}


def _hoja(ws: Worksheet, reg: xlsx_formulas.RegistroFormulas, estilos: _Estilos) -> dict[str, Any]:
    celdas: list[dict[str, Any]] = []
    por_coord: dict[str, dict[str, Any]] = {}
    for fila in ws.iter_rows():
        for cell in fila:
            if isinstance(cell, MergedCell):
                continue
            d = _celda(cell, reg)
            if d is None:
                continue
            celdas.append(d)
            por_coord[cell.coordinate] = d
    _aplicar_condicional(ws, por_coord)
    for d in celdas:
        d.pop("_cf", None)
        estilo = d.pop("_estilo", None)
        if estilo:
            d["s"] = estilos.indice(estilo)
    combinadas = [
        {"r1": rg.min_row, "c1": rg.min_col, "r2": rg.max_row, "c2": rg.max_col}
        for rg in ws.merged_cells.ranges
    ]
    combinadas.sort(key=lambda x: (x["r1"], x["c1"]))
    max_fila = max([ws.max_row, *(d["r"] for d in celdas), *(x["r2"] for x in combinadas)])
    max_col = max([ws.max_column, *(d["c"] for d in celdas), *(x["c2"] for x in combinadas)])
    tab = ws.sheet_properties.tabColor
    return {
        "nombre": ws.title,
        "pestana": _rgb(tab) if tab is not None else None,
        "cuadricula": ws.sheet_view.showGridLines is not False,
        "inmovilizar": _inmovilizar(ws),
        "max_fila": int(max_fila),
        "max_col": int(max_col),
        "columnas": _columnas(ws),
        "filas": _filas(ws),
        "combinadas": combinadas,
        "celdas": celdas,
    }


def exportar_libro(
    wb: Workbook,
    *,
    titulo: str,
    archivo: str | None = None,
    generado: str | None = None,
    vinculos_gasto: bool = False,
) -> dict[str, Any]:
    """El libro completo como JSON para la vista en pantalla. Verifica antes
    las fórmulas (las que no reproducen al API vuelven a valor, igual que al
    guardar el .xlsx): lo que se ve en pantalla es lo que se descargaría.
    `vinculos_gasto` (13-oct-2026): agrega `"vinculos_gasto": true` como
    ÚLTIMA llave de la raíz (el libro marca las celdas de los gastos); sin él,
    la raíz de siempre, byte a byte."""
    xlsx_formulas.asegurar_valores(wb)
    reg = xlsx_formulas.registro(wb)
    estilos = _Estilos()
    hojas = [_hoja(ws, reg, estilos) for ws in wb.worksheets]
    libro: dict[str, Any] = {
        "version": VERSION_LIBRO,
        "titulo": titulo,
        "archivo": archivo,
        "generado": generado,
        "ancho_columna_default": ANCHO_COLUMNA_DEFAULT,
        "alto_fila_default": ALTO_FILA_DEFAULT,
        "estilos": estilos.lista,
        "hojas": hojas,
    }
    if vinculos_gasto:
        libro["vinculos_gasto"] = True
    return libro
