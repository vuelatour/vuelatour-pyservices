"""Auditoría de las NOTAS DE CELDA de los libros (26-oct-2026).

Pedido del cliente, con la captura de la hoja «balance» del Balance general
(celda B6 y su nota): «esa nota que está en la celda B6 y B7 y todas están
actualizadas?, porque unas cosas nos causan ruido». Regla desde ese día: una
nota de celda dice UNA idea en 1–2 líneas, en español llano (qué es la cifra
y de dónde sale), sin fechas, versiones ni invariantes del código, sin
repetir el rótulo de la fila ni explicar una fórmula que ya se ve en la
celda; y dice lo VIGENTE (la reserva de overhaul no resta, los gastos
directos van sin combustible, la clave es solo el folio, el cierre es por
avión, la provisión AFAC sí resta…).

(a) Recorre TODAS las notas que generan los libros con los payloads de prueba
    (balance por avión, Balance general en sus variantes, Libro Dinero y
    reparto) y exige que ninguna diga algo ya superado ni sea ruidosa. Las
    notas con DATOS (cobros, desgloses, detalle de la IA) se miden por
    renglón: una línea por concepto, ninguna de más de 240 caracteres.
(b) Fija el texto EXACTO de los rótulos y las notas de la cascada de la hoja
    «balance» (en pesos, lo que ve el cliente; y en dólares), en el Balance
    general y en el balance por avión, y que la vista en pantalla
    (`libro_json`) lleva la misma nota.
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Callable
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.schemas.reparto import RepartoPdfRequest
from app.schemas.reportes import BalanceAvionRequest, BalanceGeneralRequest, DineroXlsxRequest
from app.services.balance_avion_xlsx import (
    libro_balance_general,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from app.services.dinero_xlsx import render_dinero_xlsx
from app.services.reparto_xlsx import render_reparto_xlsx
from tests import _afac_vuelo_payloads as pa
from tests import _cierre_mes_payloads as pay
from tests import _pesos_payloads as pp
from tests import test_balance_comisiones as t_com
from tests import test_balance_comisiones_general as t_comg
from tests import test_balance_entregas_socios as t_ent
from tests import test_balance_extension_horario as t_ext
from tests import test_balance_formulas as tf
from tests import test_balance_general_hojas as t_hojas
from tests import test_balance_inventario_xlsx as t_inv
from tests import test_balance_otros_ingresos_avion as t_oi
from tests import test_balance_otros_movimientos_afac_vuelo as t_afacv
from tests import test_balance_reserva_overhaul as t_res
from tests import test_claves_sin_apodo as t_cla
from tests import test_comision_vendedor_notas as t_cvn
from tests import test_dinero_otros_ingresos_avion as t_din
from tests import test_reparto_combustible_separado as t_rcomb
from tests import test_reparto_como_balance as t_rcb

# ---------------------------------------------------------------------------
# Libros de prueba: uno por cada fuente de notas (cada bandera que cambia un
# texto tiene aquí un libro con ella y otro sin ella).
# ---------------------------------------------------------------------------


def _ind(p) -> bytes:
    req = p if isinstance(p, BalanceAvionRequest) else BalanceAvionRequest.model_validate(p)
    return render_balance_avion_xlsx(req)


def _gen(p) -> bytes:
    req = p if isinstance(p, BalanceGeneralRequest) else BalanceGeneralRequest.model_validate(p)
    return render_balance_general_xlsx(req)


def _din(p) -> bytes:
    req = p if isinstance(p, DineroXlsxRequest) else DineroXlsxRequest.model_validate(p)
    return render_dinero_xlsx(req)


def _rep(p) -> bytes:
    req = p if isinstance(p, RepartoPdfRequest) else RepartoPdfRequest.model_validate(p)
    return render_reparto_xlsx(req)


def _cph(req: BalanceGeneralRequest) -> BalanceGeneralRequest:
    return req.model_copy(update={"variante": "general"})


_LIBROS: dict[str, Callable[[], bytes]] = {
    # Balance por avión.
    "individual XA-TST": lambda: _ind(tf._individual_tst()),
    "individual XB-DOS sin AFAC": lambda: _ind(tf._individual_dos()),
    "individual permisos con T.C.": lambda: _ind(tf._payload_tst_permisos_tc(con_usd_sin_tc=True)),
    "individual hojas": lambda: _ind(t_hojas._individual()),
    "individual extensión": lambda: _ind(t_ext._individual(con_llave=True)),
    "individual comisiones": lambda: _ind(t_com._individual()),
    "individual comisión vendedor": lambda: _ind(t_cvn._individual()),
    "individual pesos + reserva": lambda: _ind(t_res.individual()),
    "individual USD + reserva": lambda: _ind(t_res.individual(bandera=None)),
    "individual reserva sin T.C.": lambda: _ind(t_res.individual(sin_tc=True)),
    "individual USD por gasto": lambda: _ind(t_res.con_usd_por_gasto(t_res.individual())),
    "individual entregas": lambda: _ind(t_ent.individual(ex_socio=True)),
    "individual entregas sin pesos": lambda: _ind(t_ent.sin_pesos(t_ent.individual())),
    "individual con cierre": lambda: _ind(pay.individual_con_cierre()),
    "individual AFAC 0.0.89": lambda: _ind(pa.individual_0089()),
    "individual otros ingresos por ingreso": lambda: _ind(t_oi._individual(por_ingreso=True)),
    "individual claves viejas": lambda: _ind(t_cla._individual(t_cla._vuelos_viejos())),
    # Balance general (mensual y costo por hora).
    "general": lambda: _gen(tf._general()),
    "general costo por hora": lambda: _gen(_cph(tf._general())),
    "general tarifa AFAC por fila": lambda: _gen(tf._general_con_tarifa_afac()),
    "general hojas": lambda: _gen(t_hojas._general()),
    "general extensión costo por hora": lambda: _gen(_cph(t_ext._general(con_llave=True))),
    "general inventario": lambda: _gen(t_inv._general_0036(t_inv._TIENDA_SEP26_0074)),
    "general comisión vendedor v73": lambda: _gen(t_cvn._payload_regla_v73()),
    "general comisiones 0.0.74": lambda: _gen(t_comg._payload_0074()),
    "general comisiones 0.0.74 costo por hora": lambda: _gen(t_comg._payload_0074("general")),
    "general comisiones 0.0.66": lambda: _gen(t_comg._payload_om_0066()),
    "general externos": lambda: _gen(t_com._general_con_externos()),
    "general pesos + reserva": lambda: _gen(t_res.general()),
    "general pesos + reserva costo por hora": lambda: _gen(t_res.general("general")),
    "general USD + reserva": lambda: _gen(t_res.general(bandera=None)),
    "general reserva sin T.C.": lambda: _gen(t_res.general(sin_tc=True)),
    "general entregas": lambda: _gen(t_ent.general()),
    "general con cierre": lambda: _gen(pay.general_con_cierre("mensual")),
    "general AFAC 0.0.89": lambda: _gen(pa.general_0089()),
    "general AFAC 0.0.88": lambda: _gen(pa.general_0088("mensual")),
    "general AFAC por vuelo + empresa": lambda: _gen(t_afacv._req_empresa(afac_por_vuelo=True)),
    "general otros ingresos forma real": lambda: _gen(t_oi._general_forma_real(api_0077=True)),
    "general claves": lambda: _gen(t_cla._general(variante="general")),
    # Libro Dinero.
    "dinero otros ingresos": lambda: _din(t_din._req()),
    "dinero con cierre": lambda: _din(pay.dinero_con_cierre()),
    "dinero dos categorías": lambda: _din(t_cvn._dinero_dos()),
    "dinero claves": lambda: _din(t_cla._dinero()),
    # Reparto a socios (Excel).
    "reparto como balance pesos": lambda: _rep(t_rcb._req()),
    "reparto como balance USD": lambda: _rep(t_rcb._req(pesos=None)),
    "reparto combustible separado": lambda: _rep(t_rcomb._req(otros_ingresos=100.0)),
    "reparto combustible separado USD": lambda: _rep(t_rcomb._req(pesos=False)),
    "reparto pesos": lambda: _rep(pp.reparto()),
}


# ---------------------------------------------------------------------------
# (a) Ninguna nota desactualizada ni ruidosa.
# ---------------------------------------------------------------------------

LARGO_MAX = 240


def _plano(texto: str) -> str:
    """Minúsculas, sin acentos y con espacios simples (para buscar patrones)."""
    sin = unicodedata.normalize("NFKD", texto)
    sin = "".join(ch for ch in sin if not unicodedata.combining(ch))
    return " ".join(sin.lower().split())


# Lo que una nota YA NO puede decir (patrones sobre el texto plano, sin las
# negaciones «no (se) resta»): la regla vigente está en CLAUDE.md.
_DESACTUALIZADOS = (
    (r"\breserva\b[^.;\n]{0,80}\brestan?\b", "la reserva de overhaul resta"),
    (r"\bid (del|por) cliente\b|\balias del cliente\b", "la clave lleva el ID del cliente"),
    (r"\bcierre formal del mes\b", "cierre global del mes"),
    (r"\bbalance general abreviado\b", "un Balance general abreviado"),
    (r"\bdirectos\b[^.;\n]{0,40}\b(incluyen?|con|junto)\b[^.;\n]{0,20}\bcombustible\b",
     "gastos directos con combustible"),
    (r"\bafac\b[^.;\n]{0,60}\bno restan?\b", "la provisión AFAC no resta"),
    (r"\b0\.0\.\d+\b", "versión del API"),
    (r"\binvariante\b|\bhead\b", "invariante o commit del código"),
    (r"\bregla (de|desde) \w+ \d{4}\b", "la fecha de la regla de comisiones"),
    (r"\b\d{1,2}-(ene|feb|mar|abr|may|jun|jul|ago|sep|oct|nov|dic)-\d{4}\b", "una fecha"),
)
_NEGACIONES = re.compile(r"\b(no|ni|nunca) (se |la |lo )?restan?\b")
_SIGNO = re.compile(r"^\((−|\+|-)\)\s*")
# Nota con DATOS (importes o fechas): sus renglones los arma el API o el
# detalle del libro (desglose, cobros, AFAC) y pueden empezar por el nombre de
# su columna («Operaciones · …»); la regla del rótulo es para las editoriales.
_DATOS = re.compile(r"\$\d|\d{2}/\d{2}/\d{4}")


def _desactualizado(nota: str) -> list[str]:
    plano = _NEGACIONES.sub(" ", _plano(nota))
    return [que for patron, que in _DESACTUALIZADOS if re.search(patron, plano)]


def _ruido(nota: str, rotulo_fila: str | None, encabezado: str | None) -> list[str]:
    problemas: list[str] = []
    lineas = nota.split("\n")
    if len(lineas) == 1 and len(nota) > LARGO_MAX:
        problemas.append(f"{len(nota)} caracteres (más de {LARGO_MAX})")
    if any(len(linea) > LARGO_MAX for linea in lineas):
        problemas.append(f"un renglón de más de {LARGO_MAX} caracteres")
    if nota.count("(MXN)") > 1:
        problemas.append("«(MXN)» más de una vez")
    plano = _plano(nota)
    for rotulo in (rotulo_fila, encabezado):
        if not rotulo or _DATOS.search(nota):
            continue
        limpio = _plano(_SIGNO.sub("", rotulo))
        if len(limpio) >= 8 and plano.startswith(limpio):
            problemas.append(f"empieza por el rótulo «{rotulo}»")
    return problemas


def _encabezado_de(ws, cell) -> str | None:
    """Encabezado de la columna de la celda: el texto de la misma columna en
    las filas de encabezado de los libros (2, 4 o 7), si lo hay."""
    for fila in (2, 4, 7):
        if fila == cell.row:
            continue
        v = ws.cell(row=fila, column=cell.column).value
        if isinstance(v, str) and v.strip():
            return v
    return None


def _notas(data: bytes):
    wb = load_workbook(BytesIO(data))
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for cell in fila:
                if cell.comment is None:
                    continue
                a = ws.cell(row=cell.row, column=1).value if cell.column != 1 else None
                yield ws.title, cell.coordinate, cell.comment.text, (
                    a if isinstance(a, str) else None
                ), _encabezado_de(ws, cell)


@pytest.mark.parametrize("libro", list(_LIBROS))
def test_ninguna_nota_desactualizada_ni_ruidosa(libro) -> None:
    data = _LIBROS[libro]()
    fallas: list[str] = []
    n = 0
    for hoja, celda, nota, rotulo, encabezado in _notas(data):
        n += 1
        for que in _desactualizado(nota):
            fallas.append(f"{hoja}!{celda} dice algo superado ({que}): {nota[:90]!r}")
        for que in _ruido(nota, rotulo, encabezado):
            fallas.append(f"{hoja}!{celda} con ruido ({que}): {nota[:90]!r}")
    assert n > 0, "el libro no generó ninguna nota: el caso no prueba nada"
    assert not fallas, "\n".join(fallas)


# ---------------------------------------------------------------------------
# (b) La hoja «balance»: rótulos y notas EXACTOS de la cascada.
# ---------------------------------------------------------------------------

ANTES_COMISIONES_MXN = "UTILIDAD ANTES DE COMISIONES (MXN)"
DESPUES_COMISIONES_MXN = "UTILIDAD DESPUÉS DE COMISIONES (MXN)"
ANTES_COMISIONES_USD = "UTILIDAD ANTES DE COMISIONES USD"
DESPUES_COMISIONES_USD = "UTILIDAD DESPUÉS DE COMISIONES USD"

NOTA_ANTES_COMISIONES_MXN = (
    "Utilidad de los vuelos del avión antes de restar comisiones (la hoja de vuelos de "
    "este libro va antes de comisiones)."
)
NOTA_COMISIONES_MXN = (
    "Comisiones a cargo del avión (banco y vendedor): la suma de la columna COMISIONES de "
    "su Balance por avión."
)
NOTA_DESPUES_COMISIONES = (
    "Utilidad de los vuelos del avión ya con las comisiones restadas; de aquí se restan "
    "combustible, gastos indirectos, refacciones y permisos."
)
NOTA_OTROS_INGRESOS_MXN = (
    "Otros ingresos asignados a este avión en Ingresos: la suma del NETO MXN de la hoja "
    "'Otros ingresos' de su libro (ingreso menos comisión bancaria)."
)
NOTA_ANTES_COMISIONES_USD = (
    "Utilidad de los vuelos del avión antes de restar comisiones (la hoja de vuelos de "
    "este libro va antes de comisiones). Puede no cuadrar exacto con la suma de GANANCIA "
    "USD de esa hoja: cada vuelo convierte con su propio T.C."
)
NOTA_COMISIONES_USD = (
    "Comisiones a cargo del avión en dólares (su parte de la comisión bancaria de sus "
    "cobros y la provisión de la comisión del vendedor): las de la columna COMISIONES de "
    "su Balance por avión."
)
NOTA_COMISIONES_USD_GASTO_REAL = (
    "Comisiones a cargo del avión en dólares (su parte de la comisión bancaria de sus "
    "cobros y el pago «Comisión del vendedor (avión)» ligado al vuelo): las de la columna "
    "COMISIONES de su Balance por avión."
)
NOTA_OTROS_INGRESOS_USD = (
    "Otros ingresos asignados a este avión en Ingresos (hoja 'Otros ingresos' de su "
    "libro), netos de comisión bancaria, al T.C. promedio del libro."
)
NOTA_OTROS_INGRESOS_USD_POR_INGRESO = (
    "Otros ingresos asignados a este avión en Ingresos (hoja 'Otros ingresos' de su "
    "libro), netos de comisión bancaria, cada uno a dólares con su propio T.C., como en "
    "el reparto a socios."
)


def _monto(x: float) -> str:
    return f"-${abs(x):,.2f}" if x < 0 else f"${x:,.2f}"


def _nota_indirectos(otros: float | None, moneda: str, *, general: bool) -> str | None:
    if not otros:
        return None
    if general:
        return (
            f"Incluye {_monto(otros)} {moneda} repartidos a mano a este avión (detalle en la "
            "hoja repartidos a aviones)."
        )
    return (
        f"Incluye {_monto(otros)} {moneda} de gastos administrativos repartidos a mano "
        "(filas reparto manual de la hoja Gastos Indirectos)."
    )


def _cascada(ws, primero: str) -> list[tuple[str, str | None]]:
    """(rótulo, nota de B) de la cascada que empieza en `primero`, hasta
    UTILIDAD COBRADA."""
    k = next(c.row for c in ws["A"] if c.value == primero)
    filas: list[tuple[str, str | None]] = []
    for r in range(k, k + 12):
        rotulo = ws.cell(row=r, column=1).value
        nota = ws.cell(row=r, column=2).comment
        filas.append((rotulo, nota.text if nota is not None else None))
        if isinstance(rotulo, str) and rotulo.startswith("UTILIDAD COBRADA"):
            return filas
    raise AssertionError("la cascada no termina en UTILIDAD COBRADA")


def _balance(data: bytes):
    return load_workbook(BytesIO(data))["balance"]


def _textos_a(ws) -> list[str]:
    return [c.value for c in ws["A"] if isinstance(c.value, str)]


def _sin_rotulo_viejo(ws) -> bool:
    """Ningún rótulo de la columna A dice ya «ANTES DE GASTOS»."""
    return not any("ANTES DE GASTOS" in t for t in _textos_a(ws))


def _avion(p: dict, matricula: str) -> dict:
    return next(a for a in p["aviones"] if a["matricula"] == matricula)


def test_balance_general_en_pesos_rotulos_y_notas_exactos() -> None:
    p = pp.general()
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    ws = _balance(data)
    tst = _avion(p, "XA-TST")
    assert tst["balance"]["comisiones_usd"] and tst["otros_ingresos_avion"] is not None
    assert _cascada(ws, ANTES_COMISIONES_MXN) == [
        (ANTES_COMISIONES_MXN, NOTA_ANTES_COMISIONES_MXN),
        ("(−) COMISIONES BANCO + VENDEDOR (MXN)", NOTA_COMISIONES_MXN),
        (DESPUES_COMISIONES_MXN, NOTA_DESPUES_COMISIONES),
        ("(−) COMBUSTIBLE DEL MES (MXN)", None),
        ("(−) GASTOS INDIRECTOS (MXN)", _nota_indirectos(
            tst["balance"]["otros_mxn"], "MXN", general=True
        )),
        ("(−) REFACCIONES (INVENTARIO) (MXN)", None),
        ("(−) PERMISOS (MXN)", None),
        ("(+) OTROS INGRESOS (MXN)", NOTA_OTROS_INGRESOS_MXN),
        ("UTILIDAD DESPUÉS DE GASTOS (MXN)", None),
        ("(−) PENDIENTE DE PAGO (COBRANZA PENDIENTE) (MXN)", None),
        ("UTILIDAD COBRADA (MXN)", None),
    ]  # fmt: skip
    # El rótulo viejo ya no está en ninguna celda de la columna A.
    assert _sin_rotulo_viejo(ws)
    # La vista en pantalla lleva la MISMA nota (tooltip del panel).
    vista = libro_balance_general(BalanceGeneralRequest.model_validate(p))
    hoja = next(h for h in vista["hojas"] if h["nombre"] == "balance")
    celdas = {(c["r"], c["c"]): c for c in hoja["celdas"]}
    k = next(c.row for c in ws["A"] if c.value == ANTES_COMISIONES_MXN)
    assert celdas[(k, 1)]["v"] == ANTES_COMISIONES_MXN
    assert celdas[(k, 2)]["n"] == NOTA_ANTES_COMISIONES_MXN
    assert celdas[(k + 2, 1)]["v"] == DESPUES_COMISIONES_MXN
    assert celdas[(k + 2, 2)]["n"] == NOTA_DESPUES_COMISIONES


def test_balance_general_en_dolares_rotulos_y_notas_exactos() -> None:
    p = pp.general(bandera=None)
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    ws = _balance(data)
    tst = _avion(p, "XA-TST")
    comisiones = (
        NOTA_COMISIONES_USD_GASTO_REAL
        if tst.get("comision_vendedor_solo_gasto_real") is True
        else NOTA_COMISIONES_USD
    )
    otros_ingresos = (
        NOTA_OTROS_INGRESOS_USD_POR_INGRESO
        if tst["otros_ingresos_avion"].get("usd_por_ingreso") is True
        else NOTA_OTROS_INGRESOS_USD
    )
    assert _cascada(ws, ANTES_COMISIONES_USD) == [
        (ANTES_COMISIONES_USD, NOTA_ANTES_COMISIONES_USD),
        ("(−) COMISIONES (banco + vendedor) USD", comisiones),
        (DESPUES_COMISIONES_USD, NOTA_DESPUES_COMISIONES),
        ("(−) COMBUSTIBLE DEL MES USD", None),
        ("(−) GASTOS INDIRECTOS USD", _nota_indirectos(
            tst["balance"]["otros_usd"], "USD", general=True
        )),
        ("(−) REFACCIONES (INVENTARIO) USD", None),
        ("(−) PERMISOS USD", None),
        ("(+) OTROS INGRESOS USD", otros_ingresos),
        ("UTILIDAD DESPUÉS DE GASTOS USD", None),
        ("(−) PENDIENTE DE PAGO (COBRANZA PENDIENTE) USD", None),
        ("UTILIDAD COBRADA USD", None),
    ]  # fmt: skip
    assert _sin_rotulo_viejo(ws)


def test_balance_por_avion_en_pesos_rotulos_y_notas_exactos() -> None:
    p = pp.individual()
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    ws = _balance(data)
    assert p["otros_ingresos_avion"] is not None
    assert _cascada(ws, DESPUES_COMISIONES_MXN) == [
        (DESPUES_COMISIONES_MXN, NOTA_DESPUES_COMISIONES),
        ("(−) COMBUSTIBLE DEL MES (MXN)", None),
        ("(−) GASTOS INDIRECTOS (MXN)", _nota_indirectos(
            p["balance"]["otros_mxn"], "MXN", general=False
        )),
        ("(−) REFACCIONES (INVENTARIO) (MXN)", None),
        ("(−) PERMISOS (MXN)", None),
        ("(+) OTROS INGRESOS (MXN)", NOTA_OTROS_INGRESOS_MXN),
        ("UTILIDAD DESPUÉS DE GASTOS (MXN)", None),
        ("(−) PENDIENTE DE PAGO (COBRANZA PENDIENTE) (MXN)", None),
        ("UTILIDAD COBRADA (MXN)", None),
    ]  # fmt: skip
    # El individual nunca pinta las dos previas de comisiones (ya restan en su
    # hoja de vuelos) ni el rótulo viejo.
    assert ANTES_COMISIONES_MXN not in _textos_a(ws)
    assert _sin_rotulo_viejo(ws)


def test_balance_por_avion_en_dolares_rotulos_y_notas_exactos() -> None:
    p = pp.individual(bandera=None)
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    ws = _balance(data)
    otros_ingresos = (
        NOTA_OTROS_INGRESOS_USD_POR_INGRESO
        if p["otros_ingresos_avion"].get("usd_por_ingreso") is True
        else NOTA_OTROS_INGRESOS_USD
    )
    assert _cascada(ws, DESPUES_COMISIONES_USD) == [
        (DESPUES_COMISIONES_USD, NOTA_DESPUES_COMISIONES),
        ("(−) COMBUSTIBLE DEL MES USD", None),
        ("(−) GASTOS INDIRECTOS USD", _nota_indirectos(
            p["balance"]["otros_usd"], "USD", general=False
        )),
        ("(−) REFACCIONES (INVENTARIO) USD", None),
        ("(−) PERMISOS USD", None),
        ("(+) OTROS INGRESOS USD", otros_ingresos),
        ("UTILIDAD DESPUÉS DE GASTOS USD", None),
        ("(−) PENDIENTE DE PAGO (COBRANZA PENDIENTE) USD", None),
        ("UTILIDAD COBRADA USD", None),
    ]  # fmt: skip
    assert ANTES_COMISIONES_USD not in _textos_a(ws)
    assert _sin_rotulo_viejo(ws)


def test_reserva_sin_tc_nota_corta() -> None:
    """La reserva de overhaul sin T.C. oficial del cierre: «—» con una nota de
    una idea (y nunca dice que la reserva resta)."""
    ws = _balance(render_balance_avion_xlsx(
        BalanceAvionRequest.model_validate(t_res.individual(sin_tc=True))
    ))
    fila = next(c.row for c in ws["A"] if c.value == "Lo que debe aportar el avión (MXN)")
    celda = ws.cell(row=fila, column=2)
    assert celda.value == "—"
    assert celda.comment.text == (
        "Sin cifra en pesos: la reserva se pasa a pesos con el T.C. oficial del cierre del "
        "periodo y no hay T.C. oficial de ese día."
    )
