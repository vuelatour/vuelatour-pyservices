"""COMISIONES a cargo del avión en el balance (6-oct-2026, API 0.0.65).

Pedido del cliente: «En el balance, cuando hay una comisión de un banco, en la
parte de total cobrado no refleja el monto real que entró a la cuenta» y «La
comisión del banco y vendedor se puede ir a la columna de (comisiones del
vendedor) pero cambiar el nombre a "comisiones" y poner notas el tipo de
comisión y si hay más de una. Y la comisión del vendedor también entra en el
balance del avión, para que el monto real total cobrado ya sea después de
cualquier comisión». Respuestas del 6-oct: la comisión del vendedor la absorbe
el avión, lo cobrado al cliente por ese concepto se queda como ingreso de
VuelaTour y la regla va desde septiembre 2026 (por fecha del vuelo).

El API calcula todo (fuente única `comisionesDelVuelo`, la misma del reparto a
socios); aquí solo se pinta:
- «COMISIÓN VENDEDOR MXN» → «COMISIONES MXN» (misma posición; en el Balance
  general, al final del bloque COSTO POR HORA) con `comisiones_mxn` y la nota
  `comisiones_detalle` («N conceptos» arriba cuando hay más de una);
- GANANCIA MXN = ROUND(REMANENTE − COMISIONES, 2) (mensual) y REMANENTE VENTA
  MENOS COMPRA = ROUND(VENTA − PROVEEDOR − COMISIONES, 2) (general) cuando la
  fila trae comisiones; sin ellas, la fórmula de siempre;
- cada COBRO n MXN pinta el NETO (`neto_mxn`) con «Bruto … · comisión … ·
  neto …» en su nota y COBRADO REAL = Σ netos (la misma fórmula SUM);
  COBRADO AVIÓN dice «antes de comisiones»;
- 'cobranza': el detalle con neto y bruto y una nota; RESUMEN: «COMISIONES
  MXN» y GANANCIA − COMISIONES cuando el avión las trae.

Los payloads salen del ESPEJO del API de `test_balance_formulas` con la regla
aplicada fila por fila (`_vigente`, con la forma EXACTA del JSON del API: fila
vigente sin comisiones con `comisiones_mxn` null y las partes en 0, fila
EXTERNOS con el neto y sin llaves de comisiones, totales de la flota si
ALGÚN libro las trae), así cada fórmula se compara contra el número que el
API mandaría. Un API ≤ 0.0.64 (sin los campos) da el libro de siempre salvo
los encabezados y sus textos (huellas de los demás tests).

Revisión 6-oct-2026: el nombre de la regla viene del API (`regla_comisiones`,
la vigencia es configurable; sin él, «septiembre 2026»), y 'otros
movimientos' y el bloque VUELATOUR (empresa) explican la línea «a cargo del
avión» y la parte de VuelaTour de la comisión bancaria cuando el periodo
trae vuelos de la regla.
"""

from __future__ import annotations

import copy
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    REGLA_COMISIONES_RESPALDO,
    BalanceAvionCobro,
    BalanceAvionRequest,
    BalanceAvionTotales,
    BalanceAvionVuelo,
    BalanceGeneralRequest,
    BalanceGeneralResumenFila,
    regla_comisiones_liberal,
)
from app.services import xlsx_formulas
from app.services.balance_avion_xlsx import (
    _COLS,
    _COLS_GENERAL,
    _DISP_GENERAL,
    _DISP_MENSUAL,
    _EMPRESA_REGLA_COMISIONES,
    _GRUPO_COSTO_HORA,
    _NOTA_BLOQUE_EMPRESA,
    _NOTA_COBRADO_REAL,
    _NOTA_COBRANZA_NETO,
    _NOTA_COMISIONES,
    _NOTA_EMPRESA_BASE,
    _NOTA_ENCABEZADO_COMISIONES,
    _NOTA_ENCABEZADO_REMANENTE,
    _NOTA_GANANCIA,
    _OM_REGLA_COMISIONES,
    _RESUMEN_COMISIONES,
    _bruto_neto,
    _cobros_a_4,
    _comisiones,
    _con_regla,
    _monto_detalle_cobranza,
    _neto_cobro,
    _nota_cobro,
    _nota_comisiones,
    _nota_resumen_cobros,
    _notas_parcialidades,
    _valor_columna,
    _valor_total,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from tests import test_balance_empresa_vuelatour as t_empresa
from tests import test_balance_formulas as tf
from tests._evaluador_formulas import verificar_libro

r2 = tf.r2

COMISIONES = "COMISIONES\nMXN"
GANANCIA = "GANANCIA\nMXN"
REMANENTE = "REMANENTE\nVENTA−COSTO MXN"
REMANENTE_CPH = "REMANENTE VENTA\nMENOS COMPRA (PESOS)"
COBRADO_REAL = "COBRADO REAL\nMXN (Σ depósitos)"
COBRADO_AVION = "COBRADO AVIÓN MXN\n(prorrateado, antes\nde comisiones) ****"

# ---------------------------------------------------------------------------
# Payloads: el espejo del API con la regla de septiembre 2026
# ---------------------------------------------------------------------------

# Lo que el API 0.0.65 arma en `cobrado_con` cuando el cobro trae comisión
# (ejemplo del contrato): «Bruto … · comisión banco 5 % … · neto … · cómo».
COBRADO_CON_ARMADO = (
    "Bruto $30,030.00 · comisión banco 5 % $1,501.50 · neto $28,528.50 · "
    "Transferencia → HSBC Pesos · Registró: Itzi"
)
DETALLE_401 = [
    "Comisión bancaria · Transferencia 5 % · $1,501.50 (parte del avión 100 %)",
    "Comisión vendedor (Pablo Canales) · $2,360.00 · provisión = cotizado + IVA",
]
DETALLE_407 = ["Comisión vendedor (Alex Saab) · $1,180.00 · provisión = cotizado + IVA"]


def _vigente(fila: dict, spec: dict, tc_prom: float | None) -> dict:
    """La fila como la manda el API 0.0.65 (paso 2 de buildPayload con la
    regla), con la MISMA forma del JSON real (revisión 6-oct-2026: el espejo
    mandaba 0.0 / None donde el API manda null / 0). `spec["regla"]`:
    - ausente o None ⇒ la fila de siempre: API ≤ 0.0.64 o vuelo ANTERIOR a
      la vigencia (el 0.0.65 no manda ni las llaves nuevas ni el neto);
    - {"externos": True} ⇒ fila del libro EXTERNOS en la vigencia: cada cobro
      gana su neto y COBRADO REAL = Σ netos, SIN llaves de comisiones (no hay
      avión que las absorba) y GANANCIA = REMANENTE;
    - otro dict ⇒ fila de AVIÓN en la vigencia: neto por cobro; las dos
      partes como NÚMERO (0 sin comisión): la del avión de la comisión
      bancaria (comisión × venta avión ÷ total cotización, el factor del
      prorrateo) y la provisión del vendedor; `comisiones_mxn` = su suma o
      null si da 0; `comisiones_detalle` lista ([] sin conceptos); GANANCIA =
      REMANENTE − COMISIONES."""
    regla = spec.get("regla")
    if regla is None:
        return fila
    cobros = [
        {**c, "neto_mxn": r2(c["monto_mxn"] - (c.get("comision_mxn") or 0))} for c in fila["cobros"]
    ]
    fila = {
        **fila,
        "cobros": cobros,
        "cobrado_real_mxn": r2(tf._suma(c["neto_mxn"] for c in cobros)),
    }
    if regla.get("externos"):
        return {**fila, "es_externo": True}
    comision_cobros = tf._suma(c.get("comision_mxn") for c in cobros)
    factor = (
        fila["total_mxn"] / fila["total_cotizacion_mxn"]
        if fila["total_mxn"] and fila["total_cotizacion_mxn"]
        else 1
    )
    banco = r2(comision_cobros * factor)
    vendedor = r2(regla.get("vendedor", 0.0))
    total = r2(banco + vendedor)
    comisiones = total if total != 0 else None
    ganancia = r2(fila["remanente_mxn"] - (comisiones or 0))
    tc_g = fila["tc_costos"] if fila["tc_costos"] is not None else tc_prom
    return {
        **fila,
        "comisiones_mxn": comisiones,
        "comisiones_detalle": regla.get("detalle", []),
        "comision_banco_avion_mxn": banco,
        "comision_vendedor_prov_mxn": vendedor,
        "ganancia_mxn": ganancia,
        "ganancia_usd": r2(ganancia / tc_g) if tc_g is not None else None,
    }


def _specs_tst() -> list[dict]:
    """XA-TST de `test_balance_formulas` con la regla: #401 comisión bancaria
    (5 % del 2.º cobro, texto ya armado por el API) + vendedor; #402
    multi-avión con la comisión bancaria de su 2.º cobro (parte del avión
    82 %, sin detalle de vendedor); #403 vigente SIN comisiones (la forma
    real: `comisiones_mxn` null, partes en 0, detalle []); #404 y #406
    anteriores a la vigencia (sin las llaves); #405 comisión bancaria SIN
    detalle (la nota se compone con las partes)."""
    specs = copy.deepcopy(tf._SPECS_TST)
    por_clave = {s["clave"]: s for s in specs}
    s401 = por_clave["#401 · Cliente Uno"]
    s401["cobros"][1]["comision_mxn"] = 1501.5
    s401["cobros"][1]["cobrado_con"] = COBRADO_CON_ARMADO
    s401["regla"] = {"vendedor": 2360.0, "detalle": DETALLE_401}
    por_clave["#402 · Cliente Dos"]["regla"] = {
        "detalle": ["Comisión bancaria · Transferencia · $123.39 (parte del avión 82.12 %)"]
    }
    por_clave["#403 · Traslado"]["regla"] = {}
    por_clave["#404 · Cliente Tres"]["regla"] = None
    s405 = por_clave["#405 · Cliente Cuatro"]
    s405["cobros"][0]["comision_mxn"] = 500.0
    s405["regla"] = {}
    por_clave["#406 · Interno"]["regla"] = None
    return specs


def _specs_dos() -> list[dict]:
    """XB-DOS: la otra mitad del #402 con comisión en los cobros 4 y 5 (la
    4.ª celda agrega los NETOS) y el #407 con SOLO la comisión del vendedor."""
    specs = copy.deepcopy(tf._SPECS_DOS)
    s402, s407 = specs
    s402["cobros"][3]["comision_mxn"] = 100.0
    s402["cobros"][4]["comision_mxn"] = 125.01
    s402["regla"] = {}
    s407["regla"] = {"vendedor": 1180.0, "detalle": DETALLE_407}
    return specs


def _libro_tst() -> dict:
    return tf._libro_api(
        "XA-TST", 25.0, _specs_tst(), tf._HOJAS_TST, tf._SOCIOS, ajusta_fila=_vigente
    )


def _libro_dos() -> dict:
    return tf._libro_api(
        "XB-DOS", None, _specs_dos(), tf._HOJAS_DOS, [("Socio C", 100.0)], ajusta_fila=_vigente
    )


def _individual() -> BalanceAvionRequest:
    return BalanceAvionRequest.model_validate(_libro_tst())


def _general(variante: str = "mensual") -> BalanceGeneralRequest:
    payload = tf._general_payload([_libro_tst(), _libro_dos()])
    return BalanceGeneralRequest.model_validate({**payload, "variante": variante})


# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------

MAESTRA_TST = "reporte horas XA-TST"
MAESTRA = "reporte horas FLOTA"
# Las dos celdas que el espejo del API deja como VALOR a propósito (redondeo
# intermedio de COSTO TOTAL del #405 y TOTAL USD de 'Gastos Indirectos',
# `test_balance_formulas`): las comisiones no agregan ninguna.
_DEGRADADAS_INDIVIDUAL = [(MAESTRA_TST, "U7"), ("Gastos Indirectos", "C5")]


def _libros(data: bytes):
    """(con fórmulas, con caché)."""
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _col(ws, encabezado: str) -> int:
    return next(c.column for c in ws[2] if c.value == encabezado)


def _celda(ws, fila: int, encabezado: str):
    return ws.cell(row=fila, column=_col(ws, encabezado))


def _ref(ws, fila: int, encabezado: str) -> str:
    return f"{get_column_letter(_col(ws, encabezado))}{fila}"


def _tot(ws) -> int:
    return next(c.row for c in ws["A"] if c.value == "TOTALES")


def _pie(ws) -> list[str]:
    return [c.value for c in ws["A"] if isinstance(c.value, str)]


@pytest.fixture
def registros(monkeypatch) -> list[xlsx_formulas.RegistroFormulas]:
    """Registro de fórmulas de cada libro que se cierra (cuáles se quedaron
    como fórmula y cuáles volvieron a valor)."""
    capturados: list[xlsx_formulas.RegistroFormulas] = []
    original = xlsx_formulas.finalizar

    def espia(wb):
        cache = original(wb)
        capturados.append(xlsx_formulas.registro(wb))
        return cache

    monkeypatch.setattr(xlsx_formulas, "finalizar", espia)
    return capturados


# ---------------------------------------------------------------------------
# 1) Esquema: campos ADITIVOS y liberales.
# ---------------------------------------------------------------------------


def test_campos_nuevos_con_default() -> None:
    v = BalanceAvionVuelo()
    assert (v.comisiones_mxn, v.comision_banco_avion_mxn, v.comision_vendedor_prov_mxn) == (
        None,
        None,
        None,
    )
    assert v.comisiones_detalle == []
    assert BalanceAvionTotales().comisiones_mxn is None
    assert BalanceAvionCobro().neto_mxn is None


@pytest.mark.parametrize(
    "entrada,esperado",
    [
        (None, []),
        ("Comisión bancaria", []),
        ({"x": 1}, []),
        ([" Comisión bancaria · $1.00 ", "", "  ", None, 3, "Comisión vendedor"],
         ["Comisión bancaria · $1.00", "Comisión vendedor"]),
        (("una",), ["una"]),
    ],
)  # fmt: skip
def test_detalle_liberal_nunca_es_un_422(entrada, esperado) -> None:
    v = BalanceAvionVuelo.model_validate({"comisiones_detalle": entrada})
    assert v.comisiones_detalle == esperado


# ---------------------------------------------------------------------------
# 2) Helpers puros.
# ---------------------------------------------------------------------------

COBRO_CON_COMISION = {
    "fecha": "2026-09-15",
    "monto_mxn": 20400.0,
    "comision_mxn": 1020.0,
    "neto_mxn": 19380.0,
}
# El ejemplo del contrato: lo que el API arma en `cobrado_con`.
ARMADO_CONTRATO = (
    "Bruto $20,400.00 · comisión banco 5 % $1,020.00 · neto $19,380.00 · "
    "Transferencia → Scotiabank Pesos · Registró: Itzi"
)


def test_neto_y_desglose_bruto_neto() -> None:
    c = BalanceAvionCobro(**COBRO_CON_COMISION)
    assert _neto_cobro(c) == 19380.0
    assert _bruto_neto(c) == "Bruto $20,400.00 · comisión banco $1,020.00 · neto $19,380.00"
    # API ≤ 0.0.64 (sin neto): el monto de siempre, sin desglose.
    viejo = BalanceAvionCobro(monto_mxn=20400.0, comision_mxn=1020.0)
    assert (_neto_cobro(viejo), _bruto_neto(viejo)) == (20400.0, None)
    # Vuelo vigente SIN comisión (neto = bruto): sin desglose.
    assert _bruto_neto(BalanceAvionCobro(monto_mxn=100.0, neto_mxn=100.0)) is None
    # Neto distinto sin el monto de la comisión: bruto y neto, nada inventado.
    assert _bruto_neto(BalanceAvionCobro(monto_mxn=100.0, neto_mxn=95.0)) == (
        "Bruto $100.00 · neto $95.00"
    )


def test_nota_del_cobro_con_neto() -> None:
    # El API ya armó bruto · comisión · neto en `cobrado_con`: va tal cual
    # tras la fecha (el monto no se repite).
    c = BalanceAvionCobro(**COBRO_CON_COMISION, cobrado_con=ARMADO_CONTRATO)
    assert _nota_cobro(c) == f"15/09/2026 · {ARMADO_CONTRATO}"
    # `cobrado_con` sin el desglose: aquí se antepone.
    c = BalanceAvionCobro(**COBRO_CON_COMISION, cobrado_con="Transferencia → Scotiabank Pesos")
    assert _nota_cobro(c) == (
        "15/09/2026 · Bruto $20,400.00 · comisión banco $1,020.00 · neto $19,380.00 · "
        "Transferencia → Scotiabank Pesos"
    )
    # Sin método, cuenta ni registro: haber entrado neto ya es algo que contar.
    assert _nota_cobro(BalanceAvionCobro(**COBRO_CON_COMISION)) == (
        "15/09/2026 · Bruto $20,400.00 · comisión banco $1,020.00 · neto $19,380.00"
    )
    # Neto = bruto: la nota de siempre (y sin nada que contar, sin nota).
    sin = {"fecha": "2026-09-15", "monto_mxn": 12522.2, "neto_mxn": 12522.2}
    assert _nota_cobro(BalanceAvionCobro(**sin, cobrado_con="Efectivo · Itzi")) == (
        "15/09/2026 · $12,522.20 · Efectivo · Itzi"
    )
    assert _nota_cobro(BalanceAvionCobro(**sin)) is None
    # Otra redacción del API («neto: $…»): tampoco se repite.
    otra = "Depósito neto: $19,380.00 (comisión 5 %) · Transferencia"
    c = BalanceAvionCobro(**COBRO_CON_COMISION, cobrado_con=otra)
    assert _nota_cobro(c) == f"15/09/2026 · {otra}"
    # «Neto» como apodo de quien registró NO es el desglose: se antepone.
    c = BalanceAvionCobro(**COBRO_CON_COMISION, cobrado_con="Efectivo · Registró: Neto")
    assert _nota_cobro(c) == (
        "15/09/2026 · Bruto $20,400.00 · comisión banco $1,020.00 · neto $19,380.00 · "
        "Efectivo · Registró: Neto"
    )
    # Sin neto (API previo) un «Bruto» en el texto no cambia nada.
    previo = BalanceAvionCobro(fecha="2026-09-15", monto_mxn=100.0, cobrado_con="Bruto · x")
    assert _nota_cobro(previo) == "15/09/2026 · $100.00 · Bruto · x"


def test_resumen_y_cuarta_parcialidad_con_netos() -> None:
    cobros = [
        BalanceAvionCobro(
            fecha=f"2026-09-0{d}", monto_mxn=m, comision_mxn=k, neto_mxn=r2(m - (k or 0))
        )
        for d, m, k in (
            (1, 1000.0, None),
            (2, 2000.0, None),
            (3, 3000.0, None),
            (4, 4000.5, 100.0),
            (5, 5000.25, 125.01),
        )
    ]
    a4 = _cobros_a_4(cobros)
    assert (a4[3].fecha, a4[3].monto_mxn, a4[3].neto_mxn) == ("05/09/2026 (+2)", 9000.75, 8775.74)
    assert _neto_cobro(a4[3]) == 8775.74
    l4 = "04/09/2026 · Bruto $4,000.50 · comisión banco $100.00 · neto $3,900.50"
    l5 = "05/09/2026 · Bruto $5,000.25 · comisión banco $125.01 · neto $4,875.24"
    assert _notas_parcialidades(cobros) == [None, None, None, f"{l4}\n{l5}"]
    assert _nota_resumen_cobros(cobros) == "\n".join(
        ["01/09/2026 · $1,000.00", "02/09/2026 · $2,000.00", "03/09/2026 · $3,000.00", l4, l5]
    )
    # API previo: la 4.ª no trae neto.
    viejos = [BalanceAvionCobro(fecha="2026-09-01", monto_mxn=10.0) for _ in range(5)]
    assert _cobros_a_4(viejos)[3].neto_mxn is None


def test_nota_de_la_celda_comisiones() -> None:
    assert _nota_comisiones(BalanceAvionVuelo(comisiones_detalle=DETALLE_401)) == "\n".join(
        ["2 conceptos", *DETALLE_401]
    )
    assert _nota_comisiones(BalanceAvionVuelo(comisiones_detalle=DETALLE_407)) == DETALLE_407[0]
    # Si el API ya trae el encabezado, no se repite.
    con = ["Comisiones del vuelo · 2 conceptos", *DETALLE_401]
    assert _nota_comisiones(BalanceAvionVuelo(comisiones_detalle=con)) == "\n".join(con)
    # Sin detalle: una línea por parte que mande el API.
    # La provisión sin «+ IVA» (revisión 6-oct-2026): no toda cotización lo
    # cobra y el detalle del API sí lo distingue («cotizado (sin IVA)»).
    partes = BalanceAvionVuelo(comision_banco_avion_mxn=123.39, comision_vendedor_prov_mxn=2360.0)
    assert _nota_comisiones(partes) == (
        "2 conceptos\nComisión bancaria (parte del avión) · $123.39\n"
        "Comisión del vendedor (provisión) · $2,360.00"
    )
    assert _nota_comisiones(BalanceAvionVuelo(comision_vendedor_prov_mxn=1180.0)) == (
        "Comisión del vendedor (provisión) · $1,180.00"
    )
    # Sin nada que decir (0, null o vacío): sin nota.
    assert (
        _nota_comisiones(BalanceAvionVuelo(comisiones_mxn=0.0, comision_banco_avion_mxn=0)) is None
    )
    assert _nota_comisiones(BalanceAvionVuelo()) is None


def test_columna_comisiones_con_api_previo_y_nuevo() -> None:
    assert _valor_columna("comisiones_mxn", BalanceAvionVuelo(comisiones_mxn=3861.5)) == 3861.5
    # API ≤ 0.0.64: la columna se llenaba con comision_vendedor_mxn (null en
    # las filas, 0 en TOTALES desde el 28-ago-2026): el libro sale igual.
    assert _valor_columna("comisiones_mxn", BalanceAvionVuelo()) is None
    previo = BalanceAvionTotales(comision_vendedor_mxn=0.0)
    assert _valor_total(previo, "comisiones_mxn") == 0.0
    nuevo = BalanceAvionTotales(comisiones_mxn=5849.67, comision_vendedor_mxn=0.0)
    assert _valor_total(nuevo, "comisiones_mxn") == 5849.67 == _comisiones(nuevo)
    assert _valor_total(BalanceAvionTotales(ganancia_mxn=1.5), "ganancia_mxn") == 1.5
    # REMANENTE VENTA MENOS COMPRA del Balance general = GANANCIA del API.
    v = BalanceAvionVuelo(remanente_mxn=100.0, ganancia_mxn=80.0)
    assert _valor_columna("cph_remanente_mxn", v) == 80.0


def test_monto_del_detalle_de_cobranza() -> None:
    assert _monto_detalle_cobranza(BalanceAvionCobro(**COBRO_CON_COMISION)) == (
        "$19,380.00 neto (bruto $20,400.00)"
    )
    assert _monto_detalle_cobranza(BalanceAvionCobro(monto_mxn=12522.2)) == "$12,522.20"
    assert _monto_detalle_cobranza(BalanceAvionCobro(monto_mxn=50.0, neto_mxn=50.0)) == "$50.00"
    assert _monto_detalle_cobranza(BalanceAvionCobro()) is None


# ---------------------------------------------------------------------------
# 3) Juegos de columnas: COMISIONES en la misma posición y al final del bloque.
# ---------------------------------------------------------------------------


def test_comisiones_en_los_dos_juegos() -> None:
    mensual = [c[1] for c in _COLS]
    assert mensual[28] == COMISIONES == _COLS[_DISP_MENSUAL.comision_col - 1][1]
    assert _DISP_MENSUAL.letra["comisiones_mxn"] == "AC"  # misma posición de siempre
    assert "COMISIÓN\nVENDEDOR MXN" not in mensual
    assert COBRADO_AVION in mensual
    # Balance general: última columna del bloque COSTO POR HORA, después de
    # REMANENTE (las 12 del cliente siguen juntas y en su orden).
    general = [c[1] for c in _COLS_GENERAL]
    i = general.index(COMISIONES)
    assert general[i - 1] == REMANENTE_CPH and general[i + 1] == "STATUS"
    assert _COLS_GENERAL[i][0] == _GRUPO_COSTO_HORA
    assert _DISP_GENERAL.comision_col == i + 1
    assert _DISP_GENERAL.total_map["comisiones_mxn"] == "comisiones_mxn"
    assert _DISP_GENERAL.total_map["cph_remanente_mxn"] == "ganancia_mxn"


# ---------------------------------------------------------------------------
# 4) Libro individual (y Balance mensual) con la regla vigente.
# ---------------------------------------------------------------------------


def test_individual_comisiones_y_ganancia(registros) -> None:
    req = _individual()
    data = render_balance_avion_xlsx(req)
    assert registros[-1].degradadas() == _DEGRADADAS_INDIVIDUAL
    assert verificar_libro(data)
    wb, wv = _libros(data)
    ws, wsv = wb[MAESTRA_TST], wv[MAESTRA_TST]
    # Encabezados (con su nota) y la nota al pie.
    assert ws["AC2"].value == COMISIONES
    assert ws["AC2"].comment.text == _NOTA_ENCABEZADO_COMISIONES
    assert _celda(ws, 2, COBRADO_AVION).column_letter == "AR"
    assert _NOTA_COMISIONES + "(el pago apareado como PROVISIÓN a la fecha del vuelo)." + (
        _NOTA_GANANCIA
    ) in _pie(ws)
    for k, v in enumerate(req.vuelos):
        f = 3 + k
        comision = _celda(ws, f, COMISIONES)
        assert comision.value == v.comisiones_mxn, v.clave  # valor del API
        nota = _nota_comisiones(v)
        assert (comision.comment.text if comision.comment else None) == nota, v.clave
        # GANANCIA = ROUND(REMANENTE − COMISIONES, 2); sin comisiones, = REMANENTE.
        ganancia = _celda(ws, f, GANANCIA).value
        if v.comisiones_mxn is None:
            assert ganancia == f"={_ref(ws, f, REMANENTE)}", v.clave
        else:
            assert ganancia == f"=ROUND({_ref(ws, f, REMANENTE)}-AC{f},2)", v.clave
        assert _celda(wsv, f, GANANCIA).value == v.ganancia_mxn, v.clave
        assert _celda(wsv, f, "GANANCIA\nUSD").value == v.ganancia_usd, v.clave
    # Las notas por fila: dos conceptos (#401), uno (#402), ninguna sin
    # comisiones (#403: null del API, celda vacía) y compuesta con las partes
    # del API (#405, sin detalle).
    assert ws["AC3"].comment.text.startswith("2 conceptos\nComisión bancaria")
    assert ws["AC5"].value is None and ws["AC5"].comment is None
    assert ws["AC7"].comment.text == "Comisión bancaria (parte del avión) · $500.00"
    assert ws["AC6"].value is None and ws["AC6"].comment is None  # #404 anterior
    # La ganancia del #401 ya descuenta las dos comisiones.
    v401 = req.vuelos[0]
    assert v401.comisiones_mxn == r2(1501.5 + 2360.0)
    assert v401.ganancia_mxn == r2(v401.remanente_mxn - 3861.5)
    # TOTALES: Σ de la columna.
    tot = _tot(ws)
    assert ws[f"AC{tot}"].value == f"=ROUND(SUM(AC3:AC{tot - 1}),2)"
    assert wsv[f"AC{tot}"].value == req.totales.comisiones_mxn == 4484.89
    assert wsv[f"AD{tot}"].value == req.totales.ganancia_mxn


def test_individual_cobros_netos_y_cobrado_avion_antes_de_comisiones() -> None:
    req = _individual()
    wb, wv = _libros(render_balance_avion_xlsx(req))
    ws, wsv = wb[MAESTRA_TST], wv[MAESTRA_TST]
    # #401: COBRO 2 = lo que entró (30,030.00 − 1,501.50) con la nota armada
    # por el API; COBRADO REAL = Σ netos; COBRADO AVIÓN sigue en el bruto.
    assert ws["AJ3"].value == 30000.0 and ws["AL3"].value == 28528.5
    assert ws["AL3"].comment.text == f"04/09/2026 · {COBRADO_CON_ARMADO}"
    assert ws["AQ3"].value == "=ROUND(SUM(AJ3,AL3,AN3,AP3),2)"
    assert wsv["AQ3"].value == req.vuelos[0].cobrado_real_mxn == 58528.5
    assert wsv["AR3"].value == req.vuelos[0].cobrado_mxn == 60030.0
    assert wsv["AS3"].value == 0  # pagado: por cobrar $0 (antes de comisiones)
    # #402: la comisión del 2.º cobro (150.25) sale del depósito y la nota
    # la compone con bruto, comisión y neto.
    assert ws["AL4"].value == 5850.25
    assert ws["AL4"].comment.text == (
        "06/09/2026 · Bruto $6,000.50 · comisión banco $150.25 · neto $5,850.25 · "
        "TRANSFERENCIA → HSBC Pesos"
    )
    assert wsv["AQ4"].value == r2(5000 + 5850.25 + 7000.25 + 9000.1)
    # STATUS junta las líneas de todos los cobros.
    assert ws["AH3"].comment.text == (
        f"01/09/2026 · $30,000.00 · TRANSFERENCIA → HSBC Pesos\n04/09/2026 · {COBRADO_CON_ARMADO}"
    )


def test_cascada_del_balance_ya_va_despues_de_comisiones() -> None:
    """La hoja 'balance' cita la GANANCIA USD de TOTALES: la del API, ya
    después de comisiones (la utilidad de los socios baja)."""
    req = _individual()
    wb, wv = _libros(render_balance_avion_xlsx(req))
    assert wb["balance"]["B4"].value == f"='{MAESTRA_TST}'!$AE$9"
    antes = wv["balance"]["B4"].value
    assert antes == req.balance.utilidad_antes_usd == req.totales.ganancia_usd
    assert antes < tf._individual_tst().balance.utilidad_antes_usd


def test_cobranza_neto_con_su_nota() -> None:
    req = _individual()
    wb, wv = _libros(render_balance_avion_xlsx(req))
    ws = wb["cobranza"]
    assert ws["K8"].value == 58528.5  # COBRADO REAL = Σ netos
    assert ws["L8"].value == 1501.5  # la comisión sigue en su columna
    assert ws["M8"].value == (
        "01/09/2026 · $30,000.00 · TRANSFERENCIA · HSBC Pesos\n"
        "04/09/2026 · $28,528.50 neto (bruto $30,030.00) · TRANSFERENCIA · HSBC Pesos · "
        "comisión $1,501.50"
    )
    tot = _tot(ws)
    assert wv["cobranza"][f"K{tot}"].value == req.totales.cobrado_real_mxn
    assert _NOTA_COBRANZA_NETO in _pie(ws)
    # Sin netos (API previo o periodo anterior) la nota no aparece.
    viejo = load_workbook(BytesIO(render_balance_avion_xlsx(tf._individual_tst())))["cobranza"]
    assert _NOTA_COBRANZA_NETO not in _pie(viejo)


# ---------------------------------------------------------------------------
# 5) Balance general: «reporte horas FLOTA», RESUMEN y la variante general.
# ---------------------------------------------------------------------------


def test_general_mensual_resumen_ganancia_menos_comisiones(registros) -> None:
    req = _general()
    data = render_balance_general_xlsx(req)
    # Solo la degradada de siempre (COSTO TOTAL del #405).
    assert registros[-1].degradadas() == [(MAESTRA, "U9")]
    assert verificar_libro(data)
    wb, wv = _libros(data)
    rs, rv = wb["RESUMEN flota"], wv["RESUMEN flota"]
    assert rs["H3"].value == COMISIONES
    for fila, r in ((4, req.resumen[0]), (5, req.resumen[1])):
        assert rs[f"I{fila}"].value == f"=ROUND(E{fila}-F{fila}-G{fila}-H{fila},2)"
        assert rv[f"H{fila}"].value == r.comisiones_mxn
        assert rv[f"I{fila}"].value == r.ganancia_mxn
    assert rv["H6"].value == req.resumen_totales.comisiones_mxn == 5849.67
    assert any(x.startswith("COMISIONES = las de la columna COMISIONES") for x in _pie(rs))
    assert any("COMBUSTIBLE − COMISIONES" in x for x in _pie(rs))
    # La hoja de vuelos de la flota: mismas reglas que el libro individual.
    ws, wsv = wb[MAESTRA], wv[MAESTRA]
    cons = req.consolidado
    for k, v in enumerate(cons.vuelos):
        f = 3 + k
        assert _celda(ws, f, COMISIONES).value == v.comisiones_mxn
        assert _celda(wsv, f, GANANCIA).value == v.ganancia_mxn
    assert wsv[f"AC{_tot(ws)}"].value == cons.totales.comisiones_mxn == 5849.67


def test_resumen_sin_comisiones_conserva_su_formula() -> None:
    """Un avión con COMISIONES 0 o null (API previo, periodo anterior): la
    GANANCIA del RESUMEN es la fórmula de siempre."""
    req = tf._general()
    assert [r.comisiones_mxn for r in req.resumen] == [0.0, 0.0]
    rs = load_workbook(BytesIO(render_balance_general_xlsx(req)))["RESUMEN flota"]
    assert rs["I4"].value == "=ROUND(E4-F4-G4,2)"
    assert rs["I5"].value == "=ROUND(E5-F5-G5,2)"
    sin = BalanceGeneralResumenFila(matricula="X", venta_mxn=1.0)
    assert sin.comisiones_mxn is None


def test_variante_general_remanente_despues_de_comisiones(registros) -> None:
    req = _general("general")
    data = render_balance_general_xlsx(req)
    assert registros[-1].degradadas() == []
    assert verificar_libro(data)
    wb, wv = _libros(data)
    ws, wsv = wb[MAESTRA], wv[MAESTRA]
    assert ws.max_column == len(_COLS_GENERAL) == 45
    com, rem = _col(ws, COMISIONES), _col(ws, REMANENTE_CPH)
    assert com == rem + 1 == _DISP_GENERAL.comision_col
    # El título del bloque cubre las 13 columnas.
    inicio = get_column_letter(_col(ws, "TOTAL COBRADO\nS/IVA (PESOS)"))
    assert f"{inicio}1:{get_column_letter(com)}1" in {str(r) for r in ws.merged_cells.ranges}
    assert ws.cell(row=2, column=com).comment.text == _NOTA_ENCABEZADO_COMISIONES
    assert ws.cell(row=2, column=rem).comment.text == _NOTA_ENCABEZADO_REMANENTE
    k = get_column_letter(_col(ws, "VENTA AVIÓN\nMXN"))
    y = get_column_letter(_col(ws, "TOTAL PARA\nPROVEEDOR (PESOS)"))
    ae = get_column_letter(com)
    for i, v in enumerate(req.consolidado.vuelos):
        f = 3 + i
        formula = ws.cell(row=f, column=rem).value
        if v.comisiones_mxn is None:
            assert formula == f"=ROUND({k}{f}-{y}{f},2)", v.clave
        else:
            assert formula == f"=ROUND({k}{f}-{y}{f}-{ae}{f},2)", v.clave
        # El número: la GANANCIA del API (remanente − comisiones).
        assert wsv.cell(row=f, column=rem).value == v.ganancia_mxn, v.clave
        assert wsv.cell(row=f, column=com).value == v.comisiones_mxn, v.clave
        nota = ws.cell(row=f, column=com).comment
        assert (nota.text if nota else None) == _nota_comisiones(v), v.clave
    tot = _tot(ws)
    t = req.consolidado.totales
    letra = get_column_letter(rem)
    assert ws.cell(row=tot, column=rem).value == f"=ROUND(SUM({letra}3:{letra}{tot - 1}),2)"
    assert wsv.cell(row=tot, column=rem).value == t.ganancia_mxn
    assert wsv.cell(row=tot, column=com).value == t.comisiones_mxn
    # Notas al pie: COMISIONES y REMANENTE después de comisiones.
    pie = "\n".join(_pie(ws))
    assert "REMANENTE VENTA MENOS COMPRA = VENTA AVIÓN MXN − TOTAL PARA PROVEEDOR (PESOS) − " in pie
    assert _NOTA_COMISIONES in pie and "GANANCIA de la fila" not in pie


# ---------------------------------------------------------------------------
# 5b) Formas REALES del JSON del API 0.0.65 (revisión 6-oct-2026: el espejo no
#     las reproducía y ningún test las cuidaba).
# ---------------------------------------------------------------------------

# Vuelo cubierto por un operador EXTERNO (libro EXTERNOS del mes): su cobro
# entra NETO de la comisión del banco, pero ningún avión la absorbe — el API
# no manda llaves de comisiones (caso #504 de la sonda: 18,000 − 900).
_SPECS_EXTERNOS = [
    {"clave": "#504 · Externo", "fecha": "2026-09-18", "D": 2.0, "E": 500, "G": 0,
     "I": 1000.0, "J": 0.0, "K": 18.0, "O": None, "op": 9000.0, "z": 18.0,
     "cobros": [tf._cobro("2026-09-18", 18000.0, 900.0)], "cobrado": 18000.0,
     "regla": {"externos": True}},
]  # fmt: skip


def _libro_externos() -> dict:
    return tf._libro_api(
        "EXTERNOS", None, copy.deepcopy(_SPECS_EXTERNOS), {}, [], ajusta_fila=_vigente
    )


def _general_con_externos(variante: str = "mensual") -> BalanceGeneralRequest:
    """TST + DOS + EXTERNOS: el consolidado junta los tres libros; EXTERNOS no
    tiene socios ni bloque en 'balance' (no viaja en `aviones`)."""
    libros = [_libro_tst(), _libro_dos(), _libro_externos()]
    payload = tf._general_payload(libros)
    payload["aviones"] = libros[:2]
    return BalanceGeneralRequest.model_validate({**payload, "variante": variante})


def _fila_de(ws, clave: str) -> int:
    return next(c.row for c in ws["A"] if c.value == clave)


def test_fila_vigente_sin_comisiones_con_la_forma_del_api() -> None:
    """#403: vuelo de la regla SIN comisiones. El API manda `comisiones_mxn`
    null, las dos partes en 0 y el detalle [] (no 0.0 / null): COMISIONES
    vacía y sin nota, GANANCIA = REMANENTE (`=AA`) y, en la variante
    general, REMANENTE VENTA MENOS COMPRA = VENTA − PROVEEDOR."""
    fila = next(f for f in _libro_tst()["vuelos"] if f["clave"] == "#403 · Traslado")
    assert fila["comisiones_mxn"] is None
    assert (fila["comision_banco_avion_mxn"], fila["comision_vendedor_prov_mxn"]) == (0, 0)
    assert fila["comisiones_detalle"] == []
    assert fila["ganancia_mxn"] == fila["remanente_mxn"]
    ws = load_workbook(BytesIO(render_balance_avion_xlsx(_individual())))[MAESTRA_TST]
    f = _fila_de(ws, "#403 · Traslado")
    assert ws[f"AD{f}"].value == f"=AA{f}"
    assert ws[f"AC{f}"].value is None and ws[f"AC{f}"].comment is None
    gen = load_workbook(BytesIO(render_balance_general_xlsx(_general("general"))))[MAESTRA]
    f = _fila_de(gen, "#403 · Traslado")
    k = get_column_letter(_col(gen, "VENTA AVIÓN\nMXN"))
    y = get_column_letter(_col(gen, "TOTAL PARA\nPROVEEDOR (PESOS)"))
    assert _celda(gen, f, REMANENTE_CPH).value == f"=ROUND({k}{f}-{y}{f},2)"
    com = _celda(gen, f, COMISIONES)
    assert com.value is None and com.comment is None


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_fila_externos_cobro_neto_sin_llaves_de_comisiones(variante, registros) -> None:
    """#504 (libro EXTERNOS): el API manda el NETO del cobro y COBRADO REAL =
    Σ netos, pero NINGUNA llave de comisiones. COBRO 1 = neto con su nota,
    COBRADO REAL con su fórmula SUM, COMISIONES vacía y GANANCIA = REMANENTE.
    Los totales de la flota llevan COMISIONES porque ALGÚN libro las trae
    (`libros.some(...)` del API; EXTERNOS no)."""
    ext = _libro_externos()
    fila = ext["vuelos"][0]
    assert {"comisiones_mxn", "comisiones_detalle", "comision_banco_avion_mxn"}.isdisjoint(fila)
    assert fila["cobros"][0]["neto_mxn"] == 17100.0 == fila["cobrado_real_mxn"]
    assert "comisiones_mxn" not in ext["totales"]
    req = _general_con_externos(variante)
    data = render_balance_general_xlsx(req)
    assert verificar_libro(data)
    wb, wv = _libros(data)
    ws, wsv = wb[MAESTRA], wv[MAESTRA]
    # Solo la degradada de siempre (COSTO TOTAL del #405, mensual).
    esperadas = [(MAESTRA, f"U{_fila_de(ws, '#405 · Cliente Cuatro')}")]
    assert registros[-1].degradadas() == (esperadas if variante == "mensual" else [])
    disp = _DISP_GENERAL if variante == "general" else _DISP_MENSUAL
    f = _fila_de(ws, "#504 · Externo")
    cobro1 = ws[f"{disp.cobro_mxn_letras[0]}{f}"]
    assert cobro1.value == 17100.0
    assert cobro1.comment.text == (
        "18/09/2026 · Bruto $18,000.00 · comisión banco $900.00 · neto $17,100.00 · "
        "TRANSFERENCIA → HSBC Pesos"
    )
    cobros = ",".join(f"{letra}{f}" for letra in disp.cobro_mxn_letras)
    assert _celda(ws, f, COBRADO_REAL).value == f"=ROUND(SUM({cobros}),2)"
    assert _celda(wsv, f, COBRADO_REAL).value == 17100.0
    com = _celda(ws, f, COMISIONES)
    assert com.value is None and com.comment is None
    if variante == "mensual":
        assert _celda(ws, f, GANANCIA).value == f"={_ref(ws, f, REMANENTE)}"
        assert _celda(wsv, f, GANANCIA).value == 9000.0
    else:
        k = get_column_letter(_col(ws, "VENTA AVIÓN\nMXN"))
        y = get_column_letter(_col(ws, "TOTAL PARA\nPROVEEDOR (PESOS)"))
        assert _celda(ws, f, REMANENTE_CPH).value == f"=ROUND({k}{f}-{y}{f},2)"
        assert _celda(wsv, f, REMANENTE_CPH).value == 9000.0
    t = req.consolidado.totales
    assert t.comisiones_mxn == r2(
        _libro_tst()["totales"]["comisiones_mxn"] + _libro_dos()["totales"]["comisiones_mxn"]
    )
    assert wsv.cell(row=_tot(ws), column=_col(ws, COMISIONES)).value == t.comisiones_mxn
    # 'cobranza': neto con el bruto al lado y la nota de los netos, que ya no
    # dice que el avión absorbe todas las comisiones (EXTERNOS no).
    cob = wb["cobranza"]
    assert any(
        isinstance(c.value, str) and "$17,100.00 neto (bruto $18,000.00)" in c.value
        for fila_c in cob.iter_rows()
        for c in fila_c
    )
    assert _NOTA_COBRANZA_NETO in _pie(cob)
    assert "EXTERNOS" in _NOTA_COBRANZA_NETO
    assert "las comisiones las absorbe el avión" not in _NOTA_COBRANZA_NETO


# Etiqueta de la regla como la arma el API con una vigencia distinta del 1.º
# (`etiquetaReglaComisiones`).
REGLA_API = "regla desde 15-oct-2026"
_TEXTOS_CON_MES = (
    _NOTA_ENCABEZADO_COMISIONES,
    _NOTA_COMISIONES,
    _RESUMEN_COMISIONES,
    _NOTA_COBRADO_REAL,
    _NOTA_COBRANZA_NETO,
)


def test_regla_comisiones_liberal() -> None:
    assert regla_comisiones_liberal(" regla sep-2026 ") == "regla sep-2026"
    assert regla_comisiones_liberal("sep-2026") == "regla sep-2026"
    assert regla_comisiones_liberal("Regla\n desde 15-oct-2026") == "Regla desde 15-oct-2026"
    for nada in (None, "", "   ", 2026, ["regla"], {"x": 1}):
        assert regla_comisiones_liberal(nada) is None
    assert BalanceAvionRequest.model_validate({"regla_comisiones": 3}).regla_comisiones is None
    assert BalanceAvionRequest().regla_comisiones is None
    general = BalanceGeneralRequest.model_validate({"regla_comisiones": " sep-2026"})
    assert general.regla_comisiones == "regla sep-2026"


def test_textos_con_la_regla_del_api() -> None:
    """La vigencia es configurable en el API: con `regla_comisiones` ningún
    texto nombra «septiembre»; sin ella, el texto de siempre tal cual."""
    for texto in _TEXTOS_CON_MES:
        assert "septiembre 2026" in texto
        assert _con_regla(texto, None) == texto
        con = _con_regla(texto, REGLA_API)
        assert "septiembre" not in con and REGLA_API in con, texto
    assert _con_regla(_NOTA_COBRANZA_NETO, REGLA_API).startswith(
        f"Con la {REGLA_API} COBRADO REAL va NETO"
    )
    assert f"(COBRO 1..4): con la {REGLA_API} cada COBRO va NETO" in _con_regla(
        _NOTA_COBRADO_REAL, REGLA_API
    )
    assert _con_regla(_NOTA_ENCABEZADO_COMISIONES, REGLA_API).startswith(
        f"COMISIONES ({REGLA_API}, por fecha del vuelo)"
    )


def test_libros_con_la_regla_del_api() -> None:
    ind = _libro_tst()
    ind["regla_comisiones"] = REGLA_API
    wb = load_workbook(BytesIO(render_balance_avion_xlsx(BalanceAvionRequest.model_validate(ind))))
    ws = wb[MAESTRA_TST]
    assert ws["AC2"].comment.text == _con_regla(_NOTA_ENCABEZADO_COMISIONES, REGLA_API)
    pie, pie_cob = _pie(ws), _pie(wb["cobranza"])
    assert any(x.startswith(_con_regla(_NOTA_COMISIONES, REGLA_API)) for x in pie)
    assert _con_regla(_NOTA_COBRADO_REAL, REGLA_API) in pie
    assert _con_regla(_NOTA_COBRANZA_NETO, REGLA_API) in pie_cob
    assert not any("septiembre 2026" in x for x in (*pie, *pie_cob))
    # General: la etiqueta SOLO en el request también llega a las hojas de la
    # flota (el consolidado no la trae).
    payload = _general("general").model_dump(mode="json")
    payload["regla_comisiones"] = REGLA_API
    assert payload["consolidado"]["regla_comisiones"] is None
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(payload))
    assert verificar_libro(data)
    g = load_workbook(BytesIO(data))
    resumen = _pie(g["RESUMEN flota"])
    assert any(x.startswith(_con_regla(_RESUMEN_COMISIONES, REGLA_API)) for x in resumen)
    gm = g[MAESTRA]
    nota = gm.cell(row=2, column=_DISP_GENERAL.comision_col).comment.text
    assert nota == _con_regla(_NOTA_ENCABEZADO_COMISIONES, REGLA_API)
    assert _con_regla(_NOTA_COBRADO_REAL, REGLA_API) in _pie(gm)
    assert _con_regla(_NOTA_COBRANZA_NETO, REGLA_API) in _pie(g["cobranza"])
    assert _OM_REGLA_COMISIONES.format(regla=REGLA_API) in g["otros movimientos"]["A2"].value
    for hoja in ("RESUMEN flota", MAESTRA, "cobranza", "otros movimientos"):
        assert not any("septiembre 2026" in x for x in _pie(g[hoja])), hoja


def _general_dict(libros: list[dict] | None = None) -> dict:
    """Balance general (dict, forma del API) con el bloque VUELATOUR
    (empresa) de `test_balance_empresa_vuelatour`."""
    p = tf._general_payload(libros if libros is not None else [_libro_tst(), _libro_dos()])
    p["empresa"] = t_empresa._empresa_api(p)
    return p


def test_otros_movimientos_y_empresa_explican_la_regla() -> None:
    """Con vuelos de la regla (el consolidado trae COMISIONES en sus
    totales), la leyenda de 'otros movimientos' y la base del bloque
    VUELATOUR (empresa) explican la línea de INGRESO «comisión del vendedor
    a cargo del avión …» (la paga el avión, no el cliente) y que la
    comisión bancaria de VuelaTour es solo su parte; sin vuelos de la regla,
    los textos de siempre."""
    p = _general_dict()
    assert p["consolidado"]["totales"]["comisiones_mxn"] is not None
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert verificar_libro(data)
    wb = load_workbook(BytesIO(data))
    a2 = wb["otros movimientos"]["A2"].value
    om = _OM_REGLA_COMISIONES.format(regla=REGLA_COMISIONES_RESPALDO)
    assert om in a2 and a2.index(om) < a2.index(" UNA FILA POR VUELO")
    assert a2.startswith("Ingreso de VuelaTour (no del avión)")
    base = _NOTA_EMPRESA_BASE + _EMPRESA_REGLA_COMISIONES.format(regla=REGLA_COMISIONES_RESPALDO)
    pie = _pie(wb["balance"])
    assert base in pie and _NOTA_EMPRESA_BASE not in pie
    # Sin la `nota` del API: el pie fijo con la misma base aclarada.
    p["empresa"]["nota"] = None
    wb = load_workbook(
        BytesIO(render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p)))
    )
    assert _NOTA_BLOQUE_EMPRESA.replace(_NOTA_EMPRESA_BASE, base) in _pie(wb["balance"])
    # Periodo sin vuelos de la regla (o API ≤ 0.0.64): los textos de siempre.
    viejo = _general_dict([tf._payload_tst(), tf._payload_dos()])
    wb = load_workbook(
        BytesIO(render_balance_general_xlsx(BalanceGeneralRequest.model_validate(viejo)))
    )
    assert "a cargo del avión" not in wb["otros movimientos"]["A2"].value
    pie = _pie(wb["balance"])
    assert _NOTA_EMPRESA_BASE in pie and base not in pie


# ---------------------------------------------------------------------------
# 6) API previo (sin los campos): el libro de siempre salvo los encabezados.
# ---------------------------------------------------------------------------


def _miembros(data: bytes) -> dict[str, bytes]:
    """Todos los miembros del .xlsx salvo docProps/core.xml (lleva la hora)."""
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def test_api_previo_ganancia_remanente_y_cobros_de_siempre() -> None:
    req = tf._individual_tst()
    ws = load_workbook(BytesIO(render_balance_avion_xlsx(req)))[MAESTRA_TST]
    for k in range(len(req.vuelos)):
        f = 3 + k
        assert ws[f"AD{f}"].value == f"=AA{f}"  # GANANCIA = REMANENTE
        assert ws[f"AC{f}"].value is None and ws[f"AC{f}"].comment is None
    assert ws["AJ3"].value == 30000.0 and ws["AL3"].value == 30030.0  # brutos
    assert ws["AC9"].value == "=ROUND(SUM(AC3:AC8),2)"  # 0 de comision_vendedor_mxn
    # Variante general: REMANENTE con la fórmula de siempre (VENTA − PROVEEDOR).
    variante = tf._general().model_copy(update={"variante": "general"})
    gen = load_workbook(BytesIO(render_balance_general_xlsx(variante)))[MAESTRA]
    rem = _col(gen, REMANENTE_CPH)
    k = get_column_letter(_col(gen, "VENTA AVIÓN\nMXN"))
    y = get_column_letter(_col(gen, "TOTAL PARA\nPROVEEDOR (PESOS)"))
    for f in range(3, _tot(gen)):
        assert gen.cell(row=f, column=rem).value == f"=ROUND({k}{f}-{y}{f},2)", f
        assert gen.cell(row=f, column=_col(gen, COMISIONES)).value is None, f


def test_campos_nuevos_en_null_no_cambian_ni_un_byte() -> None:
    """El API 0.0.65 manda los campos en null en los vuelos anteriores a la
    vigencia (y COMISIONES de TOTALES en 0): mismo libro que sin ellos."""
    base = tf._payload_tst()
    con_null = copy.deepcopy(base)
    for fila in con_null["vuelos"]:
        fila.update(
            comisiones_mxn=None,
            comisiones_detalle=None,
            comision_banco_avion_mxn=None,
            comision_vendedor_prov_mxn=None,
        )
        for c in fila["cobros"]:
            c["neto_mxn"] = None
    con_null["totales"]["comisiones_mxn"] = 0.0
    a = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(base))
    b = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(con_null))
    assert _miembros(a) == _miembros(b)


# ---------------------------------------------------------------------------
# 7) Rutas: el JSON del API 0.0.65 entra sin 422.
# ---------------------------------------------------------------------------

TOKEN = "secreto-de-prueba"
client = TestClient(app)


def test_rutas_aceptan_el_payload_del_api(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    individual = _libro_tst()
    individual["vuelos"][2]["comisiones_detalle"] = None  # liberal
    individual["regla_comisiones"] = 2026  # liberal: no es texto ⇒ None
    res = client.post("/pdf/balance-avion-xlsx", json=individual, headers=headers)
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content))[MAESTRA_TST]
    assert ws["AC2"].value == COMISIONES and ws["AD3"].value == "=ROUND(AA3-AC3,2)"
    res = client.post(
        "/pdf/balance-general-xlsx",
        json=_general("general").model_dump(mode="json"),
        headers=headers,
    )
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content))[MAESTRA]
    assert ws.cell(row=2, column=_DISP_GENERAL.comision_col).value == COMISIONES
