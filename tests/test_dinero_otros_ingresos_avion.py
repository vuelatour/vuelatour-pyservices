"""Libro Dinero con los otros ingresos ASIGNADOS a un avión (13-oct-2026, API
0.0.75; invariante «INGRESO ASIGNADO A UN AVIÓN» del API).

Un ingreso ING-n con avión cuenta en el balance y el reparto de SU avión, NUNCA
como ingreso de VuelaTour. El API ya lo excluye de `otros_ingresos` y de
`utilidades_otros_ingresos_mxn` y lo manda aparte en `otros_ingresos_avion`
(solo con filas) y, por avión, en `utilidades_aviones[].otros_ingresos_mxn`
(solo con dato). Aquí:
- 'Otros ingresos': las filas de avión al FINAL, en una sección aparte y FUERA
  de TOTALES (la CLAVE teñida con el color del avión), y la fila 2 lo dice;
- 'utilidades': con el dato, cada avión gana la columna «OTROS INGRESOS
  <matrícula>» (la primera de sus 5) y la fila 9 dice que esos ingresos no
  están en OTROS INGRESOS de VuelaTour;
- sin los campos: el libro de siempre, miembro a miembro (huella calculada con
  el código de HEAD d58faa1, antes de este cambio).
"""

from __future__ import annotations

import copy
import hashlib
import zipfile
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reportes import DineroOtroIngresoFila, DineroUtilidadAvion, DineroXlsxRequest
from app.services.dinero_xlsx import (
    _LEYENDA_OTROS_INGRESOS_AVION,
    _SECCION_OTROS_INGRESOS_AVION,
    _UTILIDADES_NOTA_OTROS_INGRESOS_AVION,
    render_dinero_xlsx,
)

SECCION = (
    "OTROS INGRESOS DE AVIÓN (cuentan en el balance y en el reparto de su avión; "
    "no suman en este libro)"
)
LEYENDA = (
    " Al final, aparte y fuera de TOTALES, los otros ingresos asignados a un "
    "avión: cuentan en la hoja 'Otros ingresos' del balance de ese avión y en su "
    "reparto, y en 'utilidades' van en la columna OTROS INGRESOS de ese avión."
)
NOTA_FILA9 = (
    " Los ingresos registrados en Ingresos y asignados a un avión NO están en "
    "este total: van en la columna OTROS INGRESOS de su avión (netos de su "
    "comisión bancaria)."
)

# El Libro Dinero de septiembre SIN ingresos de avión, con la forma del API
# (una fila de vuelo con su comisión y un ING sin avión). Es el MISMO payload
# con el que se calculó `_HUELLA_BASE_HEAD`.
_BASE: dict = {
    "periodo_desde": "2026-09-01",
    "periodo_hasta": "2026-09-30",
    "generado": "2026-10-13",
    "leyenda_colores": [
        {"matricula": "N990GG", "color": "#3B82F6", "modelo": "Cessna 206"},
        {"matricula": "N4142R", "color": "#F97316", "modelo": "Piper Seneca"},
    ],
    "vuelos": [
        {
            "clave": "vtbeh",
            "matricula": "N990GG",
            "color": "#3B82F6",
            "fecha": "2026-09-05",
            "ruta": "CUN-PTU-CUN",
            "tiempo": 1.5,
            "total_cobrado_usd": 1200.0,
            "tc_venta": 18.0,
            "total_cobrado_mxn": 21600.0,
            "status_cobro": "COBRADO",
            "cobros": [{"fecha": "2026-09-05", "monto_mxn": 21600.0}],
            "total_cobros_mxn": 21600.0,
            "me_deben_mxn": 0.0,
        }
    ],
    "otros_ingresos": [
        {
            "clave": "vtbeh",
            "fecha_vuelo": "2026-09-05",
            "concepto_ingreso": "TUAs/extras",
            "ingreso_mxn": 1800.0,
            "fecha_ingreso": "2026-09-05",
            "remanente_mxn": 1800.0,
        },
        {
            "clave": "ING-3",
            "concepto_ingreso": "Intereses · Banco",
            "ingreso_mxn": 85.33,
            "fecha_ingreso": "2026-09-30",
            "remanente_mxn": 85.33,
        },
    ],
    "otros_gastos": [
        {"fecha": "2026-09-10", "concepto": "Renta oficina", "monto_mxn": 15000.0,
         "acumulado_mxn": 15000.0},
    ],
    "utilidades_otros_ingresos_mxn": 1885.33,
    "utilidades_otros_gastos_mxn": 15000.0,
    "utilidades_tc": 18.0,
    "utilidades_aviones": [
        {"matricula": "N990GG", "gastos_indirectos_mxn": 1200.0, "permisos_mxn": 2500.0},
        {"matricula": "N4142R", "otros_gastos_mxn": 980.0, "combustible_mxn": 6000.0},
    ],
}  # fmt: skip

# Los ingresos de avión como los arma el API: la MISMA fila de
# `filaLibroDeIngreso` + `matricula`/`avion_color` (el concepto ya trae la
# matrícula).
_ING_AVION = [
    {
        "clave": "ING-4",
        "fecha_vuelo": None,
        "concepto_egreso": None,
        "egreso_mxn": None,
        "fecha_egreso": None,
        "nota_egreso": None,
        "concepto_ingreso": (
            "Reembolsos y devoluciones recibidos · TRANSF INTERBANCARIA SPEI · "
            "angel alvarez · N990GG"
        ),
        "ingreso_mxn": 46550.8,
        "fecha_ingreso": "2026-09-04",
        "remanente_mxn": 46550.8,
        "factura": None,
        "matricula": "N990GG",
        "avion_color": "#3B82F6",
    },
    {
        "clave": "ING-9",
        "fecha_vuelo": None,
        "concepto_egreso": "comisión bancaria",
        "egreso_mxn": 1750.0,
        "fecha_egreso": "2026-09-20",
        "nota_egreso": None,
        "concepto_ingreso": "Venta de activos · Venta de una hélice usada · N4142R",
        "ingreso_mxn": 35000.0,
        "fecha_ingreso": "2026-09-20",
        "remanente_mxn": 33250.0,
        "factura": None,
        "matricula": "N4142R",
        "avion_color": "#F97316",
    },
]


def _req(*, avion: bool = True, columna: bool = True) -> DineroXlsxRequest:
    p = copy.deepcopy(_BASE)
    if avion:
        p["otros_ingresos_avion"] = copy.deepcopy(_ING_AVION)
    if columna:
        p["utilidades_aviones"][0]["otros_ingresos_mxn"] = 46550.8
    return DineroXlsxRequest.model_validate(p)


def _libro(xlsx: bytes):
    return load_workbook(BytesIO(xlsx))


def _miembros(data: bytes) -> dict[str, bytes]:
    """Todos los miembros del .xlsx salvo docProps/core.xml (lleva la hora)."""
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _huella(data: bytes) -> str:
    h = hashlib.sha256()
    for nombre, contenido in _miembros(data).items():
        h.update(nombre.encode())
        h.update(contenido)
    return h.hexdigest()


# Calculada con el código de HEAD d58faa1 (13-oct-2026), ANTES de este
# cambio, sobre `_BASE`: si se mueve, un API ≤ 0.0.74 ya no ve su libro.
_HUELLA_BASE_HEAD = "a4c952a2004a8b2d287047b106ea51746c1920ab22a5bb615037833b841685c5"


def test_esquemas_aditivos_con_default() -> None:
    assert DineroXlsxRequest().otros_ingresos_avion == []
    nulo = DineroXlsxRequest.model_validate({"otros_ingresos_avion": None})
    assert nulo.otros_ingresos_avion == []
    fila = DineroOtroIngresoFila()
    assert (fila.matricula, fila.avion_color) == (None, None)
    assert DineroUtilidadAvion().otros_ingresos_mxn is None


def test_textos_son_los_del_contrato() -> None:
    assert _SECCION_OTROS_INGRESOS_AVION == SECCION
    assert _LEYENDA_OTROS_INGRESOS_AVION == LEYENDA
    assert _UTILIDADES_NOTA_OTROS_INGRESOS_AVION == NOTA_FILA9


def test_seccion_aparte_al_final_y_fuera_de_totales() -> None:
    base = _libro(render_dinero_xlsx(_req(avion=False, columna=False)))["Otros ingresos"]
    ws = _libro(render_dinero_xlsx(_req()))["Otros ingresos"]
    # TOTALES: los MISMOS números que sin los ingresos de avión.
    fila_tot = next(c.row for c in ws["A"] if c.value == "TOTALES")
    assert fila_tot == next(c.row for c in base["A"] if c.value == "TOTALES") == 6
    for col in (4, 7, 9):
        assert ws.cell(row=fila_tot, column=col).value == base.cell(row=fila_tot, column=col).value
    assert ws.cell(row=fila_tot, column=7).value == round(1800.0 + 85.33, 2)
    # Una fila en blanco y el encabezado de sección combinado A:J.
    assert all(ws.cell(row=fila_tot + 1, column=c).value is None for c in range(1, 11))
    sec = ws.cell(row=fila_tot + 2, column=1)
    assert sec.value == SECCION
    assert sec.font.b and sec.fill.fgColor.rgb.endswith("EEF2F7")
    assert any(str(rg) == f"A{fila_tot + 2}:J{fila_tot + 2}" for rg in ws.merged_cells.ranges)
    # Las filas de avión, con el MISMO pintado de las de arriba (10 columnas,
    # montos con formato) y la CLAVE teñida con el color del avión.
    r = fila_tot + 3
    assert [ws.cell(row=r, column=c).value for c in (1, 6, 7, 8, 9)] == [
        "ING-4",
        _ING_AVION[0]["concepto_ingreso"],
        46550.8,
        "04/09/2026",
        46550.8,
    ]
    assert ws.cell(row=r, column=1).fill.fgColor.rgb.endswith("3B82F6")
    assert ws.cell(row=r, column=7).number_format == '"$"#,##0.00'
    assert ws.cell(row=r + 1, column=1).value == "ING-9"
    assert ws.cell(row=r + 1, column=1).fill.fgColor.rgb.endswith("F97316")
    assert [ws.cell(row=r + 1, column=c).value for c in (3, 4, 5, 9)] == [
        "comisión bancaria",
        1750,
        "20/09/2026",
        33250,
    ]
    # Después de ellas, nada: no hay otros TOTALES.
    assert [c.value for c in ws["A"] if c.value == "TOTALES"] == ["TOTALES"]
    assert ws.max_row == r + 1


def test_fila2_dice_donde_estan_solo_con_filas_de_avion() -> None:
    con = _libro(render_dinero_xlsx(_req()))["Otros ingresos"]["A2"].value
    sin = _libro(render_dinero_xlsx(_req(avion=False)))["Otros ingresos"]["A2"].value
    assert con == sin + LEYENDA
    assert LEYENDA not in sin


def test_utilidades_cinco_columnas_por_avion_y_fila9() -> None:
    ws = _libro(render_dinero_xlsx(_req()))["utilidades"]
    encabezados = [c.value for c in ws[3]]
    assert encabezados == [
        "PERIODO",
        "CHARTERS",
        "OTROS\nINGRESOS",
        "UTILIDAD ANTES\nGASTOS",
        "OTROS\nGASTOS",
        "COMBUSTIBLE\nDEL MES",
        "OTROS\nINGRESOS\nN990GG",
        "GASTOS\nINDIRECTOS\nN990GG",
        "OTROS GASTOS\nN990GG",
        "PERMISOS\nN990GG",
        "COMBUSTIBLE\nN990GG",
        "OTROS\nINGRESOS\nN4142R",
        "GASTOS\nINDIRECTOS\nN4142R",
        "OTROS GASTOS\nN4142R",
        "PERMISOS\nN4142R",
        "COMBUSTIBLE\nN4142R",
        "UTILIDAD\nDESPUES GASTOS",
        "TIPO\nCAMBIO",
        "UTILIDAD\nUSD",
    ]
    valores = [c.value for c in ws[4]]
    assert valores[2] == 1885.33  # OTROS INGRESOS de VuelaTour: sin los de avión
    assert valores[6:11] == [46550.8, 1200, None, 2500, None]
    assert valores[11:16] == [None, None, 980, None, 6000]
    # El T.C. sigue en su columna (la penúltima) con su formato.
    assert valores[17] == 18 and ws.cell(row=4, column=18).number_format == "0.00####"
    assert ws["A9"].value.endswith(NOTA_FILA9)
    assert ws["A9"].value.count(NOTA_FILA9) == 1


def test_utilidades_sin_el_dato_la_hoja_de_siempre() -> None:
    con = _libro(render_dinero_xlsx(_req(columna=False)))["utilidades"]
    sin = _libro(render_dinero_xlsx(_req(avion=False, columna=False)))["utilidades"]
    assert [c.value for c in con[3]] == [c.value for c in sin[3]]
    assert [c.value for c in con[4]] == [c.value for c in sin[4]]
    assert NOTA_FILA9 not in con["A9"].value


def test_sin_los_campos_libro_byte_a_byte() -> None:
    base = render_dinero_xlsx(DineroXlsxRequest.model_validate(copy.deepcopy(_BASE)))
    assert _huella(base) == _HUELLA_BASE_HEAD
    # Con las llaves en vacío/None: el mismo libro, miembro a miembro.
    p = copy.deepcopy(_BASE)
    p["otros_ingresos_avion"] = []
    for a in p["utilidades_aviones"]:
        a["otros_ingresos_mxn"] = None
    assert _miembros(render_dinero_xlsx(DineroXlsxRequest.model_validate(p))) == _miembros(base)


TOKEN = "tok-dinero-oi"
client = TestClient(app)


def test_ruta_acepta_el_payload_0075(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    res = client.post(
        "/pdf/dinero-xlsx",
        json=_req().model_dump(mode="json"),
        headers={"X-Internal-Token": TOKEN},
    )
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content))["Otros ingresos"]
    assert SECCION in [c.value for c in ws["A"]]
