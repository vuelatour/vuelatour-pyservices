"""Columna FACTURA VUELATOUR en la hoja maestra de los balances (30-sep-2026,
API 0.0.45).

Pedido de Marie (para Ale), con la foto del Excel «balance-XA-VGV-2026-09-01-
2026-09-30», hoja «reporte horas XA-VGV»: «en balance por avión el reporte de
excel se ocupa que diga el num de factura que nosotros emitimos del servicio,
no aparece en la columna, y si puede salir en el reporte general también».

El API manda por fila `factura_vuelatour` YA resuelto con su cascada única
(`etiquetasFacturaDeVuelos`: CFDI timbrado vivo → facturas EMITIDAS vigentes
«A-0424» → `vuelo.factura_folio` tecleado → etiqueta del estatus → nada).
Aquí solo se pinta, TAL CUAL, en una columna nueva AL FINAL del bloque STATUS
DE COBROS — en el libro individual y en la hoja «reporte horas FLOTA» del
Balance general (las dos salen de `_hoja_maestra`). Sin el campo (API viejo) o
en null, la celda va vacía; la fila TOTALES también. Nada más se mueve: la
columna va al final, así que COBRO 1, COMISIÓN VENDEDOR (hoy COMISIONES,
`_DISP_MENSUAL`) y las columnas de la izquierda quedan donde estaban.
Desde el 7-oct-2026 (API 0.0.66) la hoja de vuelos del Balance general no
lleva COMISIONES (`_COLS_FLOTA`): ahí FACTURA VUELATOUR es la 46 (AT).

13-oct-2026 (API 0.0.74, invariante 60 del API, decisión 4 del lote del
cliente): la columna puede decir «Pendiente de factura» / «No lleva factura»
(llegan hechos del API, se pintan tal cual como TEXTO) y vacía = sin cobros.
La bandera raíz ADITIVA `factura_vuelatour_segun_cobros` (individual,
`aviones[]` y `consolidado`) cambia ÚNICAMENTE la nota al pie de la hoja
maestra; sin ella (None/False/ausente) el libro es byte a byte el de antes
(huellas `_firma_hoja` congeladas con el código de HEAD 12b5e86, que para
esta hoja es el mismo que 4d206ab).
"""

import hashlib
import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAvionRequest,
    BalanceAvionVuelo,
    BalanceGeneralRequest,
    BalanceHojaOtrosMovimientos,
    DineroXlsxRequest,
)
from app.services.balance_avion_xlsx import (
    _COLS,
    _COLS_FLOTA,
    _DISP_MENSUAL,
    _NOTA_FACTURA_VUELATOUR,
    _NOTA_FACTURA_VUELATOUR_SEGUN_COBROS,
    FILL_COBROS,
    LIGHT,
    _nota_factura_vuelatour,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from app.services.dinero_xlsx import _COLS as _COLS_DINERO
from app.services.dinero_xlsx import render_dinero_xlsx

ENCABEZADO = "FACTURA\nVUELATOUR"
# Índice (1-based) de la columna nueva: la ÚLTIMA de la hoja maestra (47 en el
# libro individual; 46 en la hoja de vuelos del Balance general, sin
# COMISIONES desde el 7-oct-2026).
COL = len(_COLS)
COL_GENERAL = len(_COLS_FLOTA)
FILA_DATOS = 3  # filas 1-2 = encabezado de grupo / columna


def _vuelo(clave: str, **extra) -> dict:
    return {
        "clave": clave,
        "fecha": "2026-09-05",
        "ruta": "CUN-PTU-CUN",
        "estado": "COMPLETADO",
        "status_cobro": "Cobrado",
        "cobros": [{"fecha": "2026-09-05", "monto_mxn": 30000.0}],
        "cobrado_real_mxn": 30000.0,
        "cobrado_mxn": 30000.0,
        "por_cobrar_mxn": 0.0,
        "por_cobrar_usd": 0.0,
        **extra,
    }


# Vuelos como los de la foto: uno con factura emitida, uno con dos facturas
# vigentes, uno solo con el estatus, uno SIN factura y un cancelado que sí la
# trae (los cancelados también la llevan).
_VUELOS = [
    _vuelo("XA-VGV-300", factura_vuelatour="A-0424"),
    _vuelo("XA-VGV-301", factura_vuelatour="A-0425, A-0430"),
    _vuelo("XA-VGV-302", factura_vuelatour="Facturado"),
    _vuelo("XA-VGV-303", factura_vuelatour=None),
    _vuelo("XA-VGV-304", estado="CANCELADO", factura_vuelatour="A-0399"),
]
_ESPERADO = ["A-0424", "A-0425, A-0430", "Facturado", None, "A-0399"]
_TOTALES = {"cobrado_real_mxn": 150000.0, "por_cobrar_usd": 0.0}


# Centinela: «la llave no viaja» (API ≤ 0.0.73), distinto de None/False.
_AUSENTE = object()


def _bandera(bandera) -> dict:
    return {} if bandera is _AUSENTE else {"factura_vuelatour_segun_cobros": bandera}


def _individual(vuelos=None, bandera=_AUSENTE) -> BalanceAvionRequest:
    return BalanceAvionRequest(
        matricula="XA-VGV",
        periodo_desde="2026-09-01",
        periodo_hasta="2026-09-30",
        vuelos=_VUELOS if vuelos is None else vuelos,
        totales=_TOTALES,
        **_bandera(bandera),
    )


def _general(vuelos=None, bandera=_AUSENTE, variante=None) -> BalanceGeneralRequest:
    return BalanceGeneralRequest(
        periodo_desde="2026-09-01",
        periodo_hasta="2026-09-30",
        resumen=[{"matricula": "XA-VGV", "color": "#FF0000", "vuelos": 5}],
        consolidado=BalanceAvionRequest(
            matricula="FLOTA",
            periodo_desde="2026-09-01",
            periodo_hasta="2026-09-30",
            vuelos=_VUELOS if vuelos is None else vuelos,
            totales=_TOTALES,
            **_bandera(bandera),
        ),
        **({} if variante is None else {"variante": variante}),
    )


def _maestra_individual(req: BalanceAvionRequest | None = None):
    data = render_balance_avion_xlsx(req or _individual())
    wb = load_workbook(BytesIO(data), data_only=True)
    return wb["reporte horas XA-VGV"]


def _maestra_general(req: BalanceGeneralRequest | None = None):
    data = render_balance_general_xlsx(req or _general())
    wb = load_workbook(BytesIO(data), data_only=True)
    return wb["reporte horas FLOTA"]


def _col_factura(ws) -> int:
    """Columna FACTURA VUELATOUR de ESTA hoja de vuelos."""
    return COL_GENERAL if ws.title == "reporte horas FLOTA" else COL


def _columna(ws, n: int) -> list:
    return [ws.cell(row=FILA_DATOS + i, column=_col_factura(ws)).value for i in range(n)]


# ---------------------------------------------------------------------------
# Layout: la columna va al FINAL de STATUS DE COBROS y no mueve nada más.
# ---------------------------------------------------------------------------


def test_cols_la_columna_nueva_es_la_ultima_de_status_de_cobros() -> None:
    assert _COLS[-1] == ("STATUS DE COBROS", ENCABEZADO, "factura_vuelatour", None)
    # La vecina de la izquierda es la que antes cerraba la hoja.
    assert _COLS[-2][1:3] == ("POR COBRAR\nUSD", "por_cobrar_usd")
    assert [c[2] for c in _COLS].count("factura_vuelatour") == 1
    # Los índices derivados siguen apuntando a su columna (se buscan en
    # `_COLS`, pero se congela que no se corrieron).
    cobro1, comision = _DISP_MENSUAL.cobro1_col, _DISP_MENSUAL.comision_col
    assert _COLS[cobro1 - 1][1] == "COBRO 1\nFECHA"
    assert _COLS[comision - 1][2] == "comisiones_mxn"  # COMISIONES (6-oct-2026)
    assert (COL, cobro1, comision) == (47, 35, 29)


@pytest.mark.parametrize("hoja", ["individual", "general"])
def test_encabezado_grupo_relleno_y_ancho(hoja) -> None:
    ws = _maestra_individual() if hoja == "individual" else _maestra_general()
    cols = _COLS if hoja == "individual" else _COLS_FLOTA
    col = len(cols)
    assert (col, _col_factura(ws)) == ((47, 47) if hoja == "individual" else (46, 46))
    assert ws.max_column == col
    assert ws.cell(row=2, column=col).value == ENCABEZADO
    assert ws.cell(row=2, column=col - 1).value == "POR COBRAR\nUSD"
    # El título de grupo STATUS DE COBROS (fila 1, combinado) la abarca.
    grupo = next(
        m
        for m in ws.merged_cells.ranges
        if m.min_row == 1 and ws.cell(row=1, column=m.min_col).value == "STATUS DE COBROS"
    )
    assert grupo.max_col == col
    assert ws.cell(row=1, column=grupo.min_col).value == "STATUS DE COBROS"
    # Mismo verde suave del bloque en encabezado y datos.
    assert ws.cell(row=2, column=col).fill.fgColor.rgb.endswith(FILL_COBROS)
    assert ws.cell(row=FILA_DATOS, column=col).fill.fgColor.rgb.endswith(FILL_COBROS)
    assert ws.column_dimensions[get_column_letter(col)].width == 16
    # Los anchos de la izquierda no cambian (COBRADO REAL/AVIÓN 17, resto 13).
    anchos = {
        h: ws.column_dimensions[get_column_letter(i)].width
        for i, (_g, h, _a, _f) in enumerate(cols, start=1)
    }
    assert anchos["POR COBRAR\nUSD"] == 13
    assert anchos["COBRADO REAL\nMXN (Σ depósitos)"] == 17
    assert ws.freeze_panes == "D3"


# ---------------------------------------------------------------------------
# Valores: la etiqueta del API TAL CUAL; sin dato ⇒ vacía; TOTALES vacía.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("hoja", ["individual", "general"])
def test_valor_tal_cual_y_vacio_sin_factura(hoja) -> None:
    ws = _maestra_individual() if hoja == "individual" else _maestra_general()
    assert _columna(ws, len(_VUELOS)) == _ESPERADO
    # Las filas son las del payload, en orden (la clave lo confirma).
    assert [ws.cell(row=FILA_DATOS + i, column=1).value for i in range(len(_VUELOS))] == [
        v["clave"] for v in _VUELOS
    ]
    # Texto, no número ni fecha.
    celda = ws.cell(row=FILA_DATOS, column=_col_factura(ws))
    assert celda.data_type == "s"
    assert celda.number_format == "General"


@pytest.mark.parametrize("hoja", ["individual", "general"])
def test_fila_totales_va_vacia(hoja) -> None:
    ws = _maestra_individual() if hoja == "individual" else _maestra_general()
    fila_tot = FILA_DATOS + len(_VUELOS)
    assert ws.cell(row=fila_tot, column=1).value == "TOTALES"
    celda = ws.cell(row=fila_tot, column=_col_factura(ws))
    assert celda.value is None
    assert celda.fill.fgColor.rgb.endswith(LIGHT)
    # La vecina sí trae su total: la fila TOTALES no se corrió.
    assert ws.cell(row=fila_tot, column=_col_factura(ws) - 1).value == 0.0


def test_api_viejo_sin_la_clave_deja_la_columna_vacia() -> None:
    """Un API 0.0.44 no manda el campo: el libro sale igual salvo la columna
    nueva, que va vacía (con su encabezado y su relleno)."""
    viejos = [{k: v for k, v in f.items() if k != "factura_vuelatour"} for f in _VUELOS]
    for ws in (_maestra_individual(_individual(viejos)), _maestra_general(_general(viejos))):
        assert ws.cell(row=2, column=_col_factura(ws)).value == ENCABEZADO
        assert _columna(ws, len(viejos)) == [None] * len(viejos)


def test_nada_a_la_izquierda_cambia_con_o_sin_factura() -> None:
    """Con y sin facturas registradas, todas las celdas de las columnas 1..46
    (valores, formatos y rellenos) son idénticas: la columna es aditiva."""
    sin = [{**f, "factura_vuelatour": None} for f in _VUELOS]
    a, b = _maestra_individual(), _maestra_individual(_individual(sin))
    assert a.max_row == b.max_row
    for fila in range(1, a.max_row + 1):
        for col in range(1, COL):
            ca, cb = a.cell(row=fila, column=col), b.cell(row=fila, column=col)
            assert (ca.value, ca.number_format, ca.fill.fgColor.rgb) == (
                cb.value,
                cb.number_format,
                cb.fill.fgColor.rgb,
            ), (fila, col)


def test_multi_avion_todas_sus_filas_traen_la_misma_factura() -> None:
    """La factura es del VUELO: el API manda la misma etiqueta en cada fila
    COMPARTIDO del general; aquí se pinta en las dos."""
    filas = [
        _vuelo(
            "XA-VGV-310 · 50 %",
            multi_avion=True,
            participacion=0.5,
            participacion_fuente="tramos",
            avion_color="#FF0000",
            factura_vuelatour="A-0431",
        ),
        _vuelo(
            "XB-PEV-310 · 50 %",
            multi_avion=True,
            participacion=0.5,
            participacion_fuente="tramos",
            avion_color="#00FF00",
            factura_vuelatour="A-0431",
        ),
    ]
    ws = _maestra_general(_general(filas))
    assert _columna(ws, 2) == ["A-0431", "A-0431"]


def test_folio_que_empieza_con_igual_se_queda_como_texto() -> None:
    """El folio lo teclea la oficina (`vuelo.factura_folio`): openpyxl vuelve
    FÓRMULA toda cadena con «=» al inicio y Excel abriría el libro con error."""
    ws = _maestra_individual(_individual([_vuelo("XA-VGV-320", factura_vuelatour="=A-12")]))
    celda = ws.cell(row=FILA_DATOS, column=COL)
    assert celda.value == "=A-12"
    assert celda.data_type == "s"


def test_otros_movimientos_folio_con_igual_tambien_es_texto() -> None:
    """La MISMA etiqueta viaja en «factura vuelatour» de 'otros movimientos'
    del general: sin la guarda ahí, el libro entero abriría con error aunque
    la hoja maestra la tenga. Los folios normales no cambian."""
    req = _general([_vuelo("XA-VGV-321", factura_vuelatour="=A-12")])
    req.consolidado.otros_movimientos = BalanceHojaOtrosMovimientos(
        filas=[
            {
                "clave": "XA-VGV-321",
                "concepto_ingreso": "TUA",
                "ingreso_mxn": 500.0,
                "factura": "=A-12",
            },
            {
                "clave": "XA-VGV-322",
                "concepto_ingreso": "TUA",
                "ingreso_mxn": 500.0,
                "factura": "A-0424",
            },
        ],
        filas_sueltas=[{"concepto_egreso": "Renta", "egreso_mxn": 100.0, "factura": "=B-1"}],
    )
    wb = load_workbook(BytesIO(render_balance_general_xlsx(req)), data_only=True)
    ws = wb["otros movimientos"]
    celdas = {c.value: c for c in ws["J"] if c.value in ("=A-12", "A-0424", "=B-1")}
    assert set(celdas) == {"=A-12", "A-0424", "=B-1"}
    assert all(c.data_type == "s" for c in celdas.values())
    # Y la hoja maestra del mismo libro también.
    assert wb["reporte horas FLOTA"].cell(row=FILA_DATOS, column=COL_GENERAL).data_type == "s"


# ---------------------------------------------------------------------------
# Nota al pie: en los dos libros, una sola vez.
# ---------------------------------------------------------------------------


_BANDERAS = [
    pytest.param(_AUSENTE, _NOTA_FACTURA_VUELATOUR, id="sin-bandera"),
    pytest.param(None, _NOTA_FACTURA_VUELATOUR, id="None"),
    pytest.param(False, _NOTA_FACTURA_VUELATOUR, id="False"),
    pytest.param(True, _NOTA_FACTURA_VUELATOUR_SEGUN_COBROS, id="True"),
]


def _maestra(hoja: str, bandera):
    if hoja == "individual":
        return _maestra_individual(_individual(bandera=bandera))
    if hoja == "general":
        return _maestra_general(_general(bandera=bandera))
    return _maestra_general(_general(bandera=bandera, variante="general"))


@pytest.mark.parametrize("hoja", ["individual", "general", "general-variante-general"])
@pytest.mark.parametrize(("bandera", "esperada"), _BANDERAS)
def test_nota_al_pie_explica_la_columna(hoja, bandera, esperada) -> None:
    """Sin la bandera (ausente/None/False) la nota de siempre; con `True`
    (API 0.0.74) la que explica «Pendiente de factura» / «No lleva factura»
    / vacía = sin cobros — una sola vez, en la MISMA fila, en el libro
    individual, en «reporte horas FLOTA» mensual y en la variante general
    (ahí la recibe del `consolidado`)."""
    ws = _maestra(hoja, bandera)
    notas = [
        c.value
        for c in ws["A"]
        if isinstance(c.value, str) and c.value.startswith("FACTURA VUELATOUR =")
    ]
    assert notas == [esperada]
    assert notas[0].startswith(
        "FACTURA VUELATOUR = folio de la factura del servicio emitida al "
        "cliente (timbrada, registrada en Facturas emitidas o capturada en el "
        "vuelo)"
    )
    # Va justo después de la nota del ESTATUS DE COBRO (bloque de cobros).
    fila = next(c.row for c in ws["A"] if c.value == esperada)
    assert ws.cell(row=fila - 1, column=1).value.startswith("El ESTATUS DE COBRO por vuelo")


def test_nota_factura_vuelatour_exige_true_y_la_de_siempre_no_cambio() -> None:
    assert _nota_factura_vuelatour(_individual()) == _NOTA_FACTURA_VUELATOUR
    assert _nota_factura_vuelatour(_individual(bandera=None)) == _NOTA_FACTURA_VUELATOUR
    assert _nota_factura_vuelatour(_individual(bandera=False)) == _NOTA_FACTURA_VUELATOUR
    assert (
        _nota_factura_vuelatour(_individual(bandera=True))
        == _NOTA_FACTURA_VUELATOUR_SEGUN_COBROS
    )
    # Texto EXACTO del contrato (API 0.0.74).
    assert _NOTA_FACTURA_VUELATOUR_SEGUN_COBROS == (
        "FACTURA VUELATOUR = folio de la factura del servicio emitida al cliente "
        "(timbrada, registrada en Facturas emitidas o capturada en el vuelo); sin "
        "folio registrado dice el estatus del vuelo («Facturado» / «Factura "
        "elaborada y enviada»); sin factura registrada dice «Pendiente de factura» "
        "cuando algún cobro entró a HSBC o Scotiabank (o fue por transferencia, "
        "link de pago HSBC o cheque sin cuenta capturada) o se pidió factura desde "
        "el vuelo, «No lleva factura» cuando todos sus cobros fueron por "
        "BillPocket, Paywise, efectivo o dólares, y vacía = vuelo sin cobros. Es "
        "la factura del VUELO completo: en un vuelo multi-avión todas sus filas "
        "traen la misma."
    )
    # La de siempre, ni una letra (la congelan las huellas de los demás tests).
    assert _NOTA_FACTURA_VUELATOUR.endswith(
        "vacía = sin factura. Es la factura del VUELO completo: en un vuelo "
        "multi-avión todas sus filas traen la misma."
    )


# ---------------------------------------------------------------------------
# Huellas: sin la bandera, la hoja maestra es byte a byte la de antes.
# ---------------------------------------------------------------------------


def _firma_hoja(data: bytes, nombre: str) -> str:
    """Huella del LAYOUT de una hoja (misma receta que
    test_balance_extension_horario / test_balance_inventario_xlsx): valores
    (caché `data_only`), formatos, fuentes, rellenos, alineación, bordes,
    merges, anchos, altos, panel fijo y comentarios; sin las constantes del
    encabezado A1:D1."""
    ws = load_workbook(BytesIO(data), data_only=True)[nombre]
    for col in range(1, 5):
        ws._cells.pop((1, col), None)
    partes: list = []
    for fila in ws.iter_rows():
        for c in fila:
            f = c.font
            partes.append([
                c.coordinate, repr(c.value), c.number_format,
                bool(f.b), bool(f.i), f.sz,
                f.color.rgb if f.color is not None else None,
                c.fill.fgColor.rgb if c.fill.fill_type else None,
                c.alignment.horizontal, c.alignment.vertical,
                bool(c.alignment.wrap_text), c.border.left.style,
                c.comment.text if c.comment else None,
            ])
    partes.append(sorted(str(r) for r in ws.merged_cells.ranges))
    partes.append(sorted((k, d.width) for k, d in ws.column_dimensions.items()))
    partes.append(sorted((k, d.height) for k, d in ws.row_dimensions.items() if d.height))
    partes.append(ws.freeze_panes)
    return hashlib.sha256(json.dumps(partes, default=str).encode()).hexdigest()


# Calculadas con el código de HEAD 12b5e86 (13-oct-2026) ANTES de tocar el
# render, sobre `_individual()` / `_general()` SIN la bandera. Si una cambia,
# un API ≤ 0.0.73 ya no ve su libro de siempre.
_FIRMA_INDIVIDUAL_SIN_BANDERA = "a8875073e6bfd7722332b64294f1e3e162870737084a6f5fcd455df1ab38143a"
_FIRMA_GENERAL_SIN_BANDERA = "eeab475d2c1543175f2f91dbec23db1ba1333d5f86d4997302db1082e0afb728"


@pytest.mark.parametrize("bandera", [_AUSENTE, None, False], ids=["sin-bandera", "None", "False"])
def test_sin_bandera_las_dos_hojas_maestras_son_byte_a_byte_las_de_antes(bandera) -> None:
    assert (
        _firma_hoja(render_balance_avion_xlsx(_individual(bandera=bandera)), "reporte horas XA-VGV")
        == _FIRMA_INDIVIDUAL_SIN_BANDERA
    )
    assert (
        _firma_hoja(render_balance_general_xlsx(_general(bandera=bandera)), "reporte horas FLOTA")
        == _FIRMA_GENERAL_SIN_BANDERA
    )


@pytest.mark.parametrize("hoja", ["individual", "general"])
def test_con_bandera_solo_cambia_la_celda_de_la_nota(hoja) -> None:
    """Con `True` la ÚNICA diferencia contra el libro sin bandera es la celda
    de la nota (columna A de esa fila): valores, formatos y rellenos de todo
    lo demás, idénticos."""
    a, b = _maestra(hoja, _AUSENTE), _maestra(hoja, True)
    assert (a.max_row, a.max_column) == (b.max_row, b.max_column)
    diferencias = []
    for fila in range(1, a.max_row + 1):
        for col in range(1, a.max_column + 1):
            ca, cb = a.cell(row=fila, column=col), b.cell(row=fila, column=col)
            if (ca.value, ca.number_format, ca.fill.fgColor.rgb) != (
                cb.value,
                cb.number_format,
                cb.fill.fgColor.rgb,
            ):
                diferencias.append((ca.coordinate, ca.value, cb.value))
    assert len(diferencias) == 1, diferencias
    coord, antes, despues = diferencias[0]
    assert coord.startswith("A")
    assert (antes, despues) == (_NOTA_FACTURA_VUELATOUR, _NOTA_FACTURA_VUELATOUR_SEGUN_COBROS)


def test_columnas_1_a_46_intactas_con_y_sin_bandera_en_las_filas_de_datos() -> None:
    a, b = _maestra_individual(), _maestra_individual(_individual(bandera=True))
    for fila in range(1, FILA_DATOS + len(_VUELOS) + 1):  # encabezados, datos y TOTALES
        for col in range(1, COL):
            ca, cb = a.cell(row=fila, column=col), b.cell(row=fila, column=col)
            assert (ca.value, ca.number_format, ca.fill.fgColor.rgb) == (
                cb.value,
                cb.number_format,
                cb.fill.fgColor.rgb,
            ), (fila, col)


# ---------------------------------------------------------------------------
# Los textos nuevos se pintan TAL CUAL (texto, sin fórmula) en los tres
# libros que llevan la columna.
# ---------------------------------------------------------------------------

_VUELOS_COBROS = [
    _vuelo("XA-VGV-330", factura_vuelatour="Pendiente de factura"),
    _vuelo("XA-VGV-331", factura_vuelatour="No lleva factura"),
    _vuelo("XA-VGV-332", factura_vuelatour=None),
]
_ESPERADO_COBROS = ["Pendiente de factura", "No lleva factura", None]
_TEXTOS_NUEVOS = ("Pendiente de factura", "No lleva factura")


@pytest.mark.parametrize("bandera", [_AUSENTE, True], ids=["sin-bandera", "True"])
def test_textos_nuevos_en_la_maestra_individual_y_flota(bandera) -> None:
    for ws in (
        _maestra_individual(_individual(_VUELOS_COBROS, bandera=bandera)),
        _maestra_general(_general(_VUELOS_COBROS, bandera=bandera)),
    ):
        assert _columna(ws, 3) == _ESPERADO_COBROS
        for i in range(2):
            celda = ws.cell(row=FILA_DATOS + i, column=_col_factura(ws))
            assert celda.data_type == "s"
            assert celda.number_format == "General"


def test_textos_nuevos_en_otros_movimientos_col_j() -> None:
    req = _general(_VUELOS_COBROS, bandera=True)
    req.consolidado.otros_movimientos = BalanceHojaOtrosMovimientos(
        filas=[
            {
                "clave": "XA-VGV-330",
                "concepto_ingreso": "TUA",
                "ingreso_mxn": 500.0,
                "factura": "Pendiente de factura",
            },
            {
                "clave": "XA-VGV-331",
                "concepto_ingreso": "TUA",
                "ingreso_mxn": 500.0,
                "factura": "No lleva factura",
            },
        ],
        filas_sueltas=[],
    )
    wb = load_workbook(BytesIO(render_balance_general_xlsx(req)), data_only=True)
    ws = wb["otros movimientos"]
    celdas = {c.value: c for c in ws["J"] if c.value in _TEXTOS_NUEVOS}
    assert set(celdas) == set(_TEXTOS_NUEVOS)
    assert all(c.data_type == "s" for c in celdas.values())


def test_textos_nuevos_en_el_libro_dinero_hoja_1_y_otros_ingresos() -> None:
    req = DineroXlsxRequest(
        periodo_desde="2026-09-01",
        periodo_hasta="2026-09-30",
        vuelos=[
            {"clave": "vt330", "factura_vuelatour": "Pendiente de factura"},
            {"clave": "vt331", "factura_vuelatour": "No lleva factura"},
            {"clave": "vt332", "factura_vuelatour": None},
        ],
        otros_ingresos=[
            {
                "clave": "vt330",
                "concepto_ingreso": "tuas",
                "ingreso_mxn": 100.0,
                "factura": "Pendiente de factura",
            },
            {
                "clave": "vt331",
                "concepto_ingreso": "tuas",
                "ingreso_mxn": 100.0,
                "factura": "No lleva factura",
            },
        ],
    )
    wb = load_workbook(BytesIO(render_dinero_xlsx(req)), data_only=True)
    hoja1 = wb.worksheets[0]
    col = next(i for i, c in enumerate(_COLS_DINERO, start=1) if c[2] == "factura_vuelatour")
    letra = get_column_letter(col)
    celdas = {c.value: c for c in hoja1[letra] if c.value in _TEXTOS_NUEVOS}
    assert set(celdas) == set(_TEXTOS_NUEVOS)
    assert all(c.data_type == "s" for c in celdas.values())
    otros = wb["Otros ingresos"]
    celdas_j = {c.value: c for c in otros["J"] if c.value in _TEXTOS_NUEVOS}
    assert set(celdas_j) == set(_TEXTOS_NUEVOS)
    assert all(c.data_type == "s" for c in celdas_j.values())


# ---------------------------------------------------------------------------
# Esquema: ADITIVO y LIBERAL (jamás un 422 que tumbe el balance entero).
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("entrada", "esperado"),
    [
        ("A-0424", "A-0424"),
        ("  A-0424  ", "A-0424"),
        ("", None),
        ("   ", None),
        (None, None),
        (1234, "1234"),
        (1234.0, "1234"),
        (12.5, "12.5"),
        (float("nan"), None),
        (True, None),
        ({"serie": "A"}, None),
    ],
)
def test_esquema_normaliza_lo_que_llega(entrada, esperado) -> None:
    assert BalanceAvionVuelo(factura_vuelatour=entrada).factura_vuelatour == esperado


def test_esquema_sin_la_clave_es_none() -> None:
    assert BalanceAvionVuelo().factura_vuelatour is None


def test_esquema_bandera_segun_cobros_aditiva() -> None:
    assert BalanceAvionRequest().factura_vuelatour_segun_cobros is None
    assert (
        BalanceAvionRequest(factura_vuelatour_segun_cobros=True).factura_vuelatour_segun_cobros
        is True
    )
    assert (
        BalanceAvionRequest(factura_vuelatour_segun_cobros=False).factura_vuelatour_segun_cobros
        is False
    )
    assert BalanceAvionVuelo(factura_vuelatour="Pendiente de factura").factura_vuelatour == (
        "Pendiente de factura"
    )


# ---------------------------------------------------------------------------
# Rutas: el payload del API 0.0.45 entra sin 422 en los dos libros.
# ---------------------------------------------------------------------------

TOKEN = "secreto-de-prueba"
client = TestClient(app)


def test_rutas_aceptan_el_payload_nuevo(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    # Aunque una aserción falle, el token de prueba no se queda en caché.
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    # JSON crudo como lo manda NestJS (con un folio numérico de más).
    vuelos = [*_VUELOS, _vuelo("XA-VGV-305", factura_vuelatour=1234)]

    res = client.post(
        "/pdf/balance-avion-xlsx",
        json={
            "matricula": "XA-VGV",
            "periodo_desde": "2026-09-01",
            "periodo_hasta": "2026-09-30",
            "vuelos": vuelos,
            "totales": _TOTALES,
        },
        headers=headers,
    )
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content), data_only=True)["reporte horas XA-VGV"]
    assert _columna(ws, len(vuelos)) == [*_ESPERADO, "1234"]

    res = client.post(
        "/pdf/balance-general-xlsx",
        json={
            "periodo_desde": "2026-09-01",
            "periodo_hasta": "2026-09-30",
            "consolidado": {"matricula": "FLOTA", "vuelos": vuelos, "totales": _TOTALES},
        },
        headers=headers,
    )
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content), data_only=True)["reporte horas FLOTA"]
    assert _columna(ws, len(vuelos)) == [*_ESPERADO, "1234"]


def test_rutas_aceptan_la_bandera_del_api_0_0_74(monkeypatch, request) -> None:
    """JSON como lo manda el API 0.0.74: la bandera en la raíz del individual,
    en cada `aviones[]` y en el `consolidado` del general. Y el JSON del
    0.0.73 (sin la bandera) sigue entrando con la nota de siempre."""
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    vuelos = [
        {**v} for v in _VUELOS_COBROS
    ]

    def notas_de(ws) -> list:
        return [
            c.value
            for c in ws["A"]
            if isinstance(c.value, str) and c.value.startswith("FACTURA VUELATOUR =")
        ]

    casos = ((True, _NOTA_FACTURA_VUELATOUR_SEGUN_COBROS), (None, _NOTA_FACTURA_VUELATOUR))
    for bandera, esperada in casos:
        extra = {"factura_vuelatour_segun_cobros": True} if bandera else {}
        res = client.post(
            "/pdf/balance-avion-xlsx",
            json={
                "matricula": "XA-VGV",
                "periodo_desde": "2026-09-01",
                "periodo_hasta": "2026-09-30",
                "vuelos": vuelos,
                "totales": _TOTALES,
                "sin_provision_vendedor": True,
                **extra,
            },
            headers=headers,
        )
        assert res.status_code == 200, res.text
        ws = load_workbook(BytesIO(res.content), data_only=True)["reporte horas XA-VGV"]
        assert _columna(ws, len(vuelos)) == _ESPERADO_COBROS
        assert notas_de(ws) == [esperada]

        res = client.post(
            "/pdf/balance-general-xlsx",
            json={
                "periodo_desde": "2026-09-01",
                "periodo_hasta": "2026-09-30",
                "resumen": [{"matricula": "XA-VGV", "color": "#FF0000", "vuelos": 3}],
                "aviones": [
                    {
                        "matricula": "XA-VGV",
                        "periodo_desde": "2026-09-01",
                        "periodo_hasta": "2026-09-30",
                        "vuelos": vuelos,
                        "totales": _TOTALES,
                        "sin_provision_vendedor": True,
                        **extra,
                    }
                ],
                "consolidado": {
                    "matricula": "FLOTA",
                    "vuelos": vuelos,
                    "totales": _TOTALES,
                    **extra,
                },
            },
            headers=headers,
        )
        assert res.status_code == 200, res.text
        ws = load_workbook(BytesIO(res.content), data_only=True)["reporte horas FLOTA"]
        assert _columna(ws, len(vuelos)) == _ESPERADO_COBROS
        assert notas_de(ws) == [esperada]
