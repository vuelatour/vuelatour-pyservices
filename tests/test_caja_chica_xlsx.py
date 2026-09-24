"""Excel de la REPOSICIÓN de caja chica (24-sep-2026).

Pedido del cliente: «al momento de reembolsar la caja de cada uno, me puede
arrojar un Excel descargable con la información de lo que estoy
reembolsando». El API manda todo calculado; aquí se congela CÓMO se pinta:
encabezado legible (etiquetas en A:C, valor en D), tabla con fechas reales
dd/mm/aaaa y montos "$"#,##0.00, fila TOTAL, totales, avisos y resaltes.
"""

from datetime import date, datetime
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reportes import CajaChicaReposicionRequest
from app.services.caja_chica_xlsx import COLUMNAS, render_caja_chica_reposicion_xlsx

MONEY = '"$"#,##0.00'
AVISO = "1 gasto se capturó DESPUÉS de registrar esta reposición (resaltado en naranja)."
SIN_FILAS = "Sin gastos pendientes desde la última reposición"


def _req(**extra) -> CajaChicaReposicionRequest:
    base = {
        "titulo": "Reposición de caja chica · Itzi",
        "subtitulo": (
            "Reposición del 21/09/2026 por $1,070.00 MXN · "
            "generado 2026-09-24 13:54 (hora Cancún)"
        ),
        "hoja": "Reposición",
        "encabezado_titulo": "Datos de la reposición",
        "encabezado": [
            {"etiqueta": "Responsable", "valor": "Itzi"},
            {"etiqueta": "Monto del fondo", "valor": 1000},
            {"etiqueta": "Fecha de la reposición", "valor": "21/09/2026"},
            {"etiqueta": "Periodo cubierto", "valor": "del 06/08/2026 al 24/08/2026"},
        ],
        "filas": [
            {
                "fecha": "2026-08-06",
                "tipo": "Gasto en efectivo",
                "categoria": "Taxi / estacionamiento",
                "descripcion": "Taxi al FBO · MID",
                "vuelo": "#297",
                "matricula": "XA-VGV",
                "gasto": 300,
                "comprobante": "Con comprobante",
                "facturacion": "Facturada",
                "capturo": "Itzi",
                "capturado": "2026-08-06 09:10",
                "saldo": 700,
                "por_reponer": 300,
            },
            {
                "fecha": "2026-08-10",
                "tipo": "Ajuste",
                "descripcion": "faltante del conteo",
                "otro": -20,
                "capturo": "Mary Cruz",
                "capturado": "2026-08-10 10:00",
                "saldo": 680,
                "por_reponer": 320,
            },
            {
                "fecha": "2026-08-24",
                "tipo": "Gasto en efectivo",
                "categoria": "Alimentos",
                "gasto": 770,
                "comprobante": "Sin comprobante",
                "facturacion": "Pendiente",
                "capturado": "2026-09-22 18:00",
                "saldo": -90,
                "por_reponer": 1090,
                "resaltar": True,
            },
        ],
        "n_gastos": 2,
        "total_gastos": 1070,
        "total_otros": -20,
        "totales_titulo": "Totales de la reposición",
        "totales": [
            {"etiqueta": "Σ gastos del periodo", "valor": 1070},
            {"etiqueta": "Monto repuesto", "valor": 1070},
            {
                "etiqueta": "Diferencia (repuesto − por reponer)",
                "valor": -20,
                "nota": "Quedó pendiente por reponer",
            },
            {"etiqueta": "Saldo del libro después", "valor": 980, "destacado": True},
        ],
        "avisos": [AVISO],
    }
    base.update(extra)
    return CajaChicaReposicionRequest(**base)


def _hoja(req: CajaChicaReposicionRequest):
    return load_workbook(BytesIO(render_caja_chica_reposicion_xlsx(req))).active


def _fila_de(ws, texto: str, col: int = 1) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == texto:
            return r
    raise AssertionError(f"no está «{texto}» en la columna {col}")


def test_titulo_subtitulo_y_pestana() -> None:
    ws = _hoja(_req())
    assert ws.title == "Reposición"
    assert ws.cell(row=1, column=1).value == "Reposición de caja chica · Itzi"
    assert ws.cell(row=1, column=1).font.bold
    assert ws.cell(row=2, column=1).value.startswith("Reposición del 21/09/2026")


def test_encabezado_legible_etiqueta_en_a_valor_en_d() -> None:
    """La etiqueta ocupa A:C (≈ 54 caracteres): «Por reponer antes de esta
    reposición» ya no se corta como en el export genérico (columna de 12)."""
    ws = _hoja(_req())
    r = _fila_de(ws, "Monto del fondo")
    assert ws.cell(row=r, column=1).font.bold
    assert ws.cell(row=r, column=4).value == 1000
    assert ws.cell(row=r, column=4).number_format == MONEY
    assert any(str(m) == f"A{r}:C{r}" for m in ws.merged_cells.ranges)
    ancho_etiqueta = sum(ws.column_dimensions[c].width for c in "ABC")
    assert ancho_etiqueta >= len("Por reponer antes de esta reposición") + 4
    # Texto: se queda como texto, sin formato de moneda.
    r2 = _fila_de(ws, "Periodo cubierto")
    assert ws.cell(row=r2, column=4).value == "del 06/08/2026 al 24/08/2026"


def test_tabla_encabezados_fechas_reales_y_montos() -> None:
    ws = _hoja(_req())
    h = _fila_de(ws, "Fecha")
    assert [ws.cell(row=h, column=c).value for c in range(1, len(COLUMNAS) + 1)] == [
        c[0] for c in COLUMNAS
    ]
    primera = h + 1
    f = ws.cell(row=primera, column=1)
    # FECHA real de Excel (ordenable/filtrable) con formato dd/mm/aaaa.
    assert isinstance(f.value, (date, datetime))
    assert f.value.strftime("%Y-%m-%d") == "2026-08-06"
    assert f.number_format == "dd/mm/yyyy"
    assert ws.cell(row=primera, column=7).value == 300
    assert ws.cell(row=primera, column=7).number_format == MONEY
    assert ws.cell(row=primera, column=13).number_format == MONEY
    assert ws.cell(row=primera, column=5).value == "#297"
    # El ajuste va en SU columna, con signo, y la de gasto vacía.
    assert ws.cell(row=primera + 1, column=7).value is None
    assert ws.cell(row=primera + 1, column=8).value == -20


def test_resaltar_pinta_fecha_y_captura_en_naranja() -> None:
    ws = _hoja(_req())
    h = _fila_de(ws, "Fecha")
    marcada = h + 3
    for col in (1, 12):
        c = ws.cell(row=marcada, column=col)
        assert c.font.bold
        assert str(c.font.color.rgb).endswith("ED7D31")
    normal = ws.cell(row=h + 1, column=1)
    assert not normal.font.bold


def test_fila_total_totales_y_avisos() -> None:
    ws = _hoja(_req())
    r = _fila_de(ws, "TOTAL")
    assert ws.cell(row=r, column=2).value == "2 gastos"
    assert ws.cell(row=r, column=7).value == 1070
    assert ws.cell(row=r, column=8).value == -20
    assert ws.cell(row=r, column=7).number_format == MONEY
    d = _fila_de(ws, "Diferencia (repuesto − por reponer)")
    assert d > r
    assert ws.cell(row=d, column=4).value == -20
    assert ws.cell(row=d, column=5).value == "Quedó pendiente por reponer"
    dest = _fila_de(ws, "Saldo del libro después")
    assert ws.cell(row=dest, column=4).font.bold
    a = _fila_de(ws, AVISO)
    assert str(ws.cell(row=a, column=1).font.color.rgb).endswith("ED7D31")


def test_un_gasto_en_singular_y_sin_otros() -> None:
    ws = _hoja(_req(n_gastos=1, total_otros=None))
    r = _fila_de(ws, "TOTAL")
    assert ws.cell(row=r, column=2).value == "1 gasto"
    assert ws.cell(row=r, column=8).value is None


def test_sin_filas_dice_por_que() -> None:
    ws = _hoja(_req(filas=[], n_gastos=0, total_gastos=0, sin_filas=SIN_FILAS))
    assert _fila_de(ws, SIN_FILAS) > 0
    r = _fila_de(ws, "TOTAL")
    assert ws.cell(row=r, column=2).value == "0 gastos"


def test_payload_minimo_y_campos_desconocidos_no_truenan() -> None:
    """Skew tolerante: un API más nuevo puede mandar campos que aquí no
    existen, y uno viejo puede omitir casi todo."""
    req = CajaChicaReposicionRequest(
        **{"campo_del_futuro": 1, "filas": [{"fecha": "no-es-fecha", "otro_campo": 2}]}
    )
    ws = _hoja(req)
    h = _fila_de(ws, "Fecha")
    assert ws.cell(row=h + 1, column=1).value == "no-es-fecha"
    assert ws.title == "Reposición"


def test_nombre_de_hoja_saneado() -> None:
    ws = _hoja(_req(hoja="Reposición 21/09 [Itzi] con un nombre muy largo"))
    assert "/" not in ws.title and "[" not in ws.title
    assert len(ws.title) <= 31


def test_impresion_horizontal_a_una_hoja_de_ancho() -> None:
    ws = _hoja(_req())
    assert ws.page_setup.orientation == "landscape"
    assert ws.sheet_properties.pageSetUpPr.fitToPage


TOKEN = "secreto-de-prueba"
client = TestClient(app)


def test_endpoint_exige_token(monkeypatch) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    res = client.post(
        "/reportes/caja-chica-reposicion.xlsx", json=_req().model_dump(mode="json")
    )
    assert res.status_code == 401


def test_endpoint_devuelve_el_xlsx(monkeypatch) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    res = client.post(
        "/reportes/caja-chica-reposicion.xlsx",
        json=_req().model_dump(mode="json"),
        headers={"X-Internal-Token": TOKEN},
    )
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    ws = load_workbook(BytesIO(res.content)).active
    assert ws.cell(row=1, column=1).value == "Reposición de caja chica · Itzi"


def test_texto_que_empieza_con_igual_se_queda_como_texto() -> None:
    """Revisión 24-sep-2026: openpyxl vuelve FÓRMULA toda cadena que empieza
    con «=». Una nota de gasto «= 3 taxis» (la teclea el piloto) o una nota de
    la reposición así daban un libro que Excel abre con error. Se escriben
    como TEXTO literal, sin cambiar lo que dicen."""
    datos = _req().model_dump()
    datos["filas"][0]["descripcion"] = "= 3 taxis al FBO"
    datos["encabezado"].append({"etiqueta": "Notas / referencia", "valor": "=reembolso agosto"})
    datos["avisos"].append("=aviso raro")
    req = CajaChicaReposicionRequest(**datos)
    ws = load_workbook(BytesIO(render_caja_chica_reposicion_xlsx(req))).active
    celdas = [c for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]
    formulas = [c.coordinate for c in celdas if c.data_type == "f"]
    assert formulas == []
    textos = {c.value for c in celdas}
    assert {"= 3 taxis al FBO", "=reembolso agosto", "=aviso raro"} <= textos
