"""Formato condicional «es menor que 0» en las celdas de ganancia / utilidad /
remanente de los dos libros (8-oct-2026, pedido del cliente con la captura del
RESUMEN: la GANANCIA negativa del N58BT en rojo).

Es una REGLA de Excel (cfRule `lessThan 0`, relleno FFC7CE + letra 9C0006: el
preset «Relleno rojo claro con texto rojo oscuro»), no un relleno fijo: la
evalúa Excel con el número que la celda muestre, sobrevive a que la oficina
edite la fórmula y no toca la huella de layout (que no mira las reglas)."""

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from app.services.balance_avion_xlsx import (
    _EMPRESA_RESULTADO,
    render_balance_avion_xlsx,
)
from tests import test_balance_comisiones as tc
from tests import test_balance_comisiones_general as tg


def _celdas_rojas(ws) -> set[str]:
    """Celdas cubiertas por la regla, verificando de paso que CADA regla de la
    hoja es la de «menor que 0» con los colores del preset."""
    celdas: set[str] = set()
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            assert r.type == "cellIs" and r.operator == "lessThan"
            assert r.formula == ["0"]
            assert r.dxf.fill.fgColor.rgb.endswith("FFC7CE")
            assert r.dxf.fill.bgColor.rgb.endswith("FFC7CE")
            assert r.dxf.font.color.rgb.endswith("9C0006")
        for rango in str(cf.sqref).split():
            c1, r1, c2, r2 = range_boundaries(rango)
            for col in range(c1, c2 + 1):
                for fila in range(r1, r2 + 1):
                    celdas.add(f"{get_column_letter(col)}{fila}")
    return celdas


def _columna(ws, fila: int, encabezado: str) -> str:
    for c in ws[fila]:
        if c.value == encabezado:
            return c.column_letter
    raise AssertionError(f"sin encabezado {encabezado!r} en la fila {fila} de {ws.title!r}")


def _fila_totales(ws, desde: int = 1) -> int:
    for c in ws["A"]:
        if c.value == "TOTALES" and c.row >= desde:
            return c.row
    raise AssertionError(f"sin fila TOTALES en {ws.title!r}")


def _hoja_con(wb, fila: int, encabezado: str):
    for ws in wb.worksheets:
        if any(c.value == encabezado for c in ws[fila]):
            return ws
    raise AssertionError(f"ninguna hoja trae {encabezado!r} en la fila {fila}")


def _assert_maestra(ws, encabezados: tuple[str, ...], *, sin: str) -> None:
    rojas = _celdas_rojas(ws)
    tot = _fila_totales(ws, desde=3)
    assert tot > 3, "el fixture trae vuelos"
    for h in encabezados:
        letra = _columna(ws, 2, h)
        for fila in range(3, tot + 1):
            assert f"{letra}{fila}" in rojas, (h, fila)
    # Una columna cualquiera de dinero NO lleva la regla (solo ganancia/remanente).
    letra_sin = _columna(ws, 2, sin)
    assert f"{letra_sin}3" not in rojas


def _assert_balance(ws) -> None:
    rojas = _celdas_rojas(ws)
    utilidades = [
        c.row for c in ws["A"] if isinstance(c.value, str) and c.value.startswith("UTILIDAD")
    ]
    assert utilidades, "la hoja 'balance' trae la cascada"
    for fila in utilidades:
        assert f"B{fila}" in rojas, fila
    # Reparto a socios: las celdas MONTO USD (columna C, numéricas) también.
    socios = [
        c.row for c in ws["C"]
        if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
        or (isinstance(c.value, str) and c.value.startswith("="))
    ]
    assert socios, "la hoja 'balance' trae socios"
    assert all(f"C{fila}" in rojas for fila in socios)
    resultado = [c.row for c in ws["A"] if c.value == _EMPRESA_RESULTADO]
    for fila in resultado:
        assert f"B{fila}" in rojas


def test_libro_individual() -> None:
    wb = load_workbook(BytesIO(render_balance_avion_xlsx(tc._individual())))
    _assert_maestra(
        wb.worksheets[0],
        ("REMANENTE\nVENTA−COSTO MXN", "GANANCIA\nMXN", "GANANCIA\nUSD"),
        sin="COSTO TOTAL\nMXN",
    )
    _assert_balance(wb["balance"])


def test_balance_general_mensual() -> None:
    wb = load_workbook(BytesIO(tg._render(tg._payload("mensual"))))
    resumen = wb["RESUMEN flota"]
    rojas = _celdas_rojas(resumen)
    letra = _columna(resumen, 3, "GANANCIA\nMXN")
    tot = _fila_totales(resumen, desde=4)
    assert tot > 4
    for fila in range(4, tot + 1):
        assert f"{letra}{fila}" in rojas, fila
    # Las demás columnas del RESUMEN no llevan la regla.
    assert f"{_columna(resumen, 3, 'VENTA\nMXN')}4" not in rojas
    _assert_maestra(
        _hoja_con(wb, 2, "GANANCIA\nMXN"),
        ("REMANENTE\nVENTA−COSTO MXN", "GANANCIA\nMXN", "GANANCIA\nUSD"),
        sin="COSTO TOTAL\nMXN",
    )
    _assert_balance(wb["balance"])
    om = wb["otros movimientos"]
    rojas_om = _celdas_rojas(om)
    tot_om = _fila_totales(om, desde=4)
    assert all(f"I{fila}" in rojas_om for fila in range(4, tot_om + 1))
    assert "G4" not in rojas_om


def test_balance_general_variante_general() -> None:
    wb = load_workbook(BytesIO(tg._render(tg._payload("general"))))
    _assert_maestra(
        _hoja_con(wb, 2, "REMANENTE VENTA\nMENOS COMPRA (PESOS)"),
        ("REMANENTE VENTA\nMENOS COMPRA (PESOS)",),
        sin="COSTO TOTAL\nMXN",
    )
    _assert_balance(wb["balance"])
