"""Reserva de overhaul MARCADA en el balance (sin restar) + USD por gasto
(19-oct-2026, API 0.0.98).

Decisión del cliente: la reserva de overhaul NO se resta en el Balance por
avión ni en el Balance general; solo queda MARCADA, muy clara, para que se
entienda qué debe aportar el avión para el overhaul (el reparto a socios sigue
restándola). El API manda llaves ADITIVAS en el bloque `balance` de cada
avión (`reserva_overhaul_usd` es el DISCRIMINADOR) y aquí, DESPUÉS de los
socios de cada cascada (hoja USD del individual, cada bloque del general y la
hoja en pesos), va un bloque informativo con relleno ámbar: encabezado, «Lo
que debe aportar el avión», «Cómo se calcula: …» y una nota. La cascada
(filas, índices y fórmulas) es BYTE A BYTE la de siempre con o sin la llave.

Regla B del API (bandera raíz `usd_por_gasto`): el API calcula TODO dólar
gasto por gasto y aquí ninguna celda deriva un USD como fórmula de pesos ÷
T.C.: el TOTAL USD de las hojas de gastos y combustible y la cadena USD/IVA de
la hoja de vuelos van como VALOR del API con una nota en el encabezado.

Sin la llave y sin la bandera el libro es BYTE A BYTE el del HEAD 16d6e2f
(huellas `_HUELLAS_HEAD`, calculadas con el código de ese HEAD sobre los
MISMOS constructores de `tests/_pesos_payloads.py`, con y sin la bandera
`balance_en_pesos`).
"""

from __future__ import annotations

import copy
import re
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAvionBalanceBloque,
    BalanceAvionRequest,
    BalanceAvionVuelo,
    BalanceGeneralRequest,
)
from app.services import libro_json, xlsx_formulas
from app.services.balance_avion_xlsx import (
    _COLS_USD_POR_GASTO,
    _DISP_FLOTA,
    _DISP_GENERAL,
    _DISP_MENSUAL,
    _EMPRESA_PARTICIPACION,
    _FILA_RESERVA,
    _FILA_RESERVA_PESOS,
    _FILL_RESERVA,
    _NOTA_BALANCE_PESOS,
    _NOTA_BALANCE_PESOS_RESERVA,
    _NOTA_ENCABEZADO_GANANCIA_USD_POR_GASTO_FLOTA,
    _NOTA_ENCABEZADO_USD_POR_GASTO,
    _NOTA_RESERVA_OVERHAUL,
    _NOTA_RESERVA_OVERHAUL_INCOMPLETA,
    _NOTA_SIN_PESOS_RESERVA,
    _TITULO_RESERVA_OVERHAUL,
    _USD_INFO_COLOR,
    MONEY,
    _como_se_calcula_reserva,
    _formulas_cascada,
    libro_balance_avion,
    libro_balance_general,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from tests import _pesos_payloads as pp
from tests import test_balance_empresa_vuelatour as t_empresa
from tests import test_balance_en_pesos as tp
from tests import test_balance_formulas as tf
from tests import test_libro_json as tlj
from tests._evaluador_formulas import verificar_libro

r2 = tf.r2

MAESTRA_TST = "reporte horas XA-TST"
MAESTRA_DOS = "reporte horas XB-DOS"
MAESTRA_FLOTA = "reporte horas FLOTA"

# La reserva de los payloads: 18.20 h × 2 motores × 55 USD/hr (la tarifa del
# API es la Σ de los motores, 110 USD/hr) = 2,002.00 USD; en pesos al T.C.
# oficial del cierre (30-sep-2026, 18.0243) = 36,084.65 MXN.
HORAS, USD_HR, MOTORES, TC, TC_DIA = 18.2, 110.0, 2, 18.0243, "2026-09-30"
RESERVA_USD = 2002.0
RESERVA_MXN = 36084.65
LLAVES = (
    "reserva_overhaul_usd",
    "reserva_overhaul_mxn",
    "reserva_overhaul_horas_hr",
    "reserva_overhaul_usd_hr",
    "reserva_overhaul_motores",
    "reserva_overhaul_tc",
    "reserva_overhaul_tc_dia",
    "reserva_overhaul_incompleta",
)

# ---------------------------------------------------------------------------
# Textos EXACTOS del contrato.
# ---------------------------------------------------------------------------

TITULO = "RESERVA DE OVERHAUL · informativa (no se resta en este balance)"
FILA_USD = "Lo que debe aportar el avión USD"
FILA_MXN = "Lo que debe aportar el avión (MXN)"
COMO = "18.20 h voladas en el periodo × 2 motores × 55.00 USD/hr = 2,002.00 USD"
COMO_PESOS = COMO + " · en pesos al T.C. oficial del 30-sep-2026 (18.0243)"
PREFIJO_COMO = "Cómo se calcula: "
NOTA = (
    "Es la aportación del periodo para el overhaul de los motores de este avión: "
    "horas voladas × tarifa por hora de cada motor. No se resta en la utilidad de "
    "este balance; el reparto a socios sí la descuenta antes de repartir. No es "
    "dinero guardado en una cuenta ni un acumulado."
)
SIN_TARIFA = (
    "Sin tarifa de reserva de overhaul configurada para este avión (ficha del "
    "avión → reserva de overhaul): no se puede calcular lo que debe aportar."
)
NOTA_USD_POR_GASTO = (
    "USD calculado gasto por gasto: cada gasto a dólares con su T.C. propio o el "
    "oficial del día en que se pagó (criterio del reparto a socios); no es el "
    "total en pesos ÷ T.C. promedio."
)

# Huellas «sin la llave y sin la bandera»: calculadas con el código del HEAD
# 16d6e2f corriendo `tests/_pesos_payloads.py` CON `balance_en_pesos` (True) y
# sin ella (None) — .xlsx: todos los miembros salvo docProps/core.xml; vista
# en pantalla: el JSON con `vinculos_gasto: true`. Las de `None` son las de
# `test_balance_en_pesos`. Si una se mueve, un API ≤ 0.0.97 ya no ve su libro.
# Huellas regeneradas el 23-oct-2026 (API 0.0.105, apodos fuera de los documentos; pedido
# del cliente: «dejar solo las columnas de Num de cotización, nombre completo, fecha, ruta y
# consecutivo del vuelo»): entran CLIENTE y Nº VUELO CLIENTE entre CLAVE y FECHA en la hoja
# maestra, la cobranza y 'dinero-vlos', y cliente / ruta / nº vuelo cliente en 'otros
# movimientos' y 'Otros ingresos'; todas las demás columnas corren. Mismos números y las
# mismas fórmulas (citadas por llave, no por letra): solo cambia el layout.
_HUELLAS_HEAD = {
    ("individual", True): (
        "63cb2b0705b92da37c170a30699fa868c4adea22474ee33ce9cf1ee66ca9d280",
        "1ded22c0e26d3d564cd775ba7f77e61fceb0df84c1887b31c70145a1f9c4a8c6",
    ),
    ("individual", None): tp._HUELLAS_HEAD["individual"],
    ("general-mensual", True): (
        "66265494e02c5ee3cbfd9af216fdcce84c8e9195edc984dbe36f638c3b897b8e",
        "69e13b5cd9c4fa2174f7f86a9a12f99d1563692ee30230c32d8252473cc078d0",
    ),
    ("general-mensual", None): tp._HUELLAS_HEAD["general-mensual"],
    ("general-general", True): (
        "35be7ddd056ce8188a61b0e1f869c0b3556006cbff0e504ebb5cdc779e8c5549",
        "15c7a42ba91be97d675ca5b40ee33e225631ae0ee0b46817d940dac9a11bc16e",
    ),
    ("general-general", None): tp._HUELLAS_HEAD["general-general"],
}


# ---------------------------------------------------------------------------
# Payloads.
# ---------------------------------------------------------------------------


def con_reserva(libro: dict, *, sin_tc: bool = False, incompleta: bool = False) -> dict:
    """El libro de un avión con las llaves de la reserva de overhaul del API
    0.0.98. NADA más cambia: la cascada no la resta. `sin_tc`: sin T.C.
    oficial del cierre ⇒ `reserva_overhaul_mxn` None. `incompleta`: voló
    horas pero el avión no tiene tarifa ⇒ 0.0."""
    libro = copy.deepcopy(libro)
    libro["balance"].update(
        reserva_overhaul_usd=0.0 if incompleta else RESERVA_USD,
        reserva_overhaul_mxn=None if sin_tc else (0.0 if incompleta else RESERVA_MXN),
        reserva_overhaul_horas_hr=HORAS,
        reserva_overhaul_usd_hr=0.0 if incompleta else USD_HR,
        reserva_overhaul_motores=0 if incompleta else MOTORES,
        reserva_overhaul_tc=None if sin_tc else TC,
        reserva_overhaul_tc_dia=None if sin_tc else TC_DIA,
        reserva_overhaul_incompleta=incompleta,
    )
    return libro


def individual(*, bandera: bool | None = True, **kw) -> dict:
    return con_reserva(pp.individual(bandera=bandera), **kw)


def general(variante: str = "mensual", *, bandera: bool | None = True, **kw) -> dict:
    p = pp.general(variante, bandera=bandera)
    p["aviones"] = [con_reserva(a, **kw) for a in p["aviones"]]
    return p


def _bloques(p: dict) -> list[dict]:
    return [a["balance"] for a in p["aviones"]] if "aviones" in p else [p["balance"]]


def con_usd_por_gasto(libro: dict) -> dict:
    """Un libro de avión (o el consolidado FLOTA) con `usd_por_gasto: true` y
    los dólares del lado de los costos PERTURBADOS como los mandaría la regla
    B: `usd ≠ mxn ÷ T.C.`, `siva ≠ usd ÷ 1.16`, IVA PAGADO MXN ≠ IVA USD ×
    T.C., TOTAL USD de cada hoja ≠ TOTAL MXN ÷ T.C. promedio. Las fórmulas de
    pesos que citan esas celdas (DIF. IVA, Σ de TOTALES, USD X HR = TOTAL USD
    ÷ horas) se mantienen consistentes."""
    p = copy.deepcopy(libro)
    p["usd_por_gasto"] = True
    for v in p["vuelos"]:
        if v.get("costo_usd") is None:
            continue
        v["costo_usd"] = r2(v["costo_usd"] + 0.37)
        v["costo_usd_siva"] = r2(v["costo_usd"] / 1.16 + 0.11)
        v["iva_pagado_usd"] = r2(v["costo_usd"] - v["costo_usd_siva"] + 0.05)
        if v.get("iva_pagado_mxn") is not None:
            v["iva_pagado_mxn"] = r2(v["iva_pagado_mxn"] + 5.55)
            if v.get("iva_mxn") is not None:
                v["dif_iva_mxn"] = r2(v["iva_mxn"] - v["iva_pagado_mxn"])
        if v.get("ganancia_usd") is not None:
            v["ganancia_usd"] = r2(v["ganancia_usd"] - 0.37)
        if v.get("costo_hr_usd") is not None:
            v["costo_hr_usd"] = r2(v["costo_hr_usd"] + 0.33)
            v["costo_hr_usd_siva"] = r2(v["costo_hr_usd"] / 1.16 + 0.07)
    t, vs = p["totales"], p["vuelos"]
    t["dif_iva_mxn"] = r2(
        tf._suma(v["dif_iva_mxn"] for v in vs if v.get("dif_iva_mxn") is not None)
    )
    t["ganancia_usd"] = r2(
        tf._suma(v["ganancia_usd"] for v in vs if v.get("ganancia_usd") is not None)
    )
    con_costo = [v["costo_hr_usd"] for v in vs if (v.get("costo_hr_usd") or 0) > 0]
    if con_costo:
        # El promedio del API sale de sus números completos: no es el de las
        # celdas redondeadas.
        t["costo_hr_prom_usd"] = r2(sum(con_costo) / len(con_costo) + 0.05)
    horas = t.get("tiempo_vuelo") or 0
    # `otros_gastos` no se perturba: la hoja fusionada «Gastos Indirectos» suma
    # los dos USD X HR ya redondeados y su fórmula ROUND(C5/D5,2) debe seguir
    # reproduciendo esa suma.
    for nombre in ("combustible", "gastos_indirectos", "refacciones", "permisos"):
        h = p.get(nombre)
        if h and h.get("usd"):
            h["usd"] = r2(h["usd"] + 1.23)
            h["usd_hr"] = xlsx_formulas.redondear_excel(h["usd"] / horas, 2) if horas else None
    otros = p.get("otros_gastos") or {}
    if horas and p.get("gastos_indirectos") and otros.get("usd"):
        fusion = xlsx_formulas.redondear_excel(
            (p["gastos_indirectos"]["usd"] + otros["usd"]) / horas, 2
        )
        p["gastos_indirectos"]["usd_hr"] = r2(fusion - (otros.get("usd_hr") or 0))
    if p.get("balance") and p.get("permisos") is not None and p.get("gastos_indirectos"):
        b = p["balance"]
        b["utilidad_antes_usd"] = t["ganancia_usd"]
        b["combustible_usd"] = p["combustible"]["usd"]
        b["gastos_indirectos_usd"] = p["gastos_indirectos"]["usd"]
        b["otros_usd"] = p["otros_gastos"]["usd"]
        b["permisos_usd"] = p["permisos"]["usd"]
        if p.get("refacciones") is not None:
            b["refacciones_usd"] = p["refacciones"]["usd"]
        indirectos = r2((b["gastos_indirectos_usd"] or 0) + (b["otros_usd"] or 0))
        b["utilidad_despues_usd"] = r2(
            b["utilidad_antes_usd"]
            - b["combustible_usd"]
            - indirectos
            - (b.get("refacciones_usd") or 0)
            - b["permisos_usd"]
            + (b.get("otros_ingresos_avion_usd") or 0)
        )
        b["utilidad_cobrada_usd"] = r2(b["utilidad_despues_usd"] - b["por_cobrar_usd"])
        for s in b["socios"]:
            s["monto_usd"] = r2(s["porcentaje"] / 100 * b["utilidad_cobrada_usd"])
    return p


def _con_comisiones(vuelos: list[dict]) -> list[dict]:
    """Las filas con comisiones del avión y ganancia en dólares (las que en la
    hoja del general van antes de comisiones)."""
    return [
        v
        for v in vuelos
        if (v.get("comisiones_mxn") or 0) != 0 and v.get("ganancia_usd") is not None
    ]


def general_usd_por_gasto(
    variante: str, *, en_consolidado: bool, antes_comisiones: str | None = None
) -> dict:
    """El general sin pesos con la regla B: la bandera en el consolidado (y
    sus dólares perturbados) o solo en los `aviones[]`. `antes_comisiones`:
    None = sin `ganancia_antes_comisiones_usd` (API sin la llave); 'todas' =
    cada fila con comisiones la trae (= ganancia + un número cualquiera: es
    un valor del API); 'una-null' = la primera en null."""
    p = pp.general(variante, bandera=None)
    cons = con_usd_por_gasto(p["consolidado"])
    if not en_consolidado:
        cons.pop("usd_por_gasto")
        for a in p["aviones"]:
            a["usd_por_gasto"] = True
    if antes_comisiones is not None:
        filas = _con_comisiones(cons["vuelos"])
        for v in filas:
            v["ganancia_antes_comisiones_usd"] = r2(v["ganancia_usd"] + 12.34)
        if antes_comisiones == "una-null":
            filas[0]["ganancia_antes_comisiones_usd"] = None
    p["consolidado"] = cons
    ge = p["gastos_empresa"]
    ge["usd"] = r2(ge["usd"] + 1.23)
    ge["usd_hr"] = xlsx_formulas.redondear_excel(ge["usd"] / cons["totales"]["tiempo_vuelo"], 2)
    p["empresa"] = t_empresa._empresa_api(p)
    return p


@pytest.fixture
def registros(monkeypatch) -> list[xlsx_formulas.RegistroFormulas]:
    """Registro de fórmulas de cada libro que se cierra (cuáles volvieron a
    valor)."""
    capturados: list[xlsx_formulas.RegistroFormulas] = []
    original = xlsx_formulas.finalizar

    def espia(wb):
        cache = original(wb)
        capturados.append(xlsx_formulas.registro(wb))
        return cache

    monkeypatch.setattr(xlsx_formulas, "finalizar", espia)
    return capturados


def _ultimo_socio(ws, socios: list[dict]) -> int:
    nombres = {s["nombre"] for s in socios}
    return max(c.row for c in ws["A"] if c.value in nombres)


def _celdas_hasta(ws, fila: int) -> list[tuple]:
    """Lo que se ve de la fila 1 a `fila`: coordenada, valor/fórmula, formato,
    negritas y combinadas que empiezan ahí."""
    merges = sorted(str(m) for m in ws.merged_cells.ranges if m.min_row <= fila)
    celdas = [
        (c.coordinate, c.value, c.number_format, c.font.b)
        for r in ws.iter_rows(min_row=1, max_row=fila)
        for c in r
        if c.value is not None
    ]
    return [*celdas, ("merges", merges)]


def _bloque_en(ws, inicio: int, *, pesos: bool) -> None:
    """El bloque informativo que empieza en `inicio` (la fila en blanco tras
    los socios): encabezado, monto, «Cómo se calcula» y nota, todo ámbar."""
    n = 4 if pesos else 3
    assert all(ws.cell(row=inicio, column=c).value is None for c in range(1, n + 1))
    h = ws.cell(row=inicio + 1, column=1)
    assert h.value == TITULO == _TITULO_RESERVA_OVERHAUL
    assert h.font.b
    merges = {str(m) for m in ws.merged_cells.ranges}
    ultima = "D" if pesos else "C"
    assert f"A{inicio + 1}:{ultima}{inicio + 1}" in merges
    assert f"A{inicio + 3}:{ultima}{inicio + 3}" in merges
    assert f"A{inicio + 4}:{ultima}{inicio + 4}" in merges
    # Relleno ámbar: en las filas combinadas lo lleva la celda ancla (al
    # releer, las demás son MergedCell sin estilo); en la del monto, todas.
    for r in (inicio + 1, inicio + 3, inicio + 4):
        assert ws.cell(row=r, column=1).fill.fgColor.rgb.endswith(_FILL_RESERVA), r
    for c in range(1, n + 1):
        assert ws.cell(row=inicio + 2, column=c).fill.fgColor.rgb.endswith(_FILL_RESERVA), c
    assert ws.cell(row=inicio + 2, column=1).value == (FILA_MXN if pesos else FILA_USD)
    assert ws.cell(row=inicio + 2, column=1).font.b
    assert ws.cell(row=inicio + 4, column=1).value == NOTA == _NOTA_RESERVA_OVERHAUL
    assert ws.cell(row=inicio + 4, column=1).font.i


# ---------------------------------------------------------------------------
# 1) Esquema aditivo y helpers puros.
# ---------------------------------------------------------------------------


def test_esquema_aditivo_con_defaults_none_y_liberal() -> None:
    b = BalanceAvionBalanceBloque()
    for llave in LLAVES:
        assert getattr(b, llave) is None
    b = BalanceAvionBalanceBloque.model_validate(
        {
            "reserva_overhaul_usd": 0,
            "reserva_overhaul_motores": 2.0,
            "reserva_overhaul_tc_dia": 20260930,
            "reserva_overhaul_horas_hr": "x",
            "reserva_overhaul_tc": float("nan"),
            "reserva_overhaul_incompleta": "si",
            "reserva_overhaul_xx": 1,
        }
    )
    assert b.reserva_overhaul_usd == 0.0 and b.reserva_overhaul_motores == 2
    assert b.reserva_overhaul_tc_dia is None and b.reserva_overhaul_horas_hr is None
    assert b.reserva_overhaul_tc is None and b.reserva_overhaul_incompleta is None

    def bloque(**llaves) -> BalanceAvionBalanceBloque:
        return BalanceAvionBalanceBloque.model_validate(llaves)

    assert bloque(reserva_overhaul_motores="dos").reserva_overhaul_motores is None
    assert bloque(reserva_overhaul_incompleta=True).reserva_overhaul_incompleta is True
    assert bloque(reserva_overhaul_incompleta=False).reserva_overhaul_incompleta is None
    assert bloque(reserva_overhaul_tc_dia=" 2026-09-30 ").reserva_overhaul_tc_dia == "2026-09-30"


@pytest.mark.parametrize("valor", [None, False, "true", 1, {}])
def test_usd_por_gasto_liberal_solo_true(valor) -> None:
    assert BalanceAvionRequest().usd_por_gasto is None
    assert BalanceAvionRequest.model_validate({"usd_por_gasto": valor}).usd_por_gasto is None
    assert BalanceAvionRequest.model_validate({"usd_por_gasto": True}).usd_por_gasto is True


@pytest.mark.parametrize(
    ("con_oi", "esperado"),
    [
        (False, ("ROUND(B4-B5-B6-B7-B8,2)", "ROUND(B9-B10,2)", 9, 11)),
        (True, ("ROUND(B4-B5-B6-B7-B8+B9,2)", "ROUND(B10-B11,2)", 10, 12)),
    ],
    ids=["hoy", "hoy+oi"],
)
def test_formulas_cascada_son_las_de_siempre(con_oi, esperado) -> None:
    assert _formulas_cascada("B", 4, con_oi=con_oi) == esperado
    assert _formulas_cascada("C", 4, con_oi=con_oi)[0] == esperado[0].replace("B", "C")


def test_como_se_calcula_reserva() -> None:
    con_tc = BalanceAvionBalanceBloque.model_validate(individual()["balance"])
    assert _como_se_calcula_reserva(con_tc) == COMO_PESOS
    sin_tc = BalanceAvionBalanceBloque.model_validate(individual(sin_tc=True)["balance"])
    assert _como_se_calcula_reserva(sin_tc) == COMO
    inc = BalanceAvionBalanceBloque.model_validate(individual(incompleta=True)["balance"])
    assert _como_se_calcula_reserva(inc) == SIN_TARIFA == _NOTA_RESERVA_OVERHAUL_INCOMPLETA
    # Un motor: singular y la tarifa entera; T.C. sin día: sin «del …».
    uno = con_tc.model_copy(
        update={
            "reserva_overhaul_motores": 1,
            "reserva_overhaul_tc": 17.123456,
            "reserva_overhaul_tc_dia": None,
        }
    )
    assert _como_se_calcula_reserva(uno) == (
        "18.20 h voladas en el periodo × 1 motor × 110.00 USD/hr = 2,002.00 USD "
        "· en pesos al T.C. oficial (17.123456)"
    )
    # Nada de «se resta» en los textos del bloque.
    for texto in (NOTA, SIN_TARIFA, TITULO):
        assert "se resta $" not in texto and "MISMA resta" not in texto
    assert (FILA_USD, FILA_MXN) == (_FILA_RESERVA, _FILA_RESERVA_PESOS)
    assert NOTA_USD_POR_GASTO == _NOTA_ENCABEZADO_USD_POR_GASTO


# ---------------------------------------------------------------------------
# 2) Libro INDIVIDUAL en USD: cascada idéntica + bloque después de los socios.
# ---------------------------------------------------------------------------


def test_individual_usd_cascada_identica_y_bloque_tras_los_socios(registros) -> None:
    p = individual(bandera=None)
    b = p["balance"]
    con = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    verificar_libro(con)
    sin = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(pp.individual(bandera=None)))
    ws, wsv = tp._libros(con)
    ws, wsv = ws["balance"], wsv["balance"]
    ws_sin = load_workbook(BytesIO(sin))["balance"]
    fin = _ultimo_socio(ws, b["socios"])
    assert fin == _ultimo_socio(ws_sin, b["socios"])
    # Hasta el último socio, celda por celda lo mismo (valores, fórmulas,
    # formatos, negritas y combinadas): nada se resta ni se corre.
    assert _celdas_hasta(ws, fin) == _celdas_hasta(ws_sin, fin)
    assert TITULO not in [c.value for c in ws_sin["A"]]
    _bloque_en(ws, fin + 1, pesos=False)
    monto = ws.cell(row=fin + 3, column=2)
    assert monto.value == RESERVA_USD and monto.number_format == MONEY and monto.font.b
    assert ws.cell(row=fin + 4, column=1).value == PREFIJO_COMO + COMO_PESOS
    # La utilidad es la de siempre: la reserva no la toca.
    cobrada = tp._fila_de(ws, "UTILIDAD COBRADA USD")
    assert wsv.cell(row=cobrada, column=2).value == b["utilidad_cobrada_usd"]
    sin_llave = pp.individual(bandera=None)["balance"]
    assert b["utilidad_cobrada_usd"] == sin_llave["utilidad_cobrada_usd"]


def test_individual_usd_sin_otros_ingresos_tambien(registros) -> None:
    p = con_reserva(tf._payload_tst())
    con = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    sin = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(tf._payload_tst()))
    ws = load_workbook(BytesIO(con))["balance"]
    ws_sin = load_workbook(BytesIO(sin))["balance"]
    fin = _ultimo_socio(ws, p["balance"]["socios"])
    assert _celdas_hasta(ws, fin) == _celdas_hasta(ws_sin, fin)
    assert ws["B9"].value == "=ROUND(B4-B5-B6-B7-B8,2)"
    _bloque_en(ws, fin + 1, pesos=False)
    assert ws.cell(row=fin + 3, column=2).value == RESERVA_USD


# ---------------------------------------------------------------------------
# 3) Libro INDIVIDUAL en pesos (bandera `balance_en_pesos`).
# ---------------------------------------------------------------------------


def test_individual_pesos_cascada_identica_y_bloque(registros) -> None:
    p = individual()
    b = p["balance"]
    con = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    verificar_libro(con)
    sin = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(pp.individual()))
    ws, wsv = tp._libros(con)
    ws, wsv = ws["balance"], wsv["balance"]
    ws_sin = load_workbook(BytesIO(sin))["balance"]
    fin = _ultimo_socio(ws, b["socios"])
    assert _celdas_hasta(ws, fin) == _celdas_hasta(ws_sin, fin)
    _bloque_en(ws, fin + 1, pesos=True)
    mxn = ws.cell(row=fin + 3, column=2)
    assert mxn.value == RESERVA_MXN and mxn.number_format == MONEY and mxn.font.b
    usd = ws.cell(row=fin + 3, column=3)
    assert usd.value == RESERVA_USD
    assert usd.font.italic and usd.font.size == 9 and usd.font.color.rgb.endswith(_USD_INFO_COLOR)
    assert ws.cell(row=fin + 4, column=1).value == PREFIJO_COMO + COMO_PESOS
    # El pie explica la columna USD por gasto y que la reserva es informativa.
    textos = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert textos[-1] == _NOTA_BALANCE_PESOS_RESERVA
    assert _NOTA_BALANCE_PESOS not in textos
    assert "bloque informativo (no se resta en este balance)" in _NOTA_BALANCE_PESOS_RESERVA
    assert "con su T.C. propio" in _NOTA_BALANCE_PESOS_RESERVA
    assert wsv.cell(row=tp._fila_de(ws, "UTILIDAD COBRADA (MXN)"), column=2).value == (
        b["utilidad_cobrada_mxn"]
    )


def test_individual_pesos_sin_tc_raya_con_nota_y_usd_pintado() -> None:
    p = individual(sin_tc=True)
    b = p["balance"]
    assert b["reserva_overhaul_mxn"] is None and b["utilidad_despues_mxn"] is not None
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    wb, wv = tp._libros(data)
    ws = wb["balance"]
    r = tp._fila_de(ws, FILA_MXN)
    assert ws.cell(row=r, column=2).value == "—"
    assert ws.cell(row=r, column=2).comment.text == _NOTA_SIN_PESOS_RESERVA
    assert wv["balance"].cell(row=r, column=3).value == RESERVA_USD
    assert ws.cell(row=r + 1, column=1).value == PREFIJO_COMO + COMO  # sin la cola del T.C.
    # La cascada en pesos no se entera: sigue con sus números.
    for etiqueta, llave in (
        ("UTILIDAD DESPUÉS DE GASTOS (MXN)", "utilidad_despues_mxn"),
        ("UTILIDAD COBRADA (MXN)", "utilidad_cobrada_mxn"),
    ):
        assert wv["balance"].cell(row=tp._fila_de(ws, etiqueta), column=2).value == b[llave]


@pytest.mark.parametrize("bandera", [None, True], ids=["usd", "pesos"])
def test_incompleta_cero_y_sin_tarifa(bandera) -> None:
    p = individual(bandera=bandera, incompleta=True)
    assert p["balance"]["reserva_overhaul_usd"] == 0.0
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    wb, wv = tp._libros(data)
    ws = wb["balance"]
    r = tp._fila_de(ws, FILA_MXN if bandera else FILA_USD)
    assert ws.cell(row=r, column=2).value == 0 and wv["balance"].cell(row=r, column=2).value == 0
    if bandera:
        assert wv["balance"].cell(row=r, column=3).value == 0
    assert ws.cell(row=r + 1, column=1).value == PREFIJO_COMO + SIN_TARIFA
    assert ws.cell(row=r + 2, column=1).value == NOTA


# ---------------------------------------------------------------------------
# 4) Balance GENERAL: cada bloque de avión con el bloque informativo después
#    de sus socios; cascada, socios y bloque VUELATOUR intactos.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("variante", ["mensual", "general"])
@pytest.mark.parametrize("bandera", [None, True], ids=["usd", "pesos"])
def test_general_cada_bloque_lleva_el_informativo(registros, variante, bandera) -> None:
    p = general(variante, bandera=bandera)
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    verificar_libro(data)
    sin = render_balance_general_xlsx(
        BalanceGeneralRequest.model_validate(pp.general(variante, bandera=bandera))
    )
    wb, wv = tp._libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    ws_sin = load_workbook(BytesIO(sin))["balance"]
    assert TITULO not in [c.value for c in ws_sin["A"]]
    pesos = bool(bandera)
    fila_cobrada = "UTILIDAD COBRADA (MXN)" if pesos else "UTILIDAD COBRADA USD"
    cobradas = tp._filas_de(ws, fila_cobrada)
    # El primer bloque (nada se corre antes de él) es idéntico celda por
    # celda hasta su último socio (COBRADA, blanco, título, encabezado, socios).
    fin1 = cobradas[0] + 3 + len(p["aviones"][0]["balance"]["socios"])
    assert _celdas_hasta(ws, fin1) == _celdas_hasta(ws_sin, fin1)
    titulos = [c.row for c in ws["A"] if c.value == TITULO]
    assert len(titulos) == len(p["aviones"]) == 2
    for titulo, cob, avion in zip(titulos, cobradas, p["aviones"], strict=True):
        b = avion["balance"]
        # Socios justo antes del bloque, citando la COBRADA de su cascada.
        fin = cob + 3 + len(b["socios"])
        assert titulo == fin + 2
        _bloque_en(ws, fin + 1, pesos=pesos)
        for i, s in enumerate(b["socios"]):
            rs = fin - len(b["socios"]) + 1 + i
            assert ws.cell(row=rs, column=1).value == s["nombre"]
            assert ws.cell(row=rs, column=3).value == f"=ROUND(B{rs}/100*$B${cob},2)"
            monto = s["monto_mxn"] if pesos else s["monto_usd"]
            assert wsv.cell(row=rs, column=3).value == monto
        assert wsv.cell(row=cob, column=2).value == (
            b["utilidad_cobrada_mxn"] if pesos else b["utilidad_cobrada_usd"]
        )
        assert ws.cell(row=fin + 3, column=2).value == (RESERVA_MXN if pesos else RESERVA_USD)
        assert ws.cell(row=fin + 4, column=1).value == PREFIJO_COMO + COMO_PESOS
    # Bloque VUELATOUR: PARTICIPACIÓN cita las celdas de los socios empresa.
    emp = p["empresa"]
    part = tp._fila_de(
        ws,
        "(+) PARTICIPACIÓN COMO SOCIO EN LOS AVIONES (MXN)" if pesos else _EMPRESA_PARTICIPACION,
    )
    citadas = [int(x) for x in re.findall(r"C(\d+)", ws.cell(row=part, column=2).value)]
    assert len(citadas) == 2
    assert all(ws.cell(row=c, column=1).value == t_empresa._EMPRESA for c in citadas)
    assert wsv.cell(row=part, column=2).value == (
        emp["participacion_mxn"] if pesos else emp["participacion_usd"]
    )
    res = tp._fila_de(ws, "RESULTADO VUELATOUR (MXN)" if pesos else "RESULTADO VUELATOUR USD")
    assert wsv.cell(row=res, column=2).value == (
        emp["resultado_mxn"] if pesos else emp["resultado_usd"]
    )
    if pesos:
        textos = [c.value for c in ws["A"] if isinstance(c.value, str)]
        assert textos[-1] == _NOTA_BALANCE_PESOS_RESERVA


def test_general_mixto_solo_el_avion_con_la_llave_lleva_el_bloque() -> None:
    p = pp.general("mensual", bandera=None)
    p["aviones"][0] = con_reserva(p["aviones"][0])
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    ws = load_workbook(BytesIO(data))["balance"]
    assert len(tp._filas_de(ws, TITULO)) == 1
    titulo = tp._fila_de(ws, TITULO)
    cob = tp._filas_de(ws, "UTILIDAD COBRADA USD")[0]
    assert titulo == cob + 3 + len(p["aviones"][0]["balance"]["socios"]) + 2
    verificar_libro(data)


# ---------------------------------------------------------------------------
# 5) Sin la llave y sin la bandera: BYTE A BYTE (.xlsx y vista en pantalla).
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(("libro", "bandera"), list(_HUELLAS_HEAD))
def test_sin_la_llave_libro_de_hoy_byte_a_byte(libro, bandera) -> None:
    data, vista = tp._render(libro, tp._payload(libro, bandera))
    assert (tp._huella_xlsx(data), tp._huella_json(vista)) == _HUELLAS_HEAD[(libro, bandera)]


@pytest.mark.parametrize("forma", ["null", "solo-metadatos", "usd_por_gasto-apagada"])
@pytest.mark.parametrize(("libro", "bandera"), list(_HUELLAS_HEAD))
def test_llaves_sin_discriminador_tambien_es_el_libro_de_hoy(libro, bandera, forma) -> None:
    """`reserva_overhaul_usd` es el discriminador del bloque y solo un `true`
    enciende `usd_por_gasto`: en null, con solo los metadatos o con la
    bandera en false/texto, no hay bloque, ni nota, ni valor que cambie."""
    p = tp._payload(libro, bandera)
    libros = p["aviones"] + [p["consolidado"]] if "aviones" in p else [p]
    for b in _bloques(p):
        if forma == "null":
            b.update(dict.fromkeys(LLAVES))
        elif forma == "solo-metadatos":
            b.update(con_reserva({"balance": b})["balance"])
            b["reserva_overhaul_usd"] = None
    if forma == "usd_por_gasto-apagada":
        for i, lib in enumerate(libros):
            lib["usd_por_gasto"] = [False, None, "true"][i % 3]
    data, vista = tp._render(libro, p)
    assert (tp._huella_xlsx(data), tp._huella_json(vista)) == _HUELLAS_HEAD[(libro, bandera)]


# ---------------------------------------------------------------------------
# 6) Vista en pantalla (el MISMO libro) y rutas.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("bandera", [None, True], ids=["usd", "pesos"])
def test_libro_json_exporta_el_bloque(bandera) -> None:
    req = BalanceAvionRequest.model_validate(individual(bandera=bandera))
    libro = libro_balance_avion(req)
    assert tlj._comparar_con_xlsx(libro, render_balance_avion_xlsx(req)) > 0
    hoja = next(h for h in libro["hojas"] if h["nombre"] == "balance")
    celdas = {(c["r"], c["c"]): c for c in hoja["celdas"]}
    r = next(r for (r, c), x in celdas.items() if c == 1 and x["v"] == TITULO)
    assert celdas[(r + 1, 1)]["v"] == (FILA_MXN if bandera else FILA_USD)
    monto = celdas[(r + 1, 2)]
    assert monto.get("f") is None  # valor, no fórmula
    assert monto["v"] == (RESERVA_MXN if bandera else RESERVA_USD)
    assert monto["t"] == libro_json.formatear(monto["v"], MONEY)
    assert celdas[(r + 2, 1)]["v"] == PREFIJO_COMO + COMO_PESOS
    assert celdas[(r + 3, 1)]["v"] == NOTA
    if bandera:
        assert celdas[(r + 1, 3)]["v"] == RESERVA_USD


TOKEN = "tok-reserva-overhaul"
client = TestClient(app)


def test_rutas_aceptan_el_payload_0098(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    p = con_usd_por_gasto(individual())
    res = client.post("/pdf/balance-avion-xlsx", json=p, headers=headers)
    assert res.status_code == 200, res.text
    wb = load_workbook(BytesIO(res.content))
    assert tp._filas_de(wb["balance"], TITULO)
    assert wb["combustible"]["E5"].value == p["combustible"]["usd"]
    res = client.post(
        "/pdf/balance-general-xlsx", json=general("general", bandera=None), headers=headers
    )
    assert res.status_code == 200, res.text
    assert len(tp._filas_de(load_workbook(BytesIO(res.content))["balance"], TITULO)) == 2
    res = client.post("/pdf/balance-general-libro", json=general("mensual"), headers=headers)
    assert res.status_code == 200, res.text
    hoja = next(h for h in res.json()["hojas"] if h["nombre"] == "balance")
    assert sum(1 for c in hoja["celdas"] if c.get("t") == TITULO) == 2
    assert libro_balance_general(BalanceGeneralRequest.model_validate(general("mensual")))


# ---------------------------------------------------------------------------
# 7) `usd_por_gasto` (regla B del API 0.0.98): ningún dólar como fórmula de
#    pesos ÷ T.C.; valores del API con nota y 0 fórmulas degradadas.
# ---------------------------------------------------------------------------

# Columnas de la hoja de vuelos que pasan a VALOR con la bandera.
COLUMNAS_VALOR = {
    "costo_usd",
    "costo_usd_siva",
    "iva_pagado_usd",
    "iva_pagado_mxn",
    "ganancia_usd",
    "costo_hr_usd",
    "costo_hr_usd_siva",
    "cph_iva_hr_usd",
}


def _es_formula(x) -> bool:
    return isinstance(x, str) and x.startswith("=")


def _columnas_valor(ws, disp, vuelos: list[dict]) -> None:
    """Cada columna de `COLUMNAS_VALOR` del juego `disp`: filas como VALOR del
    API (jamás fórmula) y la nota en su encabezado. En una hoja sin COMISIONES
    (el general) la GANANCIA USD de un vuelo con comisiones va VACÍA."""
    assert set(_COLS_USD_POR_GASTO) == COLUMNAS_VALOR
    for attr in sorted(COLUMNAS_VALOR):
        letra = disp.letra.get(attr)
        if letra is None:
            continue
        assert ws[f"{letra}2"].comment is not None, attr
        assert ws[f"{letra}2"].comment.text.startswith(NOTA_USD_POR_GASTO), attr
        for i, v in enumerate(vuelos):
            celda = ws[f"{letra}{3 + i}"]
            assert not _es_formula(celda.value), (attr, celda.coordinate)
            esperado = v.get(attr)
            if attr == "cph_iva_hr_usd":
                con_hr = v.get("costo_hr_usd") is not None
                esperado = r2(v["costo_hr_usd"] - v["costo_hr_usd_siva"]) if con_hr else None
            if attr == "ganancia_usd" and disp.comision_col is None and v.get("comisiones_mxn"):
                esperado = v.get("ganancia_antes_comisiones_usd")
            assert celda.value == esperado, (attr, celda.coordinate)


def test_usd_por_gasto_individual_valores_con_nota_y_cero_degradadas(registros) -> None:
    """XB-DOS (libro sin degradaciones de base) con los dólares perturbados:
    con la bandera NINGUNA fórmula vuelve a valor."""
    p = con_usd_por_gasto(pp.general("mensual", bandera=None)["aviones"][1])
    assert p["matricula"] == "XB-DOS"
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert registros[-1].degradadas() == []
    verificar_libro(data)
    wb = load_workbook(BytesIO(data))
    ws = wb[MAESTRA_DOS]
    fila_tot = 3 + len(p["vuelos"])
    _columnas_valor(ws, _DISP_MENSUAL, p["vuelos"])
    # TOTALES: la Σ de GANANCIA USD sigue siendo fórmula (el API suma filas
    # redondeadas); el PROMEDIO de COSTO X HORA USD es valor (el API promedia
    # sus números completos).
    letra_g, letra_c = _DISP_MENSUAL.letra["ganancia_usd"], _DISP_MENSUAL.letra["costo_hr_usd"]
    assert ws[f"{letra_g}{fila_tot}"].value == f"=ROUND(SUM({letra_g}3:{letra_g}{fila_tot - 1}),2)"
    assert ws[f"{letra_c}{fila_tot}"].value == p["totales"]["costo_hr_prom_usd"]
    # Las de pesos y las de la VENTA siguen como fórmula.
    pesos_y_venta = (
        "costo_total_mxn",
        "remanente_mxn",
        "dif_iva_mxn",
        "ganancia_mxn",
        "por_cobrar_usd",
    )
    for attr in pesos_y_venta:
        letra = _DISP_MENSUAL.letra[attr]
        assert _es_formula(ws[f"{letra}3"].value), attr
    # TOTAL USD de las hojas: valor del API con la nota; USD X HR sigue siendo
    # TOTAL USD ÷ horas; «Otros ingresos» no cambia (ya va por ingreso).
    hojas = (
        ("Gastos Indirectos", "C"),
        ("refacciones", "C"),
        ("permisos", "C"),
        ("combustible", "E"),
    )
    for hoja, col in hojas:
        w = wb[hoja]
        assert not _es_formula(w[f"{col}5"].value), hoja
        assert w[f"{col}4"].comment.text == NOTA_USD_POR_GASTO, hoja
    assert wb["Gastos Indirectos"]["C5"].value == r2(
        p["gastos_indirectos"]["usd"] + p["otros_gastos"]["usd"]
    )
    assert wb["Gastos Indirectos"]["E5"].value == "=ROUND(C5/D5,2)"
    assert wb["combustible"]["E5"].value == p["combustible"]["usd"]
    assert wb["Otros ingresos"]["C5"].value == p["otros_ingresos_avion"]["usd"] == 0
    assert wb["Otros ingresos"]["C4"].comment is None
    # La cascada sigue citando los TOTAL USD (ahora valores) y cuadra.
    bal = wb["balance"]
    assert bal["B5"].value == "='combustible'!$E$5"
    wv = load_workbook(BytesIO(data), data_only=True)["balance"]
    assert wv.cell(row=tp._fila_de(bal, "UTILIDAD COBRADA USD"), column=2).value == (
        p["balance"]["utilidad_cobrada_usd"]
    )


def test_usd_por_gasto_individual_xa_tst_solo_la_degradacion_de_siempre(registros) -> None:
    """XA-TST: la única fórmula que vuelve a valor es W7 (COSTO TOTAL MXN del
    #405, redondeo intermedio del espejo del API, en PESOS: nada que ver con
    la regla B); 'Gastos Indirectos'!C5, que hoy se degrada, ya es valor."""
    p = con_usd_por_gasto(pp.individual(bandera=None))
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert registros[-1].degradadas() == [(MAESTRA_TST, "W7")]
    verificar_libro(data)
    wb = load_workbook(BytesIO(data))
    _columnas_valor(wb[MAESTRA_TST], _DISP_MENSUAL, p["vuelos"])
    # «Otros ingresos» (dos ingresos, por ingreso) conserva su fórmula: no es
    # un dólar derivado de pesos ÷ T.C. promedio.
    assert wb["Otros ingresos"]["C5"].value == "=ROUND(SUM(K8:K9),2)"
    assert wb["Otros ingresos"]["C4"].comment is None


def test_usd_por_gasto_general_variante_general_cero_degradadas(registros) -> None:
    p = general_usd_por_gasto("general", en_consolidado=True)
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert registros[-1].degradadas() == []
    verificar_libro(data)
    wb = load_workbook(BytesIO(data))
    ws = wb[MAESTRA_FLOTA]
    cons = p["consolidado"]
    _columnas_valor(ws, _DISP_GENERAL, cons["vuelos"])
    # TOTAL PAGADO S/IVA (PESOS − IVA PESOS) y REMANENTE siguen como fórmula.
    for attr in ("cph_pagado_siva_mxn", "cph_remanente_mxn", "cph_cobrado_siva_mxn"):
        assert _es_formula(ws[f"{_DISP_GENERAL.letra[attr]}3"].value), attr
    for hoja, col, usd in (
        ("otros gastos", "C", p["gastos_empresa"]["usd"]),
        ("repartidos a aviones", "C", cons["otros_gastos"]["usd"]),
        ("combustible", "E", cons["combustible"]["usd"]),
    ):
        w = wb[hoja]
        assert w[f"{col}5"].value == usd and not _es_formula(w[f"{col}5"].value), hoja
        assert w[f"{col}4"].comment.text == NOTA_USD_POR_GASTO, hoja
    # El bloque VUELATOUR sigue citando 'otros gastos'!$C$5 y cuadra.
    bal = wb["balance"]
    otros = tp._fila_de(bal, "(−) OTROS GASTOS DE LA EMPRESA USD")
    assert bal.cell(row=otros, column=2).value == "='otros gastos'!$C$5"


def test_usd_por_gasto_general_mensual_ganancia_usd_vacia_con_comisiones(registros) -> None:
    """Hoja de vuelos del general (antes de comisiones) con la bandera solo
    en los `aviones[]`: GANANCIA USD ya no se deriva de pesos ÷ T.C. — en los
    vuelos con comisiones queda VACÍA y el encabezado lo dice; el resto, el
    valor del API."""
    p = general_usd_por_gasto("mensual", en_consolidado=False)
    assert "usd_por_gasto" not in p["consolidado"]
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert registros[-1].degradadas() == [(MAESTRA_FLOTA, "W9")]  # la de siempre (pesos)
    verificar_libro(data)
    ws = load_workbook(BytesIO(data))[MAESTRA_FLOTA]
    vuelos = p["consolidado"]["vuelos"]
    _columnas_valor(ws, _DISP_FLOTA, vuelos)
    letra = _DISP_FLOTA.letra["ganancia_usd"]
    assert ws[f"{letra}2"].comment.text == _NOTA_ENCABEZADO_GANANCIA_USD_POR_GASTO_FLOTA
    assert "queda vacía" in _NOTA_ENCABEZADO_GANANCIA_USD_POR_GASTO_FLOTA
    con_comisiones = [v for v in vuelos if (v.get("comisiones_mxn") or 0) != 0]
    assert con_comisiones
    for i, v in enumerate(vuelos):
        celda = ws[f"{letra}{3 + i}"]
        if (v.get("comisiones_mxn") or 0) != 0:
            assert celda.value is None, celda.coordinate
        elif v.get("ganancia_usd") is not None:
            assert celda.value == v["ganancia_usd"], celda.coordinate
    assert ws[f"{letra}{3 + len(vuelos)}"].value == (
        f"=ROUND(SUM({letra}3:{letra}{2 + len(vuelos)}),2)"
    )
    assert ws[f"{_DISP_FLOTA.letra['costo_hr_usd']}{3 + len(vuelos)}"].value == (
        p["consolidado"]["totales"]["costo_hr_prom_usd"]
    )


def test_ganancia_antes_comisiones_usd_es_aditiva() -> None:
    assert BalanceAvionVuelo().ganancia_antes_comisiones_usd is None
    fila = BalanceAvionVuelo.model_validate({"ganancia_antes_comisiones_usd": 12.5})
    assert fila.ganancia_antes_comisiones_usd == 12.5
    fila = BalanceAvionVuelo.model_validate({"ganancia_antes_comisiones_usd": None})
    assert fila.ganancia_antes_comisiones_usd is None


def test_usd_por_gasto_general_mensual_ganancia_antes_comisiones_pintada(registros) -> None:
    """Con `ganancia_antes_comisiones_usd` en cada fila con comisiones: se
    pinta como VALOR, el encabezado lleva la nota genérica (ninguna celda
    quedó vacía) y TOTALES = Σ de las celdas visibles, con fórmula viva."""
    p = general_usd_por_gasto("mensual", en_consolidado=True, antes_comisiones="todas")
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert registros[-1].degradadas() == [(MAESTRA_FLOTA, "W9")]
    verificar_libro(data)
    wb, wv = tp._libros(data)
    ws = wb[MAESTRA_FLOTA]
    vuelos = p["consolidado"]["vuelos"]
    _columnas_valor(ws, _DISP_FLOTA, vuelos)
    letra = _DISP_FLOTA.letra["ganancia_usd"]
    assert ws[f"{letra}2"].comment.text == NOTA_USD_POR_GASTO
    filas = _con_comisiones(vuelos)
    assert filas
    visibles = []
    for i, v in enumerate(vuelos):
        celda = ws[f"{letra}{3 + i}"]
        if v in filas:
            assert celda.value == v["ganancia_antes_comisiones_usd"]
            assert celda.value == r2(v["ganancia_usd"] + 12.34)
        if celda.value is not None:
            visibles.append(celda.value)
    fila_tot = 3 + len(vuelos)
    assert ws[f"{letra}{fila_tot}"].value == f"=ROUND(SUM({letra}3:{letra}{fila_tot - 1}),2)"
    assert wv[MAESTRA_FLOTA][f"{letra}{fila_tot}"].value == xlsx_formulas.redondear_excel(
        sum(visibles), 2
    )


def test_usd_por_gasto_general_mensual_una_fila_sin_la_llave_vacia_y_nota(registros) -> None:
    """La llave en null en UNA fila con comisiones: solo esa celda queda
    vacía y el encabezado lo explica; las demás, el valor del API."""
    p = general_usd_por_gasto("mensual", en_consolidado=True, antes_comisiones="una-null")
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert registros[-1].degradadas() == [(MAESTRA_FLOTA, "W9")]
    verificar_libro(data)
    ws = load_workbook(BytesIO(data))[MAESTRA_FLOTA]
    vuelos = p["consolidado"]["vuelos"]
    _columnas_valor(ws, _DISP_FLOTA, vuelos)
    letra = _DISP_FLOTA.letra["ganancia_usd"]
    assert ws[f"{letra}2"].comment.text == _NOTA_ENCABEZADO_GANANCIA_USD_POR_GASTO_FLOTA
    filas = _con_comisiones(vuelos)
    vacias = [ws[f"{letra}{3 + vuelos.index(v)}"].value for v in filas]
    assert vacias[0] is None and all(x is not None for x in vacias[1:])


def test_ganancia_antes_comisiones_usd_sin_la_bandera_se_ignora() -> None:
    """Sin `usd_por_gasto` la llave no mueve nada: la hoja del general deriva
    la GANANCIA USD como hoy (fórmula) y el libro es byte a byte el de
    siempre."""
    p = pp.general("mensual", bandera=None)
    for v in _con_comisiones(p["consolidado"]["vuelos"]):
        v["ganancia_antes_comisiones_usd"] = r2(v["ganancia_usd"] + 12.34)
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert tp._huella_xlsx(data) == _HUELLAS_HEAD[("general-mensual", None)][0]
    ws = load_workbook(BytesIO(data))[MAESTRA_FLOTA]
    letra = _DISP_FLOTA.letra["ganancia_usd"]
    vuelos = p["consolidado"]["vuelos"]
    primera = vuelos.index(_con_comisiones(vuelos)[0])
    assert str(ws[f"{letra}{3 + primera}"].value).startswith("=ROUND(")
    assert ws[f"{letra}2"].comment is None


def test_usd_por_gasto_sin_bandera_las_formulas_de_siempre() -> None:
    """Las MISMAS llaves perturbadas sin la bandera: las fórmulas de siempre
    (y `finalizar` las degrada a valor porque ya no reproducen el API)."""
    p = con_usd_por_gasto(pp.general("mensual", bandera=None)["aviones"][1])
    p["usd_por_gasto"] = False
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    ws = load_workbook(BytesIO(data))[MAESTRA_DOS]
    letra = _DISP_MENSUAL.letra["costo_usd"]
    assert ws[f"{letra}2"].comment is None
    assert ws[f"{letra}3"].value == p["vuelos"][0]["costo_usd"]  # degradada a valor
    assert load_workbook(BytesIO(data))["combustible"]["E4"].comment is None
