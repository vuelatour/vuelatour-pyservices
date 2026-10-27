"""Balance general: VUELOS CON AVIONES EXTERNOS en el bloque «VUELATOUR
(empresa)» (27-oct-2026, API 0.0.112, invariante 112 del API). Pedido del
cliente sobre la hoja «balance» de septiembre: «me falta en el balance de
VuelaTour los vuelos con aviones externos … que sea claro cuánto dejaron de
ingreso, cuánto de gastos y al final la utilidad, y un resumen a un lado de
cuántos fueron».

El API manda `empresa.externos` (ya calculado con el libro EXTERNOS; su
utilidad YA suma en `resultado_*`). Aquí: tres filas entre TIENDA y
RESULTADO (la utilidad como `ROUND(ingreso − gastos, 2)` verificada), el
RESULTADO cita esa celda (`+B{ext}`), la tabla de los vuelos a un lado
(columna F) y la nota al pie. Sin la llave, el libro es byte a byte el de
antes (también con `externos: null`).
"""

from __future__ import annotations

import copy
import zipfile
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.schemas.reportes import BalanceEmpresaBloque, BalanceGeneralRequest
from app.services import xlsx_formulas
from app.services.balance_avion_xlsx import (
    _COL_TABLA_EXTERNOS,
    _EMPRESA_EXTERNOS,
    _EMPRESA_EXTERNOS_ENCABEZADOS,
    _EMPRESA_EXTERNOS_GASTOS,
    _EMPRESA_EXTERNOS_HOJAS,
    _EMPRESA_EXTERNOS_INGRESO,
    _EMPRESA_EXTERNOS_TOTAL,
    _NOTA_EMPRESA_EXTERNOS,
    _SIN_PESOS,
    _etiqueta_pesos,
    libro_balance_general,
    render_balance_general_xlsx,
)
from tests import _pesos_payloads as pp
from tests import test_balance_empresa_vuelatour as te

r2 = pp.r2

_HEAD_MXN = _etiqueta_pesos(_EMPRESA_EXTERNOS)
_RES_MXN = "RESULTADO VUELATOUR (MXN)"
_RES_USD = "RESULTADO VUELATOUR USD"


def _vuelo(**kw) -> dict:
    base = {
        "clave": "#453",
        "folio": "453",
        "fecha": "2026-10-25",
        "cliente": "Oscar Dos",
        "operador": "Simon Herrera",
        "estado": "COMPLETADO",
        "cancelado": False,
        "ruta": "CUN → HOL → CUN",
        "ingreso_mxn": 17564.36,
        "gastos_mxn": 18182.0,
        "utilidad_mxn": -617.64,
        "por_cobrar_mxn": 0,
        "status_cobro": "PAGADO",
        "ingreso_usd": 950.0,
        "gastos_usd": 727.28,
        "utilidad_usd": 222.72,
    }
    base.update(kw)
    return base


_NOTA_API = (
    "Vuelos con aviones externos (libro EXTERNOS del periodo): 3 vuelos (1 cancelado) "
    "con Simon Herrera, Aerocharter Cancun. Ingreso = la VENTA de cada vuelo en ese libro "
    "(lo cotizado al T.C. de su vuelo; en un cancelado, lo realmente cobrado); gastos = "
    "todo lo que ese libro resta (operaciones con el pago al operador —una sola vez—, "
    "combustible, comisiones, permisos y otros); utilidad = su UTILIDAD DESPUÉS DE GASTOS, "
    "que suma en el RESULTADO completa aunque esté por cobrar. De esos vuelos, $41,798.85 "
    "MXN siguen por cobrar (1 vuelo). Sus TUAs, extras, pernocta y comisión del vendedor ya "
    "están en los ingresos y egresos propios (no se suman dos veces)."
)


def _externos(*, sin_pesos: bool = False, hojas: float = 1000.0) -> dict:
    """Octubre 2026 de prod (espejo de `externosParaEmpresa` del API): #453
    (Simon Herrera, pago al operador marcado), #413 (Aerocharter Cancun, sin
    cobrar) y #234 CANCELADO sin dinero; el libro resta `hojas` de
    combustible fuera de las filas."""
    detalle = [
        _vuelo(),
        _vuelo(
            clave="#413",
            folio="413",
            fecha="2026-10-23",
            cliente="Cliente Trece",
            operador="Aerocharter Cancun",
            ingreso_mxn=41798.85,
            gastos_mxn=0.0,
            utilidad_mxn=41798.85,
            por_cobrar_mxn=41798.85,
            status_cobro="PENDIENTE",
            ingreso_usd=2325.0,
            gastos_usd=0.0,
            utilidad_usd=2325.0,
        ),
        _vuelo(
            clave="#234",
            folio="234",
            estado="CANCELADO",
            cancelado=True,
            ingreso_mxn=0.0,
            gastos_mxn=0.0,
            utilidad_mxn=0.0,
            status_cobro="",
            ingreso_usd=0.0,
            gastos_usd=0.0,
            utilidad_usd=0.0,
        ),
    ]
    ingreso = r2(17564.36 + 41798.85)
    antes = r2(-617.64 + 41798.85)
    utilidad = r2(antes - hojas)
    ingreso_usd = r2(950 + 2325)
    antes_usd = r2(222.72 + 2325)
    hojas_usd = r2(hojas / 18.5)
    utilidad_usd = r2(antes_usd - hojas_usd)
    ext = {
        "vuelos": 3,
        "cancelados": 1,
        "operadores": ["Simon Herrera", "Aerocharter Cancun"],
        "ingreso_mxn": ingreso,
        "gastos_mxn": r2(ingreso - utilidad),
        "utilidad_mxn": utilidad,
        "por_cobrar_mxn": 41798.85,
        "hojas_mxn": hojas,
        "ingreso_usd": ingreso_usd,
        "gastos_usd": r2(ingreso_usd - utilidad_usd),
        "utilidad_usd": utilidad_usd,
        "por_cobrar_usd": 2325.0,
        "hojas_usd": hojas_usd,
        "detalle": detalle,
        "sin_pesos": [],
        "nota": _NOTA_API,
    }
    if sin_pesos:
        detalle[1].update(ingreso_mxn=None, gastos_mxn=None, utilidad_mxn=None)
        ext.update(ingreso_mxn=None, gastos_mxn=None, utilidad_mxn=None, hojas_mxn=None)
        ext["sin_pesos"] = ["#413"]
        ext["nota"] = _NOTA_API + (
            " #413 sin pesos (sin T.C.): ingreso, gastos, utilidad y RESULTADO quedan "
            "vacíos hasta capturarlo."
        )
    return ext


def _con_externos(p: dict, **kw) -> dict:
    """El payload con `empresa.externos` y el RESULTADO ya sumado (como el API)."""
    p = copy.deepcopy(p)
    ext = _externos(**kw)
    emp = p["empresa"]
    emp["externos"] = ext
    if emp.get("resultado_mxn") is not None:
        emp["resultado_mxn"] = (
            None if ext["utilidad_mxn"] is None else r2(emp["resultado_mxn"] + ext["utilidad_mxn"])
        )
    if emp.get("resultado_usd") is not None:
        emp["resultado_usd"] = r2(emp["resultado_usd"] + ext["utilidad_usd"])
    return p


def _render(p: dict) -> bytes:
    return render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))


def _libros(data: bytes):
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _fila(ws, texto: str) -> int:
    filas = [c.row for c in ws["A"] if c.value == texto]
    assert len(filas) == 1, (texto, filas)
    return filas[0]


def _fila_que_empieza(ws, prefijo: str) -> int:
    return next(c.row for c in ws["A"] if isinstance(c.value, str) and c.value.startswith(prefijo))


def _textos_col_a(ws) -> list[str]:
    return [c.value for c in ws["A"] if isinstance(c.value, str)]


def _miembros(data: bytes) -> dict[str, bytes]:
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


@pytest.fixture
def registros(monkeypatch) -> list[xlsx_formulas.RegistroFormulas]:
    capturados: list[xlsx_formulas.RegistroFormulas] = []
    original = xlsx_formulas.finalizar

    def espia(wb):
        cache = original(wb)
        capturados.append(xlsx_formulas.registro(wb))
        return cache

    monkeypatch.setattr(xlsx_formulas, "finalizar", espia)
    return capturados


def _degradadas_balance(reg: xlsx_formulas.RegistroFormulas) -> list[tuple[str, str]]:
    return [d for d in reg.degradadas() if d[0] == "balance"]


def _num_fila(coord: str) -> int:
    return int("".join(ch for ch in coord if ch.isdigit()))


# ---------------------------------------------------------------------------
# 1) Esquema: aditivo, liberal, jamás un 422.
# ---------------------------------------------------------------------------


def test_esquema_aditivo_y_liberal() -> None:
    assert BalanceEmpresaBloque.model_validate({}).externos is None
    assert BalanceEmpresaBloque.model_validate({"externos": None}).externos is None
    e = BalanceEmpresaBloque.model_validate(
        {
            "externos": {
                "vuelos": None,
                "cancelados": "x",
                "operadores": None,
                "detalle": [{"clave": "#1", "folio": 453, "cancelado": None, "operador": " "}],
                "sin_pesos": None,
                "nota": "  ",
            }
        }
    ).externos
    assert e is not None
    assert (e.vuelos, e.cancelados, e.operadores, e.sin_pesos, e.nota) == (0, 0, [], [], None)
    assert e.detalle[0].folio == "453"
    assert e.detalle[0].cancelado is False
    assert e.detalle[0].operador is None
    assert e.ingreso_mxn is None and e.utilidad_mxn is None


# ---------------------------------------------------------------------------
# 2) Bloque EN PESOS (lo que ve la oficina): filas, fórmulas, tabla y pie.
# ---------------------------------------------------------------------------


def test_pesos_filas_formulas_tabla_y_pie(registros) -> None:
    p = _con_externos(pp.general("mensual"))
    emp = p["empresa"]
    ext = emp["externos"]
    data = _render(p)
    assert _degradadas_balance(registros[-1]) == []
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]

    head = _fila(ws, _HEAD_MXN)
    ing = _fila(ws, f"{_EMPRESA_EXTERNOS_INGRESO} · 3 vuelos (1 cancelado)")
    gas = _fila(ws, _EMPRESA_EXTERNOS_GASTOS)
    tienda = _fila(ws, "(+) UTILIDAD TIENDA (INVENTARIO) (MXN)")
    res = _fila(ws, _RES_MXN)
    assert (ing, gas) == (head + 1, head + 2)
    assert tienda < head < res

    # La utilidad = ROUND(ingreso − gastos, 2) en pesos y en USD informativo,
    # con el número del API en caché; las dos de detalle, como valor.
    assert ws.cell(row=head, column=2).value == f"=ROUND(B{ing}-B{gas},2)"
    assert ws.cell(row=head, column=3).value == f"=ROUND(C{ing}-C{gas},2)"
    assert wsv.cell(row=head, column=2).value == ext["utilidad_mxn"]
    assert wsv.cell(row=head, column=3).value == ext["utilidad_usd"]
    assert ws.cell(row=head, column=2).comment.text == _NOTA_EMPRESA_EXTERNOS
    assert ws.cell(row=ing, column=2).value == ext["ingreso_mxn"]
    assert ws.cell(row=ing, column=3).value == ext["ingreso_usd"]
    assert ws.cell(row=gas, column=2).value == ext["gastos_mxn"]
    assert ws.cell(row=gas, column=3).value == ext["gastos_usd"]

    # RESULTADO = la fórmula de siempre + la celda de los externos, y
    # reproduce el número del API (que ya la trae sumada).
    part = _fila(ws, "(+) PARTICIPACIÓN COMO SOCIO EN LOS AVIONES (MXN)")
    otros = _fila(ws, "(−) OTROS GASTOS DE LA EMPRESA (MXN)")
    ingp = _fila_que_empieza(ws, "(+) INGRESOS PROPIOS MXN (")
    pag = _fila_que_empieza(ws, "(−) EGRESOS PROPIOS MXN (")
    assert ws.cell(row=res, column=2).value == (
        f"=ROUND(B{part}+B{ingp}-B{pag}-B{otros}+B{tienda}+B{head},2)"
    )
    assert ws.cell(row=res, column=3).value == (
        f"=ROUND(C{part}+C{ingp}-C{pag}-C{otros}+C{tienda}+C{head},2)"
    )
    assert wsv.cell(row=res, column=2).value == emp["resultado_mxn"]
    assert wsv.cell(row=res, column=3).value == emp["resultado_usd"]

    # La tabla a un lado (columna F), a la altura de la fila de externos.
    col = _COL_TABLA_EXTERNOS
    assert ws.cell(row=head, column=col).value == (
        "Vuelos con aviones externos del periodo: 3 vuelos (1 cancelado)"
    )
    enc = [ws.cell(row=head + 1, column=col + i).value for i in range(9)]
    assert enc == _EMPRESA_EXTERNOS_ENCABEZADOS
    filas = [[ws.cell(row=head + 2 + i, column=col + j).value for j in range(9)] for i in range(3)]
    assert filas[0] == [
        "#453",
        "25/10/2026",
        "Simon Herrera",
        "Oscar Dos",
        "COMPLETADO",
        17564.36,
        18182.0,
        -617.64,
        "PAGADO",
    ]
    assert filas[1][:5] == [
        "#413",
        "23/10/2026",
        "Aerocharter Cancun",
        "Cliente Trece",
        "COMPLETADO",
    ]
    assert filas[1][5:] == [41798.85, 0.0, 41798.85, "PENDIENTE · por cobrar $41,798.85 MXN"]
    assert filas[2][:5] == ["#234", "25/10/2026", "Simon Herrera", "Oscar Dos", "CANCELADO"]
    assert filas[2][5:8] == [0.0, 0.0, 0.0]
    hojas = head + 5
    assert ws.cell(row=hojas, column=col).value == _EMPRESA_EXTERNOS_HOJAS
    assert ws.cell(row=hojas, column=col + 6).value == 1000.0
    assert ws.cell(row=hojas, column=col + 7).value == -1000.0
    total = hojas + 1
    assert ws.cell(row=total, column=col).value == f"{_EMPRESA_EXTERNOS_TOTAL} (MXN)"
    assert [ws.cell(row=total, column=col + 5 + i).value for i in range(3)] == [
        ext["ingreso_mxn"],
        ext["gastos_mxn"],
        ext["utilidad_mxn"],
    ]
    # La tabla es más alta que las tres filas: el RESULTADO va DESPUÉS de ella.
    assert res > total
    # Pie: la nota de los externos del API está entre las notas del bloque.
    assert ext["nota"] in _textos_col_a(ws)


def test_pesos_sin_pesos_pinta_raya_y_resultado_vacio() -> None:
    p = _con_externos(pp.general("mensual"), sin_pesos=True)
    assert p["empresa"]["resultado_mxn"] is None
    data = _render(p)
    wb, _ = _libros(data)
    ws = wb["balance"]
    head = _fila(ws, _HEAD_MXN)
    res = _fila(ws, _RES_MXN)
    assert ws.cell(row=head, column=2).value == _SIN_PESOS
    assert ws.cell(row=head + 1, column=2).value == _SIN_PESOS
    assert ws.cell(row=res, column=2).value == _SIN_PESOS
    # USD informativo sigue (el libro sí tiene dólares).
    assert ws.cell(row=head, column=3).value == f"=ROUND(C{head + 1}-C{head + 2},2)"
    # La fila #413 en la tabla lleva «—» en pesos.
    col = _COL_TABLA_EXTERNOS
    assert ws.cell(row=head + 3, column=col + 5).value == _SIN_PESOS
    assert "#413 sin pesos (sin T.C.)" in " ".join(_textos_col_a(ws))


def test_pesos_sin_hojas_ni_cancelados_tabla_corta(registros) -> None:
    p = _con_externos(pp.general("mensual"), hojas=0.0)
    ext = p["empresa"]["externos"]
    ext["detalle"] = ext["detalle"][:1]
    ext.update(vuelos=1, cancelados=0, operadores=["Simon Herrera"])
    data = _render(p)
    assert _degradadas_balance(registros[-1]) == []
    wb, _ = _libros(data)
    ws = wb["balance"]
    head = _fila(ws, _HEAD_MXN)
    _fila(ws, f"{_EMPRESA_EXTERNOS_INGRESO} · 1 vuelo")
    col = _COL_TABLA_EXTERNOS
    assert ws.cell(row=head, column=col).value == "Vuelos con aviones externos del periodo: 1 vuelo"
    # Sin hojas no hay renglón de hojas: el TOTAL va justo después del vuelo.
    assert ws.cell(row=head + 3, column=col).value == f"{_EMPRESA_EXTERNOS_TOTAL} (MXN)"
    assert _EMPRESA_EXTERNOS_HOJAS not in [
        ws.cell(row=r, column=col).value for r in range(head, head + 6)
    ]
    # Tabla de 4 filas = las 3 del bloque + 1: el RESULTADO va en head + 4.
    assert _fila(ws, _RES_MXN) == head + 4


# ---------------------------------------------------------------------------
# 3) Bloque en dólares (API sin `balance_en_pesos`): las mismas filas en USD.
# ---------------------------------------------------------------------------


def test_usd_filas_formulas_y_tabla(registros) -> None:
    p = _con_externos(te._payload())
    emp = p["empresa"]
    ext = emp["externos"]
    data = _render(p)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    head = _fila(ws, _EMPRESA_EXTERNOS)
    # C16 ya venía degradada en el payload USD de siempre: solo desde el bloque.
    assert [d for d in _degradadas_balance(registros[-1]) if _num_fila(d[1]) >= head] == []
    ing, gas = head + 1, head + 2
    assert ws.cell(row=ing, column=1).value == (
        f"{_EMPRESA_EXTERNOS_INGRESO} · 3 vuelos (1 cancelado)"
    )
    assert ws.cell(row=head, column=2).value == f"=ROUND(B{ing}-B{gas},2)"
    assert wsv.cell(row=head, column=2).value == ext["utilidad_usd"]
    assert ws.cell(row=ing, column=2).value == ext["ingreso_usd"]
    assert ws.cell(row=gas, column=2).value == ext["gastos_usd"]
    b = te._bloque(ws)
    assert ws.cell(row=b["res"], column=2).value == (
        f"=ROUND(B{b['part']}+B{b['ing']}-B{b['pag']}-B{b['otros']}+B{b['tienda']}+B{head},2)"
    )
    assert wsv.cell(row=b["res"], column=2).value == emp["resultado_usd"]
    col = _COL_TABLA_EXTERNOS
    assert ws.cell(row=head + 1, column=col + 5).value == "VENTA USD"
    assert ws.cell(row=head + 2, column=col + 5).value == 950.0
    assert ws.cell(row=head + 2, column=col + 7).value == 222.72
    assert ext["nota"] in _textos_col_a(ws)


# ---------------------------------------------------------------------------
# 4) Sin la llave (o null): el libro de antes, byte a byte.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("libro", ["pesos", "usd"])
def test_sin_externos_byte_a_byte(libro: str) -> None:
    p = pp.general("mensual") if libro == "pesos" else te._payload()
    sin_llave = _render(p)
    con_null = copy.deepcopy(p)
    con_null["empresa"]["externos"] = None
    assert _miembros(_render(con_null)) == _miembros(sin_llave)
    wb, _ = _libros(sin_llave)
    ws = wb["balance"]
    textos = _textos_col_a(ws)
    assert _HEAD_MXN not in textos and _EMPRESA_EXTERNOS not in textos
    assert not any(t.startswith(_EMPRESA_EXTERNOS_INGRESO) for t in textos)
    if libro == "pesos":
        res = _fila(ws, _RES_MXN)
        formula = ws.cell(row=res, column=2).value
    else:
        b = te._bloque(ws)
        formula = ws.cell(row=b["res"], column=2).value
    assert formula.count("+") == 2 and formula.count("-") == 2  # part+ing−pag−otros+tienda


# ---------------------------------------------------------------------------
# 5) La vista en pantalla (libro_json) hereda las filas y la tabla.
# ---------------------------------------------------------------------------


def test_vista_en_pantalla_hereda_los_externos() -> None:
    p = _con_externos(pp.general("mensual"))
    vista = libro_balance_general(
        BalanceGeneralRequest.model_validate({**p, "vinculos_gasto": True})
    )
    hoja = next(h for h in vista["hojas"] if h["nombre"] == "balance")
    textos = {c.get("t") for c in hoja["celdas"] if c.get("t")}
    assert _HEAD_MXN in textos
    assert "Vuelos con aviones externos del periodo: 3 vuelos (1 cancelado)" in textos
    assert "Simon Herrera" in textos
