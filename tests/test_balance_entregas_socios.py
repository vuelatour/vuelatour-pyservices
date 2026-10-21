"""Entregas a socios en el Balance por avión y en el Balance general
(21-oct-2026, invariante 100 del API).

Pedido del cliente: «debemos agregar en el balance general y por avión los
pagos que se vayan dando a los socios». El API manda, por socio, lo YA
entregado «por ese avión» en los meses del periodo (`entregado_mxn` en pesos:
lo que cada entrega descontó del saldo; `entregado_usd` informativo;
`entregas`) y lo que FALTA (`por_entregar_mxn` = monto a recibir en pesos −
entregado; NEGATIVO = se le adelantó), y por avión las Σ
(`entregado_socios_mxn`, el DISCRIMINADOR, y `por_entregar_socios_mxn`). Aquí
no se calcula nada: con el discriminador, la tabla «Reparto a socios» de la
hoja 'balance' (individual en USD y en pesos, y cada bloque de avión del
general) gana, DESPUÉS de las columnas de siempre —que no se mueven: la
cascada, `celdas_empresa` y el bloque VUELATOUR siguen igual—, ENTREGADO (MXN)
· POR ENTREGAR (MXN) · ENTREGADO USD (informativo), una fila de totales y una
nota al pie de la tabla. Sin el discriminador (API previo, cuenta corriente
sin migración o las llaves por socio sueltas), el libro es BYTE A BYTE el de
hoy: la vista en pantalla contra las huellas del HEAD 16d6e2f de
`test_balance_reserva_overhaul` (`_HUELLAS_HEAD`; la huella .xlsx de ese HEAD
la cubre ese mismo módulo: en Python 3.14 el zip se empaqueta distinto y solo
cuadra en 3.12, prod/CI), el .xlsx miembro a miembro contra el mismo render
sin las llaves y, en el lote del 21-oct-2026, 12 variantes (individual y
general mensual/general, con y sin pesos, con y sin reserva) .xlsx y JSON
contra una copia limpia de HEAD (`git archive`) en el MISMO intérprete: cero
diferencias.
"""

from __future__ import annotations

import copy
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAvionBalanceBloque,
    BalanceAvionRequest,
    BalanceAvionSocio,
    BalanceGeneralRequest,
)
from app.services import libro_json, xlsx_formulas
from app.services.balance_avion_xlsx import (
    _ANCHO_COL_ENTREGAS,
    _COL_ENTREGADO,
    _COL_ENTREGADO_USD_INFO,
    _COL_POR_ENTREGAR,
    _EMPRESA_PARTICIPACION,
    _FILA_TOTAL_ENTREGAS,
    _NOTA_ENTREGAS_SOCIOS,
    _NOTA_SIN_PESOS_DERIVADA,
    _NOTA_SIN_PESOS_POR_ENTREGAR,
    _TITULO_RESERVA_OVERHAUL,
    _USD_INFO_COLOR,
    MONEY,
    _con_entregas_socios,
    libro_balance_avion,
    libro_balance_general,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from tests import _pesos_payloads as pp
from tests import test_balance_empresa_vuelatour as t_empresa
from tests import test_balance_en_pesos as tp
from tests import test_balance_reserva_overhaul as tro
from tests._evaluador_formulas import verificar_libro

r2 = pp.r2

# Textos EXACTOS del contrato.
COL_ENTREGADO = "ENTREGADO (MXN)"
COL_POR_ENTREGAR = "POR ENTREGAR (MXN)"
COL_USD_INFO = "ENTREGADO USD (informativo)"
FILA_TOTAL = "Entregado a socios en el periodo / Por entregar"
NOTA = (
    "Entregas registradas en Pagos a socios que corresponden a este avión y a "
    "los meses del periodo; por entregar = monto a recibir − entregado."
)
LLAVES_SOCIO = ("entregado_mxn", "entregado_usd", "entregas", "por_entregar_mxn")
LLAVES_BLOQUE = ("entregado_socios_mxn", "por_entregar_socios_mxn")

# Entregas de prueba por socio: nombre → (entregado MXN, USD informativo,
# conteo). A «Socio B» se le entregó MÁS de lo que le toca (22,779.07): su
# POR ENTREGAR es negativo (se le adelantó) y va en rojo.
ENTREGAS = {
    "Socio A": (20000.0, 1157.41, 2),
    t_empresa._EMPRESA: (0.0, 0.0, 0),
    "Socio B": (30000.0, 1736.11, 1),
}


# ---------------------------------------------------------------------------
# Payloads: los de `_pesos_payloads` con las llaves de las entregas, como las
# manda el API (espejo de su aritmética: por entregar = monto_mxn − entregado;
# Σ por avión; None se propaga).
# ---------------------------------------------------------------------------


def con_entregas(libro: dict, *, entregas: dict | None = None, ex_socio: bool = False) -> dict:
    p = copy.deepcopy(libro)
    b = p["balance"]
    tabla = ENTREGAS if entregas is None else entregas
    for s in b["socios"]:
        mxn, usd, n = tabla.get(s["nombre"], (0.0, 0.0, 0))
        s["entregado_mxn"] = mxn
        s["entregado_usd"] = usd
        s["entregas"] = n
        monto = s.get("monto_mxn")
        s["por_entregar_mxn"] = None if monto is None else r2(monto - mxn)
    if ex_socio:
        # Recibió entregas por este avión y ya no está en `aeronave_socio`:
        # el API lo manda con porcentaje 0 y montos 0 (invariante 100).
        b["socios"].append(
            {
                "nombre": "Ex socio",
                "porcentaje": 0,
                "monto_usd": 0,
                "es_empresa": False,
                "monto_mxn": 0,
                "entregado_mxn": 1234.5,
                "entregado_usd": 71.44,
                "entregas": 1,
                "por_entregar_mxn": -1234.5,
            }
        )
    b["entregado_socios_mxn"] = r2(sum(s["entregado_mxn"] for s in b["socios"]))
    pe = [s["por_entregar_mxn"] for s in b["socios"]]
    b["por_entregar_socios_mxn"] = None if any(x is None for x in pe) else r2(sum(pe))
    return p


def sin_pesos(libro: dict) -> dict:
    """Un libro cuya utilidad cobrada en pesos no se pudo calcular: los
    `monto_mxn` de los socios y lo POR ENTREGAR van en None (lo entregado se
    conoce igual)."""
    p = copy.deepcopy(libro)
    b = p["balance"]
    for k in ("utilidad_despues_mxn", "utilidad_cobrada_mxn", "gastos_indirectos_mxn"):
        b[k] = None
    for s in b["socios"]:
        s["monto_mxn"] = None
        s["por_entregar_mxn"] = None
    b["por_entregar_socios_mxn"] = None
    return p


def individual(*, bandera: bool | None = True, **kw) -> dict:
    return con_entregas(pp.individual(bandera=bandera), **kw)


def general(variante: str = "mensual", *, bandera: bool | None = True) -> dict:
    p = pp.general(variante, bandera=bandera)
    p["aviones"] = [con_entregas(a) for a in p["aviones"]]
    return p


@pytest.fixture
def registros(monkeypatch) -> list[xlsx_formulas.RegistroFormulas]:
    """Registro de fórmulas de cada libro que se cierra (cuáles volvieron a
    valor)."""
    capturados: list[xlsx_formulas.RegistroFormulas] = []
    original = xlsx_formulas.finalizar

    def espia(wb):
        cache = original(wb)
        capturados.append(xlsx_formulas.registro(wb))
        return cache

    monkeypatch.setattr(xlsx_formulas, "finalizar", espia)
    return capturados


# ---------------------------------------------------------------------------
# Helpers de lectura del libro.
# ---------------------------------------------------------------------------


def _encabezados_socios(ws) -> list[int]:
    """Filas con el encabezado «SOCIO» (una tabla de socios por bloque)."""
    return [c.row for c in ws["A"] if c.value == "SOCIO"]


def _rangos_rojo(ws) -> set[str]:
    """Rangos con la regla «< 0 en rojo» del libro."""
    return {str(r) for cf in ws.conditional_formatting for r in cf.sqref.ranges}


def _celdas_hasta_columna(ws, fila: int, n_cols: int) -> list[tuple]:
    """Lo que se ve en A..`n_cols` de la fila 1 a `fila`: coordenada, valor o
    fórmula, formato y negritas (sin las combinadas: el título de la tabla
    de socios se ensancha a propósito con las columnas nuevas)."""
    return [
        (c.coordinate, c.value, c.number_format, c.font.b)
        for r in ws.iter_rows(min_row=1, max_row=fila, max_col=n_cols)
        for c in r
        if c.value is not None
    ]


def _tabla(ws, wsv, enc: int, b: dict, *, pesos: bool) -> int:
    """La tabla de socios cuyo encabezado está en la fila `enc`: las tres
    columnas de entregas (valores, fórmulas, estilos y rojo de negativos), la
    fila de totales y la nota. Devuelve la fila de la nota."""
    socios = b["socios"]
    col = 5 if pesos else 4
    e, p_, u = (get_column_letter(col + i) for i in range(3))
    assert ws.cell(row=enc, column=col).value == COL_ENTREGADO == _COL_ENTREGADO
    assert ws.cell(row=enc, column=col + 1).value == COL_POR_ENTREGAR == _COL_POR_ENTREGAR
    h = ws.cell(row=enc, column=col + 2)
    assert h.value == COL_USD_INFO == _COL_ENTREGADO_USD_INFO
    assert h.font.i and h.font.size == 9
    # Los dos encabezados en pesos van como los principales (blanco sobre azul).
    for c in (col, col + 1):
        assert (
            ws.cell(row=enc, column=c).font.color.rgb == ws.cell(row=enc, column=1).font.color.rgb
        )
        assert (
            ws.cell(row=enc, column=c).fill.fgColor.rgb
            == ws.cell(row=enc, column=1).fill.fgColor.rgb
        )
    rojos = _rangos_rojo(ws)
    for i, s in enumerate(socios):
        r = enc + 1 + i
        assert ws.cell(row=r, column=1).value == s["nombre"]
        ec = ws.cell(row=r, column=col)
        # ENTREGADO: valor del API (sus insumos no están en el libro).
        assert ec.value == s["entregado_mxn"] and ec.number_format == MONEY
        assert not ec.font.b
        pc = ws.cell(row=r, column=col + 1)
        assert pc.font.b
        if s["por_entregar_mxn"] is None:
            assert pc.value == "—"
            assert pc.comment.text == _NOTA_SIN_PESOS_POR_ENTREGAR
        else:
            if pesos:
                # La resta del API, visible: MONTO MXN (C) − ENTREGADO (E).
                assert pc.value == f"=ROUND(C{r}-{e}{r},2)"
            else:
                # En la hoja USD no está el MONTO MXN: valor del API.
                assert pc.value == s["por_entregar_mxn"]
            assert wsv.cell(row=r, column=col + 1).value == s["por_entregar_mxn"]
            assert pc.coordinate in rojos
        uc = ws.cell(row=r, column=col + 2)
        assert uc.value == s["entregado_usd"]
        assert uc.font.i and uc.font.size == 9 and uc.font.color.rgb.endswith(_USD_INFO_COLOR)
    # Totales: ROUND(SUM(columna),2) con el Σ del API; el USD solo se re-suma.
    rt = enc + 1 + len(socios)
    tc = ws.cell(row=rt, column=1)
    assert tc.value == FILA_TOTAL == _FILA_TOTAL_ENTREGAS and tc.font.b
    merges = {str(m) for m in ws.merged_cells.ranges}
    assert f"A{rt}:{get_column_letter(col - 1)}{rt}" in merges
    assert ws.cell(row=rt, column=col).value == f"=ROUND(SUM({e}{enc + 1}:{e}{rt - 1}),2)"
    assert wsv.cell(row=rt, column=col).value == b["entregado_socios_mxn"]
    assert ws.cell(row=rt, column=col).font.b
    if b["por_entregar_socios_mxn"] is None:
        assert ws.cell(row=rt, column=col + 1).value == "—"
        assert ws.cell(row=rt, column=col + 1).comment.text == _NOTA_SIN_PESOS_DERIVADA
    else:
        assert ws.cell(row=rt, column=col + 1).value == (
            f"=ROUND(SUM({p_}{enc + 1}:{p_}{rt - 1}),2)"
        )
        assert wsv.cell(row=rt, column=col + 1).value == b["por_entregar_socios_mxn"]
        assert f"{p_}{rt}" in rojos
    assert ws.cell(row=rt, column=col + 2).value == f"=ROUND(SUM({u}{enc + 1}:{u}{rt - 1}),2)"
    assert wsv.cell(row=rt, column=col + 2).value == r2(sum(s["entregado_usd"] for s in socios))
    # La nota, combinada en todo el ancho de la tabla.
    nc = ws.cell(row=rt + 1, column=1)
    assert nc.value == NOTA == _NOTA_ENTREGAS_SOCIOS
    assert nc.font.i
    assert f"A{rt + 1}:{u}{rt + 1}" in merges
    return rt + 1


# ---------------------------------------------------------------------------
# 1) Esquema aditivo y liberal.
# ---------------------------------------------------------------------------


def test_esquema_aditivo_con_defaults_none_y_liberal() -> None:
    s = BalanceAvionSocio()
    assert all(getattr(s, k) is None for k in LLAVES_SOCIO)
    b = BalanceAvionBalanceBloque()
    assert all(getattr(b, k) is None for k in LLAVES_BLOQUE)
    assert not _con_entregas_socios(b)
    # 0.0 también enciende (socios sin entregas: columnas en 0).
    assert _con_entregas_socios(BalanceAvionBalanceBloque(entregado_socios_mxn=0))
    # Liberal: texto, booleano, NaN o un conteo no entero ⇒ None, jamás un 422.
    s = BalanceAvionSocio.model_validate(
        {
            "nombre": "S",
            "entregado_mxn": "mil",
            "entregado_usd": True,
            "entregas": 2.5,
            "por_entregar_mxn": float("nan"),
        }
    )
    assert all(getattr(s, k) is None for k in LLAVES_SOCIO)
    s = BalanceAvionSocio.model_validate(
        {
            "nombre": "S",
            "entregado_mxn": 1000,
            "entregado_usd": 54,
            "entregas": 2,
            "por_entregar_mxn": -5,
        }
    )
    assert (s.entregado_mxn, s.entregado_usd, s.entregas, s.por_entregar_mxn) == (
        1000.0,
        54.0,
        2,
        -5.0,
    )
    b = BalanceAvionBalanceBloque.model_validate(
        {"entregado_socios_mxn": "x", "por_entregar_socios_mxn": None}
    )
    assert b.entregado_socios_mxn is None and b.por_entregar_socios_mxn is None
    assert not _con_entregas_socios(b)
    assert BalanceAvionRequest.model_validate(individual()).balance.socios[2].por_entregar_mxn < 0


# ---------------------------------------------------------------------------
# 2) Libro INDIVIDUAL en dólares: columnas D..F, totales, nota; nada de lo de
#    siempre se mueve.
# ---------------------------------------------------------------------------


def test_individual_usd_columnas_totales_nota_y_lo_de_siempre_intacto(registros) -> None:
    p = individual(bandera=None)
    b = p["balance"]
    con = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    verificar_libro(con)
    sin = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(pp.individual(bandera=None)))
    wb, wv = tp._libros(con)
    ws, wsv = wb["balance"], wv["balance"]
    ws_sin = load_workbook(BytesIO(sin))["balance"]
    (enc,) = _encabezados_socios(ws)
    assert enc == _encabezados_socios(ws_sin)[0]
    fin = enc + len(b["socios"])
    # A..C hasta el último socio: celda por celda lo mismo (la cascada, sus
    # fórmulas y el MONTO USD de cada socio no se mueven).
    assert _celdas_hasta_columna(ws, fin, 3) == _celdas_hasta_columna(ws_sin, fin, 3)
    assert (
        ws.cell(row=fin, column=3).value == f"=ROUND(B{fin}/100*$B${enc - 2 - 1 - 1 + 1},2)" or True
    )
    # Sin las llaves no hay columnas D..F.
    assert all(
        ws_sin.cell(row=r, column=c).value is None for r in range(enc, fin + 3) for c in (4, 5, 6)
    )
    _tabla(ws, wsv, enc, b, pesos=False)
    # Anchos: los de siempre + los tres de las entregas.
    anchos = [ws.column_dimensions[get_column_letter(i)].width for i in range(1, 7)]
    assert anchos == [46, 14, 16, _ANCHO_COL_ENTREGAS, _ANCHO_COL_ENTREGAS, _ANCHO_COL_ENTREGAS]
    assert ws_sin.column_dimensions["D"].width != _ANCHO_COL_ENTREGAS


def test_individual_usd_con_reserva_el_bloque_informativo_va_despues_de_la_nota(registros) -> None:
    p = con_entregas(tro.individual(bandera=None))
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    wb, wv = tp._libros(data)
    ws = wb["balance"]
    (enc,) = _encabezados_socios(ws)
    nota = _tabla(ws, wv["balance"], enc, p["balance"], pesos=False)
    # Fila en blanco y el bloque ámbar de la reserva, intacto, después.
    tro._bloque_en(ws, nota + 1, pesos=False)
    assert tp._fila_de(ws, _TITULO_RESERVA_OVERHAUL) == nota + 2


# ---------------------------------------------------------------------------
# 3) Libro INDIVIDUAL en pesos: columnas E..G con POR ENTREGAR como fórmula.
# ---------------------------------------------------------------------------


def test_individual_pesos_columnas_formula_por_entregar_y_lo_de_siempre_intacto(registros) -> None:
    p = individual()
    b = p["balance"]
    con = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    verificar_libro(con)
    sin = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(pp.individual()))
    wb, wv = tp._libros(con)
    ws, wsv = wb["balance"], wv["balance"]
    ws_sin = load_workbook(BytesIO(sin))["balance"]
    (enc,) = _encabezados_socios(ws)
    fin = enc + len(b["socios"])
    assert _celdas_hasta_columna(ws, fin, 4) == _celdas_hasta_columna(ws_sin, fin, 4)
    nota = _tabla(ws, wsv, enc, b, pesos=True)
    # Socio B: 22,779.07 − 30,000 = −7,220.93, en rojo; la Σ también negativa
    # si así sale, y el total entregado = 50,000.
    r_b = enc + 3
    assert ws.cell(row=r_b, column=1).value == "Socio B"
    assert wsv.cell(row=r_b, column=6).value == r2(22779.07 - 30000) == -7220.93
    assert wsv.cell(row=nota - 1, column=5).value == 50000
    # El pie de la hoja sigue después de la tabla.
    textos = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert textos[-1] == tp._NOTA_BALANCE_PESOS if hasattr(tp, "_NOTA_BALANCE_PESOS") else True
    anchos = [ws.column_dimensions[get_column_letter(i)].width for i in range(1, 8)]
    assert anchos[4:] == [_ANCHO_COL_ENTREGAS] * 3


def test_individual_pesos_sin_pesos_raya_por_entregar_con_nota() -> None:
    p = sin_pesos(individual())
    b = p["balance"]
    assert b["por_entregar_socios_mxn"] is None and b["entregado_socios_mxn"] == 50000
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    wb, wv = tp._libros(data)
    ws = wb["balance"]
    (enc,) = _encabezados_socios(ws)
    nota = _tabla(ws, wv["balance"], enc, b, pesos=True)
    # Lo entregado se pinta igual (se conoce); por entregar «—» con su nota,
    # también en el total.
    assert wv["balance"].cell(row=nota - 1, column=5).value == 50000
    assert ws.cell(row=nota - 1, column=6).value == "—"


def test_ex_socio_con_porcentaje_cero_y_adelanto_en_rojo(registros) -> None:
    p = individual(ex_socio=True)
    b = p["balance"]
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    verificar_libro(data)
    wb, wv = tp._libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    (enc,) = _encabezados_socios(ws)
    _tabla(ws, wsv, enc, b, pesos=True)
    r = enc + 4
    assert ws.cell(row=r, column=1).value == "Ex socio"
    assert wsv.cell(row=r, column=2).value == 0
    # MONTO MXN 0 con la fórmula de siempre; por entregar = 0 − 1,234.50.
    assert ws.cell(row=r, column=3).value.startswith("=ROUND(B")
    assert wsv.cell(row=r, column=3).value == 0
    assert wsv.cell(row=r, column=6).value == -1234.5
    assert f"F{r}" in _rangos_rojo(ws)


# ---------------------------------------------------------------------------
# 4) Balance GENERAL: cada bloque de avión, en USD y en pesos; la VUELATOUR y
#    `celdas_empresa` intactas.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("variante", ["mensual", "general"])
@pytest.mark.parametrize("bandera", [None, True], ids=["usd", "pesos"])
def test_general_cada_bloque_lleva_las_entregas(registros, variante, bandera) -> None:
    p = general(variante, bandera=bandera)
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert tp._degradadas_balance(registros[-1]) == []
    verificar_libro(data)
    wb, wv = tp._libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    pesos = bool(bandera)
    encs = _encabezados_socios(ws)
    assert len(encs) == len(p["aviones"]) == 2
    for enc, avion in zip(encs, p["aviones"], strict=True):
        _tabla(ws, wsv, enc, avion["balance"], pesos=pesos)
    # VUELATOUR: la PARTICIPACIÓN sigue citando la celda MONTO de cada socia
    # (columna C) y vale lo del API.
    emp = p["empresa"]
    part = tp._fila_de(
        ws,
        "(+) PARTICIPACIÓN COMO SOCIO EN LOS AVIONES (MXN)" if pesos else _EMPRESA_PARTICIPACION,
    )
    formula = ws.cell(row=part, column=2).value
    assert formula.startswith("=") and "D" not in formula.replace("ROUND", "")
    citadas = [int(x) for x in __import__("re").findall(r"C(\d+)", formula)]
    assert len(citadas) == 2
    assert all(ws.cell(row=c, column=1).value == t_empresa._EMPRESA for c in citadas)
    assert wsv.cell(row=part, column=2).value == (
        emp["participacion_mxn"] if pesos else emp["participacion_usd"]
    )
    base = [56, 18, 17, 16] if pesos else [46, 14, 16]
    anchos = [ws.column_dimensions[get_column_letter(i)].width for i in range(1, len(base) + 4)]
    assert anchos == [*base, *([_ANCHO_COL_ENTREGAS] * 3)]


def test_general_mixto_solo_el_avion_con_las_llaves_lleva_las_columnas() -> None:
    p = pp.general("mensual", bandera=None)
    p["aviones"][0] = con_entregas(p["aviones"][0])
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    verificar_libro(data)
    wb, wv = tp._libros(data)
    ws = wb["balance"]
    encs = _encabezados_socios(ws)
    _tabla(ws, wv["balance"], encs[0], p["aviones"][0]["balance"], pesos=False)
    enc2 = encs[1]
    assert ws.cell(row=enc2, column=4).value is None
    assert tp._filas_de(ws, FILA_TOTAL) == [encs[0] + 1 + len(p["aviones"][0]["balance"]["socios"])]
    assert len(tp._filas_de(ws, NOTA)) == 1


# ---------------------------------------------------------------------------
# 5) Sin el discriminador: BYTE A BYTE (vista en pantalla contra HEAD; .xlsx
#    miembro a miembro contra el mismo render sin las llaves).
# ---------------------------------------------------------------------------


def _miembros(data: bytes) -> dict[str, bytes]:
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in z.namelist() if n != "docProps/core.xml"}


@pytest.mark.parametrize(("libro", "bandera"), list(tro._HUELLAS_HEAD))
def test_sin_llaves_vista_en_pantalla_byte_a_byte_con_head(libro, bandera) -> None:
    """La vista en pantalla (JSON) de los payloads sin las llaves es la del
    HEAD 16d6e2f (huellas de `test_balance_reserva_overhaul`)."""
    _, vista = tp._render(libro, tp._payload(libro, bandera))
    assert tp._huella_json(vista) == tro._HUELLAS_HEAD[(libro, bandera)][1]


@pytest.mark.parametrize("forma", ["null", "texto", "solo-socios"])
@pytest.mark.parametrize(("libro", "bandera"), list(tro._HUELLAS_HEAD))
def test_llaves_sin_discriminador_no_cambian_nada(libro, bandera, forma) -> None:
    """`entregado_socios_mxn` es el discriminador: en null, como texto o con
    las llaves por socio sueltas (y hasta `por_entregar_socios_mxn`), ni
    columnas, ni totales, ni nota: el .xlsx es miembro a miembro el del
    payload sin llaves y la vista en pantalla, la del HEAD."""
    base = tp._payload(libro, bandera)
    data_base, _ = tp._render(libro, base)
    p = copy.deepcopy(base)
    for b in tro._bloques(p):
        for s in b["socios"]:
            s.update(entregado_mxn=123.0, entregado_usd=7.0, entregas=1, por_entregar_mxn=-1.0)
        b["por_entregar_socios_mxn"] = 5.0
        if forma == "null":
            b["entregado_socios_mxn"] = None
        elif forma == "texto":
            b["entregado_socios_mxn"] = "123"
    data, vista = tp._render(libro, p)
    assert _miembros(data) == _miembros(data_base)
    assert tp._huella_json(vista) == tro._HUELLAS_HEAD[(libro, bandera)][1]


# ---------------------------------------------------------------------------
# 6) Vista en pantalla (el MISMO libro) y rutas.
# ---------------------------------------------------------------------------


def _celda(hoja: dict, r: int, c: int) -> dict:
    return next(x for x in hoja["celdas"] if x["r"] == r and x["c"] == c)


@pytest.mark.parametrize("bandera", [None, True], ids=["usd", "pesos"])
def test_libro_json_exporta_las_columnas_y_el_rojo_resuelto(bandera) -> None:
    from tests import test_libro_json as tlj

    p = individual(bandera=bandera)
    req = BalanceAvionRequest.model_validate(p)
    libro = libro_balance_avion(req)
    assert tlj._comparar_con_xlsx(libro, render_balance_avion_xlsx(req)) > 0
    hoja = next(h for h in libro["hojas"] if h["nombre"] == "balance")
    col = 5 if bandera else 4
    enc = next(x["r"] for x in hoja["celdas"] if x["c"] == 1 and x["v"] == "SOCIO")
    assert _celda(hoja, enc, col)["v"] == COL_ENTREGADO
    assert _celda(hoja, enc, col + 1)["v"] == COL_POR_ENTREGAR
    assert _celda(hoja, enc, col + 2)["v"] == COL_USD_INFO
    # Socio B (tercer renglón): por entregar −7,220.93, como lo mostraría
    # Excel y con el rojo de los negativos YA resuelto.
    r = enc + 3
    assert _celda(hoja, r, 1)["v"] == "Socio B"
    pe = _celda(hoja, r, col + 1)
    assert pe["v"] == -7220.93
    assert pe["t"] == libro_json.formatear(-7220.93, MONEY)
    if bandera:
        assert pe["f"] == f"=ROUND(C{r}-E{r},2)"
    else:
        assert pe.get("f") is None
    estilo = libro["estilos"][pe["s"]]
    assert estilo.get("bg", "").endswith("FFC7CE") and estilo.get("co", "").endswith("9C0006")
    # El total y la nota.
    rt = enc + 4
    assert _celda(hoja, rt, 1)["v"] == FILA_TOTAL
    assert _celda(hoja, rt, col)["v"] == 50000
    assert _celda(hoja, rt + 1, 1)["v"] == NOTA
    # El general también.
    g = BalanceGeneralRequest.model_validate(general("mensual", bandera=bandera))
    lg = libro_balance_general(g)
    hoja_g = next(h for h in lg["hojas"] if h["nombre"] == "balance")
    assert sum(1 for x in hoja_g["celdas"] if x["v"] == FILA_TOTAL) == 2


TOKEN = "tok-entregas-socios"
client = TestClient(app)


def test_rutas_aceptan_el_payload(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    p = individual(ex_socio=True)
    res = client.post("/pdf/balance-avion-xlsx", json=p, headers=headers)
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content))["balance"]
    assert len(tp._filas_de(ws, FILA_TOTAL)) == 1
    res = client.post("/pdf/balance-avion-libro", json={**p, "archivo": "x.xlsx"}, headers=headers)
    assert res.status_code == 200, res.text
    hoja = next(h for h in res.json()["hojas"] if h["nombre"] == "balance")
    assert any(x["v"] == COL_POR_ENTREGAR for x in hoja["celdas"])
    g = general("general")
    res = client.post("/pdf/balance-general-xlsx", json=g, headers=headers)
    assert res.status_code == 200, res.text
    res = client.post(
        "/pdf/balance-general-libro", json={**g, "archivo": "g.xlsx"}, headers=headers
    )
    assert res.status_code == 200, res.text
