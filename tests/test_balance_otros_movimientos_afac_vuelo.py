"""AFAC POR VUELO en 'otros movimientos' del Balance general (16-oct-2026,
API 0.0.89, invariante 90 del API).

Pedido del cliente: «lo que necesito Urgente es lo del AFAC […] debe aparecer
en balance general en la hoja de Otros movimientos para poder confirmar el
monto, estaría bien que estuvieran 2 columnas juntas, una que sea lo
reservado de afac o cotizado y otra de lo que realmente costó y que en esa
celda poder poner el monto y se registre sin tenerte que ir a otro lado». Y
después: «la columna nueva que te pedí mejor que sea en la columna de
egreso». Decisiones del cliente del 16-oct: UNA fila por vuelo; el avión sigue
pagando lo RESERVADO y el COSTO REAL es egreso de VuelaTour; se escribe solo
el monto del vuelo; los gastos marcados «Pago de AFAC» dejan de restar al
avión y se comparan cada mes.

El API manda las filas `seccion: 'AFAC_VUELO'` contiguas al final de `filas`
(EGRESO = costo real, INGRESO = reservado, remanente hecho), la bandera
`afac_por_vuelo` y el resumen `afac_resumen`. Aquí no se calcula nada: las
filas se pintan con `pinta()` de siempre (y SUMAN en TOTALES con las fórmulas
de siempre); se agregan la fila separadora, el vínculo editable `afac_real`
en la celda EGRESO del libro JSON, el bloque «AFAC DEL PERIODO» bajo TOTALES
(fórmulas verificadas), la leyenda de la fila 2, los textos del bloque
VUELATOUR y, en el libro individual, la línea de los pagos de AFAC en
'permisos'. SIN las llaves nuevas: libros BYTE A BYTE (huellas del HEAD
575c816 con los payloads de forma real del API 0.0.88).
"""

from __future__ import annotations

import copy
import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAfacPago,
    BalanceAfacResumen,
    BalanceAfacVuelo,
    BalanceAvionRequest,
    BalanceGeneralRequest,
    BalanceHojaOtrosMovimientos,
    BalanceOtroMovimientoFila,
)
from app.services.balance_avion_xlsx import (
    _AFAC_PERIODO_DIFERENCIA,
    _AFAC_PERIODO_FILAS,
    _AFAC_PERIODO_PAGADO,
    _AFAC_PERIODO_PAGO,
    _AFAC_PERIODO_SIN_PAGOS,
    _EMPRESA_AFAC_EN_FILAS,
    _EMPRESA_AFAC_POR_VUELO,
    _EMPRESA_INGRESOS,
    _EMPRESA_INGRESOS_AFAC,
    _EMPRESA_INGRESOS_AFAC_VUELO,
    _EMPRESA_PAGOS,
    _EMPRESA_PAGOS_AFAC,
    _EMPRESA_PAGOS_AFAC_VUELO,
    _LEYENDA_AFAC_PERIODO,
    _LEYENDA_PROVISION,
    _MARCA_PROVISION,
    _NOTA_BLOQUE_EMPRESA,
    _NOTA_EMPRESA_BASE,
    _OM_AFAC_EN_FILAS,
    _OM_AFAC_POR_VUELO,
    _OM_SECCION_AFAC,
    _OM_SECCION_AFAC_VUELO,
    _TITULO_AFAC_OM,
    _TITULO_AFAC_PERIODO,
    FILL_AFAC_EDITABLE,
    LIGHT,
    MONEY,
    NAVY,
    RED,
    _vinculo_afac_real,
    libro_balance_avion,
    libro_balance_general,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from tests import _afac_vuelo_payloads as pa
from tests import _cierre_mes_payloads as pay
from tests import test_balance_empresa_vuelatour as tev
from tests import test_balance_formulas as tf
from tests import test_balance_otros_movimientos_afac as t_afac
from tests._evaluador_formulas import verificar_libro
from tests.test_cierre_mes import _HUELLAS_HEAD, _balance, _huella_json, _huella_xlsx
from tests.test_comision_vendedor_notas import OM_FILA2_VIEJA, _rellenos_amber

r2 = tf.r2
TOKEN = "secreto-de-prueba"
client = TestClient(app)
OM = "otros movimientos"

# ---------------------------------------------------------------------------
# Huellas «de siempre»: calculadas con el código del HEAD 575c816 (API 0.0.88
# en producción, ANTES de este cambio) corriendo los constructores de
# `_afac_vuelo_payloads` (forma real del 0.0.88: cierre de mes + filas «AFAC
# · matrícula» con `afac_en_filas`) — .xlsx: todos los miembros salvo
# docProps/core.xml; vista en pantalla: el JSON con `vinculos_gasto: true`.
# Si una se mueve, el API 0.0.88 (el de hoy) ya no ve su libro de siempre.
# ---------------------------------------------------------------------------
_HUELLAS_0088 = {
    "general-mensual": (
        "266a7bf7d10bef7858bbcff8e39e168391b14fd5e3dc66b1b075ee493a267cda",
        "562f629d32fba5f80c8373ec4891f985fa72e78ce296eb9d4b12b2b548bce1a0",
    ),
    "general-general": (
        "59de23fed99e7a2ad37f4674f5bd59507e3b11fede2d06b51bc3fc440bacb6de",
        "7a920bd714e9de249ebee0a8a3454dc54b7444db7f8ecd347d168942b885291e",
    ),
    "individual": (
        "5ff8072c668e207a77822c0463cc24ee31a19b329873b2d5412156ab57f6e7e2",
        "d4a163ef9f67f77bc96ae06032a60bc94e3e2a31b7ef12d778a36e0dd7e2a866",
    ),
}


def _payload_0088(caso: str) -> dict:
    if caso == "individual":
        return pa.individual_0088()
    return pa.general_0088(caso.split("-")[1])


def _libros(data: bytes):
    """(con fórmulas, con caché)."""
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _general(p: dict | None = None) -> tuple[bytes, object, object]:
    data = render_balance_general_xlsx(
        BalanceGeneralRequest.model_validate(pa.general_0089() if p is None else p)
    )
    wb, wv = _libros(data)
    return data, wb[OM], wv[OM]


def _filas_de(ws, texto: str) -> list[int]:
    return [c.row for c in ws["A"] if c.value == texto]


def _fila_de(ws, texto: str) -> int:
    filas = _filas_de(ws, texto)
    assert len(filas) == 1, (texto, filas)
    return filas[0]


def _textos(ws) -> list[str]:
    return [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]


def _hoja_json(libro: dict, nombre: str) -> dict:
    return next(h for h in libro["hojas"] if h["nombre"] == nombre)


def _celda_json(hoja: dict, r: int, c: int) -> dict | None:
    return next((x for x in hoja["celdas"] if x["r"] == r and x["c"] == c), None)


def _libro_general(p: dict) -> dict:
    req = BalanceGeneralRequest.model_validate({**p, "vinculos_gasto": True})
    return libro_balance_general(req)


# ---------------------------------------------------------------------------
# (a) Textos del contrato.
# ---------------------------------------------------------------------------


def test_textos_son_los_del_contrato() -> None:
    assert _OM_SECCION_AFAC_VUELO == (
        "AFAC POR VUELO · RESERVADO (lo aporta el avión) vs COSTO REAL (lo paga "
        "VuelaTour a AFAC; escríbelo en la columna egreso del Balance general en "
        "pantalla)"
    )
    assert _EMPRESA_INGRESOS_AFAC_VUELO == (
        "(+) INGRESOS PROPIOS USD (TUAs, extras, pernocta, comisión — cotizado del "
        "periodo; AFAC reservado por vuelo)"
    )
    assert _EMPRESA_PAGOS_AFAC_VUELO == (
        "(−) EGRESOS PROPIOS USD (pago al vendedor, TUAs, extensión, comisión "
        "bancaria, sueltos; costo real AFAC por vuelo)"
    )
    # Ningún texto nuevo activa el amarillo de la provisión ni lleva la marca
    # «sin TC» (que en el API vacía un lado del bloque VUELATOUR); tampoco los
    # textos de las filas del payload de forma real.
    for texto in (
        _OM_SECCION_AFAC_VUELO,
        _OM_AFAC_POR_VUELO,
        _EMPRESA_INGRESOS_AFAC_VUELO,
        _EMPRESA_PAGOS_AFAC_VUELO,
        _EMPRESA_AFAC_POR_VUELO,
        _TITULO_AFAC_PERIODO,
        _LEYENDA_AFAC_PERIODO,
        *(str(v) for f in pa.filas_afac_vuelo() for v in f.values()),
    ):
        assert not _MARCA_PROVISION.search(texto), texto
        assert "sin TC" not in texto, texto
    # Los textos nuevos NO son los de las filas por avión del 0.0.74.
    assert _OM_SECCION_AFAC_VUELO != _OM_SECCION_AFAC
    assert _TITULO_AFAC_PERIODO != _TITULO_AFAC_OM


# ---------------------------------------------------------------------------
# (b) Fila separadora y filas AFAC por vuelo pintadas como filas NORMALES:
#     EGRESO (D) = costo real, INGRESO (G) = reservado, I = ROUND(G-D,2).
# ---------------------------------------------------------------------------


def test_separador_y_filas_en_las_columnas_de_siempre() -> None:
    data, om, omv = _general()
    sep = _fila_de(om, _OM_SECCION_AFAC_VUELO)
    # La fila separadora va INMEDIATAMENTE después de la última fila por
    # vuelo (#402 de 'filas') y antes de la primera AFAC (sin fila en blanco).
    vuelo_402 = _filas_de(om, "#402 · Cliente Dos")
    assert vuelo_402[0] + 1 == sep and _filas_de(om, "#401 · Cliente Uno")[1] == sep + 1
    a = om.cell(row=sep, column=1)
    assert a.font.b is True and a.font.sz == 10 and a.font.color.rgb == f"00{NAVY}"
    for col in range(1, 11):
        celda = om.cell(row=sep, column=col)
        assert celda.fill.fill_type == "solid" and celda.fill.fgColor.rgb == f"00{LIGHT}"
        if col > 1:
            assert celda.value is None
    assert not any(r.min_row <= sep <= r.max_row for r in om.merged_cells.ranges)
    # Una sola separadora, aunque haya 4 filas AFAC.
    assert _textos(om).count(_OM_SECCION_AFAC_VUELO) == 1
    f401, f402, f407, f405 = sep + 1, sep + 2, sep + 3, sep + 4
    filas = pa.filas_afac_vuelo()
    for r, f in zip((f401, f402, f407, f405), filas, strict=True):
        assert om.cell(row=r, column=1).value == f["clave"]
        assert om.cell(row=r, column=3).value == f["concepto_egreso"]
        assert om.cell(row=r, column=4).value == f["egreso_mxn"]
        assert om.cell(row=r, column=5).value is None
        assert om.cell(row=r, column=6).value == f["concepto_ingreso"]
        assert om.cell(row=r, column=7).value == f["ingreso_mxn"]
        assert om.cell(row=r, column=8).value is None
        assert om.cell(row=r, column=9).value == f"=ROUND(G{r}-D{r},2)"
        assert omv.cell(row=r, column=9).value == f["remanente_mxn"]
        assert om.cell(row=r, column=10).value is None
    # Fecha del vuelo en B; swatch del avión en A (XB-DOS sin color: nada).
    assert om.cell(row=f401, column=2).value == "03/09/2026"
    assert om.cell(row=f401, column=1).fill.fgColor.rgb == "003B82F6"
    assert om.cell(row=f407, column=1).fill.fill_type is None
    # Sin costo real: D VACÍA (no 0) con el formato de dinero; el concepto
    # dice «por capturar».
    d402 = om.cell(row=f402, column=4)
    assert d402.value is None and d402.number_format == MONEY
    assert om.cell(row=f402, column=3).value == pa.CONCEPTO_POR_CAPTURAR
    # Notas: la del reservado en G (cómo se calculó) y las notas del costo
    # real en D.
    assert om.cell(row=f401, column=7).comment.text == filas[0]["nota_ingreso"]
    assert om.cell(row=f401, column=4).comment.text == "Recibo del gestor 0412"
    assert om.cell(row=f402, column=4).comment is None
    assert verificar_libro(data)


def test_celda_egreso_editable_en_verde_y_la_del_reservado_no() -> None:
    _, om, _ = _general()
    sep = _fila_de(om, _OM_SECCION_AFAC_VUELO)
    for r in range(sep + 1, sep + 5):
        d = om.cell(row=r, column=4)
        assert d.fill.fill_type == "solid" and d.fill.fgColor.rgb == f"00{FILL_AFAC_EDITABLE}"
        assert d.number_format == MONEY
        assert om.cell(row=r, column=7).fill.fill_type is None


def test_totales_con_las_formulas_de_siempre_suman_real_y_reservado_una_vez() -> None:
    data, om, omv = _general()
    tot = _fila_de(om, "TOTALES")
    assert om[f"D{tot}"].value == f"=ROUND(SUM(D4:D{tot - 1}),2)"
    assert om[f"G{tot}"].value == f"=ROUND(SUM(G4:G{tot - 1}),2)"
    assert om[f"I{tot}"].value == f"=ROUND(G{tot}-D{tot},2)"
    # Egresos: TUA pagado del #401 + los costos reales capturados (el #402
    # vacío no suma). Ingresos: TUA, pernocta, los reservados y la suelta.
    egresos = r2(1100.25 + 1000 + 300 + 700)
    ingresos = r2(1250.5 + 3000 + 1082.5 + 564.06 + 660 + 85.33)
    assert omv[f"D{tot}"].value == egresos
    assert omv[f"G{tot}"].value == ingresos
    assert omv[f"I{tot}"].value == r2(ingresos - egresos)
    assert verificar_libro(data)


def test_los_numeros_no_dependen_de_la_bandera() -> None:
    """Skew: sin `afac_por_vuelo` (o con un pyservices previo) las filas se
    pintan y SUMAN igual —solo cambian el separador, la fila 2, el verde y el
    vínculo—: TOTALES idénticos."""
    _, om, omv = _general()
    p = pa.general_0089()
    p["consolidado"]["otros_movimientos"].pop("afac_por_vuelo")
    _, om_sin, omv_sin = _general(p)
    tot, tot_sin = _fila_de(om, "TOTALES"), _fila_de(om_sin, "TOTALES")
    assert tot_sin == tot - 1  # sin la separadora
    for col in "DGI":
        assert omv[f"{col}{tot}"].value == omv_sin[f"{col}{tot_sin}"].value
    assert _OM_SECCION_AFAC_VUELO not in _textos(om_sin)
    primera_afac = _filas_de(om_sin, "#401 · Cliente Uno")[1]
    assert primera_afac == _filas_de(om_sin, "#402 · Cliente Dos")[0] + 1
    assert om_sin.cell(row=primera_afac, column=4).fill.fill_type is None


def test_remanente_negativo_en_rojo_y_nada_en_amarillo() -> None:
    data, om, _ = _general()
    assert _rellenos_amber(om) == []
    assert _LEYENDA_PROVISION not in _textos(om)
    vista = _hoja_json(_libro_general(pa.general_0089()), OM)
    sep = _fila_de(om, _OM_SECCION_AFAC_VUELO)
    estilos = _libro_general(pa.general_0089())["estilos"]
    rojo = _celda_json(vista, sep + 3, 9)  # #407: −300.00
    assert rojo["v"] == -300.0 and estilos[rojo["s"]]["co"] == "9C0006"
    verde = _celda_json(vista, sep + 1, 9)  # #401: 82.50
    assert estilos[verde["s"]].get("co") != "9C0006"


def test_notas_del_costo_real_con_la_marca_de_provision_no_pintan_amarillo() -> None:
    """Revisión (16-oct-2026): `nota_egreso` de una fila AFAC_VUELO son las
    NOTAS LIBRES del operador (PUT /v1/flights/:id/afac-real), no un texto
    del API. Una nota con «· PROVISIÓN (» no puede volver amarilla la celda
    editable ni sacar la leyenda del amarillo (una fila de vuelo con la
    marca del API, sí: el amarillo de siempre no cambia)."""
    p = pa.general_0089()
    om_p = p["consolidado"]["otros_movimientos"]
    afac = [f for f in om_p["filas"] if f.get("seccion") == "AFAC_VUELO"]
    afac[0]["nota_egreso"] = "Recibo · PROVISIÓN (el gestor manda la factura después)"
    afac[0]["afac"]["notas"] = afac[0]["nota_egreso"]
    data, om, _ = _general(p)
    assert _rellenos_amber(om) == []
    assert _LEYENDA_PROVISION not in _textos(om)
    sep = _fila_de(om, _OM_SECCION_AFAC_VUELO)
    d = om.cell(row=sep + 1, column=4)
    assert d.fill.fgColor.rgb == f"00{FILL_AFAC_EDITABLE}"
    assert d.comment.text == afac[0]["nota_egreso"]
    # Control: la MISMA marca en una fila de vuelo normal sí pinta amarillo.
    normal = next(f for f in om_p["filas"] if f.get("seccion") is None and f.get("egreso_mxn"))
    normal["concepto_egreso"] = "pago comisión vendedor (X) · PROVISIÓN (mismo monto)"
    _, om2, _ = _general(p)
    assert _rellenos_amber(om2) != []
    assert _LEYENDA_PROVISION in _textos(om2)


# ---------------------------------------------------------------------------
# (c) Libro JSON: el vínculo editable `afac_real` en la celda EGRESO (D).
# ---------------------------------------------------------------------------


def test_vinculo_afac_real_en_la_celda_egreso_y_no_en_el_reservado() -> None:
    _, om, _ = _general()
    sep = _fila_de(om, _OM_SECCION_AFAC_VUELO)
    vista = _hoja_json(_libro_general(pa.general_0089()), OM)
    d401 = _celda_json(vista, sep + 1, 4)
    assert d401["v"] == 1000.0 and d401["t"] == "1,000.00"
    assert d401["vinculo"] == {
        "tipo": "afac_real",
        "campo": "costo_real_mxn",
        "vuelo_id": pa.VUELO_401,
        "clave": "#401 · Cliente Uno",
        "matricula": "XA-TST",
        "valor": 1000.0,
        "reservado_mxn": 1082.5,
        "editable": True,
        "motivo": None,
        "mes": "2026-09",
        "capturado_por": "Ale",
        "capturado_at": pa.CAPTURADO_AT,
    }
    # Sin capturar: la celda viaja vacía con `valor` null.
    d402 = _celda_json(vista, sep + 2, 4)
    assert d402["v"] is None and d402["t"] == ""
    assert d402["vinculo"]["valor"] is None and d402["vinculo"]["vuelo_id"] == pa.VUELO_402
    assert d402["vinculo"]["capturado_por"] is None
    # Solo costo real (avión sin tarifa): reservado null.
    assert _celda_json(vista, sep + 3, 4)["vinculo"]["reservado_mxn"] is None
    # Ninguna otra celda de la fila lleva vínculo (G = reservado, sobre todo).
    for r in range(sep + 1, sep + 5):
        for c in (1, 2, 3, 5, 6, 7, 8, 9, 10):
            x = _celda_json(vista, r, c)
            assert x is None or ("vinculo" not in x and "vinculos" not in x), (r, c)
    # Las filas por vuelo de siempre y TOTALES no llevan `afac_real`.
    tipos = [x["vinculo"]["tipo"] for x in vista["celdas"] if "vinculo" in x]
    assert tipos == ["afac_real"] * 4
    json.dumps(vista)


def test_fila_no_editable_lleva_motivo_y_sin_verde() -> None:
    """Fila no «dueña» de un vuelo multi-avión, mes cerrado o migración sin
    aplicar: `editable: false` con su motivo; la celda NO se pinta de verde
    y el vínculo viaja igual (el panel la deja de solo lectura con el
    `motivo` como explicación)."""
    filas = pa.filas_afac_vuelo()
    motivo = "Septiembre 2026 ya está cerrado (Ale, 16-oct). Reábrelo en Reportes para corregir."
    for f in filas:
        f["afac"]["editable"] = False
        f["afac"]["motivo_no_editable"] = motivo
    p = pa.general_0089(filas=filas)
    _, om, _ = _general(p)
    sep = _fila_de(om, _OM_SECCION_AFAC_VUELO)
    vista = _hoja_json(_libro_general(p), OM)
    for r in range(sep + 1, sep + 5):
        assert om.cell(row=r, column=4).fill.fill_type is None
        v = _celda_json(vista, r, 4)["vinculo"]
        assert v["editable"] is False and v["motivo"] == motivo


def test_sin_vuelo_id_no_hay_vinculo_y_editable_solo_con_true() -> None:
    base = BalanceOtroMovimientoFila.model_validate(pa.filas_afac_vuelo()[0])
    assert _vinculo_afac_real(base)["editable"] is True
    # `editable` solo abre la celda con un `true` JSON (jamás un 422).
    raro = copy.deepcopy(pa.filas_afac_vuelo()[0])
    raro["afac"]["editable"] = "true"
    assert _vinculo_afac_real(BalanceOtroMovimientoFila.model_validate(raro))["editable"] is False
    # Sin `vuelo_id` (o en blanco) no hay celda editable.
    sin_id = copy.deepcopy(pa.filas_afac_vuelo()[0])
    sin_id["afac"]["vuelo_id"] = "   "
    assert _vinculo_afac_real(BalanceOtroMovimientoFila.model_validate(sin_id)) is None
    sin_afac = copy.deepcopy(pa.filas_afac_vuelo()[0])
    sin_afac.pop("afac")
    assert _vinculo_afac_real(BalanceOtroMovimientoFila.model_validate(sin_afac)) is None
    p = pa.general_0089(filas=[sin_id])
    _, om, _ = _general(p)
    fila = _fila_de(om, _OM_SECCION_AFAC_VUELO) + 1
    assert om.cell(row=fila, column=4).fill.fill_type is None
    vista = _hoja_json(_libro_general(p), OM)
    assert "vinculo" not in _celda_json(vista, fila, 4)


# ---------------------------------------------------------------------------
# (d) Bloque «AFAC DEL PERIODO» debajo de TOTALES.
# ---------------------------------------------------------------------------


def test_bloque_afac_del_periodo_con_formulas_verificadas() -> None:
    data, om, omv = _general()
    sep = _fila_de(om, _OM_SECCION_AFAC_VUELO)
    a, b = sep + 1, sep + 4
    tot = _fila_de(om, "TOTALES")
    titulo = _fila_de(om, _TITULO_AFAC_PERIODO)
    assert titulo == tot + 2
    assert om.cell(row=titulo, column=1).font.color.rgb == f"00{NAVY}"
    assert om.cell(row=titulo + 1, column=1).value == _LEYENDA_AFAC_PERIODO
    assert any(
        r.min_row == titulo + 1 and r.min_col == 1 and r.max_col == 10
        for r in om.merged_cells.ranges
    )
    sub = _fila_de(om, _AFAC_PERIODO_FILAS)
    assert sub == titulo + 2
    assert om[f"D{sub}"].value == f"=ROUND(SUM(D{a}:D{b}),2)"
    assert om[f"G{sub}"].value == f"=ROUND(SUM(G{a}:G{b}),2)"
    assert om[f"I{sub}"].value == f"=ROUND(SUM(I{a}:I{b}),2)"
    assert (omv[f"D{sub}"].value, omv[f"G{sub}"].value, omv[f"I{sub}"].value) == (
        2000.0,
        2306.56,
        306.56,
    )
    assert om[f"C{sub}"].value == "costo real capturado: 3 de 4 vuelos · 1 por capturar"
    assert om[f"F{sub}"].value == "reservado (lo aporta el avión) · 3 vuelos"
    # Un renglón por pago: fecha en E, detalle en C, monto en D.
    pagos = _filas_de(om, _AFAC_PERIODO_PAGO)
    assert pagos == [sub + 1, sub + 2, sub + 3]
    assert [om[f"C{r}"].value for r in pagos] == [x["detalle"] for x in pa.pagos_afac()]
    assert [om[f"D{r}"].value for r in pagos] == [1500.0, 750.0, None]
    assert [om[f"E{r}"].value for r in pagos] == ["01/09/2026", "10/09/2026", "22/09/2026"]
    assert om[f"D{pagos[2]}"].comment.text.startswith("Pago en USD 40.00 sin T.C.")
    pag = _fila_de(om, _AFAC_PERIODO_PAGADO)
    assert pag == pagos[-1] + 1
    assert om[f"D{pag}"].value == f"=ROUND(SUM(D{pagos[0]}:D{pagos[-1]}),2)"
    assert omv[f"D{pag}"].value == 2250.0
    assert om[f"C{pag}"].value == (
        "3 gastos marcados «Pago de AFAC» · 1 en USD sin T.C. (no suma)"
    )
    dif = _fila_de(om, _AFAC_PERIODO_DIFERENCIA)
    assert om[f"D{dif}"].value == f"=ROUND(D{pag}-D{sub},2)"
    assert omv[f"D{dif}"].value == 250.0
    assert om[f"D{dif}"].font.color.rgb == f"00{RED}"
    # El bloque va DEBAJO de TOTALES: sus sumas no lo incluyen.
    assert om[f"D{tot}"].value == f"=ROUND(SUM(D4:D{tot - 1}),2)"
    # Cada fórmula reproduce el número del API (ninguna se degradó).
    assert verificar_libro(data)


def test_pagos_del_bloque_llevan_su_vinculo_de_gasto() -> None:
    _, om, _ = _general()
    pagos = _filas_de(om, _AFAC_PERIODO_PAGO)
    vista = _hoja_json(_libro_general(pa.general_0089()), OM)
    assert [_celda_json(vista, r, 4)["vinculos"] for r in pagos] == [
        [{"tipo": "gasto", "id": pa.GASTO_1, "rol": "fila", "monto_mxn": 1500.0,
          "columna": "pagado a AFAC"}],
        [{"tipo": "gasto", "id": pa.GASTO_2, "rol": "fila", "monto_mxn": 750.0,
          "columna": "pagado a AFAC"}],
        [{"tipo": "gasto", "id": pa.GASTO_3, "rol": "sin_tc", "monto_mxn": None,
          "columna": "pagado a AFAC"}],
    ]  # fmt: skip


def test_bloque_sin_pagos_y_diferencia_cero_sin_rojo() -> None:
    p = pa.general_0089()
    res = p["consolidado"]["otros_movimientos"]["afac_resumen"]
    res.update(pagos=[], pagado_mxn=0.0, pagos_sin_tc=0, diferencia_pagado_mxn=-2000.0)
    data, om, omv = _general(p)
    pago = _fila_de(om, _AFAC_PERIODO_PAGO)
    assert om[f"C{pago}"].value == _AFAC_PERIODO_SIN_PAGOS and om[f"D{pago}"].value is None
    pag = _fila_de(om, _AFAC_PERIODO_PAGADO)
    assert om[f"D{pag}"].value == 0.0  # valor del API, sin fórmula
    assert om[f"C{pag}"].value == "0 gastos marcados «Pago de AFAC»"
    dif = _fila_de(om, _AFAC_PERIODO_DIFERENCIA)
    assert omv[f"D{dif}"].value == -2000.0
    assert verificar_libro(data)
    # Diferencia 0 (pagado = costo real): negro.
    p = pa.general_0089()
    res = p["consolidado"]["otros_movimientos"]["afac_resumen"]
    res["pagos"][1]["monto_mxn"] = 500.0
    res.update(pagado_mxn=2000.0, diferencia_pagado_mxn=0.0)
    data, om, omv = _general(p)
    dif = _fila_de(om, _AFAC_PERIODO_DIFERENCIA)
    assert omv[f"D{dif}"].value == 0.0
    assert om[f"D{dif}"].font.color.rgb != f"00{RED}"
    assert verificar_libro(data)


def test_bloque_solo_con_pagos_sin_filas_afac() -> None:
    """Mes con pagos marcados pero sin vuelos con AFAC: el API manda
    `afac_resumen` sin `afac_por_vuelo`. Ni separador ni leyenda en la fila
    2; el bloque se pinta con los totales del API como VALOR (no hay filas
    que sumar) y los pagos con su vínculo."""
    p = pa.general_0089(filas=[])
    res = p["consolidado"]["otros_movimientos"]["afac_resumen"]
    res.update(vuelos=0, vuelos_con_reservado=0, vuelos_con_real=0, vuelos_sin_real=0,
               reservado_mxn=0.0, costo_real_mxn=0.0, remanente_mxn=0.0,
               diferencia_pagado_mxn=2250.0)  # fmt: skip
    assert "afac_por_vuelo" not in p["consolidado"]["otros_movimientos"]
    data, om, omv = _general(p)
    assert _OM_SECCION_AFAC_VUELO not in _textos(om)
    assert _OM_AFAC_POR_VUELO not in om["A2"].value
    sub = _fila_de(om, _AFAC_PERIODO_FILAS)
    assert [om[f"{c}{sub}"].value for c in "DGI"] == [0.0, 0.0, 0.0]
    assert om[f"C{sub}"].value == "costo real capturado: 0 de 0 vuelos"
    dif = _fila_de(om, _AFAC_PERIODO_DIFERENCIA)
    assert omv[f"D{dif}"].value == 2250.0
    assert verificar_libro(data)


def test_filas_no_contiguas_sin_separador_y_bloque_como_valor() -> None:
    """Defensa (el API siempre las manda contiguas): si una fila por vuelo se
    cuela entre las AFAC, se pintan como filas normales (suman igual), sin
    separador, con su vínculo, y el bloque usa los valores del API."""
    p = pa.general_0089()
    om_p = p["consolidado"]["otros_movimientos"]
    vuelo = om_p["filas"].pop(1)  # «#402 · Cliente Dos» (pernocta)
    om_p["filas"].insert(3, vuelo)  # entre #401 AFAC y #402 AFAC
    data, om, omv = _general(p)
    assert _OM_SECCION_AFAC_VUELO not in _textos(om)
    sub = _fila_de(om, _AFAC_PERIODO_FILAS)
    assert [omv[f"{c}{sub}"].value for c in "DGI"] == [2000.0, 2306.56, 306.56]
    assert all(not str(om[f"{c}{sub}"].value).startswith("=") for c in "DGI")
    tot = _fila_de(om, "TOTALES")
    assert omv[f"D{tot}"].value == r2(1100.25 + 1000 + 300 + 700)
    vista = _hoja_json(_libro_general(p), OM)
    assert sum(1 for x in vista["celdas"] if x.get("vinculo", {}).get("tipo") == "afac_real") == 4
    assert verificar_libro(data)


def test_fuera_de_libro_solo_si_hay() -> None:
    p = pa.general_0089()
    p["consolidado"]["otros_movimientos"]["afac_resumen"]["reales_fuera_de_libro_mxn"] = 450.0
    data, om, omv = _general(p)
    fuera = _fila_de(om, "fuera de libro")
    assert fuera == _fila_de(om, _AFAC_PERIODO_DIFERENCIA) + 1
    assert omv[f"D{fuera}"].value == 450.0
    assert verificar_libro(data)
    _, om, _ = _general()
    assert _filas_de(om, "fuera de libro") == []


# ---------------------------------------------------------------------------
# (e) Fila 2 y payload mixto (las dos banderas: nunca las manda el API).
# ---------------------------------------------------------------------------


def _con_afac_vuelo_tf(*, por_vuelo: bool | None = True) -> BalanceGeneralRequest:
    """El general de `test_balance_formulas` (sin banderas de fila 2) con las
    filas AFAC por vuelo y el resumen."""
    p = tf._general().model_dump(mode="json")
    om = p["consolidado"]["otros_movimientos"]
    om["filas"].extend(pa.filas_afac_vuelo())
    if por_vuelo is not None:
        om["afac_por_vuelo"] = por_vuelo
    om["afac_resumen"] = pa.resumen_afac()
    return BalanceGeneralRequest.model_validate(p)


def test_fila2_con_bandera_lleva_la_leyenda_al_final() -> None:
    wb = load_workbook(BytesIO(render_balance_general_xlsx(_con_afac_vuelo_tf())))
    assert wb[OM]["A2"].value == OM_FILA2_VIEJA + _OM_AFAC_POR_VUELO


@pytest.mark.parametrize("bandera", [None, False], ids=["ausente", "false"])
def test_fila2_sin_bandera_la_de_siempre(bandera) -> None:
    wb = load_workbook(BytesIO(render_balance_general_xlsx(_con_afac_vuelo_tf(por_vuelo=bandera))))
    assert wb[OM]["A2"].value == OM_FILA2_VIEJA


def test_fila2_va_antes_de_otros_ingresos_por_avion() -> None:
    """Con el payload de forma real (trae `otros_ingresos_por_avion`): la
    leyenda AFAC va en el MISMO lugar que la de las filas por avión."""
    _, om, _ = _general()
    _, om88, _ = _general(pa.general_0088("mensual"))
    assert om["A2"].value == om88["A2"].value.replace(_OM_AFAC_EN_FILAS, _OM_AFAC_POR_VUELO)


def test_payload_mixto_manda_afac_por_vuelo() -> None:
    """Las dos banderas, filas «AFAC · matrícula» y bloque informativo viejo
    (ningún API lo manda): ni encabezado por avión, ni bloque informativo, ni
    la leyenda vieja; la separadora y el bloque del periodo sí."""
    p = pa.general_0089()
    om_p = p["consolidado"]["otros_movimientos"]
    om_p["afac_en_filas"] = True
    om_p["afac"] = t_afac.BLOQUE_AFAC_VIEJO
    om_p["filas"][2:2] = copy.deepcopy([t_afac.FILA_AFAC_TST])
    _, om, _ = _general(p)
    textos = _textos(om)
    assert _OM_SECCION_AFAC not in textos and _TITULO_AFAC_OM not in textos
    assert "SUBTOTAL AFAC" not in textos
    assert _OM_AFAC_EN_FILAS not in om["A2"].value and _OM_AFAC_POR_VUELO in om["A2"].value
    assert _fila_de(om, _OM_SECCION_AFAC_VUELO) and _fila_de(om, _TITULO_AFAC_PERIODO)


# ---------------------------------------------------------------------------
# (f) Bloque VUELATOUR: etiquetas y pie con la bandera.
# ---------------------------------------------------------------------------


def _req_empresa(**om_extra) -> BalanceGeneralRequest:
    p = tev._payload()
    om = p["consolidado"]["otros_movimientos"]
    om["filas"].extend(pa.filas_afac_vuelo())
    om["afac_resumen"] = pa.resumen_afac()
    om.update(om_extra)
    return BalanceGeneralRequest.model_validate(p)


def _balance_ws(req: BalanceGeneralRequest):
    ws = load_workbook(BytesIO(render_balance_general_xlsx(req)))["balance"]
    return ws, _fila_de(ws, "RESULTADO VUELATOUR USD")


def test_bloque_empresa_con_bandera_etiquetas_y_pie() -> None:
    req = _req_empresa(afac_por_vuelo=True)
    ws, res = _balance_ws(req)
    assert _fila_de(ws, _EMPRESA_INGRESOS_AFAC_VUELO)
    assert _fila_de(ws, _EMPRESA_PAGOS_AFAC_VUELO)
    for viejo in (_EMPRESA_INGRESOS, _EMPRESA_PAGOS, _EMPRESA_INGRESOS_AFAC, _EMPRESA_PAGOS_AFAC):
        assert viejo not in _textos(ws)
    assert ws[f"A{res + 2}"].value == req.empresa.nota
    assert ws[f"A{res + 3}"].value == _NOTA_EMPRESA_BASE + _EMPRESA_AFAC_POR_VUELO
    # Los números del API, tal cual.
    assert ws[f"B{_fila_de(ws, _EMPRESA_INGRESOS_AFAC_VUELO)}"].value == (
        req.empresa.ingresos_propios_usd
    )
    assert ws[f"B{_fila_de(ws, _EMPRESA_PAGOS_AFAC_VUELO)}"].value == (
        req.empresa.pagos_vendedor_usd
    )


def test_bloque_empresa_manda_sobre_afac_en_filas_y_pie_fijo() -> None:
    req = _req_empresa(afac_por_vuelo=True, afac_en_filas=True)
    req.empresa.nota = None
    ws, res = _balance_ws(req)
    assert _fila_de(ws, _EMPRESA_INGRESOS_AFAC_VUELO)
    base = _NOTA_EMPRESA_BASE + _EMPRESA_AFAC_POR_VUELO
    assert ws[f"A{res + 2}"].value == _NOTA_BLOQUE_EMPRESA.replace(_NOTA_EMPRESA_BASE, base)
    assert not any(_EMPRESA_AFAC_EN_FILAS in t for t in _textos(ws))


@pytest.mark.parametrize("bandera", [None, False], ids=["ausente", "false"])
def test_bloque_empresa_sin_bandera_el_de_siempre(bandera) -> None:
    extra = {} if bandera is None else {"afac_por_vuelo": bandera}
    ws, res = _balance_ws(_req_empresa(**extra))
    assert _fila_de(ws, _EMPRESA_INGRESOS) and _fila_de(ws, _EMPRESA_PAGOS)
    assert ws[f"A{res + 3}"].value == _NOTA_EMPRESA_BASE


# ---------------------------------------------------------------------------
# (g) Libro INDIVIDUAL: la línea de los pagos de AFAC en 'permisos'.
# ---------------------------------------------------------------------------


def test_individual_linea_de_pagos_afac_en_permisos() -> None:
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(pa.individual_0089()))
    ws = load_workbook(BytesIO(data))["permisos"]
    assert ws["A3"].value == (
        "2 pagos de AFAC ($2,250.00) los paga VuelaTour: no restan aquí (ver "
        "Balance general · otros movimientos, bloque AFAC DEL PERIODO)."
    )
    assert ws["A3"].font.i is True
    assert ws["A3"].comment.text == (
        "01/09/2026 · Permisos Vuelos Privados · Yousefh Pineda Bedolla · $1,500.00\n"
        "10/09/2026 · Derechos AFAC MMCT · vuelo #402 · $750.00"
    )
    # La hoja 'permisos' de siempre (los marcados ya no vienen en `permisos`).
    base = load_workbook(
        BytesIO(render_balance_avion_xlsx(BalanceAvionRequest.model_validate(pa.individual_0088())))
    )["permisos"]
    assert base["A3"].value is None
    for fila in base.iter_rows():
        for c in fila:
            if c.coordinate != "A3":
                assert ws[c.coordinate].value == c.value, c.coordinate
    assert verificar_libro(data)
    vista = libro_balance_avion(
        BalanceAvionRequest.model_validate({**pa.individual_0089(), "vinculos_gasto": True})
    )
    a3 = _celda_json(_hoja_json(vista, "permisos"), 3, 1)
    assert a3["t"].startswith("2 pagos de AFAC") and a3["n"].startswith("01/09/2026")


_PAGO_USD_SIN_TC = {
    "fecha": "2026-09-22", "categoria": "Permiso", "detalle": "Derechos AFAC en USD",
    "monto_mxn": None, "moneda_original": "USD", "monto_original": 40.0,
    "tc_aplicado": None, "gasto_id": pa.GASTO_3,
}  # fmt: skip


@pytest.mark.parametrize("total", [0.0, None], ids=["api-0", "null"])
def test_individual_un_pago_en_usd_sin_tc(total) -> None:
    """Revisión (16-oct-2026): el `buildHoja` del API manda `total_mxn` 0 (no
    null) cuando el único pago es USD sin T.C.: la línea NO cita un «$0.00»
    falso, dice que ese pago no suma."""
    p = pa.individual_0089()
    p["pagos_afac"]["filas"] = [dict(_PAGO_USD_SIN_TC)]
    p["pagos_afac"]["total_mxn"] = total
    p["pagos_afac"]["usd"] = None
    p["pagos_afac"]["usd_hr"] = None
    ws = load_workbook(
        BytesIO(render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p)))
    )["permisos"]
    assert ws["A3"].value == (
        "1 pago de AFAC (1 en USD sin T.C., no suma) lo paga VuelaTour: no resta "
        "aquí (ver Balance general · otros movimientos, bloque AFAC DEL PERIODO)."
    )
    assert ws["A3"].comment.text == "22/09/2026 · Derechos AFAC en USD · sin T.C. (USD 40.00)"


def test_individual_pagos_mixtos_el_total_dice_que_no_incluye_el_usd_sin_tc() -> None:
    """Revisión (16-oct-2026): con un pago en pesos y otro USD sin T.C., el
    `total_mxn` del API ($1,500.00) NO incluye el segundo: la línea lo dice
    (antes: «2 pagos de AFAC ($1,500.00)», como si fuera de los dos)."""
    p = pa.individual_0089()
    p["pagos_afac"]["filas"] = [p["pagos_afac"]["filas"][0], dict(_PAGO_USD_SIN_TC)]
    p["pagos_afac"]["total_mxn"] = 1500.0
    p["pagos_afac"]["usd"] = None
    p["pagos_afac"]["usd_hr"] = None
    ws = load_workbook(
        BytesIO(render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p)))
    )["permisos"]
    assert ws["A3"].value == (
        "2 pagos de AFAC ($1,500.00; 1 en USD sin T.C., no suma) los paga "
        "VuelaTour: no restan aquí (ver Balance general · otros movimientos, "
        "bloque AFAC DEL PERIODO)."
    )
    assert ws["A3"].comment.text == (
        "01/09/2026 · Permisos Vuelos Privados · Yousefh Pineda Bedolla · $1,500.00\n"
        "22/09/2026 · Derechos AFAC en USD · sin T.C. (USD 40.00)"
    )


@pytest.mark.parametrize("pagos", [None, {"filas": []}], ids=["null", "vacia"])
def test_individual_sin_pagos_byte_a_byte(pagos) -> None:
    p = pa.individual_0088()
    p["pagos_afac"] = pagos
    data, vista = _balance("individual", p)
    assert (_huella_xlsx(data), _huella_json(vista)) == _HUELLAS_0088["individual"]


# ---------------------------------------------------------------------------
# (h) BYTE A BYTE: sin las llaves nuevas, y con ellas en neutro.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("caso", list(_HUELLAS_0088))
def test_payload_0088_byte_a_byte(caso: str) -> None:
    libro = "individual" if caso == "individual" else "general"
    data, vista = _balance(libro, _payload_0088(caso))
    assert (_huella_xlsx(data), _huella_json(vista)) == _HUELLAS_0088[caso]
    neutro = pa.con_llaves_afac_neutras(_payload_0088(caso), libro)
    data, vista = _balance(libro, neutro)
    assert (_huella_xlsx(data), _huella_json(vista)) == _HUELLAS_0088[caso]


@pytest.mark.parametrize("caso", list(_HUELLAS_HEAD))
def test_payloads_del_cierre_con_llaves_afac_neutras_byte_a_byte(caso: str) -> None:
    """Las huellas del HEAD f0c8ad1 de `test_cierre_mes` (sin AFAC): con las
    llaves AFAC en neutro, el mismo libro."""
    libro = "individual" if caso == "individual" else "general"
    base = pay.individual_base() if caso == "individual" else pay.general_base(caso.split("-")[1])
    data, vista = _balance(libro, pa.con_llaves_afac_neutras(base, libro))
    assert (_huella_xlsx(data), _huella_json(vista)) == _HUELLAS_HEAD[caso]


# ---------------------------------------------------------------------------
# (i) Esquemas aditivos, variante «general» y rutas.
# ---------------------------------------------------------------------------


def test_esquemas_aditivos_con_default() -> None:
    assert BalanceOtroMovimientoFila().afac is None
    assert BalanceHojaOtrosMovimientos().afac_por_vuelo is None
    assert BalanceHojaOtrosMovimientos().afac_resumen is None
    assert BalanceAvionRequest().pagos_afac is None
    v = BalanceAfacVuelo()
    assert (v.vuelo_id, v.editable, v.costo_real_mxn, v.reservado_mxn) == ("", False, None, None)
    r = BalanceAfacResumen()
    assert (r.vuelos, r.pagos, r.pagado_mxn, r.diferencia_pagado_mxn) == (0, [], None, None)
    assert BalanceAfacPago().gasto_id is None
    # Metadato liberal: jamás un 422.
    v = BalanceAfacVuelo.model_validate(
        {"vuelo_id": 5, "editable": 1, "costo_real_mxn": "nan", "tc_costos": "17.03",
         "mes_contable": "  ", "capturado_por_nombre": "  Ale "}
    )  # fmt: skip
    assert (v.vuelo_id, v.editable, v.costo_real_mxn, v.tc_costos) == ("", False, None, 17.03)
    assert (v.mes_contable, v.capturado_por_nombre) == (None, "Ale")
    r = BalanceAfacResumen.model_validate(
        {"vuelos": None, "vuelos_sin_real": -2, "pagos_sin_tc": 3.0}
    )
    assert (r.vuelos, r.vuelos_sin_real, r.pagos_sin_tc) == (0, 0, 3)
    assert BalanceAfacPago.model_validate({"gasto_id": " ", "detalle": None}).gasto_id is None


def test_variante_general_pinta_la_misma_pestana() -> None:
    def pestaña(variante: str):
        data, om, omv = _general(pa.general_0089(variante))
        return (
            [(c.coordinate, c.value) for fila in om.iter_rows() for c in fila],
            [(c.coordinate, c.value) for fila in omv.iter_rows() for c in fila],
            data,
        )

    mensual, general = pestaña("mensual"), pestaña("general")
    assert mensual[0] == general[0] and mensual[1] == general[1]
    assert verificar_libro(general[2])


def test_rutas_aceptan_el_payload_0089(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    p = pa.general_0089()
    res = client.post("/pdf/balance-general-xlsx", json=p, headers=headers)
    assert res.status_code == 200, res.text
    om = load_workbook(BytesIO(res.content))[OM]
    assert _fila_de(om, _OM_SECCION_AFAC_VUELO)
    res = client.post("/pdf/balance-general-libro", json={**p, "vinculos_gasto": True},
                      headers=headers)  # fmt: skip
    assert res.status_code == 200, res.text
    vista = _hoja_json(res.json(), OM)
    vinculos = [x["vinculo"] for x in vista["celdas"] if "vinculo" in x]
    assert [v["vuelo_id"] for v in vinculos] == [
        pa.VUELO_401, pa.VUELO_402, pa.VUELO_407, pa.VUELO_405
    ]  # fmt: skip
    assert all(x["c"] == 4 for x in vista["celdas"] if "vinculo" in x)
    res = client.post("/pdf/balance-avion-xlsx", json=pa.individual_0089(), headers=headers)
    assert res.status_code == 200, res.text
    assert load_workbook(BytesIO(res.content))["permisos"]["A3"].value.startswith("2 pagos")
