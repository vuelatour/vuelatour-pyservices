"""Hoja «combustible» de los dos libros de balance (11-sep-2026).

Pedido del cliente: saber CON QUÉ se pagó cada carga de combustible. El API
manda el medio de pago YA legible en `pago` (campo ADITIVO de
`BalanceAvionGastoFila`) y la hoja lo pinta como 7.ª columna «PAGO» —
individual y general. Todo lo demás de la hoja (secciones por matrícula,
subtotales, litros, $/L) queda EXACTAMENTE igual; solo creció a 7 columnas
(encabezado, bordes/relleno del subtotal, merges de títulos y notas, anchos).

También se blinda aquí que el DETALLE se pinta ÍNTEGRO: en la hoja 'otros
gastos' del GENERAL el API puede mandar el vuelo dentro del texto
(«… · vuelo #123») y esa referencia no se recorta nunca.
"""

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.schemas.reportes import BalanceAvionRequest, BalanceGeneralRequest
from app.services.balance_avion_xlsx import (
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)

_CARGA_1 = {
    "fecha": "2026-08-03",
    "detalle": "Gasavión 100LL · Aeropuerto CUN · factura A-991",
    "litros": 180.0,
    "monto_mxn": 9000.0,
    "pago": "Tarjeta BBVA ···4321",
}
_CARGA_2 = {
    "fecha": "2026-08-19",
    "detalle": "Gasavión 100LL · Aeropuerto MID",
    "litros": 120.0,
    "monto_mxn": 6300.0,
    "pago": "Efectivo (caja chica)",
}
# API VIEJO: la misma fila sin `pago` (skew de deploy) — celda vacía.
_CARGA_SIN_PAGO = {
    "fecha": "2026-08-25",
    "detalle": "Gasavión 100LL · Aeropuerto CZM",
    "litros": 90.0,
    "monto_mxn": 4500.0,
}

_COMBUSTIBLE_INDIVIDUAL = {
    "filas": [_CARGA_1, _CARGA_2, _CARGA_SIN_PAGO],
    "total_mxn": 19800.0,
    "litros_total": 390.0,
    "precio_litro_prom": 50.77,
    "usd": 990.0,
    "usd_hr": 99.0,
}


def _hoja_combustible(wb):
    return wb["combustible"]


def _individual() -> bytes:
    return render_balance_avion_xlsx(
        BalanceAvionRequest(
            matricula="XB-ABC",
            periodo_desde="2026-08-01",
            periodo_hasta="2026-08-31",
            combustible=_COMBUSTIBLE_INDIVIDUAL,
        )
    )


def _general(**extra) -> bytes:
    """GENERAL: el consolidado trae las cargas de DOS aviones (con matrícula),
    que la hoja agrupa en secciones con subtotal."""
    cargas = [
        {**_CARGA_1, "matricula": "XB-ABC", "avion_color": "#FF0000"},
        {**_CARGA_2, "matricula": "XB-ABC", "avion_color": "#FF0000"},
        {
            "fecha": "2026-08-11",
            "detalle": "Turbosina · Aeropuerto CUN",
            "litros": 400.0,
            "monto_mxn": 22000.0,
            "pago": "Transferencia HSBC",
            "matricula": "XB-XYZ",
            "avion_color": "#0000FF",
        },
    ]
    return render_balance_general_xlsx(
        BalanceGeneralRequest(
            periodo_desde="2026-08-01",
            periodo_hasta="2026-08-31",
            consolidado=BalanceAvionRequest(
                matricula="FLOTA",
                periodo_desde="2026-08-01",
                periodo_hasta="2026-08-31",
                combustible={
                    "filas": cargas,
                    "total_mxn": 37300.0,
                    "litros_total": 700.0,
                    "precio_litro_prom": 53.28,
                    "usd": 1865.0,
                },
            ),
            **extra,
        )
    )


# ===== Columna PAGO =====


def test_individual_pinta_la_columna_pago_al_final() -> None:
    ws = _hoja_combustible(load_workbook(BytesIO(_individual())))

    # Encabezado del ledger (fila 7): la 7.ª columna es PAGO y las seis de
    # siempre no se movieron.
    assert [ws.cell(row=7, column=c).value for c in range(1, 8)] == [
        "FECHA", "DETALLE", "LITROS", "MONTO MXN",
        "MONEDA ORIGINAL", "MONTO ORIGINAL", "PAGO",
    ]

    # Las filas traen el medio de pago tal cual lo mandó el API…
    assert ws.cell(row=8, column=7).value == "Tarjeta BBVA ···4321"
    assert ws.cell(row=9, column=7).value == "Efectivo (caja chica)"
    # …y la fila de un API viejo (sin `pago`) deja la celda VACÍA.
    assert ws.cell(row=10, column=7).value is None

    # Celda con borde y wrap (mismo trato que DETALLE).
    pago = ws.cell(row=8, column=7)
    assert pago.border.left.style == "thin"
    assert pago.alignment.wrap_text is True

    # Los montos no se movieron de columna (fiabilidad: la hoja solo creció).
    assert ws.cell(row=8, column=3).value == 180.0
    assert ws.cell(row=8, column=4).value == 9000.0

    # Ancho de la nueva columna.
    assert ws.column_dimensions[get_column_letter(7)].width == 18


def test_individual_subtotal_titulo_y_nota_abarcan_las_7_columnas() -> None:
    ws = _hoja_combustible(load_workbook(BytesIO(_individual())))

    # Fila TOTAL DEL PERIODO: la 7.ª celda también lleva borde y relleno.
    total_row = next(
        r for r in range(8, 20)
        if str(ws.cell(row=r, column=2).value or "").startswith("TOTAL DEL PERIODO")
    )
    celda = ws.cell(row=total_row, column=7)
    assert celda.border.left.style == "thin"
    assert celda.fill.fgColor.rgb.endswith("EEF2F7")

    # Título (fila 1) y nota al pie mergeados hasta la columna 7 (antes 6).
    merges = {str(m) for m in ws.merged_cells.ranges}
    assert "A1:G1" in merges
    assert any(m.startswith("A") and m.endswith(f"G{total_row + 2}") for m in merges)


def test_general_secciones_por_matricula_con_pago_y_subtotales_de_7_columnas() -> None:
    ws = _hoja_combustible(load_workbook(BytesIO(_general())))

    # Encabezado de sección (matrícula) mergeado a 7 columnas.
    assert ws.cell(row=8, column=1).value == "XB-ABC"
    assert "A8:G8" in {str(m) for m in ws.merged_cells.ranges}

    # Las cargas del avión traen su PAGO.
    assert ws.cell(row=9, column=7).value == "Tarjeta BBVA ···4321"
    assert ws.cell(row=10, column=7).value == "Efectivo (caja chica)"

    # Subtotal de la sección: la 7.ª celda con borde y relleno gris.
    sub = next(
        r for r in range(9, 30)
        if str(ws.cell(row=r, column=2).value or "").startswith("Subtotal XB-ABC")
    )
    assert ws.cell(row=sub, column=7).border.left.style == "thin"
    assert ws.cell(row=sub, column=7).fill.fgColor.rgb.endswith("F3F4F6")

    # Segunda sección (otro avión) con su propio pago.
    fila_xyz = next(
        r for r in range(sub, 30)
        if ws.cell(row=r, column=7).value == "Transferencia HSBC"
    )
    assert "Turbosina" in str(ws.cell(row=fila_xyz, column=2).value)


def test_hoja_vacia_sigue_mergeando_las_7_columnas() -> None:
    data = render_balance_avion_xlsx(
        BalanceAvionRequest(matricula="XB-ABC", periodo_desde="2026-08-01")
    )
    ws = _hoja_combustible(load_workbook(BytesIO(data)))
    assert ws.cell(row=8, column=1).value == "Sin cargas de combustible en el periodo."
    assert "A8:G8" in {str(m) for m in ws.merged_cells.ranges}


# ===== Detalle íntegro: «vuelo #123» en 'otros gastos' del GENERAL =====


def test_otros_gastos_del_general_conserva_el_vuelo_del_detalle() -> None:
    detalle = "Servicio de rampa · proveedor SAESA · vuelo #123"
    data = _general(
        gastos_empresa={
            "filas": [
                {
                    "fecha": "2026-08-07",
                    "categoria": "Servicios",
                    "detalle": detalle,
                    "monto_mxn": 3500.0,
                }
            ],
            "total_mxn": 3500.0,
            "usd": 175.0,
        }
    )
    ws = load_workbook(BytesIO(data))["otros gastos"]
    # ÍNTEGRO: ni recortado ni re-armado — la referencia al vuelo sobrevive.
    assert ws.cell(row=8, column=3).value == detalle
    assert "vuelo #123" in ws.cell(row=8, column=3).value
    assert ws.cell(row=8, column=3).alignment.wrap_text is True
