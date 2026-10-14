"""Otros ingresos ASIGNADOS a un avión en los balances (13-oct-2026, API 0.0.75;
invariante «INGRESO ASIGNADO A UN AVIÓN» del API).

Pedido del cliente: un ingreso registrado en Ingresos con avión (el caso real
ING-4 de N990GG: «reembolso del pago del serv de 50 hrs en toledo, fact 53»,
$46,550.80 MXN) cuenta UNA vez, para ESE avión: pestaña «Otros ingresos» de su
libro, línea «(+) OTROS INGRESOS USD» que SUMA en la cascada (y entra al
reparto) y deja de ser ingreso de VuelaTour. El API manda todo calculado; aquí
se pinta con fórmulas que reproducen sus números:
- libro INDIVIDUAL: pestaña «Otros ingresos» entre 'permisos' y 'balance'
  (siempre que llegue la llave, aunque vacía) y la línea en la cascada, que
  cita su TOTAL USD;
- Balance GENERAL: la misma línea en cada bloque de 'balance' (VALOR, como el
  resto del bloque), hoja 'otros ingresos aviones' SOLO con filas y la cola de
  la fila 2 de 'otros movimientos' con la bandera `otros_ingresos_por_avion`;
- sin la llave (API ≤ 0.0.74): libros byte a byte (huellas de miembros
  calculadas con el código de HEAD d58faa1, antes de este cambio).
Los payloads salen del espejo del API de `tests/test_balance_formulas.py`.

14-oct-2026 (API 0.0.77; decisión del cliente: «unificar la conversión del
ingreso de un avión entre balance y reparto (ING-4 2,710.47 vs 2,746.70):
hazlo, si es lo mejor»): el API convierte CADA ingreso con la función del
reparto (`usdDeIngresoAvion`) y manda por fila `tc_usd`, `neto_usd`,
`tc_usd_origen` y `tc_usd_nota`, con la bandera de la hoja `usd_por_ingreso`.
Con ella la hoja gana T.C. A USD (J) y NETO USD (K) y TOTAL USD =
ROUND(SUM(K8:Kn),2) (individual C5, que cita la cascada; general B5), con las
tres notas `*_POR_INGRESO`. Sin ella, byte a byte: huellas con un payload de
FORMA REAL del API de producción 0.0.76 (revisión F12) calculadas con el
código de HEAD b66df98 antes del cambio.
"""

from __future__ import annotations

import copy
import hashlib
import json
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAvionBalanceBloque,
    BalanceAvionHojaOtrosIngresos,
    BalanceAvionOtroIngresoFila,
    BalanceAvionRequest,
    BalanceGeneralRequest,
    BalanceHojaOtrosMovimientos,
)
from app.services import libro_json, xlsx_formulas
from app.services.balance_avion_xlsx import (
    _ANCHOS_OTROS_INGRESOS_POR_INGRESO,
    _ENCABEZADOS_OTROS_INGRESOS,
    _ENCABEZADOS_OTROS_INGRESOS_POR_INGRESO,
    _FILA_OTROS_INGRESOS,
    _NOTA_CASCADA_OTROS_INGRESOS,
    _NOTA_CASCADA_OTROS_INGRESOS_POR_INGRESO,
    _NOTA_FACTURA_VUELATOUR_DEVOLUCION,
    _NOTA_FACTURA_VUELATOUR_SEGUN_COBROS,
    _NOTA_OTROS_INGRESOS_AVION,
    _NOTA_OTROS_INGRESOS_AVION_POR_INGRESO,
    _NOTA_OTROS_INGRESOS_AVIONES_GENERAL,
    _NOTA_OTROS_INGRESOS_AVIONES_GENERAL_POR_INGRESO,
    _OM_OTROS_INGRESOS_POR_AVION,
    _SIN_OTROS_INGRESOS,
    _TC_POR_INGRESO,
    _TITULO_OTROS_INGRESOS_GENERAL,
    MUTED,
    TC,
    TC_OFICIAL_FILL,
    libro_balance_avion,
    libro_balance_general,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from tests import test_balance_comisiones as t_comisiones
from tests import test_balance_comisiones_general as tcg
from tests import test_balance_empresa_vuelatour as t_empresa
from tests import test_balance_extension_horario as t_extension
from tests import test_balance_formulas as tf
from tests._evaluador_formulas import verificar_libro

r2 = tf.r2
MAESTRA_TST = "reporte horas XA-TST"
HOJA = "Otros ingresos"
HOJA_GENERAL = "otros ingresos aviones"
DESPUES = "UTILIDAD DESPUÉS DE GASTOS USD"
COBRADA = "UTILIDAD COBRADA USD"

# ---------------------------------------------------------------------------
# Textos EXACTOS del contrato (§8 de la especificación).
# ---------------------------------------------------------------------------

NOTA_HOJA = (
    "Ingresos registrados en Ingresos (clave ING-n) y asignados a este avión: "
    "reembolsos de su servicio o de su aseguradora, ventas de piezas o activos, "
    "intereses… NETO MXN = ingreso − comisión del banco (en dólares, con el T.C. "
    "propio del ingreso). TOTAL USD = TOTAL NETO MXN ÷ TC PROMEDIO de las hojas "
    "de este libro. Ese TOTAL USD SUMA en la hoja 'balance' («(+) OTROS "
    "INGRESOS USD») y en el reparto a socios; el reparto convierte cada "
    "ingreso por separado (en dólares cuenta su neto tal cual; en pesos, con "
    "su T.C. propio o el oficial de su día), así que su monto en USD puede "
    "diferir del de esta hoja según la distancia entre esos T.C. y el promedio "
    "del libro (no son centavos), como pasa con los gastos. No cuentan como ingreso "
    "de VuelaTour: no están en 'otros movimientos' del Balance general ni en el "
    "total de 'Otros ingresos' del Libro Dinero. Los anticipos y las "
    "aportaciones nunca llevan avión."
)
NOTA_CASCADA = (
    "Ingresos registrados en Ingresos y asignados a este avión (hoja 'Otros "
    "ingresos' de su libro), netos de comisión bancaria, al T.C. promedio del "
    "libro. Suman a la utilidad del avión y a su reparto a socios; no cuentan "
    "como ingreso de VuelaTour."
)
NOTA_GENERAL = (
    "Ingresos registrados en Ingresos y asignados a un avión (clave ING-n), "
    "netos de comisión bancaria. Cada uno YA SUMA en el bloque de su avión en "
    "'balance' («(+) OTROS INGRESOS USD») y en su reparto a socios; por eso no "
    "están en 'otros movimientos' ni en el bloque VUELATOUR (empresa). TOTAL "
    "USD = suma de lo que suma cada avión (cada libro convierte con su propio "
    "T.C. promedio)."
)
OM_COLA = (
    " Los otros ingresos asignados a un avión (por ejemplo el reembolso de su "
    "servicio) no están aquí: cuentan en la hoja 'Otros ingresos' del libro de "
    "ese avión, en su bloque de 'balance' y en su reparto a socios (si hay "
    "alguno en el periodo, su detalle está en la hoja 'otros ingresos "
    "aviones')."
)
ENCABEZADOS = [
    "FECHA",
    "CATEGORÍA",
    "DETALLE",
    "MONTO ORIGINAL",
    "MONEDA",
    "T.C.",
    "INGRESO MXN",
    "COMISIÓN MXN",
    "NETO MXN",
    "CUENTA",
    "NOTAS",
]

# Contrato del API 0.0.77 (conversión ÚNICA por ingreso, bandera
# `usd_por_ingreso`): textos EXACTOS de la especificación del 14-oct-2026.
NOTA_HOJA_POR_INGRESO = (
    "Ingresos registrados en Ingresos (clave ING-n) y asignados a este avión: "
    "reembolsos de su servicio o de su aseguradora, ventas de piezas o activos, "
    "intereses… NETO MXN = ingreso − comisión del banco (en dólares, con el T.C. "
    "propio del ingreso). Cada ingreso pasa a USD por separado, con la MISMA "
    "regla del reparto a socios: en dólares cuenta su neto tal cual; en pesos, "
    "NETO USD = NETO MXN ÷ su T.C. propio o, si no lo trae, el T.C. oficial del "
    "día del ingreso (celda azul). Un reembolso ligado a un gasto que este libro "
    "cuenta usa el T.C. con el que el libro convierte ese gasto (nota en la "
    "celda), para que reembolso y gasto se anulen. TOTAL USD = suma de NETO "
    "USD: SUMA en la hoja 'balance' («(+) OTROS INGRESOS USD») y en el reparto "
    "a socios. No cuentan como ingreso de VuelaTour: no están en 'otros "
    "movimientos' del Balance general ni en el total de 'Otros ingresos' del "
    "Libro Dinero. Los anticipos y las aportaciones nunca llevan avión."
)
NOTA_GENERAL_POR_INGRESO = (
    "Ingresos registrados en Ingresos y asignados a un avión (clave ING-n), "
    "netos de comisión bancaria. Cada uno YA SUMA en el bloque de su avión en "
    "'balance' («(+) OTROS INGRESOS USD») y en su reparto a socios; por eso no "
    "están en 'otros movimientos' ni en el bloque VUELATOUR (empresa). Cada "
    "ingreso pasa a USD con su T.C. propio o el oficial de su día (en dólares, "
    "su neto tal cual; un reembolso ligado a un gasto de su libro, con el T.C. "
    "de ese gasto), como en el reparto; TOTAL USD = suma de NETO USD."
)
NOTA_CASCADA_POR_INGRESO = (
    "Ingresos registrados en Ingresos y asignados a este avión (hoja 'Otros "
    "ingresos' de su libro), netos de comisión bancaria, cada uno a USD con su "
    "T.C. propio o el oficial de su día, como en el reparto a socios (un "
    "reembolso ligado a un gasto de este libro usa el T.C. de ese gasto para "
    "anularse con él). Suman a la utilidad del avión y a su reparto; no "
    "cuentan como ingreso de VuelaTour."
)
ENCABEZADOS_POR_INGRESO = [
    "FECHA",
    "CATEGORÍA",
    "DETALLE",
    "MONTO ORIGINAL",
    "MONEDA",
    "T.C.",
    "INGRESO MXN",
    "COMISIÓN MXN",
    "NETO MXN",
    "T.C. A USD",
    "NETO USD",
    "CUENTA",
    "NOTAS",
]
# Notas de la celda T.C. A USD que arma el API (`notaTcUsdIngreso`).
NOTA_TC_OFICIAL_ING4 = (
    "T.C. oficial de referencia del día del ingreso (open.er-api, dato del "
    "2026-09-04): el ingreso en pesos no traía tipo de cambio."
)
NOTA_TC_LIGADO = (
    "T.C. con el que este libro convierte el gasto que reembolsa (gasto "
    "70e8c3a4 del 15-sep-2026, T.C. promedio de las hojas): reembolso y gasto "
    "se anulan en este libro."
)

# ---------------------------------------------------------------------------
# Filas como las arma el API (`filaHojaOtrosIngresos` sobre
# `filaLibroDeIngreso`).
# ---------------------------------------------------------------------------

_FILA_MXN = {
    "ingreso_id": "i-4",
    "clave": "ING-4",
    "fecha": "2026-09-04",
    "categoria": "Reembolsos y devoluciones recibidos",
    "detalle": "ING-4 · TRANSF INTERBANCARIA SPEI · angel alvarez",
    "moneda_original": "MXN",
    "monto_original": 46550.8,
    "tc_aplicado": None,
    "ingreso_mxn": 46550.8,
    "comision_mxn": None,
    "neto_mxn": 46550.8,
    "cuenta": "EJE SCOTIA · Scotiabank",
    "notas": "reembolso del pago del serv de 50 hrs en toledo, fact 53",
    # API 0.0.77: sin T.C. propio ni gasto ligado ⇒ el oficial del día de su
    # fecha (open.er-api, 4-sep-2026: 16.9479) ⇒ 46,550.80 ÷ 16.9479 = 2,746.70,
    # lo mismo que el reparto y el tablero (antes el balance daba 2,710.47).
    "tc_usd": 16.9479,
    "neto_usd": 2746.70,
    "tc_usd_origen": "OFICIAL_DEL_DIA",
    "tc_usd_nota": NOTA_TC_OFICIAL_ING4,
}
# USD 100 × 17.5 con comisión de 5 USD (= 87.50 MXN).
_FILA_USD = {
    "ingreso_id": "i-9",
    "clave": "ING-9",
    "fecha": "2026-09-20",
    "categoria": "Venta de activos",
    "detalle": "ING-9 · Venta de una hélice usada · Aero Partes",
    "moneda_original": "USD",
    "monto_original": 100.0,
    "tc_aplicado": 17.5,
    "ingreso_mxn": 1750.0,
    "comision_mxn": 87.5,
    "neto_mxn": 1662.5,
    "cuenta": "Efectivo / caja",
    "notas": None,
    # API 0.0.77: en dólares cuenta su neto tal cual (100 − 5), sin T.C.
    "tc_usd": None,
    "neto_usd": 95.0,
    "tc_usd_origen": None,
    "tc_usd_nota": None,
}
# USD SIN tipo de cambio: sin pesos ⇒ la hoja (y la utilidad) en null.
_FILA_USD_SIN_TC = {
    "ingreso_id": "i-11",
    "clave": "ING-11",
    "fecha": "2026-09-25",
    "categoria": "Otros ingresos",
    "detalle": "ING-11 · Intereses",
    "moneda_original": "USD",
    "monto_original": 200.0,
    "tc_aplicado": None,
    "ingreso_mxn": None,
    "comision_mxn": None,
    "neto_mxn": None,
    "cuenta": "Efectivo / caja",
    "notas": None,
    # API 0.0.77: su NETO USD sí suma (el CHECK del API impide el caso).
    "tc_usd": None,
    "neto_usd": 200.0,
    "tc_usd_origen": None,
    "tc_usd_nota": None,
}
# API 0.0.77: reembolso en pesos LIGADO a un gasto en pesos que este libro
# cuenta en una hoja (el SERVICIOS 70e8c3a4 de N4142R, $45,483.60 del
# 15-sep): se convierte con el T.C. con que el libro convierte ese gasto
# (`tcHojas` = 17.237699) ⇒ 2,638.61 y se anula con él. Sin relleno azul.
_FILA_LIGADA = {
    "ingreso_id": "i-12",
    "clave": "ING-12",
    "fecha": "2026-09-04",
    "categoria": "Reembolsos y devoluciones recibidos",
    "detalle": "ING-12 · Reembolso del servicio de 50 hrs · Toledo",
    "moneda_original": "MXN",
    "monto_original": 45483.6,
    "tc_aplicado": None,
    "ingreso_mxn": 45483.6,
    "comision_mxn": None,
    "neto_mxn": 45483.6,
    "cuenta": "EJE SCOTIA · Scotiabank",
    "notas": None,
    "tc_usd": 17.237699,
    "neto_usd": 2638.61,
    "tc_usd_origen": "GASTO_LIGADO",
    "tc_usd_nota": NOTA_TC_LIGADO,
}
# API 0.0.77: pesos sin T.C. propio ni oficial de su día ⇒ sin NETO USD (la
# hoja y la utilidad quedan vacías; el API lo grita en pendientes).
_FILA_MXN_SIN_TC = {
    **{k: v for k, v in _FILA_MXN.items() if k not in ("tc_usd", "neto_usd")},
    "ingreso_id": "i-13",
    "clave": "ING-13",
    "tc_usd": None,
    "neto_usd": None,
    "tc_usd_origen": None,
    "tc_usd_nota": None,
}


def _tc_hojas(payload: dict) -> float | None:
    """`tcHojas` del API: promedio CRUDO de los T.C. de costos de las filas."""
    zs = [v["tc_costos"] for v in payload["vuelos"] if v["tc_costos"] is not None]
    return tf._suma(zs) / len(zs) if zs else None


def _hoja_api(filas: list[dict], tc_hojas: float | None) -> dict:
    """`hojaOtrosIngresosAvion` del API: total = Σ netos no nulos; usd con la
    regla de `buildHoja` (alguna fila sin pesos ⇒ null; total 0 ⇒ 0)."""
    total = r2(tf._suma(f["neto_mxn"] for f in filas if f["neto_mxn"] is not None))
    if any(f["neto_mxn"] is None for f in filas):
        usd = None
    elif total == 0:
        usd = 0
    else:
        usd = r2(total / tc_hojas) if tc_hojas is not None else None
    return {"filas": copy.deepcopy(filas), "total_mxn": total, "usd": usd}


def _hoja_api_por_ingreso(filas: list[dict]) -> dict:
    """`hojaOtrosIngresosAvion` del API 0.0.77 (conversión ÚNICA por ingreso):
    total = Σ netos MXN no nulos (igual que antes); usd = round2(Σ `neto_usd`)
    — null si alguna fila no lo trae (pesos sin ningún T.C.), vacía ⇒ 0; una
    fila USD ya no anula la hoja por no tener pesos — y `usd_por_ingreso`
    como ÚLTIMA llave, SIEMPRE."""
    total = r2(tf._suma(f["neto_mxn"] for f in filas if f["neto_mxn"] is not None))
    if any(f.get("neto_usd") is None for f in filas):
        usd = None
    else:
        usd = r2(tf._suma(f["neto_usd"] for f in filas)) if filas else 0
    return {"filas": copy.deepcopy(filas), "total_mxn": total, "usd": usd, "usd_por_ingreso": True}


def _con_otros_ingresos(payload: dict, filas: list[dict], *, por_ingreso: bool = False) -> dict:
    """El libro de un avión como lo manda el API 0.0.75 (o, con `por_ingreso`,
    el 0.0.77: la hoja con `usd_por_ingreso` y TOTAL USD = Σ NETO USD): la
    hoja y `balance.otros_ingresos_avion_usd`, con la utilidad después de
    gastos, la cobrada y los socios RE-ARMADOS (la línea SUMA; null se
    propaga)."""
    p = copy.deepcopy(payload)
    hoja = _hoja_api_por_ingreso(filas) if por_ingreso else _hoja_api(filas, _tc_hojas(p))
    p["otros_ingresos_avion"] = hoja
    b = p["balance"]
    b["otros_ingresos_avion_usd"] = hoja["usd"]
    restas = [
        p["combustible"]["usd"],
        *(p[n]["usd"] for n in ("gastos_indirectos", "otros_gastos", "refacciones", "permisos")),
    ]
    antes = b["utilidad_antes_usd"]
    oi = hoja["usd"]
    despues = (
        r2(antes - tf._suma(restas) + oi)
        if oi is not None and all(x is not None for x in restas)
        else None
    )
    cobrada = r2(despues - p["totales"]["por_cobrar_usd"]) if despues is not None else None
    b["utilidad_despues_usd"] = despues
    b["utilidad_cobrada_usd"] = cobrada
    for s in b["socios"]:
        s["monto_usd"] = r2((s["porcentaje"] / 100) * cobrada) if cobrada is not None else None
    return p


def _individual(
    filas: list[dict] | None = None, *, por_ingreso: bool = False
) -> BalanceAvionRequest:
    filas = [_FILA_MXN, _FILA_USD] if filas is None else filas
    return BalanceAvionRequest.model_validate(
        _con_otros_ingresos(tf._payload_tst(), filas, por_ingreso=por_ingreso)
    )


def _libros(data: bytes):
    """(con fórmulas, con caché)."""
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _miembros(data: bytes) -> dict[str, bytes]:
    """Todos los miembros del .xlsx salvo docProps/core.xml (lleva la hora)."""
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _huella(data: bytes) -> str:
    h = hashlib.sha256()
    for nombre, contenido in _miembros(data).items():
        h.update(nombre.encode())
        h.update(contenido)
    return h.hexdigest()


def _fila_de(ws, etiqueta: str) -> int:
    filas = [c.row for c in ws["A"] if c.value == etiqueta]
    assert len(filas) == 1, (etiqueta, filas)
    return filas[0]


@pytest.fixture
def registros(monkeypatch) -> list[xlsx_formulas.RegistroFormulas]:
    """Registro de fórmulas de cada libro que se cierra (cuáles volvieron a
    valor)."""
    capturados: list[xlsx_formulas.RegistroFormulas] = []
    original = xlsx_formulas.finalizar

    def espia(wb):
        cache = original(wb)
        capturados.append(xlsx_formulas.registro(wb))
        return cache

    monkeypatch.setattr(xlsx_formulas, "finalizar", espia)
    return capturados


# Las dos celdas que el espejo del API deja como VALOR a propósito en el libro
# XA-TST (redondeo intermedio de COSTO TOTAL del #405 y TOTAL USD de 'Gastos
# Indirectos', `test_balance_formulas`): los otros ingresos no agregan ninguna.
_DEGRADADAS_TST = [(MAESTRA_TST, "U7"), ("Gastos Indirectos", "C5")]


# ---------------------------------------------------------------------------
# 1) Esquema y textos.
# ---------------------------------------------------------------------------


def test_esquemas_aditivos_con_default() -> None:
    assert BalanceAvionRequest().otros_ingresos_avion is None
    assert BalanceAvionBalanceBloque().otros_ingresos_avion_usd is None
    assert BalanceHojaOtrosMovimientos().otros_ingresos_por_avion is None
    hoja = BalanceAvionHojaOtrosIngresos.model_validate({"filas": None})
    assert (hoja.filas, hoja.total_mxn, hoja.usd) == ([], None, None)
    fila = BalanceAvionOtroIngresoFila.model_validate({"clave": None, "monto_original": 5})
    assert (fila.clave, fila.monto_original, fila.neto_mxn) == ("", 5.0, None)
    # Las llaves desconocidas se ignoran (skew: un API más nuevo no da 422).
    req = BalanceAvionRequest.model_validate(
        {"otros_ingresos_avion": {"filas": [{**_FILA_MXN, "algo_nuevo": 1}], "usd": 2}}
    )
    assert req.otros_ingresos_avion.filas[0].clave == "ING-4"


def test_textos_son_los_del_contrato() -> None:
    assert _NOTA_OTROS_INGRESOS_AVION == NOTA_HOJA
    assert _NOTA_CASCADA_OTROS_INGRESOS == NOTA_CASCADA
    assert _NOTA_OTROS_INGRESOS_AVIONES_GENERAL == NOTA_GENERAL
    assert _OM_OTROS_INGRESOS_POR_AVION == OM_COLA
    assert _ENCABEZADOS_OTROS_INGRESOS == ENCABEZADOS
    assert _FILA_OTROS_INGRESOS == "(+) OTROS INGRESOS USD"
    assert _SIN_OTROS_INGRESOS == "Sin otros ingresos asignados a este avión en el periodo."
    assert _TITULO_OTROS_INGRESOS_GENERAL == "Otros ingresos asignados a aviones — flota"
    # API 0.0.77 (bandera `usd_por_ingreso`): las tres notas nuevas, las 13
    # columnas con sus anchos y el texto de B5. Las de arriba NO se tocan
    # (un API ≤ 0.0.76 sigue viendo su libro de siempre).
    assert _NOTA_OTROS_INGRESOS_AVION_POR_INGRESO == NOTA_HOJA_POR_INGRESO
    assert _NOTA_CASCADA_OTROS_INGRESOS_POR_INGRESO == NOTA_CASCADA_POR_INGRESO
    assert _NOTA_OTROS_INGRESOS_AVIONES_GENERAL_POR_INGRESO == NOTA_GENERAL_POR_INGRESO
    assert _ENCABEZADOS_OTROS_INGRESOS_POR_INGRESO == ENCABEZADOS_POR_INGRESO
    assert _ANCHOS_OTROS_INGRESOS_POR_INGRESO == [12, 22, 50, 15, 9, 11, 15, 14, 15, 11, 14, 22, 36]
    assert _TC_POR_INGRESO == "por ingreso (col. J)"
    # Ninguno lleva la marca del amarillo de 'otros movimientos'.
    for t in (
        NOTA_HOJA,
        NOTA_CASCADA,
        NOTA_GENERAL,
        OM_COLA,
        NOTA_HOJA_POR_INGRESO,
        NOTA_CASCADA_POR_INGRESO,
        NOTA_GENERAL_POR_INGRESO,
    ):
        assert "· PROVISIÓN (" not in t
    # Las notas nuevas ya no hablan del T.C. promedio (esa era la diferencia
    # con el reparto que el cliente pidió quitar).
    for t in (NOTA_HOJA_POR_INGRESO, NOTA_CASCADA_POR_INGRESO, NOTA_GENERAL_POR_INGRESO):
        assert "promedio" not in t


# ---------------------------------------------------------------------------
# 2) Libro INDIVIDUAL.
# ---------------------------------------------------------------------------


def test_individual_pestana_entre_permisos_y_balance() -> None:
    wb = load_workbook(BytesIO(render_balance_avion_xlsx(_individual())))
    assert wb.sheetnames == [
        MAESTRA_TST,
        "cobranza",
        "combustible",
        "Gastos Indirectos",
        "refacciones",
        "permisos",
        HOJA,
        "balance",
        "pendientes de captura",
    ]
    # Sin la hoja 'refacciones' (API viejo): sigue entre 'permisos' y 'balance'.
    p = _con_otros_ingresos(tf._payload_tst(), [_FILA_MXN])
    p["refacciones"] = None
    wb = load_workbook(BytesIO(render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))))
    assert wb.sheetnames[-4:] == ["permisos", HOJA, "balance", "pendientes de captura"]


def test_individual_hoja_encabezados_filas_y_formulas(registros) -> None:
    req = _individual()
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws, wsv = wb[HOJA], wv[HOJA]
    assert ws["A1"].value == "Otros ingresos — XA-TST"
    assert ws["A2"].value == "Periodo: 2026-09-01 a 2026-09-30"
    assert [ws.cell(row=4, column=c).value for c in range(1, 4)] == [
        "TOTAL NETO MXN",
        "TC PROMEDIO",
        "TOTAL USD",
    ]
    assert [ws.cell(row=7, column=c).value for c in range(1, 12)] == ENCABEZADOS
    # Fila MXN (8) y USD (9).
    assert [ws.cell(row=8, column=c).value for c in range(1, 6)] == [
        "04/09/2026",
        "Reembolsos y devoluciones recibidos",
        "ING-4 · TRANSF INTERBANCARIA SPEI · angel alvarez",
        46550.8,
        "MXN",
    ]
    assert ws["F8"].value is None  # sin T.C. en pesos
    assert ws["G8"].value == "=ROUND(D8,2)"
    assert ws["H8"].value is None  # sin comisión
    assert ws["I8"].value == "=ROUND(G8,2)"
    assert (ws["J8"].value, ws["K8"].value) == (
        "EJE SCOTIA · Scotiabank",
        "reembolso del pago del serv de 50 hrs en toledo, fact 53",
    )
    assert (ws["D9"].value, ws["E9"].value, ws["F9"].value) == (100, "USD", 17.5)
    assert ws["G9"].value == "=ROUND(D9*F9,2)"
    assert ws["H9"].value == 87.5
    assert ws["I9"].value == "=ROUND(G9-H9,2)"
    # Resumen: fórmulas que citan la hoja y la maestra.
    assert ws["A5"].value == "=ROUND(SUM(I8:I9),2)"
    assert ws["B5"].value == f"='{MAESTRA_TST}'!$V$9"
    assert ws["C5"].value == f"=ROUND(A5/AVERAGE('{MAESTRA_TST}'!$V$3:$V$8),2)"
    # Los números son los del API (caché) y ninguna fórmula volvió a valor.
    hoja = req.otros_ingresos_avion
    assert (wsv["G8"].value, wsv["I8"].value) == (46550.8, 46550.8)
    assert (wsv["G9"].value, wsv["I9"].value) == (1750, 1662.5)
    assert wsv["A5"].value == hoja.total_mxn == r2(46550.8 + 1662.5)
    assert wsv["C5"].value == hoja.usd == r2(hoja.total_mxn / _tc_hojas(tf._payload_tst()))
    assert registros[-1].degradadas() == _DEGRADADAS_TST
    assert verificar_libro(data)
    # Nota al pie (envuelta, en A:K) y panel fijo bajo los encabezados.
    pie = ws.cell(row=11, column=1)
    assert pie.value == NOTA_HOJA and pie.alignment.wrap_text
    assert any(str(rg) == "A11:K11" for rg in ws.merged_cells.ranges)
    assert ws.freeze_panes == "A8"
    assert ws.column_dimensions["C"].width == 50 and ws.column_dimensions["K"].width == 36


def test_individual_cascada_suma_los_otros_ingresos(registros) -> None:
    req = _individual()
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    k = _fila_de(ws, "UTILIDAD ANTES DE GASTOS USD")
    assert [ws.cell(row=k + i, column=1).value for i in range(9)] == [
        "UTILIDAD ANTES DE GASTOS USD",
        "(−) COMBUSTIBLE DEL MES USD",
        "(−) GASTOS INDIRECTOS USD",
        "(−) REFACCIONES (INVENTARIO) USD",
        "(−) PERMISOS USD",
        _FILA_OTROS_INGRESOS,
        DESPUES,
        "(−) PENDIENTE DE PAGO (COBRANZA PENDIENTE) USD",
        COBRADA,
    ]
    linea = ws.cell(row=k + 5, column=2)
    assert linea.value == "='Otros ingresos'!$C$5"
    assert linea.comment is not None and linea.comment.text == NOTA_CASCADA
    assert ws.cell(row=k + 6, column=2).value == (
        f"=ROUND(B{k}-B{k + 1}-B{k + 2}-B{k + 3}-B{k + 4}+B{k + 5},2)"
    )
    assert ws.cell(row=k + 8, column=2).value == f"=ROUND(B{k + 6}-B{k + 7},2)"
    b = req.balance
    assert wsv.cell(row=k + 5, column=2).value == b.otros_ingresos_avion_usd
    assert wsv.cell(row=k + 6, column=2).value == b.utilidad_despues_usd
    assert wsv.cell(row=k + 8, column=2).value == b.utilidad_cobrada_usd
    # Los socios citan la UTILIDAD COBRADA corrida.
    socio = _fila_de(ws, "Socio A")
    assert ws.cell(row=socio, column=3).value == f"=ROUND(B{socio}/100*$B${k + 8},2)"
    assert [wsv.cell(row=socio + i, column=3).value for i in range(2)] == [
        s.monto_usd for s in b.socios
    ]
    # La utilidad sube EXACTAMENTE lo que suman los otros ingresos.
    sin = tf._payload_tst()["balance"]
    assert b.utilidad_despues_usd == r2(sin["utilidad_despues_usd"] + b.otros_ingresos_avion_usd)
    assert registros[-1].degradadas() == _DEGRADADAS_TST


def test_individual_hoja_vacia_texto_y_linea_en_cero() -> None:
    req = _individual([])
    assert (req.otros_ingresos_avion.total_mxn, req.otros_ingresos_avion.usd) == (0, 0)
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws = wb[HOJA]
    assert ws["A8"].value == _SIN_OTROS_INGRESOS
    assert any(str(rg) == "A8:K8" for rg in ws.merged_cells.ranges)
    assert ws["A5"].value == 0  # sin filas no hay rango que sumar: valor
    assert wv[HOJA]["C5"].value == 0
    k = _fila_de(wb["balance"], _FILA_OTROS_INGRESOS)
    assert wb["balance"].cell(row=k, column=2).value == "='Otros ingresos'!$C$5"
    assert wv["balance"].cell(row=k, column=2).value == 0
    assert verificar_libro(data)


def test_individual_usd_sin_tc_celdas_vacias_pero_la_fila_existe() -> None:
    req = _individual([_FILA_MXN, _FILA_USD_SIN_TC])
    assert req.otros_ingresos_avion.usd is None
    assert req.balance.utilidad_despues_usd is None
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws = wb[HOJA]
    # La fila USD sin T.C.: sin pesos ni fórmula (jamás un USD como MXN).
    assert (ws["G9"].value, ws["I9"].value, ws["F9"].value) == (None, None, None)
    assert ws["D9"].value == 200 and ws["E9"].value == "USD"
    # TOTAL NETO = solo los pesos; TOTAL USD vacío.
    assert wv[HOJA]["A5"].value == 46550.8
    assert ws["C5"].value is None
    bal = wb["balance"]
    k = _fila_de(bal, _FILA_OTROS_INGRESOS)
    assert bal.cell(row=k, column=2).value is None
    assert bal.cell(row=k, column=2).comment is None
    assert bal.cell(row=k + 1, column=1).value == DESPUES
    assert bal.cell(row=k + 1, column=2).value is None
    assert verificar_libro(data)


@pytest.mark.parametrize("valor", [..., None], ids=["sin-llave", "null"])
def test_individual_sin_la_llave_libro_de_hoy(valor) -> None:
    p = tf._payload_tst()
    if valor is not ...:
        p["otros_ingresos_avion"] = valor
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    wb = load_workbook(BytesIO(data))
    assert HOJA not in wb.sheetnames
    assert _FILA_OTROS_INGRESOS not in [c.value for c in wb["balance"]["A"]]
    assert wb["balance"]["B9"].value == "=ROUND(B4-B5-B6-B7-B8,2)"
    assert _huella(data) == _HUELLA_INDIVIDUAL_HEAD


# ---------------------------------------------------------------------------
# 3) Balance GENERAL.
# ---------------------------------------------------------------------------

_COLOR_TST = "#3B82F6"


def _fusion_flota(libros: list[dict], *, por_ingreso: bool = False) -> dict | None:
    """`fusionarOtrosIngresosFlota` del API: solo libros con filas; detalle con
    «MATRÍCULA · » al frente, matrícula y color por fila; orden por fecha
    (estable); totales sumados y null que se propaga. Con `por_ingreso` (API
    0.0.77) además `usd_por_ingreso: true` como última llave."""
    con = [p for p in libros if (p.get("otros_ingresos_avion") or {}).get("filas")]
    if not con:
        return None
    filas = sorted(
        (
            {
                **f,
                "detalle": f"{p['matricula']} · {f['detalle']}",
                "matricula": p["matricula"],
                "avion_color": p.get("avion_color"),
            }
            for p in con
            for f in p["otros_ingresos_avion"]["filas"]
        ),
        key=lambda f: f["fecha"],
    )
    usds = [p["otros_ingresos_avion"]["usd"] for p in con]
    hoja = {
        "filas": filas,
        "total_mxn": r2(tf._suma(p["otros_ingresos_avion"]["total_mxn"] for p in con)),
        "usd": None if any(u is None for u in usds) else r2(tf._suma(usds)),
    }
    if por_ingreso:
        hoja["usd_por_ingreso"] = True
    return hoja


def _general_dict(
    *, bandera: bool = True, filas_tst=None, filas_dos=None, por_ingreso: bool = False
) -> dict:
    tst = tf._payload_tst()
    tst["avion_color"] = _COLOR_TST
    tst = _con_otros_ingresos(
        tst, [_FILA_MXN, _FILA_USD] if filas_tst is None else filas_tst, por_ingreso=por_ingreso
    )
    dos = _con_otros_ingresos(
        tf._payload_dos(), [] if filas_dos is None else filas_dos, por_ingreso=por_ingreso
    )
    p = tf._general_payload([tst, dos])
    oi = _fusion_flota([tst, dos], por_ingreso=por_ingreso)
    if oi is not None:
        p["consolidado"]["otros_ingresos_avion"] = oi
    if bandera:
        p["consolidado"]["otros_movimientos"]["otros_ingresos_por_avion"] = True
    return p


def _general(**kw) -> BalanceGeneralRequest:
    return BalanceGeneralRequest.model_validate(_general_dict(**kw))


def _bloque(ws, matricula: str) -> int:
    """Fila de «UTILIDAD ANTES DE GASTOS USD» del bloque de `matricula` en la
    hoja 'balance' del general."""
    inicio = _fila_de(ws, matricula)
    return next(
        c.row for c in ws["A"] if c.row > inicio and c.value == "UTILIDAD ANTES DE GASTOS USD"
    )


def test_general_hoja_otros_ingresos_aviones(registros) -> None:
    req = _general()
    data = render_balance_general_xlsx(req)
    wb, wv = _libros(data)
    nombres = wb.sheetnames
    assert HOJA_GENERAL in nombres
    assert nombres.index(HOJA_GENERAL) == nombres.index("repartidos a aviones") + 1
    assert nombres.index("inventario") == nombres.index(HOJA_GENERAL) + 1
    ws, wsv = wb[HOJA_GENERAL], wv[HOJA_GENERAL]
    assert ws["A1"].value == _TITULO_OTROS_INGRESOS_GENERAL
    assert [ws.cell(row=4, column=c).value for c in range(1, 4)] == [
        "TOTAL NETO MXN",
        "TOTAL USD (suma de los aviones)",
        None,
    ]
    assert [ws.cell(row=7, column=c).value for c in range(1, 12)] == ENCABEZADOS
    # Detalle con la matrícula al frente y teñido con el color del avión.
    assert ws["C8"].value == "XA-TST · ING-4 · TRANSF INTERBANCARIA SPEI · angel alvarez"
    assert ws["C8"].fill.fgColor.rgb.endswith("3B82F6")
    assert ws["A5"].value == "=ROUND(SUM(I8:I9),2)"
    oi = req.consolidado.otros_ingresos_avion
    assert wsv["A5"].value == oi.total_mxn
    # TOTAL USD = VALOR (cada libro convirtió con su T.C.: sin fórmula).
    assert ws["B5"].value == oi.usd == req.aviones[0].otros_ingresos_avion.usd
    assert [c.value for c in ws["A"] if c.value == NOTA_GENERAL] == [NOTA_GENERAL]
    assert all(h != HOJA_GENERAL for h, _ in registros[-1].degradadas())
    assert verificar_libro(data)


def test_general_sin_filas_no_hay_hoja_pero_los_bloques_llevan_la_linea() -> None:
    req = _general(filas_tst=[])
    assert req.consolidado.otros_ingresos_avion is None
    data = render_balance_general_xlsx(req)
    wb, wv = _libros(data)
    assert HOJA_GENERAL not in wb.sheetnames
    ws = wb["balance"]
    for mat in ("XA-TST", "XB-DOS"):
        k = _bloque(ws, mat)
        assert ws.cell(row=k + 5, column=1).value == _FILA_OTROS_INGRESOS
        assert wv["balance"].cell(row=k + 5, column=2).value == 0
    assert verificar_libro(data)


def test_general_bloques_con_la_linea_como_valor_y_comentario() -> None:
    req = _general()
    data = render_balance_general_xlsx(req)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    for mat, avion in (("XA-TST", req.aviones[0]), ("XB-DOS", req.aviones[1])):
        k = _bloque(ws, mat)
        linea = ws.cell(row=k + 5, column=2)
        assert ws.cell(row=k + 5, column=1).value == _FILA_OTROS_INGRESOS
        # VALOR (el general no tiene la hoja del avión) con su nota.
        assert linea.value == avion.balance.otros_ingresos_avion_usd
        assert linea.comment is not None and linea.comment.text == NOTA_CASCADA
        assert ws.cell(row=k + 6, column=2).value == (
            f"=ROUND(B{k}-B{k + 1}-B{k + 2}-B{k + 3}-B{k + 4}+B{k + 5},2)"
        )
        assert ws.cell(row=k + 8, column=2).value == f"=ROUND(B{k + 6}-B{k + 7},2)"
        assert wsv.cell(row=k + 6, column=2).value == avion.balance.utilidad_despues_usd
        assert wsv.cell(row=k + 8, column=2).value == avion.balance.utilidad_cobrada_usd
    assert verificar_libro(data)


def test_general_con_comisiones_la_linea_va_tras_las_dos_previas(registros) -> None:
    """Cascada del general con «antes de comisiones» y «(−) COMISIONES» (API
    0.0.66): la línea de otros ingresos sigue en k + 5 contando desde UTILIDAD
    ANTES DE GASTOS (que ya va corrida dos filas)."""
    p = tcg._payload()
    tst = next(a for a in p["aviones"] if a["matricula"] == "XA-TST")
    con = _con_otros_ingresos(tst, [_FILA_MXN])
    p["aviones"] = [con if a["matricula"] == "XA-TST" else a for a in p["aviones"]]
    p["empresa"] = t_empresa._empresa_api(p)
    req = BalanceGeneralRequest.model_validate(p)
    data = render_balance_general_xlsx(req)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    k = _bloque(ws, "XA-TST")
    assert ws.cell(row=k - 2, column=1).value == (
        "UTILIDAD ANTES DE GASTOS USD (antes de comisiones)"
    )
    assert ws.cell(row=k + 5, column=1).value == _FILA_OTROS_INGRESOS
    assert ws.cell(row=k + 6, column=2).value == (
        f"=ROUND(B{k}-B{k + 1}-B{k + 2}-B{k + 3}-B{k + 4}+B{k + 5},2)"
    )
    b = req.aviones[[a.matricula for a in req.aviones].index("XA-TST")].balance
    assert wsv.cell(row=k + 6, column=2).value == b.utilidad_despues_usd
    assert wsv.cell(row=k + 8, column=2).value == b.utilidad_cobrada_usd
    # XB-DOS no trae la llave (API previo en ese libro): su cascada de siempre.
    kd = _bloque(ws, "XB-DOS")
    assert ws.cell(row=kd + 5, column=1).value == DESPUES
    assert verificar_libro(data)


@pytest.mark.parametrize("bandera", [True, None, False], ids=["true", "null", "false"])
def test_general_fila2_de_otros_movimientos(bandera) -> None:
    p = _general_dict(bandera=False)
    if bandera is not None:
        p["consolidado"]["otros_movimientos"]["otros_ingresos_por_avion"] = bandera
    wb = load_workbook(
        BytesIO(render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p)))
    )
    a2 = wb["otros movimientos"]["A2"].value
    if bandera is True:
        assert a2.endswith(OM_COLA)
        assert a2.count(OM_COLA) == 1
    else:
        assert OM_COLA not in a2
        assert a2.endswith("antes llamada 'gastos VuelaTour').")


def test_general_sin_la_llave_libro_de_hoy() -> None:
    data = render_balance_general_xlsx(tf._general())
    assert HOJA_GENERAL not in load_workbook(BytesIO(data)).sheetnames
    assert _huella(data) == _HUELLA_GENERAL_HEAD


def test_libro_json_del_general_exporta_la_hoja() -> None:
    libro = libro_balance_general(_general())
    hojas = {h["nombre"]: h for h in libro["hojas"]}
    assert HOJA_GENERAL in hojas
    celdas = {(c["r"], c["c"]): c for c in hojas[HOJA_GENERAL]["celdas"]}
    assert celdas[(5, 1)]["f"] == "=ROUND(SUM(I8:I9),2)"
    assert celdas[(8, 7)]["f"] == "=ROUND(D8,2)"
    assert celdas[(9, 9)]["f"] == "=ROUND(G9-H9,2)"
    assert libro_json.VERSION_LIBRO == 1


# ---------------------------------------------------------------------------
# 4) Rutas con el JSON del API 0.0.75.
# ---------------------------------------------------------------------------

TOKEN = "tok-otros-ingresos"
client = TestClient(app)


def test_rutas_aceptan_el_payload_0075(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    individual = _con_otros_ingresos(tf._payload_tst(), [_FILA_MXN, _FILA_USD])
    res = client.post("/pdf/balance-avion-xlsx", json=individual, headers=headers)
    assert res.status_code == 200, res.text
    assert HOJA in load_workbook(BytesIO(res.content)).sheetnames
    res = client.post("/pdf/balance-general-xlsx", json=_general_dict(), headers=headers)
    assert res.status_code == 200, res.text
    assert HOJA_GENERAL in load_workbook(BytesIO(res.content)).sheetnames
    res = client.post("/pdf/balance-avion-libro", json=individual, headers=headers)
    assert res.status_code == 200, res.text
    assert HOJA in [h["nombre"] for h in res.json()["hojas"]]


# ---------------------------------------------------------------------------
# 5) Payload de FORMA REAL del API de producción 0.0.76 (en los libros, la
#    misma forma que el 0.0.75) — revisión del 14-oct-2026, hallazgo F12: la
#    huella de «pyservices nuevo + API ≤ 0.0.76 = libro de siempre» se toma
#    con lo que ese API manda SIEMPRE (regla de comisiones con las llaves de
#    la cascada del 0.0.66, 'permisos' con `tc_por_fila`, las tres banderas
#    raíz, socio `es_empresa`, bloque VUELATOUR y la hoja «Otros ingresos» con
#    filas SIN las llaves del 0.0.77), no con `tf._payload_tst()`, que no trae
#    `tc_por_fila`. Con `api_0077` el MISMO libro como lo manda el API 0.0.77:
#    filas con `tc_usd`/`neto_usd`/`tc_usd_origen`/`tc_usd_nota`, la hoja con
#    `usd_por_ingreso` (TOTAL USD = Σ NETO USD, la cascada re-armada) y la
#    bandera raíz `factura_vuelatour_devolucion`.
# ---------------------------------------------------------------------------

# Las llaves de fila del API 0.0.77 (conversión única por ingreso): una fila
# del 0.0.75/0.0.76 no las trae.
_LLAVES_FILA_0077 = ("tc_usd", "neto_usd", "tc_usd_origen", "tc_usd_nota")


def _fila_0076(fila: dict) -> dict:
    """La fila tal como la manda el API 0.0.75/0.0.76 (sin las llaves del 0.0.77)."""
    return {k: v for k, v in fila.items() if k not in _LLAVES_FILA_0077}


def _libro_forma_real(
    libro: dict, socios: list, filas: list[dict], *, api_0077: bool = False
) -> dict:
    """Un libro de avión como lo arma `buildPayload` del API 0.0.76 (o del
    0.0.77 con `api_0077`)."""
    p = tcg._con_api_0066(libro)
    if p["permisos"]["filas"]:
        p = tf._con_permisos(p, tf._permisos_con_tc(_tc_hojas(p)))
    else:
        p["permisos"]["tc_por_fila"] = True
    # Llaves que pyservices ignora y las tres banderas raíz que el API 0.0.74–
    # 0.0.76 manda SIEMPRE en el libro de un avión.
    p["tc_promedio"] = p["totales"]["tc_promedio"]
    p["horas_voladas_hr"] = p["totales"]["tiempo_vuelo"]
    p["sin_provision_vendedor"] = True
    p["factura_vuelatour_segun_cobros"] = True
    if api_0077:
        p["factura_vuelatour_devolucion"] = True
    p["comision_vendedor_solo_gasto_real"] = True
    t_empresa._poner_socios(p, socios)
    if api_0077:
        return _con_otros_ingresos(p, filas, por_ingreso=True)
    return _con_otros_ingresos(p, [_fila_0076(f) for f in filas])


def _individual_forma_real(*, api_0077: bool = False) -> dict:
    return _libro_forma_real(
        t_comisiones._libro_tst(),
        t_empresa._SOCIOS_CON_EMPRESA["XA-TST"],
        [_FILA_MXN, _FILA_USD],
        api_0077=api_0077,
    )


def _general_forma_real(variante: str = "mensual", *, api_0077: bool = False) -> dict:
    """El Balance general del API 0.0.76 (o 0.0.77): banderas del consolidado
    (factura y comisión solo como gasto real), las de 'otros movimientos'
    (`sin_provision_vendedor`, `comision_vendedor_solo_gasto_real`,
    `otros_ingresos_por_avion`), la hoja consolidada «otros ingresos aviones»
    (solo XA-TST trae filas) y el bloque VUELATOUR (empresa)."""
    tst = _individual_forma_real(api_0077=api_0077)
    tst["avion_color"] = _COLOR_TST
    dos = _libro_forma_real(
        t_comisiones._libro_dos(),
        t_empresa._SOCIOS_CON_EMPRESA["XB-DOS"],
        [],
        api_0077=api_0077,
    )
    p = tf._general_payload([tst, dos])
    cons = p["consolidado"]
    cons["factura_vuelatour_segun_cobros"] = True
    if api_0077:
        cons["factura_vuelatour_devolucion"] = True
    cons["comision_vendedor_solo_gasto_real"] = True
    cons["otros_ingresos_avion"] = _fusion_flota([tst, dos], por_ingreso=api_0077)
    om = cons["otros_movimientos"]
    om["sin_provision_vendedor"] = True
    om["comision_vendedor_solo_gasto_real"] = True
    om["otros_ingresos_por_avion"] = True
    p["empresa"] = t_empresa._empresa_api(p)
    p["variante"] = variante
    return p


def _huella_json(libro: dict) -> str:
    return hashlib.sha256(
        json.dumps(libro, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()


# Huellas con el payload de FORMA REAL del API 0.0.76, calculadas con el
# código de HEAD b66df98 (14-oct-2026) ANTES de la conversión por ingreso:
# miembros del .xlsx (todos salvo docProps/core.xml) y JSON de la vista en
# pantalla (el API la pide con `vinculos_gasto: true`). Si una se mueve, un API
# ≤ 0.0.76 ya no ve su libro de siempre (y una ventana de despliegue con
# pyservices nuevo y API viejo publicaría otro libro).
_HUELLAS_FORMA_REAL_0076 = {
    "individual": (
        "64dd04d5923657a9937a8e70cb3ee225a11b90fb6cfbe56581df32f7c105f7b1",
        "2a931d29ba6848d32c674d2cc8e4713565082d8a0bd3cdef601a8f611c05227d",
    ),
    "general-mensual": (
        "70d85d47474cca664d4d2b0ec9f3bfa08f12a9d794872814569978c485f0dc8f",
        "986914f70208905678ffc547d2eab138b21761bcc6d08579cae9bb6a3e974fc1",
    ),
    "general-general": (
        "51a1e9111921988dd6996c3ab21718350a3ff74405732140ba59a16828bb1388",
        "50162ff2bc4a5559cc45712cb42ee9a8557eb7a7cc76b0ef09b3f6ca33dda126",
    ),
}


def _libros_forma_real(libro: str, *, api_0077: bool = False):
    """(.xlsx, JSON de la vista) del libro `libro` con la forma real."""
    if libro == "individual":
        p = _individual_forma_real(api_0077=api_0077)
        req = BalanceAvionRequest.model_validate(p)
        vista = BalanceAvionRequest.model_validate({**p, "vinculos_gasto": True})
        return render_balance_avion_xlsx(req), libro_balance_avion(vista)
    p = _general_forma_real(libro.removeprefix("general-"), api_0077=api_0077)
    req = BalanceGeneralRequest.model_validate(p)
    vista = BalanceGeneralRequest.model_validate({**p, "vinculos_gasto": True})
    return render_balance_general_xlsx(req), libro_balance_general(vista)


@pytest.mark.parametrize("libro", list(_HUELLAS_FORMA_REAL_0076))
def test_forma_real_0076_libros_byte_a_byte(libro) -> None:
    """pyservices nuevo + API ≤ 0.0.76 (sin `usd_por_ingreso` ni
    `factura_vuelatour_devolucion`) = el libro de siempre, byte a byte, en el
    individual y en las dos variantes del general, .xlsx y vista en pantalla."""
    data, vista = _libros_forma_real(libro)
    assert verificar_libro(data)
    assert (_huella(data), _huella_json(vista)) == _HUELLAS_FORMA_REAL_0076[libro]


def test_forma_real_0076_con_las_llaves_de_fila_sin_bandera_igual() -> None:
    """Las cuatro llaves de fila SIN la bandera de la hoja (o con ella en
    null/false) no mueven nada: la bandera manda, no la presencia de llaves."""
    referencia = _HUELLAS_FORMA_REAL_0076["individual"][0]
    for bandera in (..., None, False):
        p = _individual_forma_real()
        p["otros_ingresos_avion"]["filas"] = copy.deepcopy([_FILA_MXN, _FILA_USD])
        if bandera is not ...:
            p["otros_ingresos_avion"]["usd_por_ingreso"] = bandera
        data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
        assert _huella(data) == referencia, bandera


# ---------------------------------------------------------------------------
# 6) Conversión ÚNICA por ingreso (14-oct-2026, API 0.0.77, bandera
#    `usd_por_ingreso`): T.C. A USD (J), NETO USD (K) y TOTAL USD = Σ NETO USD.
# ---------------------------------------------------------------------------


def test_esquemas_0077_aditivos_y_liberales() -> None:
    fila = BalanceAvionOtroIngresoFila()
    assert (fila.tc_usd, fila.neto_usd, fila.tc_usd_origen, fila.tc_usd_nota) == (None,) * 4
    assert BalanceAvionHojaOtrosIngresos().usd_por_ingreso is None
    fila = BalanceAvionOtroIngresoFila.model_validate(_FILA_MXN)
    assert (fila.tc_usd, fila.neto_usd, fila.tc_usd_origen) == (16.9479, 2746.7, "OFICIAL_DEL_DIA")
    assert fila.tc_usd_nota == NOTA_TC_OFICIAL_ING4
    # Origen y nota: texto recortado; vacío o no-texto ⇒ None (jamás un 422).
    casos = (("  GASTO_LIGADO ", "GASTO_LIGADO"), ("   ", None), (7, None), ({}, None))
    for valor, esperado in casos:
        fila = BalanceAvionOtroIngresoFila(tc_usd_origen=valor, tc_usd_nota=valor)
        assert (fila.tc_usd_origen, fila.tc_usd_nota) == (esperado, esperado)
    hoja = BalanceAvionHojaOtrosIngresos.model_validate(_hoja_api_por_ingreso([_FILA_MXN]))
    assert hoja.usd_por_ingreso is True and hoja.usd == 2746.7


def test_por_ingreso_hoja_individual_columnas_y_formulas(registros) -> None:
    req = _individual(por_ingreso=True)
    hoja = req.otros_ingresos_avion
    assert hoja.usd == r2(2746.70 + 95.0) == 2841.7
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws, wsv = wb[HOJA], wv[HOJA]
    assert [ws.cell(row=7, column=c).value for c in range(1, 14)] == ENCABEZADOS_POR_INGRESO
    assert ws.max_column == 13
    # Resumen: TOTAL NETO MXN | T.C. («por ingreso (col. J)») | TOTAL USD.
    assert [ws.cell(row=4, column=c).value for c in range(1, 4)] == [
        "TOTAL NETO MXN",
        "T.C.",
        "TOTAL USD",
    ]
    assert ws["A5"].value == "=ROUND(SUM(I8:I9),2)"
    assert ws["B5"].value == _TC_POR_INGRESO
    assert ws["B5"].font.i and ws["B5"].font.color.rgb.endswith(MUTED)
    assert ws["C5"].value == "=ROUND(SUM(K8:K9),2)"
    assert wsv["A5"].value == hoja.total_mxn
    assert wsv["C5"].value == hoja.usd == 2841.7
    # Pesos con el T.C. oficial del día: J azul con su nota; K = I ÷ J.
    assert ws["J8"].value == 16.9479 and ws["J8"].number_format == TC
    assert ws["J8"].fill.fgColor.rgb.endswith(TC_OFICIAL_FILL)
    assert ws["J8"].comment.text == NOTA_TC_OFICIAL_ING4
    assert ws["K8"].value == "=ROUND(I8/J8,2)"
    assert wsv["K8"].value == 2746.7
    # Dólares: J vacía (cuenta su neto tal cual) y K = I ÷ F (su T.C. propio).
    assert ws["J9"].value is None and ws["J9"].fill.fill_type is None
    assert ws["J9"].comment is None
    assert ws["K9"].value == "=ROUND(I9/F9,2)"
    assert wsv["K9"].value == 95
    # Lo demás de la fila, corrido dos columnas: CUENTA (L) y NOTAS (M).
    assert (ws["G8"].value, ws["I8"].value) == ("=ROUND(D8,2)", "=ROUND(G8,2)")
    assert (ws["G9"].value, ws["I9"].value) == ("=ROUND(D9*F9,2)", "=ROUND(G9-H9,2)")
    assert (ws["L8"].value, ws["M8"].value) == (
        "EJE SCOTIA · Scotiabank",
        "reembolso del pago del serv de 50 hrs en toledo, fact 53",
    )
    assert ws["M8"].alignment.wrap_text
    # Nota al pie (A:M), anchos y panel fijo.
    pie = ws.cell(row=11, column=1)
    assert pie.value == NOTA_HOJA_POR_INGRESO and pie.alignment.wrap_text
    assert any(str(rg) == "A11:M11" for rg in ws.merged_cells.ranges)
    assert NOTA_HOJA not in [c.value for c in ws["A"]]
    anchos = [12, 22, 50, 15, 9, 11, 15, 14, 15, 11, 14, 22, 36]
    assert [ws.column_dimensions[chr(64 + i)].width for i in range(1, 14)] == anchos
    assert ws.freeze_panes == "A8"
    # Ninguna fórmula nueva volvió a valor y todas reproducen al API.
    assert registros[-1].degradadas() == _DEGRADADAS_TST
    verificadas = {(h, c) for h, c, *_ in verificar_libro(data)}
    assert {(HOJA, c) for c in ("K8", "K9", "C5", "A5")} <= verificadas


def test_por_ingreso_cascada_cita_c5_con_la_nota_nueva(registros) -> None:
    req = _individual(por_ingreso=True)
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    k = _fila_de(ws, _FILA_OTROS_INGRESOS)
    linea = ws.cell(row=k, column=2)
    # La cascada SIGUE citando el TOTAL USD de la hoja (C5).
    assert linea.value == "='Otros ingresos'!$C$5"
    assert linea.comment.text == NOTA_CASCADA_POR_INGRESO
    b = req.balance
    assert wsv.cell(row=k, column=2).value == b.otros_ingresos_avion_usd == 2841.7
    assert ws.cell(row=k + 1, column=2).value == (
        f"=ROUND(B{k - 5}-B{k - 4}-B{k - 3}-B{k - 2}-B{k - 1}+B{k},2)"
    )
    assert wsv.cell(row=k + 1, column=2).value == b.utilidad_despues_usd
    assert wsv.cell(row=k + 3, column=2).value == b.utilidad_cobrada_usd
    socio = _fila_de(ws, "Socio A")
    assert [wsv.cell(row=socio + i, column=3).value for i in range(2)] == [
        s.monto_usd for s in b.socios
    ]
    # La utilidad sube EXACTAMENTE lo que suman los NETO USD (2,746.70 de
    # ING-4 al oficial de su día, como el reparto y el tablero, + 95.00).
    sin = tf._payload_tst()["balance"]
    assert b.utilidad_despues_usd == r2(sin["utilidad_despues_usd"] + 2841.7)
    assert registros[-1].degradadas() == _DEGRADADAS_TST
    assert verificar_libro(data)


def test_por_ingreso_reembolso_ligado_t_c_de_su_gasto_sin_azul() -> None:
    """Reembolso LIGADO a un gasto de este libro (API 0.0.77, hallazgo F3): J =
    el T.C. con que el libro convierte ese gasto, SIN relleno azul (no es el
    oficial) y con la nota del API; K = I ÷ J reproduce 2,638.61."""
    req = _individual([_FILA_LIGADA], por_ingreso=True)
    assert req.otros_ingresos_avion.usd == 2638.61
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws = wb[HOJA]
    assert ws["J8"].value == 17.237699
    assert ws["J8"].fill.fill_type is None
    assert ws["J8"].comment.text == NOTA_TC_LIGADO
    assert ws["K8"].value == "=ROUND(I8/J8,2)"
    assert wv[HOJA]["K8"].value == 2638.61
    assert ws["C5"].value == "=ROUND(SUM(K8:K8),2)"
    assert verificar_libro(data)


def test_por_ingreso_t_c_propio_sin_azul_ni_nota() -> None:
    propio = {**_FILA_MXN, "tc_usd": 17.5, "neto_usd": 2660.05, "tc_usd_origen": "PROPIO"}
    propio["tc_usd_nota"] = None
    data = render_balance_avion_xlsx(_individual([propio], por_ingreso=True))
    wb, wv = _libros(data)
    ws = wb[HOJA]
    assert ws["J8"].value == 17.5
    assert ws["J8"].fill.fill_type is None and ws["J8"].comment is None
    assert ws["K8"].value == "=ROUND(I8/J8,2)"
    assert wv[HOJA]["K8"].value == 2660.05
    assert verificar_libro(data)


def test_por_ingreso_pesos_sin_tc_k_y_c5_vacias() -> None:
    """Pesos sin T.C. propio ni oficial de su día: sin NETO USD (jamás un 0
    falso), TOTAL USD vacío y la cascada vacía, como manda el API."""
    req = _individual([_FILA_MXN_SIN_TC, _FILA_USD], por_ingreso=True)
    assert req.otros_ingresos_avion.usd is None
    assert req.balance.utilidad_despues_usd is None
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws = wb[HOJA]
    assert (ws["J8"].value, ws["K8"].value) == (None, None)
    assert ws["J8"].fill.fill_type is None
    assert ws["K9"].value == "=ROUND(I9/F9,2)"
    assert wv[HOJA]["A5"].value == r2(46550.8 + 1662.5)
    assert ws["C5"].value is None
    k = _fila_de(wb["balance"], _FILA_OTROS_INGRESOS)
    assert wb["balance"].cell(row=k, column=2).value is None
    assert wb["balance"].cell(row=k, column=2).comment is None
    assert verificar_libro(data)


def test_por_ingreso_usd_sin_tc_suma_su_neto_como_valor() -> None:
    """USD sin T.C. (el CHECK del API lo impide): sin pesos, pero su NETO USD
    sí suma (regla del reparto); K va como VALOR (no hay T.C. que citar)."""
    req = _individual([_FILA_MXN, _FILA_USD_SIN_TC], por_ingreso=True)
    assert req.otros_ingresos_avion.usd == r2(2746.70 + 200.0)
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws = wb[HOJA]
    assert (ws["G9"].value, ws["I9"].value, ws["J9"].value) == (None, None, None)
    assert ws["K9"].value == 200
    assert ws["C5"].value == "=ROUND(SUM(K8:K9),2)"
    assert wv[HOJA]["C5"].value == 2946.7
    assert wv[HOJA]["A5"].value == 46550.8
    assert verificar_libro(data)


def test_por_ingreso_neto_usd_inconsistente_se_degrada_a_valor(registros) -> None:
    """Si el NETO USD del API no es I ÷ J al centavo, la celda vuelve a VALOR
    (manda el número del API) y el TOTAL USD sigue siendo su suma."""
    raro = {**_FILA_MXN, "neto_usd": 2700.0}
    req = _individual([raro, _FILA_USD], por_ingreso=True)
    assert req.otros_ingresos_avion.usd == 2795.0
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    assert wb[HOJA]["K8"].value == 2700.0
    assert wb[HOJA]["K9"].value == "=ROUND(I9/F9,2)"
    assert wb[HOJA]["C5"].value == "=ROUND(SUM(K8:K9),2)"
    assert wv[HOJA]["C5"].value == 2795.0
    assert registros[-1].degradadas() == [*_DEGRADADAS_TST, (HOJA, "K8")]
    assert verificar_libro(data)


def test_por_ingreso_hoja_vacia() -> None:
    req = _individual([], por_ingreso=True)
    assert (req.otros_ingresos_avion.total_mxn, req.otros_ingresos_avion.usd) == (0, 0)
    data = render_balance_avion_xlsx(req)
    wb, wv = _libros(data)
    ws = wb[HOJA]
    assert ws["A8"].value == _SIN_OTROS_INGRESOS
    assert any(str(rg) == "A8:M8" for rg in ws.merged_cells.ranges)
    assert [ws.cell(row=7, column=c).value for c in range(1, 14)] == ENCABEZADOS_POR_INGRESO
    # Sin filas no hay rango que sumar: los dos totales como valor (0).
    assert (ws["A5"].value, ws["B5"].value, ws["C5"].value) == (0, _TC_POR_INGRESO, 0)
    k = _fila_de(wb["balance"], _FILA_OTROS_INGRESOS)
    assert wv["balance"].cell(row=k, column=2).value == 0
    assert wb["balance"].cell(row=k, column=2).comment.text == NOTA_CASCADA_POR_INGRESO
    assert verificar_libro(data)


def test_por_ingreso_general_hoja_y_bloques(registros) -> None:
    req = _general(por_ingreso=True)
    oi = req.consolidado.otros_ingresos_avion
    assert oi.usd_por_ingreso is True
    data = render_balance_general_xlsx(req)
    wb, wv = _libros(data)
    ws, wsv = wb[HOJA_GENERAL], wv[HOJA_GENERAL]
    assert [ws.cell(row=7, column=c).value for c in range(1, 14)] == ENCABEZADOS_POR_INGRESO
    assert [ws.cell(row=4, column=c).value for c in range(1, 4)] == [
        "TOTAL NETO MXN",
        "TOTAL USD (suma de los aviones)",
        None,
    ]
    # B5 deja de ser VALOR: la suma de NETO USD de toda la flota.
    assert ws["A5"].value == "=ROUND(SUM(I8:I9),2)"
    assert ws["B5"].value == "=ROUND(SUM(K8:K9),2)"
    assert wsv["B5"].value == oi.usd == 2841.7
    assert ws["C8"].value == "XA-TST · ING-4 · TRANSF INTERBANCARIA SPEI · angel alvarez"
    assert ws["C8"].fill.fgColor.rgb.endswith("3B82F6")
    assert ws["J8"].fill.fgColor.rgb.endswith(TC_OFICIAL_FILL)
    assert ws["J8"].comment.text == NOTA_TC_OFICIAL_ING4
    assert (ws["K8"].value, ws["K9"].value) == ("=ROUND(I8/J8,2)", "=ROUND(I9/F9,2)")
    assert [c.value for c in ws["A"] if c.value == NOTA_GENERAL_POR_INGRESO] == [
        NOTA_GENERAL_POR_INGRESO
    ]
    assert NOTA_GENERAL not in [c.value for c in ws["A"]]
    assert all(h != HOJA_GENERAL for h, _ in registros[-1].degradadas())
    # Cada bloque de 'balance': la línea como VALOR con la nota nueva (la
    # elige la hoja de SU libro; XB-DOS sin ingresos: 0 con la misma nota).
    bal, balv = wb["balance"], wv["balance"]
    for mat, avion in (("XA-TST", req.aviones[0]), ("XB-DOS", req.aviones[1])):
        k = _bloque(bal, mat)
        linea = bal.cell(row=k + 5, column=2)
        assert bal.cell(row=k + 5, column=1).value == _FILA_OTROS_INGRESOS
        assert linea.value == avion.balance.otros_ingresos_avion_usd
        assert linea.comment.text == NOTA_CASCADA_POR_INGRESO
        assert balv.cell(row=k + 8, column=2).value == avion.balance.utilidad_cobrada_usd
    assert verificar_libro(data)


def test_por_ingreso_la_nota_de_cada_bloque_la_elige_su_libro() -> None:
    """Payload mixto (un libro con la bandera y otro sin ella): cada bloque
    de 'balance' del general lleva la nota de la hoja de SU libro."""
    p = _general_dict(filas_dos=[_FILA_USD])
    tst = p["aviones"][0]
    p["aviones"][0] = _con_otros_ingresos(
        {k: v for k, v in tst.items() if k != "otros_ingresos_avion"},
        [_FILA_MXN, _FILA_USD],
        por_ingreso=True,
    )
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    ws = load_workbook(BytesIO(data))["balance"]
    notas = {
        mat: ws.cell(row=_bloque(ws, mat) + 5, column=2).comment.text
        for mat in ("XA-TST", "XB-DOS")
    }
    assert notas == {"XA-TST": NOTA_CASCADA_POR_INGRESO, "XB-DOS": NOTA_CASCADA}


def test_libro_json_por_ingreso_lleva_las_columnas_nuevas() -> None:
    req = _individual(por_ingreso=True)
    libro = libro_balance_avion(req)
    hojas = {h["nombre"]: h for h in libro["hojas"]}
    oi = {(c["r"], c["c"]): c for c in hojas[HOJA]["celdas"]}
    assert oi[(7, 10)]["v"] == "T.C. A USD" and oi[(7, 11)]["v"] == "NETO USD"
    assert oi[(8, 10)]["v"] == 16.9479 and oi[(8, 10)]["n"] == NOTA_TC_OFICIAL_ING4
    estilo = libro["estilos"][oi[(8, 10)]["s"]]
    assert estilo["bg"].upper().endswith(TC_OFICIAL_FILL)
    assert oi[(8, 11)]["f"] == "=ROUND(I8/J8,2)" and oi[(8, 11)]["v"] == 2746.7
    assert oi[(5, 2)]["v"] == _TC_POR_INGRESO
    assert oi[(5, 3)]["f"] == "=ROUND(SUM(K8:K9),2)" and oi[(5, 3)]["v"] == 2841.7


# ---------------------------------------------------------------------------
# 7) Forma real del API 0.0.77 contra la del 0.0.76: SOLO cambian la hoja
#    «Otros ingresos», los números de la cascada que la suman y la nota al pie
#    de FACTURA VUELATOUR; todo lo demás, idéntico.
# ---------------------------------------------------------------------------


def _hojas_iguales(antes: bytes, despues: bytes, *, menos: set[str]) -> list[str]:
    nombres = load_workbook(BytesIO(antes)).sheetnames
    assert load_workbook(BytesIO(despues)).sheetnames == nombres
    distintas = [
        n
        for n in nombres
        if n not in menos
        and t_extension._firma_hoja(antes, n) != t_extension._firma_hoja(despues, n)
    ]
    return distintas


def _solo_cambia_la_nota_de_factura(antes: bytes, despues: bytes, maestra: str) -> None:
    a = load_workbook(BytesIO(antes), data_only=True)[maestra]
    b = load_workbook(BytesIO(despues), data_only=True)[maestra]
    assert (a.max_row, a.max_column) == (b.max_row, b.max_column)
    diferencias = [
        (ca.coordinate, ca.value, cb.value)
        for fila_a, fila_b in zip(a.iter_rows(), b.iter_rows(), strict=True)
        for ca, cb in zip(fila_a, fila_b, strict=True)
        if (ca.value, ca.number_format, ca.fill.fgColor.rgb)
        != (cb.value, cb.number_format, cb.fill.fgColor.rgb)
    ]
    assert len(diferencias) == 1, diferencias
    _coord, va, vb = diferencias[0]
    assert (va, vb) == (_NOTA_FACTURA_VUELATOUR_SEGUN_COBROS, _NOTA_FACTURA_VUELATOUR_DEVOLUCION)


def test_forma_real_0077_individual_solo_cambia_lo_esperado(registros) -> None:
    antes, _ = _libros_forma_real("individual")
    despues, _ = _libros_forma_real("individual", api_0077=True)
    assert verificar_libro(despues)
    assert _hojas_iguales(antes, despues, menos={HOJA, "balance", MAESTRA_TST}) == []
    _solo_cambia_la_nota_de_factura(antes, despues, MAESTRA_TST)
    wb, wv = _libros(despues)
    ws = wb[HOJA]
    assert ws["C5"].value == "=ROUND(SUM(K8:K9),2)"
    assert wv[HOJA]["C5"].value == 2841.7
    k = _fila_de(wb["balance"], _FILA_OTROS_INGRESOS)
    assert wb["balance"].cell(row=k, column=2).value == "='Otros ingresos'!$C$5"
    assert wv["balance"].cell(row=k, column=2).value == 2841.7
    assert wb["balance"].cell(row=k, column=2).comment.text == NOTA_CASCADA_POR_INGRESO
    assert all(h != HOJA for h, _ in registros[-1].degradadas())


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_forma_real_0077_general_solo_cambia_lo_esperado(variante) -> None:
    antes, _ = _libros_forma_real(f"general-{variante}")
    despues, _ = _libros_forma_real(f"general-{variante}", api_0077=True)
    assert verificar_libro(despues)
    menos = {HOJA_GENERAL, "balance", "reporte horas FLOTA"}
    assert _hojas_iguales(antes, despues, menos=menos) == []
    _solo_cambia_la_nota_de_factura(antes, despues, "reporte horas FLOTA")
    wb, wv = _libros(despues)
    assert wb[HOJA_GENERAL]["B5"].value == "=ROUND(SUM(K8:K9),2)"
    assert wv[HOJA_GENERAL]["B5"].value == 2841.7


def test_rutas_aceptan_el_payload_0077(monkeypatch, request) -> None:
    """JSON crudo como lo manda el API 0.0.77 (individual, general y la vista
    en pantalla): 200 con las columnas nuevas."""
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    individual = _individual_forma_real(api_0077=True)
    res = client.post("/pdf/balance-avion-xlsx", json=individual, headers=headers)
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content))[HOJA]
    assert (ws["J7"].value, ws["K8"].value) == ("T.C. A USD", "=ROUND(I8/J8,2)")
    general = _general_forma_real(api_0077=True)
    res = client.post("/pdf/balance-general-xlsx", json=general, headers=headers)
    assert res.status_code == 200, res.text
    assert load_workbook(BytesIO(res.content))[HOJA_GENERAL]["B5"].value == "=ROUND(SUM(K8:K9),2)"
    res = client.post(
        "/pdf/balance-avion-libro", json={**individual, "vinculos_gasto": True}, headers=headers
    )
    assert res.status_code == 200, res.text
    hoja = next(h for h in res.json()["hojas"] if h["nombre"] == HOJA)
    assert any(c["r"] == 7 and c["c"] == 11 and c["v"] == "NETO USD" for c in hoja["celdas"])


# ---------------------------------------------------------------------------
# Huellas de miembros (todos salvo docProps/core.xml) de los libros SIN la
# llave, calculadas con el código de HEAD d58faa1 (13-oct-2026), ANTES de
# este cambio: si una se mueve, un API ≤ 0.0.74 ya no ve su libro de siempre.
# ---------------------------------------------------------------------------

_HUELLA_INDIVIDUAL_HEAD = "4bb86191e6fdd486fb46a332b612b3b763908f6ec8c14ff9e8d9eaee4e1a506f"
_HUELLA_GENERAL_HEAD = "dd8a3740c2ef71791c94cb35ac02a3f71123fc49979b3d15a91569cbfa56ba90"
