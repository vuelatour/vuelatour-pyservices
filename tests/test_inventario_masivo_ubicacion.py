"""Plantilla de ALTA MASIVA de inventario y el catálogo de ubicaciones
(revisión adversaria 25-sep-2026, API 0.0.35).

Desde el catálogo `inventario_ubicacion` el API liga el texto de la columna
«Ubicación» que coincide con una ubicación del catálogo y deja el resto como
ubicación «(anterior)». La fila de ejemplo decía «Bodega Cancún», que ya NO
es del catálogo: quien copiara el ejemplo daba de alta productos en ámbar.
Ahora el ejemplo usa una ubicación sembrada y las instrucciones explican la
columna; una plantilla descargada ANTES (ejemplo con «Bodega Cancún») sigue
descartando su fila de ejemplo intacta.
"""

import base64
from io import BytesIO

from openpyxl import load_workbook

from app.schemas.inventario import ParseInventarioRequest, PlantillaInventarioRequest
from app.services.inventario_masivo import (
    CAMPOS,
    EJEMPLO,
    INSTRUCCIONES,
    parse_inventario,
    render_plantilla_inventario,
)

# Semilla de la migración 20260925000001 (nombres EXACTOS del cliente).
UBICACIONES_SEMBRADAS = {
    "Oficina vieja",
    "Oficina nueva",
    "Locker del aeropuerto",
    "Bodega del taller de Mérida",
    "Bodega del taller de Cozumel",
}
COL_UBICACION = CAMPOS.index("ubicacion") + 1


def _plantilla(ubicacion_ejemplo: str | None = None, **celdas_fila_2) -> bytes:
    """La plantilla real; opcionalmente con la fila 2 (ejemplo) retocada."""
    wb = load_workbook(BytesIO(render_plantilla_inventario(PlantillaInventarioRequest())))
    ws = wb.active
    if ubicacion_ejemplo is not None:
        ws.cell(row=2, column=COL_UBICACION, value=ubicacion_ejemplo)
    for campo, valor in celdas_fila_2.items():
        ws.cell(row=2, column=CAMPOS.index(campo) + 1, value=valor)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(xlsx: bytes):
    return parse_inventario(ParseInventarioRequest(
        archivo_base64=base64.b64encode(xlsx).decode(), filename="plantilla.xlsx",
    )).filas


def test_el_ejemplo_usa_una_ubicacion_del_catalogo() -> None:
    assert EJEMPLO[CAMPOS.index("ubicacion")] in UBICACIONES_SEMBRADAS
    ws = load_workbook(BytesIO(_plantilla())).active
    assert ws.cell(row=1, column=COL_UBICACION).value == "Ubicación"
    assert ws.cell(row=2, column=COL_UBICACION).value in UBICACIONES_SEMBRADAS


def test_las_instrucciones_explican_la_ubicacion() -> None:
    linea = next(t for t in INSTRUCCIONES if t.startswith("Ubicación:"))
    assert "Vacío = sin ubicación" in linea
    assert "«(anterior)»" in linea
    assert "«Mover a…»" in linea


def test_ejemplo_intacto_se_descarta_nuevo_y_de_plantilla_vieja() -> None:
    assert _parse(_plantilla()) == []
    # Plantilla descargada ANTES del catálogo: el ejemplo decía «Bodega Cancún».
    assert _parse(_plantilla(ubicacion_ejemplo="Bodega Cancún")) == []


def test_ejemplo_editado_si_es_un_alta() -> None:
    """Editar el ejemplo lo vuelve una fila real (regla de siempre), también
    si solo se cambió la ubicación a otra del catálogo o a un texto libre."""
    filas = _parse(_plantilla(existencia_inicial=5))
    assert len(filas) == 1 and filas[0].existencia_inicial == 5
    filas = _parse(_plantilla(ubicacion_ejemplo="Locker del aeropuerto"))
    assert [f.ubicacion for f in filas] == ["Locker del aeropuerto"]
    filas = _parse(_plantilla(ubicacion_ejemplo="Corner"))
    assert [f.ubicacion for f in filas] == ["Corner"]
