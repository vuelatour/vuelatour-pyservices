"""Apodos fuera de los documentos: CLAVE = solo el folio y columnas CLIENTE /
RUTA / Nº VUELO CLIENTE junto a ella (23-oct-2026, API 0.0.105).

Pedido del cliente (con capturas del Excel «reporte horas FLOTA» y 'otros
movimientos' del Balance general): «siempre debemos ocultar los apodos con la
clave del vuelo, para dejar solo las columnas de (Num de cotización, nombre
completo, fecha, ruta y consecutivo del vuelo)» y después «mejor eliminar esos
apodos de todos lados, y de hecho faltan las columnas de nombre, ruta en
algunos otros lados».

Lo que se congela aquí:
(a) Con el API 0.0.105 (`clave` = «#219»; `cliente`, `ruta` y
    `consecutivo_cliente` en cada fila de vuelo) cada hoja trae los
    encabezados nuevos en el orden pedido y las celdas correctas; las
    fórmulas y totales corrieron con las columnas.
(b) Con un API viejo (`clave` «#219 · ana02-010926 · Ana Paty» y
    `clave_cliente`) la CLAVE se pinta TAL CUAL —no se parte— y CLIENTE /
    Nº VUELO quedan vacías; `clave_cliente` ya no cambia ni un byte.
(c) Ningún documento (balances, libro JSON, Libro Dinero, reparto PDF/XLSX)
    contiene un texto con forma de ID por cliente («beh05-011026») cuando el
    API ya no lo manda.
"""

from __future__ import annotations

import copy
import json
import re
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import get_settings
from app.main import app
from app.schemas.reparto import RepartoPdfRequest
from app.schemas.reportes import (
    BalanceAvionRequest,
    BalanceAvionVuelo,
    BalanceCierreVuelo,
    BalanceGeneralRequest,
    BalanceOtroMovimientoFila,
    DineroOtroIngresoFila,
    DineroVueloFila,
    DineroXlsxRequest,
    consecutivo_liberal,
    texto_fila_liberal,
)
from app.services import cierre_mes
from app.services._formato import (
    ANCHO_CLIENTE_BASE,
    ANCHO_CLIENTE_TOPE,
    ANCHO_NUM_VUELO_CLIENTE,
    ENCABEZADO_CLIENTE,
    ENCABEZADO_NUM_VUELO_CLIENTE,
    ancho_columna_cliente,
    ancho_columna_por_contenido,
)
from app.services.balance_avion_xlsx import (
    _COBRANZA_HEADERS,
    _COLS_FLOTA_AFAC,
    _COLUMNAS_PROVISION,
    _DISP_FLOTA,
    _DISP_GENERAL,
    _DISP_MENSUAL,
    _OM_HEADERS,
    HORAS,
    RED,
    libro_balance_avion,
    libro_balance_general,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from app.services.balance_avion_xlsx import (
    MONEY as MONEY_BALANCE,
)
from app.services.dinero_xlsx import _COBRO1_COL, _COLS, render_dinero_xlsx
from app.services.reparto_pdf import render_reparto_pdf
from app.services.reparto_xlsx import render_reparto_xlsx
from tests import _cierre_mes_payloads as pay

# Forma de un ID de vuelo por cliente («beh05-011026»: alias + nº + DDMMAA).
ID_APODO = re.compile(r"\b[a-z]+\d{2,}-\d{6}\b")
# Lo que mandaba el API 0.0.75–0.0.104 en la CLAVE (texto tal cual).
CLAVE_VIEJA_MAESTRA = "#219 · ana02-010926 · Ana Paty"
CLAVE_VIEJA_OM = "#223 · morena08-010926"
CLAVE_VIEJA_DINERO = "beh05-011026 (50 % N4142R)"
FRENTE_MAESTRA = ["CLAVE", "CLIENTE", "Nº VUELO\nCLIENTE", "FECHA", "RUTA", "ESTADO"]
FRENTE_DINERO = ["CLAVE", "CLIENTE", "Nº VUELO\nCLIENTE", "FECHA", "RUTA"]
FRENTE_OM = ["clave", "cliente", "ruta", "nº vuelo\ncliente", "fecha\nvuelo", "concepto"]

# ---------------------------------------------------------------------------
# Payloads
# ---------------------------------------------------------------------------


def _vuelo(clave: str, **extra) -> dict:
    return {
        "clave": clave,
        "fecha": "2026-10-21",
        "ruta": "CUN-HOL-CUN",
        "estado": "COMPLETADO",
        "status_cobro": "Cobrado",
        "cobros": [{"fecha": "2026-10-21", "monto_mxn": 30000.0}],
        "cobrado_real_mxn": 30000.0,
        "cobrado_mxn": 30000.0,
        "por_cobrar_mxn": 0.0,
        "por_cobrar_usd": 0.0,
        "total_mxn": 30000.0,
        "total_cotizacion_mxn": 35000.0,
        "horas_cobradas": 1.5,
        **extra,
    }


def _vuelos_nuevos() -> list[dict]:
    return [
        _vuelo("#219", cliente="Ana Paty", consecutivo_cliente=2),
        _vuelo("#223", cliente="Morena Torres de la Fuente", consecutivo_cliente=8),
        # Cliente sin número de vuelo (el API manda null): celda vacía.
        _vuelo("#300", cliente="Cliente Sin Id", consecutivo_cliente=None),
    ]


def _vuelos_viejos() -> list[dict]:
    return [
        _vuelo(CLAVE_VIEJA_MAESTRA, clave_cliente="ana02-010926"),
        _vuelo(CLAVE_VIEJA_OM, clave_cliente="morena08-010926"),
    ]


_TOTALES = {
    "cobrado_real_mxn": 90000.0,
    "por_cobrar_usd": 0.0,
    "total_mxn": 90000.0,
    "total_cotizacion_mxn": 105000.0,
}


def _individual(vuelos: list[dict] | None = None) -> dict:
    return {
        "matricula": "XA-VGV",
        "periodo_desde": "2026-10-01",
        "periodo_hasta": "2026-10-31",
        "vuelos": _vuelos_nuevos() if vuelos is None else vuelos,
        "totales": _TOTALES,
    }


def _om_nuevo() -> dict:
    return {
        "filas": [
            {
                "clave": "#223",
                "cliente": "Morena Torres de la Fuente",
                "ruta": "CUN-HOL-CUN",
                "consecutivo_cliente": 8,
                "estado": "CONFIRMADO",
                "concepto_ingreso": "TUA CUN",
                "ingreso_mxn": 1250.5,
                "remanente_mxn": 1250.5,
                "factura": "A-0424",
            },
            {
                "clave": "#115",
                "cliente": "Riviera Maya Tours",
                "ruta": "CUN-MID",
                "consecutivo_cliente": 2,
                "estado": "CANCELADO",
                "concepto_egreso": "pago comisión vendedor (Ale) · gasto real",
                "egreso_mxn": 100.0,
                "concepto_ingreso": "Pernocta",
                "ingreso_mxn": 300.0,
                "remanente_mxn": 200.0,
            },
        ],
        "filas_sueltas": [
            {
                "clave": "ING-3",
                "concepto_ingreso": "Intereses",
                "ingreso_mxn": 85.33,
                "remanente_mxn": 85.33,
            }
        ],
    }


def _om_viejo() -> dict:
    return {
        "filas": [
            {
                "clave": CLAVE_VIEJA_OM,
                "clave_cliente": "morena08-010926",
                "estado": "CONFIRMADO",
                "concepto_ingreso": "TUA CUN",
                "ingreso_mxn": 1250.5,
                "remanente_mxn": 1250.5,
            }
        ],
        "filas_sueltas": [],
    }


def _general(vuelos: list[dict] | None = None, om: dict | None = None, variante=None) -> dict:
    p = {
        "periodo_desde": "2026-10-01",
        "periodo_hasta": "2026-10-31",
        "resumen": [{"matricula": "XA-VGV", "color": "#FF0000", "vuelos": 3}],
        "consolidado": {
            "matricula": "FLOTA",
            "periodo_desde": "2026-10-01",
            "periodo_hasta": "2026-10-31",
            "vuelos": _vuelos_nuevos() if vuelos is None else vuelos,
            "totales": _TOTALES,
            "otros_movimientos": _om_nuevo() if om is None else om,
        },
    }
    for v in p["consolidado"]["vuelos"]:
        v["avion_color"] = "#FF0000"
    if variante is not None:
        p["variante"] = variante
    return p


def _dinero(nuevo: bool = True) -> dict:
    multi = {
        "clave": "#219 (50 % N4142R)" if nuevo else CLAVE_VIEJA_DINERO,
        "matricula": "N4142R",
        "color": "#1D4ED8",
        "fecha": "2026-10-21",
        "ruta": "CUN-PTU",
        "tiempo": 1.0,
        "participacion": 0.5,
        "multi_avion": True,
        "total_cobrado_mxn": 10000.0,
    }
    servicio = {
        "clave": "#220" if nuevo else "vtservicio",
        "matricula": "XA-VGV",
        "color": "#FF0000",
        "fecha": "2026-10-22",
        "ruta": "CUN-MID",
        "tiempo": 1.2,
        "total_cobrado_mxn": 0.0,
    }
    oi = {
        "clave": "#219" if nuevo else "beh05-011026",
        "fecha_vuelo": "2026-10-21",
        "concepto_ingreso": "TUA CUN",
        "ingreso_mxn": 500.0,
        "remanente_mxn": 500.0,
    }
    if nuevo:
        multi.update(cliente="Beh Kay", consecutivo_cliente=5)
        servicio.update(cliente="Vuelo de servicio")
        oi.update(cliente="Beh Kay", ruta="CUN-PTU", consecutivo_cliente=5)
    else:
        multi["clave_cliente"] = "beh05-011026"
        oi["clave_cliente"] = "beh05-011026"
    return {
        "periodo_desde": "2026-10-01",
        "periodo_hasta": "2026-10-31",
        "generado": "2026-10-31",
        "vuelos": [multi, servicio],
        "otros_ingresos": [
            oi,
            {"clave": "ING-7", "concepto_ingreso": "Intereses", "ingreso_mxn": 12.5,
             "remanente_mxn": 12.5},
        ],
    }  # fmt: skip


def _xlsx_individual(p: dict) -> bytes:
    return render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))


def _xlsx_general(p: dict) -> bytes:
    return render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))


def _xlsx_dinero(p: dict) -> bytes:
    return render_dinero_xlsx(DineroXlsxRequest.model_validate(p))


def _ws(data: bytes, nombre: str):
    return load_workbook(BytesIO(data))[nombre]


def _fila(ws, r: int, n: int) -> list:
    return [ws.cell(row=r, column=c).value for c in range(1, n + 1)]


def _ancho(ws, letra: str) -> float:
    return ws.column_dimensions[letra].width


def _textos(data: bytes) -> list[str]:
    """TODOS los textos de celda de TODAS las hojas de un .xlsx."""
    wb = load_workbook(BytesIO(data))
    return [
        c.value for ws in wb.worksheets for fila in ws.iter_rows() for c in fila
        if isinstance(c.value, str)
    ]  # fmt: skip


def _texto_pdf(data: bytes) -> str:
    return re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(BytesIO(data)).pages))


def _miembros(data: bytes) -> dict[str, bytes]:
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in z.namelist() if n != "docProps/core.xml"}


def _sin_clave_cliente(o):
    o = copy.deepcopy(o)

    def quita(x):
        if isinstance(x, dict):
            x.pop("clave_cliente", None)
            for v in x.values():
                quita(v)
        elif isinstance(x, list):
            for v in x:
                quita(v)

    quita(o)
    return o


# ---------------------------------------------------------------------------
# _formato y esquemas
# ---------------------------------------------------------------------------


def test_ancho_de_cliente_por_contenido_con_tope() -> None:
    assert ancho_columna_por_contenido([], base=14, tope=40) == 14
    assert ancho_columna_por_contenido(["x" * 90], base=14, tope=40) == 40
    assert ancho_columna_por_contenido(["ab", "a · b · c"], base=1, tope=40) == 11
    assert ancho_columna_cliente([None, "Ana Paty"]) == ANCHO_CLIENTE_BASE == 22
    assert ancho_columna_cliente(["Morena Torres de la Fuente"]) == 28
    assert ancho_columna_cliente(["Nombre Larguísimo " * 4]) == ANCHO_CLIENTE_TOPE == 40
    assert ENCABEZADO_CLIENTE == "CLIENTE" and ENCABEZADO_NUM_VUELO_CLIENTE == "Nº VUELO\nCLIENTE"
    assert ANCHO_NUM_VUELO_CLIENTE == 9


@pytest.mark.parametrize(
    "modelo", [BalanceAvionVuelo, BalanceOtroMovimientoFila, DineroVueloFila, DineroOtroIngresoFila]
)
def test_esquemas_aditivos_y_liberales(modelo) -> None:
    assert modelo().cliente is None and modelo().consecutivo_cliente is None
    x = modelo.model_validate(
        {"clave": "#219", "cliente": " Ana Paty ", "consecutivo_cliente": "02",
         "clave_cliente": "ana02-010926", "llave_futura": 1}
    )  # fmt: skip
    assert (x.clave, x.cliente, x.consecutivo_cliente) == ("#219", "Ana Paty", 2)
    # `clave_cliente` sigue aceptándose (API previo) aunque ya no se use.
    assert x.clave_cliente == "ana02-010926"
    # Basura ⇒ None, jamás un 422.
    for basura in (7, "", "   ", ["x"], None, True):
        assert modelo.model_validate({"cliente": basura}).cliente is None
    for basura in (True, False, None, "", "dos", 0, -3, 2.5, float("nan"), ["2"]):
        assert modelo.model_validate({"consecutivo_cliente": basura}).consecutivo_cliente is None
    assert modelo.model_validate({"consecutivo_cliente": 3.0}).consecutivo_cliente == 3
    if modelo in (BalanceOtroMovimientoFila, DineroOtroIngresoFila):
        # Las hojas que no tenían RUTA la ganan ahora, liberal ('' del API sin
        # tramos ⇒ None). La maestra y 'dinero-vlos' conservan la suya.
        assert modelo.model_validate({"ruta": " CUN-HOL-CUN "}).ruta == "CUN-HOL-CUN"
        assert modelo.model_validate({"ruta": ""}).ruta is None
        assert modelo.model_validate({"ruta": 5}).ruta is None


def test_validadores_liberales() -> None:
    assert texto_fila_liberal(" x ") == "x" and texto_fila_liberal(3) is None
    assert consecutivo_liberal("12") == 12 and consecutivo_liberal(1) == 1
    assert consecutivo_liberal(0) is None and consecutivo_liberal(True) is None


# ---------------------------------------------------------------------------
# (a) Hojas con el API 0.0.105
# ---------------------------------------------------------------------------


def test_disposiciones_de_la_hoja_maestra_corren_dos_columnas() -> None:
    for disp in (_DISP_MENSUAL, _DISP_FLOTA, _DISP_GENERAL):
        letra = disp.letra
        assert (letra["clave"], letra["cliente"], letra["consecutivo_cliente"]) == ("A", "B", "C")
        assert (letra["fecha"], letra["ruta"], letra["estado"]) == ("D", "E", "F")
        assert letra["horas_cobradas"] == "G"
    # PERMISO AFAC (PROVISIÓN): de T a V en el individual y la flota mensual.
    assert _DISP_MENSUAL.letra["permiso_afac_mxn"] == "V"
    assert _DISP_FLOTA.letra["permiso_afac_mxn"] == "V"
    assert [c[1] for c in _COLS_FLOTA_AFAC[:6]] == FRENTE_MAESTRA


def test_individual_hoja_maestra_y_cobranza() -> None:
    data = _xlsx_individual(_individual())
    ws = _ws(data, "reporte horas XA-VGV")
    assert _fila(ws, 2, 6) == FRENTE_MAESTRA
    assert _fila(ws, 3, 6) == ["#219", "Ana Paty", 2, "21/10/2026", "CUN-HOL-CUN", "COMPLETADO"]
    assert _fila(ws, 4, 3) == ["#223", "Morena Torres de la Fuente", 8]
    assert _fila(ws, 5, 3) == ["#300", "Cliente Sin Id", None]
    # El nº de vuelo es un entero sin formato de dinero; HORAS COBRADAS sigue
    # con el suyo en su nueva columna G.
    assert ws["C3"].number_format == "General" and ws["G3"].value == 1.5
    assert ws["G3"].number_format == HORAS
    # Las constantes del encabezado no se mueven (B1 y D1) y los paneles se
    # congelan a la derecha de RUTA.
    assert ws["A1"].value == "Permiso AFAC USD/hr" and ws["C1"].value == "Factor IVA costos"
    assert ws["D1"].value == 1.16 and ws.freeze_panes == "F3"
    assert (_ancho(ws, "A"), _ancho(ws, "B"), _ancho(ws, "C")) == (16, 28, 9)
    assert (_ancho(ws, "D"), _ancho(ws, "E"), _ancho(ws, "F")) == (12, 34, 12)
    # TOTALES en su fila, sin nada en las columnas de texto nuevas.
    assert _fila(ws, 6, 3) == ["TOTALES", None, None]

    wc = _ws(data, "cobranza")
    assert _fila(wc, 7, len(_COBRANZA_HEADERS)) == _COBRANZA_HEADERS
    assert _COBRANZA_HEADERS[:6] == FRENTE_MAESTRA
    assert _fila(wc, 8, 9) == [
        "#219", "Ana Paty", 2, "21/10/2026", "CUN-HOL-CUN", "COMPLETADO", "Cobrado", 30000, 30000,
    ]  # fmt: skip
    # Fórmulas corridas dos columnas: POR COBRAR = TOTAL (H) − COBRADO (I).
    assert wc["J8"].value == "=ROUND(H8-I8,2)"
    tot = 8 + 3
    assert wc[f"A{tot}"].value == "TOTALES"
    assert wc[f"H{tot}"].value == "=ROUND(SUM(H8:H10),2)"
    assert wc[f"K{tot}"].value == "=ROUND(SUM(K8:K10),2)"
    assert wc[f"L{tot}"].value == '=ROUND(SUMIF(F8:F10,"<>CANCELADO",L8:L10),2)'
    assert wc[f"M{tot}"].value == "=ROUND(SUM(M8:M10),2)"
    assert wc["A5"].value == f"=H{tot}" and wc["F5"].value == f"=M{tot}"
    assert (_ancho(wc, "A"), _ancho(wc, "B"), _ancho(wc, "C"), _ancho(wc, "E")) == (22, 28, 9, 26)
    assert wc.freeze_panes == "A8"


@pytest.mark.parametrize("variante", [None, "mensual", "general"])
def test_general_hoja_de_la_flota_en_sus_variantes(variante) -> None:
    p = _general(variante=variante)
    if variante == "mensual":
        # Con la tarifa AFAC por fila (API 0.0.69) el juego es `_COLS_FLOTA_AFAC`.
        p["consolidado"]["vuelos"][0]["permiso_afac_usd_hr"] = 100.0
    data = _xlsx_general(p)
    ws = _ws(data, "reporte horas FLOTA")
    assert _fila(ws, 2, 6) == FRENTE_MAESTRA
    assert _fila(ws, 3, 6) == ["#219", "Ana Paty", 2, "21/10/2026", "CUN-HOL-CUN", "COMPLETADO"]
    assert ws.freeze_panes == "F3"
    # El color del avión sigue SOLO en la celda de la CLAVE.
    assert ws["A3"].fill.fgColor.rgb.endswith("FF0000")
    assert not ws["B3"].fill.fill_type
    if variante == "mensual":
        tarifa = ws.cell(row=2, column=len(_COLS_FLOTA_AFAC)).value
        assert tarifa == "AFAC USD/HR\n(tarifa del avión)"
    wc = _ws(data, "cobranza")
    assert _fila(wc, 7, len(_COBRANZA_HEADERS)) == _COBRANZA_HEADERS
    assert _fila(wc, 8, 3) == ["#219", "Ana Paty", 2]


def test_general_otros_movimientos() -> None:
    data = _xlsx_general(_general())
    ws = _ws(data, "otros movimientos")
    assert _OM_HEADERS == [
        "clave", "cliente", "ruta", "nº vuelo\ncliente", "fecha\nvuelo", "concepto", "egreso",
        "fecha", "concepto", "ingreso", "fecha", "remanente", "factura\nvuelatour",
    ]  # fmt: skip
    assert _fila(ws, 3, 13) == _OM_HEADERS
    assert _fila(ws, 4, 13) == [
        "#223", "Morena Torres de la Fuente", "CUN-HOL-CUN", 8, None, None, None, None,
        "TUA CUN", 1250.5, None, "=ROUND(J4-G4,2)", "A-0424",
    ]  # fmt: skip
    # El estado no normal sigue junto a la clave (en rojo si CANCELADO).
    assert _fila(ws, 5, 7) == [
        "#115 · CANCELADO", "Riviera Maya Tours", "CUN-MID", 2, None,
        "pago comisión vendedor (Ale) · gasto real", 100,
    ]  # fmt: skip
    assert ws["A5"].font.color.rgb.endswith(RED)
    # Fila suelta (ING-n): sin cliente, ruta ni nº de vuelo.
    assert _fila(ws, 8, 10) == [
        "ING-3", None, None, None, None, None, None, None, "Intereses", 85.33,
    ]  # fmt: skip
    # TOTALES en egreso (G), ingreso (J) y remanente (L).
    assert _fila(ws, 9, 13) == [
        "TOTALES", None, None, None, None, None, "=ROUND(SUM(G4:G8),2)", None, None,
        "=ROUND(SUM(J4:J8),2)", None, "=ROUND(J9-G9,2)", None,
    ]  # fmt: skip
    assert ws["D4"].number_format == "General" and ws["J4"].number_format == MONEY_BALANCE
    assert _COLUMNAS_PROVISION == (6, 7)
    assert (_ancho(ws, "A"), _ancho(ws, "B"), _ancho(ws, "D")) == (18, 28, 9)
    assert (_ancho(ws, "F"), _ancho(ws, "I")) == (28, 32)


def test_dinero_vuelos_y_otros_ingresos() -> None:
    data = _xlsx_dinero(_dinero())
    ws = _ws(data, "dinero-vlos")
    assert [c[1] for c in _COLS[:5]] == FRENTE_DINERO
    assert _fila(ws, 4, 5) == FRENTE_DINERO
    assert _fila(ws, 5, 6) == ["#219 (50 % N4142R)", "Beh Kay", 5, "21/10/2026", "CUN-PTU", 1]
    # Vuelo de servicio: el API manda el nombre «Vuelo de servicio».
    assert _fila(ws, 6, 6) == ["#220", "Vuelo de servicio", None, "22/10/2026", "CUN-MID", 1.2]
    assert _fila(ws, 7, 6) == ["TOTALES", None, None, None, None, 2.2]
    assert ws.freeze_panes == "F5"
    assert ws.cell(row=4, column=_COBRO1_COL).value == "COBRO 1\nFECHA"
    assert (_ancho(ws, "A"), _ancho(ws, "B"), _ancho(ws, "C"), _ancho(ws, "E")) == (14, 22, 9, 20)
    # La nota de la CLAVE (matrícula / multi-avión) sigue en la CLAVE.
    assert ws["A5"].comment is not None and "N4142R" in ws["A5"].comment.text

    wo = _ws(data, "Otros ingresos")
    assert _fila(wo, 3, 13) == _OM_HEADERS
    assert _fila(wo, 4, 13) == [
        "#219", "Beh Kay", "CUN-PTU", 5, "21/10/2026", None, None, None, "TUA CUN", 500, None,
        500, None,
    ]  # fmt: skip
    assert _fila(wo, 5, 4) == ["ING-7", None, None, None]
    assert _fila(wo, 6, 13) == [
        "TOTALES", None, None, None, None, None, 0, None, None, 512.5, None, 512.5, None,
    ]  # fmt: skip
    assert wo["D4"].number_format == "General"
    assert (_ancho(wo, "A"), _ancho(wo, "B"), _ancho(wo, "D"), _ancho(wo, "F")) == (15, 22, 9, 26)


def test_libro_json_hereda_columnas_y_paneles() -> None:
    libro = libro_balance_avion(BalanceAvionRequest.model_validate(_individual()))
    hoja = next(h for h in libro["hojas"] if h["nombre"] == "reporte horas XA-VGV")
    assert {"c": 2, "ancho": 28.0} in hoja["columnas"]
    assert hoja["inmovilizar"] == {"filas": 2, "columnas": 5}


# ---------------------------------------------------------------------------
# (b) API viejo: la CLAVE tal cual, CLIENTE vacía, `clave_cliente` ignorada
# ---------------------------------------------------------------------------


def test_api_viejo_clave_tal_cual_y_cliente_vacia() -> None:
    ind = _xlsx_individual(_individual(_vuelos_viejos()))
    ws = _ws(ind, "reporte horas XA-VGV")
    assert _fila(ws, 3, 4) == [CLAVE_VIEJA_MAESTRA, None, None, "21/10/2026"]
    assert _fila(ws, 4, 3) == [CLAVE_VIEJA_OM, None, None]
    assert _ancho(ws, "A") == 16 and _ancho(ws, "B") == ANCHO_CLIENTE_BASE
    wc = _ws(ind, "cobranza")
    assert _fila(wc, 8, 3) == [CLAVE_VIEJA_MAESTRA, None, None] and _ancho(wc, "A") == 22

    gen = _xlsx_general(_general(_vuelos_viejos(), _om_viejo()))
    assert _fila(_ws(gen, "reporte horas FLOTA"), 3, 3) == [CLAVE_VIEJA_MAESTRA, None, None]
    wo = _ws(gen, "otros movimientos")
    assert _fila(wo, 4, 5) == [CLAVE_VIEJA_OM, None, None, None, None]
    assert _ancho(wo, "A") == 18

    din = _xlsx_dinero(_dinero(nuevo=False))
    wd = _ws(din, "dinero-vlos")
    assert _fila(wd, 5, 3) == [CLAVE_VIEJA_DINERO, None, None]
    assert _fila(wd, 6, 3) == ["vtservicio", None, None]
    assert _ancho(wd, "A") == 14
    wi = _ws(din, "Otros ingresos")
    assert _fila(wi, 4, 4) == ["beh05-011026", None, None, None] and _ancho(wi, "A") == 15


@pytest.mark.parametrize("libro", ["individual", "general", "dinero"])
def test_clave_cliente_ya_no_cambia_ni_un_byte(libro) -> None:
    con = {
        "individual": lambda: _individual(_vuelos_viejos()),
        "general": lambda: _general(_vuelos_viejos(), _om_viejo()),
        "dinero": lambda: _dinero(nuevo=False),
    }[libro]()
    render = {"individual": _xlsx_individual, "general": _xlsx_general, "dinero": _xlsx_dinero}
    assert _miembros(render[libro](con)) == _miembros(render[libro](_sin_clave_cliente(con)))


# ---------------------------------------------------------------------------
# (c) Ningún documento imprime un ID por cliente cuando el API no lo manda
# ---------------------------------------------------------------------------


def test_balances_y_dinero_sin_ids_por_cliente() -> None:
    documentos = [
        _xlsx_individual(_individual()),
        _xlsx_general(_general()),
        _xlsx_general(_general(variante="general")),
        _xlsx_dinero(_dinero()),
        # Los de forma real del lote 0.0.77 con el cierre de mes (listas de
        # vuelos que pasan / vienen con `clave` + `cliente`).
        _xlsx_individual(pay.individual_con_cierre()),
        _xlsx_general(pay.general_con_cierre("mensual")),
        _xlsx_dinero(pay.dinero_con_cierre()),
    ]
    for data in documentos:
        textos = _textos(data)
        assert textos, "libro vacío"
        assert not [t for t in textos if ID_APODO.search(t)], data[:0]
    vista = libro_balance_avion(BalanceAvionRequest.model_validate(_individual()))
    assert not ID_APODO.search(json.dumps(vista, ensure_ascii=False))
    vista_g = libro_balance_general(BalanceGeneralRequest.model_validate(_general()))
    assert not ID_APODO.search(json.dumps(vista_g, ensure_ascii=False))


def test_reparto_pdf_y_xlsx_sin_ids_y_con_folio_y_cliente() -> None:
    req = RepartoPdfRequest.model_validate(pay.reparto_con_cierre())
    texto = _texto_pdf(render_reparto_pdf(req))
    assert not ID_APODO.search(texto)
    assert "#318" in texto and "Ambulancia Internacional" in texto
    assert "#316" in texto and "Mike & Nelson <MN>" in texto
    textos = _textos(render_reparto_xlsx(req))
    assert not [t for t in textos if ID_APODO.search(t)]
    assert "#318" in textos and "Ambulancia Internacional" in textos


# ---------------------------------------------------------------------------
# Listas del cierre: «#folio · Cliente»
# ---------------------------------------------------------------------------


def test_nombre_vuelo_del_cierre_folio_y_cliente() -> None:
    # API 0.0.105: la clave es solo el folio ⇒ se agrega el cliente.
    v = BalanceCierreVuelo(clave="#316", folio="316", cliente="Mike Nelson")
    assert cierre_mes.nombre_vuelo(v) == "#316 · Mike Nelson"
    # API previo: la clave ya trae al cliente ⇒ tal cual (no se duplica).
    v = BalanceCierreVuelo(clave="#316 · Mike Nelson", folio="316", cliente="Mike Nelson")
    assert cierre_mes.nombre_vuelo(v) == "#316 · Mike Nelson"
    # Sin cliente ⇒ la clave tal cual; sin clave ⇒ «#folio · cliente».
    assert cierre_mes.nombre_vuelo(BalanceCierreVuelo(clave="#316")) == "#316"
    assert cierre_mes.nombre_vuelo(BalanceCierreVuelo(folio="316", cliente="Ana")) == "#316 · Ana"
    assert cierre_mes.nombre_vuelo(BalanceCierreVuelo()) == "—"
    linea = cierre_mes.linea_vuelo_libro(
        BalanceCierreVuelo(clave="#316", cliente="Mike Nelson", venta_mxn=36632.89, cobrado_mxn=0)
    )
    assert linea == "#316 · Mike Nelson — venta $36,632.89 MXN, cobrado $0.00"


def test_bloque_cierre_del_balance_pinta_folio_y_cliente() -> None:
    p = pay.individual_con_cierre()
    for it in (*p["cierre"]["pasan"], *p["cierre"]["vienen"]):
        it["clave"] = f"#{it['folio']}"  # como lo manda el API 0.0.105
    ws = _ws(_xlsx_individual(p), "reporte horas XA-TST")
    claves = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert "#499 · Cliente Nueve" in claves and "#403 · Traslado" in claves


# ---------------------------------------------------------------------------
# Por HTTP: las rutas aceptan el payload del API 0.0.105
# ---------------------------------------------------------------------------

TOKEN = "tok-claves-sin-apodo"
client = TestClient(app)


def test_rutas_aceptan_el_payload_nuevo(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    res = client.post("/pdf/balance-avion-xlsx", json=_individual(), headers=headers)
    assert res.status_code == 200, res.text
    assert _fila(_ws(res.content, "reporte horas XA-VGV"), 3, 3) == ["#219", "Ana Paty", 2]
    res = client.post("/pdf/balance-general-xlsx", json=_general(), headers=headers)
    assert res.status_code == 200, res.text
    assert _fila(_ws(res.content, "otros movimientos"), 4, 4) == [
        "#223", "Morena Torres de la Fuente", "CUN-HOL-CUN", 8,
    ]  # fmt: skip
    res = client.post("/pdf/dinero-xlsx", json=_dinero(), headers=headers)
    assert res.status_code == 200, res.text
    assert _fila(_ws(res.content, "dinero-vlos"), 5, 3) == ["#219 (50 % N4142R)", "Beh Kay", 5]
    res = client.post("/pdf/balance-avion-libro", json=_individual(), headers=headers)
    assert res.status_code == 200, res.text
