"""Conciliación de INGRESOS con IA (24-sep-2026): `/conciliacion/sugerir-abonos`.

La IA PROPONE y la persona confirma. Lo que se congela aquí es que el modelo
nunca pueda afirmar más de lo que los hechos comprobables permiten (monto,
nombre del ordenante, días, cuenta), que jamás proponga «registrar como otro
ingreso» dinero que ya tiene un cobro de vuelo exacto (doble conteo), y que
el payload REAL del API no reviente con 422. NINGUNA prueba llama a Claude.
"""

import json

import anthropic
import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError

from app.config import get_settings
from app.main import app
from app.schemas.conciliacion import (
    AbonoParaSugerir,
    CandidatoParaAbono,
    ConciliacionSugerirAbonosRequest,
)
from app.schemas.reportes import (
    BalanceHojaOtrosMovimientos,
    BalanceOtroMovimientoFila,
    DineroOtroIngresoFila,
    DineroXlsxRequest,
)
from app.services import conciliacion_abonos
from app.services._dominio import BLOQUE_DOMINIO, DOMINIO_VERSION
from app.services.balance_avion_xlsx import _NOTA_GASTOS_EMPRESA, _hoja_otros_movimientos
from app.services.conciliacion_abonos import (
    MAX_TOKENS,
    MOTIVO_CANDIDATO_REPETIDO,
    MOTIVO_COBRO_EXACTO,
    MOTIVO_ID_DESCARTADO,
    MOTIVO_INGRESO_EXACTO,
    TIMEOUT_S,
    monto_exacto,
    nombre_en_descripcion,
    sugerir_abonos,
)
from app.services.dinero_xlsx import render_dinero_xlsx

TOKEN = "secreto-de-prueba"
client = TestClient(app)

CATEGORIAS = [
    "OTRO_INGRESO",
    "ANTICIPO_CLIENTE",
    "INGRESO_BANCARIO",
    "REEMBOLSO_DEVOLUCION",
    "VENTA_ACTIVO",
    "APORTACION_PRESTAMO",
]

# --- Casos REALES de prod (24-sep-2026) -----------------------------------
MOV_CRISTY = "7f1c2a9e-0b1d-4c55-9a52-6a1e3c0f2b01"
MOV_LETICIA = "a9d4e1b2-7c3f-4e8a-b6d1-2f0e9c8b7a02"
MOV_DANI = "c3b2a1f0-9e8d-4c7b-a6f5-e4d3c2b1a003"
COBRO_235 = "COBRO_VUELO:5e6f7a8b-9c0d-4e1f-8a2b-3c4d5e6f7a80"
COBRO_OTRO = "COBRO_VUELO:0a1b2c3d-4e5f-4a6b-8c7d-9e0f1a2b3c4d"
ING_7 = "INGRESO:11112222-3333-4444-8555-666677778888"
COBRO_USD = "COBRO_VUELO:99998888-7777-4666-8555-444433332222"

ABONO_CRISTY = {
    "id": MOV_CRISTY,
    "fecha": "2026-09-08",
    "monto": 19380.0,
    "monto_bruto": None,
    "comision_monto": None,
    "descripcion": "MARIA CRISTINA CHAVEZ BADIOLA : vuelo cristy badiola",
    "referencia": None,
    "cuenta_alias": "Paywise",
    "cuenta_moneda": "MXN",
    "cuenta_tipo": "PASARELA",
    "candidato_ids": [COBRO_235, COBRO_OTRO],
}
# Cobro del vuelo #235: TRANSFERENCIA 20,400.00 − comisión 1,020.00 = 19,380.00.
CAND_235 = {
    "id": COBRO_235,
    "tipo": "COBRO_VUELO",
    "fecha": "2026-09-08T12:00:00-05:00",
    "monto": 20400.0,
    "comision": 1020.0,
    "neto": 19380.0,
    "moneda": "MXN",
    "metodo": "TRANSFERENCIA",
    "cliente": "Cristy Chavez",
    "folio": 235,
    "referencia": None,
}
CAND_OTRO = {
    "id": COBRO_OTRO,
    "tipo": "COBRO_VUELO",
    "fecha": "2026-09-05",
    "monto": 19000.0,
    "comision": None,
    "neto": 19000.0,
    "moneda": "MXN",
    "metodo": "TRANSFERENCIA",
    "cliente": "Hotel Xcaret",
    "folio": 240,
}


def _req(abonos: list[dict], candidatos: list[dict]) -> ConciliacionSugerirAbonosRequest:
    return ConciliacionSugerirAbonosRequest.model_validate(
        {"abonos": abonos, "candidatos": candidatos, "categorias": CATEGORIAS}
    )


def _sug(movimiento_id: str, **kw) -> dict:
    base = {
        "movimiento_id": movimiento_id,
        "candidato_id": None,
        "confianza": 0.0,
        "razon": "",
        "evidencias": [],
        "alternativas": [],
        "accion": "REVISAR",
        "categoria_sugerida": None,
        "motivo_sin_match": None,
    }
    base.update(kw)
    return base


def _payload_usuario(cliente) -> dict:
    texto = cliente.texto_usuario()
    return json.loads(texto[texto.index("{") : texto.rindex("}") + 1])


@pytest.fixture
def token(monkeypatch):
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    yield {"X-Internal-Token": TOKEN}
    get_settings.cache_clear()


# ---------------------------------------------------------------------------
# Sin abonos, payload real, contexto que viaja
# ---------------------------------------------------------------------------


def test_abonos_vacios_no_llaman_a_claude(claude_fake) -> None:
    cliente = claude_fake(conciliacion_abonos, {"sugerencias": []})
    res = sugerir_abonos(ConciliacionSugerirAbonosRequest())
    assert res.sugerencias == []
    assert res.uso_ia is None
    assert cliente.llamadas == []


def test_payload_real_del_api_no_da_422(claude_fake, token) -> None:
    """Folio como NÚMERO, ids con prefijo TIPO:uuid, montos numeric como
    texto y banderas en null: así puede serializar el API. Pydantic v2 no
    convierte int → str; con un `str` estricto la IA quedaba muerta con 422
    (ya pasó el 15-sep-2026 con `vuelo_folio`)."""
    claude_fake(
        conciliacion_abonos,
        {"sugerencias": [_sug(MOV_CRISTY, candidato_id=COBRO_235, confianza=0.9, accion="LIGAR")]},
    )
    cuerpo = {
        "abonos": [ABONO_CRISTY],
        "candidatos": [
            {**CAND_235, "monto": "20400.00", "es_anticipo": None, "otra_cuenta": None},
            {**CAND_OTRO, "folio": "240", "neto": None},
        ],
        "categorias": CATEGORIAS,
        "campo_que_no_existe": True,
    }
    res = client.post("/conciliacion/sugerir-abonos", json=cuerpo, headers=token)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["sugerencias"][0]["candidato_id"] == COBRO_235
    assert data["sugerencias"][0]["candidato_tipo"] == "COBRO_VUELO"
    assert data["uso_ia"]["input_tokens"] == 120

    req = ConciliacionSugerirAbonosRequest.model_validate(cuerpo)
    assert req.candidatos[0].folio == "235"
    assert req.candidatos[0].es_anticipo is False
    # neto null ⇒ monto − comisión (re-resta de dos campos del mismo payload).
    assert req.candidatos[1].neto == 19000.0


def test_ruta_exige_token_interno() -> None:
    res = client.post("/conciliacion/sugerir-abonos", json={"abonos": []})
    assert res.status_code in (401, 503)


def test_payload_para_el_modelo_y_system_con_dominio(claude_fake) -> None:
    cliente = claude_fake(conciliacion_abonos, {"sugerencias": []})
    opciones: dict = {}

    def con_opciones(**kw):
        opciones.update(kw)
        return cliente

    cliente.with_options = con_opciones
    sugerir_abonos(_req([ABONO_CRISTY], [CAND_235, CAND_OTRO]))

    llamada = cliente.ultima
    assert llamada["max_tokens"] == MAX_TOKENS == 6000
    assert opciones["timeout"] == TIMEOUT_S == 120.0
    # Sin reintentos del SDK: con los 2 del cliente compartido una generación
    # lenta duraba 3 × 120 s (el API aborta a los 130 s) y cada intento se
    # cobraba sin llegar a ia_uso.
    assert opciones["max_retries"] == 0

    payload = _payload_usuario(cliente)
    abono = payload["abonos"][0]
    assert abono["candidato_ids"] == [COBRO_235, COBRO_OTRO]
    # El sistema ya calculó qué cuadra al centavo: el modelo no lo recalcula.
    assert abono["exactos"] == [COBRO_235]
    assert abono["cuenta_tipo"] == "PASARELA"
    # Sin campos vacíos (el JSON viaja en cada llamada).
    assert "monto_bruto" not in abono and "referencia" not in abono
    cand = payload["candidatos"][0]
    assert cand["folio"] == "235" and cand["neto"] == 19380.0
    assert "es_anticipo" not in cand and "otra_cuenta" not in cand
    assert payload["categorias"] == sorted(CATEGORIAS)

    system = llamada["system"]
    assert len(system) == 2
    assert system[0]["text"] == BLOQUE_DOMINIO
    assert "CONCILIACIÓN DE INGRESOS" in system[1]["text"]
    assert all(b["cache_control"] == {"type": "ephemeral"} for b in system)


def test_dominio_trae_los_alias_nuevos_y_conserva_los_viejos() -> None:
    assert DOMINIO_VERSION == "2026-09-24"
    assert "POCKET DE LATINOAMERICA = BillPocket: depósito AGRUPADO" in BLOQUE_DOMINIO
    assert "REV = reverso: el banco devolvió un cargo anterior" in BLOQUE_DOMINIO
    assert "SPEI = transferencia interbancaria" in BLOQUE_DOMINIO
    assert "SEL TRASPASO ENTRE CUENTAS" in BLOQUE_DOMINIO


# ---------------------------------------------------------------------------
# Caso real #235 y topes deterministas
# ---------------------------------------------------------------------------


def test_caso_real_235_neto_exacto_y_nombre_permiten_confianza_alta(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY,
                    candidato_id=COBRO_235,
                    confianza=0.99,
                    accion="LIGAR",
                    razon="Monto neto exacto y nombre de la clienta",
                    evidencias=["el ordenante es la clienta"],
                )
            ]
        },
    )
    res = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235, CAND_OTRO]))
    s = res.sugerencias[0]
    assert s.accion == "LIGAR"
    assert s.candidato_id == COBRO_235
    assert s.confianza == 0.95  # 0.99 del modelo, topado por el hecho más fuerte
    assert any("NETO" in e and "1,020.00" in e for e in s.evidencias)
    assert any("CHAVEZ" in e and "CRISTY" in e for e in s.evidencias)
    assert any("Mismo día" in e for e in s.evidencias)
    assert s.evidencias[-1] == "el ordenante es la clienta"
    assert s.motivo_sin_match is None


def test_nombre_del_ordenante_empata_con_dos_tokens() -> None:
    assert nombre_en_descripcion(
        "MARIA CRISTINA CHAVEZ BADIOLA : vuelo cristy badiola", "Cristy Chavez"
    ) == {"CRISTY", "CHAVEZ"}
    assert nombre_en_descripcion("LETICIA LEON ALVARADO : PAGO", "Leticia León Alvarado")
    assert not nombre_en_descripcion("LETICIA LEON ALVARADO : PAGO", "Leticia Pérez")
    # «PAGO»/«VUELO» no identifican a nadie.
    assert not nombre_en_descripcion("PAGO VUELO", "Pago Vuelo")


def test_pasarela_compara_bruto_del_abono_contra_bruto_del_candidato() -> None:
    abono = AbonoParaSugerir(id="m", monto=18480.0, monto_bruto=20400.0, cuenta_moneda="MXN")
    cand = CandidatoParaAbono(id="COBRO_VUELO:x", tipo="COBRO_VUELO", monto=20400.0, neto=20400.0)
    assert monto_exacto(abono, cand)
    sin_bruto = abono.model_copy(update={"monto_bruto": None})
    assert not monto_exacto(sin_bruto, cand)


def _una(claude_fake, abono: dict, cand: dict, confianza: float = 0.99):
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(abono["id"], candidato_id=cand["id"], confianza=confianza, accion="LIGAR")
            ]
        },
    )
    return sugerir_abonos(_req([abono], [cand])).sugerencias[0]


def test_monto_no_exacto_y_sin_nombre_tope_04(claude_fake) -> None:
    abono = {**ABONO_CRISTY, "descripcion": "DEPOSITO SPEI 0012345", "candidato_ids": [COBRO_OTRO]}
    cand = {**CAND_OTRO, "fecha": "2026-09-08", "monto": 19380.5, "neto": 19380.5}
    s = _una(claude_fake, abono, cand)
    assert s.candidato_id == COBRO_OTRO
    assert s.confianza <= 0.4
    assert any("Diferencia de 0.50" in e for e in s.evidencias)


def test_diferencia_menor_a_un_peso_con_nombre_tope_085(claude_fake) -> None:
    abono = {**ABONO_CRISTY, "candidato_ids": [COBRO_235]}
    cand = {**CAND_235, "neto": 19380.6}
    s = _una(claude_fake, abono, cand)
    assert s.confianza == 0.85


def test_diferencia_de_un_por_ciento_con_nombre_es_posible_comision(claude_fake) -> None:
    abono = {**ABONO_CRISTY, "candidato_ids": [COBRO_235]}
    cand = {**CAND_235, "comision": None, "neto": 19550.0, "monto": 19550.0}
    s = _una(claude_fake, abono, cand)
    assert s.confianza == 0.7
    assert any("posible comisión" in e for e in s.evidencias)


def test_diferencia_grande_tope_03(claude_fake) -> None:
    abono = {**ABONO_CRISTY, "candidato_ids": [COBRO_235]}
    cand = {**CAND_235, "comision": None, "neto": 25000.0, "monto": 25000.0}
    s = _una(claude_fake, abono, cand)
    assert s.confianza <= 0.3


def test_mas_de_45_dias_tope_04_aunque_el_monto_cuadre(claude_fake) -> None:
    abono = {**ABONO_CRISTY, "candidato_ids": [COBRO_235]}
    cand = {**CAND_235, "fecha": "2026-07-01"}
    s = _una(claude_fake, abono, cand)
    assert s.confianza <= 0.4
    assert any("69 día(s) después" in e for e in s.evidencias)


@pytest.mark.parametrize(
    ("fecha", "tope"),
    [("2026-09-05", 0.95), ("2026-09-01", 0.8), ("2026-08-10", 0.6), ("2026-10-08", 0.6)],
)
def test_topes_por_dias(claude_fake, fecha: str, tope: float) -> None:
    abono = {**ABONO_CRISTY, "candidato_ids": [COBRO_235]}
    s = _una(claude_fake, abono, {**CAND_235, "fecha": fecha})
    assert s.confianza == tope


def test_ingreso_en_otra_cuenta_tope_05(claude_fake) -> None:
    abono = {
        **ABONO_CRISTY,
        "descripcion": "Rembolso Gastos Medicos Dani",
        "monto": 21223.10,
        "candidato_ids": [ING_7],
    }
    cand = {
        "id": ING_7,
        "tipo": "INGRESO",
        "fecha": "2026-09-08",
        "monto": 21223.10,
        "neto": 21223.10,
        "moneda": "MXN",
        "categoria": "REEMBOLSO_DEVOLUCION",
        "descripcion": "Reembolso gastos médicos",
        "cuenta_alias": "GASTOS GNRAL",
        "otra_cuenta": True,
        "folio": "ING-7",
    }
    s = _una(claude_fake, abono, cand)
    assert s.candidato_tipo == "INGRESO"
    assert s.confianza <= 0.5
    assert any("otra cuenta (GASTOS GNRAL)" in e for e in s.evidencias)


def test_evidencias_maximo_cuatro_y_textos_300(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY,
                    candidato_id=COBRO_235,
                    confianza=0.9,
                    accion="LIGAR",
                    razon="x" * 900,
                    evidencias=[f"hecho {i}" for i in range(10)],
                )
            ]
        },
    )
    s = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235])).sugerencias[0]
    assert len(s.evidencias) == 4
    assert len(s.razon) == 300


# ---------------------------------------------------------------------------
# ids: permitidos, inventados, moneda, repetidos entre abonos
# ---------------------------------------------------------------------------


ABONO_LETICIA = {
    "id": MOV_LETICIA,
    "fecha": "2026-09-10",
    "monto": 95000.0,
    "descripcion": "LETICIA LEON ALVARADO : PAGO",
    "cuenta_alias": "Paywise",
    "cuenta_moneda": "MXN",
    "cuenta_tipo": "PASARELA",
    "candidato_ids": [COBRO_OTRO],
}


def test_candidato_fuera_de_la_lista_de_su_abono_se_descarta(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                # COBRO_235 existe en el pool pero NO está permitido para Leticia.
                _sug(MOV_LETICIA, candidato_id=COBRO_235, confianza=0.9, accion="LIGAR"),
                _sug(MOV_CRISTY, candidato_id="COBRO_VUELO:inventado", confianza=0.9),
            ]
        },
    )
    res = sugerir_abonos(_req([ABONO_LETICIA, ABONO_CRISTY], [CAND_235, CAND_OTRO]))
    for s in res.sugerencias:
        assert s.candidato_id is None
        assert s.accion == "REVISAR"
        assert s.confianza == 0.0
        assert s.motivo_sin_match == MOTIVO_ID_DESCARTADO
    assert any("Se descartaron 2 candidato(s)" in a for a in res.advertencias)


def test_candidato_de_otra_moneda_ni_se_ofrece_ni_se_acepta(claude_fake) -> None:
    usd = {**CAND_235, "id": COBRO_USD, "moneda": "USD", "monto": 1020.0, "neto": 1020.0}
    abono = {**ABONO_CRISTY, "candidato_ids": [COBRO_235, COBRO_USD]}
    cliente = claude_fake(
        conciliacion_abonos,
        {"sugerencias": [_sug(MOV_CRISTY, candidato_id=COBRO_USD, confianza=0.9, accion="LIGAR")]},
    )
    s = sugerir_abonos(_req([abono], [CAND_235, usd])).sugerencias[0]
    payload = _payload_usuario(cliente)
    assert payload["abonos"][0]["candidato_ids"] == [COBRO_235]
    assert [c["id"] for c in payload["candidatos"]] == [COBRO_235]
    assert s.candidato_id is None and s.accion == "REVISAR"


def test_mismo_candidato_en_dos_abonos_se_queda_el_de_mas_confianza(claude_fake) -> None:
    otro_abono = {
        **ABONO_CRISTY,
        "id": MOV_DANI,
        "descripcion": "DEPOSITO 19380",
        "candidato_ids": [COBRO_235],
    }
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(MOV_DANI, candidato_id=COBRO_235, confianza=0.8, accion="LIGAR"),
                _sug(MOV_CRISTY, candidato_id=COBRO_235, confianza=0.95, accion="LIGAR"),
            ]
        },
    )
    res = sugerir_abonos(_req([otro_abono, ABONO_CRISTY], [CAND_235]))
    perdedor, ganador = res.sugerencias
    assert ganador.candidato_id == COBRO_235 and ganador.accion == "LIGAR"
    assert perdedor.candidato_id is None
    assert perdedor.accion == "REVISAR"
    assert perdedor.confianza == 0.0
    assert perdedor.motivo_sin_match == MOTIVO_CANDIDATO_REPETIDO
    assert any("dos abonos" in a for a in res.advertencias)


def test_alternativas_solo_permitidas_sin_repetir_y_topadas(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY,
                    candidato_id=COBRO_235,
                    confianza=0.9,
                    accion="LIGAR",
                    alternativas=[
                        {"candidato_id": COBRO_235, "confianza": 0.9, "razon": "repetido"},
                        {"candidato_id": "COBRO_VUELO:inventado", "confianza": 0.9},
                        {"candidato_id": COBRO_OTRO, "confianza": 0.9, "razon": "otro"},
                        {"candidato_id": COBRO_OTRO, "confianza": 0.5, "razon": "dup"},
                    ],
                )
            ]
        },
    )
    s = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235, CAND_OTRO])).sugerencias[0]
    assert [a.candidato_id for a in s.alternativas] == [COBRO_OTRO]
    # 19,000 vs 19,380 sin nombre (Hotel Xcaret) ⇒ tope 0.4.
    assert s.alternativas[0].confianza <= 0.4


# ---------------------------------------------------------------------------
# Acciones
# ---------------------------------------------------------------------------


def test_ligar_sin_candidato_pasa_a_revisar(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {"sugerencias": [_sug(MOV_CRISTY, candidato_id=None, confianza=0.9, accion="LIGAR")]},
    )
    s = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235])).sugerencias[0]
    assert s.accion == "REVISAR"
    assert s.candidato_id is None
    assert s.confianza == 0.0
    assert s.motivo_sin_match


def test_registrar_ingreso_con_cobro_exacto_permitido_pasa_a_revisar(claude_fake) -> None:
    """Caso real #235: si el operador lo «registrara como otro ingreso» el
    dinero contaría DOS veces (venta del vuelo + otro ingreso)."""
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY,
                    accion="REGISTRAR_INGRESO",
                    categoria_sugerida="OTRO_INGRESO",
                    confianza=0.9,
                    razon="Parece otro ingreso",
                )
            ]
        },
    )
    res = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235, CAND_OTRO]))
    s = res.sugerencias[0]
    assert s.accion == "REVISAR"
    assert s.candidato_id is None
    assert s.motivo_sin_match == MOTIVO_COBRO_EXACTO
    assert s.razon == MOTIVO_COBRO_EXACTO
    # El cobro exacto queda a la mano como PRIMERA alternativa (≤ 0.7).
    assert s.alternativas[0].candidato_id == COBRO_235
    assert s.alternativas[0].confianza == 0.7
    assert any("cuadran exacto con un cobro" in a for a in res.advertencias)


def test_registrar_ingreso_con_sobre_de_grupo_exacto_tambien_se_frena(claude_fake) -> None:
    sobre = {**CAND_235, "id": "SOBRE_GRUPO:abcd", "tipo": "SOBRE_GRUPO", "folio": "G-4"}
    abono = {**ABONO_CRISTY, "candidato_ids": ["SOBRE_GRUPO:abcd"]}
    claude_fake(
        conciliacion_abonos,
        {"sugerencias": [_sug(MOV_CRISTY, accion="REGISTRAR_INGRESO", confianza=0.9)]},
    )
    s = sugerir_abonos(_req([abono], [sobre])).sugerencias[0]
    assert s.accion == "REVISAR"
    assert s.motivo_sin_match == MOTIVO_COBRO_EXACTO
    assert s.alternativas[0].candidato_id == "SOBRE_GRUPO:abcd"


def test_registrar_ingreso_con_ingreso_exacto_ya_registrado_se_frena(claude_fake) -> None:
    """Registrar OTRO ingreso por el mismo dinero inflaría «otros ingresos»."""
    ing = {
        "id": ING_7,
        "tipo": "INGRESO",
        "fecha": "2026-09-08",
        "monto": 19380.0,
        "neto": 19380.0,
        "moneda": "MXN",
        "categoria": "REEMBOLSO_DEVOLUCION",
        "folio": "ING-7",
    }
    abono = {**ABONO_CRISTY, "candidato_ids": [ING_7]}
    claude_fake(
        conciliacion_abonos,
        {"sugerencias": [_sug(MOV_CRISTY, accion="REGISTRAR_INGRESO", confianza=0.9)]},
    )
    res = sugerir_abonos(_req([abono], [ing]))
    s = res.sugerencias[0]
    assert s.accion == "REVISAR"
    assert s.motivo_sin_match == MOTIVO_INGRESO_EXACTO
    assert s.alternativas[0].candidato_id == ING_7
    assert "revisa si es este ingreso" in s.alternativas[0].razon
    assert any("ingreso ya registrado" in a for a in res.advertencias)


def test_monto_chico_un_peso_ya_no_es_redondeo(claude_fake) -> None:
    abono = {**ABONO_CRISTY, "monto": 50.0, "candidato_ids": [COBRO_235]}
    cand = {**CAND_235, "comision": None, "monto": 49.2, "neto": 49.2}
    s = _una(claude_fake, abono, cand)
    # 0.80 de diferencia es 1.6 %: con nombre, tope 0.5 (no 0.85).
    assert s.confianza == 0.5


def test_evidencias_del_modelo_que_no_son_lista_se_ignoran(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY,
                    candidato_id=COBRO_235,
                    confianza=0.9,
                    accion="LIGAR",
                    evidencias="monto exacto",
                )
            ]
        },
    )
    s = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235])).sugerencias[0]
    assert "m" not in s.evidencias
    assert all(len(e) > 1 for e in s.evidencias)


def test_registrar_ingreso_sin_cobro_exacto_se_respeta(claude_fake) -> None:
    abono = {
        **ABONO_CRISTY,
        "id": MOV_DANI,
        "monto": 21223.10,
        "descripcion": "Rembolso Gastos Medicos Dani",
        "candidato_ids": [COBRO_235],
    }
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_DANI,
                    accion="REGISTRAR_INGRESO",
                    categoria_sugerida="REEMBOLSO_DEVOLUCION",
                    confianza=0.8,
                    razon="Reembolso de gastos médicos",
                    motivo_sin_match="Ningún cobro de vuelo cuadra",
                )
            ]
        },
    )
    s = sugerir_abonos(_req([abono], [CAND_235])).sugerencias[0]
    assert s.accion == "REGISTRAR_INGRESO"
    assert s.categoria_sugerida == "REEMBOLSO_DEVOLUCION"
    assert s.confianza == 0.8
    assert s.motivo_sin_match == "Ningún cobro de vuelo cuadra"


def test_registrar_como_anticipo_nunca_va_directo(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_LETICIA,
                    accion="REGISTRAR_INGRESO",
                    categoria_sugerida="ANTICIPO_CLIENTE",
                    confianza=0.6,
                )
            ]
        },
    )
    s = sugerir_abonos(_req([ABONO_LETICIA], [CAND_OTRO])).sugerencias[0]
    assert s.accion == "REVISAR"
    assert s.categoria_sugerida == "ANTICIPO_CLIENTE"
    assert "anticipo" in (s.motivo_sin_match or "")


def test_accion_contradictoria_con_candidato_queda_para_revisar(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY,
                    candidato_id=COBRO_235,
                    confianza=0.9,
                    accion="REGISTRAR_INGRESO",
                )
            ]
        },
    )
    s = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235])).sugerencias[0]
    assert s.accion == "REVISAR"
    assert s.candidato_id == COBRO_235


def test_categoria_fuera_de_lista_es_none_y_se_normaliza(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(MOV_CRISTY, accion="REGISTRAR_INGRESO", categoria_sugerida="PROPINA"),
                _sug(MOV_LETICIA, accion="revisar", categoria_sugerida="anticipo_cliente"),
            ]
        },
    )
    a, b = sugerir_abonos(_req([ABONO_CRISTY, ABONO_LETICIA], [CAND_235, CAND_OTRO])).sugerencias
    assert a.categoria_sugerida is None
    assert b.categoria_sugerida == "ANTICIPO_CLIENTE"
    assert b.accion == "REVISAR"


def test_accion_desconocida_es_revisar(claude_fake) -> None:
    claude_fake(conciliacion_abonos, {"sugerencias": [_sug(MOV_LETICIA, accion="BORRAR")]})
    s = sugerir_abonos(_req([ABONO_LETICIA], [CAND_OTRO])).sugerencias[0]
    assert s.accion == "REVISAR"


def test_abono_sin_respuesta_y_respuestas_ajenas(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {"sugerencias": [_sug("otro-movimiento", accion="LIGAR"), "basura"]},
    )
    res = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235]))
    s = res.sugerencias[0]
    assert s.movimiento_id == MOV_CRISTY
    assert s.accion == "REVISAR" and s.candidato_id is None
    assert any("no devolvió propuesta" in a for a in res.advertencias)
    assert any("no venían en el lote" in a for a in res.advertencias)


def test_una_sugerencia_por_abono_en_el_orden_del_request(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(MOV_LETICIA),
                _sug(MOV_CRISTY, candidato_id=COBRO_235, confianza=0.9, accion="LIGAR"),
                _sug(MOV_CRISTY, candidato_id=None, accion="REVISAR"),  # repetida: gana la 1.ª
            ]
        },
    )
    res = sugerir_abonos(_req([ABONO_CRISTY, ABONO_LETICIA], [CAND_235, CAND_OTRO]))
    assert [s.movimiento_id for s in res.sugerencias] == [MOV_CRISTY, MOV_LETICIA]
    assert res.sugerencias[0].candidato_id == COBRO_235
    assert res.uso_ia is not None and res.uso_ia.output_tokens == 45


def test_respuesta_como_lista_suelta_se_tolera(claude_fake) -> None:
    claude_fake(conciliacion_abonos, json.dumps([_sug(MOV_CRISTY, accion="REVISAR")]))
    res = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235]))
    assert res.sugerencias[0].accion == "REVISAR"


# ---------------------------------------------------------------------------
# Revisión adversaria (24-sep-2026): parseo robusto e ids del modelo
# ---------------------------------------------------------------------------

_LIGAR_235 = _sug(MOV_CRISTY, candidato_id=COBRO_235, confianza=0.9, accion="LIGAR")


@pytest.mark.parametrize(
    "texto",
    [
        # Nota DESPUÉS del JSON con una llave: antes se cortaba en el ÚLTIMO «}».
        json.dumps({"sugerencias": [_LIGAR_235]}) + "\n\nNota: revisé {1} abono.",
        # Corchete ANTES del objeto: antes se tomaba como «lista suelta».
        "Resultado [1 abono]:\n" + json.dumps({"sugerencias": [_LIGAR_235]}),
        "```json\n" + json.dumps({"sugerencias": [_LIGAR_235]}) + "\n```",
    ],
)
def test_prosa_alrededor_del_json_no_tira_una_respuesta_buena(claude_fake, texto) -> None:
    claude_fake(conciliacion_abonos, texto)
    s = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235])).sugerencias[0]
    assert s.candidato_id == COBRO_235 and s.accion == "LIGAR"


def test_json_roto_no_se_rescata_por_pedazos(claude_fake, token) -> None:
    """Un objeto mal formado es 422 aunque ADENTRO haya listas de objetos
    bien formados (las alternativas): jamás se «repara» una respuesta."""
    # Falta la coma antes de "accion": el objeto de afuera NO es JSON válido.
    roto = (
        f'{{"sugerencias":[{{"movimiento_id":"{MOV_CRISTY}","alternativas":'
        f'[{{"candidato_id":"{COBRO_235}","confianza":0.9}}] "accion":"LIGAR"}}]}}'
    )
    claude_fake(conciliacion_abonos, roto)
    res = client.post("/conciliacion/sugerir-abonos", json=_cuerpo(), headers=token)
    assert res.status_code == 422
    assert res.json()["detail"]["uso_ia"]["input_tokens"] == 120


def test_confianza_nan_no_revienta_y_conserva_el_uso(claude_fake, token) -> None:
    """`json.loads` acepta el literal NaN y `min(max(nan, 0), 1)` es NaN: antes
    tronaba el `le=1` del esquema ⇒ 422 genérico SIN `uso_ia`."""
    claude_fake(
        conciliacion_abonos,
        f'{{"sugerencias":[{{"movimiento_id":"{MOV_CRISTY}","candidato_id":"{COBRO_235}",'
        f'"confianza":NaN,"accion":"LIGAR","alternativas":[{{"candidato_id":"{COBRO_OTRO}",'
        '"confianza":Infinity}]}]}',
    )
    res = client.post(
        "/conciliacion/sugerir-abonos",
        json={**_cuerpo(), "abonos": [ABONO_CRISTY], "candidatos": [CAND_235, CAND_OTRO]},
        headers=token,
    )
    assert res.status_code == 200, res.text
    s = res.json()["sugerencias"][0]
    assert s["candidato_id"] == COBRO_235
    assert s["confianza"] == 0.0
    assert s["alternativas"][0]["confianza"] == 0.0


def test_error_inesperado_al_post_validar_es_422_con_uso(claude_fake, token, monkeypatch) -> None:
    """Claude ya contestó (créditos gastados): cualquier fallo después sale
    como IA_RESPUESTA_ILEGIBLE CON `uso_ia` para que el API lo registre."""
    claude_fake(conciliacion_abonos, {"sugerencias": [_LIGAR_235]})

    def _truena(*_a, **_kw):
        raise TypeError("tipo inesperado")

    monkeypatch.setattr(conciliacion_abonos, "_procesar", _truena)
    res = client.post("/conciliacion/sugerir-abonos", json=_cuerpo(), headers=token)
    assert res.status_code == 422
    detalle = res.json()["detail"]
    assert detalle["error"] == "IA_RESPUESTA_ILEGIBLE"
    assert detalle["uso_ia"]["output_tokens"] == 45


def test_ventana_de_contexto_agotada_tambien_es_truncada(claude_fake, token) -> None:
    claude_fake(conciliacion_abonos, '{"sugerencias":[', "model_context_window_exceeded")
    res = client.post("/conciliacion/sugerir-abonos", json=_cuerpo(), headers=token)
    assert res.status_code == 502
    assert res.json()["detail"]["error"] == "IA_RESPUESTA_TRUNCADA"


def test_ids_reescritos_por_el_modelo_se_resuelven_sin_adivinar(claude_fake) -> None:
    """uuid sin el prefijo TIPO: o en MAYÚSCULAS es el MISMO candidato de la
    lista (se devuelve el id canónico, que es lo que el API valida); lo que
    no está en la lista sigue descartado."""
    uuid_235 = COBRO_235.split(":", 1)[1]
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY.upper(),
                    candidato_id=uuid_235,
                    confianza=0.9,
                    accion="LIGAR",
                    alternativas=[{"candidato_id": COBRO_OTRO.upper(), "confianza": 0.3}],
                ),
                _sug(MOV_LETICIA, candidato_id="5e6f7a8b", confianza=0.9, accion="LIGAR"),
            ]
        },
    )
    abono_leticia = {**ABONO_LETICIA, "candidato_ids": [COBRO_OTRO]}
    res = sugerir_abonos(_req([ABONO_CRISTY, abono_leticia], [CAND_235, CAND_OTRO]))
    cristy, leticia = res.sugerencias
    assert cristy.movimiento_id == MOV_CRISTY
    assert cristy.candidato_id == COBRO_235 and cristy.accion == "LIGAR"
    assert [a.candidato_id for a in cristy.alternativas] == [COBRO_OTRO]
    # Un pedazo de uuid NO es un id: descartado.
    assert leticia.candidato_id is None
    assert leticia.motivo_sin_match == MOTIVO_ID_DESCARTADO


def test_cobro_exacto_siempre_queda_a_la_vista(claude_fake) -> None:
    """Si el modelo no elige el cobro que cuadra al centavo (lo manda a
    «anticipo», o ni contesta por ese abono), el cobro va PRIMERO en las
    alternativas: nadie debe registrar como otro ingreso dinero que ya tiene
    su cobro esperando (caso real #235)."""
    otro_abono = {**ABONO_CRISTY, "id": MOV_DANI, "candidato_ids": [COBRO_235]}
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY,
                    accion="REVISAR",
                    categoria_sugerida="ANTICIPO_CLIENTE",
                    confianza=0.5,
                    alternativas=[{"candidato_id": COBRO_OTRO, "confianza": 0.2}],
                ),
            ]
        },
    )
    res = sugerir_abonos(_req([ABONO_CRISTY, otro_abono], [CAND_235, CAND_OTRO]))
    anticipo, sin_respuesta = res.sugerencias
    assert anticipo.accion == "REVISAR" and anticipo.candidato_id is None
    assert [a.candidato_id for a in anticipo.alternativas] == [COBRO_235, COBRO_OTRO]
    assert anticipo.alternativas[0].confianza == 0.7
    assert sin_respuesta.accion == "REVISAR"
    assert [a.candidato_id for a in sin_respuesta.alternativas] == [COBRO_235]


def test_con_candidato_no_exacto_el_exacto_va_de_alternativa(claude_fake) -> None:
    claude_fake(
        conciliacion_abonos,
        {
            "sugerencias": [
                _sug(
                    MOV_CRISTY,
                    candidato_id=COBRO_OTRO,
                    confianza=0.8,
                    accion="LIGAR",
                    alternativas=[{"candidato_id": COBRO_235, "confianza": 0.4, "razon": "mío"}],
                )
            ]
        },
    )
    s = sugerir_abonos(_req([ABONO_CRISTY], [CAND_235, CAND_OTRO])).sugerencias[0]
    assert s.candidato_id == COBRO_OTRO
    assert s.confianza <= 0.4  # 19,000 vs 19,380 sin nombre
    # El exacto primero, con la confianza (más baja) y la razón del modelo.
    assert s.alternativas[0].candidato_id == COBRO_235
    assert s.alternativas[0].confianza == 0.4
    assert s.alternativas[0].razon == "mío"


def test_la_evidencia_que_explica_el_tope_no_se_recorta() -> None:
    """Con 5 hechos duros (monto, nombre, días, otra cuenta, referencia) el
    recorte a 4 no debe tirar «otra cuenta», que es la que baja el tope."""
    abono = AbonoParaSugerir(
        id="m",
        fecha="2026-09-08",
        monto=21223.10,
        descripcion="DANIEL PEREZ GOMEZ REF 7788990011",
        referencia="7788990011",
        cuenta_moneda="MXN",
    )
    cand = CandidatoParaAbono(
        id=ING_7,
        tipo="INGRESO",
        fecha="2026-09-01",
        monto=21223.10,
        neto=21223.10,
        cliente="Daniel Pérez Gómez",
        referencia="7788990011",
        otra_cuenta=True,
        cuenta_alias="GASTOS GNRAL",
    )
    ev, tope = conciliacion_abonos._evidencias_abono(abono, cand)
    assert tope == 0.5
    assert len(ev) == 5
    assert any("otra cuenta" in e for e in ev[:4])


# ---------------------------------------------------------------------------
# Errores (HTTP)
# ---------------------------------------------------------------------------


def _cuerpo() -> dict:
    return {"abonos": [ABONO_CRISTY], "candidatos": [CAND_235], "categorias": CATEGORIAS}


def test_respuesta_truncada_es_502_con_uso_ia(claude_fake, token) -> None:
    claude_fake(conciliacion_abonos, '{"sugerencias":[{"movimiento_id":"', "max_tokens")
    res = client.post("/conciliacion/sugerir-abonos", json=_cuerpo(), headers=token)
    assert res.status_code == 502
    detalle = res.json()["detail"]
    assert detalle["error"] == "IA_RESPUESTA_TRUNCADA"
    assert detalle["code"] == "IA_RESPUESTA_TRUNCADA"
    # Los créditos se gastaron: el API los registra en ia_uso.
    assert detalle["uso_ia"]["input_tokens"] == 120
    assert detalle["uso_ia"]["output_tokens"] == 45


def test_json_ilegible_es_422_con_uso_ia(claude_fake, token) -> None:
    claude_fake(conciliacion_abonos, "no sé qué responder")
    res = client.post("/conciliacion/sugerir-abonos", json=_cuerpo(), headers=token)
    assert res.status_code == 422
    detalle = res.json()["detail"]
    assert detalle["error"] == "IA_RESPUESTA_ILEGIBLE"
    assert detalle["uso_ia"]["input_tokens"] == 120


def test_json_sin_lista_de_sugerencias_es_422(claude_fake, token) -> None:
    claude_fake(conciliacion_abonos, {"propuestas": []})
    res = client.post("/conciliacion/sugerir-abonos", json=_cuerpo(), headers=token)
    assert res.status_code == 422


def test_claude_caido_es_502(monkeypatch, token) -> None:
    peticion = httpx.Request("POST", "https://api.anthropic.com/v1/messages")

    class _Caido:
        def with_options(self, **_kw):
            return self

        @property
        def messages(self):
            return self

        def create(self, **_kw):
            raise anthropic.APIStatusError(
                "overloaded", response=httpx.Response(529, request=peticion), body=None
            )

    monkeypatch.setattr(conciliacion_abonos, "_client", lambda: _Caido())
    res = client.post("/conciliacion/sugerir-abonos", json=_cuerpo(), headers=token)
    assert res.status_code == 502


def test_timeout_de_claude_es_502_no_500(monkeypatch, token) -> None:
    peticion = httpx.Request("POST", "https://api.anthropic.com/v1/messages")

    class _Lento:
        def with_options(self, **_kw):
            return self

        @property
        def messages(self):
            return self

        def create(self, **_kw):
            raise anthropic.APITimeoutError(request=peticion)

    monkeypatch.setattr(conciliacion_abonos, "_client", lambda: _Lento())
    res = client.post("/conciliacion/sugerir-abonos", json=_cuerpo(), headers=token)
    assert res.status_code == 502


def test_once_abonos_es_422(claude_fake, token) -> None:
    cliente = claude_fake(conciliacion_abonos, {"sugerencias": []})
    abonos = [{**ABONO_CRISTY, "id": f"mov-{i}"} for i in range(11)]
    res = client.post(
        "/conciliacion/sugerir-abonos",
        json={"abonos": abonos, "candidatos": [], "categorias": CATEGORIAS},
        headers=token,
    )
    assert res.status_code == 422
    assert cliente.llamadas == []
    with pytest.raises(ValidationError):
        ConciliacionSugerirAbonosRequest.model_validate({"abonos": abonos})


def test_121_candidatos_es_422() -> None:
    cands = [{**CAND_OTRO, "id": f"COBRO_VUELO:{i}"} for i in range(121)]
    with pytest.raises(ValidationError):
        ConciliacionSugerirAbonosRequest.model_validate({"abonos": [], "candidatos": cands})


# ---------------------------------------------------------------------------
# Textos de las hojas (§8.3 del contrato): las notas dicen dónde van los ING-n
# ---------------------------------------------------------------------------

FILA_ING = {
    "clave": "ING-12",
    "fecha_vuelo": None,
    "concepto_egreso": "comisión bancaria",
    "egreso_mxn": 50.0,
    "fecha_egreso": "2026-09-10",
    "concepto_ingreso": "Reembolsos y devoluciones recibidos · Gastos médicos Dani",
    "ingreso_mxn": 21223.10,
    "fecha_ingreso": "2026-09-10",
    "remanente_mxn": 21173.10,
    "factura": None,
}


def test_libro_dinero_explica_los_otros_ingresos_sin_vuelo() -> None:
    import io

    req = DineroXlsxRequest(
        otros_ingresos=[DineroOtroIngresoFila(**FILA_ING)], utilidades_otros_ingresos_mxn=1.0
    )
    wb = load_workbook(io.BytesIO(render_dinero_xlsx(req)))
    hoja = wb["Otros ingresos"]
    assert "clave ING-n" in hoja["A2"].value
    assert "anticipos de clientes no están aquí" in hoja["A2"].value
    assert hoja["A4"].value == "ING-12"
    assert hoja["G4"].value == 21223.10
    assert "ingresos sin vuelo registrados en Ingresos (ING-n)" in wb["utilidades"]["A9"].value


def test_balance_general_explica_los_ing_en_sueltas() -> None:
    wb = Workbook()
    ws = wb.active
    fila = BalanceOtroMovimientoFila(**{k: v for k, v in FILA_ING.items()}, avion_color=None)
    _hoja_otros_movimientos(ws, BalanceHojaOtrosMovimientos(filas_sueltas=[fila]))
    assert "OTROS INGRESOS registrados en Ingresos (clave ING-n" in ws["A2"].value
    assert "los anticipos y las aportaciones NO aparecen" in ws["A2"].value
    assert "son solo 'gas sin avión'" not in ws["A2"].value
    claves = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert "ING-12" in claves
    assert "clave ING-n" in _NOTA_GASTOS_EMPRESA
