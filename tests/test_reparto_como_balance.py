"""Reparto a socios (PDF y Excel) con LA CASCADA DEL BALANCE (20-oct-2026, API
0.0.101, invariante 99 del API; bandera raíz `reparto_como_balance`).

Decisión del cliente («que cuadren»): el reparto usa la misma cascada del
balance. La PROVISIÓN AFAC (tarifa de la ficha del avión × horas cobradas)
RESTA del saldo (`afac_reservado_*`; el `saldo_*` del API ya la trae); la
RESERVA DE OVERHAUL ya no resta: se informa por avión (después del saldo, en
gris) y por socio («Reserva a guardar», `reparto[].reserva_*`), y al socio con
`descuenta_reserva` VuelaTour se la retiene (`monto_* = utilidad_* −
reserva_*`, lo calcula el API; `reserva_retenida_*` por avión). Aquí no se
calcula nada: se pinta lo que manda el API. Sin la bandera (API ≤ 0.0.100, o
los campos presentes sin ella), el PDF y el Excel son BYTE A BYTE los de
siempre (huellas del HEAD 3b893f1 de `tests/test_reparto_en_pesos.py`).
"""

from __future__ import annotations

import copy
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth

from app.config import get_settings
from app.main import app
from app.schemas.reparto import RepartoAvion, RepartoPdfRequest, RepartoSocioLinea
from app.services import reparto_pdf as rp
from app.services import reparto_xlsx as rx
from app.services.reparto_pdf import render_reparto_pdf
from app.services.reparto_xlsx import render_reparto_xlsx
from tests import _pesos_payloads as pp
from tests.test_reparto_en_pesos import (
    _HUELLA_PDF_HEAD,
    _HUELLA_XLSX_HEAD,
    _huella_pdf,
    _huella_xlsx,
    _texto_pdf,
)

K = pp.K_REPARTO
r2 = pp.r2


def _d(x: float) -> str:
    return f"${x:,.2f}"


# ---------------------------------------------------------------------------
# Payloads: el reparto en pesos de `_pesos_payloads` con la cascada del
# balance (espejo de la aritmética del API 0.0.101: el saldo resta la
# provisión AFAC y NO la reserva; por socio utilidad / reserva / monto).
# ---------------------------------------------------------------------------


def _cascada_balance(
    a: dict,
    *,
    afac_usd: float,
    tarifa: float | None,
    horas: float,
    reserva_usd: float,
    socios: list[tuple[str, float, bool]],
) -> None:
    a["reserva_overhaul_usd"] = reserva_usd
    a["reserva_overhaul_mxn"] = r2(reserva_usd * K)
    a["afac_reservado_usd"] = afac_usd
    a["afac_reservado_mxn"] = r2(afac_usd * K)
    a["afac_tarifa_usd_hr"] = tarifa
    a["afac_horas_cobradas_hr"] = horas
    a["saldo_usd"] = r2(
        a["ingresos_cobrado_usd"]
        - a["comisiones_venta_usd"]
        - a["gastos_directos_usd"]
        - a["gastos_indirectos_usd"]
        - a["permisos_usd"]
        - a["otros_usd"]
        - afac_usd
        + a.get("otros_ingresos_avion_usd", 0.0)
    )
    a["saldo_mxn"] = r2(
        a["ingresos_cobrado_mxn"]
        - a["comisiones_venta_mxn"]
        - a["gastos_directos_mxn"]
        - a["gastos_indirectos_mxn"]
        - a["permisos_mxn"]
        - a["otros_mxn"]
        - a["afac_reservado_mxn"]
        + a["otros_ingresos_avion_mxn"]
    )
    a["reparto"] = []
    retenida_usd = retenida_mxn = 0.0
    for nombre, pct, descuenta in socios:
        u_usd, u_mxn = r2(a["saldo_usd"] * pct / 100), r2(a["saldo_mxn"] * pct / 100)
        rs_usd = r2(reserva_usd * pct / 100)
        rs_mxn = r2(a["reserva_overhaul_mxn"] * pct / 100)
        a["reparto"].append(
            {
                "socio_nombre": nombre,
                "porcentaje": pct,
                "monto_usd": r2(u_usd - rs_usd) if descuenta else u_usd,
                "monto_mxn": r2(u_mxn - rs_mxn) if descuenta else u_mxn,
                "utilidad_usd": u_usd,
                "utilidad_mxn": u_mxn,
                "reserva_usd": rs_usd,
                "reserva_mxn": rs_mxn,
                "descuenta_reserva": descuenta,
            }
        )
        if descuenta:
            retenida_usd, retenida_mxn = r2(retenida_usd + rs_usd), r2(retenida_mxn + rs_mxn)
    a["reserva_retenida_usd"] = retenida_usd
    a["reserva_retenida_mxn"] = retenida_mxn


def _payload(*, pesos: bool | None = True, bandera: object = True) -> dict:
    """XB-TST: tarifa AFAC 100 USD/hr × 3.5 h cobradas = 350 USD de provisión;
    reserva 140 USD; Socio A 60 % (la descuenta) y Socio B 40 % (no). XB-DOS:
    sin tarifa AFAC, sin reserva, Socio A 100 %. `pesos` = `reparto_en_pesos`
    (None = sin la llave); `bandera` = `reparto_como_balance`."""
    p = pp.reparto(bandera=pesos)
    tst, dos = p["aviones"]
    _cascada_balance(
        tst,
        afac_usd=350.0,
        tarifa=100.0,
        horas=3.5,
        reserva_usd=140.0,
        socios=[("Socio A", 60, True), ("Socio B", 40, False)],
    )
    _cascada_balance(
        dos, afac_usd=0.0, tarifa=None, horas=0.0, reserva_usd=0.0, socios=[("Socio A", 100, False)]
    )
    if bandera is not None:
        p["reparto_como_balance"] = bandera
    return p


def _neutro(p: dict) -> dict:
    """Los campos del API 0.0.101 presentes con valores que no cambian ningún
    número impreso (provisión 0, reserva por socio 0, nadie descuenta)."""
    p = copy.deepcopy(p)
    for a in p["aviones"]:
        a.update(
            afac_reservado_usd=0,
            afac_reservado_mxn=0,
            afac_horas_cobradas_hr=a["horas_voladas_hr"],
            afac_tarifa_usd_hr=None,
            reserva_retenida_usd=0,
            reserva_retenida_mxn=0,
        )
        for s in a["reparto"]:
            s.update(
                utilidad_usd=s["monto_usd"],
                utilidad_mxn=s.get("monto_mxn"),
                reserva_usd=0,
                reserva_mxn=0,
                descuenta_reserva=False,
            )
    return p


def _req(**kw) -> RepartoPdfRequest:
    return RepartoPdfRequest.model_validate(_payload(**kw))


def _flujos(data: bytes) -> list[bytes]:
    return [p.get_contents().get_data() for p in PdfReader(BytesIO(data)).pages]


def _miembros(data: bytes) -> dict[str, bytes]:
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _hoja(p: dict):
    return load_workbook(BytesIO(render_reparto_xlsx(RepartoPdfRequest.model_validate(p)))).active


def _fila(ws, texto: str, col: int = 1) -> int:
    filas = [c.row for c in ws[chr(64 + col)] if c.value == texto]
    assert len(filas) == 1, (texto, filas)
    return filas[0]


def _textos_col_a(ws) -> list[str]:
    return [c.value for c in ws["A"] if isinstance(c.value, str)]


# Números del payload (XB-TST): saldo = 2000 − 95.70 − 500 − 350.
SALDO_TST = 1054.30
SALDO_DOS = 1500.0
UTIL_A, RESERVA_A, MONTO_A = 632.58, 84.0, 548.58
UTIL_B, RESERVA_B, MONTO_B = 421.72, 56.0, 421.72


def test_los_numeros_del_payload_cuadran() -> None:
    req = _req()
    tst, dos = req.aviones
    assert (tst.saldo_usd, dos.saldo_usd) == (SALDO_TST, SALDO_DOS)
    a, b = tst.reparto
    assert (a.utilidad_usd, a.reserva_usd, a.monto_usd) == (UTIL_A, RESERVA_A, MONTO_A)
    assert (b.utilidad_usd, b.reserva_usd, b.monto_usd) == (UTIL_B, RESERVA_B, MONTO_B)
    assert a.monto_usd == r2(a.utilidad_usd - a.reserva_usd)
    assert tst.reserva_retenida_usd == RESERVA_A
    assert tst.reserva_retenida_mxn == r2(tst.reparto[0].reserva_mxn)
    assert a.monto_mxn == r2(a.utilidad_mxn - a.reserva_mxn)


# ---------------------------------------------------------------------------
# (a) Esquemas: aditivos, con default, liberales.
# ---------------------------------------------------------------------------


def test_esquemas_aditivos_con_default() -> None:
    a = RepartoAvion(matricula="X", modelo="M", ingresos_cobrado_usd=0, saldo_usd=0)
    for llave in (
        "afac_reservado_usd",
        "afac_reservado_mxn",
        "afac_horas_cobradas_hr",
        "afac_tarifa_usd_hr",
        "reserva_retenida_usd",
        "reserva_retenida_mxn",
    ):
        assert getattr(a, llave) is None
    s = RepartoSocioLinea(socio_nombre="S", porcentaje=1, monto_usd=1)
    for llave in ("utilidad_usd", "utilidad_mxn", "reserva_usd", "reserva_mxn"):
        assert getattr(s, llave) is None
    assert s.descuenta_reserva is None
    base = {"periodo_desde": "a", "periodo_hasta": "b", "generado": "c"}
    assert RepartoPdfRequest.model_validate(base).reparto_como_balance is False
    # La forma REAL del API 0.0.101 (golden `payload_oficial`): tarifa null.
    linea = RepartoSocioLinea.model_validate(
        {"socio_nombre": "Hernán Pérez", "porcentaje": 40.92, "monto_usd": 1836.82,
         "utilidad_usd": 1836.82, "reserva_usd": 57.29, "descuenta_reserva": False}
    )  # fmt: skip
    assert (linea.reserva_usd, linea.descuenta_reserva) == (57.29, False)


@pytest.mark.parametrize("valor", [None, False, "true", 1, "si"])
def test_bandera_liberal(valor) -> None:
    p = _payload(bandera=None)
    p["reparto_como_balance"] = valor
    assert RepartoPdfRequest.model_validate(p).reparto_como_balance is False


@pytest.mark.parametrize("valor", [None, "si", 1, 0, "true"])
def test_descuenta_reserva_liberal(valor) -> None:
    s = RepartoSocioLinea.model_validate(
        {"socio_nombre": "S", "porcentaje": 1, "monto_usd": 1, "descuenta_reserva": valor}
    )
    assert s.descuenta_reserva is None


# ---------------------------------------------------------------------------
# (b) Sin la bandera: byte a byte, con los campos nuevos presentes.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("pesos", [None, False])
@pytest.mark.parametrize("bandera", [None, False, "true"])
def test_sin_bandera_byte_a_byte(pesos, bandera) -> None:
    """Las mismas huellas del HEAD 3b893f1 que `test_reparto_en_pesos`: los
    campos del 0.0.101 sin la bandera no mueven nada."""
    p = _neutro(pp.reparto(bandera=pesos))
    if bandera is not None:
        p["reparto_como_balance"] = bandera
    req = RepartoPdfRequest.model_validate(p)
    assert _huella_pdf(render_reparto_pdf(req)) == _HUELLA_PDF_HEAD
    assert _huella_xlsx(render_reparto_xlsx(req)) == _HUELLA_XLSX_HEAD


@pytest.mark.parametrize("pesos", [None, True])
def test_sin_bandera_los_campos_no_cambian_el_documento(pesos) -> None:
    """Con reserva y provisión ≠ 0 pero SIN la bandera: el documento de
    siempre (la reserva resta, nada de AFAC), idéntico con la bandera ausente,
    en false o como texto (flujos del PDF y miembros del .xlsx)."""
    base_pdf = render_reparto_pdf(_req(pesos=pesos, bandera=None))
    base_xlsx = render_reparto_xlsx(_req(pesos=pesos, bandera=None))
    for valor in (False, "true"):
        req = _req(pesos=pesos, bandera=valor)
        assert _flujos(render_reparto_pdf(req)) == _flujos(base_pdf)
        assert _miembros(render_reparto_xlsx(req)) == _miembros(base_xlsx)
    texto = _texto_pdf(base_pdf)
    assert "Provisión AFAC" not in texto
    assert "(-) Reserva overhaul" in texto
    assert rp.FILA_RESERVA_INFORMATIVA not in texto
    assert "(se descuenta)" not in texto and "Reserva a guardar" not in texto
    a = _payload(pesos=pesos, bandera=None)["aviones"][0]["reparto"][0]
    if pesos:
        assert f"(-) Reserva overhaul (al T.C. {K:g}) {_d(-2415)} {_d(-140)}" in texto
        # Sin la bandera, el «Monto MXN» del API tal cual (ya con la reserva
        # descontada: lo decidió el API), sin utilidad ni reserva al lado.
        assert f"Socio A 60% {_d(a['monto_mxn'])} {_d(MONTO_A)}" in texto
        assert "Monto MXN USD (inf.)" in texto
    else:
        assert f"(-) Reserva overhaul {_d(-140)}" in texto
        assert f"Socio A 60% {_d(MONTO_A)}" in texto
    ws = load_workbook(BytesIO(base_xlsx)).active
    assert ws.cell(row=4, column=11).value == (
        "Reserva overhaul (MXN)" if pesos else "Reserva overhaul"
    )
    assert ws.cell(row=4, column=15).value is None
    assert not any("Provisión AFAC" in t or "provisión AFAC" in t for t in _textos_col_a(ws))


# ---------------------------------------------------------------------------
# (c) PDF en dólares (sin `reparto_en_pesos`).
# ---------------------------------------------------------------------------


def test_pdf_usd_cascada_del_balance() -> None:
    texto = _texto_pdf(render_reparto_pdf(_req(pesos=None)))
    afac = "(-) Provisión AFAC (100 USD/hr × 3.5 h cobradas)"
    fijos = "(-) Fijos (prorrateo + reparto manual) $-0.00"
    # La provisión AFAC resta después de los fijos; el saldo ya la trae.
    assert f"{fijos} {afac} $-350.00 Saldo disponible {_d(SALDO_TST)}" in texto
    # La reserva ya no resta: va DESPUÉS del saldo, informativa, con lo que
    # VuelaTour retiene a quien la descuenta.
    assert (
        f"Saldo disponible {_d(SALDO_TST)} {rp.FILA_RESERVA_INFORMATIVA} {_d(140)} "
        f"{rp.LINEA_RESERVA_RETENIDA.strip()} {_d(84)}"
    ) in texto
    assert "(-) Reserva overhaul" not in texto
    # XB-DOS: sin tarifa (provisión 0), reserva 0 y nadie la descuenta.
    assert (
        f"{rp.FILA_PROVISION_AFAC} $-0.00 Saldo disponible {_d(SALDO_DOS)} "
        f"{rp.FILA_RESERVA_INFORMATIVA} $0.00"
    ) in texto
    assert texto.count(rp.LINEA_RESERVA_RETENIDA.strip()) == 1
    assert "AVISO: Sin tarifa AFAC" not in texto  # 0 h cobradas: nada que avisar


def test_pdf_usd_socios_utilidad_reserva_y_a_recibir() -> None:
    texto = _texto_pdf(render_reparto_pdf(_req(pesos=None)))
    assert " ".join(rp.ENCABEZADOS_SOCIOS_BALANCE) in texto
    assert f"Socio A 60% {_d(UTIL_A)} {_d(RESERVA_A)} (se descuenta) {_d(MONTO_A)}" in texto
    assert f"Socio B 40% {_d(UTIL_B)} {_d(RESERVA_B)} {_d(MONTO_B)}" in texto
    assert f"Socio A 100% {_d(SALDO_DOS)} $0.00 {_d(SALDO_DOS)}" in texto
    assert texto.count("(se descuenta)") == 1
    assert "Socio % Monto" not in texto


def test_pdf_usd_resumen_por_socio_y_global() -> None:
    req = _req(pesos=None)
    texto = _texto_pdf(render_reparto_pdf(req))
    assert " ".join(rp.ENCABEZADOS_RESUMEN_SOCIOS_BALANCE) in texto
    assert f"Socio A {_d(RESERVA_A)} {_d(MONTO_A + SALDO_DOS)}" in texto
    assert f"Socio B {_d(RESERVA_B)} {_d(MONTO_B)}" in texto
    # «Gastos del periodo» incluye la provisión AFAC y EXCLUYE la reserva:
    # venta − gastos = saldo.
    venta = sum(a.ingresos_cobrado_usd for a in req.aviones)
    gastos = sum(
        a.gastos_directos_usd
        + a.gastos_indirectos_usd
        + a.permisos_usd
        + a.otros_usd
        + (a.afac_reservado_usd or 0)
        + a.comisiones_venta_usd
        for a in req.aviones
    )
    saldo = sum(a.saldo_usd for a in req.aviones)
    assert r2(venta - gastos) == r2(saldo) == r2(SALDO_TST + SALDO_DOS)
    assert f"{_d(venta)} $208.80 {_d(gastos)} {_d(saldo)}" in texto
    # El pie explica la cascada nueva.
    assert (
        "provisión AFAC (tarifa de la ficha del avión × horas cobradas) + otros ingresos "
        "del avión; la reserva de overhaul se informa por socio (lo que debe guardar) y "
        "solo se descuenta a quien lo pidió: VuelaTour la retiene. Montos en USD."
    ) in texto


def test_pdf_aviso_sin_tarifa_afac_con_horas_cobradas() -> None:
    p = _payload(pesos=None)
    p["aviones"][1]["afac_horas_cobradas_hr"] = 2.0
    texto = _texto_pdf(render_reparto_pdf(RepartoPdfRequest.model_validate(p)))
    assert (
        "AVISO: Sin tarifa AFAC configurada en la ficha del avión: la provisión AFAC es "
        "$0.00 pese a 2 h cobradas."
    ) in texto
    assert texto.count("Sin tarifa AFAC") == 1


def test_etiquetas_y_anchos() -> None:
    a = _req().aviones[0]
    assert rp.etiqueta_provision_afac(a) == "(-) Provisión AFAC (100 USD/hr × 3.5 h cobradas)"
    a.afac_tarifa_usd_hr, a.afac_horas_cobradas_hr = 112.5, 22.7
    assert rp.etiqueta_provision_afac(a) == "(-) Provisión AFAC (112.5 USD/hr × 22.7 h cobradas)"
    a.afac_tarifa_usd_hr = None
    assert rp.etiqueta_provision_afac(a) == rp.FILA_PROVISION_AFAC
    assert rp.aviso_sin_tarifa_afac(a).endswith("pese a 22.7 h cobradas.")
    a.afac_horas_cobradas_hr = 0
    assert rp.aviso_sin_tarifa_afac(a) is None
    # En la columna de etiquetas más angosta (pesos: 112 mm − relleno) caben
    # la fila gris a 7.5 pt Oblique y la de AFAC a 9 pt.
    ancho = 112 * mm - 12
    for t in (rp.FILA_RESERVA_INFORMATIVA, rp.LINEA_RESERVA_RETENIDA):
        assert stringWidth(t, "Helvetica-Oblique", 7.5) <= ancho, t
    a.afac_tarifa_usd_hr, a.afac_horas_cobradas_hr = 1234.56, 1234.5
    assert stringWidth(rp.etiqueta_provision_afac(a), "Helvetica", 9) <= ancho
    # Sin «−» ni «→» en lo que dibuja Helvetica.
    for t in (
        rp.FILA_PROVISION_AFAC, rp.FILA_RESERVA_INFORMATIVA, rp.LINEA_RESERVA_RETENIDA,
        rp._PIE_CASCADA_BALANCE,
    ):  # fmt: skip
        assert "−" not in t and "→" not in t


# ---------------------------------------------------------------------------
# (c) PDF en pesos (las dos banderas, como las manda el API 0.0.101).
# ---------------------------------------------------------------------------


def test_pdf_pesos_cascada_del_balance() -> None:
    req = _req()
    tst = req.aviones[0]
    texto = _texto_pdf(render_reparto_pdf(req))
    assert rp.LEYENDA_PESOS in texto
    afac = "(-) Provisión AFAC (100 USD/hr × 3.5 h cobradas)"
    assert f"{afac} {_d(-6037.5)} {_d(-350)} Saldo disponible {_d(tst.saldo_mxn)}" in texto
    assert (
        f"Saldo disponible {_d(tst.saldo_mxn)} {_d(SALDO_TST)} {rp.FILA_RESERVA_INFORMATIVA} "
        f"{_d(2415)} {_d(140)} en pesos al T.C. oficial del cierre ({K:g}) "
        f"{rp.LINEA_RESERVA_RETENIDA.strip()} {_d(1449)} {_d(84)}"
    ) in texto
    assert "(-) Reserva overhaul" not in texto
    # Socios: MXN principal y el USD informativo debajo de cada monto.
    a = tst.reparto[0]
    assert " ".join(rp.ENCABEZADOS_SOCIOS_BALANCE_PESOS) in texto
    assert (
        f"Socio A 60% {_d(a.utilidad_mxn)} {_d(UTIL_A)} USD {_d(a.reserva_mxn)} (se descuenta) "
        f"{_d(RESERVA_A)} USD {_d(a.monto_mxn)} {_d(MONTO_A)} USD"
    ) in texto
    assert a.monto_mxn == r2(a.utilidad_mxn - a.reserva_mxn)
    # Resumen por socio en pesos.
    dos = req.aviones[1].reparto[0]
    assert " ".join(rp.ENCABEZADOS_RESUMEN_SOCIOS_BALANCE_PESOS) in texto
    assert (
        f"Socio A {_d(a.reserva_mxn)} {_d(RESERVA_A)} USD {_d(a.monto_mxn + dos.monto_mxn)} "
        f"{_d(MONTO_A + SALDO_DOS)} USD"
    ) in texto
    # Resumen global: gastos con la provisión AFAC y sin la reserva.
    venta = r2(sum(x.ingresos_cobrado_mxn for x in req.aviones))
    gastos = r2(
        sum(
            x.gastos_directos_mxn
            + x.gastos_indirectos_mxn
            + x.permisos_mxn
            + x.otros_mxn
            + x.afac_reservado_mxn
            + x.comisiones_venta_mxn
            for x in req.aviones
        )
    )
    saldo = r2(sum(x.saldo_mxn for x in req.aviones))
    assert r2(venta - gastos) == saldo
    assert f"{_d(venta)} $208.80 USD {_d(gastos)} {_d(saldo)}" in texto
    assert "solo se descuenta a quien lo pidió: VuelaTour la retiene. Montos en pesos (MXN)" in (
        texto
    )


def test_pdf_pesos_sin_pesos_raya_nunca_cero() -> None:
    p = _payload()
    a = p["aviones"][0]
    a["afac_reservado_mxn"] = None
    a["reparto"][0]["reserva_mxn"] = None
    a["reparto"][0]["monto_mxn"] = None
    a["reserva_retenida_mxn"] = None
    texto = _texto_pdf(render_reparto_pdf(RepartoPdfRequest.model_validate(p)))
    assert "(-) Provisión AFAC (100 USD/hr × 3.5 h cobradas) — $-350.00" in texto
    assert f"— (se descuenta) {_d(RESERVA_A)} USD — {_d(MONTO_A)} USD" in texto
    assert f"{rp.LINEA_RESERVA_RETENIDA.strip()} — {_d(84)}" in texto
    # Con una línea sin pesos, el total del socio también es «—» (y el de
    # gastos del resumen, que lleva la provisión).
    assert f"Socio A — {_d(RESERVA_A)} USD — {_d(MONTO_A + SALDO_DOS)} USD" in texto


# ---------------------------------------------------------------------------
# (c) Excel en dólares.
# ---------------------------------------------------------------------------


def test_xlsx_usd_columna_afac_al_final_y_reserva_informativa() -> None:
    ws = _hoja(_payload(pesos=None))
    assert ws["K4"].value == rx.ENCABEZADO_RESERVA_INFORMATIVA
    assert ws["O4"].value == rx.ENCABEZADO_PROVISION_AFAC
    assert ws["O4"].comment.text == rx.NOTA_ENCABEZADO_AFAC
    assert (ws["O5"].value, ws["O6"].value) == (350.0, 0.0)
    assert ws["A7"].value == "TOTAL" and ws["O7"].value == 350.0 and ws["O7"].font.b
    assert ws["O5"].number_format == '"$"#,##0.00'
    # La reserva sigue en K con su monto, pero gris/cursiva (informativa).
    assert (ws["K5"].value, ws["K6"].value, ws["K7"].value) == (140.0, 0.0, 140.0)
    assert ws["K5"].font.i and ws["K5"].font.color.rgb.endswith("5B6470")
    assert ws["K7"].font.i and ws["K7"].font.b
    # D..J y L no se mueven; el saldo del API ya trae la provisión restada.
    assert ws["L5"].value == SALDO_TST and ws["L4"].value == "Saldo disponible"
    assert ws["D4"].value == "Venta del avión cobrada"
    assert ws.column_dimensions["O"].width == 18
    assert any(str(rg) == "A1:O1" for rg in ws.merged_cells.ranges)
    textos = _textos_col_a(ws)
    assert rx.NOTA_CASCADA_BALANCE in textos
    assert rx.NOTA_OTROS_INGRESOS_AVION not in textos


def test_xlsx_usd_reparto_por_socio_columnas_nuevas() -> None:
    ws = _hoja(_payload(pesos=None))
    t = _fila(ws, "Reparto por socio")
    assert [ws.cell(row=t + 1, column=c).value for c in range(1, 8)] == [
        "Avión", "Socio", "Porcentaje", "Monto USD (a recibir)",
        "Utilidad USD", "Reserva a guardar USD", "¿Se descuenta?",
    ]  # fmt: skip
    filas = [[ws.cell(row=t + 1 + i, column=c).value for c in range(1, 8)] for i in (1, 2, 3)]
    assert filas == [
        ["XB-TST", "Socio A", 0.6, MONTO_A, UTIL_A, RESERVA_A, "Sí"],
        ["XB-TST", "Socio B", 0.4, MONTO_B, UTIL_B, RESERVA_B, "No"],
        ["XB-DOS", "Socio A", 1.0, SALDO_DOS, SALDO_DOS, 0.0, "No"],
    ]
    assert ws.cell(row=t + 2, column=6).number_format == '"$"#,##0.00'
    assert ws.cell(row=t + 5, column=1).value == rx.NOTA_REPARTO_SOCIO_BALANCE


def test_xlsx_usd_con_otros_ingresos_afac_despues_de_la_columna_15() -> None:
    p = _payload(pesos=None)
    a = p["aviones"][0]
    a["otros_ingresos_avion_usd"] = 100.0
    a["otros_ingresos_avion_mxn"] = r2(100.0 * K)
    _cascada_balance(
        a,
        afac_usd=350.0,
        tarifa=100.0,
        horas=3.5,
        reserva_usd=140.0,
        socios=[("Socio A", 60, True), ("Socio B", 40, False)],
    )
    ws = _hoja(p)
    assert ws["O4"].value == rx.ENCABEZADO_OTROS_INGRESOS_AVION
    assert ws["P4"].value == rx.ENCABEZADO_PROVISION_AFAC
    assert (ws["O5"].value, ws["P5"].value, ws["P7"].value) == (100.0, 350.0, 350.0)
    assert ws["L5"].value == r2(SALDO_TST + 100.0)
    assert any(str(rg) == "A1:P1" for rg in ws.merged_cells.ranges)
    textos = _textos_col_a(ws)
    assert rx.NOTA_OTROS_INGRESOS_AVION_BALANCE in textos
    assert rx.NOTA_CASCADA_BALANCE not in textos and rx.NOTA_OTROS_INGRESOS_AVION not in textos
    assert "− provisión AFAC + otros ingresos del avión; la reserva de overhaul se informa" in (
        rx.NOTA_OTROS_INGRESOS_AVION_BALANCE
    )


# ---------------------------------------------------------------------------
# (c) Excel en pesos.
# ---------------------------------------------------------------------------


def test_xlsx_pesos_columna_afac_y_reserva_informativa() -> None:
    p = _payload()
    ws = _hoja(p)
    assert [ws.cell(row=4, column=c).value for c in range(1, 16)] == rx.encabezados_pesos(
        False, True
    )
    assert ws["K4"].value == rx.ENCABEZADO_RESERVA_INFORMATIVA_PESOS
    assert ws["O4"].value == rx.ENCABEZADO_PROVISION_AFAC_PESOS
    assert (ws["O5"].value, ws["O6"].value, ws["O7"].value) == (6037.5, 0.0, 6037.5)
    assert ws["K5"].value == 2415.0 and ws["K5"].font.i
    assert ws["K5"].font.color.rgb.endswith(rx.GRIS_INFO)
    assert ws["K5"].comment.text.endswith(f"{K:g}.")
    assert ws["K7"].font.i and ws["K7"].font.b
    assert ws["L5"].value == p["aviones"][0]["saldo_mxn"]
    assert ws.column_dimensions["O"].width == 20
    # La tabla informativa en dólares también lleva la provisión AFAC.
    t = _fila(ws, rx.TITULO_USD_INFORMATIVO)
    assert ws.cell(row=t + 1, column=11).value == rx.ENCABEZADO_RESERVA_INFORMATIVA
    assert ws.cell(row=t + 1, column=15).value == rx.ENCABEZADO_PROVISION_AFAC
    assert ws.cell(row=t + 2, column=15).value == 350.0
    assert ws.cell(row=t + 4, column=15).value == 350.0  # TOTAL USD
    textos = _textos_col_a(ws)
    assert rx.NOTA_CASCADA_BALANCE in textos
    assert any(t.startswith("Reserva overhaul pasada a pesos al T.C.") for t in textos)


def test_xlsx_pesos_reparto_por_socio_columnas_nuevas() -> None:
    p = _payload()
    ws = _hoja(p)
    t = _fila(ws, "Reparto por socio (en pesos)")
    assert [ws.cell(row=t + 1, column=c).value for c in range(1, 11)] == [
        "Avión", "Socio", "Porcentaje", "Monto MXN (a recibir)", "Monto USD (informativo)",
        "Utilidad MXN", "Utilidad USD (informativo)", "Reserva a guardar MXN",
        "Reserva a guardar USD (informativo)", "¿Se descuenta?",
    ]  # fmt: skip
    a = p["aviones"][0]["reparto"][0]
    fila = [ws.cell(row=t + 2, column=c).value for c in range(1, 11)]
    assert fila == [
        "XB-TST", "Socio A", 0.6, a["monto_mxn"], MONTO_A, a["utilidad_mxn"], UTIL_A,
        a["reserva_mxn"], RESERVA_A, "Sí",
    ]  # fmt: skip
    assert ws.cell(row=t + 2, column=7).font.i and ws.cell(row=t + 2, column=6).number_format == (
        '"$"#,##0.00'
    )
    assert ws.cell(row=t + 3, column=10).value == "No"
    assert ws.cell(row=t + 5, column=1).value == rx.NOTA_REPARTO_SOCIO_BALANCE
    # Una línea sin pesos: «—», jamás un 0.
    p["aviones"][0]["reparto"][0]["reserva_mxn"] = None
    ws = _hoja(p)
    assert ws.cell(row=t + 2, column=8).value == rx.SIN_PESOS


# ---------------------------------------------------------------------------
# Rutas: el JSON del API 0.0.101 (las dos banderas) entra sin 422.
# ---------------------------------------------------------------------------

TOKEN = "tok-reparto-balance"
client = TestClient(app)


def test_rutas_aceptan_el_payload_0101(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    for p in (_payload(), _payload(pesos=None)):
        res = client.post("/pdf/reparto", json=p, headers=headers)
        assert res.status_code == 200, res.text
        texto = _texto_pdf(res.content)
        assert "(-) Provisión AFAC (100 USD/hr × 3.5 h cobradas)" in texto
        assert rp.FILA_RESERVA_INFORMATIVA in texto
        res = client.post("/pdf/reparto-xlsx", json=p, headers=headers)
        assert res.status_code == 200, res.text
        ws = load_workbook(BytesIO(res.content)).active
        assert ws["O4"].value in (rx.ENCABEZADO_PROVISION_AFAC, rx.ENCABEZADO_PROVISION_AFAC_PESOS)
