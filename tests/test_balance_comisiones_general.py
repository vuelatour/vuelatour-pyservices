"""Balance GENERAL: comisiones solo en 'otros movimientos' y en la cascada; el
libro individual conserva su columna (7-oct-2026, API 0.0.66).

Pedido del cliente (7-oct): «estas comisiones se están duplicando en la
general, ya que las tenemos en el apartado de "Otros movimientos". Las
comisiones deben aparecer en ese apartado pero en el individual de los
avioncitos, ya que en el reporte de los aviones no está la pestaña "Otros
movimientos"» y «en el reporte mensual y en el balance general, en "otros
movimientos", en el ingreso estás duplicando la comisión».

El API 0.0.66 calcula todo (lo que absorbe cada avión sigue saliendo de
`comisionesDelVuelo`, la fuente única del reparto a socios); aquí solo se
pinta:
- la hoja de vuelos del general (las dos variantes) va SIN la columna
  COMISIONES y antes de comisiones (`_antes_de_comisiones`; los casos de la
  hoja están en `tests/test_balance_comisiones.py`);
- la cascada de cada avión de la hoja 'balance' del general gana dos líneas
  con `balance.utilidad_antes_comisiones_usd` y `balance.comisiones_usd`, y
  UTILIDAD DESPUÉS DE COMISIONES (hasta el 26-oct-2026 «UTILIDAD ANTES DE
  GASTOS») = ROUND(antes − comisiones, 2): el número de siempre. Sin ellas
  (API 0.0.65) o con comisiones en 0, la de siempre;
- con ese API, las leyendas de 'otros movimientos', del RESUMEN y del bloque
  VUELATOUR (empresa) explican el pago al vendedor «cubierto por el avión»
  (la línea de INGRESO «a cargo del avión» ya no viaja).
La UTILIDAD COBRADA y el reparto de cada avión NO cambian ni un centavo: se
congelan contra los números que el libro mostraba con el API 0.0.65 (HEAD
4189a1e).

13-oct-2026 (API 0.0.74, invariante 61 del API: «comisión del vendedor SOLO
como gasto real»): con la bandera `comision_vendedor_solo_gasto_real` (hoja
'otros movimientos', consolidado FLOTA y cada `aviones[]`) las leyendas de la
pestaña, del RESUMEN, del bloque VUELATOUR, de la cascada y de la hoja de
vuelos hablan del pago REAL que el avión carga (la fila llega «a cargo del
avión» con egreso $0.00). `_FILA_CUBIERTA`/`_FILA_EXCEDE`/`_CASCADA_0065`
(sin la bandera) no se tocan. Ajuste del mismo día (dos categorías): el pago
se captura como «Comisión del vendedor (VuelaTour)» —egresa en la fila— o
«Comisión del vendedor (avión)» —$0.00 «a cargo del avión …» si un avión lo
carga; si ninguno, lo paga VuelaTour y el concepto dice por qué—; las
leyendas de la bandera nombran las dos y la pestaña cuadra igual con la fila
mixta, la informativa y la que ningún avión carga.
"""

from __future__ import annotations

import copy
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    REGLA_COMISIONES_RESPALDO,
    BalanceAvionBalanceBloque,
    BalanceAvionRequest,
    BalanceAvionTotales,
    BalanceAvionVuelo,
    BalanceGeneralRequest,
)
from app.services import xlsx_formulas
from app.services.balance_avion_xlsx import (
    _COMISIONES_FUERA_DE_LA_HOJA,
    _COMISIONES_FUERA_DE_LA_HOJA_GASTO_REAL,
    _CPH_NOTA_REMANENTE,
    _EMPRESA_REGLA_COMISIONES,
    _EMPRESA_REGLA_COMISIONES_CUBRE_AVION,
    _EMPRESA_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION,
    _EMPRESA_REGLA_COMISIONES_GASTO_REAL,
    _FILA_ANTES_COMISIONES,
    _FILA_COMISIONES,
    _LEYENDA_PROVISION,
    _NOTA_CASCADA_ANTES_COMISIONES,
    _NOTA_CASCADA_COMISIONES,
    _NOTA_CASCADA_COMISIONES_GASTO_REAL,
    _NOTA_COBRADO_REAL_FLOTA,
    _NOTA_COBRANZA_NETO,
    _NOTA_COBRANZA_NETO_FLOTA,
    _NOTA_COMISIONES_FLOTA,
    _NOTA_COMISIONES_FLOTA_GASTO_REAL,
    _NOTA_EMPRESA_BASE,
    _NOTA_EMPRESA_BASE_DOS_CATEGORIAS,
    _NOTA_ENCABEZADO_GANANCIA_FLOTA,
    _NOTA_ENCABEZADO_GANANCIA_FLOTA_GASTO_REAL,
    _NOTA_ENCABEZADO_REMANENTE,
    _NOTA_ENCABEZADO_REMANENTE_GASTO_REAL,
    _NOTA_GANANCIA_FLOTA,
    _OM_PAGO_VENDEDOR_DOS_CATEGORIAS,
    _OM_REGLA_COMISIONES,
    _OM_REGLA_COMISIONES_CUBRE_AVION,
    _OM_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION,
    _OM_REGLA_COMISIONES_GASTO_REAL,
    _RESUMEN_COMISIONES,
    _RESUMEN_COMISIONES_GASTO_REAL,
    _RESUMEN_PAGO_VENDEDOR_DOS_CATEGORIAS,
    _RESUMEN_REGLA_CUBRE_AVION,
    _RESUMEN_REGLA_GASTO_REAL,
    AMBER,
    _antes_de_comisiones,
    _avion_cubre_al_vendedor,
    _cascada_con_comisiones,
    _comentario_cobro,
    _con_regla,
    _fila_antes_de_comisiones,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from tests import test_balance_comisiones as tc
from tests import test_balance_formulas as tf
from tests._evaluador_formulas import verificar_libro

r2 = tf.r2

MAESTRA = "reporte horas FLOTA"
ANTES = "UTILIDAD DESPUÉS DE COMISIONES USD"  # rótulo desde el 26-oct-2026
DESPUES = "UTILIDAD DESPUÉS DE GASTOS USD"
COBRADA = "UTILIDAD COBRADA USD"

# La cascada que el libro mostraba con el API 0.0.65 (HEAD 4189a1e, este
# mismo payload del espejo, general y libro individual): utilidad antes de
# gastos (ya después de comisiones), después de gastos, COBRADA y lo de cada
# socio. Es la regla dura del 7-oct-2026: el dinero de los socios no se mueve.
_CASCADA_0065 = {
    "XA-TST": {
        ANTES: 6750.29,
        DESPUES: 5360.13,
        COBRADA: 3975.42,
        "socios": {"Socio A": 2385.25, "Socio B": 1590.17},
    },
    "XB-DOS": {
        ANTES: 5064.86,
        DESPUES: 4425.95,
        COBRADA: 3441.11,
        "socios": {"Socio C": 3441.11},
    },
}

# ---------------------------------------------------------------------------
# Payloads: el espejo del API 0.0.65 de `test_balance_comisiones` + lo que
# agrega el 0.0.66.
# ---------------------------------------------------------------------------


def _con_api_0066(libro: dict) -> dict:
    """El libro de un avión como lo manda el API 0.0.66: `balance` gana
    `comisiones_usd` (Σ de las COMISIONES de sus filas, cada una al T.C. con
    que el API convierte su ganancia —z o el T.C. promedio del libro—,
    round2) y `utilidad_antes_comisiones_usd` (= UTILIDAD ANTES + eso), SOLO
    si alguna fila es de la regla (como `totales.comisiones_mxn`). Lo demás,
    idéntico al 0.0.65: la utilidad, la cascada y los socios no cambian."""
    libro = copy.deepcopy(libro)
    filas = libro["vuelos"]
    if not any("comisiones_detalle" in f for f in filas):
        return libro
    zs = [f["tc_costos"] for f in filas if f["tc_costos"] is not None]
    tc_prom = tf._suma(zs) / len(zs) if zs else None
    usd = tf._suma(
        f["comisiones_mxn"] / (f["tc_costos"] if f["tc_costos"] is not None else tc_prom)
        for f in filas
        if f.get("comisiones_mxn")
    )
    balance = libro["balance"]
    balance["comisiones_usd"] = r2(usd)
    balance["utilidad_antes_comisiones_usd"] = r2(
        balance["utilidad_antes_usd"] + balance["comisiones_usd"]
    )
    return libro


def _libros(api_0066: bool = True) -> list[dict]:
    libros = [tc._libro_tst(), tc._libro_dos()]
    return [_con_api_0066(x) for x in libros] if api_0066 else libros


def _payload(variante: str = "mensual", *, api_0066: bool = True) -> dict:
    """Balance general (dict con la forma del API) con el bloque VUELATOUR
    (empresa)."""
    return {**tc._general_dict(_libros(api_0066)), "variante": variante}


def _render(p: dict) -> bytes:
    return render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))


def _wb(data: bytes):
    """(con fórmulas, con caché)."""
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _pie(ws) -> list[str]:
    return [c.value for c in ws["A"] if isinstance(c.value, str)]


def _miembros(data: bytes) -> dict[str, bytes]:
    """Todos los miembros del .xlsx salvo docProps/core.xml (lleva la hora)."""
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _bloques(ws, wv) -> dict[str, dict]:
    """Bloques de la hoja 'balance' del general: matrícula → {etiqueta →
    (fila, celda con fórmula o valor, número a la vista)} de la cascada y sus
    socios ({nombre → número a la vista})."""
    matriculas = {"XA-TST", "XB-DOS"}
    bloques: dict[str, dict] = {}
    actual = None
    for c in ws["A"]:
        texto = c.value
        if texto in matriculas:
            actual = bloques.setdefault(texto, {"filas": {}, "socios": {}})
            continue
        if actual is None or not isinstance(texto, str):
            continue
        if texto.startswith("Reparto a socios"):
            actual["en_socios"] = True
            continue
        if texto in ("SOCIO", "VUELATOUR (empresa)"):
            if texto == "VUELATOUR (empresa)":
                actual = None
            continue
        if actual.get("en_socios"):
            actual["socios"][texto] = wv.cell(row=c.row, column=3).value
        else:
            celda = ws.cell(row=c.row, column=2)
            actual["filas"][texto] = (c.row, celda, wv.cell(row=c.row, column=2).value)
    return bloques


@pytest.fixture
def registros(monkeypatch) -> list[xlsx_formulas.RegistroFormulas]:
    """Registro de fórmulas de cada libro que se cierra (cuáles se quedaron
    como fórmula y cuáles volvieron a valor)."""
    capturados: list[xlsx_formulas.RegistroFormulas] = []
    original = xlsx_formulas.finalizar

    def espia(wb):
        cache = original(wb)
        capturados.append(xlsx_formulas.registro(wb))
        return cache

    monkeypatch.setattr(xlsx_formulas, "finalizar", espia)
    return capturados


# ---------------------------------------------------------------------------
# 1) Esquema: campos ADITIVOS.
# ---------------------------------------------------------------------------


def test_campos_nuevos_de_la_cascada_con_default() -> None:
    b = BalanceAvionBalanceBloque()
    assert (b.comisiones_usd, b.utilidad_antes_comisiones_usd) == (None, None)
    b = BalanceAvionBalanceBloque.model_validate(
        {"comisiones_usd": 258, "utilidad_antes_comisiones_usd": "7008.49"}
    )
    assert (b.comisiones_usd, b.utilidad_antes_comisiones_usd) == (258.0, 7008.49)
    # Las dos líneas: solo con los dos campos y comisiones ≠ 0.
    assert _cascada_con_comisiones(b)
    for sin in (
        BalanceAvionBalanceBloque(),
        BalanceAvionBalanceBloque(comisiones_usd=258.2),
        BalanceAvionBalanceBloque(utilidad_antes_comisiones_usd=7008.49),
        BalanceAvionBalanceBloque(comisiones_usd=0.0, utilidad_antes_comisiones_usd=6750.29),
    ):
        assert not _cascada_con_comisiones(sin), sin


# ---------------------------------------------------------------------------
# 2) La hoja de vuelos del general antes de comisiones (`_antes_de_comisiones`).
# ---------------------------------------------------------------------------


def test_fila_con_comisiones_vuelve_a_la_de_antes_de_la_regla() -> None:
    v = BalanceAvionVuelo(
        remanente_mxn=52896.25,
        comisiones_mxn=3861.5,
        ganancia_mxn=49034.75,
        tc_costos=17.32,
        ganancia_usd=2831.11,
        comision_vendedor_mxn=None,
    )
    antes = _fila_antes_de_comisiones(v)
    assert (antes.ganancia_mxn, antes.comisiones_mxn) == (52896.25, None)
    # Remanente ÷ T.C. de costos (la fórmula visible de la celda).
    assert antes.ganancia_usd == r2(52896.25 / 17.32) == 3054.06
    # El vuelo del API no se toca (es una copia).
    assert (v.ganancia_mxn, v.ganancia_usd, v.comisiones_mxn) == (49034.75, 2831.11, 3861.5)
    # Sin comisiones (vuelo anterior o API previo): la MISMA fila.
    sin = BalanceAvionVuelo(remanente_mxn=100.0, ganancia_mxn=100.0, ganancia_usd=5.0)
    assert _fila_antes_de_comisiones(sin) is sin
    # Regla con 0 explícito (el API manda null, pero el esquema lo admite; y
    # un API ≤ 0.0.64 con `comision_vendedor_mxn` 0): la fila SIN la llave,
    # igual a la de null — si no, GANANCIA MXN citaba la columna que el
    # general no tiene (revisión 7-oct-2026). Nada más cambia.
    for cero in (
        BalanceAvionVuelo(
            remanente_mxn=100.0, comisiones_mxn=0.0, ganancia_mxn=100.0, ganancia_usd=5.0
        ),
        BalanceAvionVuelo(
            remanente_mxn=100.0, comision_vendedor_mxn=0.0, ganancia_mxn=100.0, ganancia_usd=5.0
        ),
    ):
        limpia = _fila_antes_de_comisiones(cero)
        assert limpia is not cero
        assert limpia == sin
        assert (cero.comisiones_mxn, cero.comision_vendedor_mxn) != (None, None)  # copia
    # Sin T.C. de costos el API la convirtió con el T.C. promedio del libro de
    # SU avión, que esta hoja no tiene: vacía, jamás un número que no es.
    sin_tc = BalanceAvionVuelo(
        remanente_mxn=400.0, comisiones_mxn=50.0, ganancia_mxn=350.0, ganancia_usd=20.0
    )
    assert _fila_antes_de_comisiones(sin_tc).ganancia_usd is None


def test_totales_antes_de_comisiones() -> None:
    req = BalanceAvionRequest(
        vuelos=[
            BalanceAvionVuelo(
                remanente_mxn=1000.0,
                comisiones_mxn=100.0,
                ganancia_mxn=900.0,
                tc_costos=20.0,
                ganancia_usd=45.0,
            ),
            BalanceAvionVuelo(remanente_mxn=500.25, ganancia_mxn=500.25, ganancia_usd=25.01),
        ],
        totales=BalanceAvionTotales(
            remanente_mxn=1500.25, comisiones_mxn=100.0, ganancia_mxn=1400.25, ganancia_usd=70.01
        ),
    )
    antes = _antes_de_comisiones(req)
    t = antes.totales
    assert (t.ganancia_mxn, t.comisiones_mxn) == (1500.25, None)
    assert t.ganancia_usd == r2(50.0 + 25.01) == 75.01  # Σ de las filas como se ven
    assert [v.ganancia_usd for v in antes.vuelos] == [50.0, 25.01]
    # El request original no cambia y, sin comisiones (API previo o vuelos
    # anteriores a la regla), sale el MISMO objeto: la hoja de siempre.
    assert req.totales.ganancia_mxn == 1400.25
    previo = tf._individual_tst()
    assert _antes_de_comisiones(previo) is previo


def test_regla_con_cero_deja_los_totales_del_api() -> None:
    """Filas de la regla con COMISIONES en 0 explícito y ninguna ≠ 0: solo
    esas filas pierden la llave; los TOTALES son los del API (el MISMO
    objeto: su GANANCIA ya es el remanente, como con null)."""
    req = BalanceAvionRequest(
        vuelos=[
            BalanceAvionVuelo(
                remanente_mxn=1000.0,
                comisiones_mxn=0.0,
                ganancia_mxn=1000.0,
                tc_costos=20.0,
                ganancia_usd=50.0,
            ),
            BalanceAvionVuelo(remanente_mxn=500.25, ganancia_mxn=500.25, ganancia_usd=25.01),
        ],
        totales=BalanceAvionTotales(
            remanente_mxn=1500.25, comisiones_mxn=0.0, ganancia_mxn=1500.25, ganancia_usd=75.01
        ),
    )
    antes = _antes_de_comisiones(req)
    assert antes.totales is req.totales
    assert antes.vuelos[1] is req.vuelos[1]
    assert antes.vuelos[0].comisiones_mxn is None
    assert antes.vuelos[0].model_dump(exclude={"comisiones_mxn"}) == req.vuelos[0].model_dump(
        exclude={"comisiones_mxn"}
    )


def _con_cero(p: dict, clave: str = "#403 · Traslado") -> dict:
    """El payload con COMISIONES en 0 explícito en la fila `clave` (en el
    consolidado y en el libro de su avión) en lugar del null del API."""
    filas = [*p["consolidado"]["vuelos"], *(f for a in p["aviones"] for f in a["vuelos"])]
    tocadas = 0
    for f in filas:
        if f["clave"] == clave:
            f["comisiones_mxn"] = 0.0
            tocadas += 1
    assert tocadas >= 2, clave
    return p


def _columna(ws, encabezado: str) -> str:
    return next(c.column_letter for c in ws[2] if c.value == encabezado)


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_fila_con_comisiones_en_cero_se_pinta_como_la_de_null(variante, registros) -> None:
    """#403 (vuelo de la regla sin comisiones) con COMISIONES en 0 explícito
    en lugar de null: el general sale BYTE A BYTE igual. Antes, en el mensual,
    GANANCIA MXN citaba la columna COMISIONES que la hoja ya no tiene: la
    celda iba como valor y el encabezado perdía su nota por la de «calculado
    por el sistema» (revisión 7-oct-2026)."""
    data_null = _render(_payload(variante))
    degradadas_null = registros[-1].degradadas()
    data_cero = _render(_con_cero(_payload(variante)))
    assert registros[-1].degradadas() == degradadas_null
    assert verificar_libro(data_cero)
    assert _miembros(data_cero) == _miembros(data_null)
    if variante == "mensual":
        ws = _wb(data_cero)[0][MAESTRA]
        fila = next(c.row for c in ws["A"] if c.value == "#403 · Traslado")
        col = _columna(ws, "GANANCIA\nMXN")
        assert ws[f"{col}{fila}"].value == f"=AC{fila}"
        assert ws[f"{col}2"].comment.text == _NOTA_ENCABEZADO_GANANCIA_FLOTA


def test_libro_sin_comisiones_con_un_cero_es_el_de_siempre(registros) -> None:
    """Periodo sin comisiones ≠ 0 (el espejo del API previo a la regla) y una
    fila con COMISIONES en 0 explícito: el Balance general es el de siempre,
    byte a byte (TOTALES del API, sin fórmulas degradadas de más)."""
    base = tf._general_payload()
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(base))
    degradadas = registros[-1].degradadas()
    con_cero = tf._general_payload()
    fila = con_cero["consolidado"]["vuelos"][0]
    assert "comisiones_mxn" not in fila
    fila["comisiones_mxn"] = 0.0
    data_cero = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(con_cero))
    assert registros[-1].degradadas() == degradadas
    assert _miembros(data_cero) == _miembros(data)


# ---------------------------------------------------------------------------
# 3) La cascada del general con las dos líneas (API 0.0.66).
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_cascada_del_general_con_las_dos_lineas(variante, registros) -> None:
    p = _payload(variante)
    data = _render(p)
    # Ninguna fórmula de la cascada vuelve a valor (la de siempre: COSTO TOTAL
    # del #405 en el Balance mensual).
    degradadas = [(MAESTRA, "W9")] if variante == "mensual" else []
    assert registros[-1].degradadas() == degradadas
    assert verificar_libro(data)
    ws, wv = (x["balance"] for x in _wb(data))
    bloques = _bloques(ws, wv)
    for avion in p["aviones"]:
        b = avion["balance"]
        filas = bloques[avion["matricula"]]["filas"]
        etiquetas = list(filas)
        assert etiquetas[:3] == [_FILA_ANTES_COMISIONES, _FILA_COMISIONES, ANTES]
        fa, celda_a, valor_a = filas[_FILA_ANTES_COMISIONES]
        fc, celda_c, valor_c = filas[_FILA_COMISIONES]
        fu, celda_u, valor_u = filas[ANTES]
        assert (fc, fu) == (fa + 1, fa + 2)
        # Las dos nuevas: valores del API, con su nota.
        assert valor_a == b["utilidad_antes_comisiones_usd"]
        assert valor_c == b["comisiones_usd"] > 0
        assert celda_a.comment.text == _NOTA_CASCADA_ANTES_COMISIONES
        assert celda_c.comment.text == _con_regla(_NOTA_CASCADA_COMISIONES, None)
        # Rótulo corto desde el 26-oct-2026: cabe en la columna A sin envolver.
        etiqueta = ws.cell(row=fa, column=1)
        assert len(etiqueta.value) <= 40 and not etiqueta.alignment.wrap_text
        # UTILIDAD DESPUÉS DE COMISIONES = ROUND(antes − comisiones, 2): el
        # número de siempre del API; de ahí para abajo, las fórmulas de siempre.
        assert celda_u.value == f"=ROUND(B{fa}-B{fc},2)"
        assert valor_u == b["utilidad_antes_usd"]
        assert filas[DESPUES][1].value == f"=ROUND(B{fu}-B{fu + 1}-B{fu + 2}-B{fu + 3}-B{fu + 4},2)"
        assert filas[COBRADA][1].value == f"=ROUND(B{fu + 5}-B{fu + 6},2)"


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_utilidad_cobrada_y_socios_no_cambian_ni_un_centavo(variante) -> None:
    """Regla dura del 7-oct-2026: la utilidad cobrada por avión (y lo de cada
    socio) es la del API 0.0.65 — congelada — con y sin las líneas nuevas, en
    el general y en el libro individual."""
    for api_0066 in (False, True):
        ws, wv = (x["balance"] for x in _wb(_render(_payload(variante, api_0066=api_0066))))
        bloques = _bloques(ws, wv)
        for matricula, esperado in _CASCADA_0065.items():
            filas = bloques[matricula]["filas"]
            for etiqueta in (ANTES, DESPUES, COBRADA):
                assert filas[etiqueta][2] == esperado[etiqueta], (api_0066, matricula, etiqueta)
            assert bloques[matricula]["socios"] == esperado["socios"], (api_0066, matricula)
            assert (_FILA_COMISIONES in filas) == api_0066
    # El libro individual de cada avión dice lo mismo (con o sin los campos).
    for libro in (*_libros(api_0066=False), *_libros(api_0066=True)):
        wv = _wb(render_balance_avion_xlsx(BalanceAvionRequest.model_validate(libro)))[1]
        bal = wv["balance"]
        esperado = _CASCADA_0065[libro["matricula"]]
        vistos = {c.value: bal.cell(row=c.row, column=2).value for c in bal["A"]}
        for etiqueta in (ANTES, DESPUES, COBRADA):
            assert vistos[etiqueta] == esperado[etiqueta], (libro["matricula"], etiqueta)
        socios = {n: bal.cell(row=c.row, column=3).value for c in bal["A"] if (n := c.value)}
        for nombre, monto in esperado["socios"].items():
            assert socios[nombre] == monto


def _cascadas(libro) -> dict[str, object]:
    """Celdas de la hoja 'balance' de los bloques de los aviones (hasta el
    bloque VUELATOUR (empresa), cuyo pie depende de la versión del API)."""
    ws = libro["balance"]
    fin = next((c.row for c in ws["A"] if c.value == "VUELATOUR (empresa)"), ws.max_row + 1)
    return {c.coordinate: c.value for fila in ws.iter_rows(max_row=fin - 1) for c in fila}


def test_sin_comisiones_del_api_la_cascada_es_la_de_siempre() -> None:
    """API 0.0.65 (sin los campos), comisiones en 0 o solo uno de los dos
    campos: la cascada de cada avión es la del API 0.0.65, celda por celda
    (fórmulas y números)."""
    base = _wb(_render(_payload(api_0066=False)))
    for retoque in (
        {"comisiones_usd": 0.0, "utilidad_antes_comisiones_usd": 6750.29},
        {"comisiones_usd": 258.2},
        {"comisiones_usd": None, "utilidad_antes_comisiones_usd": None},
    ):
        p = _payload(api_0066=False)
        p["aviones"][0]["balance"].update(retoque)
        otro = _wb(_render(p))
        for libro_a, libro_b in zip(base, otro, strict=True):
            assert _cascadas(libro_a) == _cascadas(libro_b), retoque
        assert _FILA_COMISIONES not in [c.value for c in otro[0]["balance"]["A"]]


def test_el_libro_individual_no_pinta_las_lineas_nuevas() -> None:
    """El libro de un avión ya resta sus comisiones en su columna COMISIONES:
    con los campos del API 0.0.66 sale byte a byte igual."""
    for previo, nuevo in zip(_libros(api_0066=False), _libros(api_0066=True), strict=True):
        assert "comisiones_usd" in nuevo["balance"]
        a = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(previo))
        b = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(nuevo))
        assert _miembros(a) == _miembros(b), nuevo["matricula"]


# ---------------------------------------------------------------------------
# 4) 'otros movimientos', RESUMEN y bloque VUELATOUR (empresa) con el API
#    0.0.66: sin la línea de INGRESO «a cargo del avión».
# ---------------------------------------------------------------------------

# Fila del vuelo con comisión del vendedor como la arma el API 0.0.66: lo
# cobrado al cliente UNA vez y el pago al vendedor «cubierto por el avión»
# (sin pago real o sin exceso: 0), y otra con pago real que EXCEDE la
# provisión (solo el exceso resta).
_FILA_CUBIERTA = {
    "clave": "#401 · Cliente Uno",
    "avion_color": "#2563EB",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-09-03",
    # Texto REAL del API 0.0.66 (7b561fc5): sin gasto real capturado lleva la
    # marca «· PROVISIÓN (…)» que pinta el amarillo pedido el 6-oct.
    "concepto_egreso": (
        "pago comisión vendedor (Pablo Canales) · cubierto por el avión XA-TST "
        "(provisión $2,360.00 en su balance) · PROVISIÓN (sin gasto real capturado)"
    ),
    "egreso_mxn": 0.0,
    "fecha_egreso": "2026-09-03",
    "concepto_ingreso": "comisión vendedor (Pablo Canales)",
    "ingreso_mxn": 2360.0,
    "fecha_ingreso": "2026-09-03",
    "remanente_mxn": 2360.0,
}
_FILA_EXCEDE = {
    "clave": "#407 · Cliente Cinco",
    "avion_color": "#16A34A",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-09-18",
    "concepto_egreso": (
        "pago comisión vendedor (Alex Saab) · cubierto por el avión XB-DOS "
        "(provisión $1,180.00 en su balance) · excede $320.00"
    ),
    "egreso_mxn": 320.0,
    "fecha_egreso": "2026-09-25",
    "concepto_ingreso": "comisión vendedor (Alex Saab)",
    "ingreso_mxn": 1180.0,
    "fecha_ingreso": "2026-09-18",
    "remanente_mxn": 860.0,
}


def _payload_om_0066(variante: str = "mensual") -> dict:
    p = _payload(variante)
    p["consolidado"]["otros_movimientos"]["filas"] = [_FILA_CUBIERTA, _FILA_EXCEDE]
    return p


def test_la_senal_del_api_0066_es_comisiones_usd_en_los_aviones() -> None:
    assert _avion_cubre_al_vendedor(BalanceGeneralRequest.model_validate(_payload()))
    assert not _avion_cubre_al_vendedor(
        BalanceGeneralRequest.model_validate(_payload(api_0066=False))
    )
    assert not _avion_cubre_al_vendedor(BalanceGeneralRequest())


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_otros_movimientos_sin_la_linea_de_ingreso(variante) -> None:
    data = _render(_payload_om_0066(variante))
    assert verificar_libro(data)
    wb, wv = _wb(data)
    om, omv = wb["otros movimientos"], wv["otros movimientos"]
    a2 = om["A2"].value
    regla = REGLA_COMISIONES_RESPALDO
    assert _OM_REGLA_COMISIONES_CUBRE_AVION.format(regla=regla) in a2
    assert _OM_REGLA_COMISIONES.format(regla=regla) not in a2
    assert "a cargo del avión" not in a2
    assert a2.index("cubierto por el avión") < a2.index(" UNA FILA POR VUELO")
    # Las filas, tal cual las arma el API: lo cobrado UNA vez y el pago
    # cubierto por el avión (0, con la marca «· PROVISIÓN (» ⇒ amarillo) o
    # solo su exceso cuando hay gasto real (sin marca ⇒ sin amarillo).
    assert om["F4"].value == _FILA_CUBIERTA["concepto_egreso"]
    assert (om["G4"].value, om["J4"].value) == (0.0, 2360.0)
    assert (om["G5"].value, om["J5"].value) == (320.0, 1180.0)
    assert omv["L4"].value == 2360.0 and omv["L5"].value == 860.0

    def _rellenos(fila: int) -> set[str]:
        return {
            om.cell(row=fila, column=c).fill.fgColor.rgb
            for c in range(1, 11)
            if om.cell(row=fila, column=c).fill.fill_type
        }

    assert "00" + AMBER in _rellenos(4)
    assert "00" + AMBER not in _rellenos(5)
    assert _LEYENDA_PROVISION in [c.value for c in om["A"]]
    tot = next(c.row for c in om["A"] if c.value == "TOTALES")
    assert omv[f"J{tot}"].value == r2(2360.0 + 1180.0 + 85.33)
    assert omv[f"G{tot}"].value == 320.0


def test_resumen_y_empresa_con_el_pago_cubierto_por_el_avion() -> None:
    wb = _wb(_render(_payload_om_0066()))[0]
    regla = REGLA_COMISIONES_RESPALDO
    resumen = next(
        c.value
        for c in wb["RESUMEN flota"]["A"]
        if isinstance(c.value, str) and c.value.startswith("COMISIONES = ")
    )
    assert resumen.endswith(
        "(ingreso cobrado y pago apareado como PROVISIÓN a la fecha del vuelo)."
        + _RESUMEN_REGLA_CUBRE_AVION.format(regla=regla)
    )
    pie = _pie(wb["balance"])
    base = _NOTA_EMPRESA_BASE + _EMPRESA_REGLA_COMISIONES_CUBRE_AVION.format(regla=regla)
    assert base in pie
    assert not any(_EMPRESA_REGLA_COMISIONES.format(regla=regla) in x for x in pie)


def test_sin_la_senal_del_api_0066_las_leyendas_de_antes() -> None:
    """API 0.0.65 (sus filas aún traen la línea «a cargo del avión»): las
    leyendas de antes; el RESUMEN, sin la frase del pago cubierto."""
    wb = _wb(_render(_payload(api_0066=False)))[0]
    regla = REGLA_COMISIONES_RESPALDO
    a2 = wb["otros movimientos"]["A2"].value
    assert _OM_REGLA_COMISIONES.format(regla=regla) in a2
    assert "cubierto por el avión" not in a2
    # La comisión bancaria del avión ya no «resta en su columna COMISIONES de
    # la hoja de vuelos» (el general no la tiene): la de su libro individual.
    assert "resta en la columna COMISIONES de su libro individual" in a2
    resumen = [c.value for c in wb["RESUMEN flota"]["A"] if isinstance(c.value, str)]
    assert not any(_RESUMEN_REGLA_CUBRE_AVION.format(regla=regla) in x for x in resumen)
    base = _NOTA_EMPRESA_BASE + _EMPRESA_REGLA_COMISIONES.format(regla=regla)
    assert base in _pie(wb["balance"])


# ---------------------------------------------------------------------------
# 5) Textos: la regla del API y la nota de netos de 'cobranza'.
# ---------------------------------------------------------------------------

REGLA_API = "regla desde 15-oct-2026"


@pytest.mark.parametrize(
    "texto",
    [
        _NOTA_COMISIONES_FLOTA,
        _NOTA_COBRADO_REAL_FLOTA,
        _NOTA_COBRANZA_NETO_FLOTA,
        # API 0.0.74 (13-oct-2026, invariante 61): la versión «gasto real».
        _NOTA_COMISIONES_FLOTA_GASTO_REAL,
    ],
)
def test_textos_nuevos_con_la_regla_del_api(texto) -> None:
    assert "septiembre 2026" in texto
    assert _con_regla(texto, None) == texto
    con = _con_regla(texto, REGLA_API)
    assert "septiembre" not in con and REGLA_API in con


@pytest.mark.parametrize(
    "texto",
    [
        _COMISIONES_FUERA_DE_LA_HOJA,
        _NOTA_ENCABEZADO_GANANCIA_FLOTA,
        _NOTA_ENCABEZADO_REMANENTE,
        _NOTA_CASCADA_COMISIONES,
        _COMISIONES_FUERA_DE_LA_HOJA_GASTO_REAL,
        _NOTA_ENCABEZADO_GANANCIA_FLOTA_GASTO_REAL,
        _NOTA_ENCABEZADO_REMANENTE_GASTO_REAL,
        _NOTA_CASCADA_COMISIONES_GASTO_REAL,
    ],
)
def test_notas_de_celda_sin_la_fecha_de_la_regla(texto) -> None:
    """26-oct-2026 (pedido del cliente: las notas «causan ruido»): las notas
    de CELDA ya no nombran la regla ni su mes —eso lo explica el pie de la
    hoja con la etiqueta del API—, así que `_con_regla` las deja tal cual."""
    assert "septiembre" not in texto and "regla" not in texto
    assert _con_regla(texto, REGLA_API) == texto
    assert len(texto) <= 240


def test_notas_de_la_cascada_no_prometen_cuadres_que_no_siempre_se_dan() -> None:
    """Revisión 7-oct-2026. La nota de «(−) COMISIONES» decía «Son las mismas
    que descuenta el reparto a socios»: no es así cuando un vuelo que aún no
    se completa ya tiene un anticipo con comisión bancaria (caso real: #314 de
    N4142R, CONFIRMADO del 15-sep-2026, $5,000 de comisión) — el balance
    cuenta ese cobro y el reparto, que solo lee vuelos completados o
    cancelados, no. La de «antes de comisiones» tampoco cuadra siempre al
    centavo con la suma de GANANCIA USD de la hoja de vuelos. Desde el
    26-oct-2026 las notas son cortas: ninguna afirma ese cuadre."""
    for com in (_NOTA_CASCADA_COMISIONES, _NOTA_CASCADA_COMISIONES_GASTO_REAL):
        assert "Son las mismas" not in com
        assert "misma cifra que descuenta el reparto" not in com
        assert "reparto" not in com
    antes = _NOTA_CASCADA_ANTES_COMISIONES
    assert "Puede no cuadrar exacto con la suma de GANANCIA USD" in antes
    # Ninguna llega al tope de alto de la nota (no se cortan).
    for texto in (_NOTA_CASCADA_COMISIONES, _NOTA_CASCADA_COMISIONES_GASTO_REAL, antes):
        assert _comentario_cobro(texto).height < 260


def test_nota_de_la_cascada_con_la_regla_del_api() -> None:
    p = _payload()
    p["regla_comisiones"] = REGLA_API
    ws = _wb(_render(p))[0]["balance"]
    fila = next(c.row for c in ws["A"] if c.value == _FILA_COMISIONES)
    assert ws.cell(row=fila, column=2).comment.text == _con_regla(
        _NOTA_CASCADA_COMISIONES, REGLA_API
    )


def test_cobranza_del_general_nombra_el_libro_individual() -> None:
    """La nota de los netos de 'cobranza' decía que la parte del avión resta
    «en la columna COMISIONES de la hoja de vuelos»: en el general ya no está
    ahí. El libro individual conserva su texto."""
    gen = _pie(_wb(_render(_payload()))[0]["cobranza"])
    assert _NOTA_COBRANZA_NETO_FLOTA in gen and _NOTA_COBRANZA_NETO not in gen
    assert "del libro individual de cada avión" in _NOTA_COBRANZA_NETO_FLOTA
    ind = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(_libros()[0]))
    pie = _pie(_wb(ind)[0]["cobranza"])
    assert _NOTA_COBRANZA_NETO in pie and _NOTA_COBRANZA_NETO_FLOTA not in pie


# ---------------------------------------------------------------------------
# 6) Ruta: el JSON del API 0.0.66 entra sin 422.
# ---------------------------------------------------------------------------

TOKEN = "secreto-de-prueba"
client = TestClient(app)


def test_ruta_acepta_el_payload_del_api_0066(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    res = client.post(
        "/pdf/balance-general-xlsx",
        json=_payload_om_0066("general"),
        headers={"X-Internal-Token": TOKEN},
    )
    assert res.status_code == 200, res.text
    wb = load_workbook(BytesIO(res.content))
    assert [c.value for c in wb["balance"]["A"]].count(_FILA_COMISIONES) == 2
    assert "COMISIONES\nMXN" not in [c.value for c in wb[MAESTRA][2]]


# ---------------------------------------------------------------------------
# 7) API 0.0.74 (13-oct-2026, invariante 61): «comisión del vendedor SOLO como
#    gasto real». La fila del vuelo completado de la regla llega con el pago
#    real «a cargo del avión» y egreso $0.00; con la bandera cambian SOLO las
#    leyendas.
# ---------------------------------------------------------------------------

_FILA_A_CARGO = {
    "clave": "#401 · Cliente Uno",
    "avion_color": "#2563EB",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-09-03",
    "concepto_egreso": (
        "pago comisión vendedor (Pablo Canales) · gasto real · a cargo del avión "
        "XA-TST ($2,360.00 en su columna COMISIONES)"
    ),
    "egreso_mxn": 0.0,
    "fecha_egreso": "2026-09-20",
    "concepto_ingreso": "comisión vendedor (Pablo Canales)",
    "ingreso_mxn": 2360.0,
    "fecha_ingreso": "2026-09-03",
    "remanente_mxn": 2360.0,
}
# Pago real que NO carga ningún avión (vuelo aún no realizado): VuelaTour lo
# egresa con la rama de siempre.
_FILA_VT = {
    "clave": "#413 · Cliente Trece",
    "avion_color": "#16A34A",
    "estado": "CONFIRMADO",
    "fecha_vuelo": "2026-09-25",
    "concepto_egreso": "pago comisión vendedor (Alex Saab) · gasto real · excede $320.00 MXN",
    "egreso_mxn": 1500.0,
    "fecha_egreso": "2026-09-26",
    "concepto_ingreso": "comisión vendedor (Alex Saab)",
    "ingreso_mxn": 1180.0,
    "fecha_ingreso": "2026-09-25",
    "remanente_mxn": -320.0,
}
REGLA = REGLA_COMISIONES_RESPALDO


def _payload_0074(
    variante: str = "mensual", *, hoja: bool | None = True, raiz: bool | None = True
) -> dict:
    """El general como lo manda el API 0.0.74: filas de los libros con la
    forma «pago real» (`tc._forma_0074`), `comisiones_usd` del 0.0.66, la
    bandera raíz en cada `aviones[]` y en el consolidado (`raiz`) y en la
    hoja 'otros movimientos' (`hoja`) junto a `sin_provision_vendedor`
    (None ⇒ la clave no viaja)."""
    libros = [
        _con_api_0066(tc._forma_0074(x, bandera=raiz)) for x in (tc._libro_tst(), tc._libro_dos())
    ]
    p = {**tc._general_dict(libros), "variante": variante}
    om = p["consolidado"]["otros_movimientos"]
    om["filas"] = [_FILA_A_CARGO, _FILA_VT]
    om["sin_provision_vendedor"] = True
    if hoja is not None:
        om["comision_vendedor_solo_gasto_real"] = hoja
    if raiz is not None:
        p["consolidado"]["comision_vendedor_solo_gasto_real"] = raiz
    return p


def test_textos_gasto_real_son_los_del_contrato_y_sin_provision() -> None:
    """Literales del contrato §G del ajuste de las dos categorías (13-oct-2026),
    que reescribió EN SU LUGAR los de la decisión 3."""
    assert _OM_REGLA_COMISIONES_GASTO_REAL == (
        " Con la {regla} (por fecha del vuelo), la comisión bancaria de aquí es solo la "
        "parte de VuelaTour: la de cada avión resta en la columna COMISIONES de su libro "
        "individual y en la cascada de la hoja 'balance', igual que los pagos «Comisión "
        "del vendedor (avión)» que aquí salen con egreso $0.00."
    )
    assert _EMPRESA_REGLA_COMISIONES_GASTO_REAL == (
        " Con la {regla} (por fecha del vuelo), de la comisión bancaria los egresos "
        "propios solo traen la parte de VuelaTour; los pagos «Comisión del vendedor "
        "(avión)» los carga cada avión en su columna COMISIONES (aquí egresan $0.00)."
    )
    assert _RESUMEN_REGLA_GASTO_REAL == (
        " Con la {regla} (por fecha del vuelo), la comisión bancaria de 'otros "
        "movimientos' es solo la parte de VuelaTour; la del avión resta en su columna "
        "COMISIONES, igual que los pagos «Comisión del vendedor (avión)»."
    )
    # Texto corto desde el 26-oct-2026 (notas de celda: una idea, sin fecha).
    assert _NOTA_CASCADA_COMISIONES_GASTO_REAL == (
        "Comisiones a cargo del avión en dólares (su parte de la comisión bancaria de sus "
        "cobros y el pago «Comisión del vendedor (avión)» ligado al vuelo): las de la "
        "columna COMISIONES de su Balance por avión."
    )
    # La nota de la cascada crece y no llega al tope de alto (no se corta).
    assert _comentario_cobro(_NOTA_CASCADA_COMISIONES_GASTO_REAL).height < 260
    for texto in (
        _OM_REGLA_COMISIONES_GASTO_REAL,
        _EMPRESA_REGLA_COMISIONES_GASTO_REAL,
        _RESUMEN_REGLA_GASTO_REAL,
        _RESUMEN_COMISIONES_GASTO_REAL,
        _NOTA_CASCADA_COMISIONES_GASTO_REAL,
        _COMISIONES_FUERA_DE_LA_HOJA_GASTO_REAL,
        _NOTA_ENCABEZADO_GANANCIA_FLOTA_GASTO_REAL,
        _NOTA_ENCABEZADO_REMANENTE_GASTO_REAL,
        _NOTA_COMISIONES_FLOTA_GASTO_REAL,
    ):
        assert "provisi" not in texto.lower(), texto


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_general_0074_leyendas_y_fila_a_cargo_del_avion(variante) -> None:
    data = _render(_payload_0074(variante))
    assert verificar_libro(data)
    wb, wv = _wb(data)
    om, omv = wb["otros movimientos"], wv["otros movimientos"]
    a2 = om["A2"].value
    assert (
        _OM_PAGO_VENDEDOR_DOS_CATEGORIAS + _OM_REGLA_COMISIONES_GASTO_REAL.format(regla=REGLA)
        in a2
    )
    assert _OM_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION.format(regla=REGLA) not in a2
    assert "cubierto por el avión" not in a2
    # La fila «a cargo del avión»: egreso 0.0 SIN amarillo, remanente =
    # ingreso con la fórmula viva; la de VuelaTour, como siempre.
    assert om["F4"].value == _FILA_A_CARGO["concepto_egreso"]
    assert (om["G4"].value, om["J4"].value) == (0.0, 2360.0)
    assert om["L4"].value == "=ROUND(J4-G4,2)" and omv["L4"].value == 2360.0
    assert (om["G5"].value, omv["L5"].value) == (1500.0, -320.0)
    rellenos = {
        om.cell(row=f, column=c).fill.fgColor.rgb
        for f in (4, 5)
        for c in range(1, 11)
        if om.cell(row=f, column=c).fill.fill_type
    }
    assert "00" + AMBER not in rellenos
    assert _LEYENDA_PROVISION not in [c.value for c in om["A"]]
    tot = next(c.row for c in om["A"] if c.value == "TOTALES")
    assert omv[f"G{tot}"].value == 1500.0
    assert omv[f"L{tot}"].value == r2(omv[f"J{tot}"].value - 1500.0)
    # RESUMEN.
    resumen = next(
        c.value
        for c in wb["RESUMEN flota"]["A"]
        if isinstance(c.value, str) and c.value.startswith("COMISIONES = ")
    )
    assert resumen == (
        _RESUMEN_COMISIONES_GASTO_REAL
        + _RESUMEN_PAGO_VENDEDOR_DOS_CATEGORIAS
        + _RESUMEN_REGLA_GASTO_REAL.format(regla=REGLA)
    )
    # Bloque VUELATOUR (empresa): base de las dos categorías + regla nueva.
    pie = _pie(wb["balance"])
    assert (
        _NOTA_EMPRESA_BASE_DOS_CATEGORIAS
        + _EMPRESA_REGLA_COMISIONES_GASTO_REAL.format(regla=REGLA)
        in pie
    )
    assert not any(
        _EMPRESA_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION.format(regla=REGLA) in x for x in pie
    )
    # La cascada de CADA avión: la nota de «(−) COMISIONES» habla del pago real.
    ws_b = wb["balance"]
    filas_com = [c.row for c in ws_b["A"] if c.value == _FILA_COMISIONES]
    assert len(filas_com) == 2
    for fila in filas_com:
        assert ws_b.cell(row=fila, column=2).comment.text == _NOTA_CASCADA_COMISIONES_GASTO_REAL
    # Hoja de vuelos del general (bandera del consolidado): notas de
    # encabezado y del pie con el pago real.
    gm = wb[MAESTRA]
    llave = "GANANCIA\nMXN" if variante == "mensual" else "REMANENTE VENTA\nMENOS COMPRA (PESOS)"
    nota_esperada = (
        _NOTA_ENCABEZADO_GANANCIA_FLOTA_GASTO_REAL
        if variante == "mensual"
        else _NOTA_ENCABEZADO_REMANENTE_GASTO_REAL
    )
    celda = next(c for c in gm[2] if c.value == llave)
    assert celda.comment.text == nota_esperada
    cola = _NOTA_GANANCIA_FLOTA if variante == "mensual" else _CPH_NOTA_REMANENTE
    assert _NOTA_COMISIONES_FLOTA_GASTO_REAL + cola in _pie(gm)
    assert _NOTA_COMISIONES_FLOTA + cola not in _pie(gm)


@pytest.mark.parametrize("valor", [None, False], ids=["ausente", "false"])
def test_general_0074_sin_la_bandera_leyendas_de_la_0073(valor) -> None:
    """Sin la bandera (en la hoja ni en la raíz) las leyendas de hoy: las de
    la 0.0.73 (`sin_provision_vendedor`) y las notas de siempre."""
    p = _payload_0074(hoja=valor, raiz=valor)
    wb = _wb(_render(p))[0]
    a2 = wb["otros movimientos"]["A2"].value
    assert _OM_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION.format(regla=REGLA) in a2
    assert _OM_REGLA_COMISIONES_GASTO_REAL.format(regla=REGLA) not in a2
    resumen = next(
        c.value
        for c in wb["RESUMEN flota"]["A"]
        if isinstance(c.value, str) and c.value.startswith("COMISIONES = ")
    )
    assert resumen.startswith(_RESUMEN_COMISIONES)
    ws_b = wb["balance"]
    for fila in (c.row for c in ws_b["A"] if c.value == _FILA_COMISIONES):
        assert ws_b.cell(row=fila, column=2).comment.text == _NOTA_CASCADA_COMISIONES
    gm = wb[MAESTRA]
    assert next(c for c in gm[2] if c.value == "GANANCIA\nMXN").comment.text == (
        _NOTA_ENCABEZADO_GANANCIA_FLOTA
    )
    assert _NOTA_COMISIONES_FLOTA + _NOTA_GANANCIA_FLOTA in _pie(gm)


def test_bandera_de_la_hoja_y_bandera_raiz_son_independientes() -> None:
    """La de la hoja 'otros movimientos' cambia la pestaña, el RESUMEN y el
    bloque VUELATOUR; la raíz del consolidado, las notas de su hoja de
    vuelos; la de cada avión, su cascada. Ninguna toca a las demás (nunca se
    reutiliza `sin_provision_vendedor` como señal del general)."""
    solo_hoja = _wb(_render(_payload_0074(hoja=True, raiz=None)))[0]
    assert _OM_REGLA_COMISIONES_GASTO_REAL.format(regla=REGLA) in (
        solo_hoja["otros movimientos"]["A2"].value
    )
    assert _NOTA_COMISIONES_FLOTA + _NOTA_GANANCIA_FLOTA in _pie(solo_hoja[MAESTRA])
    for fila in (c.row for c in solo_hoja["balance"]["A"] if c.value == _FILA_COMISIONES):
        assert solo_hoja["balance"].cell(row=fila, column=2).comment.text == (
            _NOTA_CASCADA_COMISIONES
        )
    solo_raiz = _wb(_render(_payload_0074(hoja=None, raiz=True)))[0]
    assert _OM_REGLA_COMISIONES_GASTO_REAL.format(regla=REGLA) not in (
        solo_raiz["otros movimientos"]["A2"].value
    )
    assert _NOTA_COMISIONES_FLOTA_GASTO_REAL + _NOTA_GANANCIA_FLOTA in _pie(solo_raiz[MAESTRA])
    for fila in (c.row for c in solo_raiz["balance"]["A"] if c.value == _FILA_COMISIONES):
        assert solo_raiz["balance"].cell(row=fila, column=2).comment.text == (
            _NOTA_CASCADA_COMISIONES_GASTO_REAL
        )


def test_bloque_empresa_0074_sin_nota_del_api_pie_fijo() -> None:
    p = _payload_0074()
    p["empresa"]["nota"] = None
    pie = _pie(_wb(_render(p))[0]["balance"])
    assert any(
        _NOTA_EMPRESA_BASE_DOS_CATEGORIAS
        + _EMPRESA_REGLA_COMISIONES_GASTO_REAL.format(regla=REGLA)
        in x
        for x in pie
    )
    assert not any("o provisionado" in x for x in pie)


# Las tres filas del pago al vendedor de la gramática única del API 0.0.74
# (dos categorías, `egresoPagoVendedorDeVuelo`), como llegan: mixta (la
# «(VuelaTour)» egresa y la «(avión)» la carga el avión), informativa $0.00
# (`_FILA_A_CARGO`) y la «(avión)» que ningún avión carga (la paga VuelaTour).
_FILA_MIXTA = {
    "clave": "#402 · Cliente Dos",
    "avion_color": "#16A34A",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-09-08",
    "concepto_egreso": (
        "pago comisión vendedor (Pablo Canales) · gasto real (2 pagos) · VuelaTour "
        "$2,000.00 · a cargo del avión XB-DOS ($1,000.00 en su columna COMISIONES)"
    ),
    "egreso_mxn": 2000.0,
    "fecha_egreso": "2026-09-22",
    "concepto_ingreso": "comisión vendedor (Pablo Canales)",
    "ingreso_mxn": 3000.0,
    "fecha_ingreso": "2026-09-08",
    "remanente_mxn": 1000.0,
}
_FILA_NINGUNO = {
    "clave": "#414 · Cliente Catorce",
    "avion_color": "#2563EB",
    "estado": "CONFIRMADO",
    "fecha_vuelo": "2026-09-28",
    "concepto_egreso": (
        "pago comisión vendedor (Alex Saab) · gasto real · incluye «Comisión del vendedor "
        "(avión)» que ningún avión carga (el vuelo aún no se completa: al completarse lo "
        "cargará el avión): la paga VuelaTour"
    ),
    "egreso_mxn": 1180.0,
    "fecha_egreso": "2026-09-27",
    "concepto_ingreso": "comisión vendedor (Alex Saab)",
    "ingreso_mxn": 1180.0,
    "fecha_ingreso": "2026-09-28",
    "remanente_mxn": 0.0,
}


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_general_0074_dos_categorias_filas_mixta_informativa_y_sin_avion(variante) -> None:
    """La pestaña pinta TAL CUAL las tres filas (más la de VuelaTour de
    siempre): sin amarillo, remanente con la fórmula viva y TOTALES = Σ filas
    (la informativa suma $0.00); la fila 2 explica las dos categorías."""
    p = _payload_0074(variante)
    filas = [_FILA_MIXTA, _FILA_A_CARGO, _FILA_NINGUNO, _FILA_VT]
    p["consolidado"]["otros_movimientos"]["filas"] = filas
    data = _render(p)
    assert verificar_libro(data)
    wb, wv = _wb(data)
    om, omv = wb["otros movimientos"], wv["otros movimientos"]
    assert _OM_PAGO_VENDEDOR_DOS_CATEGORIAS in om["A2"].value
    for k, fila in enumerate(filas):
        r = 4 + k
        assert om[f"A{r}"].value.startswith(fila["clave"]), r
        assert om[f"F{r}"].value == fila["concepto_egreso"], r
        assert (om[f"G{r}"].value, om[f"J{r}"].value) == (fila["egreso_mxn"], fila["ingreso_mxn"])
        assert om[f"L{r}"].value == f"=ROUND(J{r}-G{r},2)", r
        assert omv[f"L{r}"].value == fila["remanente_mxn"], r
    # La informativa: egreso 0.0 (no vacío) y remanente = ingreso.
    assert (om["G5"].value, omv["L5"].value) == (0.0, 2360.0)
    # La que ningún avión carga: remanente 0 (nunca −0).
    assert omv["L6"].value == 0 and not str(omv["L6"].value).startswith("-")
    rellenos = {
        om.cell(row=f, column=c).fill.fgColor.rgb
        for f in range(4, 4 + len(filas))
        for c in range(1, 11)
        if om.cell(row=f, column=c).fill.fill_type
    }
    assert "00" + AMBER not in rellenos
    tot = next(c.row for c in om["A"] if c.value == "TOTALES")
    todas = filas + p["consolidado"]["otros_movimientos"].get("filas_sueltas", [])
    egresos = r2(sum(f.get("egreso_mxn") or 0.0 for f in todas))
    ingresos = r2(sum(f.get("ingreso_mxn") or 0.0 for f in todas))
    assert omv[f"G{tot}"].value == egresos == 4680.0
    assert omv[f"J{tot}"].value == ingresos
    assert omv[f"L{tot}"].value == r2(ingresos - egresos)


def test_general_0074_con_la_bandera_nadie_nombra_la_categoria_vieja() -> None:
    """Con la bandera ninguna leyenda pintada (pestaña, RESUMEN, hoja de
    vuelos, cascada y bloque VUELATOUR) nombra «Comisión del vendedor» sin
    su categoría ni habla de una provisión del vendedor (la de AFAC, «PERMISO
    AFAC (PROVISIÓN)», es otra cosa): las leyendas nombran las dos."""
    wb = _wb(_render(_payload_0074("mensual")))[0]
    textos = [
        c.value
        for hoja in ("otros movimientos", "RESUMEN flota", MAESTRA, "balance")
        for fila in wb[hoja].iter_rows()
        for c in fila
        if isinstance(c.value, str)
    ]
    notas = [
        c.comment.text
        for hoja in (MAESTRA, "balance")
        for fila in wb[hoja].iter_rows()
        for c in fila
        if c.comment is not None
    ]
    assert not any("«Comisión del vendedor»" in t for t in textos + notas)
    assert not any("provisión" in t.lower() and "vendedor" in t.lower() for t in textos)
    assert any("«Comisión del vendedor (VuelaTour)»" in t for t in textos)
    assert any("«Comisión del vendedor (avión)»" in t for t in textos)


def test_rutas_aceptan_el_payload_del_api_0074(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    p = _payload_0074("general")
    res = client.post("/pdf/balance-general-xlsx", json=p, headers=headers)
    assert res.status_code == 200, res.text
    a2 = load_workbook(BytesIO(res.content))["otros movimientos"]["A2"].value
    assert _OM_REGLA_COMISIONES_GASTO_REAL.format(regla=REGLA) in a2
    res = client.post("/pdf/balance-general-libro", json=p, headers=headers)
    assert res.status_code == 200, res.text
    res = client.post(
        "/pdf/balance-avion-xlsx", json=tc._forma_0074(tc._libro_tst()), headers=headers
    )
    assert res.status_code == 200, res.text
