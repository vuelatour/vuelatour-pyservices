"""Balance GENERAL con costo por hora y PROVISIÓN en amarillo (6-oct-2026,
API 0.0.64).

Pedido del cliente: «el que ya tenemos que se renombre a "Balance mensual" y
el nuevo botón sea el "Balance general"… el que cambiaría sería el balance
general: todos esos montos de operación, piloto, pagos AFAC se van a
eliminar y se va a hacer un resumen del total de COSTO TOTAL… se van a
agregar estas columnas para sacar el costo por hora por cada vuelo», con la
fórmula «es el total de todos los gastos, entre el tiempo volado, entre el
tipo de cambio del día, entre 1.16 (para sacar subtotal)». Y aparte: «los que
están provisión podrían estar en amarillo».

`BalanceGeneralRequest.variante`: 'mensual' (o ausente) = el libro de
siempre, BYTE-IDÉNTICO; 'general' = la hoja «reporte horas FLOTA» con el
juego de columnas `_COLS_GENERAL`. Mismos números del API en las dos: las
fórmulas se verifican con el evaluador independiente contra el número que
manda el API (payloads del espejo de `test_balance_formulas`).

Revisión del 6-oct-2026: el libro de hoy se congela también por BYTES (la
huella de layout no veía notas, bordes ni fuentes) y la regla «jamás letras
fijas» se prueba con OTROS juegos de columnas y el libro completo (sin eso,
una fórmula con letras fijas pasaba toda la suite).
"""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from io import BytesIO

import openpyxl
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAvionRequest,
    BalanceAvionVuelo,
    BalanceGeneralRequest,
    BalanceOtroMovimientoFila,
)
from app.services import balance_avion_xlsx, xlsx_formulas
from app.services.balance_avion_xlsx import (
    _COLS,
    _COLS_GENERAL,
    _CPH_NOTA_REPARTIDOS_A_AVIONES,
    _DISP_GENERAL,
    _DISP_MENSUAL,
    _GRUPO_COSTO_HORA,
    _LEYENDA_PROVISION,
    _NOTA_CALCULADO_POR_SISTEMA_CELDA,
    _NOTA_REPARTIDOS_A_AVIONES,
    _TITULO_GRUPO_COSTO_HORA,
    _TOTAL_MAP,
    _TOTAL_MAP_GENERAL,
    AMBER,
    HORAS,
    MONEY,
    TC,
    TC_OFICIAL_FILL,
    _disposicion,
    _es_provision,
    _formula_cita,
    _hoja_maestra,
    _Maestra,
    _nota_desglose_costo,
    _valor_columna,
    render_balance_general_xlsx,
)
from tests import test_balance_empresa_vuelatour as t_empresa
from tests import test_balance_extension_horario as t_extension
from tests import test_balance_formulas as tf
from tests import test_comision_vendedor_notas as t_comision
from tests._evaluador_formulas import verificar_libro

MAESTRA = "reporte horas FLOTA"

# Las 12 columnas de la hoja «utilidades» del cliente, EN SU ORDEN.
COSTO_POR_HORA = [
    "TOTAL COBRADO\nS/IVA (PESOS)",
    "TIEMPO\nCALZOS/HOBS (HR)",
    "COSTO X HORA\n(DLLS S/IVA)",
    "IVA X HR\n(DLLS)",
    "COSTO HR\nMÁS IVA (DLLS)",
    "TOTAL PARA\nPROVEEDOR (DLLS)",
    "IVA TOTAL\nPAGADO (DLLS)",
    "TIPO\nCAMBIO",
    "TOTAL PARA\nPROVEEDOR (PESOS)",
    "IVA TOTAL\nPAGADO (PESOS)",
    "TOTAL PAGADO\nS/IVA (PESOS)",
    "REMANENTE VENTA\nMENOS COMPRA (PESOS)",
]
# Columnas del Balance mensual que el general ya no lleva.
FUERA = [
    "OPERACIONES",
    "PILOTO",
    "OTROS",
    "PERMISO AFAC\n(PROVISIÓN)",
    "COSTO TOTAL\nUSD",
    "COSTO TOTAL\nUSD S/IVA",
    "IVA PAGADO\nUSD",
    "IVA PAGADO\nMXN",
    "REMANENTE\nVENTA−COSTO MXN",
    "DIF. IVA\nHACIENDA MXN",
    "COMISIÓN\nVENDEDOR MXN",
    "GANANCIA\nMXN",
    "GANANCIA\nUSD",
    "COSTO X HORA\nUSD",
    "COSTO X HORA\nUSD S/IVA",
]


# ---------------------------------------------------------------------------
# Ayudas
# ---------------------------------------------------------------------------


def _con_variante(req: BalanceGeneralRequest, variante) -> BalanceGeneralRequest:
    datos = req.model_dump()
    datos["variante"] = variante
    return BalanceGeneralRequest.model_validate(datos)


def _general(variante="general") -> BalanceGeneralRequest:
    return _con_variante(tf._general(), variante)


def _libros(data: bytes):
    """(con fórmulas, con caché)."""
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _col(ws, encabezado: str) -> int:
    return next(c.column for c in ws[2] if c.value == encabezado)


def _letra(ws, encabezado: str) -> str:
    return get_column_letter(_col(ws, encabezado))


def _celda(ws, fila: int, encabezado: str):
    return ws.cell(row=fila, column=_col(ws, encabezado))


def _fila_de(ws, clave: str) -> int:
    return next(c.row for c in ws["A"] if c.value == clave)


def _relleno(celda) -> str | None:
    return celda.fill.fgColor.rgb if celda.fill.fill_type else None


def _formulas_del_bloque(ws, f: int) -> dict[str, str]:
    """Fórmulas de la fila `f` del bloque COSTO POR HORA con la letra que
    tiene cada columna EN ESTA HOJA (se busca por su encabezado)."""
    k, l_ = _letra(ws, "VENTA AVIÓN\nMXN"), _letra(ws, "IVA VENTA\nAVIÓN MXN")
    t, z = _letra(ws, COSTO_POR_HORA[1]), _letra(ws, COSTO_POR_HORA[7])
    y, ae = _letra(ws, COSTO_POR_HORA[8]), _letra(ws, COSTO_POR_HORA[5])
    an, ao = _letra(ws, COSTO_POR_HORA[4]), _letra(ws, COSTO_POR_HORA[2])
    ag, ah = _letra(ws, COSTO_POR_HORA[6]), _letra(ws, COSTO_POR_HORA[9])
    return {
        "TOTAL COBRADO\nS/IVA (PESOS)": f"=ROUND({k}{f}-{l_}{f},2)",
        # «el total de todos los gastos, entre el tiempo volado, entre el
        # tipo de cambio del día, entre 1.16» — en el orden del API.
        "COSTO X HORA\n(DLLS S/IVA)": f"={y}{f}/{z}{f}/{t}{f}/$D$1",
        "IVA X HR\n(DLLS)": f"=ROUND(ROUND({an}{f},2)-ROUND({ao}{f},2),2)",
        "COSTO HR\nMÁS IVA (DLLS)": f"={ae}{f}/{t}{f}",
        "TOTAL PARA\nPROVEEDOR (DLLS)": f"={y}{f}/{z}{f}",
        "IVA TOTAL\nPAGADO (DLLS)": f"={ae}{f}-{ae}{f}/$D$1",
        "IVA TOTAL\nPAGADO (PESOS)": f"={ag}{f}*{z}{f}",
        "TOTAL PAGADO\nS/IVA (PESOS)": f"={y}{f}-{ah}{f}",
        "REMANENTE VENTA\nMENOS COMPRA (PESOS)": f"=ROUND({k}{f}-{y}{f},2)",
    }


def _miembros(data: bytes) -> dict[str, bytes]:
    """Todos los miembros del .xlsx salvo docProps/core.xml (lleva la hora)."""
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _firma_libro(data: bytes, *, sin_hojas: tuple[str, ...] = ()) -> str:
    """Huella del LAYOUT de todas las hojas (receta de test_balance_cobrado_con):
    valores o fórmulas, formatos, fuentes, rellenos, alineación, bordes, notas,
    merges, anchos, altos y panel fijo. `sin_hojas` deja fuera esas pestañas."""
    wb = load_workbook(BytesIO(data))
    partes: list = []
    for ws in wb.worksheets:
        if ws.title in sin_hojas:
            continue
        partes.append(ws.title)
        for fila in ws.iter_rows():
            for c in fila:
                f = c.font
                partes.append([
                    c.coordinate, repr(c.value), c.number_format,
                    bool(f.b), bool(f.i), f.sz,
                    f.color.rgb if f.color is not None else None,
                    c.fill.fgColor.rgb if c.fill.fill_type else None,
                    c.alignment.horizontal, c.alignment.vertical,
                    bool(c.alignment.wrap_text), c.border.left.style,
                    c.comment.text if c.comment else None,
                ])  # fmt: skip
        partes.append(sorted(str(r) for r in ws.merged_cells.ranges))
        partes.append(sorted((k, d.width) for k, d in ws.column_dimensions.items()))
        partes.append(sorted((k, d.height) for k, d in ws.row_dimensions.items() if d.height))
        partes.append(ws.freeze_panes)
    return hashlib.sha256(json.dumps(partes, default=str).encode()).hexdigest()


@pytest.fixture
def registros(monkeypatch) -> list[xlsx_formulas.RegistroFormulas]:
    """Registro de fórmulas de cada libro que se cierra (qué celdas se
    escribieron como fórmula y cuáles volvieron a valor)."""
    capturados: list[xlsx_formulas.RegistroFormulas] = []
    original = xlsx_formulas.finalizar

    def espia(wb):
        cache = original(wb)
        capturados.append(xlsx_formulas.registro(wb))
        return cache

    monkeypatch.setattr(xlsx_formulas, "finalizar", espia)
    return capturados


# ---------------------------------------------------------------------------
# 1) Esquema: campo ADITIVO y liberal.
# ---------------------------------------------------------------------------


def test_variante_por_omision_es_el_balance_mensual() -> None:
    assert BalanceGeneralRequest().variante == "mensual"
    assert tf._general().variante == "mensual"  # payload de un API ≤ 0.0.63


@pytest.mark.parametrize(
    "entrada,esperada",
    [
        ("general", "general"),
        (" GENERAL ", "general"),
        ("Mensual", "mensual"),
        (None, "mensual"),
        ("costo_hora", "mensual"),
        (1, "mensual"),
        ([], "mensual"),
    ],
)
def test_variante_liberal_nunca_es_un_422(entrada, esperada) -> None:
    assert BalanceGeneralRequest.model_validate({"variante": entrada}).variante == esperada


# ---------------------------------------------------------------------------
# 2) Juegos de columnas (por llave, jamás letras fijas).
# ---------------------------------------------------------------------------


def test_juego_de_columnas_del_balance_general() -> None:
    encabezados = [c[1] for c in _COLS_GENERAL]
    mensual = [c[1] for c in _COLS]
    assert encabezados == [
        *mensual[: mensual.index("TACO\nFINAL") + 1],
        "COSTO TOTAL\nMXN",
        "TIPO CAMBIO\nCOSTOS",
        *COSTO_POR_HORA,
        *mensual[mensual.index("STATUS") :],
    ]
    assert len(encabezados) == 44
    assert not set(FUERA) & set(encabezados)
    grupos = [c[0] for c in _COLS_GENERAL]
    assert grupos[16:18] == ["COSTO TOTAL (MXN)"] * 2
    assert grupos[18:30] == ["COSTO POR HORA"] * 12
    claves = [c[2] for c in _COLS_GENERAL if c[2]]
    assert len(claves) == len(set(claves))
    assert _DISP_GENERAL.cols == tuple(_COLS_GENERAL)
    assert _DISP_GENERAL.costo_por_hora and _DISP_GENERAL.comision_col is None


# Letras del juego de siempre (Balance mensual y libro individual), ESCRITAS A
# MANO: las del libro de hoy. Ya no hay mapas de letras a nivel de módulo
# (`_LETRA`, `_COBRO1_COL`… se retiraron en la revisión del 6-oct-2026: una
# fórmula nueva escrita con ellos apuntaría, en silencio, a las letras del
# mensual dentro del Balance general); todo sale de la disposición.
_LETRAS_MENSUAL = {
    "clave": "A", "fecha": "B", "ruta": "C", "estado": "D",
    "horas_cobradas": "E", "tarifa_usd": "F", "iva_hr_usd": "G", "total_usd": "H",
    "iva_usd": "I", "tc_venta": "J", "total_mxn": "K", "iva_mxn": "L", "subtotal_mxn": "M",
    "tiempo_vuelo": "N", "taco_inicio": "O", "taco_fin": "P",
    "op_mxn": "Q", "piloto_mxn": "R", "otros_mxn": "S", "permiso_afac_mxn": "T",
    "costo_total_mxn": "U", "tc_costos": "V",
    "costo_usd": "W", "costo_usd_siva": "X", "iva_pagado_usd": "Y", "iva_pagado_mxn": "Z",
    "remanente_mxn": "AA", "dif_iva_mxn": "AB", "comision_vendedor_mxn": "AC",
    "ganancia_mxn": "AD", "ganancia_usd": "AE",
    "costo_hr_usd": "AF", "costo_hr_usd_siva": "AG",
    "status_cobro": "AH",
    "cobrado_real_mxn": "AQ", "cobrado_mxn": "AR",
    "por_cobrar_mxn": "AS", "por_cobrar_usd": "AT", "factura_vuelatour": "AU",
}  # fmt: skip


def test_el_juego_de_siempre_no_se_movio() -> None:
    assert _DISP_MENSUAL.cols == tuple(_COLS)
    assert _DISP_MENSUAL.letra == _LETRAS_MENSUAL
    assert _DISP_MENSUAL.cobro1_col == 35  # COBRO 1 FECHA = AI
    assert _DISP_MENSUAL.cobro_mxn_letras == ("AJ", "AL", "AN", "AP")
    assert _DISP_MENSUAL.comision_col == 29  # COMISIÓN VENDEDOR MXN = AC
    assert _DISP_MENSUAL.total_map is _TOTAL_MAP
    assert not _DISP_MENSUAL.costo_por_hora


def test_valor_de_cada_columna_sale_del_api() -> None:
    v = BalanceAvionVuelo(
        subtotal_mxn=59999.5,
        tiempo_vuelo=2.4,
        tc_costos=17.32,
        costo_total_mxn=7133.75,
        costo_hr_usd=171.62,
        costo_hr_usd_siva=147.95,
        iva_pagado_mxn=983.97,
    )
    assert _valor_columna("cph_cobrado_siva_mxn", v) == 59999.5
    assert _valor_columna("cph_tiempo_hr", v) == 2.4
    assert _valor_columna("cph_tc", v) == 17.32
    assert _valor_columna("cph_proveedor_mxn", v) == 7133.75
    assert _valor_columna("cph_iva_hr_usd", v) == 23.67  # 171.62 − 147.95
    assert _valor_columna("cph_pagado_siva_mxn", v) == 6149.78  # 7133.75 − 983.97
    assert _valor_columna("costo_hr_usd_siva", v) == 147.95
    # Sin uno de los dos lados: vacío, nunca un 0 falso.
    assert _valor_columna("cph_iva_hr_usd", BalanceAvionVuelo(costo_hr_usd=1.0)) is None
    assert _valor_columna("cph_pagado_siva_mxn", BalanceAvionVuelo(costo_total_mxn=1.0)) is None


# ---------------------------------------------------------------------------
# 3) La hoja «reporte horas FLOTA» del Balance general.
# ---------------------------------------------------------------------------


def test_encabezado_dice_balance_general_y_conserva_las_constantes() -> None:
    wb, _ = _libros(render_balance_general_xlsx(_general()))
    ws = wb[MAESTRA]
    assert ws.max_column == 44
    assert [c.value for c in ws[2]] == [c[1] for c in _COLS_GENERAL]
    grupos = {c.value: c.coordinate for c in ws[1] if c.column > 4 and c.value}
    assert list(grupos) == [
        "VENTA",
        "TIEMPO / TACÓMETRO",
        "COSTO TOTAL (MXN)",
        _TITULO_GRUPO_COSTO_HORA,
        "STATUS DE COBROS",
    ]
    assert _TITULO_GRUPO_COSTO_HORA == "BALANCE GENERAL · COSTO POR HORA"
    inicio = _letra(ws, COSTO_POR_HORA[0])
    fin = _letra(ws, COSTO_POR_HORA[-1])
    assert f"{inicio}1:{fin}1" in {str(r) for r in ws.merged_cells.ranges}
    # A1:D1: las constantes de siempre (COSTO X HORA cita $D$1).
    assert (ws["A1"].value, ws["B1"].value) == ("Permiso AFAC USD/hr", None)
    assert (ws["C1"].value, ws["D1"].value) == ("Factor IVA costos", 1.16)
    assert "ya no lleva la columna PERMISO AFAC" in ws["B1"].comment.text
    assert "COSTO X HORA (DLLS S/IVA)" in ws["D1"].comment.text
    assert ws.freeze_panes == "D3"


def test_formulas_de_la_fila_citan_el_bloque_como_la_hoja_del_cliente() -> None:
    wb, _ = _libros(render_balance_general_xlsx(_general()))
    ws = wb[MAESTRA]
    f = _fila_de(ws, "#401 · Cliente Uno")
    for encabezado, formula in _formulas_del_bloque(ws, f).items():
        assert _celda(ws, f, encabezado).value == formula, encabezado
    # El layout congelado (si alguien mueve una columna, esto lo grita).
    assert ws[f"U{f}"].value == f"=AA{f}/Z{f}/T{f}/$D$1"
    assert ws[f"AD{f}"].value == f"=ROUND(K{f}-AA{f},2)"
    # Datos (no son aritmética de celdas del libro): valor.
    for encabezado in (
        "COSTO TOTAL\nMXN",
        "TIPO CAMBIO\nCOSTOS",
        "TIEMPO\nCALZOS/HOBS (HR)",
        "TIPO\nCAMBIO",
        "TOTAL PARA\nPROVEEDOR (PESOS)",
    ):
        valor = _celda(ws, f, encabezado).value
        assert isinstance(valor, (int, float)), encabezado


def test_cada_celda_es_el_numero_del_api_y_ninguna_formula_se_degrada(registros) -> None:
    req = _general()
    data = render_balance_general_xlsx(req)
    assert registros[-1].degradadas() == []
    assert verificar_libro(data)
    _, wv = _libros(data)
    ws = wv[MAESTRA]
    for k, v in enumerate(req.consolidado.vuelos):
        fila = 3 + k
        for i, (_g, encabezado, clave, fmt) in enumerate(_COLS_GENERAL, start=1):
            if clave is None or fmt is None:
                continue
            esperado = _valor_columna(clave, v)
            visto = ws.cell(row=fila, column=i).value
            if esperado is None:
                assert visto is None, (fila, encabezado)
            else:
                assert visto == pytest.approx(esperado, abs=1e-9), (fila, encabezado)
    # La fórmula del cliente con los datos de la fila #401 (orden del pedido:
    # gastos ÷ tiempo ÷ T.C. ÷ 1.16) da el COSTO X HORA del API.
    f = _fila_de(ws, "#401 · Cliente Uno")
    assert tf.r2(7133.75 / 2.4 / 17.32 / 1.16) == 147.95
    assert _celda(ws, f, "COSTO X HORA\n(DLLS S/IVA)").value == 147.95
    assert _celda(ws, f, "COSTO HR\nMÁS IVA (DLLS)").value == 171.62
    assert _celda(ws, f, "IVA X HR\n(DLLS)").value == 23.67


def test_iva_por_hora_es_la_resta_de_los_dos_numeros_a_la_vista() -> None:
    """#407: COSTO HR MÁS IVA = 127.95 (127.954…) y COSTO X HORA = 110.31
    (110.305…). Restando las celdas completas saldría 17.65 y la celda se
    quedaba como VALOR 17.64 (pasa en una de cada cuatro filas); con el
    redondeo donde el API redondea, la fórmula da 17.64 y se queda."""
    an = 7120 / 17.95 / 3.1
    ao = an / 1.16
    assert (tf.r2(an), tf.r2(ao)) == (127.95, 110.31)
    assert not xlsx_formulas.coincide(an - ao, 17.64, 2)
    wb, wv = _libros(render_balance_general_xlsx(_general()))
    ws = wb[MAESTRA]
    f = _fila_de(ws, "#407 · Cliente Cinco")
    w, u = _letra(ws, "COSTO HR\nMÁS IVA (DLLS)"), _letra(ws, "COSTO X HORA\n(DLLS S/IVA)")
    assert _celda(ws, f, "IVA X HR\n(DLLS)").value == f"=ROUND(ROUND({w}{f},2)-ROUND({u}{f},2),2)"
    assert _celda(wv[MAESTRA], f, "IVA X HR\n(DLLS)").value == 17.64


def test_fila_sin_tc_ni_horas_deja_vacio_lo_que_el_api_no_manda() -> None:
    wb, _ = _libros(render_balance_general_xlsx(_general()))
    ws = wb[MAESTRA]
    sin_tc = _fila_de(ws, "#403 · Traslado")
    for encabezado in COSTO_POR_HORA[2:8] + COSTO_POR_HORA[9:11]:
        assert _celda(ws, sin_tc, encabezado).value is None, encabezado
    assert _celda(ws, sin_tc, "TOTAL PARA\nPROVEEDOR (PESOS)").value == 800
    cancelado = _fila_de(ws, "#404 · Cliente Tres")  # sin horas voladas
    for encabezado in COSTO_POR_HORA[2:5]:
        assert _celda(ws, cancelado, encabezado).value is None, encabezado
    assert str(_celda(ws, cancelado, "TOTAL PARA\nPROVEEDOR (DLLS)").value).startswith("=")


def test_columnas_que_repiten_un_dato_heredan_sus_senales() -> None:
    """TOTAL COBRADO S/IVA (PESOS) repite TOTAL S/IVA MXN y TIEMPO
    CALZOS/HOBS (HR) repite TIEMPO VUELO HR —el divisor de COSTO X HORA: unas
    horas infladas lo bajan—; llevan el azul del T.C. oficial y el ámbar (con
    su nota) del salto de taco de su columna de origen (revisión 6-oct-2026:
    antes solo la columna de origen avisaba)."""
    req = _general()
    cons = req.consolidado
    marcado = cons.vuelos[0].model_copy(
        update={
            "tc_venta_oficial": True,
            "salto_taco_interno": True,
            "salto_taco_interno_detalle": "tramo 2 PTU-CUN (+0.4 h)",
        }
    )
    vuelos = [marcado, *cons.vuelos[1:]]
    req = req.model_copy(update={"consolidado": cons.model_copy(update={"vuelos": vuelos})})
    ws = _libros(render_balance_general_xlsx(req))[0][MAESTRA]
    azul, ambar = "00" + TC_OFICIAL_FILL, "00" + AMBER
    for encabezado in ("TOTAL S/IVA\nMXN", COSTO_POR_HORA[0]):
        assert _relleno(_celda(ws, 3, encabezado)) == azul, encabezado
    origen, repetida = _celda(ws, 3, "TIEMPO\nVUELO HR"), _celda(ws, 3, COSTO_POR_HORA[1])
    assert _relleno(origen) == _relleno(repetida) == ambar
    assert repetida.comment.text == origen.comment.text
    assert repetida.comment.text.endswith("tramo 2 PTU-CUN (+0.4 h)")
    # Sin señales en la fila, las dos columnas del bloque van sin relleno.
    for fila, v in enumerate(vuelos[1:], start=4):
        assert not (v.tc_venta_oficial or v.salto_taco_interno), v.clave
        for encabezado in COSTO_POR_HORA[:2]:
            assert _relleno(_celda(ws, fila, encabezado)) is None, (fila, encabezado)
            assert _celda(ws, fila, encabezado).comment is None, (fila, encabezado)


def test_desglose_del_costo_en_la_nota_de_total_para_proveedor() -> None:
    wb, _ = _libros(render_balance_general_xlsx(_general()))
    ws = wb[MAESTRA]
    nota = _celda(ws, _fila_de(ws, "#401 · Cliente Uno"), "TOTAL PARA\nPROVEEDOR (PESOS)").comment
    assert nota.text == (
        "Operación $4,500.50 · Piloto $1,200.00 · Otros $350.75 · Permiso AFAC $1,082.50"
    )
    assert nota.author == "VuelaTour"
    nota = _celda(ws, _fila_de(ws, "#403 · Traslado"), "TOTAL PARA\nPROVEEDOR (PESOS)").comment
    assert nota.text == "Operación $800.00"
    # Interno sin costos: sin nota.
    interno = _fila_de(ws, "#406 · Interno")
    assert _celda(ws, interno, "TOTAL PARA\nPROVEEDOR (PESOS)").comment is None
    # COSTO TOTAL va como valor «calculado por el sistema» (nota en su encabezado).
    encabezado = ws.cell(row=2, column=_col(ws, "COSTO TOTAL\nMXN"))
    assert encabezado.comment.text.startswith("Calculado por el sistema: COSTO TOTAL = operación")
    assert [c.coordinate for c in ws[2] if c.comment] == [encabezado.coordinate]


def test_nota_del_desglose_partes_y_detalle() -> None:
    v = BalanceAvionVuelo(
        op_mxn=4500.5,
        piloto_mxn=0.0,
        otros_mxn=350.75,
        op_detalle=["Comida · Starbucks — $206.00", "TUA $1,100.25**"],
        otros_detalle=["FBO · ASUR — $350.75"],
    )
    assert _nota_desglose_costo(v) == (
        "Operación $4,500.50 · Piloto $0.00 · Otros $350.75\n\n"
        "Operación:\nComida · Starbucks — $206.00\nTUA $1,100.25**\n\n"
        "Otros:\nFBO · ASUR — $350.75"
    )
    assert _nota_desglose_costo(BalanceAvionVuelo(op_mxn=-500.0)) == "Operación -$500.00"
    assert _nota_desglose_costo(BalanceAvionVuelo(op_detalle=["TUA $939.60**"])) == (
        "Operación:\nTUA $939.60**"
    )
    assert _nota_desglose_costo(BalanceAvionVuelo()) is None
    assert _nota_desglose_costo(BalanceAvionVuelo(op_mxn=0, piloto_mxn=0, otros_mxn=0)) is None


def test_totales_solo_los_que_manda_el_api() -> None:
    req = _general()
    wb, wv = _libros(render_balance_general_xlsx(req))
    ws, wsv = wb[MAESTRA], wv[MAESTRA]
    tot = _fila_de(ws, "TOTALES")
    t = req.consolidado.totales
    sumas = {
        "TOTAL COBRADO\nS/IVA (PESOS)": t.subtotal_mxn,
        "TIEMPO\nCALZOS/HOBS (HR)": t.tiempo_vuelo,
        "TOTAL PARA\nPROVEEDOR (PESOS)": t.costo_total_mxn,
        "COSTO TOTAL\nMXN": t.costo_total_mxn,
        "REMANENTE VENTA\nMENOS COMPRA (PESOS)": t.remanente_mxn,
        "POR COBRAR\nUSD": t.por_cobrar_usd,
    }
    for encabezado, total in sumas.items():
        letra = _letra(ws, encabezado)
        assert _celda(ws, tot, encabezado).value == f"=ROUND(SUM({letra}3:{letra}{tot - 1}),2)"
        assert _celda(wsv, tot, encabezado).value == total, encabezado
    # Promedios del consolidado: valor (promedio de los promedios de cada libro).
    assert _celda(ws, tot, "TIPO CAMBIO\nCOSTOS").value == t.tc_promedio
    assert _celda(ws, tot, "TIPO\nCAMBIO").value == t.tc_promedio
    assert _celda(ws, tot, "COSTO HR\nMÁS IVA (DLLS)").value == t.costo_hr_prom_usd
    # Lo que el API no totaliza no se inventa.
    for encabezado in (
        "COSTO X HORA\n(DLLS S/IVA)",
        "IVA X HR\n(DLLS)",
        "TOTAL PARA\nPROVEEDOR (DLLS)",
        "IVA TOTAL\nPAGADO (DLLS)",
        "IVA TOTAL\nPAGADO (PESOS)",
        "TOTAL PAGADO\nS/IVA (PESOS)",
    ):
        assert _celda(ws, tot, encabezado).value is None, encabezado


def test_notas_al_pie_hablan_de_las_columnas_del_balance_general() -> None:
    ws = _libros(render_balance_general_xlsx(_con_variante(t_extension._general(), "general")))[0][
        MAESTRA
    ]
    textos = [c.value for c in ws["A"] if isinstance(c.value, str)]
    pie = "\n".join(textos)
    assert any(x.startswith("BALANCE GENERAL · COSTO POR HORA: COSTO X HORA") for x in textos)
    assert "entre 1.16 para sacar el subtotal" in pie
    assert "REMANENTE VENTA MENOS COMPRA = VENTA AVIÓN MXN − TOTAL PARA PROVEEDOR" in pie
    assert "quedan solo como nota en el desglose de TOTAL PARA PROVEEDOR (PESOS)" in pie
    assert (
        "TUA pagado del periodo (solo nota en el desglose de TOTAL PARA PROVEEDOR, no resta en "
        "este libro):"
    ) in textos
    # Ninguna nota nombra columnas que esta hoja ya no tiene.
    assert not any(x.startswith("COMISIÓN VENDEDOR MXN va vacía") for x in textos)
    assert "celda OPERACIONES" not in pie and "GANANCIA de la fila" not in pie
    assert "COSTO X HORA USD de la fila TOTALES" not in pie
    # El Balance mensual conserva las suyas.
    ws_m = _libros(render_balance_general_xlsx(t_extension._general()))[0][MAESTRA]
    assert any(
        isinstance(c.value, str) and c.value.startswith("COMISIÓN VENDEDOR MXN va vacía")
        for c in ws_m["A"]
    )


# ---------------------------------------------------------------------------
# 4) Las demás hojas cuadran igual en las dos variantes.
# ---------------------------------------------------------------------------

_CASOS_GENERAL = {
    "fórmulas": tf._general,
    "empresa": lambda: BalanceGeneralRequest.model_validate(t_empresa._payload()),
    "extensión": lambda: t_extension._general(con_llave=True),
    "comisión": lambda: t_comision._general(True),
}


@pytest.mark.parametrize("caso", list(_CASOS_GENERAL))
def test_las_demas_hojas_tienen_los_mismos_numeros_en_las_dos_variantes(caso) -> None:
    req = _CASOS_GENERAL[caso]()
    mensual = render_balance_general_xlsx(req)
    general = render_balance_general_xlsx(_con_variante(req, "general"))
    verificar_libro(general)
    (_, vm), (fg, vg) = _libros(mensual), _libros(general)
    assert vm.sheetnames == vg.sheetnames
    for hoja in vm.sheetnames:
        if hoja == MAESTRA:
            continue
        a = {c.coordinate: c.value for fila in vm[hoja].iter_rows() for c in fila}
        b = {c.coordinate: c.value for fila in vg[hoja].iter_rows() for c in fila}
        if hoja == "repartidos a aviones":
            # Texto, no número: su nota nombra la columna de cada variante.
            b = {
                k: _NOTA_REPARTIDOS_A_AVIONES if x == _CPH_NOTA_REPARTIDOS_A_AVIONES else x
                for k, x in b.items()
            }
        assert a == b, hoja
    # Las citas a la hoja maestra apuntan a SU columna en la variante (en los
    # payloads con vuelos y T.C. promedio; sin ellos las celdas van vacías).
    t = req.consolidado.totales
    if req.consolidado.vuelos and t.tc_promedio is not None and t.tiempo_vuelo:
        maestra = fg[MAESTRA]
        tot = _fila_de(maestra, "TOTALES")
        tc = f"'{MAESTRA}'!${_letra(maestra, 'TIPO CAMBIO' + chr(10) + 'COSTOS')}${tot}"
        horas = f"'{MAESTRA}'!${_letra(maestra, 'TIEMPO' + chr(10) + 'VUELO HR')}${tot}"
        assert fg["combustible"]["D5"].value == f"={tc}"
        assert fg["combustible"]["F5"].value == f"=ROUND(E5/{horas},2)"
        if "otros gastos" in fg.sheetnames:
            assert fg["otros gastos"]["B5"].value == f"={tc}"
            assert fg["otros gastos"]["D5"].value == f"={horas}"


def test_nota_de_repartidos_a_aviones_nombra_la_columna_de_cada_variante() -> None:
    """'repartidos a aviones' decía que el TUA pagado «queda solo como nota en
    OPERACIONES»: el Balance general ya no tiene esa columna, ahí vive en la
    nota del desglose de TOTAL PARA PROVEEDOR (PESOS) (revisión 6-oct-2026).
    El Balance mensual conserva su texto (byte a byte, abajo)."""
    req = tf._general()
    for variante, nota in (
        ("mensual", _NOTA_REPARTIDOS_A_AVIONES),
        ("general", _CPH_NOTA_REPARTIDOS_A_AVIONES),
    ):
        ws = load_workbook(BytesIO(render_balance_general_xlsx(_con_variante(req, variante))))[
            "repartidos a aviones"
        ]
        notas = [
            c.value
            for c in ws["A"]
            if isinstance(c.value, str) and c.value.startswith("REPARTIDOS A AVIONES =")
        ]
        assert notas == [nota], variante
    assert _NOTA_REPARTIDOS_A_AVIONES.endswith(
        "queda solo como nota en OPERACIONES y vive en 'otros movimientos'."
    )
    assert "OPERACIONES" not in _CPH_NOTA_REPARTIDOS_A_AVIONES
    assert _CPH_NOTA_REPARTIDOS_A_AVIONES.endswith(
        "queda solo como nota en el desglose de TOTAL PARA PROVEEDOR (PESOS) de la hoja de "
        "vuelos (parte Operación) y vive en 'otros movimientos'."
    )


def test_resumen_y_balance_son_identicos_en_las_dos_variantes() -> None:
    req = BalanceGeneralRequest.model_validate(t_empresa._payload())
    mensual = load_workbook(BytesIO(render_balance_general_xlsx(req)))
    general = load_workbook(BytesIO(render_balance_general_xlsx(_con_variante(req, "general"))))
    for hoja in ("RESUMEN flota", "cobranza", "balance", "otros movimientos", "inventario"):
        a = {c.coordinate: c.value for fila in mensual[hoja].iter_rows() for c in fila}
        b = {c.coordinate: c.value for fila in general[hoja].iter_rows() for c in fila}
        assert a == b, hoja


def test_el_libro_individual_nunca_usa_la_variante() -> None:
    wb = Workbook()
    maestra = _hoja_maestra(wb.active, BalanceAvionRequest(matricula="XA-TST"), costo_por_hora=True)
    assert maestra.disp is _DISP_MENSUAL
    assert wb.active.max_column == len(_COLS)


def test_cita_a_una_columna_que_la_variante_no_tiene_va_como_valor_con_nota() -> None:
    """Regla dura del contrato: jamás una referencia a una columna que no
    existe — el VALOR del API con nota «calculado por el sistema»."""
    sin_tiempo = _disposicion(
        [c for c in _COLS if c[2] != "tiempo_vuelo"],
        total_map=_TOTAL_MAP,
        fills={},
        detalle_attr={},
    )
    maestra = _Maestra(
        titulo="reporte horas X", fila_ini=3, fila_fin=4, fila_tot=5, general=True, disp=sin_tiempo
    )
    ws = Workbook().active
    celda = _formula_cita(ws, 5, 4, maestra, "ROUND(E5/{tiempo_vuelo},2)", 12.5, HORAS)
    assert celda.value == 12.5
    assert celda.comment.text == _NOTA_CALCULADO_POR_SISTEMA_CELDA
    celda = _formula_cita(ws, 5, 2, maestra, "{tc_costos}", 17.5, TC)
    assert celda.value == f"='reporte horas X'!${sin_tiempo.letra['tc_costos']}$5"
    assert celda.comment is None
    assert _formula_cita(ws, 6, 4, maestra, "{tiempo_vuelo}", None, HORAS).comment is None
    sin_maestra = _formula_cita(ws, 7, 4, None, "{tiempo_vuelo}", 3.0, HORAS)
    assert (sin_maestra.value, sin_maestra.comment) == (3.0, None)


# ---------------------------------------------------------------------------
# 4b) Regla dura «jamás letras fijas», con el libro COMPLETO: otro juego de
#     columnas y todo sigue cuadrando (revisión 6-oct-2026). Sin esto, una
#     fórmula o una cita escrita con letras fijas pasaba la suite entera:
#     coincide con el layout de hoy.
# ---------------------------------------------------------------------------

_JUEGOS_MUTADOS = (
    "columna real antes de VENTA",
    "OPERACIONES dentro del bloque",
    "sin TIPO CAMBIO COSTOS",
)


def _juego_mutado(caso: str) -> list[tuple[str, str, str | None, str | None]]:
    cols = list(_COLS_GENERAL)
    if caso == "columna real antes de VENTA":  # corre todas las letras desde E
        i = next(i for i, c in enumerate(cols) if c[0] == "VENTA")
        cols.insert(i, ("", "OPERADOR", "operador_externo", None))
    elif caso == "OPERACIONES dentro del bloque":  # corre medio bloque
        i = next(i for i, c in enumerate(cols) if c[0] == _GRUPO_COSTO_HORA)
        cols.insert(i + 3, (_GRUPO_COSTO_HORA, "OPERACIONES", "op_mxn", MONEY))
    else:  # la columna que citan 'combustible' y las hojas de gastos
        cols = [c for c in cols if c[2] != "tc_costos"]
    return cols


@pytest.mark.parametrize("caso", _JUEGOS_MUTADOS)
def test_otro_juego_de_columnas_sigue_cuadrando(caso, monkeypatch, registros) -> None:
    cols = _juego_mutado(caso)
    monkeypatch.setattr(
        balance_avion_xlsx,
        "_DISP_GENERAL",
        _disposicion(
            cols,
            total_map=_TOTAL_MAP_GENERAL,
            fills=_DISP_GENERAL.fills,
            detalle_attr={},
            titulos_grupo=_DISP_GENERAL.titulos_grupo,
            costo_por_hora=True,
        ),
    )
    req = _general()
    data = render_balance_general_xlsx(req)
    # Ninguna fórmula volvió a valor y el evaluador independiente cuadra.
    assert registros[-1].degradadas() == []
    assert verificar_libro(data)
    wb, wv = _libros(data)
    ws = wb[MAESTRA]
    assert [c.value for c in ws[2]] == [c[1] for c in cols]
    # Cada celda numérica es el número del API de SU llave.
    for k, v in enumerate(req.consolidado.vuelos):
        for i, (_g, encabezado, clave, fmt) in enumerate(cols, start=1):
            if clave is None or fmt is None:
                continue
            esperado = _valor_columna(clave, v)
            visto = wv[MAESTRA].cell(row=3 + k, column=i).value
            if esperado is None:
                assert visto is None, (caso, 3 + k, encabezado)
            else:
                assert visto == pytest.approx(esperado, abs=1e-9), (caso, 3 + k, encabezado)
    # Las fórmulas del bloque citan las letras de ESTE juego.
    f = _fila_de(ws, "#401 · Cliente Uno")
    for encabezado, formula in _formulas_del_bloque(ws, f).items():
        assert _celda(ws, f, encabezado).value == formula, (caso, encabezado)
    # TOTALES: cada Σ suma su propia columna.
    tot = _fila_de(ws, "TOTALES")
    sumas = [c for c in ws[tot] if isinstance(c.value, str) and c.value.startswith("=")]
    assert sumas
    for c in sumas:
        assert c.value == f"=ROUND(SUM({c.column_letter}3:{c.column_letter}{tot - 1}),2)", caso

    # Las demás hojas citan la fila TOTALES con la letra de la llave EN ESTE
    # juego; si el juego no la tiene, el VALOR del API con la nota.
    def cita(encabezado: str) -> str | None:
        if encabezado not in [c.value for c in ws[2]]:
            return None
        return f"'{MAESTRA}'!${_letra(ws, encabezado)}${tot}"

    tc, horas = cita("TIPO CAMBIO\nCOSTOS"), cita("TIEMPO\nVUELO HR")
    assert horas is not None and (tc is None) == (caso == "sin TIPO CAMBIO COSTOS")
    t = req.consolidado.totales
    esperadas = [
        ("combustible", "D5", tc, f"={tc}", t.tc_promedio),
        ("combustible", "F5", horas, f"=ROUND(E5/{horas},2)", None),
        ("otros gastos", "B5", tc, f"={tc}", t.tc_promedio),
        ("otros gastos", "D5", horas, f"={horas}", None),
        ("repartidos a aviones", "B5", tc, f"={tc}", t.tc_promedio),
        ("repartidos a aviones", "D5", horas, f"={horas}", None),
    ]
    for hoja, coord, ref, formula, valor in esperadas:
        celda = wb[hoja][coord]
        if ref is not None:
            assert (celda.value, celda.comment) == (formula, None), (caso, hoja, coord)
        else:
            assert celda.value == valor, (caso, hoja, coord)
            assert celda.comment.text == _NOTA_CALCULADO_POR_SISTEMA_CELDA, (caso, hoja, coord)
    # Y ninguna otra celda del libro cita otra columna de la hoja de vuelos.
    permitidas = {r for r in (tc, horas) if r}
    for hoja in wb.sheetnames:
        if hoja == MAESTRA:
            continue
        for fila in wb[hoja].iter_rows():
            for c in fila:
                if isinstance(c.value, str) and MAESTRA in c.value:
                    refs = set(re.findall(rf"'{re.escape(MAESTRA)}'!\$[A-Z]+\$\d+", c.value))
                    assert refs and refs <= permitidas, (caso, hoja, c.coordinate, c.value)


# ---------------------------------------------------------------------------
# 5) Balance MENSUAL: el libro de hoy, byte a byte.
# ---------------------------------------------------------------------------

# Huellas calculadas con el código ANTERIOR a este cambio (HEAD 93486c1,
# 6-oct-2026) sobre los mismos payloads. Además se verificó que los 37 libros
# de los tests del balance (individuales y generales) salían byte-idénticos
# —todos los miembros del .xlsx salvo docProps/core.xml— salvo los 3 con una
# fila PROVISIÓN, donde solo cambia 'otros movimientos' (el amarillo).
_FIRMAS_HOY = {
    "fórmulas": "7dcf19e716a271f873ef46bea9a179ac5f9401d7008703103a1ad1634666bf19",
    "empresa": "a3fdaea1040978760ed02d6bc187edbad7cdbbb0da4a7a1289f8f278f18f9b72",
    "pago real": "f1b3ac2829c5da84c2db68f256308d6027387c62e833239767f02937e0145d52",
}
_FIRMA_HOY_PROVISION_SIN_OTROS_MOVIMIENTOS = (
    "7a3c12df7d64abb45e8b37ddd566a41535f122169f592cbc03811eaa51f594d6"
)
_CASOS_MENSUAL = {
    "fórmulas": tf._general,
    "empresa": lambda: BalanceGeneralRequest.model_validate(t_empresa._payload()),
    # Gasto real del vendedor: su nota dice «reemplaza la provisión» en
    # minúsculas y NO se pinta de amarillo.
    "pago real": lambda: t_comision._general(True, filas=[t_comision._FILA_PAGO_REAL]),
}
# Los BYTES del libro de hoy (revisión 6-oct-2026): `_firma_libro` es una
# huella del layout y no ve el tamaño de las notas (openpyxl no relee el VML),
# el color ni los otros lados de un borde, ni el nombre de la fuente. Aquí va
# cada miembro del .xlsx salvo docProps/core.xml (lleva la hora): los
# primeros 16 hex de su sha256 (64 bits, de sobra para notar un byte movido).
# Se calcularon renderizando estos payloads con el código ANTERIOR a la
# variante (`git archive 93486c1 app tests`, con sus propios helpers de
# tests) y openpyxl 3.1.5; el código nuevo da exactamente los mismos. Si
# falla SOLO porque cambió la versión de openpyxl (y `_FIRMAS_HOY` sigue
# igual), se regeneran con ese árbol de 93486c1 y la versión nueva.
_OPENPYXL_HUELLAS = "3.1.5"
_MIEMBROS_HOY = {
    "fórmulas": {
        "[Content_Types].xml": "90c274d39d8df938",
        "_rels/.rels": "c545941ba36c15fc",
        "docProps/app.xml": "209fca6b00afe72a",
        "xl/_rels/workbook.xml.rels": "fee7f7835a21adf6",
        "xl/comments/comment1.xml": "1f01ffd50cbe956e",
        "xl/comments/comment2.xml": "a623ddcb4e5ebb3c",
        "xl/drawings/commentsDrawing1.vml": "f3fd6ba45292a07d",
        "xl/drawings/commentsDrawing2.vml": "a3fb237fe6d57d83",
        "xl/styles.xml": "c00b3c833322f4b3",
        "xl/theme/theme1.xml": "d15e8ebf78ef7b97",
        "xl/workbook.xml": "e9b5be0ab87733e4",
        "xl/worksheets/_rels/sheet2.xml.rels": "e2d28d8e38b35f17",
        "xl/worksheets/_rels/sheet9.xml.rels": "06a3b87cfefe7f5f",
        "xl/worksheets/sheet1.xml": "ed85f7783357e4ef",
        "xl/worksheets/sheet10.xml": "b134206c4fd136b1",
        "xl/worksheets/sheet2.xml": "9db705574eaea0cd",
        "xl/worksheets/sheet3.xml": "bb3dc0f4f2b56702",
        "xl/worksheets/sheet4.xml": "b4271af6eb64e959",
        "xl/worksheets/sheet5.xml": "0122d2728a5bfe58",
        "xl/worksheets/sheet6.xml": "91c22671b5955760",
        "xl/worksheets/sheet7.xml": "d35dd9696ae4e87e",
        "xl/worksheets/sheet8.xml": "85aaef47b6b15de9",
        "xl/worksheets/sheet9.xml": "fc8fb7c900b61fab",
    },
    "empresa": {
        "[Content_Types].xml": "90c274d39d8df938",
        "_rels/.rels": "c545941ba36c15fc",
        "docProps/app.xml": "209fca6b00afe72a",
        "xl/_rels/workbook.xml.rels": "fee7f7835a21adf6",
        "xl/comments/comment1.xml": "1f01ffd50cbe956e",
        "xl/comments/comment2.xml": "51429da7b4342562",
        "xl/drawings/commentsDrawing1.vml": "f3fd6ba45292a07d",
        "xl/drawings/commentsDrawing2.vml": "14fe3cbfa808de3c",
        "xl/styles.xml": "2354a5c4769e8a1b",
        "xl/theme/theme1.xml": "d15e8ebf78ef7b97",
        "xl/workbook.xml": "e9b5be0ab87733e4",
        "xl/worksheets/_rels/sheet2.xml.rels": "e2d28d8e38b35f17",
        "xl/worksheets/_rels/sheet9.xml.rels": "06a3b87cfefe7f5f",
        "xl/worksheets/sheet1.xml": "b0d80f991bd39ac2",
        "xl/worksheets/sheet10.xml": "245576c8b1d18f98",
        "xl/worksheets/sheet2.xml": "9db705574eaea0cd",
        "xl/worksheets/sheet3.xml": "bb3dc0f4f2b56702",
        "xl/worksheets/sheet4.xml": "b4271af6eb64e959",
        "xl/worksheets/sheet5.xml": "0122d2728a5bfe58",
        "xl/worksheets/sheet6.xml": "91c22671b5955760",
        "xl/worksheets/sheet7.xml": "d35dd9696ae4e87e",
        "xl/worksheets/sheet8.xml": "85aaef47b6b15de9",
        "xl/worksheets/sheet9.xml": "fd277d78da5b85d8",
    },
    "pago real": {
        "[Content_Types].xml": "aab5d091f752ef52",
        "_rels/.rels": "c545941ba36c15fc",
        "docProps/app.xml": "209fca6b00afe72a",
        "xl/_rels/workbook.xml.rels": "7676081832cc25e3",
        "xl/comments/comment1.xml": "5da58ad45ce720b4",
        "xl/comments/comment2.xml": "4538ed5560018e12",
        "xl/drawings/commentsDrawing1.vml": "d883bb9ec4031f10",
        "xl/drawings/commentsDrawing2.vml": "ed0e12dacd2b3460",
        "xl/styles.xml": "c6184aacbdeea286",
        "xl/theme/theme1.xml": "d15e8ebf78ef7b97",
        "xl/workbook.xml": "6c579d04c9cd1373",
        "xl/worksheets/_rels/sheet2.xml.rels": "e2d28d8e38b35f17",
        "xl/worksheets/_rels/sheet3.xml.rels": "06a3b87cfefe7f5f",
        "xl/worksheets/sheet1.xml": "3affe542cc6fd992",
        "xl/worksheets/sheet2.xml": "946870294a3621b2",
        "xl/worksheets/sheet3.xml": "495b31f28f184ace",
        "xl/worksheets/sheet4.xml": "fb8f4d129040df3f",
        "xl/worksheets/sheet5.xml": "9095ac67c5861e2b",
        "xl/worksheets/sheet6.xml": "2d5e2e80fcfd0233",
        "xl/worksheets/sheet7.xml": "29118b7c10960b3e",
        "xl/worksheets/sheet8.xml": "f231507baa758a09",
    },
}


def _huellas_bytes(data: bytes) -> dict[str, str]:
    return {n: hashlib.sha256(b).hexdigest()[:16] for n, b in _miembros(data).items()}


@pytest.mark.parametrize("caso", list(_CASOS_MENSUAL))
def test_mensual_es_el_libro_de_hoy(caso) -> None:
    req = _CASOS_MENSUAL[caso]()
    data = render_balance_general_xlsx(req)
    assert _firma_libro(data) == _FIRMAS_HOY[caso]
    # Byte a byte: cada miembro del .xlsx es el del libro de hoy.
    assert _huellas_bytes(data) == _MIEMBROS_HOY[caso], (
        f"openpyxl {openpyxl.__version__} (huellas calculadas con {_OPENPYXL_HUELLAS})"
    )
    # Con la llave en cualquiera de sus formas «mensual», byte a byte igual.
    base = _miembros(data)
    for variante in ("mensual", None, " MENSUAL ", "otra"):
        assert _miembros(render_balance_general_xlsx(_con_variante(req, variante))) == base


def test_mensual_y_general_difieren_solo_en_la_hoja_maestra_y_sus_citas() -> None:
    """Fuera de la hoja maestra, lo único distinto entre las dos variantes son
    las celdas que la citan, con la letra de SU columna en cada juego, y la
    nota de 'repartidos a aviones' (nombra la columna de cada variante)."""
    req = tf._general()
    mensual = render_balance_general_xlsx(req)
    general = render_balance_general_xlsx(_con_variante(req, "general"))
    wm, wg = load_workbook(BytesIO(mensual)), load_workbook(BytesIO(general))
    assert wm.sheetnames == wg.sheetnames
    cambio = {
        f"'{MAESTRA}'!${_DISP_MENSUAL.letra[k]}$": f"'{MAESTRA}'!${_DISP_GENERAL.letra[k]}$"
        for k in ("tc_costos", "tiempo_vuelo")
    }
    citas, notas = set(), set()
    for hoja in wm.sheetnames:
        if hoja == MAESTRA:
            continue
        for fila in wm[hoja].iter_rows():
            for c in fila:
                otro = wg[hoja][c.coordinate].value
                if c.value == otro:
                    continue
                if c.value == _NOTA_REPARTIDOS_A_AVIONES:
                    assert otro == _CPH_NOTA_REPARTIDOS_A_AVIONES, (hoja, c.coordinate)
                    notas.add(hoja)
                    continue
                assert isinstance(c.value, str) and MAESTRA in c.value, (hoja, c.coordinate)
                esperado = c.value
                for antes, ahora in cambio.items():
                    esperado = esperado.replace(antes, ahora)
                assert otro == esperado, (hoja, c.coordinate)
                citas.add(hoja)
    assert citas == {"combustible", "otros gastos", "repartidos a aviones"}
    assert notas == {"repartidos a aviones"}
    resto = (MAESTRA, *citas)
    assert _firma_libro(mensual, sin_hojas=resto) == _firma_libro(general, sin_hojas=resto)


# ---------------------------------------------------------------------------
# 6) PROVISIÓN en amarillo en 'otros movimientos'.
# ---------------------------------------------------------------------------

# Vuelo con TUA pagado + provisión: el API junta los dos egresos en UNA fila
# y la marca PROVISIÓN queda en la nota (una línea por concepto).
_FILA_COLAPSADA = {
    "clave": "N4142R-320",
    "avion_color": "#FF0000",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-09-22",
    "concepto_egreso": "TUAs + pago comisión vendedor · 2 conceptos (ver nota)",
    "egreso_mxn": 3130.25,
    "fecha_egreso": "2026-09-22",
    "nota_egreso": (
        "tuas pagadas = $1,100.25\n"
        "pago comisión vendedor (Alex Saab) · PROVISIÓN (mismo monto que lo cobrado: "
        "comisión + IVA; sin gasto real capturado) = $2,030.00"
    ),
    "concepto_ingreso": "TUAs + comisión vendedor con IVA · 2 conceptos (ver nota)",
    "ingreso_mxn": 3280.5,
    "fecha_ingreso": "2026-09-22",
    "remanente_mxn": 150.25,
}


def test_es_provision_solo_con_la_marca_en_mayusculas() -> None:
    assert _es_provision(BalanceOtroMovimientoFila(**t_comision._FILA_PROVISION))
    assert _es_provision(BalanceOtroMovimientoFila(**_FILA_COLAPSADA))
    assert _es_provision(BalanceOtroMovimientoFila(concepto_egreso="pago · PROVISION (x)"))
    # El gasto real «reemplaza la provisión» (minúsculas): no es provisión.
    assert not _es_provision(BalanceOtroMovimientoFila(**t_comision._FILA_PAGO_REAL))
    assert not _es_provision(BalanceOtroMovimientoFila(concepto_egreso="TUA pagado"))
    assert not _es_provision(BalanceOtroMovimientoFila())
    # Un PROVEEDOR en mayúsculas no es la marca (revisión 6-oct-2026): la fila
    # suelta lleva «<categoría> · <proveedor>», como la arma el API, y la razón
    # social del CFDI suele venir en mayúsculas.
    for concepto in (
        "Combustible · PROVISIONES AEREAS DEL SURESTE",
        "Hangar · PROVISION AEREA SA DE CV",
        "Mantenimiento · PROVISIÓN Y SERVICIOS DEL CARIBE (USD sin TC)",
    ):
        fila = BalanceOtroMovimientoFila(concepto_egreso=concepto, egreso_mxn=5000)
        assert not _es_provision(fila), concepto
    assert not _es_provision(
        BalanceOtroMovimientoFila(nota_egreso="FBO · PROVISIONES DEL CARIBE = $350.00")
    )


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_provision_en_amarillo_con_leyenda_al_pie(variante) -> None:
    filas = [t_comision._FILA_PROVISION, t_comision._FILA_PAGO_REAL, _FILA_COLAPSADA]
    req = _con_variante(t_comision._general(True, filas=filas), variante)
    ws = load_workbook(BytesIO(render_balance_general_xlsx(req)))["otros movimientos"]
    amarillo = "00" + AMBER
    for fila in (4, 6):  # provisión sola y la fila colapsada
        assert [_relleno(ws.cell(row=fila, column=c)) for c in range(2, 11)] == [
            None,
            amarillo,
            amarillo,
            *[None] * 6,
        ], fila
    assert [_relleno(ws.cell(row=5, column=c)) for c in range(2, 11)] == [None] * 9
    assert ws["A4"].value == "N4142R-317" and _relleno(ws["A4"]) == "00FF0000"
    tot = _fila_de(ws, "TOTALES")
    assert tot == 7
    ley = ws.cell(row=tot + 2, column=1)
    assert ley.value == _LEYENDA_PROVISION == "Amarillo = provisión (sin gasto real capturado)"
    assert _relleno(ley) == amarillo
    assert f"A{tot + 2}:C{tot + 2}" in {str(r) for r in ws.merged_cells.ranges}
    # La fila 2 de la hoja no cambia (la explicación del amarillo va al pie).
    assert ws["A2"].value.startswith("Ingreso de VuelaTour (no del avión)")


def test_sin_provision_no_hay_amarillo_ni_leyenda() -> None:
    req = t_comision._general(True, filas=[t_comision._FILA_PAGO_REAL])
    ws = load_workbook(BytesIO(render_balance_general_xlsx(req)))["otros movimientos"]
    rellenos = {_relleno(c) for fila in ws.iter_rows() for c in fila}
    assert "00" + AMBER not in rellenos
    assert _LEYENDA_PROVISION not in {c.value for c in ws["A"]}


def test_con_provision_solo_cambia_otros_movimientos() -> None:
    data = render_balance_general_xlsx(t_comision._general())
    assert _firma_libro(data, sin_hojas=("otros movimientos",)) == (
        _FIRMA_HOY_PROVISION_SIN_OTROS_MOVIMIENTOS
    )


# ---------------------------------------------------------------------------
# 7) Ruta: el JSON del API 0.0.64 entra sin 422.
# ---------------------------------------------------------------------------

TOKEN = "secreto-de-prueba"
client = TestClient(app)


def test_ruta_acepta_la_variante(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    payload = tf._general().model_dump(mode="json")

    periodo = f"{payload['periodo_desde']}-{payload['periodo_hasta']}"

    res = client.post(
        "/pdf/balance-general-xlsx", json={**payload, "variante": "general"}, headers=headers
    )
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content))[MAESTRA]
    assert [c.value for c in ws[2]] == [c[1] for c in _COLS_GENERAL]
    # El nombre del archivo dice cuál de los dos libros es (como el del API).
    disposicion = res.headers["content-disposition"]
    assert f'filename="balance-general-vuelatour-{periodo}.xlsx"' in disposicion

    sin_llave = {k: v for k, v in payload.items() if k != "variante"}
    res = client.post("/pdf/balance-general-xlsx", json=sin_llave, headers=headers)
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content))[MAESTRA]
    assert [c.value for c in ws[2]] == [c[1] for c in _COLS]
    disposicion = res.headers["content-disposition"]
    assert f'filename="balance-mensual-vuelatour-{periodo}.xlsx"' in disposicion


def test_columnas_por_llave_no_por_letra() -> None:
    """Las dos disposiciones resuelven sus letras desde la lista de columnas:
    la misma llave cae en columnas distintas según la variante."""
    assert _DISP_MENSUAL.letra["tc_costos"] == "V"
    assert _DISP_GENERAL.letra["tc_costos"] == "R"
    assert _DISP_GENERAL.letra["tiempo_vuelo"] == _DISP_MENSUAL.letra["tiempo_vuelo"] == "N"
    assert _DISP_GENERAL.cobro1_col == next(
        i for i, c in enumerate(_COLS_GENERAL, start=1) if c[1] == "COBRO 1\nFECHA"
    )
    assert len(_DISP_GENERAL.cobro_mxn_letras) == 4
    assert _DISP_GENERAL.cobro_mxn_letras[0] == get_column_letter(_DISP_GENERAL.cobro1_col + 1)
