"""PERMISOS AFAC COMO FILAS QUE SUMAN en 'otros movimientos' del Balance
general (13-oct-2026, API 0.0.74, invariante 62 del API).

Decisión 2 del lote del cliente del 13-oct-2026: «AFAC en "otros movimientos"
del Balance general: filas en la tabla principal que suman en TOTALES
(provisión AFAC del periodo = ingreso de VuelaTour por avión; cada permiso
pagado = egreso), ya no bloque informativo bajo totales».

El cambio es del API: manda al FINAL de `filas` una fila «AFAC · matrícula»
por avión (`seccion: 'AFAC'`) y la bandera `afac_en_filas: true`. Aquí no se
calcula nada: montos, notas y remanente llegan hechos y se pintan con
`pinta()` de siempre, así que SUMAN en TOTALES con las fórmulas de siempre.
Solo con la bandera: un encabezado de sección antes de la primera fila AFAC
(sin montos, sin merge, sin fila en blanco), la leyenda `_OM_AFAC_EN_FILAS`
al final de la fila 2, las etiquetas `_EMPRESA_*_AFAC` del bloque VUELATOUR,
el fragmento `_EMPRESA_AFAC_EN_FILAS` en su pie y SIN el bloque informativo
`afac` aunque llegara (nunca doble). Sin la bandera, el libro de siempre byte
a byte (un API 0.0.69–0.0.73 sigue viendo su bloque informativo).
"""

from __future__ import annotations

import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAfacOtrosMovimientos,
    BalanceGeneralRequest,
    BalanceHojaOtrosMovimientos,
    BalanceOtroMovimientoFila,
)
from app.services.balance_avion_xlsx import (
    _EMPRESA_AFAC_EN_FILAS,
    _EMPRESA_INGRESOS,
    _EMPRESA_INGRESOS_AFAC,
    _EMPRESA_PAGOS,
    _EMPRESA_PAGOS_AFAC,
    _EMPRESA_REGLA_COMISIONES_GASTO_REAL,
    _LEYENDA_PROVISION,
    _MARCA_PROVISION,
    _NOTA_BLOQUE_EMPRESA,
    _NOTA_EMPRESA_BASE,
    _NOTA_EMPRESA_BASE_SIN_PROVISION,
    _OM_AFAC_EN_FILAS,
    _OM_SECCION_AFAC,
    _TITULO_AFAC_OM,
    LIGHT,
    NAVY,
    RED,
    render_balance_general_xlsx,
)
from tests import test_balance_comisiones_general as tcg
from tests import test_balance_empresa_vuelatour as tev
from tests import test_balance_formulas as tf
from tests._evaluador_formulas import verificar_libro
from tests.test_balance_general_costo_hora import (
    _FIRMAS_HOY,
    _MIEMBROS_HOY,
    _firma_libro,
    _huellas_bytes,
)
from tests.test_balance_general_costo_hora import _general as _general_cph
from tests.test_comision_vendedor_notas import OM_FILA2_VIEJA, _rellenos_amber

r2 = tf.r2

TOKEN = "secreto-de-prueba"
client = TestClient(app)

# ---------------------------------------------------------------------------
# Filas AFAC con la forma EXACTA del API 0.0.74 (`filasAfacOtrosMovimientos`).
# ---------------------------------------------------------------------------

NOTA_INGRESO_TST = "\n".join(
    [
        "Provisión AFAC de cada vuelo de XA-TST (columna PERMISO AFAC de su hoja "
        "de vuelos): tarifa 25.00 USD/hr × T.C. de costos × horas cobradas",
        "#401 · Cliente Uno = $1,082.50",
        "#403 · Traslado = $866.00",
        "#405 · Cliente Cuatro = $1,516.50",
    ]
)
NOTA_EGRESO_DOS = "\n".join(
    [
        "permiso AFAC pagado · Internación = $700.00 · 20/09/2026",
        "permiso AFAC pagado · AFAC anual = $500.00 · 02/09/2026",
    ]
)
FILA_AFAC_TST = {
    "clave": "AFAC · XA-TST",
    "avion_color": "#1F77B4",
    "fecha_vuelo": None,
    "concepto_egreso": "permiso AFAC pagado · AFAC anual",
    "egreso_mxn": 2500.0,
    "fecha_egreso": "2026-09-05",
    "concepto_ingreso": "provisión AFAC del periodo · 3 vuelos · tarifa 25.00 USD/hr",
    "ingreso_mxn": 3465.0,
    "fecha_ingreso": None,
    "remanente_mxn": 965.0,
    "factura": None,
    "nota_ingreso": NOTA_INGRESO_TST,
    "seccion": "AFAC",
}
FILA_AFAC_DOS = {
    "clave": "AFAC · XB-DOS",
    "avion_color": "#2CA02C",
    "fecha_vuelo": None,
    "concepto_egreso": "permisos AFAC pagados · 2 conceptos (ver nota)",
    "egreso_mxn": 1200.0,
    "fecha_egreso": "2026-09-02",
    "concepto_ingreso": None,
    "ingreso_mxn": None,
    "fecha_ingreso": None,
    "remanente_mxn": -1200.0,
    "factura": None,
    "nota_egreso": NOTA_EGRESO_DOS,
    "seccion": "AFAC",
}
# El bloque INFORMATIVO de un API 0.0.69–0.0.73 (el de
# `tf.test_general_otros_movimientos_bloque_permisos_afac`).
BLOQUE_AFAC_VIEJO = {
    "provision": [
        {"matricula": "XA-TST", "tarifa_usd_hr": 25.0, "vuelos": 4, "monto_mxn": 3465.0},
        {"matricula": "XB-DOS", "tarifa_usd_hr": None, "vuelos": 0, "monto_mxn": 0.0},
    ],
    "provision_total_mxn": 3465.0,
    "pagados": [
        {"matricula": "XA-TST", "fecha": "2026-09-05", "detalle": "AFAC anual", "monto_mxn": 2500.0}
    ],
    "pagados_total_mxn": 2500.0,
    "remanente_mxn": 965.0,
}


def _con_afac(
    req: BalanceGeneralRequest,
    *,
    bandera: bool | None = True,
    con_bloque: bool = False,
    filas: list[dict] | None = None,
) -> BalanceGeneralRequest:
    """El general con las filas AFAC anexadas AL FINAL de `filas` (como
    `payloadGeneral`) y la bandera (`None` ⇒ no viaja); `con_bloque` pone
    además el bloque informativo de un API previo (payload mixto)."""
    om = req.consolidado.otros_movimientos
    assert om is not None
    for f in filas if filas is not None else [FILA_AFAC_TST, FILA_AFAC_DOS]:
        om.filas.append(BalanceOtroMovimientoFila.model_validate(f))
    if bandera is not None:
        om.afac_en_filas = bandera
    if con_bloque:
        om.afac = BalanceAfacOtrosMovimientos.model_validate(BLOQUE_AFAC_VIEJO)
    return req


def _libros(data: bytes):
    """(con fórmulas, con caché)."""
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _fila_de(ws, texto: str) -> int:
    filas = [c.row for c in ws["A"] if c.value == texto]
    assert len(filas) == 1, (texto, filas)
    return filas[0]


def _textos(ws) -> list[str]:
    return [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]


# ---------------------------------------------------------------------------
# (a) Encabezado de sección, filas AFAC y TOTALES que las suman.
# ---------------------------------------------------------------------------


def test_textos_son_los_del_contrato() -> None:
    assert _OM_SECCION_AFAC == (
        "PERMISOS AFAC POR AVIÓN · provisión del periodo (ingreso de VuelaTour) vs "
        "permisos pagados (egreso)"
    )
    assert _OM_AFAC_EN_FILAS == (
        " PERMISOS AFAC POR AVIÓN: al final de las filas por vuelo hay una fila «AFAC · "
        "matrícula» por avión que SÍ suma en TOTALES y en el bloque VUELATOUR (empresa) "
        "de 'balance'. Su INGRESO es la provisión AFAC del periodo que ese avión ya "
        "carga como costo en la columna PERMISO AFAC (PROVISIÓN) de la hoja de vuelos "
        "(tarifa AFAC USD/hr del avión × tipo de cambio de costos × horas cobradas); su "
        "EGRESO son TODOS los gastos «Permiso» de ese avión en el periodo (los de la "
        "hoja 'permisos' de su libro individual, sean o no de AFAC); remanente = "
        "provisión − pagado. Los permisos de vuelos EXTERNOS no entran (se quedan en el "
        "libro EXTERNOS)."
    )
    assert _EMPRESA_INGRESOS_AFAC == (
        "(+) INGRESOS PROPIOS USD (TUAs, extras, pernocta, comisión — cotizado del "
        "periodo; provisión AFAC de los aviones)"
    )
    assert _EMPRESA_PAGOS_AFAC == (
        "(−) EGRESOS PROPIOS USD (pago al vendedor, TUAs, extensión, comisión bancaria, "
        "permisos AFAC, sueltos)"
    )
    assert _EMPRESA_AFAC_EN_FILAS == (
        " Incluyen las filas «AFAC · matrícula» de 'otros movimientos': la provisión "
        "AFAC del periodo de cada avión como ingreso (la que ese avión ya carga como "
        "costo en su columna PERMISO AFAC) y los permisos AFAC pagados como egreso, a "
        "USD con el T.C. promedio de costos del libro de su avión."
    )
    # El encabezado NO empieza como el título del bloque informativo (que
    # conserva «(informativo)» como discriminador) y ningún texto AFAC lleva
    # la marca del amarillo.
    assert _TITULO_AFAC_OM != _OM_SECCION_AFAC
    assert "(informativo)" in _TITULO_AFAC_OM
    for texto in (
        _OM_SECCION_AFAC,
        _OM_AFAC_EN_FILAS,
        _EMPRESA_INGRESOS_AFAC,
        _EMPRESA_PAGOS_AFAC,
        _EMPRESA_AFAC_EN_FILAS,
        *(str(v) for f in (FILA_AFAC_TST, FILA_AFAC_DOS) for v in f.values()),
    ):
        assert not _MARCA_PROVISION.search(texto), texto


def test_encabezado_filas_afac_y_totales_que_las_suman() -> None:
    data = render_balance_general_xlsx(_con_afac(tf._general()))
    wb, wv = _libros(data)
    om, omv = wb["otros movimientos"], wv["otros movimientos"]
    ultimo_vuelo = _fila_de(om, "#402 · Cliente Dos")
    enc = _fila_de(om, _OM_SECCION_AFAC)
    tst = _fila_de(om, "AFAC · XA-TST")
    dos = _fila_de(om, "AFAC · XB-DOS")
    sueltas = _fila_de(om, "MOVIMIENTOS SIN AVIÓN / SIN VUELO")
    tot = _fila_de(om, "TOTALES")
    # Encabezado INMEDIATAMENTE después del último vuelo y antes de la
    # primera fila AFAC (sin fila en blanco); las sueltas, después.
    assert (enc, tst, dos) == (ultimo_vuelo + 1, ultimo_vuelo + 2, ultimo_vuelo + 3)
    assert sueltas > dos and tot > sueltas
    # Estilo del encabezado: negrita NAVY tamaño 10, relleno LIGHT en A..J,
    # sin montos y sin merge.
    a = om.cell(row=enc, column=1)
    assert a.font.b is True and a.font.sz == 10 and a.font.color.rgb == f"00{NAVY}"
    for col in range(1, 11):
        celda = om.cell(row=enc, column=col)
        assert celda.fill.fill_type == "solid"
        assert celda.fill.fgColor.rgb == f"00{LIGHT}"
        if col > 1:
            assert celda.value is None
    assert not any(r.min_row <= enc <= r.max_row for r in om.merged_cells.ranges)
    # Fila del XA-TST: swatch del avión, sin sufijo de estado, sus montos,
    # remanente con la fórmula de siempre y las notas como comentario.
    assert om.cell(row=tst, column=1).fill.fgColor.rgb == "001F77B4"
    assert [om.cell(row=tst, column=c).value for c in (2, 3, 4, 5, 6, 7, 8)] == [
        None,
        "permiso AFAC pagado · AFAC anual",
        2500.0,
        "05/09/2026",
        "provisión AFAC del periodo · 3 vuelos · tarifa 25.00 USD/hr",
        3465.0,
        None,
    ]
    assert om.cell(row=tst, column=9).value == f"=ROUND(G{tst}-D{tst},2)"
    assert omv.cell(row=tst, column=9).value == 965.0
    assert om.cell(row=tst, column=10).value is None
    assert om.cell(row=tst, column=7).comment.text == NOTA_INGRESO_TST
    assert om.cell(row=tst, column=4).comment is None
    # Fila del XB-DOS: solo egreso, remanente negativo en rojo, nota del
    # egreso con un renglón por permiso.
    assert [om.cell(row=dos, column=c).value for c in (3, 4, 5, 6, 7)] == [
        "permisos AFAC pagados · 2 conceptos (ver nota)",
        1200.0,
        "02/09/2026",
        None,
        None,
    ]
    assert om.cell(row=dos, column=9).value == f"=ROUND(G{dos}-D{dos},2)"
    assert omv.cell(row=dos, column=9).value == -1200.0
    assert om.cell(row=dos, column=9).font.color.rgb == f"00{RED}"
    assert om.cell(row=dos, column=4).comment.text == NOTA_EGRESO_DOS
    assert om.cell(row=dos, column=10).value is None
    # TOTALES: las MISMAS fórmulas de siempre; las filas AFAC suman (el
    # encabezado no trae montos: SUM lo ignora). Ingresos: #401, #402, AFAC
    # XA-TST y la suelta ING-3.
    assert om[f"D{tot}"].value == f"=ROUND(SUM(D4:D{tot - 1}),2)"
    assert om[f"G{tot}"].value == f"=ROUND(SUM(G4:G{tot - 1}),2)"
    assert om[f"I{tot}"].value == f"=ROUND(G{tot}-D{tot},2)"
    egresos = r2(1100.25 + 2500 + 1200)
    ingresos = r2(1250.5 + 3000 + 3465 + 85.33)
    assert omv[f"D{tot}"].value == egresos
    assert omv[f"G{tot}"].value == ingresos
    assert omv[f"I{tot}"].value == r2(ingresos - egresos)
    # Cada fórmula del libro reproduce el número del API.
    assert verificar_libro(data)


def test_solo_un_encabezado_aunque_haya_varias_filas_afac() -> None:
    ws = load_workbook(BytesIO(render_balance_general_xlsx(_con_afac(tf._general()))))[
        "otros movimientos"
    ]
    assert sum(1 for c in ws["A"] if c.value == _OM_SECCION_AFAC) == 1


# ---------------------------------------------------------------------------
# (b) Con la bandera el bloque informativo NO se pinta aunque llegara.
# ---------------------------------------------------------------------------


def test_con_bandera_y_bloque_viejo_no_se_pinta_el_bloque() -> None:
    """Payload mixto (`afac` + bandera; el API nunca lo manda): las filas
    mandan y el bloque informativo se omite — jamás doble."""
    data = render_balance_general_xlsx(_con_afac(tf._general(), con_bloque=True))
    om = load_workbook(BytesIO(data))["otros movimientos"]
    assert not any(c.value in ("SUBTOTAL AFAC", _TITULO_AFAC_OM) for c in om["A"])
    assert _fila_de(om, _OM_SECCION_AFAC)
    # Lo de abajo de TOTALES queda vacío (ni título ni leyenda del bloque).
    tot = _fila_de(om, "TOTALES")
    assert all(c.value is None for c in om["A"] if c.row > tot)


# ---------------------------------------------------------------------------
# (c) Sin la bandera: el libro del API 0.0.69–0.0.73 (bloque informativo) y
#     filas con `seccion` pintadas como filas normales, sin encabezado.
# ---------------------------------------------------------------------------


def test_sin_bandera_el_bloque_informativo_de_siempre() -> None:
    tf.test_general_otros_movimientos_bloque_permisos_afac()


@pytest.mark.parametrize("bandera", [None, False], ids=["ausente", "false"])
def test_sin_bandera_filas_con_seccion_sin_encabezado_y_bloque_viejo(bandera) -> None:
    data = render_balance_general_xlsx(_con_afac(tf._general(), bandera=bandera, con_bloque=True))
    wb, wv = _libros(data)
    om = wb["otros movimientos"]
    assert _OM_SECCION_AFAC not in _textos(om)
    # Las filas llegan pegadas al último vuelo (sin encabezado) y suman igual.
    ultimo_vuelo = _fila_de(om, "#402 · Cliente Dos")
    assert _fila_de(om, "AFAC · XA-TST") == ultimo_vuelo + 1
    tot = _fila_de(om, "TOTALES")
    assert wv["otros movimientos"][f"D{tot}"].value == r2(1100.25 + 2500 + 1200)
    # El bloque informativo del API previo, debajo de TOTALES.
    titulo = _fila_de(om, _TITULO_AFAC_OM)
    assert titulo > tot
    assert _fila_de(om, "SUBTOTAL AFAC") > titulo
    # Sin la bandera, ni la leyenda nueva de la fila 2.
    assert _OM_AFAC_EN_FILAS not in om["A2"].value


# ---------------------------------------------------------------------------
# (d) Fila 2: la leyenda AFAC va AL FINAL, solo con la bandera.
# ---------------------------------------------------------------------------


def test_fila2_con_bandera_lleva_la_leyenda_al_final() -> None:
    con = load_workbook(BytesIO(render_balance_general_xlsx(_con_afac(tf._general()))))
    assert con["otros movimientos"]["A2"].value == OM_FILA2_VIEJA + _OM_AFAC_EN_FILAS


@pytest.mark.parametrize("bandera", [None, False], ids=["ausente", "false"])
def test_fila2_sin_bandera_la_de_siempre(bandera) -> None:
    sin = load_workbook(
        BytesIO(render_balance_general_xlsx(_con_afac(tf._general(), bandera=bandera)))
    )
    a2 = sin["otros movimientos"]["A2"].value
    assert a2 == OM_FILA2_VIEJA
    assert _OM_AFAC_EN_FILAS not in a2


@pytest.mark.parametrize(
    "banderas",
    [
        {"sin_provision_vendedor": True},
        {"hay_pago_vendedor_real": True},
        {"sin_provision_vendedor": True, "comision_vendedor_solo_gasto_real": True},
    ],
    ids=["0073", "pago_real", "0074"],
)
def test_fila2_la_leyenda_afac_se_suma_al_texto_que_toque(banderas) -> None:
    """Sea cual sea el texto de la fila 2 (provisión, pago real, sin
    provisión, gasto real), con la bandera AFAC es ESE texto + la leyenda."""

    def a2(afac: bool) -> str:
        req = tf._general()
        for k, v in banderas.items():
            setattr(req.consolidado.otros_movimientos, k, v)
        req = _con_afac(req, bandera=True if afac else None)
        wb = load_workbook(BytesIO(render_balance_general_xlsx(req)))
        return wb["otros movimientos"]["A2"].value

    assert a2(True) == a2(False) + _OM_AFAC_EN_FILAS


def test_fila2_con_regla_y_gasto_real_del_api_0074() -> None:
    p = tcg._payload_0074()
    sin = load_workbook(
        BytesIO(render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p)))
    )["otros movimientos"]["A2"].value
    req = _con_afac(BalanceGeneralRequest.model_validate(p))
    con = load_workbook(BytesIO(render_balance_general_xlsx(req)))["otros movimientos"]["A2"].value
    assert con == sin + _OM_AFAC_EN_FILAS


# ---------------------------------------------------------------------------
# (e) Bloque VUELATOUR: etiquetas y pie solo con la bandera.
# ---------------------------------------------------------------------------


def _req_empresa(**om_extra) -> BalanceGeneralRequest:
    req = BalanceGeneralRequest.model_validate(tev._payload())
    for k, v in om_extra.items():
        setattr(req.consolidado.otros_movimientos, k, v)
    return req


def _pie_y_etiquetas(req: BalanceGeneralRequest) -> tuple[object, dict[str, int]]:
    ws = load_workbook(BytesIO(render_balance_general_xlsx(req)))["balance"]
    res = _fila_de(ws, "RESULTADO VUELATOUR USD")
    return ws, {"res": res}


def test_bloque_empresa_con_bandera_etiquetas_y_pie() -> None:
    req = _con_afac(_req_empresa())
    ws, b = _pie_y_etiquetas(req)
    assert _fila_de(ws, _EMPRESA_INGRESOS_AFAC)
    assert _fila_de(ws, _EMPRESA_PAGOS_AFAC)
    assert _EMPRESA_INGRESOS not in _textos(ws)
    assert _EMPRESA_PAGOS not in _textos(ws)
    # Pie: la `nota` del API y debajo la base + el fragmento AFAC.
    assert ws[f"A{b['res'] + 2}"].value == req.empresa.nota
    assert ws[f"A{b['res'] + 3}"].value == _NOTA_EMPRESA_BASE + _EMPRESA_AFAC_EN_FILAS
    # Los números del API, tal cual (la aritmética no cambia aquí).
    fila_ing = _fila_de(ws, _EMPRESA_INGRESOS_AFAC)
    assert ws[f"B{fila_ing}"].value == req.empresa.ingresos_propios_usd


def test_bloque_empresa_con_bandera_y_sin_provision() -> None:
    req = _con_afac(_req_empresa(sin_provision_vendedor=True))
    ws, b = _pie_y_etiquetas(req)
    assert ws[f"A{b['res'] + 3}"].value == (
        _NOTA_EMPRESA_BASE_SIN_PROVISION + _EMPRESA_AFAC_EN_FILAS
    )


def test_bloque_empresa_con_bandera_y_regla_del_api_0074() -> None:
    """Con la regla de comisiones el fragmento AFAC va ANTES de la regla."""
    req = _con_afac(BalanceGeneralRequest.model_validate(tcg._payload_0074()))
    pie = _textos(load_workbook(BytesIO(render_balance_general_xlsx(req)))["balance"])
    esperado = (
        _NOTA_EMPRESA_BASE_SIN_PROVISION
        + _EMPRESA_AFAC_EN_FILAS
        + _EMPRESA_REGLA_COMISIONES_GASTO_REAL.format(regla=tcg.REGLA)
    )
    assert esperado in pie


def test_bloque_empresa_con_bandera_sin_nota_pie_fijo() -> None:
    req = _con_afac(_req_empresa())
    req.empresa.nota = None
    ws, b = _pie_y_etiquetas(req)
    base = _NOTA_EMPRESA_BASE + _EMPRESA_AFAC_EN_FILAS
    assert ws[f"A{b['res'] + 2}"].value == _NOTA_BLOQUE_EMPRESA.replace(_NOTA_EMPRESA_BASE, base)
    assert ws[f"A{b['res'] + 3}"].value is None


@pytest.mark.parametrize("bandera", [None, False], ids=["ausente", "false"])
def test_bloque_empresa_sin_bandera_el_de_siempre(bandera) -> None:
    req = _con_afac(_req_empresa(), bandera=bandera)
    ws, b = _pie_y_etiquetas(req)
    assert _fila_de(ws, _EMPRESA_INGRESOS)
    assert _fila_de(ws, _EMPRESA_PAGOS)
    assert ws[f"A{b['res'] + 3}"].value == _NOTA_EMPRESA_BASE
    assert not any(_EMPRESA_AFAC_EN_FILAS in t for t in _textos(ws))
    assert _EMPRESA_INGRESOS_AFAC not in _textos(ws)
    # Pie fijo sin `nota`: el de siempre, byte a byte.
    req.empresa.nota = None
    ws, b = _pie_y_etiquetas(req)
    assert ws[f"A{b['res'] + 2}"].value == _NOTA_BLOQUE_EMPRESA


# ---------------------------------------------------------------------------
# (f) Las filas AFAC no son provisión del vendedor: nada en amarillo.
# ---------------------------------------------------------------------------


def test_filas_afac_sin_amarillo_ni_leyenda_de_provision() -> None:
    wb = load_workbook(BytesIO(render_balance_general_xlsx(_con_afac(tf._general()))))
    om = wb["otros movimientos"]
    assert _rellenos_amber(om) == []
    assert _LEYENDA_PROVISION not in _textos(om)


# ---------------------------------------------------------------------------
# (g) USD sin T.C.: la fila se pinta tal cual y TOTALES suma lo convertido.
# ---------------------------------------------------------------------------


def test_fila_afac_parcial_usd_sin_tc_suma_solo_lo_convertido() -> None:
    parcial = {
        **FILA_AFAC_TST,
        "concepto_egreso": "permisos AFAC pagados · 2 conceptos (ver nota) (parcial: USD sin TC)",
        "nota_egreso": "permiso AFAC pagado · AFAC anual = $2,500.00 · 05/09/2026\n"
        "permiso AFAC pagado · Permiso USD = — · 07/09/2026",
    }
    data = render_balance_general_xlsx(_con_afac(tf._general(), filas=[parcial]))
    wb, wv = _libros(data)
    om = wb["otros movimientos"]
    fila = _fila_de(om, "AFAC · XA-TST")
    assert om.cell(row=fila, column=3).value.endswith("(parcial: USD sin TC)")
    assert om.cell(row=fila, column=4).value == 2500.0
    tot = _fila_de(om, "TOTALES")
    assert wv["otros movimientos"][f"D{tot}"].value == r2(1100.25 + 2500)
    assert verificar_libro(data)


# ---------------------------------------------------------------------------
# (h) Esquemas aditivos y rutas con el JSON del API 0.0.74.
# ---------------------------------------------------------------------------


def test_esquemas_aditivos() -> None:
    assert BalanceOtroMovimientoFila().seccion is None
    assert BalanceHojaOtrosMovimientos().afac_en_filas is None
    assert BalanceHojaOtrosMovimientos().afac is None


def _json_api_0074(*, mixto: bool = False) -> dict:
    """El payload como lo manda el API 0.0.74: filas con `seccion`, la
    bandera y SIN la llave `afac` (o, `mixto`, con las dos)."""
    p = _con_afac(tf._general(), con_bloque=mixto).model_dump(mode="json")
    om = p["consolidado"]["otros_movimientos"]
    if not mixto:
        om.pop("afac")
    return p


def test_el_json_del_api_0074_y_el_mixto_validan() -> None:
    for mixto in (False, True):
        req = BalanceGeneralRequest.model_validate(_json_api_0074(mixto=mixto))
        om = req.consolidado.otros_movimientos
        assert om.afac_en_filas is True
        assert [f.seccion for f in om.filas][-2:] == ["AFAC", "AFAC"]
        assert (om.afac is not None) is mixto


def test_rutas_aceptan_el_payload_afac(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    p = _json_api_0074()
    res = client.post("/pdf/balance-general-xlsx", json=p, headers=headers)
    assert res.status_code == 200, res.text
    om = load_workbook(BytesIO(res.content))["otros movimientos"]
    assert _fila_de(om, _OM_SECCION_AFAC)
    res = client.post("/pdf/balance-general-libro", json=p, headers=headers)
    assert res.status_code == 200, res.text
    texto = json.dumps(res.json(), ensure_ascii=False)
    assert _OM_SECCION_AFAC in texto
    assert "AFAC · XA-TST" in texto


# ---------------------------------------------------------------------------
# (i) Huellas: sin filas AFAC, la bandera en None/False no cambia un byte.
# ---------------------------------------------------------------------------


def _sin_filas(req: BalanceGeneralRequest, bandera) -> BalanceGeneralRequest:
    if bandera is not ...:
        req.consolidado.otros_movimientos.afac_en_filas = bandera
    return req


@pytest.mark.parametrize("bandera", [..., None, False], ids=["sin_llave", "null", "false"])
def test_huella_de_valores_del_general_intacta(bandera) -> None:
    data = render_balance_general_xlsx(_sin_filas(tf._general(), bandera))
    assert tf._firma_valores(data) == tf._FIRMAS_ANTES["general"]


@pytest.mark.parametrize("bandera", [..., None, False], ids=["sin_llave", "null", "false"])
@pytest.mark.parametrize(
    "caso,armar",
    [
        ("fórmulas", tf._general),
        ("empresa", lambda: BalanceGeneralRequest.model_validate(tev._payload())),
    ],
)
def test_huellas_de_layout_y_bytes_intactas(caso, armar, bandera) -> None:
    data = render_balance_general_xlsx(_sin_filas(armar(), bandera))
    assert _firma_libro(data) == _FIRMAS_HOY[caso]
    assert _huellas_bytes(data) == _MIEMBROS_HOY[caso]


# ---------------------------------------------------------------------------
# (j) La variante «general» (costo por hora) pinta la MISMA pestaña.
# ---------------------------------------------------------------------------


def test_variante_general_mismas_filas_encabezado_y_totales() -> None:
    def pestaña(req: BalanceGeneralRequest):
        data = render_balance_general_xlsx(req)
        wb, wv = _libros(data)
        om, omv = wb["otros movimientos"], wv["otros movimientos"]
        return (
            [(c.coordinate, c.value) for fila in om.iter_rows() for c in fila],
            [(c.coordinate, c.value) for fila in omv.iter_rows() for c in fila],
            data,
        )

    mensual = pestaña(_con_afac(_general_cph("mensual")))
    general = pestaña(_con_afac(_general_cph("general")))
    assert mensual[0] == general[0]
    assert mensual[1] == general[1]
    om = load_workbook(BytesIO(general[2]))["otros movimientos"]
    assert _fila_de(om, _OM_SECCION_AFAC) + 1 == _fila_de(om, "AFAC · XA-TST")
    assert verificar_libro(general[2])
