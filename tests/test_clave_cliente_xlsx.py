"""CLAVE con ID de vuelo por cliente (13-oct-2026, API 0.0.75, decisión 6 del
lote del cliente: «january02-211026» JUNTO al folio).

El API arma la CLAVE (el folio sigue siendo la llave del dinero) y manda
`clave_cliente` SOLO cuando el vuelo tiene ID:

| Hoja | Con ID | Sin ID (como hoy) |
|---|---|---|
| Maestra y cobranza | '#364 · beh05-011026 · Beh Kay ' | '#364 · Beh Kay ' |
| Otros movimientos | '#364 · beh05-011026' (+ ' · CANCELADO' aquí) | 'vtbeh364' |
| Libro Dinero | 'beh05-011026' (+ ' (50 % N4142R)'; servicio = 'vtservicio') | 'vtbeh' |

pyservices pinta la CLAVE tal cual llega y `clave_cliente` SOLO decide el ancho
de la columna A, POR CONTENIDO y con tope (`_formato.ancho_columna_clave`: 48
en balances, 40 en Dinero). Sin la llave (API previo o migración sin aplicar)
el libro sale byte a byte el de antes: las huellas `_firma_hoja` de abajo se
congelaron con el código de HEAD 462ce75 ANTES de tocar el render (idénticas
con d58faa1). Ningún importe, T.C. ni fórmula cambia.
"""

from __future__ import annotations

import copy
import hashlib
import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAvionRequest,
    BalanceAvionVuelo,
    BalanceGeneralRequest,
    BalanceOtroMovimientoFila,
    DineroOtroIngresoFila,
    DineroVueloFila,
    DineroXlsxRequest,
)
from app.services._formato import (
    ANCHO_CLAVE_TOPE_BALANCE,
    ANCHO_CLAVE_TOPE_DINERO,
    ancho_columna_clave,
    hay_id_cliente,
)
from app.services.balance_avion_xlsx import (
    libro_balance_avion,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from app.services.dinero_xlsx import render_dinero_xlsx

CLAVE_376 = "#376 · january02-211026 · January Davila"  # 40 caracteres
CLAVE_364 = "#364 · beh05-011026 · Beh Kay "  # con el espacio final del cliente
CLAVE_DINERO_MULTI = "january02-211026 (50 % N4142R)"

# ---------------------------------------------------------------------------
# Payloads (los MISMOS con los que se congelaron las huellas)
# ---------------------------------------------------------------------------


def _vuelo(clave, **extra):
    return {
        "clave": clave,
        "fecha": "2026-10-21",
        "ruta": "CUN-PTU-CUN",
        "estado": "COMPLETADO",
        "status_cobro": "Cobrado",
        "cobros": [{"fecha": "2026-10-21", "monto_mxn": 30000.0}],
        "cobrado_real_mxn": 30000.0,
        "cobrado_mxn": 30000.0,
        "por_cobrar_mxn": 0.0,
        "por_cobrar_usd": 0.0,
        **extra,
    }


def _vuelos():
    return [
        _vuelo(CLAVE_376, clave_cliente="january02-211026"),
        _vuelo(CLAVE_364, clave_cliente="beh05-011026"),
        # Sin ID (cliente sin alias / vuelo sin número): el API no manda la llave.
        _vuelo("#300 · Cliente Sin Id"),
    ]


_TOTALES = {"cobrado_real_mxn": 90000.0, "por_cobrar_usd": 0.0}


def _individual():
    return {
        "matricula": "XA-VGV",
        "periodo_desde": "2026-10-01",
        "periodo_hasta": "2026-10-31",
        "vuelos": _vuelos(),
        "totales": _TOTALES,
    }


def _general(variante=None):
    p = {
        "periodo_desde": "2026-10-01",
        "periodo_hasta": "2026-10-31",
        "resumen": [{"matricula": "XA-VGV", "color": "#FF0000", "vuelos": 3}],
        "consolidado": {
            "matricula": "FLOTA",
            "periodo_desde": "2026-10-01",
            "periodo_hasta": "2026-10-31",
            "vuelos": _vuelos(),
            "totales": _TOTALES,
            "otros_movimientos": {
                "filas": [
                    {
                        "clave": "#376 · january02-211026",
                        "estado": "CONFIRMADO",
                        "concepto_ingreso": "TUA CUN",
                        "ingreso_mxn": 1250.5,
                        "remanente_mxn": 1250.5,
                        "clave_cliente": "january02-211026",
                    },
                    {
                        "clave": "#115 · riviera02-050826",
                        "estado": "CANCELADO",
                        "concepto_ingreso": "Pernocta",
                        "ingreso_mxn": 300.0,
                        "remanente_mxn": 300.0,
                        "clave_cliente": "riviera02-050826",
                    },
                ],
                "filas_sueltas": [
                    {
                        "clave": "ING-3",
                        "concepto_ingreso": "Intereses",
                        "ingreso_mxn": 85.33,
                        "remanente_mxn": 85.33,
                    }
                ],
            },
        },
    }
    if variante is not None:
        p["variante"] = variante
    return p


def _dinero():
    return {
        "periodo_desde": "2026-10-01",
        "periodo_hasta": "2026-10-31",
        "generado": "2026-10-31",
        "vuelos": [
            {
                "clave": CLAVE_DINERO_MULTI,
                "matricula": "N4142R",
                "color": "#1D4ED8",
                "fecha": "2026-10-21",
                "ruta": "CUN-PTU",
                "tiempo": 1.0,
                "participacion": 0.5,
                "multi_avion": True,
                "total_cobrado_mxn": 10000.0,
                "clave_cliente": "january02-211026",
            },
            {
                # Vuelo de servicio: 'vtservicio' manda aunque haya ID.
                "clave": "vtservicio",
                "matricula": "XA-VGV",
                "color": "#FF0000",
                "fecha": "2026-10-22",
                "ruta": "CUN-MID",
                "tiempo": 1.2,
                "total_cobrado_mxn": 0.0,
                "clave_cliente": "january03-221026",
            },
        ],
        "otros_ingresos": [
            {
                "clave": "january02-211026",
                "fecha_vuelo": "2026-10-21",
                "concepto_ingreso": "TUA CUN",
                "ingreso_mxn": 500.0,
                "remanente_mxn": 500.0,
                "clave_cliente": "january02-211026",
            },
            {
                "clave": "ING-7",
                "concepto_ingreso": "Intereses",
                "ingreso_mxn": 12.5,
                "remanente_mxn": 12.5,
            },
        ],
    }


def _sin_clave_cliente(o):
    """COPIA sin ninguna llave `clave_cliente` (API ≤ 0.0.74)."""
    o = copy.deepcopy(o)

    def quita(x):
        if isinstance(x, dict):
            x.pop("clave_cliente", None)
            for v in x.values():
                quita(v)
        elif isinstance(x, list):
            for v in x:
                quita(v)

    quita(o)
    return o


def _xlsx_individual(payload):
    return render_balance_avion_xlsx(BalanceAvionRequest.model_validate(payload))


def _xlsx_general(payload):
    return render_balance_general_xlsx(BalanceGeneralRequest.model_validate(payload))


def _xlsx_dinero(payload):
    return render_dinero_xlsx(DineroXlsxRequest.model_validate(payload))


def _ws(data: bytes, nombre: str):
    return load_workbook(BytesIO(data))[nombre]


def _ancho_a(data: bytes, nombre: str) -> float:
    return _ws(data, nombre).column_dimensions["A"].width


# ---------------------------------------------------------------------------
# Huellas del LAYOUT (receta `_firma_hoja` de test_balance_factura_vuelatour,
# COMPLETA: aquí no se quitan las constantes A1:D1 — el código congelado ya
# las pintaba)
# ---------------------------------------------------------------------------


def _firma_hoja(data: bytes, nombre: str) -> str:
    ws = load_workbook(BytesIO(data), data_only=True)[nombre]
    partes: list = []
    for fila in ws.iter_rows():
        for c in fila:
            f = c.font
            partes.append([
                c.coordinate, repr(c.value), c.number_format,
                bool(f.b), bool(f.i), f.sz,
                f.color.rgb if f.color is not None else None,
                c.fill.fgColor.rgb if c.fill.fill_type else None,
                c.alignment.horizontal, c.alignment.vertical,
                bool(c.alignment.wrap_text), c.border.left.style,
                c.comment.text if c.comment else None,
            ])  # fmt: skip
    partes.append(sorted(str(r) for r in ws.merged_cells.ranges))
    partes.append(sorted((k, d.width) for k, d in ws.column_dimensions.items()))
    partes.append(sorted((k, d.height) for k, d in ws.row_dimensions.items() if d.height))
    partes.append(ws.freeze_panes)
    return hashlib.sha256(json.dumps(partes, default=str).encode()).hexdigest()


def _valores(data: bytes, nombre: str) -> list:
    """Valores (caché) y fórmulas de TODA la hoja: lo que no puede moverse."""
    con_f = load_workbook(BytesIO(data))[nombre]
    con_v = load_workbook(BytesIO(data), data_only=True)[nombre]
    return [
        (c.coordinate, c.value, con_v[c.coordinate].value, c.number_format)
        for fila in con_f.iter_rows()
        for c in fila
    ]


# Calculadas con el código de HEAD 462ce75 (13-oct-2026) ANTES de tocar el
# render, sobre los payloads de arriba SIN `clave_cliente` (con d58faa1 dan
# lo mismo). Si una cambia, un API ≤ 0.0.74 ya no ve su libro de siempre.
_HUELLAS_HEAD = {
    ("individual", "reporte horas XA-VGV"): (
        "fbe4eea70927ba2e6b26b84b53da168667cae9faeabb211d37dc748af60be6f8"
    ),
    ("individual", "cobranza"): "c892b195a7967a39b353243360bd46bb189663cd9fb68602b1910b96f23251be",
    ("general", "reporte horas FLOTA"): (
        "fc120648f1a6298df3f3915a06f8ff1908f3a295e88f3d0827d2832493cfacc5"
    ),
    ("general", "cobranza"): "a9abb850ae6c1035315842cb40232ebcd55955f02270cf5f4b514686f6ddc30c",
    ("general", "otros movimientos"): (
        "854facc0d07c53f06735e11fd320cb58eeccb01b1860cc3561f1701607188cdc"
    ),
    ("general-costo", "reporte horas FLOTA"): (
        "f6382b0629040ee7fbb403e8cb1bf6ec99b8fd5171a6cc64031c4f9b856eb3ec"
    ),
    ("dinero", "dinero-vlos"): "f9248fd7d4742c431ba453c3277c5739c1f3628f7d98dc5e1de0f1f7755c0cda",
    (
        "dinero",
        "Otros ingresos",
    ): "3b64cbbfc1b51ccfd6413cd16eddf4af8c83ec484a3f062080002f96f66b874d",
}


def _render(libro: str, *, con_id: bool) -> bytes:
    payload = {
        "individual": _individual,
        "general": _general,
        "general-costo": lambda: _general("general"),
        "dinero": _dinero,
    }[libro]()
    if not con_id:
        payload = _sin_clave_cliente(payload)
    if libro == "individual":
        return _xlsx_individual(payload)
    if libro == "dinero":
        return _xlsx_dinero(payload)
    return _xlsx_general(payload)


# ---------------------------------------------------------------------------
# _formato
# ---------------------------------------------------------------------------


def test_ancho_columna_clave_sin_id_es_la_base_exacta_como_int():
    for base in (14, 15, 16, 18, 22):
        ancho = ancho_columna_clave(["x" * 90], base=base, tope=48, con_id_cliente=False)
        assert ancho == base and type(ancho) is int


def test_ancho_columna_clave_con_id_por_contenido_con_tope():
    assert len(CLAVE_376) == 40
    tope = ANCHO_CLAVE_TOPE_BALANCE
    assert ancho_columna_clave([CLAVE_376], base=16, tope=tope, con_id_cliente=True) == 42
    assert ancho_columna_clave(["x" * 60], base=16, tope=tope, con_id_cliente=True) == 48
    # Texto corto: la base (nunca se ANGOSTA).
    assert ancho_columna_clave(["#1 · a01"], base=16, tope=tope, con_id_cliente=True) == 16
    # Sin textos (o vacíos / None): la base.
    assert ancho_columna_clave([], base=22, tope=tope, con_id_cliente=True) == 22
    assert ancho_columna_clave([None, ""], base=22, tope=tope, con_id_cliente=True) == 22
    # El texto más largo manda; «·» cuenta 1.
    assert ancho_columna_clave(["ab", "a · b · c"], base=1, tope=tope, con_id_cliente=True) == 11
    assert ANCHO_CLAVE_TOPE_BALANCE == 48 and ANCHO_CLAVE_TOPE_DINERO == 40


def test_hay_id_cliente():
    assert hay_id_cliente([BalanceAvionVuelo(), BalanceAvionVuelo(clave_cliente="x01")])
    assert not hay_id_cliente([BalanceAvionVuelo(), BalanceAvionVuelo(clave_cliente=None)])
    assert not hay_id_cliente([])
    # Filas sin el atributo (otro modelo): cuentan como no.
    assert not hay_id_cliente([object(), {"clave_cliente": "x01"}])


# ---------------------------------------------------------------------------
# Balance individual: hoja maestra y cobranza
# ---------------------------------------------------------------------------


def test_individual_clave_tal_cual_y_columna_a_por_contenido():
    data = _render("individual", con_id=True)
    ws = _ws(data, "reporte horas XA-VGV")
    assert ws["A3"].value == CLAVE_376
    assert ws["A4"].value == CLAVE_364  # el espacio final se conserva
    assert _ancho_a(data, "reporte horas XA-VGV") == 42
    assert _ancho_a(data, "cobranza") == 42
    sin = _render("individual", con_id=False)
    assert _ancho_a(sin, "reporte horas XA-VGV") == 16
    assert _ancho_a(sin, "cobranza") == 22


def test_tope_del_ancho_con_un_cliente_larguisimo():
    payload = _individual()
    payload["vuelos"][0]["clave"] = "#376 · january02-211026 · " + "Nombre Larguísimo " * 4
    data = _xlsx_individual(payload)
    assert _ancho_a(data, "reporte horas XA-VGV") == 48
    assert _ancho_a(data, "cobranza") == 48


# ---------------------------------------------------------------------------
# Balance general: otros movimientos, reporte horas FLOTA y cobranza
# ---------------------------------------------------------------------------


def test_general_otros_movimientos_mide_el_texto_pintado():
    data = _render("general", con_id=True)
    ws = _ws(data, "otros movimientos")
    pintadas = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
    assert "#376 · january02-211026" in pintadas  # CONFIRMADO: estado normal
    assert "#115 · riviera02-050826 · CANCELADO" in pintadas
    assert _ancho_a(data, "otros movimientos") == len("#115 · riviera02-050826 · CANCELADO") + 2
    assert _ancho_a(data, "otros movimientos") == 37
    # La hoja de vuelos de la flota y la cobranza del general también.
    assert _ancho_a(data, "reporte horas FLOTA") == 42
    assert _ancho_a(data, "cobranza") == 42
    sin = _render("general", con_id=False)
    assert _ancho_a(sin, "otros movimientos") == 18
    assert _ancho_a(sin, "reporte horas FLOTA") == 16
    assert _ancho_a(sin, "cobranza") == 22


def test_general_variante_costo_por_hora_tambien_se_ensancha():
    assert _ancho_a(_render("general-costo", con_id=True), "reporte horas FLOTA") == 42
    assert _ancho_a(_render("general-costo", con_id=False), "reporte horas FLOTA") == 16


# ---------------------------------------------------------------------------
# Libro Dinero
# ---------------------------------------------------------------------------


def test_dinero_vuelos_y_otros_ingresos():
    data = _render("dinero", con_id=True)
    ws = _ws(data, "dinero-vlos")
    assert ws["A5"].value == CLAVE_DINERO_MULTI
    assert ws["A6"].value == "vtservicio"
    assert _ancho_a(data, "dinero-vlos") == len(CLAVE_DINERO_MULTI) + 2 == 32
    assert _ancho_a(data, "Otros ingresos") == len("january02-211026") + 2 == 18
    sin = _render("dinero", con_id=False)
    assert _ancho_a(sin, "dinero-vlos") == 14
    assert _ancho_a(sin, "Otros ingresos") == 15
    # La nota de la celda CLAVE (matrícula / multi-avión) no cambia.
    com_con = _ws(data, "dinero-vlos")["A5"].comment
    com_sin = _ws(sin, "dinero-vlos")["A5"].comment
    assert com_con is not None and com_con.text == com_sin.text


def test_dinero_tope_40():
    payload = _dinero()
    payload["vuelos"][0]["clave"] = "january02-211026 (50 % N4142R) " + "x" * 30
    assert _ancho_a(_xlsx_dinero(payload), "dinero-vlos") == ANCHO_CLAVE_TOPE_DINERO == 40


# ---------------------------------------------------------------------------
# Sin la llave: byte a byte el libro de antes; con la llave: solo el ancho
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(("libro", "hoja"), sorted(_HUELLAS_HEAD))
def test_sin_la_llave_las_huellas_son_las_de_head(libro, hoja):
    assert _firma_hoja(_render(libro, con_id=False), hoja) == _HUELLAS_HEAD[(libro, hoja)]


@pytest.mark.parametrize(("libro", "hoja"), sorted(_HUELLAS_HEAD))
def test_con_la_llave_ningun_valor_formula_ni_formato_cambia(libro, hoja):
    con, sin = _render(libro, con_id=True), _render(libro, con_id=False)
    assert _valores(con, hoja) == _valores(sin, hoja)
    # Solo cambia el ancho de la columna A.
    dims_con = {k: d.width for k, d in _ws(con, hoja).column_dimensions.items() if k != "A"}
    dims_sin = {k: d.width for k, d in _ws(sin, hoja).column_dimensions.items() if k != "A"}
    assert dims_con == dims_sin
    assert _ancho_a(con, hoja) > _ancho_a(sin, hoja)


# ---------------------------------------------------------------------------
# Libro JSON de la vista en pantalla: hereda el ancho sin tocarlo
# ---------------------------------------------------------------------------


def test_libro_json_hereda_el_ancho():
    libro = libro_balance_avion(BalanceAvionRequest.model_validate(_individual()))
    hoja = next(h for h in libro["hojas"] if h["nombre"] == "reporte horas XA-VGV")
    assert {"c": 1, "ancho": 42.0} in hoja["columnas"]


# ---------------------------------------------------------------------------
# Esquema: aditivo y liberal
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "modelo", [BalanceAvionVuelo, BalanceOtroMovimientoFila, DineroVueloFila, DineroOtroIngresoFila]
)
def test_esquema_acepta_la_llave_default_none_y_liberal(modelo):
    assert modelo().clave_cliente is None
    assert modelo.model_validate({"clave_cliente": "beh05-011026"}).clave_cliente == "beh05-011026"
    assert modelo.model_validate({"clave_cliente": " beh05 "}).clave_cliente == "beh05"
    # Metadato: jamás un 422 (no-texto o vacío ⇒ None) y las llaves
    # desconocidas se siguen ignorando.
    for basura in (7, "", "   ", ["x"], None):
        assert modelo.model_validate({"clave_cliente": basura}).clave_cliente is None
    assert modelo.model_validate({"clave_cliente": "a01", "llave_futura": 1}).clave_cliente == "a01"


# ---------------------------------------------------------------------------
# Por HTTP: las rutas aceptan el payload del API 0.0.75
# ---------------------------------------------------------------------------

TOKEN = "tok-clave-cliente"
client = TestClient(app)


def test_rutas_aceptan_clave_cliente(monkeypatch, request):
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    res = client.post("/pdf/balance-avion-xlsx", json=_individual(), headers=headers)
    assert res.status_code == 200, res.text
    assert _ancho_a(res.content, "reporte horas XA-VGV") == 42
    res = client.post("/pdf/balance-general-xlsx", json=_general(), headers=headers)
    assert res.status_code == 200, res.text
    assert _ancho_a(res.content, "otros movimientos") == 37
    res = client.post("/pdf/dinero-xlsx", json=_dinero(), headers=headers)
    assert res.status_code == 200, res.text
    assert _ancho_a(res.content, "dinero-vlos") == 32
    res = client.post("/pdf/balance-avion-libro", json=_individual(), headers=headers)
    assert res.status_code == 200, res.text
