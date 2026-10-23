"""Reparto a socios (PDF y Excel) con el COMBUSTIBLE SEPARADO de los gastos
directos, en el orden del Excel «Balance» del cliente (23-oct-2026, API
0.0.103, invariante 104 del API).

Pedido del cliente: «debemos separar gastos directos del combustible, porque
los está uniendo, porque no nos está coincidiendo con el excel». Con las
llaves ADITIVAS `combustible_usd/_mxn` y `utilidad_antes_gastos_usd/_mxn`
(`gastos_directos_*` ya SIN combustible) la cascada se pinta como el Excel:
venta − comisiones − directos − provisión AFAC = utilidad antes de gastos −
combustible − indirectos − permisos − fijos + otros ingresos = saldo. Caso
real N621TX sep-2026: 787,851.13 − 2,591.00 − 40,588.14 − 42,575.49 =
702,096.50; combustible 147,277.58 (antes la card decía «Gastos directos
187,865.72» = 147,277.58 + 40,588.14). Aquí no se calcula nada: se pintan los
números del API (la utilidad antes de gastos solo se re-suma si el API no la
manda). Sin `combustible_usd` (API ≤ 0.0.102, o la llave en null), PDF y
Excel BYTE A BYTE los de siempre (huellas del HEAD 3b893f1 de
`test_reparto_en_pesos` y flujos/miembros contra el payload sin las llaves).
"""

from __future__ import annotations

import copy
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth

from app.config import get_settings
from app.main import app
from app.schemas.reparto import RepartoAvion, RepartoPdfRequest
from app.services import reparto_combustible as rc
from app.services import reparto_pdf as rp
from app.services import reparto_xlsx as rx
from app.services.reparto_comisiones import ETIQUETA_COLUMNA, ETIQUETA_FILA
from app.services.reparto_pdf import render_reparto_pdf
from app.services.reparto_xlsx import render_reparto_xlsx
from tests import _pesos_payloads as pp
from tests import test_reparto_como_balance as tcb
from tests.test_reparto_en_pesos import (
    _HUELLA_PDF_HEAD,
    _HUELLA_XLSX_HEAD,
    _huella_pdf,
    _huella_xlsx,
    _texto_pdf,
)

r2 = pp.r2
LLAVES = (
    "combustible_usd",
    "combustible_mxn",
    "utilidad_antes_gastos_usd",
    "utilidad_antes_gastos_mxn",
)

# Caso real N621TX, septiembre 2026 (USD): la cascada del Excel del cliente.
VENTA, COMISIONES, DIRECTOS, AFAC = 787_851.13, 2_591.00, 40_588.14, 42_575.49
UTILIDAD_ANTES = 702_096.50
COMBUSTIBLE, INDIRECTOS, PERMISOS, FIJOS = 147_277.58, 12_345.67, 1_000.00, 0.0
SALDO = 541_473.25  # = 702,096.50 − 147,277.58 − 12,345.67 − 1,000.00
RESERVA = 15_000.0
TARIFA_AFAC, HORAS_COBRADAS = 310.77, 137.0  # 310.77 × 137 = 42,575.49
ETIQUETA_AFAC = "(-) Provisión AFAC (310.77 USD/hr × 137 h cobradas)"
FIJOS_PDF = "(-) Fijos (prorrateo + reparto manual) $-0.00"
K = 18.0  # T.C. con que se arma el espejo en pesos (los pesos los manda el API)


def _d(x: float) -> str:
    return f"${x:,.2f}"


def _mxn(x: float) -> float:
    return r2(x * K)


def _n621tx(*, otros_ingresos: float = 0.0) -> dict:
    """N621TX con la forma del API 0.0.103 (USD y pesos, cascada del balance):
    `gastos_directos_*` SIN combustible, `combustible_*` aparte y
    `utilidad_antes_gastos_*` = venta − comisiones − directos − AFAC."""
    saldo = r2(UTILIDAD_ANTES - COMBUSTIBLE - INDIRECTOS - PERMISOS - FIJOS + otros_ingresos)
    utilidad_mxn = r2(_mxn(VENTA) - _mxn(COMISIONES) - _mxn(DIRECTOS) - _mxn(AFAC))
    saldo_mxn = r2(
        utilidad_mxn
        - _mxn(COMBUSTIBLE)
        - _mxn(INDIRECTOS)
        - _mxn(PERMISOS)
        - _mxn(FIJOS)
        + _mxn(otros_ingresos)
    )
    reparto: list[dict] = []
    retenida_usd = retenida_mxn = 0.0
    for nombre, pct, descuenta in (("Socio A", 60, True), ("Socio B", 40, False)):
        u_usd, u_mxn = r2(saldo * pct / 100), r2(saldo_mxn * pct / 100)
        rs_usd, rs_mxn = r2(RESERVA * pct / 100), r2(_mxn(RESERVA) * pct / 100)
        reparto.append(
            {
                "socio_nombre": nombre,
                "porcentaje": pct,
                "monto_usd": r2(u_usd - rs_usd) if descuenta else u_usd,
                "monto_mxn": r2(u_mxn - rs_mxn) if descuenta else u_mxn,
                "utilidad_usd": u_usd,
                "utilidad_mxn": u_mxn,
                "reserva_usd": rs_usd,
                "reserva_mxn": rs_mxn,
                "descuenta_reserva": descuenta,
            }
        )
        if descuenta:
            retenida_usd, retenida_mxn = r2(retenida_usd + rs_usd), r2(retenida_mxn + rs_mxn)
    detalle = (
        [
            {
                "clave": "ING-9", "fecha": "2026-09-12",
                "categoria": "Reembolsos y devoluciones recibidos",
                "concepto": "Reembolso de servicio", "monto": otros_ingresos, "moneda": "USD",
                "comision": 0, "usd": otros_ingresos, "mxn": _mxn(otros_ingresos),
            }
        ]
        if otros_ingresos
        else []
    )  # fmt: skip
    return {
        "matricula": "N621TX", "modelo": "Cessna Citation",
        "ingresos_cobrado_usd": VENTA, "ingresos_cobrado_mxn": _mxn(VENTA),
        "otros_ingresos_vuelatour_usd": 0, "otros_ingresos_vuelatour_desglose": None,
        "comisiones_venta_usd": COMISIONES, "comisiones_venta_mxn": _mxn(COMISIONES),
        "pendiente_cobro_usd": 0, "pendiente_cobro_mxn": 0, "pendiente_bruto_usd": 0,
        "horas_voladas_hr": 139.4,
        "gastos_directos_usd": DIRECTOS, "gastos_directos_mxn": _mxn(DIRECTOS),
        "combustible_usd": COMBUSTIBLE, "combustible_mxn": _mxn(COMBUSTIBLE),
        "utilidad_antes_gastos_usd": UTILIDAD_ANTES, "utilidad_antes_gastos_mxn": utilidad_mxn,
        "gastos_indirectos_usd": INDIRECTOS, "gastos_indirectos_mxn": _mxn(INDIRECTOS),
        "permisos_usd": PERMISOS, "permisos_mxn": _mxn(PERMISOS),
        "otros_usd": FIJOS, "otros_mxn": _mxn(FIJOS),
        "reserva_overhaul_usd": RESERVA, "reserva_overhaul_mxn": _mxn(RESERVA),
        "reserva_tc": K,
        "afac_reservado_usd": AFAC, "afac_reservado_mxn": _mxn(AFAC),
        "afac_tarifa_usd_hr": TARIFA_AFAC, "afac_horas_cobradas_hr": HORAS_COBRADAS,
        "reserva_retenida_usd": retenida_usd, "reserva_retenida_mxn": retenida_mxn,
        "otros_ingresos_avion_usd": otros_ingresos,
        "otros_ingresos_avion_mxn": _mxn(otros_ingresos),
        "otros_ingresos_avion_sin_tc_mxn": 0, "otros_ingresos_avion": detalle,
        "saldo_usd": saldo, "saldo_mxn": saldo_mxn, "pesos_sin_tc_usd": 0,
        "gastos_sin_tc_mxn": 0, "cobros_sin_tc_mxn": 0, "gastos_tc_oficial": None,
        "cobros_tc_oficial_count": 0, "reserva_incompleta": False,
        "reparto_porcentaje_total": 100, "reparto": reparto, "vuelos": [],
        "gastos_por_categoria": [],
    }  # fmt: skip


def _payload(*, pesos: bool = True, como_balance: bool = True, **avion) -> dict:
    """El reparto del API 0.0.103 (las dos banderas en la raíz) con N621TX."""
    return {
        "periodo_desde": "2026-09-01", "periodo_hasta": "2026-09-30", "generado": "2026-10-23",
        "otros_ingresos_vuelatour_total_usd": 0, "otros_ingresos_vuelatour_desglose": None,
        "tc_oficial": None, "comision_vendedor_solo_gasto_real": True,
        "reparto_como_balance": como_balance, "reparto_en_pesos": pesos,
        "aviones": [_n621tx(**avion)],
    }  # fmt: skip


def _req(**kw) -> RepartoPdfRequest:
    return RepartoPdfRequest.model_validate(_payload(**kw))


def _sin_llaves(p: dict, *, nulos: bool) -> dict:
    """El mismo payload sin las 4 llaves (API previo) o con ellas en null."""
    p = copy.deepcopy(p)
    for a in p["aviones"]:
        for k in LLAVES:
            if nulos:
                a[k] = None
            else:
                a.pop(k, None)
    return p


def _flujos(data: bytes) -> list[bytes]:
    return [p.get_contents().get_data() for p in PdfReader(BytesIO(data)).pages]


def _miembros(data: bytes) -> dict[str, bytes]:
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _hoja(p: dict):
    return load_workbook(BytesIO(render_reparto_xlsx(RepartoPdfRequest.model_validate(p)))).active


def _fila(ws, texto: str, col: int = 1) -> int:
    filas = [c.row for c in ws[chr(64 + col)] if c.value == texto]
    assert len(filas) == 1, (texto, filas)
    return filas[0]


def _textos_col_a(ws) -> list[str]:
    return [c.value for c in ws["A"] if isinstance(c.value, str)]


def _encabezados(ws, fila: int = 4) -> list[str]:
    return [c.value for c in ws[fila] if c.value is not None]


def test_los_numeros_del_caso_cuadran() -> None:
    """El caso del cliente: 787,851.13 − 2,591.00 − 40,588.14 − 42,575.49 =
    702,096.50 y el saldo = utilidad antes − combustible − indirectos −
    permisos − fijos; la card unía 147,277.58 + 40,588.14 = 187,865.72."""
    assert r2(VENTA - COMISIONES - DIRECTOS - AFAC) == UTILIDAD_ANTES
    assert r2(UTILIDAD_ANTES - COMBUSTIBLE - INDIRECTOS - PERMISOS - FIJOS) == SALDO
    assert r2(TARIFA_AFAC * HORAS_COBRADAS) == AFAC
    assert r2(COMBUSTIBLE + DIRECTOS) == 187_865.72
    a = _req().aviones[0]
    assert (a.saldo_usd, a.utilidad_antes_gastos_usd, a.combustible_usd) == (
        SALDO, UTILIDAD_ANTES, COMBUSTIBLE,
    )  # fmt: skip
    assert a.saldo_mxn == r2(
        a.utilidad_antes_gastos_mxn
        - a.combustible_mxn
        - a.gastos_indirectos_mxn
        - a.permisos_mxn
        - a.otros_mxn
    )


# ---------------------------------------------------------------------------
# (a) Esquema y helpers: aditivos, con default None, la PRESENCIA decide.
# ---------------------------------------------------------------------------


def test_esquema_aditivo_con_default_none() -> None:
    a = RepartoAvion(matricula="X", modelo="M", ingresos_cobrado_usd=0, saldo_usd=0)
    for llave in LLAVES:
        assert getattr(a, llave) is None
    assert not rc.combustible_separado(a)
    # 0.0 también cuenta: la PRESENCIA de la llave decide.
    a.combustible_usd = 0.0
    assert rc.combustible_separado(a)
    sin = RepartoPdfRequest.model_validate(_sin_llaves(_payload(), nulos=False))
    assert not rc.hay_combustible_separado(sin)
    assert rc.hay_combustible_separado(_req())


def test_utilidad_antes_de_gastos_del_api_o_resumada() -> None:
    a = _req().aviones[0]
    assert rc.utilidad_antes_gastos_usd(a, como_balance=True) == UTILIDAD_ANTES
    assert rc.utilidad_antes_gastos_mxn(a, como_balance=True) == a.utilidad_antes_gastos_mxn
    # Sin la del API: la resta de los componentes (la provisión AFAC solo
    # con la cascada del balance).
    a.utilidad_antes_gastos_usd = a.utilidad_antes_gastos_mxn = None
    assert rc.utilidad_antes_gastos_usd(a, como_balance=True) == UTILIDAD_ANTES
    assert rc.utilidad_antes_gastos_usd(a, como_balance=False) == r2(UTILIDAD_ANTES + AFAC)
    assert rc.utilidad_antes_gastos_mxn(a, como_balance=True) == r2(
        _mxn(VENTA) - _mxn(COMISIONES) - _mxn(DIRECTOS) - _mxn(AFAC)
    )
    # Un componente en pesos sin dato ⇒ None (jamás un peso inventado).
    a.gastos_directos_mxn = None
    assert rc.utilidad_antes_gastos_mxn(a, como_balance=True) is None


def test_etiquetas_caben_y_sin_caracteres_que_helvetica_no_dibuja() -> None:
    ancho = 112 * mm - 12  # la columna de etiquetas más angosta (pesos)
    for t in (
        rc.FILA_GASTOS_DIRECTOS_SIN_COMBUSTIBLE,
        rc.FILA_UTILIDAD_ANTES_GASTOS,
        rc.FILA_COMBUSTIBLE,
    ):
        assert stringWidth(t, "Helvetica-Bold", 9) <= ancho, t
        assert "−" not in t and "→" not in t
    pie = rc.pie_combustible(_req())
    assert "−" not in pie and "→" not in pie
    assert pie.endswith("- comisiones - gastos directos - provisión AFAC.")
    assert rc.pie_combustible(_req(como_balance=False)).endswith("- gastos directos.")


# ---------------------------------------------------------------------------
# (b) Sin la llave (ausente o en null): byte a byte.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("pesos", [None, False, "si"])
def test_sin_la_llave_las_huellas_de_head(pesos) -> None:
    """Las huellas congeladas de `test_reparto_en_pesos` (HEAD 3b893f1) con
    las llaves del 0.0.103 presentes en null: nada se mueve."""
    p = pp.reparto(bandera=pesos)
    for a in p["aviones"]:
        a.update({k: None for k in LLAVES})
    req = RepartoPdfRequest.model_validate(p)
    assert _huella_pdf(render_reparto_pdf(req)) == _HUELLA_PDF_HEAD
    assert _huella_xlsx(render_reparto_xlsx(req)) == _HUELLA_XLSX_HEAD


@pytest.mark.parametrize("pesos", [None, True])
def test_llaves_en_null_no_cambian_la_cascada_del_balance(pesos) -> None:
    """El reparto del API 0.0.101 (cascada del balance, USD y pesos) con las
    llaves en null: flujos del PDF y miembros del .xlsx idénticos a sin
    ellas."""
    base = tcb._payload(pesos=pesos)
    con_nulos = copy.deepcopy(base)
    for a in con_nulos["aviones"]:
        a.update({k: None for k in LLAVES})
    pdf_base = render_reparto_pdf(RepartoPdfRequest.model_validate(base))
    xlsx_base = render_reparto_xlsx(RepartoPdfRequest.model_validate(base))
    req = RepartoPdfRequest.model_validate(con_nulos)
    assert _flujos(render_reparto_pdf(req)) == _flujos(pdf_base)
    assert _miembros(render_reparto_xlsx(req)) == _miembros(xlsx_base)
    texto = _texto_pdf(pdf_base)
    for t in (rc.FILA_COMBUSTIBLE, rc.FILA_UTILIDAD_ANTES_GASTOS, "(sin combustible)"):
        assert t not in texto


def test_sin_la_llave_el_n621tx_sale_con_la_cascada_de_siempre() -> None:
    """Un API ≤ 0.0.102 mandaba los directos CON el combustible (187,865.72):
    el PDF y el Excel de siempre, con la provisión AFAC tras los fijos y al
    final de la tabla."""
    p = _sin_llaves(_payload(pesos=False), nulos=False)
    a = p["aviones"][0]
    a["gastos_directos_usd"] = r2(DIRECTOS + COMBUSTIBLE)
    texto = _texto_pdf(render_reparto_pdf(RepartoPdfRequest.model_validate(p)))
    assert (
        f"(-) Gastos directos {_d(-187_865.72)} (-) Gastos indirectos {_d(-INDIRECTOS)} "
        f"(-) Permisos {_d(-PERMISOS)} {FIJOS_PDF} {ETIQUETA_AFAC} {_d(-AFAC)} "
        f"Saldo disponible {_d(SALDO)}"
    ) in texto
    assert rc.FILA_COMBUSTIBLE not in texto and rc.FILA_UTILIDAD_ANTES_GASTOS not in texto
    assert "va en su propia línea" not in texto
    ws = _hoja(p)
    assert ws["G4"].value == "Gastos directos" and ws["O4"].value == rx.ENCABEZADO_PROVISION_AFAC
    assert ws["G5"].value == 187_865.72
    assert rc.ENCABEZADO_COMBUSTIBLE not in _encabezados(ws)
    assert rc.NOTA_COMBUSTIBLE_SEPARADO not in _textos_col_a(ws)


# ---------------------------------------------------------------------------
# (c) PDF en dólares.
# ---------------------------------------------------------------------------


def test_pdf_usd_cascada_en_el_orden_del_excel() -> None:
    texto = _texto_pdf(render_reparto_pdf(_req(pesos=False)))
    assert (
        f"Venta del avión cobrada {_d(VENTA)} "
        f"{ETIQUETA_FILA} {_d(-COMISIONES)} "
        f"{rc.FILA_GASTOS_DIRECTOS_SIN_COMBUSTIBLE} {_d(-DIRECTOS)} "
        f"{ETIQUETA_AFAC} {_d(-AFAC)} "
        f"{rc.FILA_UTILIDAD_ANTES_GASTOS} {_d(UTILIDAD_ANTES)} "
        f"{rc.FILA_COMBUSTIBLE} {_d(-COMBUSTIBLE)} "
        f"(-) Gastos indirectos {_d(-INDIRECTOS)} "
        f"(-) Permisos {_d(-PERMISOS)} "
        f"{FIJOS_PDF} "
        f"Saldo disponible {_d(SALDO)} "
        f"{rp.FILA_RESERVA_INFORMATIVA} {_d(RESERVA)}"
    ) in texto
    # Ni la fila vieja (con el combustible unido) ni su número.
    assert "(-) Gastos directos $" not in texto
    assert "187,865.72" not in texto
    # Resumen global: «Gastos del periodo» suma el combustible (venta −
    # gastos = saldo).
    gastos = r2(COMISIONES + DIRECTOS + AFAC + COMBUSTIBLE + INDIRECTOS + PERMISOS + FIJOS)
    assert r2(VENTA - gastos) == SALDO
    assert f"{_d(VENTA)} $0.00 {_d(gastos)} {_d(SALDO)}" in texto
    # El pie lo explica, entre la cláusula de la cascada del balance y
    # «Montos en USD».
    assert (
        "VuelaTour la retiene. El combustible (gasavión/turbosina) va en su propia línea, "
        "fuera de los gastos directos, como en el Balance: Utilidad antes de gastos = venta "
        "del avión cobrada - comisiones - gastos directos - provisión AFAC. Montos en USD."
    ) in texto


def test_pdf_usd_otros_ingresos_del_avion_tras_los_fijos() -> None:
    texto = _texto_pdf(render_reparto_pdf(_req(pesos=False, otros_ingresos=1_000.0)))
    assert (
        f"{FIJOS_PDF} {rp.FILA_OTROS_INGRESOS_AVION} {_d(1000)} "
        f"ING-9 · Reembolsos y devoluciones recibidos · Reembolso de servicio {_d(1000)} "
        f"Saldo disponible {_d(SALDO + 1000)}"
    ) in texto
    # Ya no va antes de las comisiones (la primera «(-) Comisiones…» del
    # texto es la fila de la cascada).
    assert texto.index(f"{ETIQUETA_FILA} ") < texto.index(rp.FILA_OTROS_INGRESOS_AVION)
    # Resumen de 5 columnas: venta + otros ingresos − gastos = saldo.
    gastos = r2(COMISIONES + DIRECTOS + AFAC + COMBUSTIBLE + INDIRECTOS + PERMISOS + FIJOS)
    assert f"{_d(VENTA)} {_d(1000)} $0.00 {_d(gastos)} {_d(SALDO + 1000)}" in texto


def test_pdf_utilidad_antes_de_gastos_resumada_si_el_api_no_la_manda() -> None:
    p = _payload(pesos=False)
    p["aviones"][0]["utilidad_antes_gastos_usd"] = None
    texto = _texto_pdf(render_reparto_pdf(RepartoPdfRequest.model_validate(p)))
    assert f"{ETIQUETA_AFAC} {_d(-AFAC)} {rc.FILA_UTILIDAD_ANTES_GASTOS} {_d(UTILIDAD_ANTES)}" in (
        texto
    )


def test_pdf_sin_cascada_del_balance_la_reserva_sigue_restando() -> None:
    """Combustible aparte SIN `reparto_como_balance` (no lo manda el API
    0.0.103, pero el contrato lo tolera): el orden del Excel sin la fila de
    provisión AFAC, la utilidad antes de gastos sin restarla y la reserva
    de overhaul restando antes del saldo, como siempre."""
    p = _payload(pesos=False, como_balance=False)
    a = p["aviones"][0]
    a["utilidad_antes_gastos_usd"] = None
    a["saldo_usd"] = r2(
        VENTA - COMISIONES - DIRECTOS - COMBUSTIBLE - INDIRECTOS - PERMISOS - FIJOS - RESERVA
    )
    texto = _texto_pdf(render_reparto_pdf(RepartoPdfRequest.model_validate(p)))
    assert (
        f"{rc.FILA_GASTOS_DIRECTOS_SIN_COMBUSTIBLE} {_d(-DIRECTOS)} "
        f"{rc.FILA_UTILIDAD_ANTES_GASTOS} {_d(r2(VENTA - COMISIONES - DIRECTOS))} "
        f"{rc.FILA_COMBUSTIBLE} {_d(-COMBUSTIBLE)} (-) Gastos indirectos"
    ) in texto
    assert f"(-) Reserva overhaul {_d(-RESERVA)} Saldo disponible {_d(a['saldo_usd'])}" in texto
    assert "Provisión AFAC" not in texto
    assert "- comisiones - gastos directos. Montos en USD." in texto


# ---------------------------------------------------------------------------
# (c) PDF en pesos (las dos banderas, como las manda el API 0.0.103).
# ---------------------------------------------------------------------------


def test_pdf_pesos_cascada_en_el_orden_del_excel() -> None:
    req = _req()
    a = req.aviones[0]
    texto = _texto_pdf(render_reparto_pdf(req))
    assert rp.LEYENDA_PESOS in texto
    assert (
        f"{ETIQUETA_FILA} {_d(-_mxn(COMISIONES))} {_d(-COMISIONES)} "
        f"{rc.FILA_GASTOS_DIRECTOS_SIN_COMBUSTIBLE} {_d(-_mxn(DIRECTOS))} {_d(-DIRECTOS)} "
        f"{ETIQUETA_AFAC} {_d(-_mxn(AFAC))} {_d(-AFAC)} "
        f"{rc.FILA_UTILIDAD_ANTES_GASTOS} {_d(a.utilidad_antes_gastos_mxn)} {_d(UTILIDAD_ANTES)} "
        f"{rc.FILA_COMBUSTIBLE} {_d(-_mxn(COMBUSTIBLE))} {_d(-COMBUSTIBLE)} "
        f"(-) Gastos indirectos {_d(-_mxn(INDIRECTOS))} {_d(-INDIRECTOS)} "
        f"(-) Permisos {_d(-_mxn(PERMISOS))} {_d(-PERMISOS)} "
        f"{FIJOS_PDF} $-0.00 "
        f"Saldo disponible {_d(a.saldo_mxn)} {_d(SALDO)} "
        f"{rp.FILA_RESERVA_INFORMATIVA} {_d(_mxn(RESERVA))} {_d(RESERVA)}"
    ) in texto
    # Resumen global en pesos: venta − gastos (con combustible) = saldo.
    gastos_mxn = r2(
        _mxn(COMISIONES)
        + _mxn(DIRECTOS)
        + _mxn(AFAC)
        + _mxn(COMBUSTIBLE)
        + _mxn(INDIRECTOS)
        + _mxn(PERMISOS)
        + _mxn(FIJOS)
    )
    assert r2(_mxn(VENTA) - gastos_mxn) == a.saldo_mxn
    assert f"{_d(_mxn(VENTA))} $0.00 USD {_d(gastos_mxn)} {_d(a.saldo_mxn)}" in texto
    assert "gastos directos - provisión AFAC. Montos en pesos (MXN)" in texto


def test_pdf_pesos_sin_pesos_raya_nunca_cero() -> None:
    p = _payload()
    a = p["aviones"][0]
    a["combustible_mxn"] = None
    a["utilidad_antes_gastos_mxn"] = None
    a["gastos_directos_mxn"] = None
    texto = _texto_pdf(render_reparto_pdf(RepartoPdfRequest.model_validate(p)))
    assert f"{rc.FILA_GASTOS_DIRECTOS_SIN_COMBUSTIBLE} — {_d(-DIRECTOS)}" in texto
    assert f"{rc.FILA_UTILIDAD_ANTES_GASTOS} — {_d(UTILIDAD_ANTES)}" in texto
    assert f"{rc.FILA_COMBUSTIBLE} — {_d(-COMBUSTIBLE)}" in texto
    # Con un dato sin pesos, los gastos del resumen también son «—».
    assert f"{_d(_mxn(VENTA))} $0.00 USD — {_d(a['saldo_mxn'])}" in texto
    # Solo sin la del API y con los componentes completos se re-suma.
    p = _payload()
    p["aviones"][0]["utilidad_antes_gastos_mxn"] = None
    texto = _texto_pdf(render_reparto_pdf(RepartoPdfRequest.model_validate(p)))
    esperado = r2(_mxn(VENTA) - _mxn(COMISIONES) - _mxn(DIRECTOS) - _mxn(AFAC))
    assert f"{rc.FILA_UTILIDAD_ANTES_GASTOS} {_d(esperado)} {_d(UTILIDAD_ANTES)}" in texto


# ---------------------------------------------------------------------------
# (c) Excel en dólares.
# ---------------------------------------------------------------------------


def test_xlsx_usd_bloque_del_combustible_en_el_orden_del_excel() -> None:
    ws = _hoja(_payload(pesos=False))
    assert _encabezados(ws) == [
        "Matrícula", "Modelo", "Horas voladas", "Venta del avión cobrada", ETIQUETA_COLUMNA,
        "Pendiente de cobro (parte avión)", rc.ENCABEZADO_GASTOS_DIRECTOS_SIN_COMBUSTIBLE,
        rx.ENCABEZADO_PROVISION_AFAC, rc.ENCABEZADO_UTILIDAD_ANTES_GASTOS,
        rc.ENCABEZADO_COMBUSTIBLE, "Gastos indirectos", "Permisos",
        "Fijos (prorrateo + reparto)", rx.ENCABEZADO_RESERVA_INFORMATIVA, "Saldo disponible",
        "Otros ingr. VuelaTour (informativo)", "Deuda total del cliente (informativo)",
    ]  # fmt: skip
    assert [ws.cell(row=5, column=c).value for c in range(4, 16)] == [
        VENTA, COMISIONES, 0, DIRECTOS, AFAC, UTILIDAD_ANTES, COMBUSTIBLE, INDIRECTOS,
        PERMISOS, FIJOS, RESERVA, SALDO,
    ]  # fmt: skip
    # TOTAL (un avión): re-suma de cada columna; comisiones + directos + AFAC
    # + combustible + indirectos + permisos + fijos = venta − saldo.
    assert ws["A6"].value == "TOTAL"
    tot = {c: ws.cell(row=6, column=c).value for c in range(4, 16)}
    assert (tot[7], tot[8], tot[9], tot[10]) == (DIRECTOS, AFAC, UTILIDAD_ANTES, COMBUSTIBLE)
    assert r2(tot[5] + tot[7] + tot[8] + tot[10] + tot[11] + tot[12] + tot[13]) == r2(
        tot[4] - tot[15]
    )
    for c in (8, 9, 10):
        assert ws.cell(row=5, column=c).number_format == '"$"#,##0.00'
        assert ws.cell(row=6, column=c).font.b
    # Comentarios: la provisión AFAC (como siempre) y la utilidad antes de gastos.
    assert ws["H4"].comment.text == rx.NOTA_ENCABEZADO_AFAC
    assert ws["I4"].comment.text == rc.NOTA_ENCABEZADO_UTILIDAD_ANTES_GASTOS
    # La provisión AFAC ya no va al final; la reserva (N) sigue informativa.
    assert ws.cell(row=4, column=18).value is None
    assert ws["N5"].font.i and ws["N6"].font.i and ws["N6"].font.b
    # Título y notas combinan las 17 columnas; anchos del bloque.
    assert any(str(rg) == "A1:Q1" for rg in ws.merged_cells.ranges)
    assert [ws.column_dimensions[chr(64 + c)].width for c in (7, 8, 9, 10, 11)] == [
        15, 18, 18, 16, 16,
    ]  # fmt: skip
    textos = _textos_col_a(ws)
    assert rc.NOTA_COMBUSTIBLE_SEPARADO in textos and rx.NOTA_CASCADA_BALANCE in textos
    assert any(str(rg) == f"A{_fila(ws, rc.NOTA_COMBUSTIBLE_SEPARADO)}:Q{_fila(ws, rc.NOTA_COMBUSTIBLE_SEPARADO)}" for rg in ws.merged_cells.ranges)  # noqa: E501


def test_xlsx_usd_otros_ingresos_del_avion_tras_el_corrimiento() -> None:
    ws = _hoja(_payload(pesos=False, otros_ingresos=1_000.0))
    assert ws.cell(row=4, column=18).value == rx.ENCABEZADO_OTROS_INGRESOS_AVION
    assert (ws.cell(row=5, column=18).value, ws.cell(row=6, column=18).value) == (1000.0, 1000.0)
    assert ws["O5"].value == r2(SALDO + 1000)
    assert any(str(rg) == "A1:R1" for rg in ws.merged_cells.ranges)
    assert ws.column_dimensions["R"].width == 18
    textos = _textos_col_a(ws)
    assert rx.NOTA_OTROS_INGRESOS_AVION_BALANCE in textos and rc.NOTA_COMBUSTIBLE_SEPARADO in (
        textos
    )


def test_xlsx_usd_sin_cascada_del_balance_bloque_de_dos_columnas() -> None:
    ws = _hoja(_payload(pesos=False, como_balance=False))
    assert _encabezados(ws)[6:10] == [
        rc.ENCABEZADO_GASTOS_DIRECTOS_SIN_COMBUSTIBLE, rc.ENCABEZADO_UTILIDAD_ANTES_GASTOS,
        rc.ENCABEZADO_COMBUSTIBLE, "Gastos indirectos",
    ]  # fmt: skip
    assert ws["M4"].value == "Reserva overhaul" and ws["N4"].value == "Saldo disponible"
    assert ws["H5"].value == UTILIDAD_ANTES and ws["I5"].value == COMBUSTIBLE
    assert ws["H4"].comment.text == rc.NOTA_ENCABEZADO_UTILIDAD_ANTES_GASTOS
    assert ws.cell(row=4, column=17).value is None
    assert "Provisión AFAC" not in _encabezados(ws)
    assert any(str(rg) == "A1:P1" for rg in ws.merged_cells.ranges)


def test_xlsx_avion_sin_la_llave_en_un_reparto_que_la_trae() -> None:
    """Un avión de un API previo junto a uno del 0.0.103 (no pasa, pero no
    debe reventar): sus celdas nuevas quedan vacías y el total suma el otro."""
    p = _payload(pesos=False)
    viejo = _sin_llaves(_payload(pesos=False), nulos=True)["aviones"][0]
    viejo["matricula"] = "XB-VIE"
    p["aviones"].append(viejo)
    ws = _hoja(p)
    assert ws["G6"].value == DIRECTOS and ws["I6"].value is None and ws["J6"].value is None
    assert ws["A7"].value == "TOTAL" and ws["J7"].value == COMBUSTIBLE
    assert ws["I7"].value == UTILIDAD_ANTES


# ---------------------------------------------------------------------------
# (c) Excel en pesos.
# ---------------------------------------------------------------------------


def test_xlsx_pesos_bloque_del_combustible() -> None:
    p = _payload()
    a = p["aviones"][0]
    ws = _hoja(p)
    assert _encabezados(ws) == rx.encabezados_pesos(False, True, True)
    assert _encabezados(ws)[6:10] == [
        rc.ENCABEZADO_GASTOS_DIRECTOS_SIN_COMBUSTIBLE_PESOS, rx.ENCABEZADO_PROVISION_AFAC_PESOS,
        rc.ENCABEZADO_UTILIDAD_ANTES_GASTOS_PESOS, rc.ENCABEZADO_COMBUSTIBLE_PESOS,
    ]  # fmt: skip
    assert [ws.cell(row=5, column=c).value for c in range(7, 16)] == [
        a["gastos_directos_mxn"], a["afac_reservado_mxn"], a["utilidad_antes_gastos_mxn"],
        a["combustible_mxn"], a["gastos_indirectos_mxn"], a["permisos_mxn"], a["otros_mxn"],
        a["reserva_overhaul_mxn"], a["saldo_mxn"],
    ]  # fmt: skip
    assert ws["A6"].value == "TOTAL"
    assert ws["I6"].value == a["utilidad_antes_gastos_mxn"]
    assert ws["J6"].value == a["combustible_mxn"]
    assert ws["O5"].font.b  # saldo en negritas
    # La reserva (N) sigue informativa, con el T.C. en su comentario.
    assert ws["N5"].font.i and ws["N5"].comment.text.endswith(f"{K:g}.")
    assert ws["H4"].comment.text == rx.NOTA_ENCABEZADO_AFAC
    assert ws["I4"].comment.text == rc.NOTA_ENCABEZADO_UTILIDAD_ANTES_GASTOS
    # Las informativas en dólares (P, Q) siguen grises.
    assert ws["P4"].value == "Otros ingr. VuelaTour (USD, informativo)" and ws["P5"].font.i
    # La tabla USD informativa lleva el mismo bloque.
    t = _fila(ws, rx.TITULO_USD_INFORMATIVO)
    assert [ws.cell(row=t + 1, column=c).value for c in range(7, 11)] == [
        rc.ENCABEZADO_GASTOS_DIRECTOS_SIN_COMBUSTIBLE, rx.ENCABEZADO_PROVISION_AFAC,
        rc.ENCABEZADO_UTILIDAD_ANTES_GASTOS, rc.ENCABEZADO_COMBUSTIBLE,
    ]  # fmt: skip
    assert [ws.cell(row=t + 2, column=c).value for c in range(7, 11)] == [
        DIRECTOS, AFAC, UTILIDAD_ANTES, COMBUSTIBLE,
    ]  # fmt: skip
    assert ws.cell(row=t + 3, column=1).value == "TOTAL USD"
    assert ws.cell(row=t + 3, column=10).value == COMBUSTIBLE
    assert ws.cell(row=t + 3, column=15).value == SALDO
    assert any(str(rg) == "A1:Q1" for rg in ws.merged_cells.ranges)
    assert [ws.column_dimensions[chr(64 + c)].width for c in (7, 8, 9, 10, 11)] == [
        16, 18, 18, 16, 17,
    ]  # fmt: skip
    assert rc.NOTA_COMBUSTIBLE_SEPARADO in _textos_col_a(ws)
    # Sin pesos en el combustible: «—» y el total también «—», jamás un 0.
    a["combustible_mxn"] = None
    ws = _hoja(p)
    assert ws["J5"].value == rx.SIN_PESOS and ws["J6"].value == rx.SIN_PESOS
    assert ws.cell(row=t + 2, column=10).value == COMBUSTIBLE  # el USD informativo sigue


# ---------------------------------------------------------------------------
# Rutas: el JSON del API 0.0.103 entra sin 422.
# ---------------------------------------------------------------------------

TOKEN = "tok-reparto-combustible"
client = TestClient(app)


def test_rutas_aceptan_el_payload_0103(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    for p in (_payload(), _payload(pesos=False)):
        res = client.post("/pdf/reparto", json=p, headers=headers)
        assert res.status_code == 200, res.text
        texto = _texto_pdf(res.content)
        assert rc.FILA_UTILIDAD_ANTES_GASTOS in texto and rc.FILA_COMBUSTIBLE in texto
        res = client.post("/pdf/reparto-xlsx", json=p, headers=headers)
        assert res.status_code == 200, res.text
        ws = load_workbook(BytesIO(res.content)).active
        assert ws["J4"].value in (rc.ENCABEZADO_COMBUSTIBLE, rc.ENCABEZADO_COMBUSTIBLE_PESOS)
