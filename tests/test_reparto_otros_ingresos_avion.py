"""Reparto a socios (PDF y Excel) con los otros ingresos ASIGNADOS a un avión
(13-oct-2026, API 0.0.75; invariante «INGRESO ASIGNADO A UN AVIÓN» del API).

El API suma al saldo del avión los ingresos registrados en Ingresos con ese
avión, netos de comisión bancaria (`otros_ingresos_avion_usd`; `saldo_usd` YA
los trae) y manda el detalle por ingreso y los MXN que no tuvieron tipo de
cambio (`otros_ingresos_avion_sin_tc_mxn`: NO entraron al saldo). Aquí:
- PDF: fila «(+) Otros ingresos del avión…» antes de las deducciones (SUMA,
  no es gris), hasta 5 renglones grises de detalle y «y N más», el renglón
  gris de los MXN sin T.C., y el resumen con 5 columnas para que venta +
  otros ingresos de aviones − gastos = saldo;
- Excel: columna 15 «Otros ingresos del avión (suman al saldo)» con su total y
  una nota, solo con dato;
- sin los campos (API ≤ 0.0.74): el PDF y el Excel de siempre (flujos de
  contenido / miembros comparados contra el código de HEAD d58faa1).
"""

from __future__ import annotations

import copy
import hashlib
import re
import zipfile
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth

from app.config import get_settings
from app.main import app
from app.schemas.reparto import RepartoAvion, RepartoOtroIngresoLinea, RepartoPdfRequest
from app.services.reparto_pdf import (
    ENCABEZADOS_RESUMEN_CON_OTROS_INGRESOS,
    FILA_OTROS_INGRESOS_AVION,
    LINEA_OTROS_INGRESOS_SIN_TC,
    MAX_LINEAS_OTROS_INGRESOS,
    _linea_otro_ingreso,
    render_reparto_pdf,
)
from app.services.reparto_xlsx import (
    ENCABEZADO_OTROS_INGRESOS_AVION,
    NOTA_OTROS_INGRESOS_AVION,
    render_reparto_xlsx,
)
from tests import test_reparto_comisiones as trc

FILA = "(+) Otros ingresos del avión (registrados en Ingresos, netos de comisión)"
SIN_TC = "Otros ingresos del avión en MXN sin tipo de cambio — no entraron al saldo"
RESUMEN = [
    "Venta del avión cobrada",
    "Otros ingresos de aviones",
    "Otros ingresos VuelaTour",
    "Gastos del periodo",
    "Saldo a repartir",
]
COLUMNA = "Otros ingresos del avión (suman al saldo)"
NOTA_XLSX = (
    "Otros ingresos del avión = reembolsos, ventas e intereses registrados en "
    "Ingresos con ese avión, netos de comisión bancaria; SUMAN al saldo (Saldo "
    "disponible = venta cobrada − comisiones − gastos − fijos − reserva + otros "
    "ingresos del avión)."
)

# ING-4 (MXN 46,550.80) al T.C. oficial 18 ⇒ 2,586.16 USD.
_ING_4 = {
    "clave": "ING-4",
    "fecha": "2026-09-04",
    "categoria": "Reembolsos y devoluciones recibidos",
    "concepto": "TRANSF INTERBANCARIA SPEI · angel alvarez",
    "monto": 46550.8,
    "moneda": "MXN",
    "comision": None,
    "usd": 2586.16,
}


def _linea(n: int, usd: float | None = 10.0) -> dict:
    return {
        "clave": f"ING-{n}",
        "fecha": "2026-09-10",
        "categoria": "Otros ingresos",
        "concepto": f"Intereses {n}",
        "monto": 100.0,
        "moneda": "USD",
        "comision": None,
        "usd": usd,
    }


def _con_otros_ingresos(
    avion: dict, lineas: list[dict], *, usd: float, sin_tc_mxn: float = 0.0
) -> dict:
    """El avión del API 0.0.75: el saldo ya trae los otros ingresos y el
    reparto se re-arma con él."""
    a = copy.deepcopy(avion)
    a["otros_ingresos_avion_usd"] = usd
    a["otros_ingresos_avion_sin_tc_mxn"] = sin_tc_mxn
    a["otros_ingresos_avion"] = copy.deepcopy(lineas)
    a["saldo_usd"] = round(a["saldo_usd"] + usd, 2)
    a["reparto"] = [
        {"socio_nombre": "Socio A", "porcentaje": 100, "monto_usd": a["saldo_usd"]}
    ]
    return a


def _payload(*, lineas=None, usd=2586.16, sin_tc_mxn=0.0) -> dict:
    p = trc._payload()
    lineas = [_ING_4] if lineas is None else lineas
    p["aviones"][0] = _con_otros_ingresos(
        p["aviones"][0], lineas, usd=usd, sin_tc_mxn=sin_tc_mxn
    )
    return p


def _req(**kw) -> RepartoPdfRequest:
    return RepartoPdfRequest.model_validate(_payload(**kw))


def _texto_pdf(data: bytes) -> str:
    return re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(BytesIO(data)).pages))


def _flujos_pdf(data: bytes) -> list[bytes]:
    """Flujos de contenido de cada página (lo que se DIBUJA; la fecha de
    creación del documento queda fuera)."""
    return [p.get_contents().get_data() for p in PdfReader(BytesIO(data)).pages]


def _huella_pdf(data: bytes) -> str:
    h = hashlib.sha256()
    for flujo in _flujos_pdf(data):
        h.update(flujo)
    return h.hexdigest()


def _miembros(data: bytes) -> dict[str, bytes]:
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _huella_xlsx(data: bytes) -> str:
    h = hashlib.sha256()
    for nombre, contenido in _miembros(data).items():
        h.update(nombre.encode())
        h.update(contenido)
    return h.hexdigest()


# Calculadas con el código de HEAD d58faa1 (13-oct-2026), ANTES de este
# cambio, sobre `trc._payload()` (el reparto de siempre, sin los campos).
_HUELLA_PDF_HEAD = "b14c5b7a69142320770fe425695ec52b39c761137134e6490850f32aad0dcb2f"
_HUELLA_XLSX_HEAD = "bbad3d6fb93b4f18b138e4758d6056b9b469c6eddb51fbf4d1e2955e1732b6c8"


# ---------------------------------------------------------------------------
# Esquema y textos
# ---------------------------------------------------------------------------


def test_esquema_aditivo_con_defaults() -> None:
    a = trc._req().aviones[0]
    assert (a.otros_ingresos_avion_usd, a.otros_ingresos_avion_sin_tc_mxn) == (0.0, 0.0)
    assert a.otros_ingresos_avion == []
    nulos = RepartoAvion.model_validate(
        {
            **trc._avion("XB-TST", 0),
            "otros_ingresos_avion_usd": None,
            "otros_ingresos_avion_sin_tc_mxn": None,
            "otros_ingresos_avion": None,
        }
    )
    assert (nulos.otros_ingresos_avion_usd, nulos.otros_ingresos_avion) == (0.0, [])
    linea = RepartoOtroIngresoLinea.model_validate({"clave": None, "usd": None, "x": 1})
    assert (linea.clave, linea.usd) == ("", None)


def test_textos_son_los_del_contrato_y_caben() -> None:
    assert FILA_OTROS_INGRESOS_AVION == FILA
    assert LINEA_OTROS_INGRESOS_SIN_TC == SIN_TC
    assert ENCABEZADOS_RESUMEN_CON_OTROS_INGRESOS == RESUMEN
    assert ENCABEZADO_OTROS_INGRESOS_AVION == COLUMNA
    assert NOTA_OTROS_INGRESOS_AVION == NOTA_XLSX
    assert MAX_LINEAS_OTROS_INGRESOS == 5
    # En la celda de 120 mm (6 pt de relleno por lado): la fila a 9 pt y la
    # gris a 7.5 pt Oblique no se enciman con el monto.
    assert stringWidth(FILA, "Helvetica", 9) <= 120 * mm - 12
    assert stringWidth(SIN_TC, "Helvetica-Oblique", 7.5) <= 120 * mm - 12


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_pdf_fila_que_suma_antes_de_las_deducciones() -> None:
    texto = _texto_pdf(render_reparto_pdf(_req()))
    assert f"{FILA} $2,586.16" in texto
    # Después de la venta y antes de las comisiones y los gastos.
    assert texto.index("Venta del avión cobrada $2,000.00") < texto.index(FILA)
    # (La etiqueta de comisiones también sale antes, en la nota del resumen:
    # se busca la fila de la cascada, con su monto.)
    assert texto.index(FILA) < texto.index("(-) Comisiones (banco + vendedor) $-95.70")
    assert texto.index(FILA) < texto.index("(-) Gastos directos")
    # Su detalle gris y el saldo con los otros ingresos (2000 − 500 − 95.70 +
    # 2586.16).
    assert "ING-4 · Reembolsos y devoluciones recibidos · TRANSF INTERBANCARIA SPEI" in texto
    assert "Saldo disponible $3,990.46" in texto


def test_pdf_resumen_de_cinco_columnas_cuadra() -> None:
    req = _req()
    texto = _texto_pdf(render_reparto_pdf(req))
    assert " ".join(RESUMEN) in texto
    venta = sum(a.ingresos_cobrado_usd for a in req.aviones)
    otros = sum(a.otros_ingresos_avion_usd for a in req.aviones)
    gastos = 2 * 500.0 + 95.7
    saldo = sum(a.saldo_usd for a in req.aviones)
    assert round(venta + otros - gastos, 2) == round(saldo, 2) == 5490.46
    assert "$4,000.00 $2,586.16 $208.80 $1,095.70 $5,490.46" in texto


def test_pdf_detalle_maximo_cinco_y_y_n_mas() -> None:
    lineas = [_linea(n) for n in range(10, 17)]  # 7 ingresos
    texto = _texto_pdf(render_reparto_pdf(_req(lineas=lineas, usd=70.0)))
    for n in range(10, 15):
        assert f"ING-{n} · Otros ingresos · Intereses {n} $10.00" in texto
    assert "ING-15" not in texto and "ING-16" not in texto
    assert "y 2 más" in texto


def test_pdf_ingreso_sin_tc_y_renglon_gris() -> None:
    lineas = [_ING_4, {**_linea(20, usd=None), "moneda": "MXN", "monto": 1234.5}]
    texto = _texto_pdf(
        render_reparto_pdf(_req(lineas=lineas, usd=2586.16, sin_tc_mxn=1234.5))
    )
    assert "ING-20 · Otros ingresos · Intereses 20 sin T.C." in texto
    assert f"{SIN_TC} $1,234.50 MXN" in texto


def test_pdf_solo_sin_tc_no_imprime_la_fila_que_suma() -> None:
    lineas = [{**_linea(20, usd=None), "moneda": "MXN", "monto": 1234.5}]
    texto = _texto_pdf(render_reparto_pdf(_req(lineas=lineas, usd=0.0, sin_tc_mxn=1234.5)))
    assert FILA not in texto
    assert f"{SIN_TC} $1,234.50 MXN" in texto
    # Sin otros ingresos que sumen, el resumen de 4 columnas de siempre.
    assert "Otros ingresos de aviones" not in texto


def test_pdf_detalle_largo_se_recorta_a_su_columna() -> None:
    """Un concepto largo se recorta con «...» al ancho de la celda de 120 mm
    (Helvetica no envuelve dentro de una tabla: se encimaría con el monto)."""
    larga = {**_linea(30), "concepto": "Reembolso " + "muy largo " * 40}
    etiqueta, monto = _linea_otro_ingreso(RepartoOtroIngresoLinea.model_validate(larga))
    assert etiqueta.startswith("   ING-30 · Otros ingresos · Reembolso muy largo")
    assert etiqueta.endswith("...") and monto == "$10.00"
    assert stringWidth(etiqueta, "Helvetica-Oblique", 7.5) <= 120 * mm - 12
    texto = _texto_pdf(render_reparto_pdf(_req(lineas=[larga], usd=10.0)))
    assert re.search(r"ING-30 · Otros ingresos · Reembolso[^$]*\.\.\. \$10\.00", texto)


def test_pdf_sin_los_campos_el_de_siempre() -> None:
    base = render_reparto_pdf(trc._req())
    assert _huella_pdf(base) == _HUELLA_PDF_HEAD
    # Con los campos en cero / vacíos: el mismo dibujo.
    ceros = trc._payload()
    for a in ceros["aviones"]:
        a.update(
            {
                "otros_ingresos_avion_usd": 0,
                "otros_ingresos_avion_sin_tc_mxn": 0,
                "otros_ingresos_avion": [],
            }
        )
    assert _flujos_pdf(render_reparto_pdf(RepartoPdfRequest.model_validate(ceros))) == _flujos_pdf(
        base
    )
    assert "Otros ingresos de aviones" not in _texto_pdf(base)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def test_xlsx_columna_15_con_total_y_nota() -> None:
    ws = load_workbook(BytesIO(render_reparto_xlsx(_req()))).active
    assert ws["O4"].value == COLUMNA
    assert (ws["O5"].value, ws["O6"].value) == (2586.16, 0)
    assert ws["A7"].value == "TOTAL" and ws["O7"].value == 2586.16
    assert ws["O7"].font.b
    assert ws["O5"].number_format == '"$"#,##0.00'
    assert ws.column_dimensions["O"].width == 18
    # El saldo (L) ya trae los otros ingresos; D..L no se movieron.
    assert ws["L5"].value == round(2000.0 - 500.0 - 95.7 + 2586.16, 2)
    assert ws["D4"].value == "Venta del avión cobrada" and ws["L4"].value == "Saldo disponible"
    # El título y las notas combinan las 15 columnas; la nota nueva va justo
    # después de la de la tabla.
    assert any(str(rg) == "A1:O1" for rg in ws.merged_cells.ranges)
    assert ws["A8"].value.startswith("Venta del avión cobrada = ")
    assert ws["A9"].value == NOTA_XLSX
    assert any(str(rg) == "A9:O9" for rg in ws.merged_cells.ranges)


def test_xlsx_sin_dato_catorce_columnas_de_siempre() -> None:
    base = render_reparto_xlsx(trc._req())
    ws = load_workbook(BytesIO(base)).active
    assert all(ws.cell(row=r, column=15).value is None for r in range(1, 10))
    assert ws.column_dimensions["O"].width != 18
    assert NOTA_XLSX not in [c.value for c in ws["A"]]
    assert _huella_xlsx(base) == _HUELLA_XLSX_HEAD
    ceros = trc._payload()
    for a in ceros["aviones"]:
        a.update({"otros_ingresos_avion_usd": 0, "otros_ingresos_avion": []})
    assert _miembros(render_reparto_xlsx(RepartoPdfRequest.model_validate(ceros))) == _miembros(
        base
    )


# ---------------------------------------------------------------------------
# API 0.0.77 (14-oct-2026, conversión ÚNICA del ingreso de un avión): cada
# línea del detalle gana `tc_origen` AL FINAL (ADITIVO). El reparto no cambia
# ni un número y pyservices ignora la llave: PDF y Excel idénticos.
# ---------------------------------------------------------------------------


def test_tc_origen_del_api_0077_se_ignora() -> None:
    sin = _req(lineas=[_ING_4, _linea(5)])
    con = _req(
        lineas=[{**_ING_4, "tc_origen": "OFICIAL_DEL_DIA"}, {**_linea(5), "tc_origen": None}]
    )
    assert _flujos_pdf(render_reparto_pdf(con)) == _flujos_pdf(render_reparto_pdf(sin))
    assert _miembros(render_reparto_xlsx(con)) == _miembros(render_reparto_xlsx(sin))


# ---------------------------------------------------------------------------
# Rutas con el JSON del API 0.0.75
# ---------------------------------------------------------------------------

TOKEN = "tok-reparto-oi"
client = TestClient(app)


def test_rutas_aceptan_el_payload_0075(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    payload = _payload(sin_tc_mxn=12.5)
    res = client.post("/pdf/reparto", json=payload, headers=headers)
    assert res.status_code == 200, res.text
    assert FILA in _texto_pdf(res.content)
    res = client.post("/pdf/reparto-xlsx", json=payload, headers=headers)
    assert res.status_code == 200, res.text
    assert load_workbook(BytesIO(res.content)).active["O4"].value == COLUMNA
