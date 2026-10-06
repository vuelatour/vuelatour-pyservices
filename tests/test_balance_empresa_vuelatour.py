"""Balance general: bloque «VUELATOUR (empresa)» al final de la hoja
'balance' (6-oct-2026). Pedido del cliente sobre la hoja «balance» del libro
`balance-general-vuelatour-…xlsx`: «en la hoja de balance falta, hasta el
final, el balance de la empresa VuelaTour».

El bloque lo arma el API (`empresa`, API 0.0.59): aquí solo se pinta. La
participación como socio CITA la celda MONTO USD del socio `es_empresa` de
cada bloque de avión de arriba (aunque los bloques tengan distinto número de
socios), los otros gastos citan el TOTAL USD de la hoja 'otros gastos' y el
resultado es la fórmula de las cinco líneas; todo verificado con el
evaluador independiente de los tests. Sin `empresa` (API viejo) el libro es
el de siempre, byte a byte.
"""

from __future__ import annotations

import hashlib
import json
import zipfile
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.schemas.reportes import BalanceGeneralRequest
from app.services.balance_avion_xlsx import (
    GREEN,
    MUTED,
    NAVY,
    RED,
    render_balance_general_xlsx,
)
from tests import test_balance_formulas as tf
from tests._evaluador_formulas import verificar_libro

r2 = tf.r2

_EMPRESA = "Aero Charter Cancún S.A. de C.V."

# Socios por avión: (nombre, %, es_empresa). Bloques con DISTINTO número de
# socios y la empresa en medio del primero: las referencias no pueden salir
# de una cuenta fija de filas.
_SOCIOS_CON_EMPRESA = {
    "XA-TST": [("Socio A", 50.0, False), (_EMPRESA, 30.0, True), ("Socio B", 20.0, False)],
    "XB-DOS": [(_EMPRESA, 100.0, True)],
}
_SOCIOS_SIN_EMPRESA = {
    "XA-TST": [("Socio A", 60.0, False), ("Socio B", 40.0, False)],
    "XB-DOS": [("Socio C", 100.0, False)],
}

# TC de SU vuelo para cada fila de 'otros movimientos' del consolidado de
# test_balance_formulas (K de la cotización; la suelta ING-3, TC oficial del
# día del cobro) — espejo de lo que el API convierte.
_INGRESOS_MXN_TC = [(1250.5, 17.25), (3000.0, 18.123456), (85.33, 17.5)]
_EGRESOS_MXN_TC = [(1100.25, 17.25)]


def _poner_socios(libro: dict, socios: list[tuple[str, float, bool]]) -> None:
    """Socios de un avión como los manda el API (monto = cobrada × % / 100)."""
    cobrada = libro["balance"]["utilidad_cobrada_usd"]
    libro["balance"]["socios"] = [
        {
            "nombre": nombre,
            "porcentaje": pct,
            "monto_usd": r2((pct / 100) * cobrada),
            **({"es_empresa": True} if es_empresa else {}),
        }
        for nombre, pct, es_empresa in socios
    ]


def _empresa_api(p: dict) -> dict:
    """Espejo de `balance-empresa.util.ts` (API 0.0.59)."""
    tc = p["consolidado"]["totales"]["tc_promedio"]
    participaciones = [
        {"matricula": a["matricula"], "porcentaje": s["porcentaje"], "monto_usd": s["monto_usd"]}
        for a in p["aviones"]
        for s in a["balance"]["socios"]
        if s.get("es_empresa")
    ]
    participacion = r2(tf._suma(x["monto_usd"] for x in participaciones))
    ingresos = r2(tf._suma(mxn / k for mxn, k in _INGRESOS_MXN_TC))
    pagos = r2(tf._suma(mxn / k for mxn, k in _EGRESOS_MXN_TC))
    gastos = p.get("gastos_empresa")
    otros = gastos["usd"] if gastos is not None else 0
    inv = p.get("inventario")
    tienda = r2(inv["total_utilidad_mxn"] / tc) if inv is not None else None
    resultado = r2(participacion + ingresos - pagos - otros + (tienda or 0))
    return {
        "participaciones": participaciones,
        "participacion_usd": participacion,
        "ingresos_propios_usd": ingresos,
        "pagos_vendedor_usd": pagos,
        "otros_gastos_empresa_usd": otros,
        "tc_usado": tc,
        "tienda_utilidad_usd": tienda,
        "resultado_usd": resultado,
        "tc_promedio": tc,
        "nota": "Ingresos cobrados y pagos al vendedor pagados del periodo; TC de cada vuelo.",
    }


def _payload(
    *,
    socios: dict | None = None,
    empresa: bool = True,
    otros_gastos: bool = True,
    inventario: bool = True,
) -> dict:
    p = tf._general_payload()
    socios = _SOCIOS_CON_EMPRESA if socios is None else socios
    for libro in p["aviones"]:
        _poner_socios(libro, socios[libro["matricula"]])
    if not otros_gastos:
        p.pop("gastos_empresa")
    if not inventario:
        p.pop("inventario")
    if empresa:
        p["empresa"] = _empresa_api(p)
    return p


def _render(p: dict) -> bytes:
    return render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))


def _miembros(data: bytes) -> dict[str, bytes]:
    """Todos los miembros del .xlsx salvo docProps/core.xml (lleva la hora)."""
    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _firma_libro(data: bytes) -> str:
    """Huella del LAYOUT de TODAS las hojas (misma receta que
    test_balance_extension_horario, con las fórmulas tal cual): valores o
    fórmulas, formatos, fuentes, rellenos, alineación, bordes, comentarios,
    merges, anchos, altos y panel fijo. No depende de cómo serialice el XML
    la versión de openpyxl instalada."""
    wb = load_workbook(BytesIO(data))
    partes: list = []
    for ws in wb.worksheets:
        partes.append(ws.title)
        for fila in ws.iter_rows():
            for c in fila:
                f = c.font
                partes.append([
                    c.coordinate, repr(c.value), c.number_format,
                    bool(f.b), bool(f.i), f.sz,
                    f.color.rgb if f.color is not None else None,
                    c.fill.fgColor.rgb if c.fill.fill_type else None,
                    c.alignment.horizontal, c.alignment.vertical,
                    bool(c.alignment.wrap_text), c.border.left.style,
                    c.comment.text if c.comment else None,
                ])  # fmt: skip
        partes.append(sorted(str(r) for r in ws.merged_cells.ranges))
        partes.append(sorted((k, d.width) for k, d in ws.column_dimensions.items()))
        partes.append(sorted((k, d.height) for k, d in ws.row_dimensions.items() if d.height))
        partes.append(ws.freeze_panes)
    return hashlib.sha256(json.dumps(partes, default=str).encode()).hexdigest()


# Huellas del libro SIN `empresa` calculadas con el código ANTERIOR a este
# cambio (HEAD f509d1a, 6-oct-2026) sobre los mismos payloads (con los socios
# `es_empresa`, que ese código ignoraba). Además se verificó que los .xlsx
# completos salían byte-idénticos (todos los miembros del zip salvo
# docProps/core.xml). Si una cambia, un API 0.0.58 ya no ve su libro de
# siempre.
_FIRMAS_SIN_EMPRESA = {
    "socio empresa en dos aviones": (
        {},
        "28b3b186a50dd9665222fcabf3965a441012dccdf4c8336b6c3697e0c7861f3d",
    ),
    "sin socio empresa": (
        {"socios": _SOCIOS_SIN_EMPRESA},
        "4822200b323eeb1deba59522f88b74e130cfc47695691dd0e1c16c2ea415cf2d",
    ),
    "sin hoja otros gastos": (
        {"otros_gastos": False},
        "2b9fae9dec1a78494dadec56c52d02f4228a070d5f8450dc01ff2e3cbbad63e2",
    ),
    "sin inventario": (
        {"inventario": False},
        "95e4d0be2052107bfb28cfd5aa896f0d7bb6e0669ba37215016bdaf31c19a3a4",
    ),
}


def _libros(data: bytes):
    """(con fórmulas, con caché)."""
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _fila(ws, texto: str) -> int:
    return next(c.row for c in ws["A"] if c.value == texto)


def _filas_socio_empresa(ws) -> list[int]:
    """Filas (en orden) de los socios EMPRESA en los bloques de los aviones."""
    return [c.row for c in ws["A"] if c.value == _EMPRESA]


def _bloque(ws) -> dict[str, int]:
    """Fila de cada renglón del bloque «VUELATOUR (empresa)»."""
    return {
        "titulo": _fila(ws, "VUELATOUR (empresa)"),
        "part": _fila(ws, "(+) PARTICIPACIÓN COMO SOCIO EN LOS AVIONES USD"),
        "ing": _fila(ws, "(+) INGRESOS PROPIOS COBRADOS USD (TUAs, extras, pernocta, comisión)"),
        "pag": _fila(ws, "(−) PAGOS AL VENDEDOR USD"),
        "otros": _fila(ws, "(−) OTROS GASTOS DE LA EMPRESA USD"),
        "tienda": _fila(ws, "(+) UTILIDAD TIENDA (INVENTARIO) USD"),
        "res": _fila(ws, "RESULTADO VUELATOUR USD"),
    }


def _formula_resultado(b: dict[str, int]) -> str:
    return f"=ROUND(B{b['part']}+B{b['ing']}-B{b['pag']}-B{b['otros']}+B{b['tienda']},2)"


def _verificadas_del_bloque(data: bytes, desde: int) -> dict[str, str]:
    """Fórmulas de la hoja 'balance' desde la fila `desde`, verificadas por el
    evaluador INDEPENDIENTE de los tests (celda → fórmula)."""
    return {
        celda: formula
        for hoja, celda, formula, _calc, _cache in verificar_libro(data)
        if hoja == "balance" and int("".join(ch for ch in celda if ch.isdigit())) >= desde
    }


# ---------------------------------------------------------------------------
# 1) El bloque, al final de la hoja, con sus fórmulas.
# ---------------------------------------------------------------------------


def test_bloque_empresa_al_final_con_formulas_y_caches_del_api() -> None:
    p = _payload()
    emp = p["empresa"]
    data = _render(p)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    b = _bloque(ws)
    socios = _filas_socio_empresa(ws)
    assert len(socios) == 2
    # Después del ÚLTIMO bloque de avión, en el orden del contrato.
    assert b["titulo"] > max(socios)
    assert [b["part"], b["ing"], b["pag"], b["otros"], b["tienda"], b["res"]] == [
        b["titulo"] + 1,
        b["titulo"] + 4,  # dos líneas grises (una por avión) bajo la participación
        b["titulo"] + 5,
        b["titulo"] + 6,
        b["titulo"] + 7,
        b["titulo"] + 8,
    ]
    titulo = ws.cell(row=b["titulo"], column=1)
    assert titulo.font.b and titulo.font.color.rgb.endswith("FFFFFF")
    assert titulo.fill.fgColor.rgb.endswith(NAVY)

    # Participación = Σ de las celdas MONTO USD reales de los socios empresa.
    c1, c2 = (f"C{f}" for f in socios)
    assert ws[f"B{b['part']}"].value == f"=ROUND(SUM({c1},{c2}),2)"
    assert wsv[f"B{b['part']}"].value == emp["participacion_usd"]
    assert ws[f"A{b['part'] + 1}"].value == "   XA-TST · 30.00 %"
    assert ws[f"B{b['part'] + 1}"].value == f"={c1}"
    assert ws[f"A{b['part'] + 2}"].value == "   XB-DOS · 100.00 %"
    assert ws[f"B{b['part'] + 2}"].value == f"={c2}"
    assert [wsv[f"B{b['part'] + 1}"].value, wsv[f"B{b['part'] + 2}"].value] == [
        x["monto_usd"] for x in emp["participaciones"]
    ]
    gris = ws[f"A{b['part'] + 1}"].font
    assert gris.i and gris.color.rgb.endswith(MUTED)

    # Ingresos propios, pagos al vendedor y tienda: VALOR del API con nota.
    nota_om = "MXN de la hoja 'otros movimientos' convertidos con el TC de cada vuelo"
    for clave, campo in (
        ("ing", "ingresos_propios_usd"),
        ("pag", "pagos_vendedor_usd"),
        ("tienda", "tienda_utilidad_usd"),
    ):
        assert ws[f"B{b[clave]}"].value == emp[campo], clave
    assert ws[f"B{b['ing']}"].comment.text.startswith(nota_om)
    assert ws[f"B{b['pag']}"].comment.text.startswith(nota_om)
    assert ws[f"B{b['tienda']}"].comment.text.startswith(
        "utilidad de la hoja 'inventario' / TC promedio"
    )

    # Otros gastos: cita el TOTAL USD de la hoja 'otros gastos'.
    assert ws[f"B{b['otros']}"].value == "='otros gastos'!$C$5"
    assert wb["otros gastos"]["C5"].value == "=ROUND(A5/B5,2)"
    assert wsv[f"B{b['otros']}"].value == emp["otros_gastos_empresa_usd"]

    # Resultado: fórmula de las cinco líneas, negrita y verde (≥ 0).
    res = ws[f"B{b['res']}"]
    assert res.value == _formula_resultado(b)
    assert wsv[f"B{b['res']}"].value == emp["resultado_usd"]
    assert res.font.b and res.font.color.rgb.endswith(GREEN)
    assert ws[f"A{b['res']}"].font.b

    # Nota al pie del bloque + la nota del API.
    textos = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert any(t.startswith("Participación = utilidad COBRADA de cada avión") for t in textos)
    assert any("gastos personales del dueño no entran" in t for t in textos)
    assert emp["nota"] in textos

    # El evaluador independiente reproduce CADA fórmula del bloque.
    verificadas = _verificadas_del_bloque(data, b["titulo"])
    assert set(verificadas) == {
        f"B{b['part']}",
        f"B{b['part'] + 1}",
        f"B{b['part'] + 2}",
        f"B{b['otros']}",
        f"B{b['res']}",
    }


@pytest.mark.parametrize(
    "socios",
    [
        # Empresa al FINAL de un bloque de 4 socios y al inicio de uno de 2.
        {
            "XA-TST": [
                ("A", 40.0, False),
                ("B", 20.0, False),
                ("C", 15.0, False),
                (_EMPRESA, 25.0, True),
            ],
            "XB-DOS": [(_EMPRESA, 70.0, True), ("D", 30.0, False)],
        },
        # Solo un avión con la empresa (el segundo).
        {
            "XA-TST": [("A", 100.0, False)],
            "XB-DOS": [("D", 45.5, False), (_EMPRESA, 54.5, True)],
        },
    ],  # fmt: skip
    ids=["cuatro-y-dos-socios", "un-solo-avion"],
)
def test_las_referencias_son_las_celdas_reales_del_socio_empresa(socios) -> None:
    p = _payload(socios=socios)
    data = _render(p)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    b = _bloque(ws)
    filas = _filas_socio_empresa(ws)
    refs = ",".join(f"C{f}" for f in filas)
    assert ws[f"B{b['part']}"].value == f"=ROUND(SUM({refs}),2)"
    for i, f in enumerate(filas, start=1):
        assert ws[f"B{b['part'] + i}"].value == f"=C{f}"
        # La celda citada ES el monto del socio empresa de su bloque.
        assert ws[f"A{f}"].value == _EMPRESA
        assert wsv[f"C{f}"].value == p["empresa"]["participaciones"][i - 1]["monto_usd"]
    assert wsv[f"B{b['part']}"].value == p["empresa"]["participacion_usd"]
    assert ws[f"B{b['res']}"].value == _formula_resultado(b)
    assert wsv[f"B{b['res']}"].value == p["empresa"]["resultado_usd"]
    assert f"B{b['part']}" in _verificadas_del_bloque(data, b["titulo"])


# ---------------------------------------------------------------------------
# 2) Sin `empresa` (API viejo): el libro de siempre, byte a byte.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("caso", list(_FIRMAS_SIN_EMPRESA))
def test_sin_empresa_el_libro_es_el_de_antes(caso) -> None:
    kwargs, firma = _FIRMAS_SIN_EMPRESA[caso]
    data = _render(_payload(empresa=False, **kwargs))
    assert _firma_libro(data) == firma
    textos = [c.value for c in load_workbook(BytesIO(data))["balance"]["A"]]
    assert "VUELATOUR (empresa)" not in textos


def test_sin_empresa_ni_la_bandera_es_empresa_cambia_un_byte() -> None:
    """La bandera `es_empresa` de los socios y `empresa: null` no cambian
    nada: el .xlsx es idéntico al de un payload sin ellas."""
    con_bandera = _payload(empresa=False)
    sin_bandera = _payload(empresa=False)
    for libro in sin_bandera["aviones"]:
        for s in libro["balance"]["socios"]:
            s.pop("es_empresa", None)
    nulo = {**con_bandera, "empresa": None}
    base = _miembros(_render(sin_bandera))
    assert _miembros(_render(con_bandera)) == base
    assert _miembros(_render(nulo)) == base


# ---------------------------------------------------------------------------
# 3) Casos de borde.
# ---------------------------------------------------------------------------


def test_sin_hoja_otros_gastos_la_linea_va_como_valor() -> None:
    p = _payload(otros_gastos=False)
    data = _render(p)
    wb, wv = _libros(data)
    assert "otros gastos" not in wb.sheetnames
    ws, wsv = wb["balance"], wv["balance"]
    b = _bloque(ws)
    assert ws[f"B{b['otros']}"].value == p["empresa"]["otros_gastos_empresa_usd"]
    assert ws[f"B{b['res']}"].value == _formula_resultado(b)
    assert wsv[f"B{b['res']}"].value == p["empresa"]["resultado_usd"]
    assert f"B{b['res']}" in _verificadas_del_bloque(data, b["titulo"])


def test_sin_inventario_la_tienda_queda_vacia_y_el_resultado_cuadra() -> None:
    p = _payload(inventario=False)
    assert p["empresa"]["tienda_utilidad_usd"] is None
    data = _render(p)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    b = _bloque(ws)
    tienda = ws[f"B{b['tienda']}"]
    assert tienda.value is None and tienda.comment is None  # vacía, nunca un 0 falso
    assert ws[f"B{b['res']}"].value == _formula_resultado(b)
    assert wsv[f"B{b['res']}"].value == p["empresa"]["resultado_usd"]
    assert f"B{b['res']}" in _verificadas_del_bloque(data, b["titulo"])


def test_sin_socio_empresa_participacion_cero_y_linea_sin_participacion() -> None:
    p = _payload(socios=_SOCIOS_SIN_EMPRESA)
    assert p["empresa"]["participaciones"] == []
    assert p["empresa"]["participacion_usd"] == 0
    data = _render(p)
    wb, wv = _libros(data)
    ws, wsv = wb["balance"], wv["balance"]
    b = _bloque(ws)
    assert ws[f"B{b['part']}"].value == 0  # valor: no hay celda que sumar
    assert ws[f"A{b['part'] + 1}"].value == "   sin participación registrada"
    assert b["ing"] == b["part"] + 2
    assert ws[f"B{b['res']}"].value == _formula_resultado(b)
    assert wsv[f"B{b['res']}"].value == p["empresa"]["resultado_usd"]
    assert f"B{b['res']}" in _verificadas_del_bloque(data, b["titulo"])


def test_resultado_negativo_en_rojo() -> None:
    p = _payload(otros_gastos=False)
    emp = p["empresa"]
    emp["otros_gastos_empresa_usd"] = 9876.54
    emp["resultado_usd"] = r2(
        emp["participacion_usd"]
        + emp["ingresos_propios_usd"]
        - emp["pagos_vendedor_usd"]
        - 9876.54
        + emp["tienda_utilidad_usd"]
    )
    assert emp["resultado_usd"] < 0
    data = _render(p)
    wb, wv = _libros(data)
    b = _bloque(wb["balance"])
    res = wb["balance"][f"B{b['res']}"]
    assert res.value == _formula_resultado(b)
    assert res.font.color.rgb.endswith(RED)
    assert wv["balance"][f"B{b['res']}"].value == emp["resultado_usd"]


def test_participaciones_sin_bandera_en_los_socios_van_como_valor() -> None:
    """Skew: el API manda `participaciones` pero los socios no traen
    `es_empresa` ⇒ no hay celda que citar: total y líneas como valor."""
    p = _payload()
    for libro in p["aviones"]:
        for s in libro["balance"]["socios"]:
            s.pop("es_empresa", None)
    data = _render(p)
    ws = load_workbook(BytesIO(data))["balance"]
    b = _bloque(ws)
    emp = p["empresa"]
    assert ws[f"B{b['part']}"].value == emp["participacion_usd"]
    assert ws[f"A{b['part'] + 1}"].value == "   XA-TST · 30.00 %"
    assert ws[f"B{b['part'] + 1}"].value == emp["participaciones"][0]["monto_usd"]
    assert ws[f"B{b['part'] + 2}"].value == emp["participaciones"][1]["monto_usd"]
    assert ws[f"B{b['res']}"].value == _formula_resultado(b)
    verificar_libro(data)


def test_bloque_a_medias_no_tumba_el_balance() -> None:
    """Todos los campos son opcionales (y `participaciones` acepta null): un
    bloque vacío se pinta con celdas vacías, sin fórmulas sin número."""
    p = _payload(empresa=False)
    p["empresa"] = {"participaciones": None}
    data = _render(p)
    ws = load_workbook(BytesIO(data))["balance"]
    b = _bloque(ws)
    assert ws[f"B{b['part']}"].value is None
    assert ws[f"B{b['res']}"].value is None
    assert ws[f"B{b['ing']}"].value is None
    # Las líneas por avión sí citan su celda (el monto del socio existe).
    assert ws[f"B{b['part'] + 1}"].value.startswith("=C")
    verificar_libro(data)
