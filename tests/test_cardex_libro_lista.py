"""Cardex como LISTA por fecha y RESUMEN POR MES (18-oct-2026, API 0.0.92).

Pedido del cliente: «cambiar el formato del Cardex del producto en forma de
lista por orden de fecha … con sus entradas, salidas y movimientos» y «en el
Excel que se descarga debería venir stock inicio del mes y stock final del
mes y cantidad de unidades vendidas». Con `movimientos` el libro trae tres
hojas (lista, resumen por mes y la de siempre); sin ellos, el libro de
siempre byte a byte.
"""

from __future__ import annotations

import hashlib
import zipfile
from datetime import datetime
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.cardex_libro import CardexLibroRequest
from app.services.cardex_libro_xlsx import render_cardex_libro_xlsx
from tests.test_cardex_libro_xlsx import _req

# Huella de los miembros del .xlsx del libro de SIEMPRE (payload `_req()`,
# sin `movimientos`), calculada con el código ANTERIOR a este cambio
# (HEAD 46b2a50). docProps/core.xml lleva la hora: fuera.
_MIEMBROS_SIN_MOVIMIENTOS = {
    "[Content_Types].xml": "ba5af858089dd335",
    "_rels/.rels": "c545941ba36c15fc",
    "docProps/app.xml": "209fca6b00afe72a",
    "xl/_rels/workbook.xml.rels": "26ad8fcc38d41229",
    "xl/styles.xml": "ea40c58c417f4e70",
    "xl/theme/theme1.xml": "d15e8ebf78ef7b97",
    "xl/workbook.xml": "b4cbc52a3880a472",
    "xl/worksheets/sheet1.xml": "8ef6c2307b13db39",
}


def _miembros(xlsx: bytes) -> dict[str, str]:
    z = zipfile.ZipFile(BytesIO(xlsx))
    return {
        n: hashlib.sha256(z.read(n)).hexdigest()[:16]
        for n in sorted(z.namelist())
        if n != "docProps/core.xml"
    }


def _payload() -> dict:
    """La forma que manda el API 0.0.92 (`cardexLibroXlsx`): compra de agosto
    (10 a $90), salida a costo de agosto (3), compra de septiembre (5 a $108),
    venta de octubre a XA-VGV (2 a $144) y noviembre sin movimientos."""
    return {
        "titulo": "Cardex — Aceite Semisintético 15W-50",
        "item_nombre": "Aceite Semisintético 15W-50",
        "unidad": "cuarto (qt)",
        "generado": "2026-11-05",
        "moneda": "MXN",
        "nota": "Compras al T.C. oficial del día de la compra; ventas al del día de la venta.",
        "entradas": [
            {"fecha": "2026-08-29", "cantidad": 10, "descripcion": "Aceite · Entrada a bodega",
             "valor_compra_unitario": 90, "valor_compra_total": 900, "stock_despues": 10},
            {"fecha": "2026-09-10", "cantidad": 5, "descripcion": "Aceite · Entrada a bodega",
             "valor_compra_unitario": 108, "valor_compra_total": 540, "stock_despues": 12},
        ],
        "salidas": [
            {"fecha": "2026-08-30", "cantidad": 3, "descripcion": "Aceite · a costo",
             "venta_unitaria": 90, "venta_total": 270, "remanente": 7, "ganancia": 0,
             "vendido_a": "—"},
            {"fecha": "2026-10-02", "cantidad": 2, "descripcion": "Aceite",
             "venta_unitaria": 144, "venta_total": 288, "remanente": 10, "ganancia": 72,
             "vendido_a": "XA-VGV"},
        ],
        "movimientos": [
            {"fecha": "2026-08-29", "mes": "2026-08", "tipo": "ENTRADA",
             "detalle": "Entrada a bodega", "entra": 10, "sale": None, "existencia": 10,
             "precio_unitario": 90, "total": 900, "ganancia": None, "a_costo": False},
            {"fecha": "2026-08-30", "mes": "2026-08", "tipo": "SALIDA",
             "detalle": "Sin avión · a costo", "entra": None, "sale": 3, "existencia": 7,
             "precio_unitario": 90, "total": 270, "ganancia": 0, "a_costo": True},
            {"fecha": "2026-09-10", "mes": "2026-09", "tipo": "ENTRADA",
             "detalle": "Entrada a bodega", "entra": 5, "sale": None, "existencia": 12,
             "precio_unitario": 108, "total": 540, "ganancia": None, "a_costo": False},
            {"fecha": "2026-10-02", "mes": "2026-10", "tipo": "SALIDA",
             "detalle": "XA-VGV", "entra": None, "sale": 2, "existencia": 10,
             "precio_unitario": 144, "total": 288, "ganancia": 72, "a_costo": False},
        ],
        "meses": [
            {"mes": "2026-08", "etiqueta": "Agosto 2026", "stock_inicio": 0, "entradas": 10,
             "compras": 10, "unidades_vendidas": 3, "stock_final": 7, "movimientos": 2,
             "vendido": None, "ganancia": None, "en_curso": False},
            {"mes": "2026-09", "etiqueta": "Septiembre 2026", "stock_inicio": 7, "entradas": 5,
             "compras": 5, "unidades_vendidas": 0, "stock_final": 12, "movimientos": 1,
             "vendido": None, "ganancia": None, "en_curso": False},
            {"mes": "2026-10", "etiqueta": "Octubre 2026", "stock_inicio": 12, "entradas": 0,
             "compras": 0, "unidades_vendidas": 2, "stock_final": 10, "movimientos": 1,
             "vendido": 288, "ganancia": 72, "en_curso": False},
            {"mes": "2026-11", "etiqueta": "Noviembre 2026 (en curso)", "stock_inicio": 10,
             "entradas": 0, "compras": 0, "unidades_vendidas": 0, "stock_final": 10,
             "movimientos": 0, "vendido": None, "ganancia": None, "en_curso": True},
        ],
        "total_compra": 1440,
        "total_venta": 558,
        "total_ganancia": 72,
    }


def _libro(payload: dict):
    req = CardexLibroRequest.model_validate(payload)
    return load_workbook(BytesIO(render_cardex_libro_xlsx(req)))


def _fila(ws, r: int, n: int) -> list:
    return [ws.cell(row=r, column=c).value for c in range(1, n + 1)]


def test_sin_movimientos_el_libro_de_siempre_byte_a_byte() -> None:
    assert _miembros(render_cardex_libro_xlsx(_req())) == _MIEMBROS_SIN_MOVIMIENTOS
    # `meses` sin `movimientos` (o en null) no enciende el formato nuevo.
    for extra in ({"meses": None}, {"meses": _payload()["meses"]}, {"movimientos": None}):
        req = CardexLibroRequest.model_validate({**_req().model_dump(), **extra})
        assert _miembros(render_cardex_libro_xlsx(req)) == _MIEMBROS_SIN_MOVIMIENTOS


def test_con_movimientos_tres_hojas() -> None:
    wb = _libro(_payload())
    assert wb.sheetnames == ["Cardex", "Resumen por mes", "Entradas y salidas"]


def test_lista_por_fecha_con_banda_por_mes() -> None:
    ws = _libro(_payload())["Cardex"]
    assert ws.cell(row=1, column=1).value == "Cardex — Aceite Semisintético 15W-50"
    assert ws.cell(row=4, column=1).value == "Fecha"
    assert ws.cell(row=4, column=6).value == "Existencia\n(después)"
    assert ws.cell(row=5, column=1).value == (
        "Agosto 2026 · Stock al inicio: 0 · Entradas: 10 · Vendidas: 3 · Stock al final: 7"
    )
    # Fecha real (Excel la filtra), cantidades enteras a la vista.
    assert _fila(ws, 6, 9) == [
        datetime(2026, 8, 29), "Compra", "Entrada a bodega", 10, None, 10, 90, 900, None,
    ]
    assert ws.cell(row=6, column=1).number_format == "dd/mm/yyyy"
    assert _fila(ws, 7, 9) == [
        datetime(2026, 8, 30), "Salida a costo", "Sin avión · a costo", None, 3, 7, 90, 270, 0,
    ]
    assert ws.cell(row=8, column=1).value.startswith("Septiembre 2026 · Stock al inicio: 7")
    assert _fila(ws, 9, 6) == [datetime(2026, 9, 10), "Compra", "Entrada a bodega", 5, None, 12]
    assert ws.cell(row=10, column=1).value.startswith("Octubre 2026 · Stock al inicio: 12")
    assert _fila(ws, 11, 9) == [
        datetime(2026, 10, 2), "Venta", "XA-VGV", None, 2, 10, 144, 288, 72,
    ]
    # Un mes sin movimientos también tiene su banda (el stock se arrastra).
    assert ws.cell(row=12, column=1).value == (
        "Noviembre 2026 (en curso) · Stock al inicio: 10 · Entradas: 0 · Vendidas: 0"
        " · Stock al final: 10 · sin movimientos"
    )
    # Totales del API y existencia de hoy.
    assert _fila(ws, 14, 9)[2] == "TOTAL COMPRAS" and ws.cell(row=14, column=8).value == 1440
    assert [ws.cell(row=15, column=c).value for c in (3, 8, 9)] == [
        "TOTAL VENTAS / GANANCIA", 558, 72,
    ]
    assert [ws.cell(row=16, column=c).value for c in (3, 6)] == ["EXISTENCIA ACTUAL", 10]
    assert ws.freeze_panes == "A5"


def test_resumen_por_mes() -> None:
    ws = _libro(_payload())["Resumen por mes"]
    assert _fila(ws, 4, 8) == [
        "Mes", "Stock al inicio\ndel mes", "Entradas", "Unidades\nvendidas",
        "Stock al final\ndel mes", "Movimientos", "Vendido", "Utilidad",
    ]
    assert _fila(ws, 5, 8) == ["Agosto 2026", 0, 10, 3, 7, 2, None, None]
    assert _fila(ws, 6, 8) == ["Septiembre 2026", 7, 5, 0, 12, 1, None, None]
    assert _fila(ws, 7, 8) == ["Octubre 2026", 12, 0, 2, 10, 1, 288, 72]
    assert _fila(ws, 8, 8) == ["Noviembre 2026 (en curso)", 10, 0, 0, 10, 0, None, None]
    # TOTAL: re-suma de columnas; el stock final es el del último mes.
    assert _fila(ws, 9, 8) == ["TOTAL", None, 15, 5, 10, 4, 288, 72]
    # Cuadre que pidió el cliente, fila por fila.
    for r in range(5, 9):
        ini, ent, ven, fin = (ws.cell(row=r, column=c).value for c in (2, 3, 4, 5))
        assert fin == ini + ent - ven


def test_la_hoja_de_siempre_va_al_final_identica() -> None:
    nueva = _libro(_payload())["Entradas y salidas"]
    vieja_payload = {k: v for k, v in _payload().items() if k not in ("movimientos", "meses")}
    vieja = _libro(vieja_payload).active
    assert vieja.title == "Cardex libro"
    celdas = lambda ws: [  # noqa: E731
        (c.coordinate, c.value, c.number_format)
        for fila in ws.iter_rows() for c in fila
    ]
    assert celdas(nueva) == celdas(vieja)
    rangos = lambda ws: sorted(map(str, ws.merged_cells.ranges))  # noqa: E731
    assert rangos(nueva) == rangos(vieja)


def test_lista_vacia() -> None:
    p = _payload() | {"movimientos": [], "meses": [], "entradas": [], "salidas": []}
    wb = _libro(p)
    assert wb["Cardex"].cell(row=5, column=1).value == "Sin movimientos."
    assert wb["Resumen por mes"].cell(row=5, column=1).value == "Sin movimientos."


def test_ruta_acepta_el_payload_del_api(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", "secreto-de-prueba")
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    res = TestClient(app).post(
        "/pdf/cardex-libro-xlsx",
        json=_payload(),
        headers={"X-Internal-Token": "secreto-de-prueba"},
    )
    assert res.status_code == 200, res.text
    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["Cardex", "Resumen por mes", "Entradas y salidas"]
