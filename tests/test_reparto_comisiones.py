"""Reparto a socios (PDF y Excel) con las COMISIONES a cargo del avión (regla
de septiembre 2026, API 0.0.65, 6-oct-2026).

El API suma en `comisiones_venta_usd` —el campo de siempre, que la cascada ya
resta antes del saldo— la parte del avión de la comisión bancaria de sus
cobros y la PROVISIÓN de la comisión del vendedor. Revisión 6-oct-2026: el
PDF y el Excel que reciben los socios imprimían a la vez «(-) Comisiones de
venta» y «su pago no es costo del avión» / «no del avión». Ahora, SOLO
cuando un avión absorbe comisiones, la fila y la columna se llaman
«Comisiones (banco + vendedor)» y las frases dicen qué absorbe el avión; sin
comisiones absorbidas el documento es el de siempre (comparado contra el
código anterior: flujos de contenido del PDF y miembros del .xlsx
idénticos, con y sin `regla_comisiones`).

13-oct-2026 (API 0.0.74, invariante 61 del API; decisión 3 del lote del
cliente: «comisión del vendedor SOLO como gasto real»): lo del vendedor en
`comisiones_venta_usd` es su PAGO REAL (el gasto «Comisión del vendedor»
capturado), ya no la provisión. Con la bandera ADITIVA
`comision_vendedor_solo_gasto_real` las frases dicen «pago real»; sin ella
(ausente/None/False) los documentos de siempre. Ajuste del mismo día (dos
categorías, misma bandera): lo del vendedor es el pago capturado como
«Comisión del vendedor (avión)» —en cualquier fecha— y el de «Comisión del
vendedor (VuelaTour)» lo paga VuelaTour: las frases nombran las dos y la
línea gris dice «pago «(avión)» resta en Comisiones» (88 caracteres, cabe).
"""

from __future__ import annotations

import re
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth

from app.config import get_settings
from app.main import app
from app.schemas.reparto import RepartoPdfRequest
from app.services.reparto_comisiones import (
    ETIQUETA_COLUMNA,
    ETIQUETA_FILA,
    LINEA_COMISION_VENDEDOR,
    LINEA_COMISION_VENDEDOR_GASTO_REAL,
    absorbe_comisiones,
    comisiones_al_avion,
    frase_regla,
    linea_comision_vendedor,
    solo_gasto_real,
)
from app.services.reparto_pdf import render_reparto_pdf
from app.services.reparto_xlsx import render_reparto_xlsx

# Frases de antes que, con comisiones absorbidas, dejaron de ser verdad.
_VIEJAS = ("su pago no es costo del avión", "no del avión", "Comisiones de venta")
_DESGLOSE = {"tuas_usd": 100.0, "extras_usd": 0, "pernocta_usd": 0, "comision_usd": 80.0,
             "iva_usd": 28.8}  # fmt: skip


def _avion(matricula: str, comisiones: float, *, desglose: bool = True) -> dict:
    """Un avión con la forma de `buildRepartoPayload` del API 0.0.65."""
    saldo = round(2000.0 - 500.0 - comisiones, 2)
    return {
        "matricula": matricula,
        "modelo": "C206",
        "ingresos_cobrado_usd": 2000.0,
        "otros_ingresos_vuelatour_usd": 208.8 if desglose else 0,
        "otros_ingresos_vuelatour_desglose": _DESGLOSE if desglose else None,
        "comisiones_venta_usd": comisiones,
        "pendiente_cobro_usd": 0,
        "pendiente_bruto_usd": 0,
        "horas_voladas_hr": 3.2,
        "gastos_directos_usd": 500.0,
        "gastos_indirectos_usd": 0,
        "permisos_usd": 0,
        "otros_usd": 0,
        "reserva_overhaul_usd": 0,
        "saldo_usd": saldo,
        "gastos_sin_tc_mxn": 0,
        "cobros_sin_tc_mxn": 0,
        "gastos_tc_oficial": None,
        "cobros_tc_oficial_count": 0,
        "reserva_incompleta": False,
        "reparto_porcentaje_total": 100,
        "reparto": [{"socio_nombre": "Socio A", "porcentaje": 100, "monto_usd": saldo}],
        "vuelos": [],
    }


def _payload(comisiones: float = 95.7, **extra) -> dict:
    """XB-TST absorbe `comisiones` (parte de la comisión bancaria + provisión
    del vendedor); XB-DOS no trae ninguna."""
    return {
        "periodo_desde": "2026-09-01",
        "periodo_hasta": "2026-09-30",
        "generado": "2026-10-06",
        "otros_ingresos_vuelatour_total_usd": 208.8,
        "otros_ingresos_vuelatour_desglose": _DESGLOSE,
        "tc_oficial": None,
        "aviones": [_avion("XB-TST", comisiones), _avion("XB-DOS", 0, desglose=False)],
        **extra,
    }


def _req(comisiones: float = 95.7, **extra) -> RepartoPdfRequest:
    return RepartoPdfRequest.model_validate(_payload(comisiones, **extra))


def _texto_pdf(data: bytes) -> str:
    """Texto del PDF en una sola línea (los párrafos se parten en renglones)."""
    texto = " ".join(p.extract_text() for p in PdfReader(BytesIO(data)).pages)
    return re.sub(r"\s+", " ", texto)


def _textos_xlsx(ws) -> list[str]:
    return [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]


# ---------------------------------------------------------------------------
# Cuándo aplican los textos nuevos
# ---------------------------------------------------------------------------


def test_cuando_un_avion_absorbe_comisiones() -> None:
    """Mismo criterio que la fila del PDF y la columna E del Excel."""
    assert comisiones_al_avion(_req())
    assert not comisiones_al_avion(_req(0))
    req = _req()
    assert [absorbe_comisiones(a) for a in req.aviones] == [True, False]
    # Payload sin la llave (API viejo): 0 por default.
    sin_llave = _payload()
    for avion in sin_llave["aviones"]:
        avion.pop("comisiones_venta_usd")
    assert not comisiones_al_avion(RepartoPdfRequest.model_validate(sin_llave))


def test_frase_con_la_regla_del_api_o_la_de_siempre() -> None:
    assert frase_regla(_req()) == (
        "Con la regla de septiembre 2026 (por fecha del vuelo), el avión absorbe la "
        "provisión de la comisión del vendedor y su parte de la comisión bancaria de sus cobros"
    )
    req = _req(regla_comisiones="regla desde 15-oct-2026")
    assert frase_regla(req, "con").startswith("con la regla desde 15-oct-2026 (por fecha")
    # Liberal: sin texto ⇒ el respaldo; sin «regla» ⇒ se antepone.
    assert _req(regla_comisiones=7).regla_comisiones is None
    assert _req(regla_comisiones="  ").regla_comisiones is None
    assert _req(regla_comisiones="sep-2026").regla_comisiones == "regla sep-2026"


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def test_pdf_con_comisiones_dice_lo_que_absorbe_el_avion() -> None:
    texto = _texto_pdf(render_reparto_pdf(_req()))
    assert f"{ETIQUETA_FILA} $-95.70" in texto
    for vieja in _VIEJAS:
        assert vieja not in texto, vieja
    frase = frase_regla(_req())
    # Bajo el resumen, al pie y en la línea gris del avión que las absorbe.
    assert f"{frase}: restan en {ETIQUETA_FILA}." in texto
    assert f"{frase}: restan antes del saldo en {ETIQUETA_FILA}" in texto
    assert LINEA_COMISION_VENDEDOR.strip() in texto
    # «Gastos del periodo» del resumen incluye las comisiones: venta − gastos
    # = saldo (2 × 2000 − 2 × 500 − 95.70).
    assert "$1,095.70" in texto and "$2,904.30" in texto


def test_pdf_sin_comisiones_conserva_los_textos_de_siempre() -> None:
    texto = _texto_pdf(render_reparto_pdf(_req(0)))
    assert "Comisiones (banco + vendedor)" not in texto
    assert "Comisiones de venta" not in texto  # con 0, la fila no se imprime
    assert "incluye comisión del vendedor cotizada (pre-IVA) — su pago no es costo del avión" in (
        texto
    )
    assert "El pago de la comisión al vendedor sale de VuelaTour (otros movimientos), no del" in (
        texto
    )
    assert "y el pago de la comisión al vendedor sale de VuelaTour, no del avión" in texto


def test_pdf_nombra_la_regla_del_api() -> None:
    texto = _texto_pdf(render_reparto_pdf(_req(regla_comisiones="regla desde 15-oct-2026")))
    assert "Con la regla desde 15-oct-2026 (por fecha del vuelo)" in texto
    assert "septiembre" not in texto


def test_linea_del_avion_cabe_en_su_columna() -> None:
    """La línea gris va en una celda de 120 mm (6 pt de relleno por lado) a
    7.5 pt: más larga se encimaría con el monto. También la del API 0.0.74."""
    for linea in (LINEA_COMISION_VENDEDOR, LINEA_COMISION_VENDEDOR_GASTO_REAL):
        ancho = stringWidth(linea, "Helvetica-Oblique", 7.5)
        assert ancho <= 120 * mm - 12, linea
    assert abs(len(LINEA_COMISION_VENDEDOR_GASTO_REAL) - len(LINEA_COMISION_VENDEDOR)) <= 2


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def test_xlsx_con_comisiones_columna_visible_y_notas_nuevas() -> None:
    ws = load_workbook(BytesIO(render_reparto_xlsx(_req()))).active
    assert ws["E4"].value == ETIQUETA_COLUMNA
    assert not ws.column_dimensions["E"].hidden
    assert (ws["E5"].value, ws["E6"].value) == (95.7, 0)
    assert ws["A7"].value == "TOTAL" and ws["E7"].value == 95.7
    textos = _textos_xlsx(ws)
    for vieja in _VIEJAS[:2]:
        assert not any(vieja in t for t in textos), vieja
    nota = ws["A8"].value
    assert f". {ETIQUETA_COLUMNA}: {frase_regla(_req(), 'con')}; restan antes del saldo" in nota
    assert nota.startswith("Venta del avión cobrada = ")
    assert nota.endswith("solo cuando es mayor a la parte del avión.")
    assert ws["A10"].value == (
        "Comisión del vendedor cotizada por avión (ingreso de VuelaTour; con la regla de "
        "septiembre 2026, por fecha del vuelo, su provisión la absorbe el avión en Comisiones): "
        "XB-TST $80.00."
    )


def test_xlsx_sin_comisiones_igual_que_siempre() -> None:
    ws = load_workbook(BytesIO(render_reparto_xlsx(_req(0)))).active
    assert ws["E4"].value == "Comisiones venta"
    assert ws.column_dimensions["E"].hidden
    assert "; el pago de la comisión al vendedor sale de VuelaTour, no del avión" in ws["A8"].value
    assert ws["A10"].value == (
        "Comisión del vendedor cotizada por avión (ingreso de VuelaTour, su pago no es costo "
        "del avión): XB-TST $80.00."
    )


def test_xlsx_nombra_la_regla_del_api() -> None:
    req = _req(regla_comisiones="regla desde 15-oct-2026")
    textos = _textos_xlsx(load_workbook(BytesIO(render_reparto_xlsx(req))).active)
    assert any("con la regla desde 15-oct-2026 (por fecha del vuelo)" in t for t in textos)
    assert not any("septiembre" in t for t in textos)


# ---------------------------------------------------------------------------
# Rutas: el JSON del API 0.0.65 entra sin 422
# ---------------------------------------------------------------------------

TOKEN = "secreto-de-prueba"
client = TestClient(app)


def test_rutas_aceptan_el_payload_del_api(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    payload = _payload(regla_comisiones=None)
    res = client.post("/pdf/reparto-xlsx", json=payload, headers=headers)
    assert res.status_code == 200, res.text
    assert load_workbook(BytesIO(res.content)).active["E4"].value == ETIQUETA_COLUMNA
    res = client.post("/pdf/reparto", json={**payload, "regla_comisiones": ["x"]}, headers=headers)
    assert res.status_code == 200, res.text
    assert ETIQUETA_FILA in _texto_pdf(res.content)


# ===========================================================================
# 13-oct-2026 (API 0.0.74, invariante 61): «comisión del vendedor SOLO como
# gasto real». Bandera `comision_vendedor_solo_gasto_real`. Los literales
# son los del ajuste «dos categorías» del mismo día (contrato del ajuste),
# que reescribió EN SU LUGAR los de la decisión 3.
# ===========================================================================

FRASE_GASTO_REAL = (
    "Con la regla de septiembre 2026 (por fecha del vuelo), el avión absorbe su parte de "
    "la comisión bancaria de sus cobros y, en cualquier fecha, el pago al vendedor "
    "capturado como «Comisión del vendedor (avión)» (el de «Comisión del vendedor "
    "(VuelaTour)» lo paga VuelaTour)"
)
LINEA_GASTO_REAL = (
    "   incluye comisión del vendedor cotizada (pre-IVA) — pago «(avión)» resta en Comisiones"
)
NOTA_OTROS_INGRESOS_COLA = (
    "Lo cobrado al cliente por la comisión sigue siendo de VuelaTour; el pago capturado "
    "como «Comisión del vendedor (avión)» lo carga el avión ('otros movimientos' del "
    "Balance general lo muestra con egreso $0.00) y el de «Comisión del vendedor "
    "(VuelaTour)» lo paga VuelaTour ahí mismo."
)
NOTA_POR_AVION_GASTO_REAL = (
    "Comisión del vendedor cotizada por avión (ingreso de VuelaTour; su pago, si se "
    "captura como «Comisión del vendedor (avión)», lo absorbe el avión en Comisiones (el "
    "de «(VuelaTour)» lo paga VuelaTour)): XB-TST $80.00."
)


def test_bandera_aditiva_con_default_none_y_textos_del_contrato() -> None:
    assert _req().comision_vendedor_solo_gasto_real is None
    assert not solo_gasto_real(_req())
    for valor in (None, False):
        assert not solo_gasto_real(_req(comision_vendedor_solo_gasto_real=valor))
    req = _req(comision_vendedor_solo_gasto_real=True)
    assert solo_gasto_real(req)
    assert frase_regla(req) == FRASE_GASTO_REAL
    assert LINEA_COMISION_VENDEDOR_GASTO_REAL == LINEA_GASTO_REAL
    assert linea_comision_vendedor(req) == LINEA_GASTO_REAL
    assert linea_comision_vendedor(_req()) == LINEA_COMISION_VENDEDOR
    for t in (
        FRASE_GASTO_REAL,
        LINEA_GASTO_REAL,
        NOTA_OTROS_INGRESOS_COLA,
        NOTA_POR_AVION_GASTO_REAL,
    ):
        assert "provisi" not in t.lower()
        # Ninguna nombra la categoría vieja (sin «(VuelaTour)» / «(avión)»).
        assert "«Comisión del vendedor»" not in t
    # La frase con «con» (nota de la tabla del Excel) es la misma.
    assert frase_regla(req, "con") == "c" + FRASE_GASTO_REAL[1:]


def test_linea_gris_de_las_dos_categorias_medida() -> None:
    """La línea del ajuste mide 88 caracteres (|Δ| = 2 con la de siempre) y
    ~301.4 pt en Helvetica-Oblique 7.5: cabe en la celda de 120 mm (tope
    328.2 pt)."""
    assert len(LINEA_GASTO_REAL) == 88
    assert abs(len(LINEA_GASTO_REAL) - len(LINEA_COMISION_VENDEDOR)) == 2
    ancho = stringWidth(LINEA_GASTO_REAL, "Helvetica-Oblique", 7.5)
    assert round(ancho, 1) == 301.4
    assert ancho <= 328.2 and ancho <= 120 * mm - 12


def test_pdf_con_la_bandera_dice_pago_real_y_nada_de_provision() -> None:
    req = _req(comision_vendedor_solo_gasto_real=True)
    texto = _texto_pdf(render_reparto_pdf(req))
    assert f"{ETIQUETA_FILA} $-95.70" in texto
    assert f"{FRASE_GASTO_REAL}: restan en {ETIQUETA_FILA}. {NOTA_OTROS_INGRESOS_COLA}" in texto
    assert f"{FRASE_GASTO_REAL}: restan antes del saldo en {ETIQUETA_FILA}" in texto
    assert LINEA_GASTO_REAL.strip() in texto
    assert "provisi" not in texto.lower()
    for vieja in _VIEJAS:
        assert vieja not in texto, vieja
    # Sin comisiones absorbidas la bandera no cambia nada: el de siempre.
    assert _texto_pdf(
        render_reparto_pdf(_req(0, comision_vendedor_solo_gasto_real=True))
    ) == _texto_pdf(render_reparto_pdf(_req(0)))


def test_pdf_sin_la_bandera_el_de_siempre() -> None:
    base = _texto_pdf(render_reparto_pdf(_req()))
    for valor in (None, False):
        assert _texto_pdf(render_reparto_pdf(_req(comision_vendedor_solo_gasto_real=valor))) == base
    assert "su provisión resta en Comisiones" in base
    assert "que le paga al vendedor (otros movimientos)" in base


def test_xlsx_con_la_bandera_dice_pago_real() -> None:
    req = _req(comision_vendedor_solo_gasto_real=True)
    ws = load_workbook(BytesIO(render_reparto_xlsx(req))).active
    assert ws["E4"].value == ETIQUETA_COLUMNA
    assert f". {ETIQUETA_COLUMNA}: {frase_regla(req, 'con')}; restan antes del saldo" in (
        ws["A8"].value
    )
    frase = (
        "con la regla de septiembre 2026 (por fecha del vuelo), el avión absorbe su parte "
        "de la comisión bancaria de sus cobros y, en cualquier fecha, el pago al vendedor "
        "capturado como «Comisión del vendedor (avión)»"
    )
    assert frase in ws["A8"].value
    assert ws["A10"].value == NOTA_POR_AVION_GASTO_REAL
    assert not any("provisi" in t.lower() for t in _textos_xlsx(ws))


def test_xlsx_sin_la_bandera_miembros_identicos() -> None:
    """Sin la bandera (ausente/None/False) el .xlsx es el de siempre, miembro
    a miembro (salvo `docProps/core.xml`, que lleva la hora)."""
    import zipfile

    def miembros(data: bytes) -> dict[str, bytes]:
        z = zipfile.ZipFile(BytesIO(data))
        return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}

    base = miembros(render_reparto_xlsx(_req()))
    for valor in (None, False):
        assert miembros(render_reparto_xlsx(_req(comision_vendedor_solo_gasto_real=valor))) == base


def test_con_la_bandera_y_la_regla_del_api_nadie_dice_septiembre() -> None:
    req = _req(comision_vendedor_solo_gasto_real=True, regla_comisiones="regla desde 15-oct-2026")
    texto = _texto_pdf(render_reparto_pdf(req))
    frase = (
        "Con la regla desde 15-oct-2026 (por fecha del vuelo), el avión absorbe su parte "
        "de la comisión bancaria de sus cobros y, en cualquier fecha, el pago al vendedor"
    )
    assert frase in texto
    assert "septiembre" not in texto
    textos = _textos_xlsx(load_workbook(BytesIO(render_reparto_xlsx(req))).active)
    assert not any("septiembre" in t for t in textos)
    assert any(
        "si se captura como «Comisión del vendedor (avión)», lo absorbe el avión en Comisiones"
        in t
        for t in textos
    )


def test_rutas_aceptan_el_payload_del_api_0074(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    payload = _payload(comision_vendedor_solo_gasto_real=True)
    res = client.post("/pdf/reparto-xlsx", json=payload, headers=headers)
    assert res.status_code == 200, res.text
    assert load_workbook(BytesIO(res.content)).active["A10"].value == NOTA_POR_AVION_GASTO_REAL
    res = client.post("/pdf/reparto", json=payload, headers=headers)
    assert res.status_code == 200, res.text
    assert LINEA_GASTO_REAL.strip() in _texto_pdf(res.content)
