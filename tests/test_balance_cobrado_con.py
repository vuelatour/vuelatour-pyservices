"""Nota «cómo se cobró» en cada parcialidad del balance (6-oct-2026, API 0.0.60).

Pedido del cliente, con las capturas del Excel «reporte horas FLOTA» y del
panel (que muestra por cobro «$136,856.80 MXN · Conciliado · Transferencia →
Scotiabank Pesos · Registró: Itzi»): «al lado de la columna STATUS, si ya se
pagó, que venga la misma información de cómo se cobró, quién lo cobró y, si es
posible, a qué cuenta». Y luego: «mejor, después de cada cobro venga la nota
con los detalles, para no tener tantas columnas nuevas».

SIN columnas nuevas: cada celda COBRO n MXN lleva una NOTA de Excel
«fecha · monto · cómo se cobró» y la celda STATUS (mismo texto de siempre)
una nota resumen con todas las líneas. El «cómo se cobró» llega ARMADO del API
(`cobrado_con`, fuente única `etiquetaCobradoCon`); un API ≤ 0.0.59 no lo
manda y la nota se compone con `metodo`/`cuenta`. Sin método, cuenta ni
registro no hay nota, y un libro sin cobros es el de siempre, byte a byte.
"""

from __future__ import annotations

import hashlib
import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.config import get_settings
from app.main import app
from app.schemas.reportes import (
    BalanceAvionCobro,
    BalanceAvionRequest,
    BalanceGeneralRequest,
)
from app.services.balance_avion_xlsx import (
    _COLS,
    _DISP_MENSUAL,
    _comentario_cobro,
    _nota_cobro,
    _nota_resumen_cobros,
    _notas_parcialidades,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)

FILA = 3  # primera fila de vuelo (1-2 = encabezado de grupo / columna)
COL_STATUS = next(i for i, c in enumerate(_COLS, start=1) if c[2] == "status_cobro")
COLS_FECHA = [_DISP_MENSUAL.cobro1_col + 2 * k for k in range(4)]
COLS_MXN = [_DISP_MENSUAL.cobro1_col + 1 + 2 * k for k in range(4)]
COLS_NOTA = {COL_STATUS, *COLS_MXN}

# Cobros con la forma que manda el API 0.0.60 (nombres reales de quienes
# registran en prod; cuentas tal como las teclea la oficina).
COBRO_TRANSFER = {
    "fecha": "2026-09-15",
    "monto_mxn": 12522.2,
    "metodo": "TRANSFERENCIA",
    "comision_mxn": None,
    "cuenta": "Scotiabank Pesos",
    "metodo_etiqueta": "Transferencia",
    "registro": "Itzi",
    "cobrado_con": "Transferencia → Scotiabank Pesos · Itzi",
}
COBRO_PAYWISE = {
    "fecha": "2026-09-22",
    "monto_mxn": 136856.8,
    "metodo": "PAYWISE",
    "comision_mxn": None,
    "cuenta": "Paywise",
    "metodo_etiqueta": "Link de pago (Paywise)",
    "registro": "Pablo Canales",
    "cobrado_con": "Link de pago (Paywise) → Paywise · Pablo Canales",
}
COBRO_EFECTIVO = {
    "fecha": "2026-09-28",
    "monto_mxn": 8000.0,
    "metodo": "EFECTIVO",
    "comision_mxn": None,
    "cuenta": None,
    "metodo_etiqueta": "Efectivo",
    "registro": "Alejandro Canales",
    "cobrado_con": "Efectivo · Alejandro Canales",
}
LINEA_TRANSFER = "15/09/2026 · $12,522.20 · Transferencia → Scotiabank Pesos · Itzi"
LINEA_PAYWISE = "22/09/2026 · $136,856.80 · Link de pago (Paywise) → Paywise · Pablo Canales"
LINEA_EFECTIVO = "28/09/2026 · $8,000.00 · Efectivo · Alejandro Canales"

_CLAVES_NUEVAS = ("metodo_etiqueta", "registro", "cobrado_con")


def _api_viejo(cobro: dict) -> dict:
    """El mismo cobro como lo manda un API ≤ 0.0.59 (sin las tres llaves)."""
    return {k: v for k, v in cobro.items() if k not in _CLAVES_NUEVAS}


def _sin_detalle(cobro: dict) -> dict:
    """Solo fecha y monto (fixtures viejos: sin método, cuenta ni registro)."""
    return {"fecha": cobro["fecha"], "monto_mxn": cobro["monto_mxn"]}


def _vuelo(clave: str, fecha: str, cobros: list[dict], status: str, **extra) -> dict:
    cobrado = round(sum(c["monto_mxn"] for c in cobros), 2) if cobros else None
    return {
        "clave": clave,
        "fecha": fecha,
        "ruta": "CUN-CZM-CUN",
        "estado": "COMPLETADO",
        "status_cobro": status,
        "cobros": cobros,
        "cobrado_real_mxn": cobrado,
        "cobrado_mxn": cobrado,
        **extra,
    }


def _vuelos(transforma=lambda c: c, *, con_cobros: bool = True) -> list[dict]:
    """Septiembre: #301 con DOS cobros (transferencia + Paywise), #302 sin
    cobros y #303 con un cobro en efectivo (sin cuenta)."""
    if not con_cobros:
        return [
            _vuelo("#301 · Cliente Uno", "2026-09-15", [], "Pendiente"),
            _vuelo("#302 · Cliente Dos", "2026-09-18", [], "Pendiente"),
            _vuelo("#303 · Cliente Tres", "2026-09-28", [], "Pendiente"),
        ]
    return [
        _vuelo(
            "#301 · Cliente Uno",
            "2026-09-15",
            [transforma(COBRO_TRANSFER), transforma(COBRO_PAYWISE)],
            "Cobrado",
        ),
        _vuelo("#302 · Cliente Dos", "2026-09-18", [], "Pendiente"),
        _vuelo("#303 · Cliente Tres", "2026-09-28", [transforma(COBRO_EFECTIVO)], "Parcial"),
    ]


def _totales(vuelos: list[dict]) -> dict:
    cobrado = [v["cobrado_real_mxn"] for v in vuelos if v["cobrado_real_mxn"] is not None]
    total = round(sum(cobrado), 2) if cobrado else None
    return {"cobrado_real_mxn": total, "cobrado_mxn": total}


def _individual(vuelos: list[dict]) -> BalanceAvionRequest:
    return BalanceAvionRequest(
        matricula="XA-VGV",
        periodo_desde="2026-09-01",
        periodo_hasta="2026-09-30",
        vuelos=vuelos,
        totales=_totales(vuelos),
    )


def _general(vuelos: list[dict]) -> BalanceGeneralRequest:
    pintados = [{**v, "avion_color": "#2563EB"} for v in vuelos]
    return BalanceGeneralRequest(
        periodo_desde="2026-09-01",
        periodo_hasta="2026-09-30",
        resumen=[{"matricula": "XA-VGV", "color": "#2563EB", "vuelos": len(vuelos)}],
        consolidado=BalanceAvionRequest(
            matricula="FLOTA",
            periodo_desde="2026-09-01",
            periodo_hasta="2026-09-30",
            vuelos=pintados,
            totales=_totales(pintados),
        ),
    )


def _maestra(data: bytes, nombre: str = "reporte horas XA-VGV"):
    return load_workbook(BytesIO(data))[nombre]


def _nota(ws, fila: int, col: int) -> str | None:
    c = ws.cell(row=fila, column=col)
    return c.comment.text if c.comment else None


def _notas_de_cobro(ws) -> dict[tuple[int, int], str]:
    """Notas en STATUS / COBRO n MXN de las filas de vuelo."""
    return {
        (c.row, c.column): c.comment.text
        for fila in ws.iter_rows(min_row=FILA)
        for c in fila
        if c.column in COLS_NOTA and c.comment
    }


def _firma_libro(data: bytes, *, sin_notas_cobro: bool = False) -> str:
    """Huella del LAYOUT de TODAS las hojas (misma receta que
    test_balance_empresa_vuelatour): valores o fórmulas, formatos, fuentes,
    rellenos, alineación, bordes, comentarios, merges, anchos, altos y panel
    fijo. `sin_notas_cobro` deja fuera SOLO las notas de STATUS / COBRO n
    MXN de las filas de vuelo de la hoja maestra (las de este cambio)."""
    wb = load_workbook(BytesIO(data))
    partes: list = []
    for ws in wb.worksheets:
        partes.append(ws.title)
        maestra = ws.title.startswith("reporte horas")
        for fila in ws.iter_rows():
            for c in fila:
                f = c.font
                nota = c.comment.text if c.comment else None
                if sin_notas_cobro and maestra and c.row >= FILA and c.column in COLS_NOTA:
                    nota = None
                partes.append([
                    c.coordinate, repr(c.value), c.number_format,
                    bool(f.b), bool(f.i), f.sz,
                    f.color.rgb if f.color is not None else None,
                    c.fill.fgColor.rgb if c.fill.fill_type else None,
                    c.alignment.horizontal, c.alignment.vertical,
                    bool(c.alignment.wrap_text), c.border.left.style,
                    nota,
                ])  # fmt: skip
        partes.append(sorted(str(r) for r in ws.merged_cells.ranges))
        partes.append(sorted((k, d.width) for k, d in ws.column_dimensions.items()))
        partes.append(sorted((k, d.height) for k, d in ws.row_dimensions.items() if d.height))
        partes.append(ws.freeze_panes)
    return hashlib.sha256(json.dumps(partes, default=str).encode()).hexdigest()


def _comentarios(data: bytes) -> dict[tuple[str, str], str]:
    wb = load_workbook(BytesIO(data))
    return {
        (ws.title, c.coordinate): c.comment.text
        for ws in wb.worksheets
        for fila in ws.iter_rows()
        for c in fila
        if c.comment
    }


# ---------------------------------------------------------------------------
# 1) Helpers puros.
# ---------------------------------------------------------------------------


def test_nota_cobro_pinta_el_texto_del_api_tras_fecha_y_monto() -> None:
    assert _nota_cobro(BalanceAvionCobro(**COBRO_TRANSFER)) == LINEA_TRANSFER
    assert _nota_cobro(BalanceAvionCobro(**COBRO_PAYWISE)) == LINEA_PAYWISE
    assert _nota_cobro(BalanceAvionCobro(**COBRO_EFECTIVO)) == LINEA_EFECTIVO


def test_cobrado_con_del_api_gana_sobre_metodo_y_cuenta() -> None:
    """El texto armado del API se pinta TAL CUAL, aunque método y cuenta
    dijeran otra cosa: aquí no se decide nada."""
    c = BalanceAvionCobro(**{**COBRO_TRANSFER, "cobrado_con": "Cheque → HSBC Pesos · Itzi"})
    assert _nota_cobro(c) == "15/09/2026 · $12,522.20 · Cheque → HSBC Pesos · Itzi"


@pytest.mark.parametrize(
    "cobro,esperado",
    [
        (
            {"metodo": "TRANSFERENCIA", "cuenta": "Scotiabank Pesos"},
            "15/09/2026 · $12,522.20 · TRANSFERENCIA → Scotiabank Pesos",
        ),
        ({"metodo": "EFECTIVO"}, "15/09/2026 · $12,522.20 · EFECTIVO"),
        ({"cuenta": "HSBC Dólares"}, "15/09/2026 · $12,522.20 · HSBC Dólares"),
        # Con la etiqueta del API 0.0.60 pero sin el texto armado: la etiqueta
        # gana sobre el código crudo.
        (
            {"metodo": "TRANSFERENCIA", "metodo_etiqueta": "Transferencia",
             "cuenta": "Scotiabank Pesos"},
            "15/09/2026 · $12,522.20 · Transferencia → Scotiabank Pesos",
        ),
        (
            {"metodo": "TRANSFERENCIA", "cuenta": "Scotiabank Pesos", "registro": "Itzi"},
            "15/09/2026 · $12,522.20 · TRANSFERENCIA → Scotiabank Pesos · Registró: Itzi",
        ),
        ({"registro": "Itzi"}, "15/09/2026 · $12,522.20 · Registró: Itzi"),
        # Multi-avión: el sufijo de la parte viaja en el método, tal cual.
        (
            {"metodo": "TRANSFERENCIA · parte de esta fila (50 % de la venta del "
             "avión + ingreso VuelaTour)", "cuenta": "HSBC Pesos"},
            "15/09/2026 · $12,522.20 · TRANSFERENCIA · parte de esta fila (50 % de "
            "la venta del avión + ingreso VuelaTour) → HSBC Pesos",
        ),
    ],
)  # fmt: skip
def test_api_viejo_compone_con_metodo_y_cuenta(cobro, esperado) -> None:
    c = BalanceAvionCobro(fecha="2026-09-15", monto_mxn=12522.2, **cobro)
    assert _nota_cobro(c) == esperado


@pytest.mark.parametrize(
    "cobro",
    [
        {},
        {"metodo": None, "cuenta": None, "registro": None, "cobrado_con": None},
        {"metodo": "—"},  # el API pinta «—» cuando el cobro no trae método
        {"metodo": "  ", "cuenta": "", "cobrado_con": "   "},
        {"comision_mxn": 150.25},
    ],
)
def test_sin_metodo_cuenta_ni_registro_no_hay_nota(cobro) -> None:
    c = BalanceAvionCobro(fecha="2026-09-15", monto_mxn=12522.2, **cobro)
    assert _nota_cobro(c) is None
    assert _nota_resumen_cobros([c, c]) is None
    assert _notas_parcialidades([c]) == [None]


def test_partes_que_faltan_y_reembolso() -> None:
    base = {"metodo_etiqueta": "Transferencia", "cuenta": "HSBC Pesos"}
    # Cobro USD sin TC: monto None ⇒ la nota no inventa un $0.
    assert _nota_cobro(BalanceAvionCobro(fecha="2026-09-15", **base)) == (
        "15/09/2026 · Transferencia → HSBC Pesos"
    )
    assert _nota_cobro(BalanceAvionCobro(monto_mxn=100.0, **base)) == (
        "$100.00 · Transferencia → HSBC Pesos"
    )
    # Reembolso = cobro negativo: el signo va antes del «$».
    assert _nota_cobro(BalanceAvionCobro(fecha="2026-09-30", monto_mxn=-500.0, **base)) == (
        "30/09/2026 · -$500.00 · Transferencia → HSBC Pesos"
    )


def test_resumen_una_linea_por_cobro_en_orden() -> None:
    cobros = [BalanceAvionCobro(**c) for c in (COBRO_TRANSFER, COBRO_PAYWISE)]
    assert _nota_resumen_cobros(cobros) == f"{LINEA_TRANSFER}\n{LINEA_PAYWISE}"
    assert _nota_resumen_cobros([]) is None
    # Si al menos uno dice cómo se cobró, el resumen lleva TODOS (los que no,
    # con su fecha y monto).
    mixto = [BalanceAvionCobro(**_sin_detalle(COBRO_TRANSFER)), cobros[1]]
    assert _nota_resumen_cobros(mixto) == f"15/09/2026 · $12,522.20\n{LINEA_PAYWISE}"


def test_comentario_del_tamano_de_las_de_desglose() -> None:
    com = _comentario_cobro(LINEA_EFECTIVO)
    assert com.author == "VuelaTour"
    assert (com.width, com.height) == (320, 70)
    # Muchas líneas largas (multi-avión, 4.ª agregada): crece y tiene tope.
    largo = "\n".join([LINEA_PAYWISE * 2] * 3)
    assert 70 < _comentario_cobro(largo).height <= 260
    assert _comentario_cobro("\n".join([LINEA_PAYWISE] * 40)).height == 260


def test_esquema_aditivo_y_liberal() -> None:
    c = BalanceAvionCobro()
    assert (c.metodo_etiqueta, c.registro, c.cobrado_con) == (None, None, None)
    c = BalanceAvionCobro(cobrado_con="  Efectivo · Itzi  ", registro="", metodo_etiqueta=123)
    assert (c.cobrado_con, c.registro, c.metodo_etiqueta) == ("Efectivo · Itzi", None, None)
    c = BalanceAvionCobro.model_validate({"cobrado_con": {"x": 1}, "registro": True})
    assert (c.cobrado_con, c.registro) == (None, None)


# ---------------------------------------------------------------------------
# 2) En el libro: dos notas + resumen, sin mover nada.
# ---------------------------------------------------------------------------


def test_fila_con_dos_cobros_lleva_dos_notas_y_el_resumen_en_status() -> None:
    ws = _maestra(render_balance_avion_xlsx(_individual(_vuelos())))
    # #301: COBRO 1 y COBRO 2 MXN con su nota; 3 y 4 sin cobro ni nota.
    assert _nota(ws, FILA, COLS_MXN[0]) == LINEA_TRANSFER
    assert _nota(ws, FILA, COLS_MXN[1]) == LINEA_PAYWISE
    assert ws.cell(row=FILA, column=COLS_MXN[0]).comment.author == "VuelaTour"
    assert [_nota(ws, FILA, c) for c in COLS_MXN[2:]] == [None, None]
    # Las celdas de FECHA no llevan nota (la nota va en el monto).
    assert all(_nota(ws, FILA, c) is None for c in COLS_FECHA)
    # STATUS: el mismo texto de siempre + nota resumen con las dos líneas.
    status = ws.cell(row=FILA, column=COL_STATUS)
    assert status.value == "Cobrado"
    assert status.comment.text == f"{LINEA_TRANSFER}\n{LINEA_PAYWISE}"
    # #302 sin cobros: ni una nota en el bloque.
    assert all(_nota(ws, FILA + 1, c) is None for c in COLS_NOTA)
    # #303 un cobro en efectivo (sin cuenta).
    assert _nota(ws, FILA + 2, COLS_MXN[0]) == LINEA_EFECTIVO
    assert _nota(ws, FILA + 2, COL_STATUS) == LINEA_EFECTIVO
    assert ws.cell(row=FILA + 2, column=COL_STATUS).value == "Parcial"
    # Montos y fechas en su celda, como siempre.
    assert ws.cell(row=FILA, column=COLS_FECHA[1]).value == "22/09/2026"
    assert ws.cell(row=FILA, column=COLS_MXN[1]).value == pytest.approx(136856.8)


# Huellas de los libros CON cobros (payload del API 0.0.60) calculadas con el
# código ANTERIOR a este cambio (HEAD 4125973), que ignoraba las llaves
# nuevas y no pintaba notas de cobro: quitando SOLO esas notas, el libro de
# hoy es el de antes celda por celda en TODAS las hojas (comentarios de
# desglose incluidos). Además se verificó que, fuera de los comentarios
# (xl/comments/comment1.xml y su dibujo VML), los miembros del .xlsx salían
# byte-idénticos.
_FIRMA_INDIVIDUAL_CON_COBROS = "95b5d2c61bc2ba04a382f9392c87f91eccd2a571464fd3e318f50ba10e2c9d95"
_FIRMA_GENERAL_CON_COBROS = "ec4e8713c82b0d65f701823219db281d10238f80c501a805f90a17f9a24b03e7"


def test_nada_se_mueve_solo_cambian_las_notas_de_status_y_cobros() -> None:
    """Con cobros: quitando las notas de STATUS / COBRO n MXN, los libros son
    los de antes (valores, fórmulas, formatos, merges, anchos, altos y las
    demás notas) y las únicas notas nuevas son esas."""
    con = render_balance_avion_xlsx(_individual(_vuelos()))
    assert _firma_libro(con, sin_notas_cobro=True) == _FIRMA_INDIVIDUAL_CON_COBROS
    gen = render_balance_general_xlsx(_general(_vuelos()))
    assert _firma_libro(gen, sin_notas_cobro=True) == _FIRMA_GENERAL_CON_COBROS
    sin = render_balance_avion_xlsx(_individual(_vuelos(_sin_detalle)))
    nuevas = set(_comentarios(con)) - set(_comentarios(sin))
    assert nuevas == {
        ("reporte horas XA-VGV", f"{get_column_letter(c)}{r}")
        for r, c in [
            (FILA, COLS_MXN[0]),
            (FILA, COLS_MXN[1]),
            (FILA, COL_STATUS),
            (FILA + 2, COLS_MXN[0]),
            (FILA + 2, COL_STATUS),
        ]
    }
    # Las demás notas, idénticas.
    comunes = set(_comentarios(sin))
    assert {k: v for k, v in _comentarios(con).items() if k in comunes} == _comentarios(sin)
    # Sin columnas nuevas: 47 columnas y los encabezados de siempre.
    ws = _maestra(con)
    assert ws.max_column == len(_COLS) == 47
    assert [ws.cell(row=2, column=i).value for i in range(1, 48)] == [c[1] for c in _COLS]


# Huellas del libro SIN cobros calculadas con el código ANTERIOR a este
# cambio (HEAD 4125973, 6-oct-2026) sobre los mismos payloads. Además se
# verificó que los .xlsx completos salían byte-idénticos (todos los miembros
# del zip salvo docProps/core.xml). Si una cambia, un periodo sin cobros ya
# no ve su libro de siempre.
_FIRMA_INDIVIDUAL_SIN_COBROS = "f027e88f4a21822e699281c82f3b5e2634a56fb20b064f00e474a0b735bb1b58"
_FIRMA_GENERAL_SIN_COBROS = "5a3dda62c5eac432792c5e6593f2f3cb182e9b9951c7a8d93243d3fd6296b980"


def test_sin_cobros_el_libro_es_el_de_siempre() -> None:
    ind = render_balance_avion_xlsx(_individual(_vuelos(con_cobros=False)))
    assert _firma_libro(ind) == _FIRMA_INDIVIDUAL_SIN_COBROS
    assert _notas_de_cobro(_maestra(ind)) == {}
    gen = render_balance_general_xlsx(_general(_vuelos(con_cobros=False)))
    assert _firma_libro(gen) == _FIRMA_GENERAL_SIN_COBROS
    assert _notas_de_cobro(_maestra(gen, "reporte horas FLOTA")) == {}


def test_cobros_sin_metodo_cuenta_ni_registro_no_llevan_nota() -> None:
    """Los payloads de prueba de siempre (solo fecha y monto) no ganan notas:
    sus huellas congeladas en los demás tests siguen valiendo."""
    ws = _maestra(render_balance_avion_xlsx(_individual(_vuelos(_sin_detalle))))
    assert _notas_de_cobro(ws) == {}


def test_api_viejo_sin_cobrado_con_compone_metodo_y_cuenta() -> None:
    ws = _maestra(render_balance_avion_xlsx(_individual(_vuelos(_api_viejo))))
    l1 = "15/09/2026 · $12,522.20 · TRANSFERENCIA → Scotiabank Pesos"
    l2 = "22/09/2026 · $136,856.80 · PAYWISE → Paywise"
    l3 = "28/09/2026 · $8,000.00 · EFECTIVO"
    assert _notas_de_cobro(ws) == {
        (FILA, COLS_MXN[0]): l1,
        (FILA, COLS_MXN[1]): l2,
        (FILA, COL_STATUS): f"{l1}\n{l2}",
        (FILA + 2, COLS_MXN[0]): l3,
        (FILA + 2, COL_STATUS): l3,
    }


def test_balance_general_reporte_horas_flota() -> None:
    data = render_balance_general_xlsx(_general(_vuelos()))
    ws = _maestra(data, "reporte horas FLOTA")
    assert _notas_de_cobro(ws) == {
        (FILA, COLS_MXN[0]): LINEA_TRANSFER,
        (FILA, COLS_MXN[1]): LINEA_PAYWISE,
        (FILA, COL_STATUS): f"{LINEA_TRANSFER}\n{LINEA_PAYWISE}",
        (FILA + 2, COLS_MXN[0]): LINEA_EFECTIVO,
        (FILA + 2, COL_STATUS): LINEA_EFECTIVO,
    }
    # Ni la 'cobranza' ni otra hoja ganan notas: solo la maestra.
    con = _comentarios(data)
    sin = _comentarios(render_balance_general_xlsx(_general(_vuelos(_sin_detalle))))
    assert {hoja for hoja, _ in set(con) - set(sin)} == {"reporte horas FLOTA"}


def test_mas_de_cuatro_cobros_la_cuarta_agrega_una_linea_por_cobro() -> None:
    def cobro(dia: int, monto: float, **extra) -> dict:
        return {"fecha": f"2026-09-{dia:02d}", "monto_mxn": monto, **extra}

    efectivo = {"metodo": "EFECTIVO", "metodo_etiqueta": "Efectivo", "registro": "Itzi",
                "cobrado_con": "Efectivo · Itzi"}  # fmt: skip
    cobros = [
        cobro(1, 1000.0, **efectivo),
        cobro(2, 2000.0, **efectivo),
        cobro(3, 3000.0, **efectivo),
        cobro(4, 4000.0, **efectivo),
        cobro(5, 5000.5),  # sin detalle: su línea va con fecha y monto
        cobro(6, 6000.25, metodo="TRANSFERENCIA", cuenta="HSBC Pesos"),  # API viejo
    ]
    vuelos = [_vuelo("#310 · Cliente Seis", "2026-09-06", cobros, "Cobrado")]
    ws = _maestra(render_balance_avion_xlsx(_individual(vuelos)))
    lineas = [
        "01/09/2026 · $1,000.00 · Efectivo · Itzi",
        "02/09/2026 · $2,000.00 · Efectivo · Itzi",
        "03/09/2026 · $3,000.00 · Efectivo · Itzi",
        "04/09/2026 · $4,000.00 · Efectivo · Itzi",
        "05/09/2026 · $5,000.50",
        "06/09/2026 · $6,000.25 · TRANSFERENCIA → HSBC Pesos",
    ]
    assert [_nota(ws, FILA, c) for c in COLS_MXN[:3]] == lineas[:3]
    # La 4.ª parcialidad: fecha «(+3)», monto agregado y una línea por cobro.
    assert ws.cell(row=FILA, column=COLS_FECHA[3]).value == "06/09/2026 (+3)"
    assert ws.cell(row=FILA, column=COLS_MXN[3]).value == pytest.approx(15000.75)
    assert _nota(ws, FILA, COLS_MXN[3]) == "\n".join(lineas[3:])
    # STATUS: los SEIS cobros.
    assert _nota(ws, FILA, COL_STATUS) == "\n".join(lineas)
    # Cuatro agrupados sin ningún detalle ⇒ la 4.ª no lleva nota.
    sin_detalle = [BalanceAvionCobro(**cobro(d, 10.0)) for d in range(1, 7)]
    assert _notas_parcialidades(sin_detalle) == [None, None, None, None]


# ---------------------------------------------------------------------------
# 3) Rutas: el JSON del API 0.0.60 entra sin 422.
# ---------------------------------------------------------------------------

TOKEN = "secreto-de-prueba"
client = TestClient(app)


def test_rutas_aceptan_el_payload_nuevo(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    vuelos = _vuelos()

    res = client.post(
        "/pdf/balance-avion-xlsx",
        json={
            "matricula": "XA-VGV",
            "periodo_desde": "2026-09-01",
            "periodo_hasta": "2026-09-30",
            "vuelos": vuelos,
            "totales": _totales(vuelos),
        },
        headers=headers,
    )
    assert res.status_code == 200, res.text
    assert _nota(_maestra(res.content), FILA, COL_STATUS) == f"{LINEA_TRANSFER}\n{LINEA_PAYWISE}"

    res = client.post(
        "/pdf/balance-general-xlsx",
        json={
            "periodo_desde": "2026-09-01",
            "periodo_hasta": "2026-09-30",
            "consolidado": {"matricula": "FLOTA", "vuelos": vuelos, "totales": _totales(vuelos)},
        },
        headers=headers,
    )
    assert res.status_code == 200, res.text
    ws = _maestra(res.content, "reporte horas FLOTA")
    assert _nota(ws, FILA, COLS_MXN[1]) == LINEA_PAYWISE
