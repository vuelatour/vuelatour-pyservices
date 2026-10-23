"""Comisión del vendedor como GASTO (28-sep-2026, API 0.0.39, invariante 31).

Pedido del cliente: «¿cómo registro un gasto para que aparezca en la hoja de
otros movimientos? Como pagarle una comisión a Saab… si lo agrego como "Otros
gastos VuelaTour" me lo manda a la hoja de Otros gastos: queda duplicado.»

El API agrega la categoría «Comisión del vendedor» y el gasto real REEMPLAZA a
la PROVISIÓN en su fila de «otros movimientos» (Balance general) y de «Otros
ingresos» (Libro Dinero). Aquí solo se pinta: los conceptos y las notas de
celda llegan hechos. Lo único que cambia en pyservices son 5 leyendas, y SOLO
con las banderas que el API manda cuando hay pagos reales:

- Balance general: `otros_movimientos.hay_pago_vendedor_real` (fila 2 de
  'otros movimientos', nota del RESUMEN y nota de la hoja maestra del general).
- Libro Dinero: `utilidades_comision_vendedor_pagada_mxn` (fila 2 de 'Otros
  ingresos' y la frase nueva de la fila 9 de 'utilidades').

Sin bandera el Excel es IDÉNTICO al de antes (las leyendas de hoy siguen
siendo verdaderas mientras no haya gasto real): aquí se compara contra el
texto viejo LITERAL. El libro INDIVIDUAL nunca cambia.

6-oct-2026 (API 0.0.65, comisiones a cargo del avión): las notas del RESUMEN
y de la hoja maestra hablan ahora de la columna COMISIONES (antes «COMISIÓN
VENDEDOR MXN va vacía a propósito…»); la bandera sigue cambiando SOLO el
fragmento del pago al vendedor.

7-oct-2026 (API 0.0.66): la hoja de vuelos del GENERAL ya no lleva la columna
COMISIONES; su nota dice dónde están y remite a la leyenda de 'otros
movimientos' para el pago al vendedor, así que la bandera ya no la cambia
(sigue cambiando la fila 2 de 'otros movimientos' y la nota del RESUMEN, que
ahora remite al libro individual). La del libro INDIVIDUAL, intacta.

12-oct-2026 (API 0.0.73, invariante 58): «provisión: los gastos que todavía
no se hacen no deben aparecer» (las filas en amarillo de la captura). El API
ya no asigna NINGÚN egreso al vendedor mientras no exista el gasto real
«Comisión del vendedor» ligado al vuelo (ni concepto, ni monto, ni fecha, ni
marca amarilla: la fila queda con su ingreso cobrado y remanente = ingreso) y
lo avisa con la bandera ADITIVA `sin_provision_vendedor` (en
`otros_movimientos` del general; al nivel raíz del payload individual). Con
ella, 6 textos nuevos EXACTOS del contrato más, de la revisión adversaria del
mismo día, la base del pie del bloque VUELATOUR (empresa) sin «o
provisionado» (`_NOTA_EMPRESA_BASE_SIN_PROVISION`); sin ella, los de hoy
byte a byte (un API previo sigue provisionando). El amarillo
(`_MARCA_PROVISION`) no se toca: con el API nuevo ya no llega la marca y no
se pinta nada.

13-oct-2026 (API 0.0.74, invariante 61 del API; decisión 3 del lote del
cliente: «comisión del vendedor SOLO como gasto real»): ya no hay provisión en
NINGÚN libro (tampoco en COMISIONES del avión, el reparto ni el Libro Dinero).
Con la bandera ADITIVA `comision_vendedor_solo_gasto_real` (hoja 'otros
movimientos', raíz de cada libro —individual, `aviones[]` Y consolidado
FLOTA—, Libro Dinero y reparto) cambian SOLO textos (`*_GASTO_REAL`, ninguno
dice «provisi…»); manda sobre `sin_provision_vendedor` y
`hay_pago_vendedor_real`. Sin ella, los de hoy byte a byte.

13-oct-2026 · dos categorías (ajuste del cliente de las 17:10, mismo API
0.0.74 y MISMA bandera: «necesitamos una opción que sea Comisión del vendedor
pero que esa opción sí se cargue al balance del avión, y dejamos el otro que
se va a balance de vuelatour»): el pago al vendedor se captura como
«Comisión del vendedor (VuelaTour)» (lo paga VuelaTour en 'otros
movimientos') o «Comisión del vendedor (avión)» (lo carga el avión en su
columna COMISIONES, en cualquier fecha). Los textos de la bandera se
reescriben EN SU LUGAR (contrato §G del ajuste) y nacen los `*_DOS_CATEGORIAS`
(fila 2, RESUMEN y base del bloque VUELATOUR), que `_fragmento_*` y
`_bloque_empresa` eligen PRIMERO; con solo `sin_provision_vendedor` (API
0.0.73) los `*_SIN_PROVISION` byte a byte. Libro Dinero: filas 2, 6 y 9 y la
llave ADITIVA `utilidades_comision_vendedor_avion_mxn` (solo > 0).
"""

import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import app
from app.schemas.reparto import RepartoPdfRequest
from app.schemas.reportes import (
    REGLA_COMISIONES_RESPALDO,
    BalanceAvionRequest,
    BalanceAvionTotales,
    BalanceAvionVuelo,
    BalanceGeneralRequest,
    BalanceHojaOtrosMovimientos,
    DineroXlsxRequest,
)
from app.services.balance_avion_xlsx import (
    _COMISIONES_FUERA_DE_LA_HOJA_GASTO_REAL,
    _EMPRESA_AFAC_EN_FILAS,
    _EMPRESA_REGLA_COMISIONES,
    _EMPRESA_REGLA_COMISIONES_CUBRE_AVION,
    _EMPRESA_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION,
    _EMPRESA_REGLA_COMISIONES_GASTO_REAL,
    _LEYENDA_PROVISION,
    _MAESTRA_PAGO_VENDEDOR_GASTO_REAL,
    _MAESTRA_PAGO_VENDEDOR_SIN_PROVISION,
    _NOTA_BLOQUE_EMPRESA,
    _NOTA_CASCADA_COMISIONES_GASTO_REAL,
    _NOTA_COMISIONES_FLOTA_GASTO_REAL,
    _NOTA_COMISIONES_GASTO_REAL,
    _NOTA_EMPRESA_BASE,
    _NOTA_EMPRESA_BASE_DOS_CATEGORIAS,
    _NOTA_EMPRESA_BASE_SIN_PROVISION,
    _NOTA_ENCABEZADO_COMISIONES_GASTO_REAL,
    _NOTA_ENCABEZADO_GANANCIA_FLOTA_GASTO_REAL,
    _NOTA_ENCABEZADO_REMANENTE_GASTO_REAL,
    _NOTA_PARTICIPACION_PAGO_VENDEDOR_AVION,
    _OM_PAGO_VENDEDOR_DOS_CATEGORIAS,
    _OM_PAGO_VENDEDOR_PROVISION,
    _OM_PAGO_VENDEDOR_REAL,
    _OM_PAGO_VENDEDOR_SIN_PROVISION,
    _OM_REGLA_COMISIONES_CUBRE_AVION,
    _OM_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION,
    _OM_REGLA_COMISIONES_GASTO_REAL,
    _RESUMEN_COMISIONES_GASTO_REAL,
    _RESUMEN_PAGO_VENDEDOR_DOS_CATEGORIAS,
    _RESUMEN_PAGO_VENDEDOR_PROVISION,
    _RESUMEN_PAGO_VENDEDOR_REAL,
    _RESUMEN_PAGO_VENDEDOR_SIN_PROVISION,
    _RESUMEN_REGLA_CUBRE_AVION,
    _RESUMEN_REGLA_CUBRE_AVION_SIN_PROVISION,
    _RESUMEN_REGLA_GASTO_REAL,
    AMBER,
    _fragmento_om_pago_vendedor,
    _fragmento_resumen_pago_vendedor,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from app.services.dinero_xlsx import (
    _LEYENDA_PAGO_VENDEDOR_GASTO_REAL,
    _UTILIDADES_FILA6_GASTO_REAL,
    _UTILIDADES_NOTA_COMISION_AVION,
    _UTILIDADES_OTROS_INGRESOS_GASTO_REAL,
    render_dinero_xlsx,
)
from tests._evaluador_formulas import verificar_libro

# ---------------------------------------------------------------------------
# Textos de HOY (copiados literal del código anterior al 28-sep-2026).
# ---------------------------------------------------------------------------

OM_FILA2_VIEJA = (
    "Ingreso de VuelaTour (no del avión): TUAs, extras, viáticos de "
    "pernocta y comisión del vendedor cobrados al cliente (con su IVA) vs "
    "lo pagado — el pago de la comisión al vendedor va apareado en la "
    "misma fila como PROVISIÓN a la fecha del vuelo mientras no exista el "
    "gasto real (lo dice la nota de la celda). UNA FILA POR "
    "VUELO: los ingresos van sumados en una celda y los egresos en otra; "
    "el desglose concepto por concepto está en el COMENTARIO de cada "
    "celda (pasa el cursor sobre el triángulo rojo). Una fila puede traer "
    "solo ingreso o solo egreso. Todo en MXN. Incluye todos "
    "los estados del periodo (igual que la hoja maestra); los cancelados "
    "se marcan (clave · CANCELADO en rojo) y los demás estados no "
    "normales llevan su estado junto a la clave. Los MOVIMIENTOS SIN "
    "AVIÓN / SIN VUELO de abajo son 'gas sin avión', 'tuas sin vuelo' y "
    "los OTROS INGRESOS registrados en Ingresos (clave ING-n: intereses, "
    "reembolsos, ventas de activos…); los anticipos y las aportaciones NO "
    "aparecen: no son resultado (29-ago: los gastos de empresa viven en "
    "la hoja 'otros gastos' — antes llamada 'gastos VuelaTour')."
)
# Desde el 6-oct-2026 (API 0.0.65) la base de estas dos notas es la de la
# columna COMISIONES; el fragmento final (pago apareado) es el de siempre.
# 7-oct-2026 (API 0.0.66): la del RESUMEN remite a la columna COMISIONES del
# libro individual (la hoja de vuelos del general ya no la lleva).
RESUMEN_NOTA_VIEJA = (
    "COMISIONES = las de la columna COMISIONES del libro individual de cada "
    "avión (regla de septiembre 2026, por fecha del vuelo: su parte de la "
    "comisión bancaria de los cobros y la comisión del vendedor como "
    "PROVISIÓN; en vuelos anteriores no hay); en este libro no van en la hoja "
    "de vuelos (ahí la GANANCIA va antes de comisiones): restan aquí y en la "
    "cascada de cada avión de la hoja 'balance'. Lo cobrado al cliente por la "
    "comisión del vendedor sigue siendo INGRESO de VuelaTour; ese ingreso y el "
    "pago al vendedor viven en 'otros movimientos' "
    "(ingreso cobrado y pago apareado como PROVISIÓN a la fecha del vuelo)."
)
MAESTRA_NOTA_VIEJA = (
    "COMISIONES MXN (regla de septiembre 2026, por fecha del vuelo) = lo que "
    "absorbe el avión: su parte de la comisión bancaria de los cobros y la "
    "comisión del vendedor como PROVISIÓN (lo cobrado al cliente por ese "
    "concepto, comisión + IVA); la nota de cada celda dice cuáles son y en los "
    "vuelos anteriores va vacía. Lo cobrado al cliente por la comisión del "
    "vendedor sigue siendo ingreso de VuelaTour y su pago al vendedor sale de "
    "VuelaTour; ambos viven en 'otros movimientos' (el pago apareado como "
    "PROVISIÓN a la fecha del vuelo). GANANCIA de la fila = REMANENTE (VENTA − "
    "COSTO TOTAL) − COMISIONES."
)
# Nota de la hoja de vuelos del GENERAL desde el 7-oct-2026: sin la columna
# COMISIONES y sin el fragmento del pago apareado (remite a la leyenda de
# 'otros movimientos'): con y sin bandera, la misma.
MAESTRA_GENERAL_NOTA = (
    "COMISIONES (regla de septiembre 2026, por fecha del vuelo): las que "
    "absorbe cada avión —su parte de la comisión bancaria de sus cobros y la "
    "comisión del vendedor como PROVISIÓN— no van en esta hoja: están en la "
    "columna COMISIONES del libro individual de cada avión (Reportes › Balance "
    "por avión) y ya van restadas en la cascada de cada avión de la hoja "
    "'balance'. Lo cobrado al cliente por la comisión del vendedor es ingreso "
    "de VuelaTour: ese ingreso, el pago al vendedor y la parte de VuelaTour de "
    "la comisión bancaria viven en 'otros movimientos' (su leyenda explica "
    "cada línea). GANANCIA de la fila = REMANENTE (VENTA − COSTO TOTAL), antes "
    "de comisiones."
)
# Prefijos con que se encuentran en la columna A.
RESUMEN_PREFIJO = "COMISIONES = las de la columna COMISIONES"
MAESTRA_PREFIJO = "COMISIONES MXN (regla de septiembre 2026"
# La del individual CON la bandera del API 0.0.74 (dos categorías): la regla
# va dentro, solo para la parte bancaria.
MAESTRA_PREFIJO_GASTO_REAL = "COMISIONES MXN = lo que absorbe el avión"
MAESTRA_GENERAL_PREFIJO = "COMISIONES (regla de septiembre 2026"
DINERO_FILA2_VIEJA = (
    "Ingreso de VuelaTour (no del avión): TUAs, extras, pernocta y "
    "comisión del vendedor cobrados (con su IVA) vs lo pagado. El pago de "
    "la comisión al vendedor va apareado en su fila como PROVISIÓN a la "
    "fecha del vuelo mientras no exista el gasto real (regla 28-ago-2026). "
    "Al final, los otros ingresos que no son de un vuelo (clave ING-n, "
    "registrados en Ingresos). Los anticipos de clientes no están aquí: "
    "cuentan como cobro del vuelo al aplicarse."
)
UTILIDADES_FILA9_BASE = (
    "OTROS INGRESOS = ingreso de VuelaTour (TUAs/extras/pernocta/"
    "comisión del vendedor + su IVA), no del avión (regla 28-ago-2026), "
    "NETO de la provisión del pago al vendedor (comisión + su IVA), más "
    "los ingresos sin vuelo registrados en Ingresos (ING-n), netos de su "
    "comisión bancaria: ver hoja 'Otros ingresos'."
)

# ---------------------------------------------------------------------------
# Textos NUEVOS (contrato del 28-sep-2026, sección 5) — fragmento viejo ⇒ nuevo.
# ---------------------------------------------------------------------------

OM_FRAG_VIEJO = (
    "el pago de la comisión al vendedor va apareado en la misma fila como "
    "PROVISIÓN a la fecha del vuelo mientras no exista el gasto real (lo dice "
    "la nota de la celda)."
)
OM_FRAG_NUEVO = (
    "el pago de la comisión al vendedor va apareado en la misma fila: el GASTO "
    "REAL cuando se captura en Gastos con la categoría «Comisión del vendedor» "
    "ligado al vuelo (con «faltan $…» o «excede $…» si no cuadra con lo "
    "cobrado) o, mientras no se capture, una PROVISIÓN por el mismo monto a la "
    "fecha del vuelo (lo dice la nota de la celda). Ese pago NO se captura como "
    "«Otros gastos VuelaTour»: quedaría duplicado en la hoja 'otros gastos'."
)
RESUMEN_FRAG_VIEJO = "(ingreso cobrado y pago apareado como PROVISIÓN a la fecha del vuelo)."
RESUMEN_FRAG_NUEVO = (
    "(ingreso cobrado y pago apareado: el gasto real «Comisión del vendedor» o, "
    "mientras no se capture, la PROVISIÓN a la fecha del vuelo)."
)
# Fragmento de la nota de la hoja maestra del libro INDIVIDUAL (la única que
# lo lleva desde el 7-oct-2026; antes, con la bandera, también la del general).
MAESTRA_FRAG_VIEJO = "(el pago apareado como PROVISIÓN a la fecha del vuelo)."
DINERO_FRAG_VIEJO = (
    "El pago de la comisión al vendedor va apareado en su fila como PROVISIÓN a "
    "la fecha del vuelo mientras no exista el gasto real (regla 28-ago-2026)."
)
DINERO_FRAG_NUEVO = (
    "El pago de la comisión al vendedor va apareado en su fila: el gasto real "
    "(categoría «Comisión del vendedor» ligada al vuelo) o, mientras no se "
    "capture, una PROVISIÓN por el mismo monto a la fecha del vuelo (regla "
    "28-ago-2026)."
)


def _cambia(viejo: str, frag_viejo: str, frag_nuevo: str) -> str:
    """El texto nuevo = el viejo con SOLO el fragmento sustituido."""
    assert viejo.count(frag_viejo) == 1
    return viejo.replace(frag_viejo, frag_nuevo)


# ---------------------------------------------------------------------------
# Payloads (como los arma el API: conceptos y notas YA hechos).
# ---------------------------------------------------------------------------

CONCEPTO_PARCIAL = "pago comisión vendedor (Alex Saab) · gasto real · parcial: faltan $400.00 MXN"
NOTA_PAGO_REAL = (
    "GASTO REAL: pago al vendedor capturado como «Comisión del vendedor» ligado "
    "a este vuelo; reemplaza la provisión. En la hoja utilidades ya está "
    "descontado de \"otros ingresos\"."
)

_FILA_PROVISION = {
    "clave": "N4142R-317",
    "avion_color": "#FF0000",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-09-20",
    "concepto_egreso": (
        "pago comisión vendedor (Alex Saab) · PROVISIÓN (mismo monto que lo "
        "cobrado: comisión + IVA; sin gasto real capturado)"
    ),
    "egreso_mxn": 2030.0,
    "fecha_egreso": "2026-09-20",
    "concepto_ingreso": "comisión vendedor (Alex Saab)",
    "ingreso_mxn": 2030.0,
    "fecha_ingreso": "2026-09-20",
    "remanente_mxn": 0.0,
}
_FILA_PAGO_REAL = {
    "clave": "N4142R-260",
    "avion_color": "#FF0000",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-09-10",
    "concepto_egreso": CONCEPTO_PARCIAL,
    "egreso_mxn": 1200.0,
    "fecha_egreso": "2026-09-25",
    "concepto_ingreso": "comisión vendedor (Alex Saab)",
    "ingreso_mxn": 1600.0,
    "fecha_ingreso": "2026-09-10",
    "remanente_mxn": 400.0,
    "nota_egreso": NOTA_PAGO_REAL,
}


def _general(bandera=..., filas=None, sin_provision=...) -> BalanceGeneralRequest:
    """`bandera=...` ⇒ la clave NO viaja (así la manda el API sin pagos);
    `sin_provision` = `otros_movimientos.sin_provision_vendedor` (API 0.0.73,
    12-oct-2026; `...` = no viaja, como en un API previo)."""
    om: dict = {"filas": filas if filas is not None else [_FILA_PROVISION]}
    if bandera is not ...:
        om["hay_pago_vendedor_real"] = bandera
    if sin_provision is not ...:
        om["sin_provision_vendedor"] = sin_provision
    return BalanceGeneralRequest(
        periodo_desde="2026-09-01",
        periodo_hasta="2026-09-30",
        resumen=[{"matricula": "N4142R", "color": "#FF0000", "vuelos": 2}],
        consolidado=BalanceAvionRequest(
            matricula="FLOTA",
            periodo_desde="2026-09-01",
            periodo_hasta="2026-09-30",
            otros_movimientos=om,
        ),
    )


def _dinero(pagada=..., provision: float | None = 2030.0,
            filas=None) -> DineroXlsxRequest:
    datos: dict = {
        "periodo_desde": "2026-09-01",
        "periodo_hasta": "2026-09-30",
        "otros_ingresos": filas if filas is not None else [_FILA_PROVISION],
        "utilidades_otros_ingresos_mxn": 5000.0,
        "utilidades_comision_vendedor_provisionada_mxn": provision,
    }
    if pagada is not ...:
        datos["utilidades_comision_vendedor_pagada_mxn"] = pagada
    return DineroXlsxRequest(**datos)


def _libro(xlsx: bytes):
    return load_workbook(BytesIO(xlsx), data_only=True)


def _celda_que_empieza(ws, prefijo: str) -> str:
    """Valor de la ÚNICA celda de la columna A que empieza con `prefijo`."""
    hallados = [
        c.value
        for c in ws["A"]
        if isinstance(c.value, str) and c.value.startswith(prefijo)
    ]
    assert len(hallados) == 1, hallados
    return hallados[0]


def _leyendas_general(xlsx: bytes) -> tuple[str, str, str]:
    wb = _libro(xlsx)
    om = wb["otros movimientos"]["A2"].value
    resumen = _celda_que_empieza(wb["RESUMEN flota"], RESUMEN_PREFIJO)
    maestra = _celda_que_empieza(wb["reporte horas FLOTA"], MAESTRA_GENERAL_PREFIJO)
    return om, resumen, maestra


def _leyendas_dinero(xlsx: bytes) -> tuple[str, str]:
    wb = _libro(xlsx)
    return wb["Otros ingresos"]["A2"].value, wb["utilidades"]["A9"].value


# ---------------------------------------------------------------------------
# Esquemas: campos ADITIVOS (skew tolerante en ambos sentidos).
# ---------------------------------------------------------------------------


def test_esquemas_aceptan_los_campos_nuevos_con_default_none() -> None:
    assert BalanceHojaOtrosMovimientos().hay_pago_vendedor_real is None
    assert BalanceHojaOtrosMovimientos(hay_pago_vendedor_real=True).hay_pago_vendedor_real is True
    assert DineroXlsxRequest().utilidades_comision_vendedor_pagada_mxn is None
    assert (
        DineroXlsxRequest(utilidades_comision_vendedor_pagada_mxn=1600).utilidades_comision_vendedor_pagada_mxn
        == 1600.0
    )


# ---------------------------------------------------------------------------
# Balance general: SIN bandera ⇒ las 3 leyendas son las de HOY (literal).
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("bandera", [..., None, False], ids=["ausente", "null", "false"])
def test_general_sin_bandera_leyendas_identicas_a_hoy(bandera) -> None:
    om, resumen, maestra = _leyendas_general(render_balance_general_xlsx(_general(bandera)))
    assert om == OM_FILA2_VIEJA
    assert resumen == RESUMEN_NOTA_VIEJA
    assert maestra == MAESTRA_GENERAL_NOTA


def test_general_sin_otros_movimientos_resumen_y_maestra_de_hoy() -> None:
    """API viejo sin la pestaña: nada que marcar, textos de siempre."""
    req = _general()
    req.consolidado.otros_movimientos = None
    wb = _libro(render_balance_general_xlsx(req))
    assert "otros movimientos" not in wb.sheetnames
    assert _celda_que_empieza(wb["RESUMEN flota"], RESUMEN_PREFIJO) == RESUMEN_NOTA_VIEJA
    maestra = _celda_que_empieza(wb["reporte horas FLOTA"], MAESTRA_GENERAL_PREFIJO)
    assert maestra == MAESTRA_GENERAL_NOTA


# ---------------------------------------------------------------------------
# Balance general: CON bandera ⇒ los textos nuevos tal cual.
# ---------------------------------------------------------------------------


def test_general_con_bandera_pinta_las_leyendas_nuevas() -> None:
    om, resumen, maestra = _leyendas_general(
        render_balance_general_xlsx(_general(True, filas=[_FILA_PROVISION, _FILA_PAGO_REAL]))
    )
    assert om == _cambia(OM_FILA2_VIEJA, OM_FRAG_VIEJO, OM_FRAG_NUEVO)
    assert resumen == _cambia(RESUMEN_NOTA_VIEJA, RESUMEN_FRAG_VIEJO, RESUMEN_FRAG_NUEVO)
    # 7-oct-2026: la nota de la hoja de vuelos del general ya no habla del
    # pago al vendedor (remite a la leyenda de arriba): no cambia con la bandera.
    assert maestra == MAESTRA_GENERAL_NOTA
    # El camino que confundió al cliente queda escrito en la hoja.
    assert "«Comisión del vendedor»" in om
    assert "quedaría duplicado en la hoja 'otros gastos'" in om
    assert "PROVISIÓN" in om  # la provisión sigue existiendo sin gasto real


def test_individual_nunca_cambia_aunque_traiga_la_bandera() -> None:
    """El libro INDIVIDUAL no pinta 'otros movimientos'; su nota de la hoja
    maestra es la de siempre pase lo que pase."""
    req = BalanceAvionRequest(
        matricula="N4142R",
        periodo_desde="2026-09-01",
        periodo_hasta="2026-09-30",
        otros_movimientos={"filas": [_FILA_PAGO_REAL], "hay_pago_vendedor_real": True},
    )
    wb = _libro(render_balance_avion_xlsx(req))
    assert "otros movimientos" not in wb.sheetnames
    maestra = _celda_que_empieza(wb["reporte horas N4142R"], MAESTRA_PREFIJO)
    assert maestra == MAESTRA_NOTA_VIEJA
    assert maestra.count(MAESTRA_FRAG_VIEJO) == 1


def test_otros_movimientos_pinta_el_concepto_del_gasto_real_tal_cual() -> None:
    wb = _libro(render_balance_general_xlsx(_general(True, filas=[_FILA_PAGO_REAL])))
    ws = wb["otros movimientos"]
    assert ws["A4"].value == "N4142R-260"
    assert ws["F4"].value == CONCEPTO_PARCIAL
    assert ws["G4"].value == 1200.0
    assert ws["H4"].value == "25/09/2026"
    assert ws["J4"].value == 1600.0
    assert ws["L4"].value == 400.0
    assert ws["G4"].comment.text == NOTA_PAGO_REAL


# ---------------------------------------------------------------------------
# Libro Dinero.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("provision", [2030.0, 0.0, None])
def test_dinero_sin_campo_leyendas_identicas_a_hoy(provision) -> None:
    fila2, fila9 = _leyendas_dinero(render_dinero_xlsx(_dinero(provision=provision)))
    assert fila2 == DINERO_FILA2_VIEJA
    nota = " Provisión restada en este periodo: $2,030.00 MXN." if provision else ""
    assert fila9 == UTILIDADES_FILA9_BASE + nota


def test_dinero_con_pagado_en_cero_fila9_de_hoy_y_leyenda_nueva() -> None:
    """El campo presente (aun en 0) dice que hubo gasto real ⇒ la leyenda de
    la hoja cambia; la frase de la fila 9 solo aparece con monto > 0."""
    fila2, fila9 = _leyendas_dinero(render_dinero_xlsx(_dinero(pagada=0.0)))
    assert fila2 == _cambia(DINERO_FILA2_VIEJA, DINERO_FRAG_VIEJO, DINERO_FRAG_NUEVO)
    assert fila9 == UTILIDADES_FILA9_BASE + " Provisión restada en este periodo: $2,030.00 MXN."


def test_dinero_con_pagado_agrega_la_frase_despues_de_la_provision() -> None:
    fila2, fila9 = _leyendas_dinero(
        render_dinero_xlsx(_dinero(pagada=1600.0, filas=[_FILA_PROVISION, _FILA_PAGO_REAL]))
    )
    assert fila2 == _cambia(DINERO_FILA2_VIEJA, DINERO_FRAG_VIEJO, DINERO_FRAG_NUEVO)
    assert fila9 == (
        UTILIDADES_FILA9_BASE
        + " Provisión restada en este periodo: $2,030.00 MXN."
        + " Pagado al vendedor con gasto real en este periodo: $1,600.00 MXN."
    )


def test_dinero_solo_pagos_reales_sin_provision_viva() -> None:
    """Todos los vuelos del periodo con gasto real: provisión 0 ⇒ solo la
    frase de lo pagado (dinero con 2 decimales y moneda)."""
    _, fila9 = _leyendas_dinero(
        render_dinero_xlsx(_dinero(pagada=12345.5, provision=0.0, filas=[_FILA_PAGO_REAL]))
    )
    assert fila9 == (
        UTILIDADES_FILA9_BASE
        + " Pagado al vendedor con gasto real en este periodo: $12,345.50 MXN."
    )


def test_dinero_pinta_el_concepto_y_la_nota_del_gasto_real_tal_cual() -> None:
    wb = _libro(render_dinero_xlsx(_dinero(pagada=1200.0, filas=[_FILA_PAGO_REAL])))
    ws = wb["Otros ingresos"]
    assert ws["F4"].value == CONCEPTO_PARCIAL
    assert ws["G4"].value == 1200.0
    assert ws["L4"].value == 400.0
    assert ws["G4"].comment.text == NOTA_PAGO_REAL


# ---------------------------------------------------------------------------
# Rutas: el payload del API 0.0.39 entra sin 422 y respeta las banderas.
# ---------------------------------------------------------------------------

TOKEN = "secreto-de-prueba"
client = TestClient(app)


def test_rutas_aceptan_el_payload_nuevo(monkeypatch) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    headers = {"X-Internal-Token": TOKEN}

    res = client.post(
        "/pdf/balance-general-xlsx",
        json=_general(True, filas=[_FILA_PAGO_REAL]).model_dump(mode="json"),
        headers=headers,
    )
    assert res.status_code == 200, res.text
    om, _, _ = _leyendas_general(res.content)
    assert om == _cambia(OM_FILA2_VIEJA, OM_FRAG_VIEJO, OM_FRAG_NUEVO)

    res = client.post(
        "/pdf/dinero-xlsx",
        json=_dinero(pagada=1200.0, filas=[_FILA_PAGO_REAL]).model_dump(mode="json"),
        headers=headers,
    )
    assert res.status_code == 200, res.text
    _, fila9 = _leyendas_dinero(res.content)
    assert fila9.endswith(" Pagado al vendedor con gasto real en este periodo: $1,200.00 MXN.")


# ===========================================================================
# 12-oct-2026 (API 0.0.73, invariante 58): «provisión: los gastos que todavía
# no se hacen no deben aparecer». Con la bandera ADITIVA
# `sin_provision_vendedor` los textos EXACTOS del contrato; sin ella, los de
# hoy byte a byte. El amarillo no se toca.
# ===========================================================================

# Textos del CONTRATO, literales (partidos en otros puntos que en el código:
# lo que se compara es la cadena completa).
OM_FRAG_SIN_PROVISION = (
    "el pago de la comisión al vendedor va apareado en la misma fila SOLO cuando "
    "se captura en Gastos el gasto real «Comisión del vendedor» ligado al vuelo "
    "(con «faltan $…» o «excede $…» si no cuadra con lo cobrado); mientras no se "
    "capture, la fila no lleva egreso ni provisión (pedido del 12-oct-2026: los "
    "gastos que todavía no se hacen no aparecen). Ese pago NO se captura como "
    "«Otros gastos VuelaTour»: quedaría duplicado en la hoja 'otros gastos'."
)
RESUMEN_FRAG_SIN_PROVISION = (
    "(ingreso cobrado; el pago al vendedor se aparea solo cuando se captura el "
    "gasto real «Comisión del vendedor»: sin él no se provisiona nada en esa hoja)."
)
MAESTRA_FRAG_SIN_PROVISION = (
    "(el pago se aparea ahí cuando se captura el gasto real «Comisión del vendedor»)."
)
OM_REGLA_CUBRE_SIN_PROVISION = (
    " Con la {regla} (por fecha del vuelo), en los vuelos COMPLETADOS la comisión "
    "del vendedor la cubre el avión con su PROVISIÓN en su columna COMISIONES; "
    "aquí el pago al vendedor aparece solo cuando se captura el gasto real y resta "
    "únicamente lo que ese pago exceda la provisión del avión; la comisión "
    "bancaria de aquí es solo la parte de VuelaTour. Lo que absorbe cada avión (la "
    "provisión y su parte de la comisión bancaria) resta en la columna COMISIONES "
    "de su libro individual y en la cascada de la hoja 'balance'."
)
EMPRESA_REGLA_CUBRE_SIN_PROVISION = (
    " Con la {regla} (por fecha del vuelo), en los vuelos completados el pago al "
    "vendedor lo cubre cada avión con su provisión (los egresos propios solo traen "
    "lo que un pago real capturado la exceda) y, de la comisión bancaria, los "
    "egresos solo traen la parte de VuelaTour."
)
RESUMEN_REGLA_CUBRE_SIN_PROVISION = (
    " Con la {regla} (por fecha del vuelo), en los vuelos completados ese pago lo "
    "cubre el avión con su provisión: en 'otros movimientos' solo queda lo que un "
    "pago real capturado la exceda."
)
# Base del pie del bloque VUELATOUR (empresa) con la bandera (revisión
# adversaria del 12-oct-2026): la de hoy dice «lo pagado o provisionado» y con
# el API 0.0.73 los egresos propios ya no traen provisión alguna.
EMPRESA_BASE_SIN_PROVISION = (
    "Ingresos y egresos propios = hoja 'otros movimientos': lo cotizado de los "
    "vuelos del periodo (no cancelados) más los otros ingresos registrados, y lo "
    "pagado (el pago al vendedor solo cuando se captura el gasto real «Comisión "
    "del vendedor») — no es lo cobrado: el RESULTADO no es dinero en caja."
)
REGLA = REGLA_COMISIONES_RESPALDO

# Filas como las arma el API 0.0.73. Sin gasto real NO hay egreso (ni
# concepto, ni monto, ni fecha, ni marca); con gasto real, igual que hoy.
_FILA_V73_VIGENCIA_SIN_GASTO = {  # la captura del cliente: $12,148.35 cobrados
    "clave": "#412 · Cliente Doce",
    "avion_color": "#2563EB",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-10-03",
    "concepto_egreso": None,
    "egreso_mxn": None,
    "fecha_egreso": None,
    "concepto_ingreso": "comisión vendedor (Pablo Canales)",
    "ingreso_mxn": 12148.35,
    "fecha_ingreso": "2026-10-03",
    "remanente_mxn": 12148.35,
}
_FILA_V73_CONFIRMADO_SIN_GASTO = {  # vuelo no realizado sin gasto: tampoco egreso
    "clave": "#415 · Cliente Quince",
    "avion_color": "#2563EB",
    "estado": "CONFIRMADO",
    "fecha_vuelo": "2026-10-20",
    "concepto_ingreso": "comisión vendedor (Alex Saab)",
    "ingreso_mxn": 2030.0,
    "fecha_ingreso": "2026-10-20",
    "remanente_mxn": 2030.0,
}
_FILA_V73_CUBIERTO_PARCIAL = {  # gasto real MENOR que la provisión: $0.00 sin marca
    "clave": "#409 · Cliente Nueve",
    "avion_color": "#16A34A",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-10-01",
    "concepto_egreso": (
        "pago comisión vendedor (Pablo Canales) · cubierto por el avión XA-TST "
        "(provisión $2,360.00 en su balance) · gasto real $2,000.00 · parcial: "
        "faltan $360.00 MXN"
    ),
    "egreso_mxn": 0.0,
    "fecha_egreso": "2026-10-09",
    "concepto_ingreso": "comisión vendedor (Pablo Canales)",
    "ingreso_mxn": 2360.0,
    "fecha_ingreso": "2026-10-01",
    "remanente_mxn": 2360.0,
}
_FILA_V73_CUBIERTO_EXCEDE = {  # gasto real MAYOR que la provisión: solo el exceso
    "clave": "#407 · Cliente Cinco",
    "avion_color": "#16A34A",
    "estado": "COMPLETADO",
    "fecha_vuelo": "2026-09-18",
    "concepto_egreso": (
        "pago comisión vendedor (Alex Saab) · cubierto por el avión XB-DOS "
        "(provisión $1,180.00 en su balance) · gasto real $1,500.00 · excede "
        "$320.00 MXN"
    ),
    "egreso_mxn": 320.0,
    "fecha_egreso": "2026-09-25",
    "concepto_ingreso": "comisión vendedor (Alex Saab)",
    "ingreso_mxn": 1180.0,
    "fecha_ingreso": "2026-09-18",
    "remanente_mxn": 860.0,
}
# Los tres sabores con gasto real que siguen EXACTAMENTE como hoy: gasto real
# pre-vigencia (`_FILA_PAGO_REAL`, faltan $400), cubierto con exceso y
# cubierto con parcial (arriba).
_FILAS_V73 = [
    _FILA_V73_VIGENCIA_SIN_GASTO,
    _FILA_V73_CONFIRMADO_SIN_GASTO,
    _FILA_V73_CUBIERTO_PARCIAL,
    _FILA_V73_CUBIERTO_EXCEDE,
    _FILA_PAGO_REAL,
]


def _tg():
    """`tests.test_balance_comisiones_general` (fixtures del API 0.0.66: vuelos
    de la regla, `comisiones_usd` en los aviones, bloque VUELATOUR), importado
    PEREZOSO: `test_balance_formulas` importa este módulo y aquel importa
    `test_balance_formulas` — a nivel de módulo sería un import circular al
    recolectar la suite completa."""
    from tests import test_balance_comisiones_general as tg

    return tg


def _payload_regla_v73(variante: str = "mensual", *, sin_provision=True, filas=None) -> dict:
    """El general del API 0.0.66 (vuelos de la regla, `comisiones_usd` en los
    aviones, bloque VUELATOUR) con 'otros movimientos' como lo manda el
    0.0.73 (`sin_provision=...` ⇒ la clave no viaja)."""
    p = _tg()._payload_om_0066(variante)
    om = p["consolidado"]["otros_movimientos"]
    om["filas"] = filas if filas is not None else list(_FILAS_V73)
    if sin_provision is not ...:
        om["sin_provision_vendedor"] = sin_provision
    return p


def _rellenos_amber(ws) -> list[str]:
    """Coordenadas de las celdas de DATOS con relleno AMBER (la celda de la
    leyenda «Amarillo = provisión…» también va en ámbar y se verifica aparte
    por su texto)."""
    return [
        c.coordinate
        for fila in ws.iter_rows()
        for c in fila
        if c.value != _LEYENDA_PROVISION
        and c.fill is not None
        and c.fill.fill_type
        and c.fill.fgColor.rgb == "00" + AMBER
    ]


def _textos(ws) -> list[str]:
    return [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]


def _suma(filas: list[dict], llave: str) -> float:
    return round(sum(f.get(llave) or 0.0 for f in filas), 2)


def test_esquemas_aceptan_sin_provision_vendedor_con_default_none() -> None:
    assert BalanceHojaOtrosMovimientos().sin_provision_vendedor is None
    assert BalanceHojaOtrosMovimientos(sin_provision_vendedor=True).sin_provision_vendedor is True
    assert BalanceAvionRequest().sin_provision_vendedor is None
    assert BalanceAvionRequest(sin_provision_vendedor=True).sin_provision_vendedor is True
    # El payload del API 0.0.73 entra tal cual: la bandera viaja en la hoja
    # 'otros movimientos' del general y al nivel raíz de cada avión; el
    # consolidado FLOTA (armado a mano en el API) NO la lleva — aquí nadie la
    # lee de ahí (`_sin_provision_vendedor(cons.otros_movimientos)`).
    req = BalanceGeneralRequest.model_validate(
        {
            "consolidado": {
                "otros_movimientos": {"filas": [], "sin_provision_vendedor": True},
            },
            "aviones": [{"matricula": "XA-TST", "sin_provision_vendedor": True}],
        }
    )
    assert req.consolidado.otros_movimientos.sin_provision_vendedor is True
    assert req.consolidado.sin_provision_vendedor is None
    assert req.aviones[0].sin_provision_vendedor is True


def test_textos_sin_provision_son_los_del_contrato() -> None:
    assert _OM_PAGO_VENDEDOR_SIN_PROVISION == OM_FRAG_SIN_PROVISION
    assert _RESUMEN_PAGO_VENDEDOR_SIN_PROVISION == RESUMEN_FRAG_SIN_PROVISION
    assert _MAESTRA_PAGO_VENDEDOR_SIN_PROVISION == MAESTRA_FRAG_SIN_PROVISION
    assert _OM_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION == OM_REGLA_CUBRE_SIN_PROVISION
    assert _EMPRESA_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION == EMPRESA_REGLA_CUBRE_SIN_PROVISION
    assert _RESUMEN_REGLA_CUBRE_AVION_SIN_PROVISION == RESUMEN_REGLA_CUBRE_SIN_PROVISION
    assert _NOTA_EMPRESA_BASE_SIN_PROVISION == EMPRESA_BASE_SIN_PROVISION
    # Ninguno promete una PROVISIÓN apareada en 'otros movimientos'.
    for t in (OM_FRAG_SIN_PROVISION, RESUMEN_FRAG_SIN_PROVISION, MAESTRA_FRAG_SIN_PROVISION):
        assert "PROVISIÓN" not in t
    assert "provisionado" not in EMPRESA_BASE_SIN_PROVISION
    assert "o provisionado" in _NOTA_EMPRESA_BASE  # la de hoy, intacta


# --- Fila 2 de 'otros movimientos' y nota del RESUMEN (sin vuelos de la regla)


@pytest.mark.parametrize("hay_pago_real", [..., True, False], ids=["sin", "true", "false"])
def test_general_con_bandera_fila2_y_resumen_nuevos(hay_pago_real) -> None:
    """Con `sin_provision_vendedor` UN solo texto para los dos casos (con o sin
    pago real en el periodo): ya no hay PROVISIÓN que explicar."""
    om, resumen, maestra = _leyendas_general(
        render_balance_general_xlsx(
            _general(hay_pago_real, filas=list(_FILAS_V73), sin_provision=True)
        )
    )
    assert om == _cambia(OM_FILA2_VIEJA, OM_FRAG_VIEJO, OM_FRAG_SIN_PROVISION)
    assert resumen == _cambia(RESUMEN_NOTA_VIEJA, RESUMEN_FRAG_VIEJO, RESUMEN_FRAG_SIN_PROVISION)
    assert maestra == MAESTRA_GENERAL_NOTA
    assert "PROVISIÓN" not in om
    assert "los gastos que todavía no se hacen no aparecen" in om


@pytest.mark.parametrize("sin_provision", [..., None, False], ids=["ausente", "null", "false"])
@pytest.mark.parametrize("hay_pago_real", [..., True], ids=["sin pago real", "con pago real"])
def test_general_sin_bandera_fila2_y_resumen_de_hoy(sin_provision, hay_pago_real) -> None:
    """Bandera apagada (API previo que sigue provisionando): los textos de hoy
    byte a byte, elegidos por `hay_pago_vendedor_real` como siempre."""
    om, resumen, maestra = _leyendas_general(
        render_balance_general_xlsx(_general(hay_pago_real, sin_provision=sin_provision))
    )
    if hay_pago_real is True:
        assert om == _cambia(OM_FILA2_VIEJA, OM_FRAG_VIEJO, OM_FRAG_NUEVO)
        assert resumen == _cambia(RESUMEN_NOTA_VIEJA, RESUMEN_FRAG_VIEJO, RESUMEN_FRAG_NUEVO)
    else:
        assert om == OM_FILA2_VIEJA
        assert resumen == RESUMEN_NOTA_VIEJA
    assert maestra == MAESTRA_GENERAL_NOTA
    assert OM_FRAG_SIN_PROVISION not in om


# --- Con vuelos de la regla (API 0.0.66+): «cubierto por el avión» ------------


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_general_con_bandera_y_regla_render_sin_provision(variante) -> None:
    """Payload del API 0.0.73 (sin marca, con la bandera): fila 2, RESUMEN y
    bloque VUELATOUR con los textos nuevos; ninguna celda amarilla ni la
    leyenda del amarillo; la fila sin gasto real SIN egreso y remanente =
    ingreso; las filas con gasto real igual que hoy; TOTALES = Σ filas."""
    p = _payload_regla_v73(variante)
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert verificar_libro(data)
    wb, wv = load_workbook(BytesIO(data)), _libro(data)
    om, omv = wb["otros movimientos"], wv["otros movimientos"]

    # Fila 2: fragmento nuevo + regla «cubierto por el avión» nueva.
    assert om["A2"].value == _cambia(
        OM_FILA2_VIEJA,
        OM_FRAG_VIEJO,
        OM_FRAG_SIN_PROVISION + OM_REGLA_CUBRE_SIN_PROVISION.format(regla=REGLA),
    )
    assert _OM_REGLA_COMISIONES_CUBRE_AVION.format(regla=REGLA) not in om["A2"].value
    # RESUMEN y bloque VUELATOUR (empresa).
    resumen = _celda_que_empieza(wb["RESUMEN flota"], RESUMEN_PREFIJO)
    assert resumen == (
        _cambia(RESUMEN_NOTA_VIEJA, RESUMEN_FRAG_VIEJO, RESUMEN_FRAG_SIN_PROVISION)
        + RESUMEN_REGLA_CUBRE_SIN_PROVISION.format(regla=REGLA)
    )
    assert _RESUMEN_REGLA_CUBRE_AVION.format(regla=REGLA) not in resumen
    pie = _textos(wb["balance"])
    base = EMPRESA_BASE_SIN_PROVISION + EMPRESA_REGLA_CUBRE_SIN_PROVISION.format(regla=REGLA)
    assert base in pie  # una celda: la base nueva + la regla nueva
    vieja = _NOTA_EMPRESA_BASE + _EMPRESA_REGLA_COMISIONES_CUBRE_AVION.format(regla=REGLA)
    assert not any(vieja in x for x in pie)
    # La base de hoy («lo pagado o provisionado») no se pinta en ninguna celda.
    assert not any(_NOTA_EMPRESA_BASE in x or "o provisionado" in x for x in pie)

    # Sin marca ⇒ sin amarillo, sin leyenda y sin «PROVISIÓN (» en la pestaña.
    assert _rellenos_amber(om) == []
    assert _LEYENDA_PROVISION not in _textos(om)
    assert not any("PROVISIÓN (" in t for t in _textos(om))

    # (a) vigencia, COMPLETADO, sin gasto: SIN concepto/egreso/fecha, ingreso
    # igual, remanente = ingreso (fórmula viva: el egreso vacío cuenta 0).
    assert om["A4"].value == "#412 · Cliente Doce"
    assert (om["F4"].value, om["G4"].value, om["H4"].value) == (None, None, None)
    assert om["J4"].value == 12148.35
    assert om["L4"].value == "=ROUND(J4-G4,2)" and omv["L4"].value == 12148.35
    # (b) CONFIRMADO sin gasto: igual, sin egreso.
    assert (om["F5"].value, om["G5"].value, om["H5"].value) == (None, None, None)
    assert omv["L5"].value == 2030.0
    # (c) con gasto real, idéntico a hoy: cubierto con parcial ($0.00, sin
    # amarillo), cubierto con exceso ($320) y gasto real pre-vigencia (faltan).
    assert om["F6"].value == _FILA_V73_CUBIERTO_PARCIAL["concepto_egreso"]
    assert (om["G6"].value, om["H6"].value, omv["L6"].value) == (0.0, "09/10/2026", 2360.0)
    assert om["F7"].value == _FILA_V73_CUBIERTO_EXCEDE["concepto_egreso"]
    assert (om["G7"].value, omv["L7"].value) == (320.0, 860.0)
    assert om["F8"].value == CONCEPTO_PARCIAL
    assert (om["G8"].value, omv["L8"].value) == (1200.0, 400.0)
    assert om["G8"].comment.text == NOTA_PAGO_REAL
    # (d) TOTALES = Σ de las filas (vuelos + sueltas), remanente = Σ ing − Σ egr.
    todas = p["consolidado"]["otros_movimientos"]["filas"] + p["consolidado"][
        "otros_movimientos"
    ].get("filas_sueltas", [])
    tot = next(c.row for c in om["A"] if c.value == "TOTALES")
    assert omv[f"G{tot}"].value == _suma(todas, "egreso_mxn")
    assert omv[f"J{tot}"].value == _suma(todas, "ingreso_mxn")
    assert omv[f"L{tot}"].value == round(
        _suma(todas, "ingreso_mxn") - _suma(todas, "egreso_mxn"), 2
    )


@pytest.mark.parametrize("sin_provision", [..., None, False], ids=["ausente", "null", "false"])
def test_general_con_regla_sin_bandera_textos_cubre_avion_de_hoy(sin_provision) -> None:
    """API previo (sigue mandando la marca, sin la bandera): las leyendas
    «cubierto por el avión» de hoy y el amarillo de siempre."""
    p = _payload_regla_v73(
        sin_provision=sin_provision, filas=[_tg()._FILA_CUBIERTA, _FILA_V73_CUBIERTO_EXCEDE]
    )
    wb = load_workbook(
        BytesIO(render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p)))
    )
    om = wb["otros movimientos"]
    assert _OM_REGLA_COMISIONES_CUBRE_AVION.format(regla=REGLA) in om["A2"].value
    assert OM_REGLA_CUBRE_SIN_PROVISION.format(regla=REGLA) not in om["A2"].value
    assert OM_FRAG_SIN_PROVISION not in om["A2"].value
    resumen = _celda_que_empieza(wb["RESUMEN flota"], RESUMEN_PREFIJO)
    assert resumen.endswith(RESUMEN_FRAG_VIEJO + _RESUMEN_REGLA_CUBRE_AVION.format(regla=REGLA))
    pie = _textos(wb["balance"])
    vieja = _NOTA_EMPRESA_BASE + _EMPRESA_REGLA_COMISIONES_CUBRE_AVION.format(regla=REGLA)
    assert vieja in pie
    assert not any(EMPRESA_REGLA_CUBRE_SIN_PROVISION.format(regla=REGLA) in x for x in pie)
    assert not any(EMPRESA_BASE_SIN_PROVISION in x for x in pie)
    # La marca sigue pintando (fila 4 = `_FILA_CUBIERTA`), con su leyenda.
    assert _rellenos_amber(om) == ["F4", "G4"]
    assert _LEYENDA_PROVISION in _textos(om)


# --- Base del pie del bloque VUELATOUR (empresa) ------------------------------


def _tev():
    """`tests.test_balance_empresa_vuelatour` (general del API 0.0.59 con el
    bloque VUELATOUR y SIN vuelos de la regla), importado perezoso por la
    misma razón que `_tg`."""
    from tests import test_balance_empresa_vuelatour as tev

    return tev


def _payload_empresa(caso: str, sin_provision) -> dict:
    """Un general con el bloque VUELATOUR en los tres sabores del pie:
    «cubre avión» (API 0.0.66), regla sin «cubre avión» (API 0.0.65) y sin
    vuelos de la regla (API 0.0.59). `sin_provision` = la bandera de la hoja
    (`...` ⇒ no viaja)."""
    if caso == "cubre_avion":
        p = _tg()._payload_om_0066()
    elif caso == "regla_sin_cubre":
        p = _tg()._payload(api_0066=False)
    else:
        p = _tev()._payload()
    om = p["consolidado"].setdefault("otros_movimientos", {"filas": [_FILA_V73_VIGENCIA_SIN_GASTO]})
    om.pop("sin_provision_vendedor", None)
    if sin_provision is not ...:
        om["sin_provision_vendedor"] = sin_provision
    return p


_REGLA_EMPRESA_POR_CASO = {
    "cubre_avion": (
        _EMPRESA_REGLA_COMISIONES_CUBRE_AVION_SIN_PROVISION,
        _EMPRESA_REGLA_COMISIONES_CUBRE_AVION,
    ),
    "regla_sin_cubre": (_EMPRESA_REGLA_COMISIONES, _EMPRESA_REGLA_COMISIONES),
    "sin_regla": ("", ""),
}


def _pie_balance(p: dict) -> list[str]:
    wb = load_workbook(
        BytesIO(render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p)))
    )
    return _textos(wb["balance"])


@pytest.mark.parametrize("caso", ["cubre_avion", "regla_sin_cubre", "sin_regla"])
def test_bloque_empresa_con_bandera_base_sin_provisionado(caso) -> None:
    """Con la bandera, la base del pie del bloque VUELATOUR ya no dice «lo
    pagado o provisionado» (los egresos propios no traen provisión alguna),
    con «cubre avión», con regla sin «cubre avión» y sin regla."""
    con_bandera, _ = _REGLA_EMPRESA_POR_CASO[caso]
    pie = _pie_balance(_payload_empresa(caso, True))
    assert EMPRESA_BASE_SIN_PROVISION + con_bandera.format(regla=REGLA) in pie
    assert not any(_NOTA_EMPRESA_BASE in x or "o provisionado" in x for x in pie)


@pytest.mark.parametrize("sin_provision", [..., None, False], ids=["ausente", "null", "false"])
@pytest.mark.parametrize("caso", ["cubre_avion", "regla_sin_cubre", "sin_regla"])
def test_bloque_empresa_sin_bandera_base_de_hoy(caso, sin_provision) -> None:
    """Sin la bandera (API previo, que sigue provisionando) la base de hoy
    byte a byte, con la regla de hoy."""
    _, sin_bandera = _REGLA_EMPRESA_POR_CASO[caso]
    pie = _pie_balance(_payload_empresa(caso, sin_provision))
    assert _NOTA_EMPRESA_BASE + sin_bandera.format(regla=REGLA) in pie
    assert not any(EMPRESA_BASE_SIN_PROVISION in x for x in pie)


def test_bloque_empresa_sin_nota_del_api_pie_fijo_con_la_base_sin_provision() -> None:
    """Payload sin `nota` (pie fijo de respaldo): la base nueva sustituye a la
    de hoy dentro de `_NOTA_BLOQUE_EMPRESA`."""
    p = _payload_empresa("sin_regla", True)
    p["empresa"]["nota"] = None
    pie = _pie_balance(p)
    assert _NOTA_BLOQUE_EMPRESA.replace(_NOTA_EMPRESA_BASE, EMPRESA_BASE_SIN_PROVISION) in pie
    assert _NOTA_BLOQUE_EMPRESA not in pie
    assert not any("o provisionado" in x for x in pie)


def test_el_amarillo_no_depende_de_la_bandera() -> None:
    """`_MARCA_PROVISION` no se toca: una fila que TRAIGA la marca se pinta
    aunque viaje la bandera (tolerancia; el API 0.0.73 ya no la manda)."""
    p = _payload_regla_v73(filas=[_tg()._FILA_CUBIERTA, _FILA_V73_VIGENCIA_SIN_GASTO])
    om = load_workbook(
        BytesIO(render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p)))
    )["otros movimientos"]
    assert OM_FRAG_SIN_PROVISION in om["A2"].value
    assert _rellenos_amber(om) == ["F4", "G4"]
    assert _LEYENDA_PROVISION in _textos(om)


# --- Libro INDIVIDUAL: la bandera RAÍZ ----------------------------------------


def _individual(**extra) -> BalanceAvionRequest:
    return BalanceAvionRequest(
        matricula="N4142R", periodo_desde="2026-10-01", periodo_hasta="2026-10-31", **extra
    )


def _maestra_individual(req: BalanceAvionRequest, prefijo: str = MAESTRA_PREFIJO) -> str:
    return _celda_que_empieza(
        _libro(render_balance_avion_xlsx(req))["reporte horas N4142R"], prefijo
    )


def test_individual_con_bandera_raiz_cambia_solo_el_fragmento() -> None:
    maestra = _maestra_individual(_individual(sin_provision_vendedor=True))
    assert maestra == _cambia(MAESTRA_NOTA_VIEJA, MAESTRA_FRAG_VIEJO, MAESTRA_FRAG_SIN_PROVISION)
    assert "PROVISIÓN a la fecha del vuelo" not in maestra


@pytest.mark.parametrize("bandera", [..., None, False], ids=["ausente", "null", "false"])
def test_individual_sin_bandera_raiz_nota_de_hoy(bandera) -> None:
    extra = {} if bandera is ... else {"sin_provision_vendedor": bandera}
    assert _maestra_individual(_individual(**extra)) == MAESTRA_NOTA_VIEJA


def test_individual_ignora_la_bandera_de_la_hoja_otros_movimientos() -> None:
    """El libro individual no pinta 'otros movimientos': solo la bandera RAÍZ
    de su payload cambia su nota (la de la hoja es del general)."""
    req = _individual(
        otros_movimientos={"filas": [_FILA_V73_VIGENCIA_SIN_GASTO], "sin_provision_vendedor": True}
    )
    wb = _libro(render_balance_avion_xlsx(req))
    assert "otros movimientos" not in wb.sheetnames
    assert _celda_que_empieza(wb["reporte horas N4142R"], MAESTRA_PREFIJO) == MAESTRA_NOTA_VIEJA


# --- Rutas: el payload del API 0.0.73 entra sin 422 y activa los textos ------


def test_rutas_aceptan_el_payload_0073(monkeypatch) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    headers = {"X-Internal-Token": TOKEN}
    p = _payload_regla_v73()

    res = client.post("/pdf/balance-general-xlsx", json=p, headers=headers)
    assert res.status_code == 200, res.text
    om, resumen, _ = _leyendas_general(res.content)
    assert OM_FRAG_SIN_PROVISION in om and OM_REGLA_CUBRE_SIN_PROVISION.format(regla=REGLA) in om
    assert RESUMEN_FRAG_SIN_PROVISION in resumen
    assert _rellenos_amber(load_workbook(BytesIO(res.content))["otros movimientos"]) == []

    # La vista en pantalla es el MISMO libro (libro JSON, 12-oct-2026).
    res = client.post("/pdf/balance-general-libro", json=p, headers=headers)
    assert res.status_code == 200, res.text
    assert OM_FRAG_SIN_PROVISION in json.dumps(res.json(), ensure_ascii=False)

    res = client.post(
        "/pdf/balance-avion-xlsx",
        json=_individual(sin_provision_vendedor=True).model_dump(mode="json"),
        headers=headers,
    )
    assert res.status_code == 200, res.text
    maestra = _celda_que_empieza(_libro(res.content)["reporte horas N4142R"], MAESTRA_PREFIJO)
    assert maestra == _cambia(MAESTRA_NOTA_VIEJA, MAESTRA_FRAG_VIEJO, MAESTRA_FRAG_SIN_PROVISION)



# ===========================================================================
# 13-oct-2026 (API 0.0.74, invariante 61): «comisión del vendedor SOLO como
# gasto real». Bandera `comision_vendedor_solo_gasto_real`.
# ===========================================================================

# Textos del CONTRATO, literales. Reescritos EN SU LUGAR por el ajuste «dos
# categorías» del mismo día (contrato §G del ajuste; sección final): la
# bandera es la misma y estas pruebas siguen cubriendo esquemas, individual,
# independencia hoja/raíz, Libro Dinero y rutas.
ENCABEZADO_GASTO_REAL = (
    "COMISIONES = lo que absorbe el avión: su parte de la comisión bancaria de los cobros "
    "del vuelo (regla de septiembre 2026, por fecha del vuelo; en la misma proporción con "
    "que el cobro se prorratea al avión) + el pago al vendedor capturado en Gastos como "
    "«Comisión del vendedor (avión)» y ligado al vuelo (en cualquier fecha, en vuelos "
    "completados o cancelados con cobros retenidos; en un vuelo de varios aviones, su "
    "parte de la venta). La «Comisión del vendedor (VuelaTour)» no va aquí: la paga "
    "VuelaTour en 'otros movimientos'. Sin pago capturado no se resta nada. La nota de "
    "cada celda dice cuáles son."
)
MAESTRA_GASTO_REAL = (
    "COMISIONES MXN = lo que absorbe el avión: su parte de la comisión bancaria de los "
    "cobros (regla de septiembre 2026, por fecha del vuelo) y el pago al vendedor "
    "capturado en Gastos como «Comisión del vendedor (avión)» y ligado al vuelo (en "
    "cualquier fecha, en vuelos completados o cancelados con cobros retenidos; en un vuelo "
    "de varios aviones, su parte de la venta; sin ese pago no se resta nada); la nota de "
    "cada celda dice cuáles son. Lo cobrado al cliente por la comisión del vendedor sigue "
    "siendo ingreso de VuelaTour y vive en 'otros movimientos' del Balance general (ahí el "
    "pago capturado como «Comisión del vendedor (VuelaTour)» egresa apareado con lo "
    "cobrado y el capturado como «Comisión del vendedor (avión)» sale con egreso $0.00 «a "
    "cargo del avión», porque ya resta en esta columna). GANANCIA de la fila = REMANENTE "
    "(VENTA − COSTO TOTAL) − COMISIONES."
)
RESUMEN_GASTO_REAL = (
    "COMISIONES = las de la columna COMISIONES del libro individual de cada avión (su "
    "parte de la comisión bancaria de los cobros —regla de septiembre 2026, por fecha del "
    "vuelo— y el pago al vendedor capturado como «Comisión del vendedor (avión)» ligado al "
    "vuelo; sin ese pago no se resta nada del vendedor); en este libro no van en la hoja "
    "de vuelos (ahí la GANANCIA va antes de comisiones): restan aquí y en la cascada de "
    "cada avión de la hoja 'balance'. Lo cobrado al cliente por la comisión del vendedor "
    "sigue siendo INGRESO de VuelaTour; ese ingreso y el pago capturado como «Comisión del "
    "vendedor (VuelaTour)» viven en 'otros movimientos' "
)
DINERO_LEYENDA_GASTO_REAL = (
    "El pago de la comisión al vendedor va apareado en su fila SOLO cuando se captura en "
    "Gastos ligado al vuelo, como «Comisión del vendedor (VuelaTour)» o «Comisión del "
    "vendedor (avión)» (con «faltan $…» o «excede $…» si no cuadra con lo cobrado); "
    "mientras no se capture, la fila no lleva egreso (sin gasto real no se resta nada). "
    "Este libro es la caja de VuelaTour: resta los pagos de las dos categorías, aunque en "
    "el Balance general y en el reparto la «(avión)» la cargue el avión en su columna "
    "COMISIONES (la nota de la celda lo dice). "
)
DINERO_FILA9_GASTO_REAL = (
    "OTROS INGRESOS = ingreso de VuelaTour (TUAs/extras/pernocta/comisión del vendedor + "
    "su IVA), no del avión (regla 28-ago-2026), NETO del pago REAL al vendedor (los gastos "
    "«Comisión del vendedor (VuelaTour)» y «Comisión del vendedor (avión)» ligados al "
    "vuelo; sin gasto capturado no se resta nada; este libro resta los dos porque el "
    "dinero salió de VuelaTour, aunque en el Balance general y en el reparto la «(avión)» "
    "la cargue el avión), más los ingresos sin vuelo registrados en Ingresos (ING-n), "
    "netos de su comisión bancaria: ver hoja 'Otros ingresos'."
)


def test_textos_gasto_real_son_los_del_contrato() -> None:
    assert _NOTA_ENCABEZADO_COMISIONES_GASTO_REAL == ENCABEZADO_GASTO_REAL
    assert (
        _NOTA_COMISIONES_GASTO_REAL
        + _MAESTRA_PAGO_VENDEDOR_GASTO_REAL
        + " GANANCIA de la fila = REMANENTE (VENTA − COSTO TOTAL) − COMISIONES."
    ) == MAESTRA_GASTO_REAL
    assert _RESUMEN_COMISIONES_GASTO_REAL == RESUMEN_GASTO_REAL
    assert _LEYENDA_PAGO_VENDEDOR_GASTO_REAL == DINERO_LEYENDA_GASTO_REAL
    assert _UTILIDADES_OTROS_INGRESOS_GASTO_REAL == DINERO_FILA9_GASTO_REAL
    for t in (
        ENCABEZADO_GASTO_REAL,
        MAESTRA_GASTO_REAL,
        RESUMEN_GASTO_REAL,
        DINERO_LEYENDA_GASTO_REAL,
        DINERO_FILA9_GASTO_REAL,
    ):
        assert "provisi" not in t.lower(), t


def test_esquemas_aceptan_la_bandera_y_las_llaves_nuevas_con_default_none() -> None:
    """Cinco modelos con la bandera (hoja 'otros movimientos', raíz del libro
    individual, consolidado del general, Libro Dinero y reparto) y las llaves
    del pago real de la fila y de los totales: default None, sin 422."""
    assert BalanceHojaOtrosMovimientos().comision_vendedor_solo_gasto_real is None
    assert BalanceAvionRequest().comision_vendedor_solo_gasto_real is None
    assert DineroXlsxRequest().comision_vendedor_solo_gasto_real is None
    assert (
        RepartoPdfRequest(periodo_desde="a", periodo_hasta="b", generado="c")
        .comision_vendedor_solo_gasto_real
        is None
    )
    assert BalanceAvionVuelo().comision_vendedor_real_mxn is None
    assert BalanceAvionVuelo().comision_vendedor_real_pagos is None
    assert BalanceAvionTotales().comision_vendedor_real_mxn is None
    req = BalanceGeneralRequest.model_validate(
        {
            "consolidado": {
                "comision_vendedor_solo_gasto_real": True,
                "totales": {"comision_vendedor_real_mxn": 3000.0, "comision_vendedor_prov_mxn": 0},
                "vuelos": [
                    {"comision_vendedor_real_mxn": 3000.0, "comision_vendedor_real_pagos": 1}
                ],
                "otros_movimientos": {
                    "filas": [],
                    "sin_provision_vendedor": True,
                    "comision_vendedor_solo_gasto_real": True,
                },
            },
            "aviones": [{"matricula": "XA-TST", "comision_vendedor_solo_gasto_real": True}],
        }
    )
    assert req.consolidado.comision_vendedor_solo_gasto_real is True
    assert req.consolidado.otros_movimientos.comision_vendedor_solo_gasto_real is True
    assert req.consolidado.vuelos[0].comision_vendedor_real_pagos == 1
    assert req.consolidado.totales.comision_vendedor_real_mxn == 3000.0
    assert req.aviones[0].comision_vendedor_solo_gasto_real is True
    dinero = DineroXlsxRequest.model_validate({"comision_vendedor_solo_gasto_real": True})
    assert dinero.comision_vendedor_solo_gasto_real is True


# --- Libro INDIVIDUAL: la bandera raíz manda sobre `sin_provision_vendedor` --


@pytest.mark.parametrize("sin_provision", [..., True], ids=["sin 0.0.73", "con 0.0.73"])
def test_individual_con_la_bandera_raiz_0074_nota_del_pago_real(sin_provision) -> None:
    extra = {"comision_vendedor_solo_gasto_real": True}
    if sin_provision is not ...:
        extra["sin_provision_vendedor"] = sin_provision
    assert _maestra_individual(_individual(**extra), MAESTRA_PREFIJO_GASTO_REAL) == (
        MAESTRA_GASTO_REAL
    )


@pytest.mark.parametrize("valor", [None, False], ids=["null", "false"])
def test_individual_con_la_bandera_apagada_nota_de_hoy(valor) -> None:
    assert _maestra_individual(_individual(comision_vendedor_solo_gasto_real=valor)) == (
        MAESTRA_NOTA_VIEJA
    )
    assert _maestra_individual(
        _individual(comision_vendedor_solo_gasto_real=valor, sin_provision_vendedor=True)
    ) == _cambia(MAESTRA_NOTA_VIEJA, MAESTRA_FRAG_VIEJO, MAESTRA_FRAG_SIN_PROVISION)


def test_individual_ignora_la_bandera_de_la_hoja_y_el_general_la_raiz_del_individual() -> None:
    """La raíz del individual solo cambia SU libro; la de la hoja 'otros
    movimientos' no toca el individual (no la pinta)."""
    req = _individual(
        otros_movimientos={
            "filas": [_FILA_V73_VIGENCIA_SIN_GASTO],
            "comision_vendedor_solo_gasto_real": True,
        }
    )
    assert _maestra_individual(req) == MAESTRA_NOTA_VIEJA
    # Un general SIN bandera en la hoja ni en el consolidado, con aviones que
    # sí la traen en su raíz: fila 2 y RESUMEN de hoy.
    p = _general(...).model_dump(mode="json")
    p["aviones"] = [{"matricula": "N4142R", "comision_vendedor_solo_gasto_real": True}]
    om, resumen, maestra = _leyendas_general(
        render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    )
    assert om == OM_FILA2_VIEJA
    assert resumen == RESUMEN_NOTA_VIEJA
    assert maestra == MAESTRA_GENERAL_NOTA


# --- General SIN vuelos de la regla: fila 2, RESUMEN y bloque sin la regla ----


def test_general_con_la_bandera_sin_vuelos_de_la_regla() -> None:
    """Con la bandera de la hoja y sin vuelos de la regla: fila 2 = el
    fragmento de las DOS categorías (sin frase de regla; manda sobre el de la
    0.0.73), RESUMEN = la nota del pago capturado + el fragmento de las dos
    categorías, y la base del bloque VUELATOUR de las dos categorías (sin «o
    provisionado»)."""
    req = _general(..., filas=[_FILA_V73_VIGENCIA_SIN_GASTO, _FILA_PAGO_REAL], sin_provision=True)
    req.consolidado.otros_movimientos.comision_vendedor_solo_gasto_real = True
    om, resumen, maestra = _leyendas_general(render_balance_general_xlsx(req))
    assert om == _cambia(OM_FILA2_VIEJA, OM_FRAG_VIEJO, OM_FRAG_DOS_CATEGORIAS)
    assert resumen == RESUMEN_GASTO_REAL + RESUMEN_FRAG_DOS_CATEGORIAS
    assert maestra == MAESTRA_GENERAL_NOTA  # el consolidado no trae la raíz
    for caso in ("sin_regla", "regla_sin_cubre"):
        p = _payload_empresa(caso, True)
        p["consolidado"]["otros_movimientos"]["comision_vendedor_solo_gasto_real"] = True
        pie = _pie_balance(p)
        assert any(x.startswith(EMPRESA_BASE_DOS_CATEGORIAS) for x in pie), caso
        assert not any("o provisionado" in x for x in pie), caso
        assert not any(EMPRESA_BASE_SIN_PROVISION in x for x in pie), caso


# --- Libro Dinero ---------------------------------------------------------------


def _dinero_0074(pagada=..., filas=None, bandera=True) -> DineroXlsxRequest:
    """El Libro Dinero del API 0.0.74: provisión SIEMPRE 0 y la bandera."""
    req = _dinero(pagada=pagada, provision=0.0, filas=filas)
    return req.model_copy(update={"comision_vendedor_solo_gasto_real": bandera})


_FILA_V74_SIN_GASTO = {  # sin gasto real: la fila de comisión SIN egreso
    "clave": "N4142R-317",
    "fecha_vuelo": "2026-09-20",
    "concepto_egreso": None,
    "egreso_mxn": None,
    "fecha_egreso": None,
    "nota_egreso": None,
    "concepto_ingreso": "comisión vendedor (Alex Saab)",
    "ingreso_mxn": 2030.0,
    "fecha_ingreso": "2026-09-20",
    "remanente_mxn": 2030.0,
}


def test_dinero_con_la_bandera_leyendas_nuevas_sin_provision() -> None:
    fila2, fila9 = _leyendas_dinero(render_dinero_xlsx(_dinero_0074(filas=[_FILA_V74_SIN_GASTO])))
    assert fila2 == _cambia(DINERO_FILA2_VIEJA, DINERO_FRAG_VIEJO + " ", DINERO_LEYENDA_GASTO_REAL)
    assert fila9 == DINERO_FILA9_GASTO_REAL
    assert "provisi" not in (fila2 + fila9).lower()
    # La fila sin gasto real: sin egreso, remanente = ingreso, sin nota.
    ws = _libro(render_dinero_xlsx(_dinero_0074(filas=[_FILA_V74_SIN_GASTO])))["Otros ingresos"]
    assert (ws["F4"].value, ws["G4"].value, ws["H4"].value) == (None, None, None)
    assert ws["G4"].comment is None
    assert (ws["J4"].value, ws["L4"].value) == (2030.0, 2030.0)


@pytest.mark.parametrize("pagada", [0.0, 1600.0], ids=["pagada 0", "pagada 1600"])
def test_dinero_con_la_bandera_y_pagos_reales(pagada) -> None:
    """Con pagos reales: la leyenda de la bandera MANDA sobre la del 0.0.39 y la
    fila 9 agrega lo pagado (solo con monto > 0); nunca «Provisión restada»."""
    fila2, fila9 = _leyendas_dinero(
        render_dinero_xlsx(
            _dinero_0074(pagada=pagada, filas=[_FILA_V74_SIN_GASTO, _FILA_PAGO_REAL])
        )
    )
    assert fila2 == _cambia(DINERO_FILA2_VIEJA, DINERO_FRAG_VIEJO + " ", DINERO_LEYENDA_GASTO_REAL)
    cola = (
        f" Pagado al vendedor con gasto real en este periodo: ${pagada:,.2f} MXN."
        if pagada
        else ""
    )
    assert fila9 == DINERO_FILA9_GASTO_REAL + cola
    assert "Provisión restada" not in fila9


@pytest.mark.parametrize("bandera", [None, False], ids=["null", "false"])
def test_dinero_con_la_bandera_apagada_leyendas_de_hoy(bandera) -> None:
    fila2, fila9 = _leyendas_dinero(
        render_dinero_xlsx(_dinero(provision=2030.0).model_copy(
            update={"comision_vendedor_solo_gasto_real": bandera}
        ))
    )
    assert fila2 == DINERO_FILA2_VIEJA
    assert fila9 == UTILIDADES_FILA9_BASE + " Provisión restada en este periodo: $2,030.00 MXN."


# --- Rutas: el payload del API 0.0.74 entra sin 422 ---------------------------


def test_rutas_aceptan_el_payload_0074(monkeypatch) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    headers = {"X-Internal-Token": TOKEN}
    res = client.post(
        "/pdf/balance-avion-xlsx",
        json=_individual(comision_vendedor_solo_gasto_real=True).model_dump(mode="json"),
        headers=headers,
    )
    assert res.status_code == 200, res.text
    maestra = _celda_que_empieza(
        _libro(res.content)["reporte horas N4142R"], MAESTRA_PREFIJO_GASTO_REAL
    )
    assert maestra == MAESTRA_GASTO_REAL
    res = client.post(
        "/pdf/balance-avion-libro",
        json=_individual(comision_vendedor_solo_gasto_real=True).model_dump(mode="json"),
        headers=headers,
    )
    assert res.status_code == 200, res.text
    assert "«Comisión del vendedor (VuelaTour)» egresa apareado con lo cobrado" in json.dumps(
        res.json(), ensure_ascii=False
    )
    res = client.post(
        "/pdf/dinero-xlsx",
        json=_dinero_0074(filas=[_FILA_V74_SIN_GASTO]).model_dump(mode="json"),
        headers=headers,
    )
    assert res.status_code == 200, res.text
    fila2, fila9 = _leyendas_dinero(res.content)
    assert DINERO_LEYENDA_GASTO_REAL in fila2 and fila9 == DINERO_FILA9_GASTO_REAL


# ===========================================================================
# 13-oct-2026 · dos categorías (ajuste del cliente de las 17:10; API 0.0.74,
# invariante 61 del API; MISMA bandera `comision_vendedor_solo_gasto_real`):
# «Comisión del vendedor (VuelaTour)» la paga VuelaTour en 'otros
# movimientos' y «Comisión del vendedor (avión)» la carga el avión en su
# columna COMISIONES. Cambian SOLO textos; sin la bandera, byte a byte.
# ===========================================================================

# Textos del CONTRATO §G del ajuste, literales (los reescritos de la sección
# anterior —ENCABEZADO_GASTO_REAL, MAESTRA_GASTO_REAL, RESUMEN_GASTO_REAL,
# DINERO_LEYENDA_GASTO_REAL, DINERO_FILA9_GASTO_REAL— también lo son).
NOTA_COMISIONES_DOS = (
    "COMISIONES MXN = lo que absorbe el avión: su parte de la comisión bancaria de los "
    "cobros (regla de septiembre 2026, por fecha del vuelo) y el pago al vendedor "
    "capturado en Gastos como «Comisión del vendedor (avión)» y ligado al vuelo (en "
    "cualquier fecha, en vuelos completados o cancelados con cobros retenidos; en un vuelo "
    "de varios aviones, su parte de la venta; sin ese pago no se resta nada); la nota de "
    "cada celda dice cuáles son. Lo cobrado al cliente por la comisión del vendedor sigue "
    "siendo ingreso de VuelaTour y vive en 'otros movimientos' del Balance general "
)
MAESTRA_FRAG_DOS = (
    "(ahí el pago capturado como «Comisión del vendedor (VuelaTour)» egresa apareado con "
    "lo cobrado y el capturado como «Comisión del vendedor (avión)» sale con egreso $0.00 "
    "«a cargo del avión», porque ya resta en esta columna)."
)
FUERA_DE_LA_HOJA_DOS = (
    "Las comisiones que absorbe cada avión (su parte de la comisión bancaria de sus cobros "
    "—regla de septiembre 2026, por fecha del vuelo— y el pago al vendedor capturado como "
    "«Comisión del vendedor (avión)» ligado al vuelo; sin ese pago no se resta nada del "
    "vendedor) no van en esta hoja: están en la columna COMISIONES del libro individual de "
    "cada avión y ya van restadas en la cascada de cada avión de la hoja 'balance'."
)
FLOTA_DOS = (
    "COMISIONES (regla de septiembre 2026, por fecha del vuelo, para la parte bancaria): "
    "las que absorbe cada avión —su parte de la comisión bancaria de sus cobros y el pago "
    "al vendedor capturado como «Comisión del vendedor (avión)» ligado al vuelo— no van en "
    "esta hoja: están en la columna COMISIONES del libro individual de cada avión "
    "(Reportes › Balance por avión) y ya van restadas en la cascada de cada avión de la "
    "hoja 'balance'. Lo cobrado al cliente por la comisión del vendedor es ingreso de "
    "VuelaTour: ese ingreso, el pago capturado como «Comisión del vendedor (VuelaTour)» y "
    "la parte de VuelaTour de la comisión bancaria viven en 'otros movimientos' (su "
    "leyenda explica cada línea)."
)
RESUMEN_FRAG_DOS_CATEGORIAS = (
    "(ingreso cobrado; el pago al vendedor se aparea solo cuando se captura ligado al "
    "vuelo: «Comisión del vendedor (VuelaTour)» como egreso y «Comisión del vendedor "
    "(avión)» con egreso $0.00 porque lo carga el avión; sin pago capturado no se "
    "provisiona nada en esa hoja)."
)
RESUMEN_REGLA_DOS = (
    " Con la {regla} (por fecha del vuelo), la comisión bancaria de 'otros movimientos' es "
    "solo la parte de VuelaTour; la del avión resta en su columna COMISIONES, igual que "
    "los pagos «Comisión del vendedor (avión)»."
)
OM_FRAG_DOS_CATEGORIAS = (
    "el pago de la comisión al vendedor va apareado en la misma fila SOLO cuando se "
    "captura en Gastos ligado al vuelo, con una de dos categorías: la «Comisión del "
    "vendedor (VuelaTour)» la paga VuelaTour y aquí egresa (con «faltan $…» o «excede $…» "
    "si no cuadra con lo cobrado); la «Comisión del vendedor (avión)» la carga el avión en "
    "la columna COMISIONES de su libro y aquí sale con egreso $0.00 «a cargo del avión …» "
    "(si ningún avión la carga —vuelo aún no completado, cancelado sin cobros retenidos o "
    "sin avión de la flota— la paga VuelaTour y el concepto lo dice). Mientras no se "
    "capture, la fila no lleva egreso. Ese pago NO se captura como «Otros gastos "
    "VuelaTour»: quedaría duplicado en la hoja 'otros gastos'."
)
OM_REGLA_DOS = (
    " Con la {regla} (por fecha del vuelo), la comisión bancaria de aquí es solo la parte "
    "de VuelaTour: la de cada avión resta en la columna COMISIONES de su libro individual "
    "y en la cascada de la hoja 'balance', igual que los pagos «Comisión del vendedor "
    "(avión)» que aquí salen con egreso $0.00."
)
EMPRESA_REGLA_DOS = (
    " Con la {regla} (por fecha del vuelo), de la comisión bancaria los egresos propios "
    "solo traen la parte de VuelaTour; los pagos «Comisión del vendedor (avión)» los carga "
    "cada avión en su columna COMISIONES (aquí egresan $0.00)."
)
EMPRESA_BASE_DOS_CATEGORIAS = (
    "Ingresos y egresos propios = hoja 'otros movimientos': lo cotizado de los vuelos del "
    "periodo (no cancelados) más los otros ingresos registrados, y lo pagado (el pago al "
    "vendedor solo cuando se captura: la «Comisión del vendedor (VuelaTour)» egresa aquí; "
    "la «(avión)» la carga el avión y aquí egresa $0.00) — no es lo cobrado: el RESULTADO "
    "no es dinero en caja."
)
CASCADA_DOS = (
    "Comisiones que absorbe el avión, en USD — las de la columna COMISIONES de su libro "
    "individual: su parte de la comisión bancaria de sus cobros (regla de septiembre 2026, "
    "por fecha del vuelo) y el pago al vendedor capturado como «Comisión del vendedor "
    "(avión)» ligado al vuelo. Por vuelo es la misma cifra que descuenta el reparto a "
    "socios, pero el total puede no coincidir con el del reparto: por ejemplo, si un vuelo "
    "que aún no se completa ya tiene un anticipo con comisión bancaria, este balance "
    "cuenta ese cobro y su comisión, y el reparto no (solo lee vuelos completados o "
    "cancelados)."
)
PARTICIPACION_DOS = (
    " El pago «Comisión del vendedor (avión)» es la excepción: se reparte entre los "
    "aviones como la venta (va en COMISIONES)."
)
DINERO_FILA6_DOS = (
    "CHARTERS y UTILIDAD dependen del costo por hora del proveedor (regla pendiente de "
    "definir con el equipo): esas celdas van vacías a propósito — no se inventan números. "
    "Las comisiones del vendedor tienen regla: lo cobrado al cliente es ingreso de "
    "VuelaTour (hoja 'Otros ingresos') y su pago real (las dos categorías) se descuenta "
    "ahí; en el Balance general y en el reparto la «Comisión del vendedor (avión)» la "
    "resta el avión."
)
# Cola de la fila 9 (solo con la llave > 0); «$X» = el monto con formato.
DINERO_NOTA_AVION_DOS = (
    " De lo pagado al vendedor, $X MXN se capturó como «Comisión del vendedor (avión)» en "
    "vuelos que carga un avión: el Balance general y el reparto se lo restan al avión "
    "(columna COMISIONES); aquí sigue restando porque el dinero salió de VuelaTour."
)
# La fila 6 de hoy (sin la bandera), literal.
DINERO_FILA6_VIEJA = (
    "CHARTERS y UTILIDAD dependen del costo por hora del proveedor (regla pendiente de "
    "definir con el equipo): esas celdas van vacías a propósito — no se inventan números. "
    "Las comisiones del vendedor ya tienen regla (28-ago-2026): ingreso y pago de "
    "VuelaTour en la hoja 'Otros ingresos', no restan al avión."
)


def test_dos_categorias_literales_del_contrato() -> None:
    assert _NOTA_ENCABEZADO_COMISIONES_GASTO_REAL == ENCABEZADO_GASTO_REAL
    assert _NOTA_COMISIONES_GASTO_REAL == NOTA_COMISIONES_DOS
    assert _MAESTRA_PAGO_VENDEDOR_GASTO_REAL == MAESTRA_FRAG_DOS
    assert NOTA_COMISIONES_DOS.endswith("del Balance general ")  # termina con un espacio
    assert _COMISIONES_FUERA_DE_LA_HOJA_GASTO_REAL == FUERA_DE_LA_HOJA_DOS
    assert _NOTA_ENCABEZADO_GANANCIA_FLOTA_GASTO_REAL.endswith(FUERA_DE_LA_HOJA_DOS)
    assert _NOTA_ENCABEZADO_REMANENTE_GASTO_REAL.endswith(FUERA_DE_LA_HOJA_DOS)
    assert _NOTA_COMISIONES_FLOTA_GASTO_REAL == FLOTA_DOS
    assert _RESUMEN_COMISIONES_GASTO_REAL == RESUMEN_GASTO_REAL
    assert _RESUMEN_PAGO_VENDEDOR_DOS_CATEGORIAS == RESUMEN_FRAG_DOS_CATEGORIAS
    assert _RESUMEN_REGLA_GASTO_REAL == RESUMEN_REGLA_DOS
    assert _OM_PAGO_VENDEDOR_DOS_CATEGORIAS == OM_FRAG_DOS_CATEGORIAS
    assert _OM_REGLA_COMISIONES_GASTO_REAL == OM_REGLA_DOS
    assert _EMPRESA_REGLA_COMISIONES_GASTO_REAL == EMPRESA_REGLA_DOS
    assert _NOTA_EMPRESA_BASE_DOS_CATEGORIAS == EMPRESA_BASE_DOS_CATEGORIAS
    assert _NOTA_CASCADA_COMISIONES_GASTO_REAL == CASCADA_DOS
    assert _NOTA_PARTICIPACION_PAGO_VENDEDOR_AVION == PARTICIPACION_DOS
    assert _LEYENDA_PAGO_VENDEDOR_GASTO_REAL == DINERO_LEYENDA_GASTO_REAL
    assert _UTILIDADES_OTROS_INGRESOS_GASTO_REAL == DINERO_FILA9_GASTO_REAL
    assert _UTILIDADES_FILA6_GASTO_REAL == DINERO_FILA6_DOS
    assert _UTILIDADES_NOTA_COMISION_AVION.format(monto="$X") == DINERO_NOTA_AVION_DOS


# Todo lo que se pinta con la bandera para el pago al vendedor.
_TEXTOS_DOS_CATEGORIAS = (
    ENCABEZADO_GASTO_REAL,
    MAESTRA_GASTO_REAL,
    RESUMEN_GASTO_REAL,
    FUERA_DE_LA_HOJA_DOS,
    FLOTA_DOS,
    RESUMEN_FRAG_DOS_CATEGORIAS,
    RESUMEN_REGLA_DOS,
    OM_FRAG_DOS_CATEGORIAS,
    OM_REGLA_DOS,
    EMPRESA_REGLA_DOS,
    EMPRESA_BASE_DOS_CATEGORIAS,
    CASCADA_DOS,
    PARTICIPACION_DOS,
    DINERO_LEYENDA_GASTO_REAL,
    DINERO_FILA9_GASTO_REAL,
    DINERO_FILA6_DOS,
    DINERO_NOTA_AVION_DOS,
)


def test_dos_categorias_sin_el_nombre_viejo_ni_provision() -> None:
    """Con la bandera ningún texto nombra la categoría vieja («Comisión del
    vendedor» seguido de la comilla de cierre) ni habla de una provisión. La
    ÚNICA mención es la negación literal del contrato en el fragmento del
    RESUMEN («sin pago capturado no se provisiona nada en esa hoja»)."""
    for t in _TEXTOS_DOS_CATEGORIAS:
        assert "«Comisión del vendedor»" not in t, t[:70]
        if t is RESUMEN_FRAG_DOS_CATEGORIAS:
            assert t.lower().count("provisi") == 1
            assert "no se provisiona nada en esa hoja" in t
        else:
            assert "provisi" not in t.lower(), t[:70]
    # Las leyendas de 'otros movimientos' y del bloque VUELATOUR nombran las
    # DOS categorías; las del avión, solo la que el avión carga.
    for t in (OM_FRAG_DOS_CATEGORIAS, RESUMEN_FRAG_DOS_CATEGORIAS, MAESTRA_FRAG_DOS):
        assert "«Comisión del vendedor (VuelaTour)»" in t
        assert "«Comisión del vendedor (avión)»" in t
    for t in (FUERA_DE_LA_HOJA_DOS, CASCADA_DOS, RESUMEN_REGLA_DOS, OM_REGLA_DOS):
        assert "«Comisión del vendedor (avión)»" in t
        assert "«Comisión del vendedor (VuelaTour)»" not in t


@pytest.mark.parametrize("hay_pago_real", [..., True, False], ids=["sin", "true", "false"])
@pytest.mark.parametrize("sin_provision", [..., True, False], ids=["ausente", "0073", "false"])
def test_fragmentos_con_la_bandera_evaluan_dos_categorias_primero(
    sin_provision, hay_pago_real
) -> None:
    datos: dict = {"comision_vendedor_solo_gasto_real": True}
    if sin_provision is not ...:
        datos["sin_provision_vendedor"] = sin_provision
    if hay_pago_real is not ...:
        datos["hay_pago_vendedor_real"] = hay_pago_real
    om = BalanceHojaOtrosMovimientos(**datos)
    assert _fragmento_om_pago_vendedor(om) == OM_FRAG_DOS_CATEGORIAS
    assert _fragmento_resumen_pago_vendedor(om) == RESUMEN_FRAG_DOS_CATEGORIAS


@pytest.mark.parametrize("bandera", [..., None, False], ids=["ausente", "null", "false"])
def test_fragmentos_sin_la_bandera_los_de_siempre_byte_a_byte(bandera) -> None:
    """Con solo `sin_provision_vendedor` (API 0.0.73) los `*_SIN_PROVISION`;
    con solo `hay_pago_vendedor_real`, los del 0.0.39; sin nada, los de la
    provisión. Ninguno es el de las dos categorías."""

    def om(**datos) -> BalanceHojaOtrosMovimientos:
        if bandera is not ...:
            datos["comision_vendedor_solo_gasto_real"] = bandera
        return BalanceHojaOtrosMovimientos(**datos)

    sin = om(sin_provision_vendedor=True, hay_pago_vendedor_real=True)
    assert _fragmento_om_pago_vendedor(sin) == OM_FRAG_SIN_PROVISION
    assert _fragmento_resumen_pago_vendedor(sin) == RESUMEN_FRAG_SIN_PROVISION
    real = om(hay_pago_vendedor_real=True)
    assert _fragmento_om_pago_vendedor(real) == _OM_PAGO_VENDEDOR_REAL == OM_FRAG_NUEVO
    assert _fragmento_resumen_pago_vendedor(real) == _RESUMEN_PAGO_VENDEDOR_REAL
    nada = om()
    assert _fragmento_om_pago_vendedor(nada) == _OM_PAGO_VENDEDOR_PROVISION
    assert _fragmento_resumen_pago_vendedor(nada) == _RESUMEN_PAGO_VENDEDOR_PROVISION
    assert _fragmento_resumen_pago_vendedor(None) == RESUMEN_FRAG_VIEJO


# --- Base del pie del bloque VUELATOUR, con y sin `afac_en_filas` ---------------


def _payload_empresa_dos(caso: str, *, afac: bool) -> dict:
    p = _payload_empresa(caso, True)
    om = p["consolidado"]["otros_movimientos"]
    om["comision_vendedor_solo_gasto_real"] = True
    if afac:
        om["afac_en_filas"] = True
    return p


@pytest.mark.parametrize("afac", [False, True], ids=["sin AFAC", "afac_en_filas"])
@pytest.mark.parametrize("caso", ["cubre_avion", "regla_sin_cubre", "sin_regla"])
def test_bloque_empresa_base_dos_categorias(caso, afac) -> None:
    """Con la bandera la base es la de las DOS categorías (manda sobre la de
    la 0.0.73), con el fragmento AFAC justo después si viaja `afac_en_filas`
    y, con «cubre avión» (API 0.0.66+), la regla nueva al final."""
    pie = _pie_balance(_payload_empresa_dos(caso, afac=afac))
    base = EMPRESA_BASE_DOS_CATEGORIAS + (_EMPRESA_AFAC_EN_FILAS if afac else "")
    if caso == "cubre_avion":
        assert base + EMPRESA_REGLA_DOS.format(regla=REGLA) in pie
    else:
        assert any(x.startswith(base) for x in pie), caso
    assert not any(EMPRESA_BASE_SIN_PROVISION in x or "o provisionado" in x for x in pie)
    if not afac:
        assert not any(_EMPRESA_AFAC_EN_FILAS in x for x in pie)


@pytest.mark.parametrize("afac", [False, True], ids=["sin AFAC", "afac_en_filas"])
def test_bloque_empresa_dos_categorias_sin_nota_pie_fijo(afac) -> None:
    p = _payload_empresa_dos("sin_regla", afac=afac)
    p["empresa"]["nota"] = None
    pie = _pie_balance(p)
    base = EMPRESA_BASE_DOS_CATEGORIAS + (_EMPRESA_AFAC_EN_FILAS if afac else "")
    assert _NOTA_BLOQUE_EMPRESA.replace(_NOTA_EMPRESA_BASE, base) in pie


# --- Libro Dinero: filas 2, 6 y 9 y la llave nueva ------------------------------


def _miembros(data: bytes) -> dict[str, bytes]:
    """Miembros del .xlsx salvo docProps/core.xml (lleva la hora)."""
    import zipfile

    z = zipfile.ZipFile(BytesIO(data))
    return {n: z.read(n) for n in sorted(z.namelist()) if n != "docProps/core.xml"}


def _dinero_dos(pagada=..., avion=...) -> DineroXlsxRequest:
    """El Libro Dinero del API 0.0.74 con la llave
    `utilidades_comision_vendedor_avion_mxn` (`...` ⇒ no viaja)."""
    req = _dinero_0074(pagada=pagada, filas=[_FILA_V74_SIN_GASTO, _FILA_PAGO_REAL])
    if avion is not ...:
        req = req.model_copy(update={"utilidades_comision_vendedor_avion_mxn": avion})
    return req


def test_dinero_esquema_acepta_la_llave_nueva_con_default_none() -> None:
    assert DineroXlsxRequest().utilidades_comision_vendedor_avion_mxn is None
    req = DineroXlsxRequest.model_validate(
        {"comision_vendedor_solo_gasto_real": True, "utilidades_comision_vendedor_avion_mxn": 900}
    )
    assert req.utilidades_comision_vendedor_avion_mxn == 900.0


def test_dinero_dos_categorias_filas_2_6_y_9_con_la_llave() -> None:
    wb = _libro(render_dinero_xlsx(_dinero_dos(pagada=1600.0, avion=1000.0)))
    assert wb["Otros ingresos"]["A2"].value == _cambia(
        DINERO_FILA2_VIEJA, DINERO_FRAG_VIEJO + " ", DINERO_LEYENDA_GASTO_REAL
    )
    assert wb["utilidades"]["A6"].value == DINERO_FILA6_DOS
    assert wb["utilidades"]["A9"].value == (
        DINERO_FILA9_GASTO_REAL
        + " Pagado al vendedor con gasto real en este periodo: $1,600.00 MXN."
        + DINERO_NOTA_AVION_DOS.replace("$X", "$1,000.00")
    )


@pytest.mark.parametrize("avion", [..., None, 0.0], ids=["ausente", "null", "cero"])
def test_dinero_dos_categorias_sin_la_llave_sin_cola_del_avion(avion) -> None:
    wb = _libro(render_dinero_xlsx(_dinero_dos(pagada=1600.0, avion=avion)))
    assert wb["utilidades"]["A6"].value == DINERO_FILA6_DOS
    assert wb["utilidades"]["A9"].value == (
        DINERO_FILA9_GASTO_REAL
        + " Pagado al vendedor con gasto real en este periodo: $1,600.00 MXN."
    )


@pytest.mark.parametrize("bandera", [None, False], ids=["null", "false"])
def test_dinero_sin_la_bandera_la_llave_no_mueve_un_byte(bandera) -> None:
    """Sin la bandera (API ≤ 0.0.73) la llave nueva se ignora y las filas 2, 6
    y 9 son las de hoy."""
    base = _dinero(provision=2030.0).model_copy(
        update={"comision_vendedor_solo_gasto_real": bandera}
    )
    con = base.model_copy(update={"utilidades_comision_vendedor_avion_mxn": 1000.0})
    a, b = render_dinero_xlsx(base), render_dinero_xlsx(con)
    assert _miembros(a) == _miembros(b)
    wb = _libro(b)
    assert wb["Otros ingresos"]["A2"].value == DINERO_FILA2_VIEJA
    assert wb["utilidades"]["A6"].value == DINERO_FILA6_VIEJA
    assert "Comisión del vendedor (avión)" not in wb["utilidades"]["A9"].value


def test_rutas_aceptan_el_dinero_con_la_llave_nueva(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    res = client.post(
        "/pdf/dinero-xlsx",
        json=_dinero_dos(pagada=1600.0, avion=1000.0).model_dump(mode="json"),
        headers={"X-Internal-Token": TOKEN},
    )
    assert res.status_code == 200, res.text
    fila9 = _libro(res.content)["utilidades"]["A9"].value
    assert fila9.endswith(DINERO_NOTA_AVION_DOS.replace("$X", "$1,000.00"))
