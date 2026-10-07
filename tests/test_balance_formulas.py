"""Balance por avión: FÓRMULAS visibles con el número del API en caché
(5-oct-2026). Pedido del cliente sobre la hoja «reporte horas»: «¿me puedes
apoyar poniendo con fórmulas las celdas que sean por suma? Por ejemplo los
permisos de AFAC, queremos que tenga las fórmulas visibles, no que sea solo un
texto con la cantidad».

Los payloads de aquí salen de un ESPEJO de la aritmética del API
(`aircraft-balance.service.ts`: fila de vuelo, totales, buildHoja, cascada y
consolidado del general), así cada fórmula se compara contra el número que el
API de verdad mandaría. Casos: fila sin TC de costos (z null), cancelada, sin
permiso AFAC (horas cobradas 0 y avión sin tarifa), multi-avión, 4 y 5
cobros, costo cero (excluido del promedio) y un redondeo intermedio del API
que la fórmula NO puede reproducir (la celda se queda como valor).
"""

from __future__ import annotations

import hashlib
import json
import math
import zipfile
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.schemas.reportes import BalanceAvionRequest, BalanceGeneralRequest
from app.services import xlsx_formulas
from app.services.balance_avion_xlsx import (
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from tests import test_balance_combustible_pago as t_combustible
from tests import test_balance_extension_horario as t_extension
from tests import test_balance_factura_vuelatour as t_factura
from tests import test_balance_general_hojas as t_hojas
from tests import test_balance_inventario_xlsx as t_inventario
from tests import test_comision_vendedor_notas as t_comision
from tests._evaluador_formulas import Libro, verificar_libro

# ---------------------------------------------------------------------------
# Espejo de la aritmética del API
# ---------------------------------------------------------------------------


def r2(x: float) -> float:
    """`round2` del API: Math.round(x × 100) / 100 (empate hacia +∞)."""
    return math.floor(x * 100 + 0.5) / 100


def _suma(valores) -> float:
    """reduce secuencial del API (sin compensar)."""
    total = 0.0
    for v in valores:
        total += v or 0
    return total


def _promedio(valores) -> float | None:
    vals = [v for v in valores if v is not None]
    return _suma(vals) / len(vals) if vals else None


def _fila_api(s: dict, *, afac: float | None, tc_prom: float | None) -> dict:
    """Una fila de vuelo como la arma el paso 2 de buildPayload."""
    cancelado = s.get("estado") == "CANCELADO"
    d, i, j, k, z, o = s["D"], s["I"], s["J"], s["K"], s["z"], s.get("O")
    total_mxn = r2(i * k) if i is not None and k is not None else None  # L
    iva_mxn = r2(j * k) if j is not None and k is not None else None  # M
    subtotal = r2(total_mxn - (iva_mxn or 0)) if total_mxn is not None else None  # N
    afac_mxn = r2(afac * z * d) if afac is not None and z is not None and d > 0 else None
    op, piloto, otros = s.get("op"), s.get("piloto"), s.get("otros")
    costo = r2((op or 0) + (piloto or 0) + (otros or 0) + (afac_mxn or 0))  # Y
    ae = costo / z if z is not None else None
    af = ae / 1.16 if ae is not None else None
    ag = ae - af if ae is not None else None
    ah = ag * z if ag is not None else None
    remanente = r2((total_mxn or 0) - costo)  # AI (todas las filas con parte)
    dif_iva = r2((iva_mxn or 0) - (ah or 0))
    tc_g = z if z is not None else tc_prom
    ganancia_usd = r2(remanente / tc_g) if tc_g is not None else None
    an = ae / o if ae is not None and o else None
    ao = an / 1.16 if an is not None else None
    cobros = s.get("cobros", [])
    cobrado_real = r2(_suma(c["monto_mxn"] for c in cobros))
    if cancelado:
        cobrado = total_mxn if total_mxn is not None else cobrado_real
        por_cobrar, por_cobrar_usd = 0, 0
    else:
        cobrado = s["cobrado"]
        por_cobrar = r2((total_mxn or 0) - cobrado)
        por_cobrar_usd = r2(por_cobrar / k) if k is not None else None
    return {
        "clave": s["clave"],
        "fecha": s["fecha"],
        "orden_ts": s["fecha"] + "T12:00:00Z",
        "ruta": s.get("ruta", "CUN-MID-CUN"),
        "estado": s.get("estado", "COMPLETADO"),
        "multi_avion": s.get("multi", False),
        "participacion": 0.5 if s.get("multi") else 1,
        "participacion_fuente": "tramos" if s.get("multi") else "unico",
        "horas_cobradas": r2(d),
        "tarifa_usd": s.get("E"),
        "iva_hr_usd": s.get("G"),
        "total_usd": i,
        "iva_usd": j,
        "tc_venta": k,
        "total_mxn": total_mxn,
        "iva_mxn": iva_mxn,
        "subtotal_mxn": subtotal,
        "total_cotizacion_mxn": s.get("cotizacion", total_mxn),
        "tiempo_vuelo": o,
        "taco_inicio": s.get("P"),
        "taco_fin": s.get("Q"),
        "op_mxn": r2(op) if op is not None else None,
        "piloto_mxn": r2(piloto) if piloto is not None else None,
        "otros_mxn": r2(otros) if otros is not None else None,
        "permiso_afac_mxn": afac_mxn,
        "costo_total_mxn": costo,
        "tc_costos": z,
        "costo_usd": r2(ae) if ae is not None else None,
        "costo_usd_siva": r2(af) if af is not None else None,
        "iva_pagado_usd": r2(ag) if ag is not None else None,
        "iva_pagado_mxn": r2(ah) if ah is not None else None,
        "remanente_mxn": remanente,
        "dif_iva_mxn": dif_iva,
        "comision_vendedor_mxn": None,
        "ganancia_mxn": remanente,
        "ganancia_usd": ganancia_usd,
        "costo_hr_usd": r2(an) if an is not None else None,
        "costo_hr_usd_siva": r2(ao) if ao is not None else None,
        "status_cobro": "Cobrado" if cancelado or por_cobrar == 0 else "Parcial",
        "cobros": cobros,
        "cobrado_mxn": cobrado,
        "cobrado_real_mxn": cobrado_real,
        "por_cobrar_mxn": por_cobrar,
        "por_cobrar_usd": por_cobrar_usd,
        # Insumo del promedio COSTO X HORA (anValues: AN con costo > 0).
        "_an": an if an is not None and costo > 0 else None,
    }


_SUMAS = [
    "horas_cobradas", "tiempo_vuelo", "total_mxn", "iva_mxn", "subtotal_mxn",
    "op_mxn", "piloto_mxn", "otros_mxn", "permiso_afac_mxn", "costo_total_mxn",
    "remanente_mxn", "dif_iva_mxn", "comision_vendedor_mxn", "ganancia_mxn",
    "ganancia_usd", "cobrado_mxn", "cobrado_real_mxn", "por_cobrar_mxn",
    "por_cobrar_usd",
]  # fmt: skip
# Totales de la regla de comisiones a cargo del avión (API 0.0.65).
_SUMAS_COMISIONES = ["comisiones_mxn", "comision_banco_avion_mxn", "comision_vendedor_prov_mxn"]


def _hoja_api(filas: list[tuple], tc_prom: float | None, horas: float, *, litros=False) -> dict:
    """buildHoja del API (+ litros/$ por litro en la hoja de combustible)."""
    salida = [
        {
            "fecha": f[0],
            "categoria": f[1],
            "detalle": f[2],
            "monto_mxn": f[3],
            **({"litros": f[4]} if litros else {}),
        }
        for f in filas
    ]
    total = r2(_suma(f[3] for f in filas))
    usd = 0 if total == 0 else (r2(total / tc_prom) if tc_prom is not None else None)
    hoja = {
        "filas": salida,
        "total_mxn": total,
        "usd": usd,
        "usd_hr": r2(usd / horas) if usd is not None and horas > 0 else None,
    }
    if litros:
        lt = r2(_suma(f[4] for f in filas))
        hoja["litros_total"] = lt
        hoja["precio_litro_prom"] = r2(total / lt) if lt > 0 and total > 0 else None
    return hoja


def _libro_api(
    matricula: str,
    afac: float | None,
    specs: list[dict],
    hojas: dict,
    socios: list,
    *,
    ajusta_fila=None,
) -> dict:
    """El payload de UN avión (buildPayload) con totales, hojas y cascada.
    `ajusta_fila(fila, spec, tc_prom)` (6-oct-2026, comisiones a cargo del
    avión, `test_balance_comisiones`) retoca cada fila ANTES de los totales
    y la cascada; sin él, el payload de siempre."""
    zs = [s["z"] for s in specs if s["z"] is not None]
    tc_prom = _suma(zs) / len(zs) if zs else None
    filas = [_fila_api(s, afac=afac, tc_prom=tc_prom) for s in specs]
    if ajusta_fila is not None:
        filas = [ajusta_fila(f, s, tc_prom) for f, s in zip(filas, specs, strict=True)]
    an = [f.pop("_an") for f in filas]
    totales = {campo: r2(_suma(f[campo] for f in filas)) for campo in _SUMAS}
    # API 0.0.65: las tres llaves SOLO si alguna fila trae la regla
    # (`filasVuelo.some(r => r.comisiones_detalle !== undefined)`).
    if any("comisiones_detalle" in f for f in filas):
        for campo in _SUMAS_COMISIONES:
            totales[campo] = r2(_suma(f.get(campo) for f in filas))
    totales["tc_promedio"] = r2(tc_prom) if tc_prom is not None else None
    an_validos = [x for x in an if x is not None]
    totales["costo_hr_prom_usd"] = r2(_promedio(an_validos)) if an_validos else None
    totales["total_cotizacion_mxn"] = r2(
        _suma(f["total_cotizacion_mxn"] for f in filas if f["estado"] != "CANCELADO")
    )
    totales["comision_banco_mxn"] = r2(
        _suma(c.get("comision_mxn") for f in filas for c in f["cobros"])
    )
    horas = totales["tiempo_vuelo"]
    h = {
        nombre: _hoja_api(hojas.get(nombre, []), tc_prom, horas)
        for nombre in ("gastos_indirectos", "otros_gastos", "refacciones", "permisos")
    }
    comb = _hoja_api(hojas.get("combustible", []), tc_prom, horas, litros=True)
    antes = totales["ganancia_usd"]
    restas = [comb["usd"], *(h[n]["usd"] for n in h)]
    despues = r2(antes - _suma(restas)) if all(x is not None for x in restas) else None
    cobrada = r2(despues - totales["por_cobrar_usd"]) if despues is not None else None
    return {
        "matricula": matricula,
        "periodo_desde": "2026-09-01",
        "periodo_hasta": "2026-09-30",
        "permiso_afac_usd_hr": afac,
        "vuelos": filas,
        "totales": totales,
        **h,
        "combustible": comb,
        "balance": {
            "utilidad_antes_usd": antes,
            "combustible_usd": comb["usd"],
            "gastos_indirectos_usd": h["gastos_indirectos"]["usd"],
            "refacciones_usd": h["refacciones"]["usd"],
            "otros_usd": h["otros_gastos"]["usd"],
            "permisos_usd": h["permisos"]["usd"],
            "utilidad_despues_usd": despues,
            "por_cobrar_usd": totales["por_cobrar_usd"],
            "utilidad_cobrada_usd": cobrada,
            "socios": [
                {
                    "nombre": n,
                    "porcentaje": p,
                    "monto_usd": r2((p / 100) * cobrada) if cobrada is not None else None,
                }
                for n, p in socios
            ],
        },
        "pendientes": ["Pendiente de prueba"],
    }


def _cobro(fecha: str, monto: float, comision: float | None = None) -> dict:
    return {
        "fecha": fecha,
        "monto_mxn": monto,
        "metodo": "TRANSFERENCIA",
        "comision_mxn": comision,
        "cuenta": "HSBC Pesos",
    }


# XA-TST: tarifa AFAC 25 USD/hr. Seis vuelos que cubren los casos del contrato.
_SPECS_TST = [
    # Normal, pagado completo.
    {"clave": "#401 · Cliente Uno", "fecha": "2026-09-03", "D": 2.5, "E": 1200,
     "G": 192, "I": 3480.0, "J": 480.0, "K": 17.25, "O": 2.4, "P": 1000.0,
     "Q": 1002.4, "op": 4500.5, "piloto": 1200.0, "otros": 350.75, "z": 17.32,
     "cobros": [_cobro("2026-09-01", 30000.0), _cobro("2026-09-04", 30030.0)],
     "cobrado": 60030.0},
    # MULTI-AVIÓN (50 %) con 4 cobros y comisión bancaria.
    {"clave": "#402 · Cliente Dos", "fecha": "2026-09-08", "multi": True,
     "ruta": "CUN-MID · 50 % COMPARTIDO", "D": 1.25, "E": 1250, "G": 200,
     "I": 1812.5, "J": 250.0, "K": 18.123456, "O": 1.31, "op": 2000.0,
     "z": 18.05, "cobros": [_cobro("2026-09-05", 5000.0),
                            _cobro("2026-09-06", 6000.5, 150.25),
                            _cobro("2026-09-07", 7000.25),
                            _cobro("2026-09-09", 9000.1)],
     "cobrado": 25000.0, "cotizacion": 40000.0},
    # SIN TC (traslado sin cotizar): z y K null ⇒ indicadores USD vacíos y
    # GANANCIA USD al TC promedio del libro.
    {"clave": "#403 · Traslado", "fecha": "2026-09-12", "D": 1.0, "E": 1200,
     "G": 0, "I": 1200.0, "J": 0.0, "K": None, "O": 0.95, "op": 800.0,
     "z": None, "cobrado": 0},
    # CANCELADO: venta = lo retenido, sin horas cobradas (sin AFAC).
    {"clave": "#404 · Cliente Tres", "fecha": "2026-09-15", "estado": "CANCELADO",
     "D": 0, "E": 1200, "G": 192, "I": 574.71, "J": 79.27, "K": 17.4,
     "O": None, "op": 3000.0, "z": 17.4, "cobros": [_cobro("2026-09-14", 10000.0)],
     "cotizacion": 60000.0},
    # Redondeo INTERMEDIO del API: op y piloto traen fracciones de centavo
    # (gastos USD × TC) y COSTO TOTAL = round2(crudos) ≠ ROUND(redondeados).
    {"clave": "#405 · Cliente Cuatro", "fecha": "2026-09-20", "D": 1.5, "E": 1200,
     "G": 192, "I": 2088.0, "J": 288.0, "K": 17.6, "O": 1.45, "op": 1500.004,
     "piloto": 800.004, "z": 17.6, "cobros": [_cobro("2026-09-19", 20000.0)],
     "cobrado": 20000.0},
    # Cliente INTERNO sin costos: COSTO X HORA = 0 y NO entra al promedio.
    {"clave": "#406 · Interno", "fecha": "2026-09-25", "D": 0, "E": None,
     "G": None, "I": 0.0, "J": 0.0, "K": 17.5, "O": 1.2, "z": 17.5, "cobrado": 0},
]  # fmt: skip

_HOJAS_TST = {
    "gastos_indirectos": [("2026-09-03", "Hangar", "Renta hangar", 5000.0),
                          ("2026-09-15", "Seguro", "Póliza", 3250.5)],
    "otros_gastos": [("2026-09-30", "Nómina",
                      "reparto manual: $980.00 de $4,000.00 MXN", 980.0)],
    "refacciones": [("2026-09-10", "Refacción", "Filtro de aceite", 1200.0)],
    "permisos": [("2026-09-05", "Permiso", "AFAC anual", 2500.0)],
    "combustible": [("2026-09-02", "Gas", "Carga CUN", 6000.0, 200.5),
                    ("2026-09-20", "Gas", "Carga MID", 4500.25, 150.25),
                    ("2026-09-25", "Gas", "Carga sin litros", 1000.0, None)],
}  # fmt: skip

# XB-DOS: SIN tarifa AFAC. La otra mitad del multi-avión (con 5 cobros: el
# 4º agrega el resto) y un vuelo normal.
_SPECS_DOS = [
    {"clave": "#402 · Cliente Dos", "fecha": "2026-09-08", "multi": True,
     "ruta": "MID-CUN · 50 % COMPARTIDO", "D": 1.25, "E": 1250, "G": 200,
     "I": 1812.5, "J": 250.0, "K": 18.123456, "O": 1.28, "op": 1800.0,
     "z": 18.05, "cobros": [_cobro("2026-09-05", 1000.0), _cobro("2026-09-06", 2000.0),
                            _cobro("2026-09-07", 3000.0), _cobro("2026-09-08", 4000.5),
                            _cobro("2026-09-09", 5000.25)],
     "cobrado": 15000.0, "cotizacion": 40000.0},
    {"clave": "#407 · Cliente Cinco", "fecha": "2026-09-18", "D": 3.0, "E": 1100,
     "G": 176, "I": 3828.0, "J": 528.0, "K": 17.9, "O": 3.1, "op": 5200.0,
     "piloto": 1500.0, "otros": 420.0, "z": 17.95,
     "cobros": [_cobro("2026-09-17", 68521.2)], "cobrado": 68521.2},
]  # fmt: skip

_HOJAS_DOS = {
    "gastos_indirectos": [("2026-09-11", "Hangar", "Renta hangar", 4200.0)],
    "combustible": [("2026-09-18", "Gas", "Carga CUN", 7300.4, 240.0)],
}

_SOCIOS = [("Socio A", 60.0), ("Socio B", 40.0)]


def _payload_tst() -> dict:
    return _libro_api("XA-TST", 25.0, _SPECS_TST, _HOJAS_TST, _SOCIOS)


def _payload_dos() -> dict:
    return _libro_api("XB-DOS", None, _SPECS_DOS, _HOJAS_DOS, [("Socio C", 100.0)])


def _individual_tst() -> BalanceAvionRequest:
    return BalanceAvionRequest.model_validate(_payload_tst())


# % COBRADO de 17 cifras (revisión 5-oct-2026): 18,500.50 ÷ 52,158.40 =
# 0.35469838031841466 en `repr`, pero openpyxl escribía «%.16g» =
# 0.3546983803184147. La caché tiene que escribirse igual que openpyxl o el
# libro, leído con data_only=True, difiere del anterior en el último bit.
_SPECS_PCT17 = [
    {"clave": "#501 · Cliente Seis", "fecha": "2026-09-10", "D": 2.0, "E": 1124.1,
     "G": 179.86, "I": 2607.92, "J": 359.71, "K": 20.0, "O": 2.1, "op": 3000.0,
     "z": 20.0, "cobros": [_cobro("2026-09-09", 18500.5)], "cobrado": 18500.5},
]  # fmt: skip


def _individual_pct17() -> BalanceAvionRequest:
    return BalanceAvionRequest.model_validate(
        _libro_api("XA-PCT", 25.0, _SPECS_PCT17, {}, [("Socio A", 100.0)])
    )


def _individual_dos() -> BalanceAvionRequest:
    return BalanceAvionRequest.model_validate(_payload_dos())


def _flota(libros: list[dict]) -> dict:
    """Hojas de FLOTA del consolidado (hojaFlota del API)."""

    def sum_t(campo):
        return r2(_suma(p["totales"].get(campo) for p in libros))

    def avg_t(campo):
        vals = [p["totales"][campo] for p in libros if p["totales"][campo] is not None]
        return r2(_suma(vals) / len(vals)) if vals else None

    totales = {
        campo: sum_t(campo) for campo in [*_SUMAS, "total_cotizacion_mxn", "comision_banco_mxn"]
    }
    # API 0.0.65: Σ de los libros SOLO si ALGUNO las trae
    # (`libros.some(p => p.totales.comisiones_mxn !== undefined)`): el libro
    # EXTERNOS o un avión sin vuelos de la regla no las traen.
    if any("comisiones_mxn" in p["totales"] for p in libros):
        for campo in _SUMAS_COMISIONES:
            totales[campo] = sum_t(campo)
    totales["tc_promedio"] = avg_t("tc_promedio")
    totales["costo_hr_prom_usd"] = avg_t("costo_hr_prom_usd")
    horas = totales["tiempo_vuelo"]

    def hoja_flota(nombre: str, litros: bool = False) -> dict:
        filas = sorted(
            (
                {**f, "detalle": f"{p['matricula']} · {f['detalle']}", "matricula": p["matricula"]}
                for p in libros
                for f in p[nombre]["filas"]
            ),
            key=lambda f: f["fecha"],
        )
        usd = r2(_suma(p[nombre]["usd"] for p in libros))
        hoja = {
            "filas": filas,
            "total_mxn": r2(_suma(p[nombre]["total_mxn"] for p in libros)),
            "usd": usd,
            "usd_hr": r2(usd / horas) if horas > 0 and usd != 0 else None,
        }
        if litros:
            lt = r2(_suma(p[nombre]["litros_total"] for p in libros))
            hoja["litros_total"] = lt
            hoja["precio_litro_prom"] = (
                r2(hoja["total_mxn"] / lt) if lt > 0 and hoja["total_mxn"] > 0 else None
            )
        return hoja

    return {
        "matricula": "FLOTA",
        "periodo_desde": "2026-09-01",
        "periodo_hasta": "2026-09-30",
        "permiso_afac_usd_hr": None,
        "vuelos": sorted((v for p in libros for v in p["vuelos"]), key=lambda v: v["orden_ts"]),
        "totales": totales,
        "gastos_indirectos": hoja_flota("gastos_indirectos"),
        "refacciones": hoja_flota("refacciones"),
        "otros_gastos": hoja_flota("otros_gastos"),
        "permisos": hoja_flota("permisos"),
        "combustible": hoja_flota("combustible", litros=True),
        "otros_movimientos": {
            "filas": [
                {
                    "clave": "#401 · Cliente Uno",
                    "estado": "COMPLETADO",
                    "concepto_ingreso": "TUA CUN",
                    "ingreso_mxn": 1250.5,
                    "concepto_egreso": "TUA pagado",
                    "egreso_mxn": 1100.25,
                    "remanente_mxn": r2(1250.5 - 1100.25),
                },
                {
                    "clave": "#402 · Cliente Dos",
                    "estado": "COMPLETADO",
                    "concepto_ingreso": "Pernocta",
                    "ingreso_mxn": 3000.0,
                    "remanente_mxn": 3000.0,
                },
            ],
            "filas_sueltas": [
                {
                    "clave": "ING-3",
                    "concepto_ingreso": "Intereses",
                    "ingreso_mxn": 85.33,
                    "remanente_mxn": 85.33,
                },
            ],
        },
        "pendientes": [],
    }


def _general_payload(libros: list[dict] | None = None) -> dict:
    libros = libros if libros is not None else [_payload_tst(), _payload_dos()]
    cons = _flota(libros)
    resumen, acc = (
        [],
        dict.fromkeys(
            [
                "horas",
                "horas_cobradas",
                "venta",
                "costo",
                "combustible",
                "comisiones",
                "ganancia",
                "cobrado",
                "por_cobrar",
            ],
            0.0,
        ),
    )
    pendientes = 0
    for p in libros:
        t, comb = p["totales"], p["combustible"]["total_mxn"]
        fila = {
            "matricula": p["matricula"],
            "vuelos": len(p["vuelos"]),
            "horas": t["tiempo_vuelo"],
            "horas_cobradas": t["horas_cobradas"],
            "venta_mxn": t["total_mxn"],
            "costo_mxn": t["costo_total_mxn"],
            "combustible_mxn": comb,
            # API 0.0.65: Σ COMISIONES del libro; antes, comision_vendedor_mxn.
            "comisiones_mxn": t.get("comisiones_mxn", t["comision_vendedor_mxn"]),
            "ganancia_mxn": r2(t["ganancia_mxn"] - comb),
            "cobrado_mxn": t["cobrado_mxn"],
            "por_cobrar_mxn": t["por_cobrar_mxn"],
            "pendientes": len(p["pendientes"]),
        }
        resumen.append(fila)
        for k, campo in [
            ("horas", "horas"),
            ("horas_cobradas", "horas_cobradas"),
            ("venta", "venta_mxn"),
            ("costo", "costo_mxn"),
            ("combustible", "combustible_mxn"),
            ("comisiones", "comisiones_mxn"),
            ("ganancia", "ganancia_mxn"),
            ("cobrado", "cobrado_mxn"),
            ("por_cobrar", "por_cobrar_mxn"),
        ]:
            acc[k] += fila[campo] or 0
        pendientes += fila["pendientes"]
    empresa = _hoja_api(
        [("2026-09-10", "Renta", "Oficina", 15000.0), ("2026-09-28", "Luz", "CFE", 2345.67)],
        cons["totales"]["tc_promedio"],
        0,
    )
    return {
        "periodo_desde": "2026-09-01",
        "periodo_hasta": "2026-09-30",
        "resumen": resumen,
        "resumen_totales": {
            "matricula": "TOTALES",
            "vuelos": 7,  # vuelos DISTINTOS: el multi-avión cuenta una vez
            "horas": r2(acc["horas"]),
            "horas_cobradas": r2(acc["horas_cobradas"]),
            "venta_mxn": r2(acc["venta"]),
            "costo_mxn": r2(acc["costo"]),
            "combustible_mxn": r2(acc["combustible"]),
            "comisiones_mxn": r2(acc["comisiones"]),
            "ganancia_mxn": r2(acc["ganancia"]),
            "cobrado_mxn": r2(acc["cobrado"]),
            "por_cobrar_mxn": r2(acc["por_cobrar"]),
            "pendientes": pendientes,
        },
        "consolidado": cons,
        "aviones": libros,
        "gastos_empresa": empresa,
        "inventario": {
            "filas": [
                {
                    "nombre": "Aceite 15w50",
                    "existencia": 12.5,
                    "valor_costo_mxn": 6250.75,
                    "compradas_cant": 20,
                    "compradas_costo_mxn": 10000.0,
                    "salidas_cant": 7.5,
                    "vendido_mxn": 4687.5,
                    "utilidad_mxn": 937.5,
                    "matriculas": "XA-TST",
                },
                {
                    "nombre": "Filtro",
                    "existencia": 3.25,
                    "valor_costo_mxn": 1950.33,
                    "compradas_cant": 4,
                    "compradas_costo_mxn": 2400.4,
                    "salidas_cant": 1,
                    "vendido_mxn": 750.1,
                    "utilidad_mxn": 150.02,
                    "matriculas": "XB-DOS",
                },
            ],
            "total_piezas": 15.75,
            "total_valor_mxn": r2(6250.75 + 1950.33),
            "total_compras_mxn": r2(10000.0 + 2400.4),
            "total_vendido_mxn": r2(4687.5 + 750.1),
            "total_utilidad_mxn": r2(937.5 + 150.02),
        },
    }


def _general() -> BalanceGeneralRequest:
    return BalanceGeneralRequest.model_validate(_general_payload())


# ---------------------------------------------------------------------------
# Lectura
# ---------------------------------------------------------------------------


def _libros(data: bytes):
    """(con fórmulas, con caché)."""
    return load_workbook(BytesIO(data)), load_workbook(BytesIO(data), data_only=True)


def _col(ws, encabezado: str) -> int:
    return next(c.column for c in ws[2] if c.value == encabezado)


def _formula(ws, fila: int, encabezado: str):
    return ws.cell(row=fila, column=_col(ws, encabezado)).value


# ---------------------------------------------------------------------------
# 1) Cada fórmula reproduce al API (evaluador independiente), en los
#    fixtures nuevos Y en los de las demás pruebas del balance.
# ---------------------------------------------------------------------------

_LIBROS = {
    "nuevo individual XA-TST": lambda: render_balance_avion_xlsx(_individual_tst()),
    "nuevo individual XB-DOS (sin AFAC)": lambda: render_balance_avion_xlsx(_individual_dos()),
    "nuevo individual XA-PCT (% de 17 cifras)": lambda: render_balance_avion_xlsx(
        _individual_pct17()
    ),
    "nuevo general": lambda: render_balance_general_xlsx(_general()),
    "hojas general": lambda: render_balance_general_xlsx(t_hojas._general()),
    "hojas individual": lambda: render_balance_avion_xlsx(t_hojas._individual()),
    "factura individual": lambda: render_balance_avion_xlsx(t_factura._individual()),
    "factura general": lambda: render_balance_general_xlsx(t_factura._general()),
    "combustible individual": t_combustible._individual,
    "combustible general": t_combustible._general,
    "extensión individual": lambda: render_balance_avion_xlsx(
        t_extension._individual(con_llave=True)
    ),
    "extensión general": lambda: render_balance_general_xlsx(t_extension._general(con_llave=True)),
    "inventario general": lambda: render_balance_general_xlsx(
        t_inventario._general(inventario=t_inventario._TIENDA_SEP26)
    ),
    # API sin `inventario`: hoja 'refacciones' con GANANCIA = venta − costo.
    "refacciones general (sin inventario)": lambda: render_balance_general_xlsx(
        t_inventario._general()
    ),
    "comisión general": lambda: render_balance_general_xlsx(t_comision._general(bandera=True)),
    # Variante «Balance general» (6-oct-2026, API 0.0.64): la hoja maestra con
    # costo total + costo por hora, mismos números del API.
    "nuevo general · costo por hora": lambda: render_balance_general_xlsx(
        _general().model_copy(update={"variante": "general"})
    ),
    "extensión general · costo por hora": lambda: render_balance_general_xlsx(
        t_extension._general(con_llave=True).model_copy(update={"variante": "general"})
    ),
}


@pytest.mark.parametrize("nombre", list(_LIBROS))
def test_cada_formula_reproduce_el_numero_del_api(nombre) -> None:
    verificadas = verificar_libro(_LIBROS[nombre]())
    assert verificadas, "el libro no trae ni una fórmula"


def test_el_libro_nuevo_trae_formulas_en_todas_las_hojas_con_sumas() -> None:
    hojas = {h for h, *_ in verificar_libro(render_balance_avion_xlsx(_individual_tst()))}
    assert hojas == {
        "reporte horas XA-TST",
        "cobranza",
        "combustible",
        "Gastos Indirectos",
        "refacciones",
        "permisos",
        "balance",
    }
    hojas_gen = {h for h, *_ in verificar_libro(render_balance_general_xlsx(_general()))}
    assert hojas_gen == {
        "RESUMEN flota",
        "reporte horas FLOTA",
        "otros movimientos",
        "cobranza",
        "combustible",
        "otros gastos",
        "repartidos a aviones",
        "inventario",
        "balance",
    }


# ---------------------------------------------------------------------------
# 2) Hoja maestra individual: constantes, fila por fila y TOTALES.
# ---------------------------------------------------------------------------


def test_constantes_del_encabezado_permiso_afac_y_factor_iva() -> None:
    wb, wv = _libros(render_balance_avion_xlsx(_individual_tst()))
    ws = wb["reporte horas XA-TST"]
    assert ws["A1"].value == "Permiso AFAC USD/hr"
    assert ws["B1"].value == 25.0
    assert ws["C1"].value == "Factor IVA costos"
    assert ws["D1"].value == 1.16
    assert "ROUND(esta tarifa × TIPO CAMBIO COSTOS × HORAS COBRADAS, 2)" in ws["B1"].comment.text
    # Los encabezados de grupo y de columna no se movieron.
    assert ws["E1"].value == "VENTA"
    assert ws.cell(row=2, column=1).value == "CLAVE"
    assert ws.freeze_panes == "D3"
    # La AFAC de cada vuelo cita la constante con $.
    assert _formula(ws, 3, "PERMISO AFAC\n(PROVISIÓN)") == "=ROUND($B$1*V3*E3,2)"
    assert wv["reporte horas XA-TST"]["T3"].value == r2(25 * 17.32 * 2.5)


def test_fila_normal_formulas_exactas() -> None:
    wb, _ = _libros(render_balance_avion_xlsx(_individual_tst()))
    ws = wb["reporte horas XA-TST"]
    esperadas = {
        "VENTA AVIÓN\nMXN": "=ROUND(H3*J3,2)",
        "IVA VENTA\nAVIÓN MXN": "=ROUND(I3*J3,2)",
        "TOTAL S/IVA\nMXN": "=ROUND(K3-L3,2)",
        "COSTO TOTAL\nMXN": "=ROUND(Q3+R3+S3+T3,2)",
        "COSTO TOTAL\nUSD": "=U3/V3",
        "COSTO TOTAL\nUSD S/IVA": "=W3/$D$1",
        "IVA PAGADO\nUSD": "=W3-X3",
        "IVA PAGADO\nMXN": "=Y3*V3",
        "REMANENTE\nVENTA−COSTO MXN": "=ROUND(K3-U3,2)",
        "DIF. IVA\nHACIENDA MXN": "=ROUND(L3-Z3,2)",
        "GANANCIA\nMXN": "=AA3",
        "GANANCIA\nUSD": "=ROUND(AD3/V3,2)",
        "COSTO X HORA\nUSD": "=W3/N3",
        "COSTO X HORA\nUSD S/IVA": "=AF3/$D$1",
        "COBRADO REAL\nMXN (Σ depósitos)": "=ROUND(SUM(AJ3,AL3,AN3,AP3),2)",
        "POR COBRAR\nMXN": "=ROUND(K3-AR3,2)",
        "POR COBRAR\nUSD": "=ROUND(AS3/J3,2)",
    }
    for encabezado, formula in esperadas.items():
        assert _formula(ws, 3, encabezado) == formula, encabezado
    # Insumos que NO están en el libro: siguen como valor.
    for encabezado in (
        "HORAS\nCOBRADAS",
        "TARIFA\nUSD/HR S/IVA",
        "IVA\nUSD/HR",
        "VENTA AVIÓN\nUSD *",
        "IVA VENTA\nAVIÓN USD",
        "TIPO CAMBIO\nVENTA",
        "TIEMPO\nVUELO HR",
        "OPERACIONES",
        "PILOTO",
        "OTROS",
        "TIPO CAMBIO\nCOSTOS",
        "COBRADO AVIÓN MXN\n(prorrateado, antes\nde comisiones) ****",
    ):
        valor = _formula(ws, 3, encabezado)
        assert not (isinstance(valor, str) and valor.startswith("=")), encabezado


def test_fila_sin_tc_de_costos_indicadores_vacios_y_ganancia_al_tc_promedio() -> None:
    wb, wv = _libros(render_balance_avion_xlsx(_individual_tst()))
    ws, wsv = wb["reporte horas XA-TST"], wv["reporte horas XA-TST"]
    fila = 5  # #403 · Traslado
    assert ws.cell(row=fila, column=1).value == "#403 · Traslado"
    for encabezado in (
        "COSTO TOTAL\nUSD",
        "COSTO TOTAL\nUSD S/IVA",
        "IVA PAGADO\nUSD",
        "IVA PAGADO\nMXN",
        "COSTO X HORA\nUSD",
        "COSTO X HORA\nUSD S/IVA",
        "PERMISO AFAC\n(PROVISIÓN)",
        "VENTA AVIÓN\nMXN",
        "POR COBRAR\nUSD",
    ):
        assert _formula(ws, fila, encabezado) is None, encabezado
    assert _formula(ws, fila, "GANANCIA\nUSD") == "=ROUND(AD5/AVERAGE($V$3:$V$8),2)"
    tc_prom = _suma([17.32, 18.05, 17.4, 17.6, 17.5]) / 5
    assert wsv.cell(row=fila, column=_col(ws, "GANANCIA\nUSD")).value == r2(-800 / tc_prom)


def test_cancelado_sin_afac_y_por_cobrar_cero() -> None:
    wb, wv = _libros(render_balance_avion_xlsx(_individual_tst()))
    ws, wsv = wb["reporte horas XA-TST"], wv["reporte horas XA-TST"]
    fila = 6  # #404 · Cliente Tres (CANCELADO)
    assert ws.cell(row=fila, column=4).value == "CANCELADO"
    assert _formula(ws, fila, "PERMISO AFAC\n(PROVISIÓN)") is None
    assert _formula(ws, fila, "POR COBRAR\nMXN") == "=ROUND(K6-AR6,2)"
    assert wsv.cell(row=fila, column=_col(ws, "POR COBRAR\nMXN")).value == 0
    assert wsv.cell(row=fila, column=_col(ws, "POR COBRAR\nUSD")).value == 0


def test_redondeo_intermedio_del_api_deja_la_celda_como_valor() -> None:
    """COSTO TOTAL del #405 = round2(1500.004 + 800.004 + 660) = 2960.01 en el
    API; con las celdas a la vista (1500.00 + 800.00 + 660.00) la fórmula daría
    2960.00 — la celda se queda como VALOR y las que la citan siguen siendo
    fórmula con el número correcto."""
    wb, wv = _libros(render_balance_avion_xlsx(_individual_tst()))
    ws = wb["reporte horas XA-TST"]
    fila = 7
    assert ws.cell(row=fila, column=1).value == "#405 · Cliente Cuatro"
    assert _formula(ws, fila, "COSTO TOTAL\nMXN") == 2960.01
    assert _formula(ws, fila, "COSTO TOTAL\nUSD") == "=U7/V7"
    assert _formula(ws, fila, "REMANENTE\nVENTA−COSTO MXN") == "=ROUND(K7-U7,2)"
    assert wv["reporte horas XA-TST"]["AA7"].value == r2(r2(2088 * 17.6) - 2960.01)


def test_multi_avion_y_cuatro_cobros() -> None:
    wb, wv = _libros(render_balance_avion_xlsx(_individual_tst()))
    ws, wsv = wb["reporte horas XA-TST"], wv["reporte horas XA-TST"]
    fila = 4  # #402 multi-avión, 4 cobros
    assert ws.cell(row=fila, column=8).comment is not None  # nota de participación
    assert _formula(ws, fila, "COBRADO REAL\nMXN (Σ depósitos)") == (
        "=ROUND(SUM(AJ4,AL4,AN4,AP4),2)"
    )
    assert wsv["AQ4"].value == r2(5000 + 6000.5 + 7000.25 + 9000.1)
    assert _formula(ws, fila, "VENTA AVIÓN\nMXN") == "=ROUND(H4*J4,2)"
    assert wsv["K4"].value == r2(1812.5 * 18.123456)


def test_cinco_cobros_el_cuarto_agrega_y_la_suma_cuadra() -> None:
    wb, wv = _libros(render_balance_avion_xlsx(_individual_dos()))
    ws = wb["reporte horas XB-DOS"]
    assert ws["AP3"].value == r2(4000.5 + 5000.25)
    assert ws["AQ3"].value == "=ROUND(SUM(AJ3,AL3,AN3,AP3),2)"
    assert wv["reporte horas XB-DOS"]["AQ3"].value == r2(15000.75)


def test_totales_sumas_y_promedios_como_el_api() -> None:
    req = _individual_tst()
    wb, wv = _libros(render_balance_avion_xlsx(req))
    ws, wsv = wb["reporte horas XA-TST"], wv["reporte horas XA-TST"]
    tot = 9
    assert ws.cell(row=tot, column=1).value == "TOTALES"
    assert ws["E9"].value == "=ROUND(SUM(E3:E8),2)"
    assert ws["T9"].value == "=ROUND(SUM(T3:T8),2)"
    assert ws["V9"].value == "=ROUND(AVERAGE(V3:V8),2)"
    assert ws["AF9"].value == '=ROUND(AVERAGEIF(AF3:AF8,">0"),2)'
    assert wsv["V9"].value == req.totales.tc_promedio
    assert wsv["AF9"].value == req.totales.costo_hr_prom_usd
    # El interno (#406) tiene COSTO X HORA = 0 y queda fuera del promedio.
    assert wsv["AF8"].value == 0
    for col in ("K", "U", "AA", "AE", "AQ", "AR", "AS", "AT"):
        assert ws[f"{col}9"].value == f"=ROUND(SUM({col}3:{col}8),2)", col


def test_avion_sin_tarifa_afac_constante_vacia_y_columna_sin_formula() -> None:
    wb, _ = _libros(render_balance_avion_xlsx(_individual_dos()))
    ws = wb["reporte horas XB-DOS"]
    assert ws["A1"].value == "Permiso AFAC USD/hr"
    assert ws["B1"].value is None
    assert "no tiene configurada" in ws["B1"].comment.text
    assert ws["D1"].value == 1.16
    for fila in (3, 4):
        assert ws[f"T{fila}"].value is None
        assert ws[f"U{fila}"].value == f"=ROUND(Q{fila}+R{fila}+S{fila}+T{fila},2)"


# ---------------------------------------------------------------------------
# 3) Demás hojas del libro individual.
# ---------------------------------------------------------------------------


def test_hojas_de_gastos_y_combustible_citan_la_maestra() -> None:
    wb, wv = _libros(render_balance_avion_xlsx(_individual_tst()))
    maestra = "'reporte horas XA-TST'"
    ws = wb["permisos"]
    assert ws["A5"].value == "=ROUND(SUM(D8:D8),2)"
    assert ws["B5"].value == f"={maestra}!$V$9"
    assert ws["C5"].value == f"=ROUND(A5/AVERAGE({maestra}!$V$3:$V$8),2)"
    assert ws["D5"].value == f"={maestra}!$N$9"
    assert ws["E5"].value == "=ROUND(C5/D5,2)"
    wc = wb["combustible"]
    assert wc["A5"].value == "=D11"
    assert wc["B5"].value == "=ROUND(SUM(C8:C10),2)"
    assert wc["C5"].value == "=ROUND(A5/B5,2)"
    assert wc["E5"].value == f"=ROUND(A5/AVERAGE({maestra}!$V$3:$V$8),2)"
    assert wc["F5"].value == f"=ROUND(E5/{maestra}!$N$9,2)"
    assert wc["C11"].value == "=ROUND(SUM(C8:C10),1)"
    assert wc["D11"].value == "=ROUND(SUM(D8:D10),2)"
    assert wv["combustible"]["C5"].value == r2(11500.25 / 350.75)


def test_cobranza_sumas_y_cotizacion_sin_cancelados() -> None:
    req = _individual_tst()
    wb, wv = _libros(render_balance_avion_xlsx(req))
    ws, wsv = wb["cobranza"], wv["cobranza"]
    tot = 14
    assert ws.cell(row=tot, column=1).value == "TOTALES"
    assert ws["J14"].value == '=ROUND(SUMIF(D8:D13,"<>CANCELADO",J8:J13),2)'
    assert wsv["J14"].value == req.totales.total_cotizacion_mxn
    assert ws["H8"].value == "=ROUND(F8-G8,2)"
    assert ws["A5"].value == "=F14"
    assert ws["G5"].value == "=B5/A5"
    assert ws["L14"].value == "=ROUND(SUM(L8:L13),2)"
    assert wsv["L14"].value == 150.25


def test_balance_cascada_y_reparto_con_formulas() -> None:
    req = _individual_tst()
    wb, wv = _libros(render_balance_avion_xlsx(req))
    ws, wsv = wb["balance"], wv["balance"]
    assert ws["B4"].value == "='reporte horas XA-TST'!$AE$9"
    assert ws["B5"].value == "='combustible'!$E$5"
    assert ws["B7"].value == "='refacciones'!$C$5"
    assert ws["B8"].value == "='permisos'!$C$5"
    assert ws["B9"].value == "=ROUND(B4-B5-B6-B7-B8,2)"
    assert ws["B10"].value == "='reporte horas XA-TST'!$AT$9"
    assert ws["B11"].value == "=ROUND(B9-B10,2)"
    assert ws["C15"].value == "=ROUND(B15/100*$B$11,2)"
    b = req.balance
    assert wsv["B9"].value == b.utilidad_despues_usd
    assert wsv["B11"].value == b.utilidad_cobrada_usd
    assert [wsv["C15"].value, wsv["C16"].value] == [s.monto_usd for s in b.socios]


def test_gastos_indirectos_fusionados_no_se_inventan_centavos() -> None:
    """TOTAL USD de la hoja fusionada = round(usd indirectos + usd otros) del
    API: 469.47 + 55.76 = 525.23, mientras TOTAL MXN ÷ TC promedio da 525.24.
    La celda queda como VALOR; la cascada la cita y cuadra."""
    wb, wv = _libros(render_balance_avion_xlsx(_individual_tst()))
    assert wb["Gastos Indirectos"]["C5"].value == 525.23
    assert wb["Gastos Indirectos"]["A5"].value == "=ROUND(SUM(D8:D10),2)"
    assert wb["balance"]["B6"].value == "='Gastos Indirectos'!$C$5"
    assert wv["balance"]["B6"].value == 525.23


# ---------------------------------------------------------------------------
# 4) Balance general.
# ---------------------------------------------------------------------------


def test_general_maestra_tarifa_por_avion_y_promedios_como_valor() -> None:
    req = _general()
    wb, wv = _libros(render_balance_general_xlsx(req))
    ws = wb["reporte horas FLOTA"]
    assert ws["B1"].value is None
    assert "POR AVIÓN" in ws["B1"].comment.text
    assert ws["D1"].value == 1.16
    filas = range(3, 3 + len(req.consolidado.vuelos))
    for fila in filas:
        valor = ws[f"T{fila}"].value
        assert not (isinstance(valor, str) and valor.startswith("=")), fila
    sin_tc = next(f for f in filas if ws[f"A{f}"].value == "#403 · Traslado")
    # Sin la columna COMISIONES (AC) desde el 7-oct-2026: GANANCIA USD = AD y
    # COSTO X HORA USD = AE en la hoja de vuelos del general.
    assert ws[f"AD{sin_tc}"].value == req.consolidado.vuelos[sin_tc - 3].ganancia_usd
    tot = 3 + len(req.consolidado.vuelos)
    assert ws[f"V{tot}"].value == req.consolidado.totales.tc_promedio
    assert ws[f"AE{tot}"].value == req.consolidado.totales.costo_hr_prom_usd
    assert ws[f"K{tot}"].value == f"=ROUND(SUM(K3:K{tot - 1}),2)"
    assert wv["reporte horas FLOTA"][f"K{tot}"].value == req.consolidado.totales.total_mxn


def test_general_resumen_ganancia_y_totales() -> None:
    req = _general()
    wb, wv = _libros(render_balance_general_xlsx(req))
    ws = wb["RESUMEN flota"]
    assert ws["I4"].value == "=ROUND(E4-F4-G4,2)"
    assert ws["I5"].value == "=ROUND(E5-F5-G5,2)"
    assert ws["B6"].value == 7  # vuelos distintos: valor
    assert ws["E6"].value == "=ROUND(SUM(E4:E5),2)"
    assert ws["L6"].value == "=ROUND(SUM(L4:L5),0)"
    assert wv["RESUMEN flota"]["I6"].value == req.resumen_totales.ganancia_mxn


def test_general_otros_gastos_balance_inventario_y_otros_movimientos() -> None:
    req = _general()
    wb, wv = _libros(render_balance_general_xlsx(req))
    og = wb["otros gastos"]
    assert og["C5"].value == "=ROUND(A5/B5,2)"
    assert og["E5"].value is None  # sin horas: el API no calcula USD X HR
    assert wv["otros gastos"]["C5"].value == req.gastos_empresa.usd
    rep = wb["repartidos a aviones"]
    assert not str(rep["C5"].value).startswith("=")  # Σ de los USD de cada avión
    bal = wb["balance"]
    # Bloque de cada avión: lo de arriba es valor; después, cobrada y socios
    # son fórmula.
    assert not str(bal["B5"].value).startswith("=")
    assert bal["B10"].value == "=ROUND(B5-B6-B7-B8-B9,2)"
    assert bal["B12"].value == "=ROUND(B10-B11,2)"
    inv = wb["inventario"]
    assert inv["B7"].value == "=ROUND(SUM(B5:B6),3)"
    assert inv["C7"].value == "=ROUND(SUM(C5:C6),2)"
    om = wb["otros movimientos"]
    assert om["I4"].value == "=ROUND(G4-D4,2)"
    tot = next(c.row for c in om["A"] if c.value == "TOTALES")
    assert om[f"G{tot}"].value == f"=ROUND(SUM(G4:G{tot - 1}),2)"
    assert om[f"I{tot}"].value == f"=ROUND(G{tot}-D{tot},2)"
    assert wv["otros movimientos"][f"I{tot}"].value == r2(1250.5 + 3000 + 85.33 - 1100.25)


# ---------------------------------------------------------------------------
# 5) Caché: data_only=True da los MISMOS números que el generador anterior.
# ---------------------------------------------------------------------------


def _firma_valores(data: bytes) -> str:
    """Huella de TODOS los valores a la vista (data_only=True), hoja por hoja.
    Los números se normalizan a float (openpyxl lee «7000» como int); las
    constantes nuevas del encabezado (A1:D1 de las hojas maestras) quedan
    fuera porque el generador anterior no las tenía."""
    wb = load_workbook(BytesIO(data), data_only=True)
    partes = []
    for ws in wb.worksheets:
        maestra = ws.title.startswith("reporte horas")
        for fila in ws.iter_rows():
            for c in fila:
                if c.value is None:
                    continue
                if maestra and c.row == 1 and c.column <= 4:
                    continue
                v = c.value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    v = repr(float(v))
                partes.append([ws.title, c.coordinate, v, c.number_format])
    return hashlib.sha256(json.dumps(partes, default=str).encode()).hexdigest()


# Calculadas con el generador ANTERIOR a este cambio (HEAD 8560f57, sin
# fórmulas) sobre los mismos payloads: el libro con fórmulas, leído con su
# caché, muestra exactamente los números de antes.
# Recalculadas el 6-oct-2026 (API 0.0.65) SOLO por los textos de COMISIONES
# (encabezados «COMISIONES MXN» y «COBRADO AVIÓN MXN (prorrateado, antes de
# comisiones)» y las notas al pie que los explican): comparadas celda por
# celda contra las de antes, ningún número cambió.
# 7-oct-2026 (API 0.0.66): recalculada SOLO la del general — su hoja de
# vuelos ya no lleva COMISIONES (lo de su derecha corre una columna) y
# cambian las notas que lo explican (pie, encabezado de GANANCIA, RESUMEN).
# Comparada celda por celda contra la de HEAD 4189a1e quitando esa columna:
# ningún número cambió (este payload no trae comisiones). Las tres del
# libro individual no se movieron.
_FIRMAS_ANTES = {
    "tst": "468b0e1d198e97f6677649f69c8251e6087c6de1abd46c5b9831cc1905c3b2a1",
    "dos": "0e66f2168336235fe45bb987d2906a2a05e5bed0a72053e145f8eca6489c88b6",
    "general": "82ca976de801490f91d6c4f1b82975542d1c887006b85a0f0cd5c027890491d7",
    # Con la caché en `repr` (17 cifras) esta daba 5472c557…: % COBRADO
    # difería del generador anterior en el último bit.
    "pct17": "179b9207b41bbe28a7c09ef3fd6481a7ddf156e7e91478843469814974408b1b",
}


@pytest.mark.parametrize(
    "clave,render",
    [
        ("tst", lambda: render_balance_avion_xlsx(_individual_tst())),
        ("dos", lambda: render_balance_avion_xlsx(_individual_dos())),
        ("general", lambda: render_balance_general_xlsx(_general())),
        ("pct17", lambda: render_balance_avion_xlsx(_individual_pct17())),
    ],
)
def test_cache_da_los_numeros_del_generador_anterior(clave, render) -> None:
    assert _firma_valores(render()) == _FIRMAS_ANTES[clave]


def test_sin_cache_la_celda_tiene_formula_y_con_cache_el_numero() -> None:
    data = render_balance_avion_xlsx(_individual_tst())
    wb, wv = _libros(data)
    assert wb["reporte horas XA-TST"]["K3"].value == "=ROUND(H3*J3,2)"
    assert wv["reporte horas XA-TST"]["K3"].value == r2(3480 * 17.25)
    # Excel recalcula al abrir (las fórmulas dan lo mismo que la caché).
    calc = zipfile.ZipFile(BytesIO(data)).read("xl/workbook.xml").decode()
    assert 'fullCalcOnLoad="1"' in calc


def test_full_calc_on_load_lo_pone_guardar_no_el_default_de_openpyxl() -> None:
    """openpyxl ya trae fullCalcOnLoad=True por omisión: la prueba de arriba
    pasaría aunque `finalizar` dejara de pedirlo (revisión 5-oct-2026). Aquí
    el libro llega con el default APAGADO."""
    wb = _libro_dos_hojas()
    wb.calculation.fullCalcOnLoad = False
    data = xlsx_formulas.guardar(wb)
    calc = zipfile.ZipFile(BytesIO(data)).read("xl/workbook.xml").decode()
    assert 'fullCalcOnLoad="1"' in calc


# ---------------------------------------------------------------------------
# 6) xlsx_formulas: inyección de la caché, verificación y redondeo.
# ---------------------------------------------------------------------------


def _libro_dos_hojas() -> Workbook:
    wb = Workbook()
    a = wb.active
    a.title = "Alfa"
    a["A1"], a["A2"] = 10.25, 4.5
    xlsx_formulas.escribir(a["A3"], "ROUND(SUM(A1:A2),2)", 14.75)
    b = wb.create_sheet("Beta Dos")
    b["B1"] = 3
    xlsx_formulas.escribir(b["B2"], "'Alfa'!A3*B1", 44.25)
    return wb


def test_inyeccion_data_only_true_y_false() -> None:
    data = xlsx_formulas.guardar(_libro_dos_hojas())
    wb, wv = _libros(data)
    assert wb["Alfa"]["A3"].value == "=ROUND(SUM(A1:A2),2)"
    assert wb["Beta Dos"]["B2"].value == "='Alfa'!A3*B1"
    assert wv["Alfa"]["A3"].value == 14.75
    assert wv["Beta Dos"]["B2"].value == 44.25


def test_inyeccion_conserva_el_resto_del_zip_byte_a_byte() -> None:
    wb = _libro_dos_hojas()
    cache = xlsx_formulas.finalizar(wb)
    buf = BytesIO()
    wb.save(buf)
    antes = buf.getvalue()
    despues = xlsx_formulas.inyectar_valores_cache(antes, cache)
    za, zd = zipfile.ZipFile(BytesIO(antes)), zipfile.ZipFile(BytesIO(despues))
    assert [i.filename for i in za.infolist()] == [i.filename for i in zd.infolist()]
    hojas = set(xlsx_formulas.rutas_de_hojas(za).values())
    for info in za.infolist():
        if info.filename in hojas:
            continue
        assert za.read(info.filename) == zd.read(info.filename), info.filename
        assert zd.getinfo(info.filename).compress_type == info.compress_type
    # En las hojas solo cambió el <v> de las fórmulas.
    xml_a = za.read("xl/worksheets/sheet1.xml").decode()
    xml_d = zd.read("xl/worksheets/sheet1.xml").decode()
    assert xml_d == xml_a.replace(
        "<f>ROUND(SUM(A1:A2),2)</f><v></v>", "<f>ROUND(SUM(A1:A2),2)</f><v>14.75</v>"
    )


def _cambiar_zip(data: bytes, cambios: dict[str, bytes], renombres: dict[str, str]) -> bytes:
    entrada = zipfile.ZipFile(BytesIO(data))
    salida = BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for info in entrada.infolist():
            nombre = renombres.get(info.filename, info.filename)
            z.writestr(nombre, cambios.get(info.filename, entrada.read(info.filename)))
    return salida.getvalue()


def test_la_hoja_se_resuelve_por_rels_no_por_orden() -> None:
    """Se intercambian los archivos sheet1/sheet2 y sus Target (uno relativo,
    otro absoluto): la caché tiene que caer en la hoja por su NOMBRE."""
    wb = _libro_dos_hojas()
    cache = xlsx_formulas.finalizar(wb)
    buf = BytesIO()
    wb.save(buf)
    z = zipfile.ZipFile(BytesIO(buf.getvalue()))
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rels = (
        rels.replace('Target="/xl/worksheets/sheet1.xml"', 'Target="TMP"')
        .replace('Target="/xl/worksheets/sheet2.xml"', 'Target="worksheets/sheet1.xml"')
        .replace('Target="TMP"', 'Target="/xl/worksheets/sheet2.xml"')
    )
    revuelto = _cambiar_zip(
        buf.getvalue(),
        {"xl/_rels/workbook.xml.rels": rels.encode()},
        {
            "xl/worksheets/sheet1.xml": "xl/worksheets/sheet2.xml",
            "xl/worksheets/sheet2.xml": "xl/worksheets/sheet1.xml",
        },
    )
    rutas = xlsx_formulas.rutas_de_hojas(zipfile.ZipFile(BytesIO(revuelto)))
    assert rutas == {"Alfa": "xl/worksheets/sheet2.xml", "Beta Dos": "xl/worksheets/sheet1.xml"}
    final = xlsx_formulas.inyectar_valores_cache(revuelto, cache)
    wv = load_workbook(BytesIO(final), data_only=True)
    assert wv["Alfa"]["A3"].value == 14.75
    assert wv["Beta Dos"]["B2"].value == 44.25


def test_inyeccion_acepta_v_vacia_autocerrada_y_quita_el_tipo() -> None:
    xml = (
        '<sheetData><row r="1"><c r="A1" s="2" t="str"><f>B1*2</f><v/></c>'
        '<c r="B1" t="n"><v>2</v></c></row></sheetData>'
    )
    out = xlsx_formulas._inyectar_en_hoja(xml, {"A1": 4.0}, "Hoja")
    assert '<c r="A1" s="2"><f>B1*2</f><v>4</v></c>' in out
    with pytest.raises(ValueError, match="no se encontró la fórmula"):
        xlsx_formulas._inyectar_en_hoja(xml, {"C9": 1.0}, "Hoja")


def test_inyeccion_hoja_inexistente_revienta() -> None:
    data = xlsx_formulas.guardar(_libro_dos_hojas())
    with pytest.raises(ValueError, match="no tiene la hoja"):
        xlsx_formulas.inyectar_valores_cache(data, {"Gamma": {"A1": 1.0}})


@pytest.mark.parametrize(
    "valor,texto",
    [
        (14.75, "14.75"),
        (7000.0, "7000"),
        (3, "3"),
        (-0.5, "-0.5"),
        (-0.0, "0"),
        (1e-05, "0.00001"),
        (1.5e16, "15000000000000000"),
        (17.123456, "17.123456"),
        # Como openpyxl («%.16g»), no repr: 0.30000000000000004 → «0.3».
        (0.1 + 0.2, "0.3"),
        (18500.50 / 52158.40, "0.3546983803184147"),
        (1.2345678901234567e-7, "0.0000001234567890123457"),
    ],
)
def test_numero_xml_sin_notacion_cientifica(valor, texto) -> None:
    assert xlsx_formulas.numero_xml(valor) == texto


def test_round_de_excel_medio_hacia_afuera() -> None:
    assert xlsx_formulas.redondear_excel(1.005, 2) == 1.01  # el binario es 1.00499…
    assert xlsx_formulas.redondear_excel(-2.345, 2) == -2.35
    assert xlsx_formulas.redondear_excel(2.5, 0) == 3
    assert xlsx_formulas.redondear_excel(15.7505, 3) == 15.751


def test_formula_que_no_reproduce_al_api_vuelve_a_valor() -> None:
    wb = Workbook()
    ws = wb.active
    # A la vista 1500.00 y 800.00; el API sumó los crudos 1500.004 + 800.004.
    ws["A1"], ws["A2"] = 1500.0, 800.0
    xlsx_formulas.escribir(ws["A3"], "ROUND(A1+A2,2)", 2300.01)
    xlsx_formulas.escribir(ws["A4"], "A3*2", 4600.02)  # cita a la degradada
    xlsx_formulas.escribir(ws["A5"], "A1/B9", 5.0)  # #DIV/0! (B9 vacía)
    # Empate: el API (binario) da 1081.87 y Excel 1081.88 ⇒ valor.
    ws["A6"] = 17.31
    xlsx_formulas.escribir(ws["A7"], "ROUND(25*A6*2.5,2)", 1081.87)
    cache = xlsx_formulas.finalizar(wb)
    assert ws["A3"].value == 2300.01
    assert ws["A4"].value == "=A3*2"
    assert ws["A5"].value == 5.0
    assert ws["A7"].value == 1081.87
    assert cache == {"Sheet": {"A4": 4600.02}}
    assert sorted(xlsx_formulas.registro(wb).degradadas()) == [
        ("Sheet", "A3"),
        ("Sheet", "A5"),
        ("Sheet", "A7"),
    ]
    # La verificación no crea celdas fantasma (B9 sigue sin existir).
    assert (9, 2) not in ws._cells


def test_formula_sin_numero_del_api_no_se_escribe() -> None:
    wb = Workbook()
    ws = wb.active
    with pytest.raises(TypeError):
        xlsx_formulas.escribir(ws["A1"], "1+1", None)  # type: ignore[arg-type]


def test_criterios_de_sumif_y_averageif() -> None:
    assert xlsx_formulas.cumple_criterio(5.0, ">0")
    assert not xlsx_formulas.cumple_criterio(0.0, ">0")
    assert not xlsx_formulas.cumple_criterio(xlsx_formulas.VACIO, ">0")
    assert xlsx_formulas.cumple_criterio("COMPLETADO", "<>CANCELADO")
    assert not xlsx_formulas.cumple_criterio("cancelado", "<>CANCELADO")
    assert xlsx_formulas.cumple_criterio(xlsx_formulas.VACIO, "<>CANCELADO")


def test_dos_libros_no_comparten_registro() -> None:
    a, b = Workbook(), Workbook()
    a.active["A1"] = 1
    xlsx_formulas.escribir(a.active["A2"], "A1+1", 2.0)
    assert xlsx_formulas.registro(b).por_hoja == {}
    assert xlsx_formulas.finalizar(b) == {}


def test_inyeccion_tolera_otro_orden_de_atributos() -> None:
    xml = '<row r="1"><c s="2" r="A1"><f>B1*2</f><v></v></c><c r="B1" t="n"><v>2</v></c></row>'
    out = xlsx_formulas._inyectar_en_hoja(xml, {"A1": 4.0}, "Hoja")
    assert '<c s="2" r="A1"><f>B1*2</f><v>4</v></c>' in out


def test_si_la_cache_no_se_puede_escribir_el_libro_sale_con_valores(monkeypatch) -> None:
    """Otra versión de openpyxl que serialice distinto no debe tumbar el
    balance ni dejar celdas vacías en el teléfono: el libro sale como antes,
    con VALORES."""

    def revienta(*_a, **_k):
        raise ValueError("formato inesperado")

    monkeypatch.setattr(xlsx_formulas, "inyectar_valores_cache", revienta)
    data = render_balance_avion_xlsx(_individual_tst())
    wb, _ = _libros(data)
    assert not Libro(data).formulas()
    assert wb["reporte horas XA-TST"]["K3"].value == r2(3480 * 17.25)
    assert _firma_valores(data) == _FIRMAS_ANTES["tst"]


# ---------------------------------------------------------------------------
# 7) Revisión 5-oct-2026: ROUND perdido, caché como openpyxl y «jamás un 500».
# ---------------------------------------------------------------------------


def _libro_con_encabezado(encabezado: str, formula: str, valor: float) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = encabezado
    ws["A2"], ws["A3"] = 1250.5, 3085.33
    xlsx_formulas.escribir(ws["A4"], formula, valor)
    return xlsx_formulas.guardar(wb)


def test_el_verificador_de_tests_detecta_un_round_perdido() -> None:
    """Quitar el ROUND exterior de una suma seguía cuadrando «a la vista»
    (tolerancia de medio centavo) y la regresión llegaba a producción."""
    with pytest.raises(AssertionError, match="no lleva ROUND exterior"):
        verificar_libro(_libro_con_encabezado("EGRESO\nMXN", "SUM(A2:A3)", 4335.83))
    # Con su ROUND, la misma suma pasa.
    assert verificar_libro(_libro_con_encabezado("EGRESO\nMXN", "ROUND(SUM(A2:A3),2)", 4335.83))
    # La lista blanca solo cubre la aritmética de fila de sus columnas…
    assert verificar_libro(_libro_con_encabezado("COSTO TOTAL\nUSD", "A2/A3", 1250.5 / 3085.33))
    # …no un TOTALES con función de esas mismas columnas.
    with pytest.raises(AssertionError, match="no lleva ROUND exterior"):
        verificar_libro(_libro_con_encabezado("COSTO TOTAL\nUSD", "SUM(A2:A3)", 4335.83))


def test_numero_xml_igual_que_openpyxl() -> None:
    """El <v> se lee como el número que openpyxl escribía antes (safe_string,
    «%.16g»), en cualquier magnitud."""
    from random import Random

    from openpyxl.compat.strings import safe_string

    rnd = Random(20261005)
    for _ in range(5000):
        x = rnd.uniform(-1e6, 1e6) * rnd.choice([1e-9, 1e-3, 1.0, 1e9, 1e17])
        assert float(xlsx_formulas.numero_xml(x)) == float(safe_string(x)), x


def test_un_error_del_evaluador_no_tumba_el_balance(monkeypatch) -> None:
    """Otra versión de openpyxl sin `ws._cells` (AttributeError) no es un 500:
    cada fórmula vuelve a su número del API y el libro es el de siempre."""

    def sin_cells(*_a, **_k):
        raise AttributeError("'Worksheet' object has no attribute '_cells'")

    monkeypatch.setattr(xlsx_formulas.Evaluador, "valor_celda", sin_cells)
    data = render_balance_avion_xlsx(_individual_tst())
    assert not Libro(data).formulas()
    assert _firma_valores(data) == _FIRMAS_ANTES["tst"]


def test_un_error_inesperado_solo_regresa_esa_celda_a_valor(monkeypatch) -> None:
    """El fallo del evaluador en UNA fórmula no se lleva las demás."""
    original = xlsx_formulas.Evaluador.valor_celda

    def falla_en_beta_b1(self, ws, fila, col):
        if ws.title == "Beta Dos" and (fila, col) == (1, 2):
            raise IndexError("fallo raro de openpyxl")
        return original(self, ws, fila, col)

    monkeypatch.setattr(xlsx_formulas.Evaluador, "valor_celda", falla_en_beta_b1)
    wb, wv = _libros(xlsx_formulas.guardar(_libro_dos_hojas()))
    assert wb["Alfa"]["A3"].value == "=ROUND(SUM(A1:A2),2)"
    assert wb["Beta Dos"]["B2"].value == 44.25  # degradada a su número
    assert wv["Alfa"]["A3"].value == 14.75


def test_si_la_verificacion_revienta_el_libro_sale_con_valores(monkeypatch) -> None:
    def revienta(*_a, **_k):
        raise RecursionError("maximum recursion depth exceeded")

    monkeypatch.setattr(xlsx_formulas, "finalizar", revienta)
    data = render_balance_general_xlsx(_general())
    assert not Libro(data).formulas()
    assert _firma_valores(data) == _FIRMAS_ANTES["general"]
