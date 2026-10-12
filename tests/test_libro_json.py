"""Libro de Excel → JSON para la vista en pantalla (12-oct-2026).

Pedido del cliente: ver el balance como el Excel sin descargarlo, con los
mismos colores, formato y estructura. La regla que se congela aquí: el JSON
sale del MISMO libro openpyxl que el .xlsx — mismas hojas, mismos números
(los de la caché de las fórmulas), mismos estilos — y ningún texto de celda
es una fórmula sin su número.
"""

from __future__ import annotations

import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.schemas.reportes import BalanceAvionRequest
from app.services import libro_json, xlsx_formulas
from app.services.balance_avion_xlsx import (
    libro_balance_avion,
    libro_balance_general,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from tests import test_balance_formulas as t_formulas

# ---------------------------------------------------------------------------
# Formato de número «como Excel»
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("valor", "fmt", "esperado"),
    [
        (1234.5, "#,##0.00", "1,234.50"),
        (-1234.5, "#,##0.00", "-1,234.50"),
        (0, "#,##0.00", "0.00"),
        (1234.5, '"$"#,##0.00', "$1,234.50"),
        (-1234.5, '"$"#,##0.00', "-$1,234.50"),
        (16.991632, "0.00####", "16.991632"),
        (17, "0.00####", "17.00"),
        (16.9916, "0.00####", "16.9916"),
        (2.35, "0.0", "2.4"),
        (1.005, "0.00", "1.01"),
        (8.857, '0.00"%"', "8.86%"),
        (0.25, "0.0%", "25.0%"),
        (3.7, "0", "4"),
        (1234.5, "General", "1234.5"),
        (1234.0, "General", "1234"),
        (0.1 + 0.2, "General", "0.3"),
        (-5, "#,##0.00;[Red](#,##0.00)", "(5.00)"),
        (0, "#,##0.00;-#,##0.00;\"-\"", "-"),
        (None, "#,##0.00", ""),
        ("texto", "#,##0.00", "texto"),
        (True, "General", "VERDADERO"),
        (float("nan"), "#,##0.00", ""),
    ],
)
def test_formatear_como_excel(valor, fmt, esperado):
    assert libro_json.formatear(valor, fmt) == esperado


# ---------------------------------------------------------------------------
# Libro sintético: cada rasgo del contrato
# ---------------------------------------------------------------------------


def _libro_sintetico() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "hoja uno"
    ws["A1"] = "Título"
    ws["A1"].font = Font(bold=True, size=14, color="0F4C81")
    ws.merge_cells("A1:C1")
    ws["A2"] = "Concepto"
    ws["A2"].fill = PatternFill("solid", fgColor="0F4C81")
    ws["A2"].font = Font(bold=True, color="FFFFFF")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="D5DBE3")
    ws["A2"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["B3"] = 1234.5
    ws["B3"].number_format = "#,##0.00"
    ws["B4"] = -20
    ws["B4"].number_format = "#,##0.00"
    ws["B5"] = 8.857
    ws["B5"].number_format = '0.00"%"'
    # Fórmula que reproduce al API: se queda viva, con su número.
    xlsx_formulas.escribir(ws["B6"], "=SUM(B3:B4)", 1214.5)
    ws["B6"].number_format = "#,##0.00"
    # Fórmula que NO reproduce al API: vuelve a ser valor.
    xlsx_formulas.escribir(ws["B7"], "=B3*2", 999)
    ws["B7"].number_format = "#,##0.00"
    ws["C3"] = "texto"
    ws["C3"].comment = Comment("nota de la celda", "VuelaTour")
    ws["D3"] = 1.5
    ws["D3"].number_format = "0.0"
    ws["D3"].alignment = Alignment(horizontal="right")
    ws.conditional_formatting.add(
        "B3:B7",
        CellIsRule(
            operator="lessThan",
            formula=["0"],
            fill=PatternFill("solid", fgColor="FFC7CE", bgColor="FFC7CE"),
            font=Font(color="9C0006"),
        ),
    )
    ws.column_dimensions["A"].width = 30
    ws.row_dimensions[2].height = 38
    ws.freeze_panes = "B3"
    ws.sheet_properties.tabColor = "15803D"
    wb.create_sheet("vacía")
    return wb


def _celda(hoja: dict, r: int, c: int) -> dict | None:
    return next((x for x in hoja["celdas"] if x["r"] == r and x["c"] == c), None)


def test_estructura_del_libro():
    libro = libro_json.exportar_libro(_libro_sintetico(), titulo="Prueba", archivo="prueba.xlsx")
    assert libro["version"] == 1
    assert libro["titulo"] == "Prueba" and libro["archivo"] == "prueba.xlsx"
    assert [h["nombre"] for h in libro["hojas"]] == ["hoja uno", "vacía"]
    hoja = libro["hojas"][0]
    assert hoja["pestana"] == "15803D"
    assert hoja["cuadricula"] is True
    assert hoja["inmovilizar"] == {"filas": 2, "columnas": 1}
    assert hoja["combinadas"] == [{"r1": 1, "c1": 1, "r2": 1, "c2": 3}]
    assert {"c": 1, "ancho": 30.0} in hoja["columnas"]
    assert {"r": 2, "alto": 38.0} in hoja["filas"]
    assert hoja["max_fila"] >= 7 and hoja["max_col"] >= 4
    vacia = libro["hojas"][1]
    assert vacia["celdas"] == [] and vacia["combinadas"] == [] and vacia["inmovilizar"] is None
    # Serializable tal cual (FastAPI lo manda como JSON).
    json.dumps(libro)


def test_celdas_estilos_formulas_y_condicional():
    libro = libro_json.exportar_libro(_libro_sintetico(), titulo="Prueba")
    hoja = libro["hojas"][0]
    estilos = libro["estilos"]

    titulo = _celda(hoja, 1, 1)
    assert titulo["v"] == "Título" and titulo["t"] == "Título"
    assert estilos[titulo["s"]] == {"b": True, "sz": 14.0, "co": "0F4C81"}
    # Las celdas tapadas por la combinada no viajan.
    assert _celda(hoja, 1, 2) is None and _celda(hoja, 1, 3) is None

    encabezado = estilos[_celda(hoja, 2, 1)["s"]]
    assert encabezado == {
        "bg": "0F4C81",
        "co": "FFFFFF",
        "b": True,
        "ha": "center",
        "va": "center",
        "wr": True,
        "bd": "lrtb",
    }

    b3 = _celda(hoja, 3, 2)
    assert b3["v"] == 1234.5 and b3["t"] == "1,234.50" and "f" not in b3
    assert estilos[b3["s"]] == {"nf": "#,##0.00"}  # positivo: sin rojo

    b4 = _celda(hoja, 4, 2)
    assert b4["v"] == -20 and b4["t"] == "-20.00"
    assert estilos[b4["s"]] == {"nf": "#,##0.00", "bg": "FFC7CE", "co": "9C0006"}

    assert _celda(hoja, 5, 2)["t"] == "8.86%"

    b6 = _celda(hoja, 6, 2)
    assert b6["f"] == "=SUM(B3:B4)" and b6["v"] == 1214.5 and b6["t"] == "1,214.50"

    b7 = _celda(hoja, 7, 2)
    assert "f" not in b7 and b7["v"] == 999 and b7["t"] == "999.00"

    c3 = _celda(hoja, 3, 3)
    assert c3["n"] == "nota de la celda" and c3["v"] == "texto"

    d3 = _celda(hoja, 3, 4)
    assert d3["t"] == "1.5" and estilos[d3["s"]] == {"ha": "right", "nf": "0.0"}

    # Ningún texto es una fórmula cruda.
    assert all(not str(x["t"]).startswith("=") for x in hoja["celdas"])
    # Los estilos no se repiten.
    assert len(estilos) == len({json.dumps(e, sort_keys=True) for e in estilos})


def test_vinculo_de_fase_2_viaja_si_la_celda_lo_trae():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 10
    libro_json.vincular(ws["A1"], {"tipo": "gasto", "id": "abc", "campo": "monto"})
    libro = libro_json.exportar_libro(wb, titulo="x")
    assert _celda(libro["hojas"][0], 1, 1)["vinculo"] == {
        "tipo": "gasto",
        "id": "abc",
        "campo": "monto",
    }


# ---------------------------------------------------------------------------
# Los balances reales: el JSON es el mismo libro que el .xlsx
# ---------------------------------------------------------------------------


def _comparar_con_xlsx(libro: dict, xlsx: bytes) -> int:
    """Hoja por hoja y celda por celda: mismos nombres, mismos valores (los
    de la caché de las fórmulas en el .xlsx), misma fórmula visible. Devuelve
    cuántas fórmulas vivas se compararon."""
    con_formulas = load_workbook(BytesIO(xlsx))
    con_valores = load_workbook(BytesIO(xlsx), data_only=True)
    assert [h["nombre"] for h in libro["hojas"]] == con_formulas.sheetnames
    formulas = 0
    for hoja in libro["hojas"]:
        ws_f = con_formulas[hoja["nombre"]]
        ws_v = con_valores[hoja["nombre"]]
        por_coord = {(x["r"], x["c"]): x for x in hoja["celdas"]}
        for fila in ws_f.iter_rows():
            for cell in fila:
                if cell.value is None:
                    continue
                d = por_coord.get((cell.row, cell.column))
                assert d is not None, f"{hoja['nombre']}!{cell.coordinate} falta en el JSON"
                assert not str(d["t"]).startswith("="), f"{hoja['nombre']}!{cell.coordinate}"
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                    assert d.get("f") == cell.value
                    cache = ws_v[cell.coordinate].value
                    donde = f"{hoja['nombre']}!{cell.coordinate}"
                    assert d["v"] == cache, f"{donde}: {d['v']} ≠ {cache}"
                else:
                    assert d["v"] == cell.value, f"{hoja['nombre']}!{cell.coordinate}"
        assert hoja["combinadas"] == sorted(
            (
                {"r1": rg.min_row, "c1": rg.min_col, "r2": rg.max_row, "c2": rg.max_col}
                for rg in ws_f.merged_cells.ranges
            ),
            key=lambda x: (x["r1"], x["c1"]),
        )
    return formulas


def test_balance_general_json_es_el_mismo_libro_que_el_xlsx():
    req = t_formulas._general()
    libro = libro_balance_general(req, archivo="balance-mensual-vuelatour-x.xlsx")
    xlsx = render_balance_general_xlsx(req)
    formulas = _comparar_con_xlsx(libro, xlsx)
    assert formulas > 0
    assert libro["archivo"] == "balance-mensual-vuelatour-x.xlsx"
    assert libro["titulo"].startswith("Balance mensual VuelaTour")
    assert any(h["inmovilizar"] for h in libro["hojas"])
    json.dumps(libro)


def test_balance_avion_json_es_el_mismo_libro_que_el_xlsx():
    req = BalanceAvionRequest.model_validate(t_formulas._payload_tst())
    libro = libro_balance_avion(req)
    xlsx = render_balance_avion_xlsx(req)
    formulas = _comparar_con_xlsx(libro, xlsx)
    assert formulas > 0
    assert libro["titulo"].startswith("Balance ")
    # Una nota de celda del libro real viaja al JSON (las notas de cobros).
    assert any("n" in x for h in libro["hojas"] for x in h["celdas"])
