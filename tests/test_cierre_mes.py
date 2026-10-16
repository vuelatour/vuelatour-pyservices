"""Cierre de mes formal y vuelos que pasan al mes siguiente (16-oct-2026, API
0.0.88, invariante 89 del API).

Pedido del cliente: «los vuelos que aún no se cobren se marquen en el balance
y se muevan para el siguiente mes hasta que se cobre y complete». El API
decide TODO; aquí solo se PINTAN los campos aditivos del contrato:

- Balance por avión / general (`BalanceAvionVuelo.cierre_marca`,
  `fila_sombra`, `diferido_a/de`; raíz `cierre`; general `resumen_cierre`):
  ESTADO con la marca (ámbar «Pasa a …» / azul «Viene de …»), la fila
  SOMBRA del mes de origen con sus horas y el dinero vacío (ninguna SUM la
  cuenta), la fila que VIENE con el dinero y sin horas, el bloque «CIERRE
  DEL MES» bajo TOTALES (individual) y al pie del RESUMEN (general), la
  hoja cobranza y la raíz `cierre` de la vista en pantalla.
- Reparto PDF/XLSX (`RepartoVueloLinea.cierre_marca`, raíz `cierre`): sello
  del cierre, vuelos que salen / entran por avión, externos y el aviso de
  `coincide_con_snapshot = false`.
- Libro Dinero (`DineroVueloFila.cierre_marca`, raíz `cierre`).
- SIN los campos (API ≤ 0.0.87), con ellos en neutro o con un mes ABIERTO
  sin vuelos que pasarían: los libros de siempre, BYTE A BYTE contra las
  huellas calculadas con el código del HEAD f0c8ad1 sobre los MISMOS
  payloads (`tests/_cierre_mes_payloads.py`).
"""

from __future__ import annotations

import copy
import hashlib
import json
import re
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.pdfbase.pdfmetrics import stringWidth

from app.config import get_settings
from app.main import app
from app.schemas.reparto import RepartoCierre, RepartoPdfRequest, RepartoVueloLinea
from app.schemas.reportes import (
    BalanceAvionRequest,
    BalanceAvionVuelo,
    BalanceCierre,
    BalanceGeneralRequest,
    DineroVueloFila,
    DineroXlsxRequest,
)
from app.services import cierre_mes, xlsx_formulas
from app.services.balance_avion_xlsx import (
    libro_balance_avion,
    libro_balance_general,
    render_balance_avion_xlsx,
    render_balance_general_xlsx,
)
from app.services.dinero_xlsx import render_dinero_xlsx
from app.services.reparto_pdf import _ANCHO_FOLIO_PT, _folio_vuelo, render_reparto_pdf
from app.services.reparto_xlsx import ENCABEZADOS_CIERRE, render_reparto_xlsx
from tests import _cierre_mes_payloads as pay
from tests._evaluador_formulas import verificar_libro

MAESTRA_TST = "reporte horas XA-TST"
MAESTRA_FLOTA = "reporte horas FLOTA"
AMBER = "FCD34D"  # ámbar del balance (`AMBER` de balance_avion_xlsx)

# ---------------------------------------------------------------------------
# Huellas «de siempre»: calculadas con el código del HEAD f0c8ad1 (antes del
# cierre de mes) corriendo los constructores de `_cierre_mes_payloads` —
# .xlsx: todos los miembros salvo docProps/core.xml; vista en pantalla: el
# JSON con `vinculos_gasto: true`; PDF: los flujos de contenido de cada
# página. Si una se mueve, un API ≤ 0.0.87 ya no ve su libro de siempre.
# ---------------------------------------------------------------------------
_HUELLAS_HEAD = {
    "individual": (
        "5ff8072c668e207a77822c0463cc24ee31a19b329873b2d5412156ab57f6e7e2",
        "68324b8e36a303b0a4620088433230e8bda79ba8868ce8041cd2034e3c1cb3cf",
    ),
    "general-mensual": (
        "00a974b56179ed1e8ffc1d0e4dc2d27711c5ffc76a714f186d5fec8c2bc8913c",
        "9abf85a8fa27a7200d4b642240a74110a9f84190a57a289c3331cd0a81ad2a90",
    ),
    "general-general": (
        "818589cf5528e33b95a2c5c72c635be9593a774f6886e119b426cd4ddca4d8b1",
        "ad7064d3103b9c562d0983c415c6cb020b95314d69bc3997816d130601041c9b",
    ),
}
_HUELLA_REPARTO_PDF_HEAD = "9448b553c276bcd2bd8ddb878136bf888bf2d68f85d2fc5a1a3aaee2da91d5c0"
_HUELLA_REPARTO_XLSX_HEAD = "4e1474b790df44f408f8634d116156b512c0c71824952734238ff42a82039dcc"
_HUELLA_DINERO_HEAD = "a4c952a2004a8b2d287047b106ea51746c1920ab22a5bb615037833b841685c5"


def _huella_xlsx(data: bytes) -> str:
    z = zipfile.ZipFile(BytesIO(data))
    h = hashlib.sha256()
    for nombre in sorted(z.namelist()):
        if nombre != "docProps/core.xml":
            h.update(nombre.encode())
            h.update(z.read(nombre))
    return h.hexdigest()


def _huella_json(libro: dict) -> str:
    return hashlib.sha256(
        json.dumps(libro, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()


def _huella_pdf(data: bytes) -> str:
    h = hashlib.sha256()
    for p in PdfReader(BytesIO(data)).pages:
        h.update(p.get_contents().get_data())
    return h.hexdigest()


def _texto_pdf(data: bytes) -> str:
    return re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(BytesIO(data)).pages))


def _balance(libro: str, p: dict) -> tuple[bytes, dict]:
    """(.xlsx, vista en pantalla) del libro `libro` con el payload `p`."""
    if libro == "individual":
        return (
            render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p)),
            libro_balance_avion(BalanceAvionRequest.model_validate({**p, "vinculos_gasto": True})),
        )
    return (
        render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p)),
        libro_balance_general(BalanceGeneralRequest.model_validate({**p, "vinculos_gasto": True})),
    )


def _base(libro: str) -> dict:
    if libro == "individual":
        return pay.individual_base()
    return pay.general_base(libro.removeprefix("general-"))


@pytest.fixture
def registros(monkeypatch) -> list[xlsx_formulas.RegistroFormulas]:
    """Registro de fórmulas de cada libro que se cierra (cuáles volvieron a
    valor porque no reprodujeron el número del API)."""
    capturados: list[xlsx_formulas.RegistroFormulas] = []
    original = xlsx_formulas.finalizar

    def espia(wb):
        cache = original(wb)
        capturados.append(xlsx_formulas.registro(wb))
        return cache

    monkeypatch.setattr(xlsx_formulas, "finalizar", espia)
    return capturados


def _fila_de(ws, prefijo: str) -> int:
    filas = [c.row for c in ws["A"] if isinstance(c.value, str) and c.value.startswith(prefijo)]
    assert filas, prefijo
    return filas[0]


def _encabezados(ws, fila: int = 2) -> dict[str, int]:
    return {str(c.value).replace("\n", " "): c.column for c in ws[fila] if c.value is not None}


# ---------------------------------------------------------------------------
# 1) Esquemas aditivos y liberales (skew en los dos sentidos, jamás un 422).
# ---------------------------------------------------------------------------


def test_esquemas_aditivos_con_default_y_liberales() -> None:
    v = BalanceAvionVuelo()
    assert (v.cierre_marca, v.fila_sombra, v.diferido_a, v.diferido_de) == (None, False, None, None)
    assert BalanceAvionRequest().cierre is None
    g = BalanceGeneralRequest()
    assert (
        (g.cierre, g.resumen_cierre) is not None and g.cierre is None and g.resumen_cierre is None
    )
    assert RepartoVueloLinea().cierre_marca is None
    assert RepartoPdfRequest(periodo_desde="a", periodo_hasta="b", generado="c").cierre is None
    assert DineroVueloFila().cierre_marca is None and DineroXlsxRequest().cierre is None

    raro = BalanceAvionVuelo.model_validate(
        {"cierre_marca": "   ", "fila_sombra": "true", "diferido_a": 7, "diferido_de": ["x"]}
    )
    assert (raro.cierre_marca, raro.fila_sombra, raro.diferido_a, raro.diferido_de) == (
        None, False, "7", None,
    )  # fmt: skip
    c = BalanceCierre.model_validate(
        {"estado": None, "pasan": None, "vienen": None, "pasarian": None,
         "periodo": "", "cerrado_por_nombre": 12}
    )  # fmt: skip
    assert (c.pasan, c.vienen, c.pasarian, c.periodo, c.cerrado_por_nombre) == (
        [], [], [], None, "12",
    )  # fmt: skip
    assert not c.tiene_algo()
    assert BalanceCierre(estado="cerrado").tiene_algo()
    # El folio llega como número del API (bigint): se pinta como texto.
    r = RepartoCierre.model_validate(
        {"version": True, "coincide_con_snapshot": "no", "diferidos_salen": [{"folio": 316}],
         "diferidos_entran": None, "pendientes_que_pasarian": "x"}
    )  # fmt: skip
    assert (r.version, r.coincide_con_snapshot) == (None, None)
    assert r.diferidos_salen[0].folio == "316"
    assert (r.diferidos_entran, r.pendientes_que_pasarian) == ([], [])


def test_textos_de_presentacion() -> None:
    assert cierre_mes.etiqueta_mes("2026-10") == "octubre 2026"
    assert cierre_mes.etiqueta_mes("2026-13") == "2026-13"
    assert cierre_mes.etiqueta_mes(None) is None
    # Sello en hora de Cancún: 00:30 UTC del 17 = 19:30 del 16 en Cancún.
    assert cierre_mes.fecha_corta("2026-10-17T00:30:00Z") == "16-oct-2026"
    assert cierre_mes.texto_estado_cierre(BalanceCierre.model_validate(pay.cierre_cerrado())) == (
        "Mes cerrado el 16-oct-2026 por Ale"
    )
    assert cierre_mes.texto_estado_cierre(BalanceCierre(estado="ABIERTO")) is None
    assert cierre_mes.texto_estado_cierre(None) is None
    reab = BalanceCierre(
        estado="REABIERTO",
        cerrado_at="2026-10-16T19:00:00Z",
        cerrado_por_nombre="Ale",
        reabierto_at="2026-10-17T15:00:00Z",
        reabierto_por_nombre="Dani",
    )
    assert cierre_mes.texto_estado_cierre(reab) == "Mes reabierto el 17-oct-2026 por Dani"
    # Color de la marca: azul solo lo que VIENE; lo que sale o pasará, ámbar
    # (también un vuelo que vino y vuelve a pasar: lo urgente es el cobro).
    entrada = cierre_mes.marca_es_entrada
    assert entrada(BalanceAvionVuelo(cierre_marca="Viene de agosto 2026", diferido_de="2026-08"))
    assert entrada(DineroVueloFila(cierre_marca="Viene de agosto 2026"))
    assert not entrada(BalanceAvionVuelo(cierre_marca=pay.MARCA_PASA, fila_sombra=True))
    assert not entrada(
        BalanceAvionVuelo(
            cierre_marca="Viene de septiembre 2026 · pasará a noviembre 2026 al cerrar",
            diferido_de="2026-09",
        )
    )


# ---------------------------------------------------------------------------
# 2) BYTE A BYTE: sin los campos, con ellos en neutro y con un mes ABIERTO
#    sin vuelos que pasarían.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("libro", list(_HUELLAS_HEAD))
def test_balance_sin_cierre_byte_a_byte(libro: str) -> None:
    data, vista = _balance(libro, _base(libro))
    assert (_huella_xlsx(data), _huella_json(vista)) == _HUELLAS_HEAD[libro]
    # Llaves nuevas explícitas en neutro (null / false): el mismo libro.
    data, vista = _balance(libro, pay.con_llaves_neutras(_base(libro), libro))
    assert (_huella_xlsx(data), _huella_json(vista)) == _HUELLAS_HEAD[libro]
    # Mes ABIERTO sin vuelos que pasarían (lo que manda el API 0.0.88 casi
    # siempre): el .xlsx idéntico; la vista solo gana la llave `cierre` (el
    # panel pinta el estado del mes en el encabezado).
    p = _base(libro)
    p["cierre"] = pay.cierre_abierto_vacio()
    data, vista = _balance(libro, p)
    assert _huella_xlsx(data) == _HUELLAS_HEAD[libro][0]
    assert vista.pop("cierre")["estado"] == "ABIERTO"
    assert _huella_json(vista) == _HUELLAS_HEAD[libro][1]


def test_reparto_sin_cierre_byte_a_byte() -> None:
    for p in (
        pay.reparto_base(),
        pay.con_llaves_neutras(pay.reparto_base(), "reparto"),
        {**pay.reparto_base(), "cierre": pay.reparto_cierre_abierto_vacio()},
    ):
        req = RepartoPdfRequest.model_validate(p)
        assert _huella_pdf(render_reparto_pdf(req)) == _HUELLA_REPARTO_PDF_HEAD
        assert _huella_xlsx(render_reparto_xlsx(req)) == _HUELLA_REPARTO_XLSX_HEAD


def test_dinero_sin_cierre_byte_a_byte() -> None:
    for p in (
        pay.dinero_base(),
        pay.con_llaves_neutras(pay.dinero_base(), "dinero"),
        {**pay.dinero_base(), "cierre": pay.cierre_abierto_vacio()},
    ):
        data = render_dinero_xlsx(DineroXlsxRequest.model_validate(p))
        assert _huella_xlsx(data) == _HUELLA_DINERO_HEAD


# ---------------------------------------------------------------------------
# 3) Balance por avión CON cierre.
# ---------------------------------------------------------------------------


def test_individual_sombra_no_suma_y_viene_suma(registros) -> None:
    render_balance_avion_xlsx(BalanceAvionRequest.model_validate(pay.individual_base()))
    degradadas_base = registros[-1].degradadas()
    data = render_balance_avion_xlsx(
        BalanceAvionRequest.model_validate(pay.individual_con_cierre())
    )
    # Ninguna fórmula nueva vuelve a valor: las de TOTALES (Σ de las filas,
    # sombra incluida en el rango) reproducen los totales del API, que no
    # cuentan la sombra.
    assert registros[-1].degradadas() == degradadas_base
    assert verificar_libro(data)
    wb = load_workbook(BytesIO(data))
    ws = wb[MAESTRA_TST]
    enc = _encabezados(ws)
    sombra = _fila_de(ws, "#499")
    viene = _fila_de(ws, "#403")
    tot = _fila_de(ws, "TOTALES")
    assert sombra < tot and viene < tot
    # ESTADO: la marca del API, ámbar la que pasa y azul la que viene; el
    # estado real va en la nota.
    est = ws.cell(row=sombra, column=enc["ESTADO"])
    assert est.value == pay.MARCA_PASA
    assert est.fill.fgColor.rgb.endswith(AMBER)
    assert "Estado del vuelo: COMPLETADO." in est.comment.text
    assert "conserva las horas" in est.comment.text
    ev = ws.cell(row=viene, column=enc["ESTADO"])
    assert ev.value == pay.MARCA_VIENE
    assert ev.fill.fgColor.rgb.endswith("DBEAFE")
    # Fila SOMBRA: horas reales y TODO el dinero vacío (sin fórmula).
    assert ws.cell(row=sombra, column=enc["TIEMPO VUELO HR"]).value == 0.95
    for titulo, col in enc.items():
        if col > enc["ESTADO"] and titulo not in ("TIEMPO VUELO HR", "STATUS"):
            assert ws.cell(row=sombra, column=col).value is None, titulo
    assert ws.cell(row=sombra, column=enc["STATUS"]).value == "Pasa a oct-2026"
    # Fila que VIENE: dinero completo y horas vacías con su nota.
    hv = ws.cell(row=viene, column=enc["TIEMPO VUELO HR"])
    assert hv.value is None and "agosto 2026" in hv.comment.text
    assert ws.cell(row=viene, column=enc["VENTA AVIÓN USD *"]).value == 1200
    # ESTADO se ensancha y la fila de la marca crece a sus renglones.
    assert ws.column_dimensions["D"].width == 30
    assert ws.row_dimensions[sombra].height == 40
    # TOTALES del libro con cierre = los del libro de siempre (mismo payload
    # de totales del API): la sombra no sumó nada.
    base = load_workbook(
        BytesIO(
            render_balance_avion_xlsx(BalanceAvionRequest.model_validate(pay.individual_base()))
        ),
        data_only=True,
    )[MAESTRA_TST]
    con = load_workbook(BytesIO(data), data_only=True)[MAESTRA_TST]
    tot_base = _fila_de(base, "TOTALES")
    for col in range(5, ws.max_column + 1):
        assert con.cell(row=tot, column=col).value == base.cell(row=tot_base, column=col).value, col


def test_individual_bloque_cierre_bajo_totales_y_nota() -> None:
    data = render_balance_avion_xlsx(
        BalanceAvionRequest.model_validate(pay.individual_con_cierre())
    )
    ws = load_workbook(BytesIO(data))[MAESTRA_TST]
    tit = _fila_de(ws, "CIERRE DEL MES")
    assert ws.cell(row=tit, column=1).value == "CIERRE DEL MES — Mes cerrado el 16-oct-2026 por Ale"
    assert tit > _fila_de(ws, "TOTALES")
    assert ws.cell(row=tit + 1, column=1).value == "Mes: septiembre 2026"
    pasan = _fila_de(ws, "Pasan a octubre 2026")
    assert ws.cell(row=pasan, column=1).value == (
        "Pasan a octubre 2026 (pendientes de cobro al cierre: no suman en este mes)"
    )
    assert ws.cell(row=pasan, column=1).fill.fgColor.rgb.endswith(AMBER)
    assert [ws.cell(row=pasan + 1, column=c).value for c in range(1, 8)] == [
        "CLAVE", "FECHA", "VENTA\nMXN", "COBRADO\nMXN", "POR COBRAR\nMXN", "VENTA\nUSD", "LEYENDA",
    ]  # fmt: skip
    assert [ws.cell(row=pasan + 2, column=c).value for c in range(1, 8)] == [
        "#499 · Cliente Nueve", "28/09/2026", 36632.89, 0, 36632.89, 2076.4, pay.LEYENDA_PASA,
    ]  # fmt: skip
    # Valores, nunca fórmulas: el bloque no entra a ninguna suma.
    assert not any(
        isinstance(ws.cell(row=r, column=c).value, str)
        and ws.cell(row=r, column=c).value.startswith("=")
        for r in range(tit, pasan + 8)
        for c in range(1, 8)
    )
    vienen = _fila_de(ws, "Vienen de agosto 2026")
    assert ws.cell(row=vienen + 2, column=1).value == "#403 · Traslado"
    assert ws.cell(row=vienen + 2, column=6).value is None  # venta USD no vino
    notas = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert any(n.startswith("CIERRE DE MES: una fila con ESTADO «Pasa a <mes>»") for n in notas)


def test_individual_cobranza_marca_y_celdas_vacias() -> None:
    data = render_balance_avion_xlsx(
        BalanceAvionRequest.model_validate(pay.individual_con_cierre())
    )
    assert verificar_libro(data)
    ws = load_workbook(BytesIO(data))["cobranza"]
    sombra = _fila_de(ws, "#499")
    assert ws.cell(row=sombra, column=4).value == pay.MARCA_PASA
    assert ws.cell(row=sombra, column=5).value == "Pasa a oct-2026"
    assert ws.cell(row=sombra, column=5).fill.fgColor.rgb.endswith(AMBER)
    # TOTAL A COBRAR … COMISIÓN BANCO vacías (POR COBRAR sin fórmula).
    assert [ws.cell(row=sombra, column=c).value for c in range(6, 13)] == [None] * 7
    viene = _fila_de(ws, "#403")
    assert ws.cell(row=viene, column=4).value == pay.MARCA_VIENE
    assert ws.cell(row=viene, column=5).value == "Cobrado"
    assert ws.column_dimensions["D"].width == 30
    notas = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert any(n.startswith("CIERRE DE MES: las filas «Pasa a <mes>»") for n in notas)


def test_individual_vista_en_pantalla_trae_el_cierre() -> None:
    p = pay.individual_con_cierre()
    vista = libro_balance_avion(BalanceAvionRequest.model_validate({**p, "vinculos_gasto": True}))
    assert list(vista)[-1] == "cierre"
    assert vista["cierre"]["estado"] == "CERRADO"
    assert vista["cierre"]["cerrado_por_nombre"] == "Ale"
    assert vista["cierre"]["pasan"][0]["folio"] == "499"
    assert vista["vinculos_gasto"] is True


# ---------------------------------------------------------------------------
# 4) Balance general CON cierre: hoja de vuelos, RESUMEN y cobranza.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("variante", ["mensual", "general"])
def test_general_marcas_resumen_y_cobranza(variante: str, registros) -> None:
    render_balance_general_xlsx(BalanceGeneralRequest.model_validate(pay.general_base(variante)))
    degradadas_base = registros[-1].degradadas()
    p = pay.general_con_cierre(variante)
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    assert registros[-1].degradadas() == degradadas_base
    assert verificar_libro(data)
    wb = load_workbook(BytesIO(data))
    # Hoja de vuelos de la flota: las dos marcas en ESTADO; la sombra sin
    # dinero.
    ws = wb[MAESTRA_FLOTA]
    enc = _encabezados(ws)
    sombra = _fila_de(ws, "#499")
    assert ws.cell(row=sombra, column=enc["ESTADO"]).value == pay.MARCA_PASA
    assert ws.cell(row=_fila_de(ws, "#403"), column=enc["ESTADO"]).value == pay.MARCA_VIENE
    for titulo, col in enc.items():
        if col > enc["ESTADO"] and "TIEMPO" not in titulo and "STATUS" not in titulo:
            assert ws.cell(row=sombra, column=col).value is None, titulo
    # El bloque NO va en la hoja de vuelos del general: va en el RESUMEN.
    assert not any(
        isinstance(c.value, str) and c.value.startswith("CIERRE DEL MES") for c in ws["A"]
    )
    rs = wb["RESUMEN flota"]
    tit = _fila_de(rs, "CIERRE DEL MES")
    assert rs.cell(row=tit, column=1).value == "CIERRE DEL MES — Mes cerrado el 16-oct-2026 por Ale"
    pasan = _fila_de(rs, "Pasan a octubre 2026")
    # CLAVE en A:B (la A del RESUMEN es angosta) y la columna AVIÓN.
    assert [rs.cell(row=pasan + 1, column=c).value for c in (1, 3, 4, 5, 6, 7, 8, 9)] == [
        "CLAVE", "FECHA", "AVIÓN", "VENTA\nMXN", "COBRADO\nMXN", "POR COBRAR\nMXN",
        "VENTA\nUSD", "LEYENDA",
    ]  # fmt: skip
    assert f"A{pasan + 2}:B{pasan + 2}" in {str(r) for r in rs.merged_cells.ranges}
    assert [rs.cell(row=pasan + 2, column=c).value for c in (1, 3, 4, 5, 7)] == [
        "#499 · Cliente Nueve", "28/09/2026", "XA-TST", 36632.89, 36632.89,
    ]  # fmt: skip
    # Renglón por avión de `resumen_cierre` (montos del API, en gris).
    assert rs.cell(row=pasan + 3, column=1).value == "XA-TST: 1 vuelo(s)"
    assert (rs.cell(row=pasan + 3, column=5).value, rs.cell(row=pasan + 3, column=7).value) == (
        36632.89, 36632.89,
    )  # fmt: skip
    vienen = _fila_de(rs, "Vienen de agosto 2026")
    assert rs.cell(row=vienen + 2, column=4).value == "XA-TST"
    # Cobranza de la flota: la sombra con su marca y celdas vacías.
    cb = wb["cobranza"]
    fs = _fila_de(cb, "#499")
    assert cb.cell(row=fs, column=4).value == pay.MARCA_PASA
    assert [cb.cell(row=fs, column=c).value for c in range(6, 13)] == [None] * 7
    # La vista en pantalla lleva el cierre de la raíz.
    vista = libro_balance_general(
        BalanceGeneralRequest.model_validate({**p, "vinculos_gasto": True})
    )
    assert vista["cierre"]["periodo"] == "2026-09"


def test_general_cierre_solo_en_el_consolidado_tambien_se_pinta() -> None:
    """Si el API manda `cierre` solo dentro del consolidado, el RESUMEN y la
    vista lo toman de ahí (no se pierde el bloque)."""
    p = pay.general_con_cierre("mensual")
    p["consolidado"]["cierre"] = p.pop("cierre")
    data = render_balance_general_xlsx(BalanceGeneralRequest.model_validate(p))
    rs = load_workbook(BytesIO(data))["RESUMEN flota"]
    assert _fila_de(rs, "CIERRE DEL MES")
    vista = libro_balance_general(BalanceGeneralRequest.model_validate(p))
    assert vista["cierre"]["estado"] == "CERRADO"


def test_mes_abierto_lista_lo_que_pasaria_y_hoy_suma() -> None:
    """Mes ABIERTO: la fila sigue con su dinero, con la marca «pasará … al
    cerrar» en ámbar, y el bloque lista los vuelos que pasarían."""
    p = pay.individual_base()
    fila = next(v for v in p["vuelos"] if v["clave"].startswith("#405"))
    fila["cierre_marca"] = "Pendiente de cobro: pasará a octubre 2026 al cerrar"
    fila["diferido_a"] = None
    p["cierre"] = {
        **pay.cierre_abierto_vacio(),
        "pasarian": [{"clave": "#405 · Cliente Cuatro", "fecha": "2026-09-20",
                      "venta_mxn": 36748.8, "cobrado_mxn": 20000.0,
                      "por_cobrar_mxn": 16748.8, "destino": "2026-10",
                      "leyenda": "Pendiente de cobro: pasará a octubre 2026 al cerrar"}],
    }  # fmt: skip
    data = render_balance_avion_xlsx(BalanceAvionRequest.model_validate(p))
    assert verificar_libro(data)
    ws = load_workbook(BytesIO(data))[MAESTRA_TST]
    enc = _encabezados(ws)
    f = _fila_de(ws, "#405")
    assert ws.cell(row=f, column=enc["ESTADO"]).value.startswith("Pendiente de cobro: pasará")
    assert ws.cell(row=f, column=enc["ESTADO"]).fill.fgColor.rgb.endswith(AMBER)
    assert ws.cell(row=f, column=enc["VENTA AVIÓN USD *"]).value == 2088
    assert ws.cell(row=_fila_de(ws, "CIERRE DEL MES"), column=1).value == (
        "CIERRE DEL MES (mes abierto)"
    )
    assert ws.cell(row=_fila_de(ws, "Pendientes de cobro"), column=1).value == (
        "Pendientes de cobro: pasarán a octubre 2026 al cerrar el mes (hoy sí suman)"
    )


# ---------------------------------------------------------------------------
# 5) Reparto PDF y Excel CON cierre.
# ---------------------------------------------------------------------------


def test_reparto_pdf_sello_listas_por_avion_y_externos() -> None:
    texto = _texto_pdf(
        render_reparto_pdf(RepartoPdfRequest.model_validate(pay.reparto_con_cierre()))
    )
    assert "Mes cerrado el 16-oct-2026 por Ale" in texto
    # El texto del API se escapa («&» y «<» no rompen el marcado).
    assert (
        "Pasan a octubre 2026 (pendientes de cobro, no se reparten): #316 Mike & Nelson <MN> "
        "$2,076.40 (cobrado $0.00, por cobrar $2,076.40)"
    ) in texto
    assert (
        "Vienen de agosto 2026 (cobrados: se reparten en este mes): #318 Ambulancia "
        "Internacional $800.00 (cobrado $800.00, por cobrar $0.00)"
    ) in texto
    # La marca del vuelo que viene, debajo del folio en el detalle.
    assert "#318 Viene de agosto 2026 Ambulancia Internacional" in texto
    # El externo (sin avión en el reparto) no se pierde.
    assert "Vuelos del cierre sin avión en este reparto" in texto
    assert "#990 Operador Externo $500.00 (cobrado $100.00, por cobrar $400.00)" in texto
    assert "AVISO:" not in texto


def test_reparto_pdf_aviso_si_no_coincide_con_el_snapshot() -> None:
    req = RepartoPdfRequest.model_validate(pay.reparto_con_cierre(coincide=False))
    texto = _texto_pdf(render_reparto_pdf(req))
    assert f"AVISO: {cierre_mes.AVISO_SNAPSHOT}" in texto


def test_reparto_pdf_marca_larga_se_parte_al_ancho_de_la_columna() -> None:
    v = RepartoVueloLinea(
        folio=405, cierre_marca="Pendiente de cobro: pasará a octubre 2026 al cerrar"
    )
    renglones = _folio_vuelo(v).split("\n")
    assert renglones[0] == "#405" and len(renglones) >= 3
    assert all(stringWidth(r, "Helvetica", 7.5) <= _ANCHO_FOLIO_PT for r in renglones)


def test_reparto_xlsx_sello_seccion_y_aviso() -> None:
    req = RepartoPdfRequest.model_validate(pay.reparto_con_cierre(coincide=False))
    ws = load_workbook(BytesIO(render_reparto_xlsx(req))).active
    assert ws["A2"].value.endswith("  ·  Mes cerrado el 16-oct-2026 por Ale")
    textos = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert f"AVISO: {cierre_mes.AVISO_SNAPSHOT}" in textos
    # Detalle de vuelos: la marca junto al folio.
    assert "#318 · Viene de agosto 2026" in [c.value for c in ws["B"]]
    tit = _fila_de(ws, "Cierre del mes")
    assert ws.cell(row=tit, column=1).value == "Cierre del mes — Mes cerrado el 16-oct-2026 por Ale"
    assert [ws.cell(row=tit + 1, column=c).value for c in range(1, 9)] == ENCABEZADOS_CIERRE
    salen = _fila_de(ws, "Pasan a octubre 2026")
    assert ws.cell(row=salen, column=1).fill.fgColor.rgb.endswith(AMBER)
    assert [ws.cell(row=salen + 1, column=c).value for c in range(1, 9)] == [
        "XB-TST", "#316", "Mike & Nelson <MN>", 2076.4, 0, 2076.4, "octubre 2026",
        "Pasa a octubre 2026 · pendiente de cobro: cobrado $0.00 de $2,076.40 USD",
    ]  # fmt: skip
    assert ws.cell(row=salen + 2, column=1).value is None  # el externo
    assert ws.cell(row=salen + 2, column=2).value == "#990"
    entran = _fila_de(ws, "Vienen de agosto 2026")
    assert ws.cell(row=entran + 1, column=2).value == "#318"
    assert ws.cell(row=entran + 1, column=7).value == "agosto 2026"
    # La sección no toca la tabla de arriba (TOTAL de siempre).
    base = load_workbook(
        BytesIO(render_reparto_xlsx(RepartoPdfRequest.model_validate(pay.reparto_base())))
    ).active
    fila_total = _fila_de(ws, "TOTAL")
    assert [ws.cell(row=fila_total, column=c).value for c in range(1, 15)] == [
        base.cell(row=_fila_de(base, "TOTAL"), column=c).value for c in range(1, 15)
    ]


# ---------------------------------------------------------------------------
# 6) Libro Dinero CON cierre.
# ---------------------------------------------------------------------------


def test_dinero_marca_junto_a_la_clave_y_lineas_al_final() -> None:
    p = pay.dinero_con_cierre()
    data = render_dinero_xlsx(DineroXlsxRequest.model_validate(p))
    ws = load_workbook(BytesIO(data))["dinero-vlos"]
    assert ws["A5"].value == "vtbeh · Viene de septiembre 2026"
    assert ws["A5"].font.color.rgb.endswith("1E3A8A")
    assert ws.column_dimensions["A"].width == len("vtbeh · Viene de septiembre 2026") + 2
    # La fila que viene suma (TOTALES de siempre: el API ya la cuenta).
    base = load_workbook(
        BytesIO(render_dinero_xlsx(DineroXlsxRequest.model_validate(pay.dinero_base())))
    )["dinero-vlos"]
    assert [c.value for c in ws[6]] == [c.value for c in base[6]]
    tit = _fila_de(ws, "CIERRE DEL MES")
    assert ws.cell(row=tit, column=1).value == "CIERRE DEL MES — Mes cerrado el 16-oct-2026 por Ale"
    assert ws.cell(row=tit + 1, column=1).value == (
        "Pasan a octubre 2026 (pendientes de cobro al cierre: no suman en este mes): "
        "#499 · Cliente Nueve — venta $36,632.89 MXN, cobrado $0.00, por cobrar $36,632.89."
    )
    assert ws.cell(row=tit + 2, column=1).value.startswith("Vienen de agosto 2026")


# ---------------------------------------------------------------------------
# 7) Rutas: el payload del API 0.0.88 entra (jamás un 422).
# ---------------------------------------------------------------------------

TOKEN = "tok-cierre-mes"
client = TestClient(app)


def test_rutas_aceptan_los_payloads_del_cierre(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    h = {"X-Internal-Token": TOKEN}
    individual = pay.individual_con_cierre()
    general = pay.general_con_cierre("mensual")
    reparto = pay.reparto_con_cierre()
    for ruta, payload in (
        ("/pdf/balance-avion-xlsx", individual),
        ("/pdf/balance-general-xlsx", general),
        ("/pdf/balance-avion-libro", individual),
        ("/pdf/balance-general-libro", general),
        ("/pdf/reparto", reparto),
        ("/pdf/reparto-xlsx", reparto),
        ("/pdf/dinero-xlsx", pay.dinero_con_cierre()),
    ):
        res = client.post(ruta, json=copy.deepcopy(payload), headers=h)
        assert res.status_code == 200, (ruta, res.text[:300])
    res = client.post("/pdf/balance-avion-libro", json=individual, headers=h)
    assert res.json()["cierre"]["estado"] == "CERRADO"
