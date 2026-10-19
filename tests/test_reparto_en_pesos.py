"""Reparto a socios (PDF y Excel) EN PESOS (18-oct-2026, API 0.0.93).

Pedido del cliente: «… todo en pesos porque casi siempre se les transfiere a
los socios en pesos mexicanos, y que los dólares solo sean informativos pero
no los que llamen primero la atención.» Con la bandera raíz
`reparto_en_pesos` la cascada de cada avión, el saldo, el reparto por socio,
los vuelos y los otros ingresos van en PESOS (los `*_mxn` del API) y los
dólares en gris, cursiva y más chicos; `pesos_sin_tc_usd > 0` avisa que el
saldo en pesos está incompleto y la reserva dice su T.C. Sin la bandera, el
PDF y el Excel son BYTE A BYTE los del HEAD 3b893f1 (`_HUELLAS_HEAD`,
calculadas con ese HEAD sobre `tests/_pesos_payloads.py`).
"""

from __future__ import annotations

import copy
import hashlib
import re
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import get_settings
from app.main import app
from app.schemas.reparto import (
    RepartoAvion,
    RepartoGastoCategoriaLinea,
    RepartoOtroIngresoLinea,
    RepartoPdfRequest,
    RepartoSocioLinea,
    RepartoVueloLinea,
)
from app.services.reparto_pdf import (
    ENCABEZADO_PESOS,
    ENCABEZADO_USD_INFO,
    LEYENDA_PESOS,
    LINEA_OTROS_INGRESOS_SIN_TC_PESOS,
    aviso_pesos_sin_tc,
    render_reparto_pdf,
)
from app.services.reparto_xlsx import (
    GRIS_INFO,
    SIN_PESOS,
    TITULO_USD_INFORMATIVO,
    encabezados_pesos,
    render_reparto_xlsx,
)
from tests import _pesos_payloads as pp

# Huellas «sin la bandera» (HEAD 3b893f1, `tests/_pesos_payloads.reparto`):
# flujos de contenido de cada página del PDF y miembros del .xlsx salvo
# docProps/core.xml. Coinciden con las de `test_cierre_mes`: las llaves en
# pesos sin la bandera no mueven nada.
_HUELLA_PDF_HEAD = "9448b553c276bcd2bd8ddb878136bf888bf2d68f85d2fc5a1a3aaee2da91d5c0"
_HUELLA_XLSX_HEAD = "4e1474b790df44f408f8634d116156b512c0c71824952734238ff42a82039dcc"


def _huella_pdf(data: bytes) -> str:
    h = hashlib.sha256()
    for p in PdfReader(BytesIO(data)).pages:
        h.update(p.get_contents().get_data())
    return h.hexdigest()


def _huella_xlsx(data: bytes) -> str:
    z = zipfile.ZipFile(BytesIO(data))
    h = hashlib.sha256()
    for nombre in sorted(z.namelist()):
        if nombre != "docProps/core.xml":
            h.update(nombre.encode())
            h.update(z.read(nombre))
    return h.hexdigest()


def _texto_pdf(data: bytes) -> str:
    return re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(BytesIO(data)).pages))


def _req(p: dict) -> RepartoPdfRequest:
    return RepartoPdfRequest.model_validate(p)


def _dinero(x: float) -> str:
    return f"${x:,.2f}"


# ---------------------------------------------------------------------------
# 1) Esquemas.
# ---------------------------------------------------------------------------


def test_esquemas_aditivos_con_default() -> None:
    a = RepartoAvion(matricula="X", modelo="M", ingresos_cobrado_usd=0, saldo_usd=0)
    for llave in (
        "ingresos_cobrado_mxn",
        "comisiones_venta_mxn",
        "pendiente_cobro_mxn",
        "gastos_directos_mxn",
        "gastos_indirectos_mxn",
        "permisos_mxn",
        "otros_mxn",
        "reserva_overhaul_mxn",
        "otros_ingresos_avion_mxn",
        "saldo_mxn",
        "pesos_sin_tc_usd",
        "reserva_tc",
    ):
        assert getattr(a, llave) is None
    assert RepartoSocioLinea(socio_nombre="S", porcentaje=1, monto_usd=1).monto_mxn is None
    v = RepartoVueloLinea()
    assert v.cobrado_mxn is None and v.pendiente_mxn is None
    assert RepartoOtroIngresoLinea().mxn is None
    assert RepartoGastoCategoriaLinea().mxn is None
    base = {"periodo_desde": "a", "periodo_hasta": "b", "generado": "c"}
    assert RepartoPdfRequest.model_validate(base).reparto_en_pesos is False


@pytest.mark.parametrize("valor", [None, False, "true", 1, "si"])
def test_bandera_liberal(valor) -> None:
    p = pp.reparto(bandera=None)
    p["reparto_en_pesos"] = valor
    assert _req(p).reparto_en_pesos is False


# ---------------------------------------------------------------------------
# 2) Sin la bandera: byte a byte.
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("bandera", [None, False, "si"])
def test_sin_bandera_byte_a_byte(bandera) -> None:
    req = _req(pp.reparto(bandera=bandera))
    assert _huella_pdf(render_reparto_pdf(req)) == _HUELLA_PDF_HEAD
    assert _huella_xlsx(render_reparto_xlsx(req)) == _HUELLA_XLSX_HEAD


# ---------------------------------------------------------------------------
# 3) PDF en pesos.
# ---------------------------------------------------------------------------


def test_pdf_en_pesos_principal_y_usd_informativo() -> None:
    p = pp.reparto()
    texto = _texto_pdf(render_reparto_pdf(_req(p)))
    assert LEYENDA_PESOS in texto
    assert f"{ENCABEZADO_PESOS} {ENCABEZADO_USD_INFO}" in texto
    for a in p["aviones"]:
        # Saldo: pesos primero, dólares después.
        assert f"Saldo disponible {_dinero(a['saldo_mxn'])} {_dinero(a['saldo_usd'])}" in texto
        assert (
            f"Venta del avión cobrada {_dinero(a['ingresos_cobrado_mxn'])} "
            f"{_dinero(a['ingresos_cobrado_usd'])}"
        ) in texto
        assert f"(-) Reserva overhaul (al T.C. {pp.K_REPARTO:g})" in texto
        for linea in a["reparto"]:
            assert (
                f"{linea['socio_nombre']} {linea['porcentaje']:g}% "
                f"{_dinero(linea['monto_mxn'])} {_dinero(linea['monto_usd'])}"
            ) in texto
    assert "Monto MXN USD (inf.)" in texto
    # Resumen por socio en pesos.
    total_mxn = round(sum(a["saldo_mxn"] for a in p["aviones"]), 2)
    assert "Total a recibir (MXN)" in texto
    assert f"Socio A {_dinero(total_mxn)}" in texto
    # Vuelos con pesos y dólares.
    assert "#316 Mike Nelson $20,700.00 $0.00 $1,200.00 $0.00" in texto
    # El pie lo dice; el de dólares ya no.
    assert "Montos en pesos (MXN); las cifras en dólares (USD) son solo informativas." in texto
    assert "Montos en USD." not in texto
    # Lo que solo existe en dólares, rotulado USD.
    assert "TUAs $100.00 USD" in texto
    # Sin dólares sin pesos, ningún aviso.
    assert "no se pudieron pasar a pesos" not in texto


def test_pdf_aviso_pesos_sin_tc() -> None:
    p = pp.reparto()
    p["aviones"][0]["pesos_sin_tc_usd"] = 120.5
    texto = _texto_pdf(render_reparto_pdf(_req(p)))
    assert f"AVISO: {aviso_pesos_sin_tc(120.5)}" in texto
    assert aviso_pesos_sin_tc(120.5) == (
        "$120.50 USD no se pudieron pasar a pesos (sin tipo de cambio): el saldo en "
        "pesos está incompleto."
    )
    assert texto.count("no se pudieron pasar a pesos") == 1


def test_pdf_sin_pesos_raya_nunca_cero() -> None:
    p = pp.reparto()
    a = p["aviones"][0]
    a["saldo_mxn"] = None
    a["reparto"][0]["monto_mxn"] = None
    a["vuelos"][0]["cobrado_mxn"] = None
    texto = _texto_pdf(render_reparto_pdf(_req(p)))
    assert f"Saldo disponible — {_dinero(a['saldo_usd'])}" in texto
    assert "#316 Mike Nelson — $0.00" in texto
    # El total a recibir del socio con una línea sin pesos también es «—».
    assert "Socio A — " in texto


def test_pdf_otros_ingresos_del_avion_en_pesos() -> None:
    p = pp.reparto()
    a = p["aviones"][0]
    a["otros_ingresos_avion_usd"] = 100.0
    a["otros_ingresos_avion_mxn"] = 1725.0
    a["otros_ingresos_avion_sin_tc_mxn"] = 300.0
    a["otros_ingresos_avion"] = [
        {"clave": "ING-4", "categoria": "Reembolso", "concepto": "serv 50 hrs",
         "monto": 1725.0, "moneda": "MXN", "usd": 100.0, "mxn": 1725.0},
        {"clave": "ING-5", "categoria": "Venta", "concepto": "pieza", "monto": 50.0,
         "moneda": "USD", "usd": None, "mxn": None},
    ]  # fmt: skip
    texto = _texto_pdf(render_reparto_pdf(_req(p)))
    assert "ING-4 · Reembolso · serv 50 hrs $1,725.00 $100.00" in texto
    assert "ING-5 · Venta · pieza sin T.C. sin T.C." in texto
    assert f"{LINEA_OTROS_INGRESOS_SIN_TC_PESOS} $300.00" in texto
    # El resumen gana la columna de otros ingresos de aviones.
    assert "Otros ingresos de aviones" in texto


# ---------------------------------------------------------------------------
# 4) Excel en pesos.
# ---------------------------------------------------------------------------


def _hoja(p: dict):
    data = render_reparto_xlsx(_req(p))
    return load_workbook(BytesIO(data)).active


def _fila(ws, texto: str, col: int = 1) -> int:
    filas = [c.row for c in ws[chr(64 + col)] if c.value == texto]
    assert len(filas) == 1, (texto, filas)
    return filas[0]


def test_xlsx_tabla_principal_en_pesos() -> None:
    p = pp.reparto()
    ws = _hoja(p)
    assert [ws.cell(row=4, column=c).value for c in range(1, 15)] == encabezados_pesos(False)
    assert "Cifras en pesos (MXN)" in ws["A2"].value
    for i, a in enumerate(p["aviones"]):
        r = 5 + i
        assert ws.cell(row=r, column=1).value == a["matricula"]
        assert ws.cell(row=r, column=4).value == a["ingresos_cobrado_mxn"]
        assert ws.cell(row=r, column=12).value == a["saldo_mxn"]
        assert ws.cell(row=r, column=11).comment.text.endswith(f"{pp.K_REPARTO:g}.")
        # Otros ingr. VuelaTour: solo en dólares, informativo.
        info = ws.cell(row=r, column=13)
        assert info.font.italic and info.font.color.rgb.endswith(GRIS_INFO)
    total = 5 + len(p["aviones"])
    assert ws.cell(row=total, column=1).value == "TOTAL"
    assert ws.cell(row=total, column=12).value == round(
        sum(a["saldo_mxn"] for a in p["aviones"]), 2
    )


def test_xlsx_dolares_debajo_informativos() -> None:
    p = pp.reparto()
    ws = _hoja(p)
    t = _fila(ws, TITULO_USD_INFORMATIVO)
    assert ws.cell(row=t + 1, column=12).value == "Saldo disponible"
    for i, a in enumerate(p["aviones"]):
        c = ws.cell(row=t + 2 + i, column=12)
        assert c.value == a["saldo_usd"]
        assert c.font.italic and c.font.size == 9 and c.font.color.rgb.endswith(GRIS_INFO)


def test_xlsx_socios_y_vuelos_en_pesos() -> None:
    p = pp.reparto()
    p["aviones"][0]["vuelos"][1]["pendiente_mxn"] = None
    ws = _hoja(p)
    t = _fila(ws, "Reparto por socio (en pesos)")
    assert [ws.cell(row=t + 1, column=c).value for c in range(1, 6)] == [
        "Avión", "Socio", "Porcentaje", "Monto MXN", "Monto USD (informativo)",
    ]  # fmt: skip
    linea = p["aviones"][0]["reparto"][0]
    assert ws.cell(row=t + 2, column=4).value == linea["monto_mxn"]
    assert ws.cell(row=t + 2, column=5).value == linea["monto_usd"]
    assert ws.cell(row=t + 2, column=5).font.italic
    v = _fila(ws, "Detalle de vuelos (venta del avión cobrada)")
    assert ws.cell(row=v + 1, column=4).value == "Cobrado avión MXN"
    assert ws.cell(row=v + 2, column=4).value == p["aviones"][0]["vuelos"][0]["cobrado_mxn"]
    assert ws.cell(row=v + 3, column=5).value == SIN_PESOS
    assert ws.cell(row=v + 2, column=6).value == p["aviones"][0]["vuelos"][0]["cobrado_usd"]


def test_xlsx_aviso_pesos_sin_tc_y_raya() -> None:
    p = pp.reparto()
    p["aviones"][1]["pesos_sin_tc_usd"] = 55
    p["aviones"][1]["saldo_mxn"] = None
    ws = _hoja(p)
    textos = [c.value for c in ws["A"] if isinstance(c.value, str)]
    aviso = next(t for t in textos if "no se pudieron pasar a pesos" in t)
    assert aviso == (
        f"AVISO {p['aviones'][1]['matricula']}: $55.00 USD no se pudieron pasar a pesos "
        "(sin tipo de cambio): el saldo en pesos está incompleto."
    )
    r = 5 + 1
    assert ws.cell(row=r, column=12).value == SIN_PESOS
    # Con un «—» en la columna, el total también es «—» (no una suma a medias).
    assert ws.cell(row=r + 1, column=12).value == SIN_PESOS


def test_xlsx_gastos_por_categoria_con_pesos() -> None:
    p = pp.reparto()
    p["aviones"][0]["gastos_por_categoria"] = [
        {"categoria": "GAS", "etiqueta": "Gasavión", "grupo": "DIRECTO", "count": 2,
         "usd": 500.0, "mxn": 8625.0},
    ]  # fmt: skip
    ws = _hoja(p)
    t = _fila(ws, "Gastos por categoría")
    assert ws.cell(row=t + 1, column=5).value == "MXN"
    assert ws.cell(row=t + 2, column=5).value == 8625.0
    assert ws.cell(row=t + 2, column=6).value == 500.0


def test_cierre_del_mes_se_conserva_en_pesos() -> None:
    """La sección «Cierre del mes» (dólares del API, al corte) y el sello del
    cierre siguen en el documento en pesos: no hay `*_mxn` de esas listas."""
    p = pp.reparto()
    con_cierre = pp.pay.reparto_con_cierre(coincide=False)
    p["cierre"] = copy.deepcopy(con_cierre["cierre"])
    ws = _hoja(p)
    assert ws.cell(row=4, column=12).value == "Saldo disponible (MXN)"
    textos = [c.value for c in ws["A"] if isinstance(c.value, str)]
    assert any(t.startswith("Cierre del mes — Mes cerrado") for t in textos)
    assert any(t.startswith("AVISO: ") and "no se pudieron" not in t for t in textos)
    assert "Mes cerrado" in ws["A2"].value
    texto = _texto_pdf(render_reparto_pdf(_req(p)))
    assert "Mes cerrado" in texto and LEYENDA_PESOS in texto


# ---------------------------------------------------------------------------
# 5) Rutas.
# ---------------------------------------------------------------------------

TOKEN = "tok-reparto-pesos"
client = TestClient(app)


def test_rutas_aceptan_el_payload_0093(monkeypatch, request) -> None:
    monkeypatch.setenv("INTERNAL_SHARED_TOKEN", TOKEN)
    get_settings.cache_clear()
    request.addfinalizer(get_settings.cache_clear)
    headers = {"X-Internal-Token": TOKEN}
    p = pp.reparto()
    res = client.post("/pdf/reparto", json=p, headers=headers)
    assert res.status_code == 200, res.text
    assert LEYENDA_PESOS in _texto_pdf(res.content)
    res = client.post("/pdf/reparto-xlsx", json=p, headers=headers)
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content)).active
    assert ws.cell(row=4, column=12).value == "Saldo disponible (MXN)"
